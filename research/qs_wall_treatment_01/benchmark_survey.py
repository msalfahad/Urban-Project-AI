"""BENCHMARK UNSEAL, step 1 (§38): survey the sealed spreadsheets to find
the human wall-treatment benchmark - AFTER the A22 freeze, never before.

Precondition, checked from disk: FREEZE_A22.json exists and every artifact
it hashed is byte-identical now. If not, the survey refuses.

What is read: sheet names, dimensions, and the first rows/columns of each
sheet, as text. Every read is logged. Nothing read here flows into the
estimate: the estimate is frozen and its hash is in the log.

    python3 -m research.qs_wall_treatment_01.benchmark_survey
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from research.qs_wall_treatment_01 import protocol as P
from research.qs_wall_treatment_01 import source_inventory as SI

OUT = Path(P.OUT_DIR)
ROWS, COLS = 12, 10
KEYWORDS = {
    "PLASTER": re.compile(r"plaster|لياسة|محارة|بياض|قصارة|لياسه|مساح", re.I),
    "P7757": re.compile(r"7757", re.I),
    "WALL": re.compile(r"\bwall|جدار|حوائط|حائط", re.I),
    "M2": re.compile(r"m2|م2|متر مربع|sqm", re.I),
}


def _sha(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _verify_freeze() -> dict:
    fz = json.loads((OUT / "FREEZE_A22.json").read_text("utf-8"))
    bad = {}
    for a, h in fz["ARTIFACT_SHA256"].items():
        cur = _sha(OUT / a) if (OUT / a).exists() else None
        if cur != h:
            bad[a] = {"FROZEN": h, "NOW": cur}
    if bad:
        raise SystemExit(f"A22 freeze does not match disk: {bad}")
    return {"FREEZE_A22_DIGEST": fz["FREEZE_DIGEST_SHA256"], "VERIFIED": True}


def _cells_xlsx(path: Path) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ROWS, max_col=COLS,
                                             values_only=True)):
            rows.append([("" if v is None else str(v))[:40] for v in row])
        out.append({"SHEET": ws.title, "DIMENSIONS": ws.calculate_dimension()
                    if hasattr(ws, "calculate_dimension") else None,
                    "MAX_ROW": ws.max_row, "MAX_COL": ws.max_column, "HEAD": rows})
    wb.close()
    return out


def _cells_xls(path: Path) -> list:
    import xlrd
    wb = xlrd.open_workbook(path, on_demand=True)
    out = []
    for ws in wb.sheets():
        rows = []
        for r in range(min(ROWS, ws.nrows)):
            rows.append([str(ws.cell_value(r, c))[:40] for c in range(min(COLS, ws.ncols))])
        out.append({"SHEET": ws.name, "MAX_ROW": ws.nrows, "MAX_COL": ws.ncols, "HEAD": rows})
    return out


def run() -> dict:
    ver = _verify_freeze()
    inv = json.loads((OUT / "SOURCE_INVENTORY.json").read_text("utf-8"))
    log = {"PHASE_ID": P.PHASE_ID, "ARTIFACT": "BENCHMARK_ACCESS_LOG",
           "OPENED_AFTER": ver, "WHAT_WAS_READ": f"sheet names, sizes, first {ROWS} rows x {COLS} cols",
           "ESTIMATE_FROZEN_SHA256": json.loads(
               (OUT / "FREEZE_A22.json").read_text("utf-8"))["ARTIFACT_SHA256"][
               "P7757_WALL_TREATMENT_ESTIMATE.json"],
           "FILES": []}
    for s in inv["SPREADSHEETS_SEALED"]:
        p = SI.UPLOAD_DIR / s["NAME"]
        rec = {"NAME": s["NAME"], "SHA256": s["SHA256"], "SIZE": s["SIZE"]}
        try:
            sheets = _cells_xls(p) if p.suffix.lower() == ".xls" else _cells_xlsx(p)
            text = json.dumps(sheets, ensure_ascii=False)
            rec["SHEETS"] = sheets
            rec["KEYWORD_HITS"] = {k: len(rx.findall(text)) for k, rx in KEYWORDS.items()}
        except Exception as e:  # noqa: BLE001 - the log must say why
            rec["ERROR"] = f"{type(e).__name__}: {e}"[:200]
        log["FILES"].append(rec)
    (OUT / "BENCHMARK_ACCESS_LOG.json").write_text(
        json.dumps(log, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return {"FILES": [(f["NAME"], f.get("KEYWORD_HITS"), [sh["SHEET"] for sh in f.get("SHEETS", [])][:12],
                       f.get("ERROR")) for f in log["FILES"]]}


if __name__ == "__main__":
    print(json.dumps(run(), indent=2, ensure_ascii=False))
