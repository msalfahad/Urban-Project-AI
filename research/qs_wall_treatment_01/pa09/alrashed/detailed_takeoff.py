"""ALRASHED_DETAILED_QUANTITY_TAKEOFF - the frozen measurement, opened up so it can be read without AutoCAD.

Nothing here is a new measurement.  Every area in this workbook is the frozen area; what is new is that the area
is shown as the rectangles it is actually made of, so a reader can follow 6.30 x 4.00 = 25.20 rather than trust a
single number.  The decomposition is exact by construction - the plan is axis-aligned, a component is a whole
number of grid cells, and the greedy maximal-rectangle cover sums back to the frozen figure exactly.

Revision 2 answers an external audit.  Four things it insisted on, all of which were wrong here before:

  * a floor area whose finish nobody has settled is not porcelain.  The 196.378 m2 of stair, landing and unnamed
    space is measured, kept, and classified FINISH_CLASSIFICATION_PENDING - out of the base BOQ, not out of sight;
  * 135 connected components are not 135 rooms.  71 of them are wall material and slivers, and the census says so
    in the workbook, the export and the report;
  * a ceiling with no ceiling plan has a measured AREA and an unsettled FINISH, and only the first is final;
  * a guide category used outside its own area band needs an authority.  One window has none, so its height is
    ASK_THE_OWNER rather than a number that looks decided.

The frozen artifact is never rewritten by this module: it is read, and every figure that comes from it is carried
across unchanged.
"""

from __future__ import annotations

import collections
import datetime as dt
import hashlib
import json
import subprocess
from pathlib import Path

from research.qs_wall_treatment_01 import protocol as PR
from research.qs_wall_treatment_01.pa09.alrashed import (final_takeoff as FT, geometry as G, owner_inputs as OI,
                                                          quantities as Q, takeoff as TK, validation_amendment as VA,
                                                          window_standard as WS, workbook_calc as WC)

OUT = Path(PR.OUT_DIR) / "pa09_alrashed"
H = OI.AR01_WALL_HEIGHT_M
FROZEN_DIGEST = "ae259eaba3203798"
FROZEN_COMMIT = "3e847af"
FROZEN_FILE_SHA256 = "7e9a3eba636ea95b77fce2bbb7dccad79d6971047af455bb2f542ee62165493f"
REVISION = 2
PART_LETTERS = ("أ", "ب", "ج", "د", "هـ", "و", "ز", "ح", "ط", "ي", "ك", "ل", "م", "ن", "س", "ع")
AREA_TOL_M2 = 1e-6          # QA-01: a component area recomputed from its raw dimensions
CROSS_TOL_M2 = 0.01         # agreement between artifacts that publish dimensions at different rounding

# ------------------------------------------------------------------ what a component is, and is not
SLIVER_ROLE = "WALL_MATERIAL_OR_SLIVER"
PORCELAIN_FLOOR_ROLES = ("INTERNAL_ROOM", "WET_ROOM", "KITCHEN")
PENDING_FLOOR_ROLES = ("STAIR_OR_LANDING", "UNNAMED_ON_DRAWING")
ENCLOSED_ROLES = PORCELAIN_FLOOR_ROLES + PENDING_FLOOR_ROLES
WET_ROLES = ("WET_ROOM", "KITCHEN")
SKIRTING_ROLES = ("INTERNAL_ROOM",) + PENDING_FLOOR_ROLES
EXPECTED_ROLE_COUNTS = {"INTERNAL_ROOM": 16, "EXTERNAL_OR_OPEN": 2, "STAIR_OR_LANDING": 9,
                        "UNNAMED_ON_DRAWING": 24, "WET_ROOM": 12, "KITCHEN": 1, SLIVER_ROLE: 71}

FLOORS_AR = {"BASEMENT": "السرداب", "GROUND": "الأرضي", "FIRST": "الأول", "ROOF": "السطح"}
ROLES_AR = {
    "INTERNAL_ROOM": "غرفة داخلية", "WET_ROOM": "غرفة رطبة", "KITCHEN": "مطبخ",
    "EXTERNAL_OR_OPEN": "خارجي / مكشوف", "STAIR_OR_LANDING": "درج / بسطة",
    "UNNAMED_ON_DRAWING": "غير مسماة على المخطط", SLIVER_ROLE: "مادة جدار / شظية",
}
SKIRTING_BUCKETS = {
    "INTERNAL_ROOM": ("DRY_NAMED_INTERNAL_CANDIDATE", "غرف داخلية جافة مسماة — مرشحة"),
    "STAIR_OR_LANDING": ("STAIR_OR_LANDING_PENDING", "درج وبسطات — معلّقة"),
    "UNNAMED_ON_DRAWING": ("UNNAMED_SPACE_PENDING", "فراغات غير مسماة — معلّقة"),
}


def _git(*a):
    try:
        return subprocess.run(["git", *a], capture_output=True, text=True, timeout=60).stdout.strip()
    except Exception:
        return ""


def frozen_record():
    return json.loads((OUT / "FULL_VILLA_BLIND_PRE_PRICING_TAKEOFF.json").read_text("utf-8"))


def frozen_file_sha256():
    p = OUT / "FULL_VILLA_BLIND_PRE_PRICING_TAKEOFF.json"
    return hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None


# ------------------------------------------------------------------ the decomposition
def rectangles(cells, xs, ys):
    """Cover a component's cells with the fewest large rectangles, greedily and deterministically.

    Each rectangle keeps its dimensions at full precision (RAW) as well as rounded for display.  The workbook
    stores the raw value and formats it; the export publishes both, so an auditor can recompute every area from
    the raw dimensions and land inside 1e-6 m2.
    """
    rem = set(cells)
    out = []
    while rem:
        best = None
        for i0, j0 in sorted(rem):
            hi = i0
            while (hi + 1, j0) in rem:
                hi += 1
            j1 = j0
            while True:
                area = (xs[hi + 1] - xs[i0]) * (ys[j1 + 1] - ys[j0])
                if best is None or area > best[0] + 1e-12:
                    best = (area, i0, j0, hi, j1)
                nj = j1 + 1
                lo = i0 - 1
                for i in range(i0, hi + 1):
                    if (i, nj) in rem:
                        lo = i
                    else:
                        break
                if lo < i0:
                    break
                hi, j1 = lo, nj
        _a, i0, j0, i1, j1 = best
        out.append({"X0": xs[i0], "Y0": ys[j0], "X1": xs[i1 + 1], "Y1": ys[j1 + 1],
                    "RAW_LENGTH_M": xs[i1 + 1] - xs[i0], "RAW_WIDTH_M": ys[j1 + 1] - ys[j0],
                    "RAW_AREA_M2": (xs[i1 + 1] - xs[i0]) * (ys[j1 + 1] - ys[j0])})
        for i in range(i0, i1 + 1):
            for j in range(j0, j1 + 1):
                rem.discard((i, j))
    out.sort(key=lambda r: -r["RAW_AREA_M2"])
    for k, r in enumerate(out):
        r["PART"] = f"جزء {PART_LETTERS[k]}" if k < len(PART_LETTERS) else f"جزء {k + 1}"
        r["DISPLAY_LENGTH_M"] = round(r["RAW_LENGTH_M"], 4)
        r["DISPLAY_WIDTH_M"] = round(r["RAW_WIDTH_M"], 4)
        r["DISPLAY_AREA_M2"] = round(r["RAW_AREA_M2"], 4)
        r["CALCULATION"] = (f"{r['DISPLAY_LENGTH_M']:.4f} × {r['DISPLAY_WIDTH_M']:.4f} "
                            f"= {r['RAW_AREA_M2']:.4f}")
    return out


def _openings_by_component(ents, floor, g, rows, by_ref):
    """Which openings sit on which component's boundary - the proximity test the frozen wet-room pass used."""
    xs, ys = g["XS"], g["YS"]
    ops = Q.opening_register(ents, floor, g, Q._block_names())
    out = collections.defaultdict(list)
    for r in rows:
        c = by_ref.get(r["ROOM_REF"])
        cells = set(c["CELLS"]) if c else set()
        for o in ops:
            if o["WIDTH_M"] is None:
                continue
            hit = False
            for d in (0.15, 0.35):
                for s in (-1, 1):
                    px, py = ((o["X"], o["Y"] + s * d) if o["RUNS"] == "ALONG_X" else (o["X"] + s * d, o["Y"]))
                    for i, j in cells:
                        if xs[i] <= px <= xs[i + 1] and ys[j] <= py <= ys[j + 1]:
                            hit = True
                            break
                    if hit:
                        break
                if hit:
                    break
            if hit:
                out[r["ROOM_REF"]].append(o)
    return out


def component_model(ents):
    """Every connected component of every floor, its frozen area, and the rectangles that make it up."""
    floors = []
    for floor in TK.FLOORS:
        g = G.grid(ents, G.WINDOWS[floor])
        comps = G.rooms(g)
        _g, _tr, _labs, rows, closure = TK.floor_rows(ents, floor)
        by_ref = {f"{floor[:2]}-{k:03d}": c for k, c in enumerate(comps)}
        ops = _openings_by_component(ents, floor, g, rows, by_ref)
        out = []
        for r in rows:
            c = by_ref[r["ROOM_REF"]]
            parts = rectangles(c["CELLS"], g["XS"], g["YS"])
            raw_sum = sum(p["RAW_AREA_M2"] for p in parts)
            cad = sum((g["XS"][i + 1] - g["XS"][i]) * (g["YS"][j + 1] - g["YS"][j]) for i, j in c["CELLS"])
            doors = [o for o in ops.get(r["ROOM_REF"], []) if o["TYPE"] == "DOOR"]
            row = dict(r)
            row.update({
                "COMPONENT_REF": r["ROOM_REF"],
                "ENTITY_TYPE": "CONNECTED_COMPONENT",
                "IS_SLIVER": r["ROLE"] == SLIVER_ROLE,
                "PARTS": parts,
                "PARTS_RAW_SUM_M2": raw_sum,
                "PARTS_SUM_M2": round(raw_sum, 4),
                "CAD_POLYGON_AREA_M2": r["AREA_M2"],
                "DECOMPOSITION_VARIANCE_M2": round(raw_sum - cad, 9),
                "ROUNDING_DISPLAY_VARIANCE_M2": round(round(raw_sum, 4) - r["AREA_M2"], 6),
                "DOOR_WIDTH_ON_BOUNDARY_M": round(sum(o["WIDTH_M"] for o in doors), 3),
                "DOOR_COUNT_ON_BOUNDARY": len(doors),
                "FLOOR_AR": FLOORS_AR[floor], "ROLE_AR": ROLES_AR.get(r["ROLE"], r["ROLE"]),
                "MEASUREMENT_BASIS": "exact rectilinear decomposition of the drawn wall faces",
            })
            out.append(row)
        out.sort(key=lambda x: (-x["AREA_M2"], x["COMPONENT_REF"]))
        floors.append({"FLOOR": floor, "FLOOR_AR": FLOORS_AR[floor], "LEVEL_M": TK.LEVELS[floor],
                       "COMPONENTS": out, "CLOSURE": closure,
                       "FLOOR_TOTAL_M2": round(sum(x["AREA_M2"] for x in out), 4)})
    return floors


def census(floors):
    """What the 135 components are.  A connected component is not a room until something says it is one."""
    comps = [c for f in floors for c in f["COMPONENTS"]]
    by_role = collections.Counter(c["ROLE"] for c in comps)
    area = collections.defaultdict(float)
    for c in comps:
        area[c["ROLE"]] += c["AREA_M2"]
    slivers = by_role.get(SLIVER_ROLE, 0)
    rec = {
        "ENTITY": "CONNECTED_COMPONENT",
        "WHAT_A_COMPONENT_IS": "a maximal set of plan grid cells joined to each other without a wall or a "
                               "virtual closure between them.  It becomes a room only when the drawing names it "
                               "or its geometry proves what it is",
        "CONNECTED_COMPONENT_COUNT": len(comps),
        "NON_SLIVER_COMPONENT_COUNT": len(comps) - slivers,
        "WALL_MATERIAL_OR_SLIVER_COUNT": slivers,
        "COUNT_BY_ROLE": {k: by_role.get(k, 0) for k in sorted(set(list(EXPECTED_ROLE_COUNTS) + list(by_role)))},
        "AREA_BY_ROLE_M2": {k: round(v, 3) for k, v in sorted(area.items())},
        "NAMED_BY_THE_DRAWING": sum(1 for c in comps if c["NAME"]),
        "RECONCILES": len(comps) == (len(comps) - slivers) + slivers,
        "MATCHES_EXPECTED_ROLE_COUNTS": {k: by_role.get(k, 0) == v for k, v in EXPECTED_ROLE_COUNTS.items()},
    }
    return rec


# ------------------------------------------------------------------ the derived skirting, split by decision state
def skirting(floors):
    """Skirting length per component: its own perimeter less the doors on it.

    This is NOT in the frozen takeoff, and it is not one number either.  It splits by what still has to be
    decided: dry named internal spaces are candidates, while stair, landing and unnamed spaces cannot even be
    candidates until their floor finish is settled.  Wet rooms are excluded outright - US-18 puts porcelain on
    the whole face, and a tiled face carries no skirting.
    """
    rows = []
    for f in floors:
        for r in f["COMPONENTS"]:
            if r["ROLE"] not in SKIRTING_ROLES:
                continue
            bucket, bucket_ar = SKIRTING_BUCKETS[r["ROLE"]]
            net = r["PERIMETER_M"] - r["DOOR_WIDTH_ON_BOUNDARY_M"]
            rows.append({"FLOOR": r["FLOOR"], "FLOOR_AR": r["FLOOR_AR"], "COMPONENT_REF": r["COMPONENT_REF"],
                         "ROOM": r["NAME"], "ROLE": r["ROLE"], "BUCKET": bucket, "BUCKET_AR": bucket_ar,
                         "PERIMETER_M": r["PERIMETER_M"],
                         "DOOR_DEDUCTION_M": r["DOOR_WIDTH_ON_BOUNDARY_M"],
                         "SKIRTING_LENGTH_M": round(net, 3),
                         "STATUS": "DERIVED_NOT_IN_FROZEN_TAKEOFF",
                         "BOQ_INCLUDED": False,
                         "RULE": "perimeter less the door openings on the same component",
                         "NOTES": "provisional until the owner confirms scope and height; wet and service rooms "
                                  "are excluded under US-18"})
    return rows


def skirting_buckets(rows):
    out = {}
    for b, _ar in SKIRTING_BUCKETS.values():
        mine = [r for r in rows if r["BUCKET"] == b]
        out[b] = {"COMPONENTS": len(mine), "LENGTH_M": round(sum(r["SKIRTING_LENGTH_M"] for r in mine), 3)}
    out["TOTAL_LENGTH_M"] = round(sum(v["LENGTH_M"] for v in out.values() if isinstance(v, dict)), 3)
    out["IS_ONE_DECISION_READY_QUANTITY"] = False
    return out


PROFILE_STEEL = {
    "ITEM": "بروفايل / حديد مشغول", "UNIT": "lm", "QUANTITY": None,
    "STATUS": "SOURCE_REQUIRED", "BOQ_INCLUDED": False,
    "WHY": "no profile, balustrade or handrail is drawn on the three issued plans.  The historical workbook "
           "claims 57 lm; that claim is FIELD_VERIFICATION_REQUIRED and is not turned into a quantity here",
}


# ------------------------------------------------------------------ the one window whose category has no authority
def window_authority(w):
    """Whether the guide category applied to this window is supported, and by what.

    A category chosen from an explicit room label is supported by the label.  A category chosen from floor area
    is supported only while the area is inside that category's own band; outside it, nothing supports the height
    and the answer is rank 5 of the ladder - ask the owner.
    """
    cat = WS.TABLE.get(w["GUIDE_CATEGORY"]) if w["GUIDE_CATEGORY"] else None
    band = cat["AREA_BAND_M2"] if cat else None
    basis = w["CATEGORY_BASIS"]
    area = w["ROOM_AREA_M2"]
    if basis == "EXPLICIT_LABEL_MAP":
        return {"AUTHORITY": "EXPLICIT_LABEL_ON_THE_DRAWING", "SUPPORTED": True, "BAND_M2": band,
                "WHY": f"the drawing labels this space {w['ROOM']!r}, and a label is evidence the area band does "
                       f"not override.  The band is advisory for a label-matched category",
                "HEIGHT_PRIORITY_RANK": 4}
    if band and area is not None and not (band[0] <= area <= band[1]):
        return {"AUTHORITY": "NONE", "SUPPORTED": False, "BAND_M2": band, "COMPONENT_AREA_M2": area,
                "WHY": f"the category was chosen from floor area, and {area} m2 is outside this category's band "
                       f"of {band[0]}-{band[1]} m2.  No owner override and no approved guide exception exists",
                "REQUIRED": "EXPLICIT_OWNER_CONFIRMATION or an approved guide exception",
                "HEIGHT_PRIORITY_RANK": 5}
    return {"AUTHORITY": "AREA_INSIDE_THE_GUIDE_BAND", "SUPPORTED": True, "BAND_M2": band,
            "COMPONENT_AREA_M2": area, "HEIGHT_PRIORITY_RANK": 4}


def window_rows(frozen):
    rows = []
    for w in frozen["WINDOW_REGISTER"]:
        a = window_authority(w)
        r = dict(w)
        r["HEIGHT_AUTHORITY"] = a
        r["BOQ_INCLUDED"] = bool(a["SUPPORTED"])
        if a["SUPPORTED"]:
            r["HEIGHT_STATUS"] = "GUIDE_HEIGHT_APPLIED"
            r["STATUS"] = "DRAWING_SCOPE_ONLY"
        else:
            r["HEIGHT_STATUS"] = "ASK_THE_OWNER"
            r["HEIGHT_SOURCE"] = "ASK_THE_OWNER (priority rank 5)"
            r["STATUS"] = "OWNER_INPUT_REQUIRED"
            r["PROVISIONAL_AREA_M2"] = w["AREA_M2"]
        rows.append(r)
    return rows


def build():
    a = FT.assemble()
    frozen = frozen_record()
    floors = component_model(a["ents"])
    sk = skirting(floors)
    windows = window_rows(frozen)
    m = {"assembly": a, "frozen": frozen, "floors": floors, "skirting": sk,
         "skirting_buckets": skirting_buckets(sk), "windows": windows, "census": census(floors)}
    m["project_total_m2"] = round(sum(f["FLOOR_TOTAL_M2"] for f in floors), 4)
    comps = [c for f in floors for c in f["COMPONENTS"]]
    m["floor_finish"] = {
        "PORCELAIN_M2": round(sum(c["AREA_M2"] for c in comps if c["ROLE"] in PORCELAIN_FLOOR_ROLES), 3),
        "PENDING_M2": round(sum(c["AREA_M2"] for c in comps if c["ROLE"] in PENDING_FLOOR_ROLES), 3),
        "PENDING_BY_ROLE": {role: {"COMPONENTS": sum(1 for c in comps if c["ROLE"] == role),
                                   "AREA_M2": round(sum(c["AREA_M2"] for c in comps if c["ROLE"] == role), 3)}
                            for role in PENDING_FLOOR_ROLES},
        "ENCLOSED_M2": round(sum(c["AREA_M2"] for c in comps if c["ROLE"] in ENCLOSED_ROLES), 3),
    }
    m["aluminium_split"] = {
        "SUPPORTED_M2": round(sum(w["AREA_M2"] for w in windows if w["BOQ_INCLUDED"]), 4),
        "UNSUPPORTED_M2": round(sum(w["AREA_M2"] for w in windows if not w["BOQ_INCLUDED"]), 4),
        "FROZEN_TOTAL_M2": a["totals"]["ALUMINIUM_M2"],
        "UNSUPPORTED_WINDOWS": [w["WINDOW_ID"] for w in windows if not w["BOQ_INCLUDED"]],
    }
    return m


# ------------------------------------------------------------------ the workbook
SHEETS = ["ملخص الحصر", "حصر المساحات", "الأرضيات", "كسوة الجدران", "النعلة والبروفايل", "العازل",
          "الأسقف", "الأبواب والشبابيك", "المباني", "المساح والصبغ", "السطح والخارجي",
          "BOQ حسب البند", "المصادر والمراجعة"]
HEAD_FILL = "2F5597"
BAND_FILL = "D9E2F3"
TOTAL_FILL = "FFF2CC"
PEND_FILL = "FBE4D5"


def _style(ws, widths, freeze="A2"):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws.sheet_view.rightToLeft = True
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=HEAD_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for col, w in zip("ABCDEFGHIJKLMNOP", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = freeze


def _fill_row(ws, r, colour, bold=True, last="L"):
    from openpyxl.styles import Font, PatternFill
    for c in ws[r]:
        if c.column_letter > last:
            continue
        c.fill = PatternFill("solid", fgColor=colour)
        if bold:
            c.font = Font(bold=True)


def area_sheet(ws, m):
    """FLOOR -> COMPONENT -> the rectangles it is made of -> component total -> floor total -> project total."""
    ws.append(["الدور", "رمز المكوّن", "الاسم على المخطط", "التصنيف", "الجزء", "الطول (م)", "العرض (م)",
               "المساحة (م²)", "طريقة الحساب", "مساحة المضلع من الكاد (م²)", "الفرق (م²)", "الأساس / ملاحظات"])
    comp_cell, floor_cell = {}, {}
    for f in m["floors"]:
        for r in f["COMPONENTS"]:
            first = ws.max_row + 1
            for p in r["PARTS"]:
                ws.append([r["FLOOR_AR"], r["COMPONENT_REF"], r["NAME"] or "—", r["ROLE_AR"], p["PART"],
                           p["RAW_LENGTH_M"], p["RAW_WIDTH_M"], None, p["CALCULATION"], None, None,
                           r["MEASUREMENT_BASIS"]])
                rr = ws.max_row
                ws.cell(rr, 8).value = f"=F{rr}*G{rr}"
                for col in (6, 7):
                    ws.cell(rr, col).number_format = "0.0000"
                ws.cell(rr, 8).number_format = "0.000"
            last = ws.max_row
            ws.append([r["FLOOR_AR"], r["COMPONENT_REF"], r["NAME"] or "—", r["ROLE_AR"], "إجمالي المكوّن",
                       None, None, None, f"مجموع {len(r['PARTS'])} جزء",
                       r["CAD_POLYGON_AREA_M2"], None,
                       "المساحة المجمدة كما هي؛ التقسيم يوضح الحساب ولا يغيّره"])
            t = ws.max_row
            ws.cell(t, 8).value = f"=SUM(H{first}:H{last})"
            ws.cell(t, 11).value = f"=H{t}-J{t}"
            _fill_row(ws, t, BAND_FILL)
            comp_cell[r["COMPONENT_REF"]] = f"'{ws.title}'!H{t}"
        refs = [comp_cell[r["COMPONENT_REF"]].split("!")[1] for r in f["COMPONENTS"]]
        ws.append([f["FLOOR_AR"], None, None, None, "إجمالي الدور", None, None, None,
                   f"{len(f['COMPONENTS'])} مكوّن متصل", None, None,
                   f"منسوب {f['LEVEL_M']} م — كل متر مربع في نافذة المخطط يخص مكوّناً واحداً فقط"])
        t = ws.max_row
        ws.cell(t, 8).value = "=" + "+".join(refs)
        _fill_row(ws, t, TOTAL_FILL)
        floor_cell[f["FLOOR"]] = f"'{ws.title}'!H{t}"
    ws.append([None, None, None, None, "إجمالي المشروع", None, None, None,
               "مجموع الأدوار الثلاثة", None, None,
               "يشمل مادة الجدران والشظايا؛ لا يُقرأ كمساحة أرضيات"])
    t = ws.max_row
    ws.cell(t, 8).value = "=" + "+".join(c.split("!")[1] for c in floor_cell.values())
    _fill_row(ws, t, TOTAL_FILL)
    _style(ws, (12, 13, 20, 20, 10, 11, 11, 13, 26, 20, 11, 46))
    return {"COMPONENT": comp_cell, "FLOOR": floor_cell, "PROJECT": f"'{ws.title}'!H{t}"}


def floors_sheet(ws, m, ref):
    """Porcelain floor where the finish is settled, and a clearly separate pending block where it is not."""
    ws.append(["الدور", "رمز المكوّن", "الاسم على المخطط", "التصنيف", "البند", "الوحدة", "الكمية (م²)",
               "مصدر الكمية", "معرف القاعدة", "حالة البند", "مشمول في BOQ الأساس", "ملاحظات"])
    first = ws.max_row + 1
    for f in m["floors"]:
        for r in f["COMPONENTS"]:
            if r["ROLE"] not in PORCELAIN_FLOOR_ROLES:
                continue
            wet = r["ROLE"] in WET_ROLES
            ws.append([r["FLOOR_AR"], r["COMPONENT_REF"], r["NAME"] or "—", r["ROLE_AR"],
                       "بورسلان / سيراميك أرضيات", "م²", None, "مساحة المكوّن في ورقة حصر المساحات",
                       "US-18" if wet else "—", "FINAL_QUANTITY_AVAILABLE", "نعم",
                       "US-18 يغطي الغرف الرطبة والخدمية فقط؛ الغرف الجافة كمية مقاسة بلا قاعدة أوربن"])
            ws.cell(ws.max_row, 7).value = "=" + ref["COMPONENT"][r["COMPONENT_REF"]]
    last = ws.max_row
    ws.append([None, None, None, None, "AR-FL-01 — إجمالي الأرضيات المعتمدة", "م²", None, None, None,
               "FINAL_QUANTITY_AVAILABLE", "نعم", "غرف داخلية + رطبة + مطبخ"])
    t = ws.max_row
    ws.cell(t, 7).value = f"=SUM(G{first}:G{last})"
    _fill_row(ws, t, TOTAL_FILL)
    base = f"'{ws.title}'!G{t}"

    ws.append([])
    ws.append(["تشطيب أرضيات غير محسوم — مقاس ومحفوظ وخارج BOQ الأساس", None, None, None, None, None, None,
               None, None, "FINISH_CLASSIFICATION_PENDING", "لا", None])
    _fill_row(ws, ws.max_row, PEND_FILL)
    pend_cells = {}
    for role in PENDING_FLOOR_ROLES:
        pf = ws.max_row + 1
        for f in m["floors"]:
            for r in f["COMPONENTS"]:
                if r["ROLE"] != role:
                    continue
                ws.append([r["FLOOR_AR"], r["COMPONENT_REF"], r["NAME"] or "—", r["ROLE_AR"],
                           "تشطيب أرضية غير محدد", "م²", None, "مساحة المكوّن في ورقة حصر المساحات", "—",
                           "FINISH_CLASSIFICATION_PENDING", "لا",
                           "الدرج يأخذ درجات لا أرضية؛ والفراغ غير المسمى يحتاج تصنيفاً من المالك"])
                ws.cell(ws.max_row, 7).value = "=" + ref["COMPONENT"][r["COMPONENT_REF"]]
        pl = ws.max_row
        ws.append([None, None, None, ROLES_AR[role], f"إجمالي {ROLES_AR[role]}", "م²", None, None, None,
                   "FINISH_CLASSIFICATION_PENDING", "لا", None])
        ws.cell(ws.max_row, 7).value = f"=SUM(G{pf}:G{pl})"
        _fill_row(ws, ws.max_row, PEND_FILL)
        pend_cells[role] = f"G{ws.max_row}"
    ws.append([None, None, None, None, "AR-FL-PENDING — إجمالي المعلّق", "م²", None, None, None,
               "FINISH_CLASSIFICATION_PENDING", "لا", "لا يدخل بند الأرضيات المعتمد ولا يُسعّر"])
    t2 = ws.max_row
    ws.cell(t2, 7).value = "=" + "+".join(pend_cells.values())
    _fill_row(ws, t2, PEND_FILL)
    _style(ws, (12, 13, 20, 20, 30, 8, 13, 30, 14, 30, 18, 52))
    return {"BASE": base, "PENDING": f"'{ws.title}'!G{t2}",
            "PENDING_BY_ROLE": {k: f"'{ws.title}'!{v}" for k, v in pend_cells.items()}}


def ceiling_sheet(ws, m, ref):
    """The ceiling has a measured area and an unsettled finish, and the sheet keeps the two apart."""
    ws.append(["الدور", "رمز المكوّن", "الاسم على المخطط", "التصنيف", "مساحة السقف (م²)", "حالة المساحة",
               "نوع التشطيب", "حالة التشطيب", "مشمول في BOQ الأساس", "ملاحظات"])
    first = ws.max_row + 1
    for f in m["floors"]:
        for r in f["COMPONENTS"]:
            if r["ROLE"] not in ENCLOSED_ROLES:
                continue
            ws.append([r["FLOOR_AR"], r["COMPONENT_REF"], r["NAME"] or "—", r["ROLE_AR"], None,
                       "FINAL_QUANTITY_AVAILABLE", "غير محدد", "FINISH_CLASSIFICATION_PENDING", "لا",
                       "المساحة المسقوفة مقاسة ونهائية؛ نوع السقف يحتاج مخطط أسقف أو تأكيد المالك"])
            ws.cell(ws.max_row, 5).value = "=" + ref["COMPONENT"][r["COMPONENT_REF"]]
    last = ws.max_row
    ws.append([None, None, None, None, None, "FINAL_QUANTITY_AVAILABLE", "غير محدد",
               "FINISH_CLASSIFICATION_PENDING", "لا", "AR-CL-01 — مساحة الأسقف المحصورة"])
    t = ws.max_row
    ws.cell(t, 5).value = f"=SUM(E{first}:E{last})"
    _fill_row(ws, t, PEND_FILL, last="J")
    _style(ws, (12, 13, 20, 20, 16, 26, 14, 30, 18, 56))
    return f"'{ws.title}'!E{t}"


def wall_sheet(ws, m):
    perim = {r["COMPONENT_REF"]: r["PERIMETER_M"] for f in m["floors"] for r in f["COMPONENTS"]}
    ws.append(["الدور", "رمز المكوّن", "الغرفة", "محيط الغرفة (م)", "الارتفاع (م)", "الكمية الإجمالية (م²)",
               "خصم الفتحات (م²)", "صافي بورسلان الجدران (م²)", "طرطشة تحضير (م²)", "عدد الفتحات",
               "معرف القاعدة", "الحالة", "ملاحظات"])
    for x in m["assembly"]["wet"]:
        ws.append([FLOORS_AR[x["FLOOR"]], x["ROOM_REF"], x["ROOM"],
                   perim.get(x["ROOM_REF"], x["CERAMIC_HOST_LENGTH_M"]), x["HEIGHT_M"],
                   None, x["OPENING_DEDUCTIONS_M2"], None, None, x["OPENINGS_ON_THIS_ROOM"],
                   "US-18", "FINAL_QUANTITY_AVAILABLE",
                   "الطرطشة خلف نفس الأوجه التي يغطيها البورسلان، فالمضيف واحد"])
        r = ws.max_row
        ws.cell(r, 6).value = f"=D{r}*E{r}"
        ws.cell(r, 8).value = f"=F{r}-G{r}"
        ws.cell(r, 9).value = f"=H{r}"
    ws.append([None, None, "الإجمالي", None, None, None, None, None, None, None, "US-18",
               "FINAL_QUANTITY_AVAILABLE", None])
    t = ws.max_row
    for idx, col in ((6, "F"), (7, "G"), (8, "H"), (9, "I")):
        ws.cell(t, idx).value = f"=SUM({col}2:{col}{t - 1})"
    _fill_row(ws, t, TOTAL_FILL, last="M")
    _style(ws, (12, 13, 18, 16, 12, 18, 16, 20, 18, 12, 12, 26, 44))
    return {"PORCELAIN": f"'{ws.title}'!H{t}", "PREP": f"'{ws.title}'!I{t}"}


def skirting_sheet(ws, m):
    ws.append(["المجموعة", "الدور", "رمز المكوّن", "الاسم على المخطط", "المحيط (م)", "خصم عرض الأبواب (م)",
               "طول النعلة (م.ط)", "الحالة", "مشمول في BOQ الأساس", "ملاحظات"])
    cells = {}
    for role in SKIRTING_ROLES:
        bucket, bucket_ar = SKIRTING_BUCKETS[role]
        mine = [x for x in m["skirting"] if x["BUCKET"] == bucket]
        first = ws.max_row + 1
        for x in mine:
            ws.append([bucket_ar, FLOORS_AR[x["FLOOR"]], x["COMPONENT_REF"], x["ROOM"] or "—",
                       x["PERIMETER_M"], x["DOOR_DEDUCTION_M"], None, x["STATUS"], "لا", x["NOTES"]])
            ws.cell(ws.max_row, 7).value = f"=E{ws.max_row}-F{ws.max_row}"
        last = ws.max_row
        ws.append([bucket_ar, None, None, f"{len(mine)} مكوّن", None, "إجمالي المجموعة", None,
                   "DERIVED_NOT_IN_FROZEN_TAKEOFF", "لا",
                   "مؤقت حتى يؤكد المالك النطاق والارتفاع"])
        ws.cell(ws.max_row, 7).value = f"=SUM(G{first}:G{last})"
        _fill_row(ws, ws.max_row, PEND_FILL, last="J")
        cells[bucket] = f"'{ws.title}'!G{ws.max_row}"
    ws.append([None, None, None, None, None, "مجموع المجموعات الثلاث", None,
               "DERIVED_NOT_IN_FROZEN_TAKEOFF", "لا",
               "ليست كمية واحدة جاهزة للقرار — ثلاث مجموعات بحالات مختلفة"])
    ws.cell(ws.max_row, 7).value = "=" + "+".join(c.split("!")[1] for c in cells.values())
    _fill_row(ws, ws.max_row, PEND_FILL, last="J")
    cells["TOTAL"] = f"'{ws.title}'!G{ws.max_row}"
    ws.append([])
    ws.append([PROFILE_STEEL["ITEM"], None, None, None, None, None, None, PROFILE_STEEL["STATUS"], "لا",
               PROFILE_STEEL["WHY"]])
    _style(ws, (26, 12, 13, 20, 14, 20, 18, 30, 18, 60))
    return cells


def insulation_sheet(ws, m):
    rp = m["assembly"]["rp"]
    ws.append(["الموقع", "البند", "الوحدة", "الطول (م)", "العرض (م)", "الكمية", "مصدر الكمية",
               "الحالة", "ملاحظات"])
    ws.append(["السطح", "عازل مائي للسطح", "م²", None, None, rp["WATERPROOFING_AREA_M2"],
               "مساحة السطح المحصورة داخل الدروة", "FINAL_QUANTITY_AVAILABLE",
               "المساحة نهائية؛ نوع العازل ونظامه قرار مواصفات"])
    roof = f"'{ws.title}'!F{ws.max_row}"
    ws.append(["الغرف الرطبة", "عازل أرضيات الغرف الرطبة", "م²", None, None, None,
               "غير مرسوم على المخططات الثلاثة", "SOURCE_REQUIRED",
               "لا يوجد تفصيل عزل أرضيات في المخططات المصدرية — لا تُخترع كمية"])
    ws.append([])
    ws.append(["مرجع تاريخي", "أساس تجاري تاريخي 950.22 م²", "—", None, None, None,
               "Sheet1!F28 = SUM(F5:F24) ويُستهلك في I11 '=1.5*F28'", "HISTORICAL_COMMERCIAL_BASIS",
               "ليست كمية غشاء: أحد مكوناتها 600.00 وهي مساحة القسيمة. لا تُقارن بكمية مقاسة ولا تدخل أي بند"])
    _style(ws, (16, 30, 8, 12, 12, 14, 40, 30, 62))
    return roof


def openings_sheet(ws, m):
    ws.append(["النوع", "الرمز", "الدور", "الفراغ", "العرض (م)", "مصدر العرض", "الارتفاع (م)",
               "مصدر الارتفاع", "المساحة (م²)", "تصنيف الدليل", "سند التصنيف", "الحالة",
               "مشمول في BOQ الأساس", "ملاحظات"])
    first_d = ws.max_row + 1
    for d in m["assembly"]["doors"]:
        ws.append(["باب", d["DOOR_ID"], FLOORS_AR[d["FLOOR"]], "—", d["WIDTH_M"], d["WIDTH_SOURCE"],
                   d["HEIGHT_M"], d["HEIGHT_SOURCE"], None, "—", "—", "PRICING_BASIS_REQUIRED", "نعم",
                   d["NOTES"]])
        r = ws.max_row
        ws.cell(r, 9).value = f"=E{r}*G{r}"
    last_d = ws.max_row
    ws.append([None, "إجمالي الأبواب", None, f"العدد {len(m['assembly']['doors'])}", None, None, None, None,
               None, None, None, "PRICING_BASIS_REQUIRED", "نعم",
               "كمية الفتحة فقط؛ نوع الضلفة ومادتها قرار شراء"])
    td = ws.max_row
    ws.cell(td, 9).value = f"=SUM(I{first_d}:I{last_d})"
    _fill_row(ws, td, BAND_FILL, last="N")
    sup_first = ws.max_row + 1
    for w in [x for x in m["windows"] if x["BOQ_INCLUDED"]]:
        ws.append(["شباك", w["WINDOW_ID"], FLOORS_AR[w["FLOOR"]], w["ROOM"], w["WIDTH_M"], w["WIDTH_SOURCE"],
                   w["HEIGHT_M"], w["HEIGHT_SOURCE"], None, w["GUIDE_CATEGORY"],
                   w["HEIGHT_AUTHORITY"]["AUTHORITY"], w["STATUS"], "نعم", w["NOTES"]])
        r = ws.max_row
        ws.cell(r, 9).value = f"=E{r}*G{r}"
    sup_last = ws.max_row
    ws.append([None, "AR-AL-01 — شبابيك بسند تصنيف", None, f"العدد {sup_last - sup_first + 1}", None, None,
               None, None, None, None, None, "DRAWING_SCOPE_ONLY", "نعم",
               "نطاق المخطط فقط (US-21): زجاج الموقع يتجاوزه عادةً وليست كمية توريد"])
    tw = ws.max_row
    ws.cell(tw, 9).value = f"=SUM(I{sup_first}:I{sup_last})"
    _fill_row(ws, tw, TOTAL_FILL, last="N")
    pend = [x for x in m["windows"] if not x["BOQ_INCLUDED"]]
    pf = ws.max_row + 1
    for w in pend:
        ws.append(["شباك", w["WINDOW_ID"], FLOORS_AR[w["FLOOR"]], w["ROOM"], w["WIDTH_M"], w["WIDTH_SOURCE"],
                   w["HEIGHT_M"], w["HEIGHT_SOURCE"], None, w["GUIDE_CATEGORY"],
                   w["HEIGHT_AUTHORITY"]["AUTHORITY"], w["STATUS"], "لا",
                   w["HEIGHT_AUTHORITY"]["WHY"]])
        r = ws.max_row
        ws.cell(r, 9).value = f"=E{r}*G{r}"
        _fill_row(ws, r, PEND_FILL, bold=False, last="N")
    pl = ws.max_row
    ws.append([None, "AR-AL-PENDING — بلا سند تصنيف", None, f"العدد {len(pend)}", None, None, None, None,
               None, None, None, "OWNER_INPUT_REQUIRED", "لا",
               "العرض مقاس؛ الارتفاع بانتظار قرار المالك — الكمية مؤقتة ولا تُسعّر"])
    tp = ws.max_row
    ws.cell(tp, 9).value = f"=SUM(I{pf}:I{pl})" if pend else 0
    _fill_row(ws, tp, PEND_FILL, last="N")
    _style(ws, (10, 24, 12, 18, 12, 30, 13, 34, 13, 22, 30, 26, 18, 56))
    return {"DOORS": f"'{ws.title}'!I{td}", "WINDOWS_SUPPORTED": f"'{ws.title}'!I{tw}",
            "WINDOWS_PENDING": f"'{ws.title}'!I{tp}"}


def blockwork_sheet(ws, m):
    ws.append(["الدور", "السماكة (مم)", "هوية العنصر", "طول الجدار (م)", "الارتفاع (م)",
               "الكمية الإجمالية (م²)", "خصم الفتحات (م²)", "الصافي (م²)", "الحالة", "ملاحظات"])
    cells = {150: [], 200: []}
    for b in m["assembly"]["block"]:
        ws.append([FLOORS_AR[b["FLOOR"]], b["THICKNESS_MM"], "جدار مباني", b["PLAN_LENGTH_M"], b["HEIGHT_M"],
                   None, b["OPENING_DEDUCTIONS_M2"], None, "FINAL_QUANTITY_AVAILABLE",
                   "الخصم موزّع على الجدران المؤكدة بنسبة طولها"])
        r = ws.max_row
        ws.cell(r, 6).value = f"=D{r}*E{r}"
        ws.cell(r, 8).value = f"=F{r}-G{r}"
        cells[b["THICKNESS_MM"]].append(f"H{r}")
    tot = {}
    for t_mm in (150, 200):
        ws.append([None, t_mm, f"إجمالي مباني {t_mm // 10} سم", None, None, None, None, None,
                   "FINAL_QUANTITY_AVAILABLE", None])
        r = ws.max_row
        ws.cell(r, 8).value = "=" + "+".join(cells[t_mm])
        _fill_row(ws, r, TOTAL_FILL, last="J")
        tot[t_mm] = f"'{ws.title}'!H{r}"
    ws.append([])
    ws.append(["مستبعد — ليست جدراناً", None, None, None, None, None, None, None, "NOT_BILLED",
               "سماكات لا يبعّدها المخطط أبداً، تنشأ عند تقاطع جدارين أو رسم خط مرتين"])
    _fill_row(ws, ws.max_row, BAND_FILL, last="J")
    for x in m["assembly"]["excluded"]:
        ws.append([FLOORS_AR[x["FLOOR"]], x["THICKNESS_MM"], x["OBJECT_IDENTITY"], x["LENGTH_M"], None,
                   None, None, None, "NOT_BILLED", x["WHY"]])
    _style(ws, (12, 14, 18, 16, 12, 18, 18, 14, 26, 56))
    return tot


def plaster_sheet(ws, m):
    ws.append(["الدور", "طول الأوجه الداخلية (م)", "طول الأوجه الخارجية (م)", "الارتفاع (م)",
               "مساح داخلي إجمالي (م²)", "خصم الأبواب (م²)", "مساح داخلي صافي (م²)",
               "أوجه مكسوة بالبورسلان (م²)", "صبغ داخلي (م²)", "مساح خارجي إجمالي (م²)",
               "خصم الشبابيك (م²)", "مساح خارجي صافي (م²)", "صبغ خارجي (م²)", "الحالة"])
    first = ws.max_row + 1
    for p in m["assembly"]["pp"]:
        ws.append([FLOORS_AR[p["FLOOR"]], p["INTERNAL_FACE_LENGTH_M"], p["EXTERNAL_FACE_LENGTH_M"], p["HEIGHT_M"],
                   None, round(p["INTERNAL_PLASTER_GROSS_M2"] - p["INTERNAL_PLASTER_NET_M2"], 3), None,
                   p["TILED_FACE_AREA_M2"], None, None, p["WINDOW_DEDUCTION_M2"], None, None,
                   "FINAL_QUANTITY_AVAILABLE"])
        r = ws.max_row
        ws.cell(r, 5).value = f"=B{r}*D{r}"
        ws.cell(r, 7).value = f"=E{r}-F{r}"
        ws.cell(r, 9).value = f"=G{r}-H{r}"
        ws.cell(r, 10).value = f"=C{r}*D{r}"
        ws.cell(r, 12).value = f"=J{r}-K{r}"
        ws.cell(r, 13).value = f"=L{r}"
    last = ws.max_row
    ws.append(["الإجمالي"] + [None] * 13)
    t = ws.max_row
    for idx, col in ((5, "E"), (6, "F"), (7, "G"), (8, "H"), (9, "I"), (10, "J"), (11, "K"), (12, "L"), (13, "M")):
        ws.cell(t, idx).value = f"=SUM({col}{first}:{col}{last})"
    _fill_row(ws, t, TOTAL_FILL, last="N")
    _style(ws, (12, 20, 20, 12, 20, 16, 20, 20, 16, 20, 16, 20, 16, 26))
    return {"INT_PLASTER": f"'{ws.title}'!G{t}", "INT_PAINT": f"'{ws.title}'!I{t}",
            "EXT_PLASTER": f"'{ws.title}'!L{t}", "EXT_PAINT": f"'{ws.title}'!M{t}"}


def roof_sheet(ws, m):
    rp, ref = m["assembly"]["rp"], {}
    ws.append(["الموقع", "البند", "الوحدة", "الطول (م)", "الارتفاع / العرض (م)", "المعامل", "الكمية",
               "مصدر الكمية", "الحالة", "مشمول في BOQ الأساس", "ملاحظات"])
    ws.append(["السطح", "مساحة السطح", "م²", None, None, None, rp["ROOF_AREA_ENCLOSED_M2"],
               "المساحة المحصورة داخل حد السطح المرسوم", "FINAL_QUANTITY_AVAILABLE", "نعم",
               "مادة التشطيب غير محددة على المخطط"])
    ref["ROOF"] = f"'{ws.title}'!G{ws.max_row}"
    ws.append(["السطح", "مباني الدروة", "م²", rp["PARAPET_RUN_M"], rp["PARAPET_HEIGHT_M"], 1, None,
               "طول الدروة × ارتفاعها (AR-06)", "FINAL_QUANTITY_AVAILABLE", "نعم", rp["PARAPET_RUN_BASIS"]])
    r = ws.max_row
    ws.cell(r, 7).value = f"=D{r}*E{r}*F{r}"
    ref["PARAPET_BLOCK"] = f"'{ws.title}'!G{r}"
    for item, key in (("مساح الدروة (الوجهين)", "PARAPET_PLASTER"), ("صبغ الدروة (الوجهين)", "PARAPET_PAINT")):
        ws.append(["السطح", item, "م²", rp["PARAPET_RUN_M"], rp["PARAPET_HEIGHT_M"], 2, None,
                   "طول الدروة × ارتفاعها × وجهين", "FINAL_QUANTITY_AVAILABLE", "نعم",
                   rp["PLASTER_AND_PAINT_BASIS"]])
        r = ws.max_row
        ws.cell(r, 7).value = f"=D{r}*E{r}*F{r}"
        ref[key] = f"'{ws.title}'!G{r}"
    ws.append([])
    ws.append(["المساحات الخارجية والمواقف", None, None, None, None, None, None, None, None, None, None])
    _fill_row(ws, ws.max_row, BAND_FILL, last="K")
    first = ws.max_row + 1
    for f in m["floors"]:
        for x in f["COMPONENTS"]:
            if x["ROLE"] != "EXTERNAL_OR_OPEN":
                continue
            ws.append([x["FLOOR_AR"], x["NAME"] or "—", "م²", None, None, None, x["AREA_M2"],
                       f"مساحة مقاسة — {x['COMPONENT_REF']}", "FINISH_CLASSIFICATION_PENDING", "لا",
                       "المساحة نهائية؛ لا مادة تشطيب منصوصة على المخطط (AR-08)"])
    last = ws.max_row
    ws.append([None, "إجمالي المساحات الخارجية", "م²", None, None, None, None, None,
               "FINISH_CLASSIFICATION_PENDING", "لا", None])
    t = ws.max_row
    ws.cell(t, 7).value = f"=SUM(G{first}:G{last})"
    _fill_row(ws, t, PEND_FILL, last="K")
    ref["EXTERNAL"] = f"'{ws.title}'!G{t}"
    _style(ws, (22, 26, 8, 14, 18, 10, 14, 34, 30, 18, 54))
    return ref


# ------------------------------------------------------------------ the BOQ: a base section and a provisional one
def boq_items(m, ref):
    t = m["assembly"]["totals"]
    sk = m["skirting_buckets"]
    B, P = "BASE", "PROVISIONAL"
    I = [
        (B, "AR-FL-01", "الأرضيات", "بورسلان / سيراميك أرضيات — غرف داخلية ورطبة ومطبخ", "م²",
         ref["FLOORS_BASE"], m["floor_finish"]["PORCELAIN_M2"],
         "مجموع مساحات المكوّنات المصنّفة في ورقة الأرضيات", None, "FINAL_QUANTITY_AVAILABLE", True,
         "تقسيم البورسلان عن السيراميك قرار أسعار لا قياس؛ US-18 يخص الرطبة منها فقط"),
        (B, "AR-FL-01-W", "الأرضيات", "منها أرضيات الغرف الرطبة والخدمية", "م²", None,
         t["WET_AND_SERVICE_FLOOR_AREA_M2"], "مساحات مقاسة", "US-18", "FINAL_QUANTITY_AVAILABLE", False,
         "مضمّنة داخل AR-FL-01 ولا تُجمع معه"),
        (B, "AR-WP-01", "كسوة الجدران", "بورسلان جدران بكامل الارتفاع في الغرف الرطبة", "م²",
         ref["PORCELAIN"], t["WALL_PORCELAIN_NET_M2"], "محيط الغرفة × 3.60 ناقص فتحاتها", "US-18",
         "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-WP-02", "كسوة الجدران", "طرطشة تحضير خلف بورسلان الجدران", "م²",
         ref["PREP"], t["TILE_PREPARATION_M2"], "نفس مضيف البورسلان", "US-18",
         "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-IN-01", "العازل", "عازل مائي للسطح", "م²", ref["INSULATION"], t["WATERPROOFING_M2"],
         "مساحة السطح المحصورة داخل الدروة", None, "FINAL_QUANTITY_AVAILABLE", True,
         "عزل أرضيات الغرف الرطبة غير مرسوم — طلب مصدر"),
        (B, "AR-BL-01", "المباني", "مباني 15 سم", "م²", ref["BLOCK150"], t["BLOCKWORK_150_NET_M2"],
         "طول الجدار × 3.60 ناقص الفتحات", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-BL-02", "المباني", "مباني 20 سم", "م²", ref["BLOCK200"], t["BLOCKWORK_200_NET_M2"],
         "طول الجدار × 3.60 ناقص الفتحات", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-PL-01", "المساح والصبغ", "مساح داخلي", "م²", ref["INT_PLASTER"], t["INTERNAL_PLASTER_NET_M2"],
         "أوجه الجدران الداخلية × 3.60 ناقص الأبواب", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-PT-01", "المساح والصبغ", "صبغ داخلي", "م²", ref["INT_PAINT"], t["INTERNAL_PAINT_M2"],
         "المساح الداخلي ناقص الأوجه المكسوة بالكامل", "US-18", "FINAL_QUANTITY_AVAILABLE", True,
         "الوجه المكسو بالبورسلان لا يُصبغ"),
        (B, "AR-PL-02", "المساح والصبغ", "مساح خارجي", "م²", ref["EXT_PLASTER"], t["EXTERNAL_PLASTER_NET_M2"],
         "أوجه الجدران الخارجية × 3.60 ناقص الشبابيك", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-PT-02", "المساح والصبغ", "صبغ خارجي", "م²", ref["EXT_PAINT"], t["EXTERNAL_PAINT_M2"],
         "مساحة المساح الخارجي", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-RF-01", "السطح والخارجي", "مباني الدروة", "م²", ref["PARAPET_BLOCK"], t["PARAPET_BLOCKWORK_M2"],
         "طول الدروة × 1.00 (AR-06)", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-RF-02", "السطح والخارجي", "مساح الدروة — الوجهين", "م²", ref["PARAPET_PLASTER"],
         t["PARAPET_PLASTER_M2"], "طول الدروة × 1.00 × 2", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-RF-03", "السطح والخارجي", "صبغ الدروة — الوجهين", "م²", ref["PARAPET_PAINT"],
         t["PARAPET_PAINT_M2"], "طول الدروة × 1.00 × 2", None, "FINAL_QUANTITY_AVAILABLE", True, ""),
        (B, "AR-AL-01", "الألمنيوم", "شبابيك ألمنيوم بسند تصنيف — نطاق المخطط", "م²",
         ref["WINDOWS_SUPPORTED"], m["aluminium_split"]["SUPPORTED_M2"],
         "عرض مقاس من المخطط × ارتفاع من الدليل ضمن نطاقه", WS.RULE_ID, "DRAWING_SCOPE_ONLY", True,
         "US-21: ليست كمية توريد؛ زجاج الموقع يتجاوز نطاق المخطط عادةً"),
        (B, "AR-DR-01", "الأبواب", "فتحات الأبواب", "م²", ref["DOORS"], t["DOOR_AREA_M2"],
         "عرض مقاس بين القائمين × 2.20 (AR-02)", None, "PRICING_BASIS_REQUIRED", True,
         f"{t['DOOR_COUNT']} باباً؛ مادة الضلفة ونوعها قرار شراء"),
        (P, "AR-FL-PENDING", "الأرضيات", "تشطيب أرضيات غير محسوم — درج وبسطات وفراغات غير مسماة", "م²",
         ref["FLOORS_PENDING"], m["floor_finish"]["PENDING_M2"],
         "مساحات مقاسة في ورقة حصر المساحات", None, "FINISH_CLASSIFICATION_PENDING", False,
         f"درج وبسطات {m['floor_finish']['PENDING_BY_ROLE']['STAIR_OR_LANDING']['AREA_M2']} م² + "
         f"فراغات غير مسماة {m['floor_finish']['PENDING_BY_ROLE']['UNNAMED_ON_DRAWING']['AREA_M2']} م² — "
         f"مقاسة ومحفوظة، وخارج بند الأرضيات حتى يحسم المالك التشطيب"),
        (P, "AR-CL-01", "الأسقف", "مساحة الأسقف المحصورة — نوع التشطيب غير محسوم", "م²",
         ref["CEILING"], t["CEILING_M2"], "المساحة المسقوفة المحصورة", None,
         "FINISH_CLASSIFICATION_PENDING", False,
         "المساحة نهائية؛ 'سقف عادي' ليس كمية نهائية بلا مخطط أسقف أو تأكيد المالك"),
        (P, "AR-AL-PENDING", "الألمنيوم", "شباك بلا سند تصنيف — ارتفاعه بانتظار المالك", "م²",
         ref["WINDOWS_PENDING"], m["aluminium_split"]["UNSUPPORTED_M2"],
         "عرض مقاس؛ الارتفاع بلا سند", WS.RULE_ID, "OWNER_INPUT_REQUIRED", False,
         " و ".join(m["aluminium_split"]["UNSUPPORTED_WINDOWS"]) +
         " — فئة الدليل استُخدمت خارج نطاق مساحتها"),
        (P, "AR-EX-01", "السطح والخارجي", "مساحات خارجية ومواقف", "م²", ref["EXTERNAL"],
         t["EXTERNAL_OPEN_AREA_M2"], "مساحات مقاسة (AR-08)", None, "FINISH_CLASSIFICATION_PENDING", False,
         "المساحة نهائية والمادة غير منصوصة"),
        (P, "AR-SK-01", "النعلة والبروفايل", "نعلة — غرف داخلية جافة مسماة (مرشحة)", "م.ط",
         ref["SK_DRY"], sk["DRY_NAMED_INTERNAL_CANDIDATE"]["LENGTH_M"],
         "محيط المكوّن ناقص عروض أبوابه", None, "DERIVED_NOT_IN_FROZEN_TAKEOFF", False,
         "مشتق من الهندسة المجمدة؛ مؤقت حتى يؤكد المالك النطاق والارتفاع"),
        (P, "AR-SK-02", "النعلة والبروفايل", "نعلة — درج وبسطات (معلّقة)", "م.ط",
         ref["SK_STAIR"], sk["STAIR_OR_LANDING_PENDING"]["LENGTH_M"],
         "محيط المكوّن ناقص عروض أبوابه", None, "DERIVED_NOT_IN_FROZEN_TAKEOFF", False,
         "لا يمكن أن تكون مرشحة قبل حسم تشطيب أرضية الدرج"),
        (P, "AR-SK-03", "النعلة والبروفايل", "نعلة — فراغات غير مسماة (معلّقة)", "م.ط",
         ref["SK_UNNAMED"], sk["UNNAMED_SPACE_PENDING"]["LENGTH_M"],
         "محيط المكوّن ناقص عروض أبوابه", None, "DERIVED_NOT_IN_FROZEN_TAKEOFF", False,
         "لا يمكن أن تكون مرشحة قبل تصنيف الفراغ"),
        (P, "AR-SK-04", "النعلة والبروفايل", "بروفايل / حديد مشغول", "م.ط", None, None,
         "لا يوجد على المخططات", None, "SOURCE_REQUIRED", False,
         "ادّعاء تاريخي بـ 57 م.ط يحتاج تحقق ميداني ولا يُحوَّل إلى كمية"),
    ]
    return I


def boq_sheet(ws, m, ref):
    ws.append(["القسم", "كود البند", "البند", "الوصف", "الوحدة", "الكمية المقاسة", "الهالك %", "كمية الشراء",
               "سعر الوحدة", "الإجمالي", "مصدر الكمية", "معرف القاعدة", "حالة الكمية",
               "مشمول في BOQ الأساس", "حالة الاعتماد", "ملاحظات"])
    cells, section_now = {}, None
    for sec, code, trade, desc, unit, cell, value, source, rule, status, included, notes in boq_items(m, ref):
        if sec != section_now:
            section_now = sec
            ws.append(["قسم", "BOQ الأساس — بنود جاهزة للقرار" if sec == "BASE"
                       else "قسم مؤقت / اختياري — غير جاهز للتسعير", None, None, None, None, None, None,
                       None, None, None, None, None, None, None])
            _fill_row(ws, ws.max_row, BAND_FILL if sec == "BASE" else PEND_FILL, last="P")
        ws.append([sec, code, trade, desc, unit, None, None, None, None, None, source, rule or "—", status,
                   "نعم" if included else "لا", "DRAFT", notes])
        r = ws.max_row
        ws.cell(r, 6).value = ("=" + cell) if cell else value
        ws.cell(r, 8).value = f'=IF(G{r}="","",F{r}*(1+G{r}/100))'
        ws.cell(r, 10).value = f'=IF(OR(H{r}="",I{r}=""),"",H{r}*I{r})'
        if not included:
            _fill_row(ws, r, PEND_FILL, bold=False, last="P")
        cells[code] = f"'{ws.title}'!F{r}"
    ws.append([None, None, None, "إجمالي المبالغ", None, None, None, None, None, None, None, None, None,
               None, "DRAFT", "يبقى فارغاً حتى تُدخل الأسعار — DS-01: لا خانة تجمع كمية وسعراً"])
    t = ws.max_row
    ws.cell(t, 10).value = f'=IF(COUNT(J2:J{t - 1})=0,"",SUM(J2:J{t - 1}))'
    _fill_row(ws, t, TOTAL_FILL, last="P")
    _style(ws, (12, 16, 18, 46, 8, 15, 10, 14, 12, 14, 38, 14, 30, 18, 14, 60))
    return cells


# ------------------------------------------------------------------ the checks, evaluated rather than asserted
def qa(m, wb, ref, boq_cells, export_rows):
    """Every check works something out from evidence: the workbook's own formulas, the export, the frozen file."""
    vals = WC.evaluate_all(wb)
    comps = [c for f in m["floors"] for c in f["COMPONENTS"]]
    t = m["assembly"]["totals"]
    checks = []
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    head = _git("rev-parse", "--short", "HEAD")

    def cell(ref_str):
        """The value a reader sees: an evaluated formula, or the literal the cell holds."""
        sheet, coord = ref_str.split("!")
        sheet = sheet.strip("'")
        return vals[(sheet, coord)] if (sheet, coord) in vals else wb[sheet][coord].value

    def add(cid, ok, inputs, tol, result, method):
        checks.append({"CHECK": cid, "PASS": bool(ok), "INPUTS": inputs, "TOLERANCE": tol,
                       "RESULT": result, "METHOD": method, "COMMIT": head, "EVALUATED_AT_UTC": stamp})

    worst = max(abs(sum(p["RAW_LENGTH_M"] * p["RAW_WIDTH_M"] for p in c["PARTS"]) - c["PARTS_RAW_SUM_M2"])
                for c in comps)
    add("QA-01_COMPONENT_AREA_RECOMPUTED_FROM_RAW_DIMENSIONS", worst <= AREA_TOL_M2,
        {"COMPONENTS": len(comps), "RECTANGLES": sum(len(c["PARTS"]) for c in comps)}, AREA_TOL_M2,
        {"WORST_ABSOLUTE_DIFFERENCE_M2": worst},
        "each rectangle's raw length x raw width re-multiplied and compared with the published raw area")

    frozen_area = {r["ROOM_REF"]: r["AREA_M2"] for f in m["frozen"]["FLOORS"] for r in f["ROOMS"]}
    diffs = {c["COMPONENT_REF"]: abs(cell(ref["COMPONENT"][c["COMPONENT_REF"]]) - frozen_area[c["COMPONENT_REF"]])
             for c in comps}
    add("QA-02_WORKBOOK_COMPONENT_TOTALS_EQUAL_THE_FROZEN_AREAS", max(diffs.values()) < 5e-5,
        {"COMPONENTS_COMPARED": len(diffs)}, 5e-5,
        {"WORST_DIFFERENCE_M2": max(diffs.values()),
         "WORST_COMPONENT": max(diffs, key=diffs.get)},
        "the workbook's own =SUM() over its rectangle rows, evaluated, against the frozen component area")

    fl = {f["FLOOR"]: (cell(ref["FLOOR"][f["FLOOR"]]),
                       sum(cell(ref["COMPONENT"][c["COMPONENT_REF"]]) for c in f["COMPONENTS"]))
          for f in m["floors"]}
    add("QA-03_FLOOR_TOTALS_EQUAL_THE_SUM_OF_THEIR_COMPONENTS",
        all(abs(a - b) < 1e-6 for a, b in fl.values()), {k: round(v[0], 4) for k, v in fl.items()}, 1e-6,
        {k: {"FLOOR_CELL": round(v[0], 6), "SUM_OF_COMPONENTS": round(v[1], 6)} for k, v in fl.items()},
        "evaluated floor-total formula against the evaluated component totals it references")

    proj = cell(ref["PROJECT"])
    add("QA-04_PROJECT_TOTAL_EQUALS_THE_SUM_OF_FLOORS",
        abs(proj - sum(v[0] for v in fl.values())) < 1e-6, {"FLOORS": len(fl)}, 1e-6,
        {"PROJECT_TOTAL_M2": round(proj, 4), "SUM_OF_FLOOR_CELLS": round(sum(v[0] for v in fl.values()), 4)},
        "evaluated project-total formula against the evaluated floor totals")

    refs = [c["COMPONENT_REF"] for c in comps]
    add("QA-05_COMPONENT_REFERENCES_ARE_UNIQUE", len(refs) == len(set(refs)), {"COMPONENTS": len(refs)}, 0,
        {"DISTINCT": len(set(refs))}, "set comparison over every component reference in the workbook")

    frozen_refs = set(frozen_area)
    add("QA-06_SAME_COMPONENT_SET_AS_THE_FROZEN_TAKEOFF", set(refs) == frozen_refs,
        {"WORKBOOK": len(refs), "FROZEN": len(frozen_refs)}, 0,
        {"ONLY_HERE": sorted(set(refs) - frozen_refs), "ONLY_FROZEN": sorted(frozen_refs - set(refs))},
        "set difference in both directions")

    cen = m["census"]
    add("QA-07_COMPONENT_CENSUS_RECONCILES",
        cen["CONNECTED_COMPONENT_COUNT"] == cen["NON_SLIVER_COMPONENT_COUNT"]
        + cen["WALL_MATERIAL_OR_SLIVER_COUNT"] and all(cen["MATCHES_EXPECTED_ROLE_COUNTS"].values()),
        {"EXPECTED_ROLE_COUNTS": EXPECTED_ROLE_COUNTS}, 0,
        {"CONNECTED_COMPONENT_COUNT": cen["CONNECTED_COMPONENT_COUNT"],
         "NON_SLIVER_COMPONENT_COUNT": cen["NON_SLIVER_COMPONENT_COUNT"],
         "WALL_MATERIAL_OR_SLIVER_COUNT": cen["WALL_MATERIAL_OR_SLIVER_COUNT"],
         "COUNT_BY_ROLE": cen["COUNT_BY_ROLE"]},
        "counted from the component model and checked against the audited role counts")

    json_fl = round(sum(r["MEASURED_QUANTITY"] for r in export_rows if r["TRADE"] == "PORCELAIN_FLOOR"), 3)
    wb_fl = cell(ref["FLOORS_BASE"])
    boq_fl = cell(boq_cells["AR-FL-01"])
    froz_fl = t["INTERNAL_FLOOR_AREA_M2"]
    add("QA-08_AR_FL_01_AGREES_ACROSS_JSON_WORKBOOK_BOQ_AND_FROZEN",
        max(abs(json_fl - wb_fl), abs(wb_fl - boq_fl), abs(boq_fl - froz_fl)) < CROSS_TOL_M2,
        {"JSON_SUM_OF_PORCELAIN_FLOOR_RECORDS": json_fl, "WORKBOOK_FLOORS_SHEET_TOTAL": round(wb_fl, 3),
         "BOQ_LINE_AR_FL_01": round(boq_fl, 3), "FROZEN_INTERNAL_FLOOR_AREA_M2": froz_fl}, CROSS_TOL_M2,
        {"WORST_DIFFERENCE_M2": round(max(abs(json_fl - wb_fl), abs(wb_fl - boq_fl),
                                          abs(boq_fl - froz_fl)), 6)},
        "the JSON records summed, and the workbook and BOQ formulas evaluated, all against the frozen total")

    pending_json = round(sum(r["MEASURED_QUANTITY"] for r in export_rows
                             if r["TRADE"] == "FLOOR_FINISH_UNCLASSIFIED"), 3)
    bad = [r["COMPONENT_REF"] for r in export_rows
           if r["TRADE"] == "PORCELAIN_FLOOR" and r["COMPONENT_ROLE"] in PENDING_FLOOR_ROLES]
    add("QA-09_PENDING_FLOOR_AREA_IS_KEPT_OUT_OF_PORCELAIN",
        not bad and abs(pending_json - cell(ref["FLOORS_PENDING"])) < CROSS_TOL_M2,
        {"PENDING_ROLES": list(PENDING_FLOOR_ROLES),
         "JSON_PENDING_M2": pending_json, "WORKBOOK_PENDING_M2": round(cell(ref["FLOORS_PENDING"]), 3),
         "BY_ROLE": m["floor_finish"]["PENDING_BY_ROLE"]}, CROSS_TOL_M2,
        {"PORCELAIN_RECORDS_WITH_A_PENDING_ROLE": bad}, "role filter over the export plus the evaluated sheet total")

    neg = [r["ITEM_CODE"] for r in export_rows
           if isinstance(r["MEASURED_QUANTITY"], (int, float)) and r["MEASURED_QUANTITY"] <= 0]
    add("QA-10_NO_QUANTITY_IS_ZERO_OR_NEGATIVE", not neg, {"RECORDS": len(export_rows)}, 0,
        {"OFFENDING_ITEM_CODES": neg}, "scan of every measured quantity in the export")

    ded_ok = all(x["OPENING_DEDUCTIONS_M2"] <= x["GROSS_WALL_PORCELAIN_M2"] for x in m["assembly"]["wet"]) \
        and all(b["DEDUCTION_GUARD_OK"] for b in m["assembly"]["block"]) \
        and all(x["DOOR_DEDUCTION_M"] < x["PERIMETER_M"] for x in m["skirting"])
    add("QA-11_NO_DEDUCTION_EXCEEDS_WHAT_IT_IS_DEDUCTED_FROM", ded_ok,
        {"WET_ROOMS": len(m["assembly"]["wet"]), "BLOCKWORK_ROWS": len(m["assembly"]["block"]),
         "SKIRTING_ROWS": len(m["skirting"])}, 0, {"ALL_WITHIN_THEIR_GROSS": ded_ok},
        "each deduction compared with its own gross")

    boq = wb["BOQ حسب البند"]
    literal = [(r, c) for r in range(2, boq.max_row + 1) for c in (7, 9)
               if boq.cell(r, c).value is not None and not str(boq.cell(r, c).value).startswith("=")]
    json_ds01 = all(r["WASTE_PERCENT"] is None and r["PROCUREMENT_QUANTITY"] is None
                    and r["UNIT_RATE"] is None and r["AMOUNT"] is None for r in export_rows)
    add("QA-12_DS01_PRICING_FIELDS_ARE_SEPARATE_AND_EMPTY", not literal and json_ds01,
        {"WORKBOOK_COLUMNS_CHECKED": ["الهالك %", "سعر الوحدة"],
         "JSON_FIELDS_CHECKED": ["WASTE_PERCENT", "PROCUREMENT_QUANTITY", "UNIT_RATE", "AMOUNT"]}, 0,
        {"WORKBOOK_CELLS_WITH_A_LITERAL_VALUE": literal, "JSON_ALL_NULL": json_ds01},
        "cell scan of the BOQ sheet and a field scan of every export record")

    add("QA-13_EVERY_RECORD_LEAVES_AS_DRAFT",
        all(r["APPROVAL_STATUS"] == "DRAFT" for r in export_rows), {"RECORDS": len(export_rows)}, 0,
        {"APPROVAL_STATUS_VALUES": sorted({r["APPROVAL_STATUS"] for r in export_rows})},
        "value set over the export")

    hist = {950.22, 57.0, 1375.0, 600.16, 0.75, 1.82, 1.715}
    leaked = sorted({r["MEASURED_QUANTITY"] for r in export_rows if r["MEASURED_QUANTITY"] in hist})
    add("QA-14_NO_HISTORICAL_FIGURE_ENTERED_A_QUANTITY_FIELD", not leaked,
        {"HISTORICAL_FIGURES_CHECKED": sorted(hist)}, 0,
        {"FOUND_IN_A_QUANTITY_FIELD": leaked,
         "NOTE": "950.22 appears only as labelled text on the insulation sheet, never as a quantity"},
        "membership test of every measured quantity against the historical figures")

    wet_refs = {x["ROOM_REF"] for x in m["assembly"]["wet"]}
    all_wet = {c["COMPONENT_REF"] for c in comps if c["ROLE"] in WET_ROLES}
    dry_us18 = [r["ITEM_CODE"] for r in export_rows
                if r["RULE_ID"] == "US-18" and r["COMPONENT_ROLE"] not in (None,) + WET_ROLES]
    add("QA-15_US18_IS_CLAIMED_ONLY_WHERE_IT_APPLIES", wet_refs == all_wet and not dry_us18,
        {"WET_AND_KITCHEN_COMPONENTS": len(all_wet), "IN_THE_CLADDING_SHEET": len(wet_refs)}, 0,
        {"DRY_RECORDS_CLAIMING_US18": dry_us18},
        "US-18's own room-type scope checked against every record that cites it")

    ceil = [r for r in export_rows if r["TRADE"] == "CEILING"]
    add("QA-16_NO_UNRESOLVED_CEILING_FINISH_IS_FINAL",
        all(r["STATUS"] == "FINISH_CLASSIFICATION_PENDING" and r["BOQ_INCLUDED"] is False for r in ceil),
        {"CEILING_RECORDS": len(ceil), "CEILING_AREA_M2": t["CEILING_M2"]}, 0,
        {"STATUSES": sorted({r["STATUS"] for r in ceil}),
         "AREA_STATUS": "FINAL_QUANTITY_AVAILABLE - the measurement is final, the finish is not"},
        "status scan over every ceiling record")

    unsupported = [w["WINDOW_ID"] for w in m["windows"] if not w["HEIGHT_AUTHORITY"]["SUPPORTED"]]
    still_final = [w["WINDOW_ID"] for w in m["windows"]
                   if not w["HEIGHT_AUTHORITY"]["SUPPORTED"] and w["BOQ_INCLUDED"]]
    add("QA-17_NO_GUIDE_CATEGORY_IS_USED_OUTSIDE_ITS_BAND_WITHOUT_AUTHORITY", not still_final,
        {"WINDOWS": len(m["windows"]),
         "AUTHORITIES": {w["WINDOW_ID"]: w["HEIGHT_AUTHORITY"]["AUTHORITY"] for w in m["windows"]}}, 0,
        {"WITHOUT_AUTHORITY": unsupported, "STILL_TREATED_AS_FINAL": still_final},
        "each window's category basis tested against that category's own area band")

    sk = m["skirting_buckets"]
    add("QA-18_SKIRTING_IS_DERIVED_SPLIT_AND_OUT_OF_THE_BASE_BOQ",
        all(r["STATUS"] == "DERIVED_NOT_IN_FROZEN_TAKEOFF" and r["BOQ_INCLUDED"] is False
            for r in m["skirting"])
        and abs(sum(v["LENGTH_M"] for k, v in sk.items() if isinstance(v, dict)) - sk["TOTAL_LENGTH_M"]) < 1e-6,
        {"BUCKETS": {k: v for k, v in sk.items() if isinstance(v, dict)}}, 1e-6,
        {"TOTAL_LENGTH_M": sk["TOTAL_LENGTH_M"], "IS_ONE_DECISION_READY_QUANTITY": False},
        "bucket sums against the published total, and a status scan of every skirting row")

    am_counts = VA.HISTORICAL_COUNTS
    ok_counts = (am_counts["SCHEDULE_ROW_COUNT"]["TOTAL"] == 31
                 and am_counts["PHYSICAL_OBJECT_COUNT"]["TOTAL"] == 34
                 and am_counts["PHYSICAL_OBJECT_COUNT"]["WINDOWS"] == 26)
    add("QA-19_SCHEDULE_ROWS_AND_PHYSICAL_OBJECTS_ARE_SEPARATE_POPULATIONS", ok_counts,
        {"ROW_COUNT": am_counts["SCHEDULE_ROW_COUNT"], "PHYSICAL_OBJECT_COUNT": am_counts["PHYSICAL_OBJECT_COUNT"]},
        0, {"ROWS_TO_OBJECTS": am_counts["RECONCILIATION"]},
        "the amendment's ROW_COUNT, MULTIPLICITY and PHYSICAL_OBJECT_COUNT fields re-added independently")

    add("QA-20_THE_FROZEN_TAKEOFF_IS_BYTE_FOR_BYTE_UNCHANGED",
        frozen_file_sha256() == FROZEN_FILE_SHA256 and m["frozen"]["DIGEST"] == FROZEN_DIGEST
        and m["frozen"]["GIT_HEAD"] == FROZEN_COMMIT,
        {"EXPECTED_SHA256": FROZEN_FILE_SHA256, "EXPECTED_DIGEST": FROZEN_DIGEST,
         "EXPECTED_COMMIT": FROZEN_COMMIT}, 0,
        {"SHA256": frozen_file_sha256(), "DIGEST": m["frozen"]["DIGEST"], "GIT_HEAD": m["frozen"]["GIT_HEAD"]},
        "SHA-256 of the frozen file on disk, plus its own internal digest and commit fields")

    add("QA-21_EVERY_FORMULA_EVALUATES_WITHOUT_ERROR", len(vals) > 0,
        {"FORMULA_CELLS": len(vals)}, 0,
        {"EVALUATED": len(vals), "ERRORS": 0,
         "BLANK_BY_DESIGN": sum(1 for v in vals.values() if v is WC.BLANK)},
        "every formula in the book evaluated by the workbook evaluator; any unsupported formula raises")

    json_rules = {r["COMPONENT_REF"]: r["RULE_ID"] for r in export_rows
                  if r["TRADE"] in ("PORCELAIN_FLOOR", "FLOOR_FINISH_UNCLASSIFIED")}
    pairs, mismatch = [], []
    fl_sheet = wb["الأرضيات"]
    for r in range(2, fl_sheet.max_row + 1):
        refc, rule = fl_sheet.cell(r, 2).value, fl_sheet.cell(r, 9).value
        if not refc or refc not in json_rules or not rule:
            continue
        sheet_rule = None if rule == "—" else rule
        pairs.append((refc, sheet_rule))
        if json_rules[refc] != sheet_rule:
            mismatch.append({"COMPONENT_REF": refc, "WORKBOOK": sheet_rule, "JSON": json_rules[refc]})
    add("QA-22_RULE_IDS_ARE_IDENTICAL_IN_THE_WORKBOOK_AND_THE_JSON", not mismatch,
        {"FLOOR_ROWS_COMPARED": len(pairs)}, 0, {"MISMATCHES": mismatch},
        "component-by-component comparison of RULE_ID between the floors sheet and the export")

    return {"CHECKS": checks, "ALL_PASS": all(c["PASS"] for c in checks),
            "PASSED": sum(1 for c in checks if c["PASS"]), "OF": len(checks),
            "EVALUATED_AT_UTC": stamp, "COMMIT": head,
            "METHOD": "every check is computed from evidence at generation time; none is a stored PASS"}


# ------------------------------------------------------------------ the structured export
PROJECT_NAME = "Ahmad Abdullah Ali Al Rashed - Sabah Al Ahmad, Block D4, Plot 247"
DRAWING_REVISION = "16-11-2025, sheet row R3 - the row the issued PDF plots"
SOURCE_FILES = ["16-11-2025.dwg", "16-11-2025.pdf"]


def _rec(**kw):
    """One export record.  Quantity, unit, waste, procurement, rate and amount are six separate fields (DS-01)."""
    base = {
        "PROJECT_ID": OI.PROJECT_ID, "PROJECT_NAME": PROJECT_NAME,
        "DRAWING_SET": "ARCHITECTURAL_PERMIT_SET", "DRAWING_REVISION": DRAWING_REVISION,
        "SOURCE_FILES": SOURCE_FILES,
        "FLOOR": None, "FLOOR_AR": None, "LEVEL_M": None, "ZONE": None,
        "ENTITY_TYPE": None, "COMPONENT_REF": None, "COMPONENT_ROLE": None, "DRAWING_LABEL": None,
        "TRADE": None, "ITEM_CODE": None, "ITEM": None, "SUBITEM": None,
        "COMPONENTS": None, "CALCULATION": None, "MEASUREMENT_BASIS": None, "DIMENSION_INPUTS": None,
        "MEASURED_QUANTITY": None, "MEASURED_UNIT": None,
        "WASTE_PERCENT": None, "PROCUREMENT_QUANTITY": None,
        "UNIT_RATE": None, "AMOUNT": None, "CURRENCY": "KWD",
        "QUANTITY_SOURCE": None, "SOURCE_TYPE": "PROJECT_DRAWING",
        "RULE_ID": None, "RULE_LEVEL": None, "AUTHORITY": None, "SCHEDULE_BUCKET": None,
        "STATUS": "FINAL_QUANTITY_AVAILABLE", "FINISH_STATUS": None, "BOQ_INCLUDED": True,
        "APPROVAL_STATUS": "DRAFT",
        "VERIFIED_AGAINST_FROZEN": True, "FROZEN_DIGEST": FROZEN_DIGEST,
        "HISTORICAL_DATA_USED": False,
        "GENERATED_BY": f"URBAN_QS_ENGINE / ALRASHED_DETAILED_QUANTITY_TAKEOFF rev {REVISION}",
        "NOTES": None,
    }
    base.update(kw)
    return base


def export_rows(m):
    t, rows = m["assembly"]["totals"], []
    for f in m["floors"]:
        for c in f["COMPONENTS"]:
            comp = [{"PART": p["PART"], "RAW_LENGTH_M": p["RAW_LENGTH_M"], "RAW_WIDTH_M": p["RAW_WIDTH_M"],
                     "RAW_AREA_M2": p["RAW_AREA_M2"], "DISPLAY_LENGTH_M": p["DISPLAY_LENGTH_M"],
                     "DISPLAY_WIDTH_M": p["DISPLAY_WIDTH_M"], "DISPLAY_AREA_M2": p["DISPLAY_AREA_M2"]}
                    for p in c["PARTS"]]
            common = dict(FLOOR=c["FLOOR"], FLOOR_AR=c["FLOOR_AR"], LEVEL_M=f["LEVEL_M"],
                          ENTITY_TYPE="CONNECTED_COMPONENT", COMPONENT_REF=c["COMPONENT_REF"],
                          COMPONENT_ROLE=c["ROLE"], DRAWING_LABEL=c["NAME"],
                          COMPONENTS=comp,
                          CALCULATION=" + ".join(p["CALCULATION"] for p in c["PARTS"]),
                          MEASUREMENT_BASIS=c["MEASUREMENT_BASIS"],
                          MEASURED_QUANTITY=c["AREA_M2"], MEASURED_UNIT="m2",
                          QUANTITY_SOURCE="exact rectilinear decomposition of the drawn wall faces")
            if c["ROLE"] in PORCELAIN_FLOOR_ROLES:
                wet = c["ROLE"] in WET_ROLES
                rows.append(_rec(ZONE="WET_ROOM" if wet else "INTERNAL_SPACE", TRADE="PORCELAIN_FLOOR",
                                 ITEM_CODE="AR-FL-01", ITEM="Floor finish", SUBITEM=c["COMPONENT_REF"],
                                 RULE_ID="US-18" if wet else None,
                                 RULE_LEVEL="URBAN_STANDARD" if wet else None,
                                 STATUS="FINAL_QUANTITY_AVAILABLE",
                                 NOTES="US-18 covers the wet and service rooms only; a dry room's floor area is "
                                       "a measured quantity under no Urban standard",
                                 **common))
            elif c["ROLE"] in PENDING_FLOOR_ROLES:
                rows.append(_rec(ZONE="UNRESOLVED_SPACE", TRADE="FLOOR_FINISH_UNCLASSIFIED",
                                 ITEM_CODE="AR-FL-PENDING", ITEM="Floor finish, type not established",
                                 SUBITEM=c["COMPONENT_REF"],
                                 STATUS="FINISH_CLASSIFICATION_PENDING",
                                 FINISH_STATUS="FINISH_CLASSIFICATION_PENDING", BOQ_INCLUDED=False,
                                 NOTES="area is measured and kept; a stair takes treads rather than floor "
                                       "finish, and an unnamed space needs the owner's classification",
                                 **common))
            elif c["ROLE"] == "EXTERNAL_OR_OPEN":
                rows.append(_rec(ZONE="EXTERNAL", TRADE="EXTERNAL_AREAS", ITEM_CODE="AR-EX-01",
                                 ITEM="External / parking area", SUBITEM=c["COMPONENT_REF"],
                                 RULE_ID="AR-08", RULE_LEVEL="OWNER_CONFIRMED_PROJECT_INPUT",
                                 STATUS="FINISH_CLASSIFICATION_PENDING",
                                 FINISH_STATUS="FINISH_CLASSIFICATION_PENDING", BOQ_INCLUDED=False,
                                 NOTES="area is final; the drawing states no finish material", **common))
            if c["ROLE"] in ENCLOSED_ROLES:
                rows.append(_rec(ZONE="CEILING", TRADE="CEILING", ITEM_CODE="AR-CL-01",
                                 ITEM="Ceiling plan area", SUBITEM=c["COMPONENT_REF"],
                                 STATUS="FINISH_CLASSIFICATION_PENDING",
                                 FINISH_STATUS="FINISH_CLASSIFICATION_PENDING", BOQ_INCLUDED=False,
                                 NOTES="the enclosed area is measured and final; the ceiling TYPE is not "
                                       "established - no ceiling plan was supplied and the owner has not "
                                       "confirmed a plain ceiling",
                                 **common))
    for x in m["assembly"]["wet"]:
        for trade, code, item in (("WALL_PORCELAIN", "AR-WP-01", "Wall porcelain, full height"),
                                  ("TILE_PREPARATION", "AR-WP-02", "Tartousha behind the porcelain")):
            rows.append(_rec(FLOOR=x["FLOOR"], FLOOR_AR=FLOORS_AR[x["FLOOR"]], ZONE="WET_ROOM",
                             ENTITY_TYPE="CONNECTED_COMPONENT", COMPONENT_REF=x["ROOM_REF"],
                             COMPONENT_ROLE="WET_ROOM", DRAWING_LABEL=x["ROOM"],
                             TRADE=trade, ITEM_CODE=code, ITEM=item, SUBITEM=x["ROOM_REF"],
                             DIMENSION_INPUTS={"PERIMETER_M": x["CERAMIC_HOST_LENGTH_M"],
                                               "HEIGHT_M": x["HEIGHT_M"],
                                               "HEIGHT_SOURCE": "OWNER_CONFIRMED_PROJECT_INPUT (AR-01)",
                                               "OPENINGS": x["OPENINGS_ON_THIS_ROOM"]},
                             CALCULATION=f"{x['CERAMIC_HOST_LENGTH_M']} × {x['HEIGHT_M']} − "
                                         f"{x['OPENING_DEDUCTIONS_M2']} = {x['NET_WALL_PORCELAIN_M2']}",
                             MEASUREMENT_BASIS="room perimeter less that room's own openings",
                             MEASURED_QUANTITY=x["NET_WALL_PORCELAIN_M2"], MEASURED_UNIT="m2",
                             QUANTITY_SOURCE="measured room perimeter and its own openings",
                             RULE_ID="US-18", RULE_LEVEL="URBAN_STANDARD", NOTES=x["REASON"]))
    for x in m["skirting"]:
        rows.append(_rec(FLOOR=x["FLOOR"], FLOOR_AR=FLOORS_AR[x["FLOOR"]], ZONE="SKIRTING",
                         ENTITY_TYPE="CONNECTED_COMPONENT", COMPONENT_REF=x["COMPONENT_REF"],
                         COMPONENT_ROLE=x["ROLE"], DRAWING_LABEL=x["ROOM"],
                         TRADE="SKIRTING",
                         ITEM_CODE={"DRY_NAMED_INTERNAL_CANDIDATE": "AR-SK-01",
                                    "STAIR_OR_LANDING_PENDING": "AR-SK-02",
                                    "UNNAMED_SPACE_PENDING": "AR-SK-03"}[x["BUCKET"]],
                         ITEM="Skirting", SUBITEM=x["COMPONENT_REF"], SCHEDULE_BUCKET=x["BUCKET"],
                         DIMENSION_INPUTS={"PERIMETER_M": x["PERIMETER_M"],
                                           "DOOR_DEDUCTION_M": x["DOOR_DEDUCTION_M"]},
                         CALCULATION=f"{x['PERIMETER_M']} − {x['DOOR_DEDUCTION_M']} = {x['SKIRTING_LENGTH_M']}",
                         MEASUREMENT_BASIS=x["RULE"],
                         MEASURED_QUANTITY=x["SKIRTING_LENGTH_M"], MEASURED_UNIT="lm",
                         QUANTITY_SOURCE="frozen component perimeter and the doors on it",
                         STATUS=x["STATUS"], BOQ_INCLUDED=False, VERIFIED_AGAINST_FROZEN=False,
                         NOTES=x["NOTES"]))
    for b in m["assembly"]["block"]:
        rows.append(_rec(FLOOR=b["FLOOR"], FLOOR_AR=FLOORS_AR[b["FLOOR"]], ZONE="WALLS",
                         ENTITY_TYPE="WALL_BAND",
                         TRADE=f"BLOCKWORK_{b['THICKNESS_MM']}",
                         ITEM_CODE="AR-BL-01" if b["THICKNESS_MM"] == 150 else "AR-BL-02",
                         ITEM=f"Blockwork {b['THICKNESS_MM']} mm",
                         SUBITEM=f"{b['FLOOR']}-{b['THICKNESS_MM']}",
                         DIMENSION_INPUTS={"PLAN_LENGTH_M": b["PLAN_LENGTH_M"], "HEIGHT_M": b["HEIGHT_M"]},
                         CALCULATION=f"{b['PLAN_LENGTH_M']} × {b['HEIGHT_M']} − {b['OPENING_DEDUCTIONS_M2']}"
                                     f" = {b['NET_AREA_M2']}",
                         MEASUREMENT_BASIS="wall cells of a thickness the drawing dimensions",
                         MEASURED_QUANTITY=b["NET_AREA_M2"], MEASURED_UNIT="m2",
                         QUANTITY_SOURCE="measured wall lengths less the openings in them",
                         NOTES="junction artefacts are excluded; only 150 and 200 mm are billed"))
    for p in m["assembly"]["pp"]:
        for trade, code, item, qty in (
                ("INTERNAL_PLASTER", "AR-PL-01", "Internal plaster", p["INTERNAL_PLASTER_NET_M2"]),
                ("INTERNAL_PAINT", "AR-PT-01", "Internal paint", p["INTERNAL_PAINT_M2"]),
                ("EXTERNAL_PLASTER", "AR-PL-02", "External plaster", p["EXTERNAL_PLASTER_NET_M2"]),
                ("EXTERNAL_PAINT", "AR-PT-02", "External paint", p["EXTERNAL_PAINT_M2"])):
            rows.append(_rec(FLOOR=p["FLOOR"], FLOOR_AR=FLOORS_AR[p["FLOOR"]], ZONE="WALL_FACES",
                             ENTITY_TYPE="WALL_FACE_SET",
                             TRADE=trade, ITEM_CODE=code, ITEM=item, SUBITEM=p["FLOOR"],
                             DIMENSION_INPUTS={"INTERNAL_FACE_LENGTH_M": p["INTERNAL_FACE_LENGTH_M"],
                                               "EXTERNAL_FACE_LENGTH_M": p["EXTERNAL_FACE_LENGTH_M"],
                                               "HEIGHT_M": p["HEIGHT_M"], "HEIGHT_SOURCE": p["HEIGHT_SOURCE"]},
                             MEASUREMENT_BASIS="wall-face lengths × the owner's wall height, less openings",
                             MEASURED_QUANTITY=qty, MEASURED_UNIT="m2",
                             QUANTITY_SOURCE="measured wall faces and measured openings",
                             RULE_ID="US-18" if trade == "INTERNAL_PAINT" else None,
                             RULE_LEVEL="URBAN_STANDARD" if trade == "INTERNAL_PAINT" else None,
                             NOTES="a fully tiled face takes no paint" if trade == "INTERNAL_PAINT" else None))
    rp = m["assembly"]["rp"]
    for code, trade, item, qty, calc in (
            ("AR-IN-01", "WATERPROOFING", "Roof waterproofing", rp["WATERPROOFING_AREA_M2"],
             "the roof area enclosed by the parapet"),
            ("AR-RF-01", "PARAPET_BLOCKWORK", "Parapet blockwork", rp["PARAPET_BLOCKWORK_M2"],
             f"{rp['PARAPET_RUN_M']} × {rp['PARAPET_HEIGHT_M']}"),
            ("AR-RF-02", "PARAPET_PLASTER", "Parapet plaster, both faces", rp["PARAPET_PLASTER_M2"],
             f"{rp['PARAPET_RUN_M']} × {rp['PARAPET_HEIGHT_M']} × 2"),
            ("AR-RF-03", "PARAPET_PAINT", "Parapet paint, both faces", rp["PARAPET_PAINT_M2"],
             f"{rp['PARAPET_RUN_M']} × {rp['PARAPET_HEIGHT_M']} × 2")):
        rows.append(_rec(FLOOR="ROOF", FLOOR_AR=FLOORS_AR["ROOF"], ZONE="ROOF", ENTITY_TYPE="ROOF_ELEMENT",
                         TRADE=trade, ITEM_CODE=code, ITEM=item, CALCULATION=calc,
                         DIMENSION_INPUTS={"PARAPET_RUN_M": rp["PARAPET_RUN_M"],
                                           "PARAPET_HEIGHT_M": rp["PARAPET_HEIGHT_M"],
                                           "HEIGHT_SOURCE": rp["PARAPET_HEIGHT_SOURCE"]},
                         MEASUREMENT_BASIS=rp["PARAPET_RUN_BASIS"],
                         MEASURED_QUANTITY=qty, MEASURED_UNIT="m2",
                         QUANTITY_SOURCE="the roof edge the drawing draws",
                         RULE_ID="AR-06", RULE_LEVEL="OWNER_CONFIRMED_PROJECT_INPUT"))
    for d in m["assembly"]["doors"]:
        rows.append(_rec(FLOOR=d["FLOOR"], FLOOR_AR=FLOORS_AR[d["FLOOR"]], ZONE="OPENING",
                         ENTITY_TYPE="OPENING",
                         TRADE="DOORS", ITEM_CODE="AR-DR-01", ITEM="Door opening", SUBITEM=d["DOOR_ID"],
                         DIMENSION_INPUTS={"WIDTH_M": d["WIDTH_M"], "WIDTH_SOURCE": d["WIDTH_SOURCE"],
                                           "HEIGHT_M": d["HEIGHT_M"], "HEIGHT_SOURCE": d["HEIGHT_SOURCE"]},
                         CALCULATION=f"{d['WIDTH_M']} × {d['HEIGHT_M']} = {d['AREA_M2']}",
                         MEASUREMENT_BASIS="clear width jamb to jamb",
                         MEASURED_QUANTITY=d["AREA_M2"], MEASURED_UNIT="m2",
                         QUANTITY_SOURCE="measured from the project geometry",
                         RULE_ID="AR-02", RULE_LEVEL="OWNER_CONFIRMED_PROJECT_INPUT",
                         STATUS="PRICING_BASIS_REQUIRED", NOTES=d["NOTES"]))
    for w in m["windows"]:
        auth = w["HEIGHT_AUTHORITY"]
        rows.append(_rec(FLOOR=w["FLOOR"], FLOOR_AR=FLOORS_AR[w["FLOOR"]], ZONE="EXTERNAL_WALL",
                         ENTITY_TYPE="OPENING", COMPONENT_REF=w["ROOM_REF"], DRAWING_LABEL=w["ROOM"],
                         TRADE="ALUMINIUM",
                         ITEM_CODE="AR-AL-01" if w["BOQ_INCLUDED"] else "AR-AL-PENDING",
                         ITEM="Window", SUBITEM=w["WINDOW_ID"],
                         DIMENSION_INPUTS={"WIDTH_M": w["WIDTH_M"], "WIDTH_SOURCE": w["WIDTH_SOURCE"],
                                           "HEIGHT_M": w["HEIGHT_M"], "HEIGHT_SOURCE": w["HEIGHT_SOURCE"],
                                           "GUIDE_CATEGORY": w["GUIDE_CATEGORY"],
                                           "CATEGORY_BASIS": w["CATEGORY_BASIS"],
                                           "COMPONENT_AREA_M2": w["ROOM_AREA_M2"]},
                         CALCULATION=f"{w['WIDTH_M']} × {w['HEIGHT_M']} = {w['AREA_M2']}",
                         MEASUREMENT_BASIS="width measured from the plot, height from the Urban guide",
                         MEASURED_QUANTITY=w["AREA_M2"], MEASURED_UNIT="m2",
                         QUANTITY_SOURCE=f"width {w['WIDTH_SOURCE']}, height {w['HEIGHT_SOURCE']}",
                         RULE_ID=WS.RULE_ID, RULE_LEVEL="URBAN_STANDARD",
                         AUTHORITY=auth, STATUS=w["STATUS"], BOQ_INCLUDED=w["BOQ_INCLUDED"],
                         NOTES=("US-21: an aluminium quantity from architectural drawings is not a supply "
                                "quantity") if w["BOQ_INCLUDED"] else auth["WHY"]))
    return rows


# ------------------------------------------------------------------ sources, review and the summary
def sources_sheet(ws, m):
    ws.append(["القسم", "البند", "التفصيل", "الحالة"])

    def sec(name, item, detail, status):
        ws.append([name, item, detail, status])

    for f in SOURCE_FILES:
        sec("المصادر", f, "ملف مشروع مقدَّم من المالك — القياس كله منه", "USED")
    sec("المصادر", "مراجعة المخطط", DRAWING_REVISION, "IDENTIFIED_BY_DIMENSION_CONTAINMENT_SCORE_1.000")
    sec("المصادر", "الملف التاريخي (إكسل)", "لم تُؤخذ منه أي كمية أو سعر أو معادلة في هذا الملف",
        "NOT_A_SOURCE_OF_QUANTITY")
    sec("المصادر", "دليل مقاسات الشبابيك", WS.VERSION + " — ارتفاعات فقط، والعرض من المخطط دائماً",
        "APPLIED_FOR_HEIGHT_ONLY")
    ws.append([])
    cen = m["census"]
    sec("تعداد المكوّنات", "مكوّنات متصلة", str(cen["CONNECTED_COMPONENT_COUNT"]),
        "CONNECTED_COMPONENT_COUNT")
    sec("تعداد المكوّنات", "مكوّنات غير شظايا", str(cen["NON_SLIVER_COMPONENT_COUNT"]),
        "NON_SLIVER_COMPONENT_COUNT")
    sec("تعداد المكوّنات", "مادة جدار / شظية", str(cen["WALL_MATERIAL_OR_SLIVER_COUNT"]),
        "WALL_MATERIAL_OR_SLIVER_COUNT")
    for role, n in cen["COUNT_BY_ROLE"].items():
        sec("تعداد المكوّنات", ROLES_AR.get(role, role), f"{n} مكوّن — {cen['AREA_BY_ROLE_M2'].get(role, 0)} م²",
            role)
    sec("تعداد المكوّنات", "ملاحظة", "المكوّن المتصل ليس غرفة حتى يسميه المخطط أو تثبت هندسته ما هو",
        "ENTITY_SEMANTICS")
    ws.append([])
    for k, v in OI.PROJECT_INPUTS.items():
        sec("مدخلات المالك", f"{v['FROM']} — {k}", f"{v['VALUE_M']} م — {v['KIND']}",
            "TRAVELS" if v.get("TRAVELS") else "PROJECT_VALUE_ONLY")
    sec("مدخلات المالك", "AR-07 — النطاق",
        "الإنشائي والصحي والكهربائي والتكييف خارج النطاق بقرار المالك", "OUT_OF_SCOPE_FOR_CURRENT_VALIDATION")
    ws.append([])
    sec("القواعد", "US-18", OI.US_WET_ROOM_FINISH["STATEMENT"], "APPLIED_TO_WET_AND_SERVICE_ROOMS_ONLY")
    sec("القواعد", WS.RULE_ID, "دليل مقاسات الشبابيك — مرتبة أولوية 4، ولا يُستعمل خارج نطاق مساحة الفئة "
                               "إلا بسند موثّق", "APPLIED_WITH_AUTHORITY_CHECK")
    for d in VA.RULE_DECISIONS:
        if d["DECISION"] == "PROMOTE":
            sec("القواعد", d["RULE_ID"], d["STATEMENT"], "PROMOTED_BY_OWNER")
    sec("القواعد", "DS-01 — فصل الكمية عن السعر",
        "الكمية والوحدة والهالك وكمية الشراء وسعر الوحدة والإجمالي حقول منفصلة، لا تجتمع في خانة واحدة",
        "ENFORCED_IN_THIS_WORKBOOK")
    ws.append([])
    for w in m["windows"]:
        a = w["HEIGHT_AUTHORITY"]
        sec("سند تصنيف الشبابيك", f"{w['WINDOW_ID']} — {w['ROOM']} ({w['ROOM_AREA_M2']} م²)",
            f"{w['GUIDE_CATEGORY']} — {a.get('WHY', 'ضمن نطاق مساحة الفئة')}", a["AUTHORITY"])
    ws.append([])
    sec("الحالة المجمدة", "Commit", FROZEN_COMMIT, "UNCHANGED")
    sec("الحالة المجمدة", "Digest", FROZEN_DIGEST, "UNCHANGED")
    sec("الحالة المجمدة", "SHA-256 للملف المجمد", FROZEN_FILE_SHA256, "UNCHANGED")
    sec("الحالة المجمدة", "Artifact", "FULL_VILLA_BLIND_PRE_PRICING_TAKEOFF",
        "هذا الملف يعيد عرض القياس المجمد ولا يعيد حسابه")
    sec("الحالة المجمدة", "فروق التقريب",
        "ورقة حصر المساحات تحمل الأبعاد بدقتها الكاملة فتطابق المجمد تماماً؛ أما أوراق البنود فتعيد الجمع من "
        "الأبعاد كما تُنشر، فقد يظهر فرق تقريب أقصاه 0.003 م² في بند مساحته 963 م²", "ROUNDING_ONLY")
    ws.append([])
    sec("ما لم يُقَس", "عزل أرضيات الغرف الرطبة", "لا تفصيل على المخططات الثلاثة", "SOURCE_REQUIRED")
    sec("ما لم يُقَس", "بروفايل / درابزين", PROFILE_STEEL["WHY"], "SOURCE_REQUIRED")
    sec("ما لم يُقَس", "نوع السقف", "لا مخطط أسقف ولا تأكيد من المالك — المساحة نهائية والنوع لا",
        "FINISH_CLASSIFICATION_PENDING")
    sec("ما لم يُقَس", "تشطيب الدرج والفراغات غير المسماة",
        f"{m['floor_finish']['PENDING_M2']} م² مقاسة ومحفوظة خارج بند الأرضيات", "FINISH_CLASSIFICATION_PENDING")
    sec("ما لم يُقَس", "الإنشائي والصحي والكهربائي والتكييف", "خارج النطاق بقرار المالك AR-07",
        "OUT_OF_SCOPE_FOR_CURRENT_VALIDATION")
    ws.append([])
    _style(ws, (22, 40, 86, 36))
    return ws.max_row


def summary_sheet(ws, m, ref, boq_cells):
    from openpyxl.styles import Font
    cen = m["census"]
    ws.append(["ملخص حصر الكميات — فيلا أحمد عبدالله علي الراشد", None, None, None])
    ws.append(["صباح الأحمد — قطعة D4، قسيمة 247، مساحة القسيمة 600.00 م²", None, None, None])
    ws.append([f"المخطط: {DRAWING_REVISION}", None, None, None])
    ws.append([f"الحالة المجمدة: {FROZEN_COMMIT} / {FROZEN_DIGEST} — لم تتغير", None, None, None])
    ws.append([f"مراجعة هذا الملف: {REVISION} — تصحيحات تدقيق خارجي", None, None, None])
    ws.append(["الأسعار والهالك: فارغة عمداً — هذا حصر كميات لا تسعيرة", None, None, None])
    ws.append([])
    ws.append(["تعداد المكوّنات المتصلة", "العدد", "المساحة (م²)", "ملاحظة"])
    _fill_row(ws, ws.max_row, HEAD_FILL, last="D")
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
    ws.append(["مكوّنات متصلة (الإجمالي)", cen["CONNECTED_COMPONENT_COUNT"], None,
               "ليست 135 غرفة — المكوّن المتصل وحدة هندسية لا وحدة استعمال"])
    ws.append(["منها مكوّنات غير شظايا", cen["NON_SLIVER_COMPONENT_COUNT"], None, "فراغات فعلية"])
    ws.append(["منها مادة جدار / شظية", cen["WALL_MATERIAL_OR_SLIVER_COUNT"], None,
               "شرائح جدران لا تُقرأ كفراغات"])
    for role, n in cen["COUNT_BY_ROLE"].items():
        ws.append([f"   {ROLES_AR.get(role, role)}", n, cen["AREA_BY_ROLE_M2"].get(role, 0), role])
    ws.append([])
    ws.append(["الدور", "عدد المكوّنات", "المساحة (م²)", "المنسوب (م)"])
    _fill_row(ws, ws.max_row, HEAD_FILL, last="D")
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
    for f in m["floors"]:
        ws.append([f["FLOOR_AR"], len(f["COMPONENTS"]), None, f["LEVEL_M"]])
        ws.cell(ws.max_row, 3).value = "=" + ref["FLOOR"][f["FLOOR"]]
    ws.append(["إجمالي المشروع (يشمل الشظايا)", cen["CONNECTED_COMPONENT_COUNT"], None, None])
    ws.cell(ws.max_row, 3).value = "=" + ref["PROJECT"]
    _fill_row(ws, ws.max_row, TOTAL_FILL, last="D")
    ws.append([])
    ws.append(["البند", "الوحدة", "الكمية", "الحالة"])
    _fill_row(ws, ws.max_row, HEAD_FILL, last="D")
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
    for sec, code, _trade, desc, unit, cell, value, _source, _rule, status, included, _notes in boq_items(m, ref):
        ws.append([f"{code} — {desc}", unit, None, f"{status} — {'أساس' if included else 'مؤقت'}"])
        ws.cell(ws.max_row, 3).value = ("=" + boq_cells[code]) if code in boq_cells else value
        if not included:
            _fill_row(ws, ws.max_row, PEND_FILL, bold=False, last="D")
    ws.append([])
    ws.sheet_view.rightToLeft = True
    for col, w in zip("ABCD", (62, 16, 18, 52)):
        ws.column_dimensions[col].width = w
    ws["A1"].font = Font(bold=True, size=14)


# ------------------------------------------------------------------ assembly of the file
def workbook(m, path):
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = {name: wb.create_sheet(name) for name in SHEETS}
    ref = area_sheet(ws["حصر المساحات"], m)
    fl = floors_sheet(ws["الأرضيات"], m, ref)
    ref["FLOORS_BASE"], ref["FLOORS_PENDING"] = fl["BASE"], fl["PENDING"]
    ref["CEILING"] = ceiling_sheet(ws["الأسقف"], m, ref)
    ref.update(wall_sheet(ws["كسوة الجدران"], m))
    sk = skirting_sheet(ws["النعلة والبروفايل"], m)
    ref["SK_DRY"] = sk["DRY_NAMED_INTERNAL_CANDIDATE"]
    ref["SK_STAIR"] = sk["STAIR_OR_LANDING_PENDING"]
    ref["SK_UNNAMED"] = sk["UNNAMED_SPACE_PENDING"]
    ref["SK_TOTAL"] = sk["TOTAL"]
    ref["INSULATION"] = insulation_sheet(ws["العازل"], m)
    ref.update(openings_sheet(ws["الأبواب والشبابيك"], m))
    blk = blockwork_sheet(ws["المباني"], m)
    ref["BLOCK150"], ref["BLOCK200"] = blk[150], blk[200]
    ref.update(plaster_sheet(ws["المساح والصبغ"], m))
    ref.update(roof_sheet(ws["السطح والخارجي"], m))
    boq_cells = boq_sheet(ws["BOQ حسب البند"], m, ref)
    summary_sheet(ws["ملخص الحصر"], m, ref, boq_cells)
    sources_sheet(ws["المصادر والمراجعة"], m)

    rows = export_rows(m)
    checks = qa(m, wb, ref, boq_cells, rows)

    src = ws["المصادر والمراجعة"]
    src.append(["فحوص الجودة", "الطريقة", checks["METHOD"], f"{checks['PASSED']}/{checks['OF']}"])
    _fill_row(src, src.max_row, BAND_FILL, last="D")
    for c in checks["CHECKS"]:
        src.append(["فحوص الجودة", c["CHECK"],
                    json.dumps({"INPUTS": c["INPUTS"], "RESULT": c["RESULT"], "TOLERANCE": c["TOLERANCE"]},
                               ensure_ascii=False, default=str)[:400],
                    "PASS" if c["PASS"] else "FAIL"])
    summ = ws["ملخص الحصر"]
    summ.append(["فحوص الجودة", f"{checks['PASSED']} / {checks['OF']}",
                 "PASS" if checks["ALL_PASS"] else "FAIL", "محسوبة من الأدلة وقت التوليد، لا نص ثابت"])
    _fill_row(summ, summ.max_row, TOTAL_FILL, last="D")
    summ.append(["حالة الاعتماد", "DRAFT", str(len(rows)), "المحرك لا يعتمد حصره بنفسه"])
    summ.append(["بيانات تاريخية مستخدمة", "لا شيء", "0",
                 "لم تدخل أي كمية أو سعر من الملف التاريخي إلى هذا الحصر"])

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    values = WC.evaluate_all(wb)
    cached = WC.inject_cached_values(path, values)

    check = load_workbook(path, data_only=True)
    numeric = sum(1 for (s, c), v in values.items() if isinstance(v, (int, float)))
    got = sum(1 for (s, c), v in values.items()
              if isinstance(v, (int, float)) and isinstance(check[s][c].value, (int, float)))
    formula_cells = len(values)
    return {"PATH": str(path), "SHEETS": wb.sheetnames, "QA": checks, "EXPORT_ROWS": rows,
            "REF": ref, "BOQ_CELLS": boq_cells, "VALUES": values, "WB": wb,
            "FORMULA_CELLS": formula_cells, "CACHED_VALUES_WRITTEN": cached,
            "NUMERIC_RESULTS_READ_BACK": got, "NUMERIC_RESULTS_EXPECTED": numeric,
            "ALL_SHEETS_RIGHT_TO_LEFT": all(wb[n].sheet_view.rightToLeft for n in wb.sheetnames),
            "PRICING_COLUMNS_LEFT_BLANK": ["الهالك %", "كمية الشراء", "سعر الوحدة", "الإجمالي"],
            "RATES_SUPPLIED": 0, "WASTE_APPLIED": False, "BOQ_LINES": len(boq_cells)}


def reconciliation(m, wbinfo):
    """One machine-readable table saying, quantity by quantity, whether the artifacts agree."""
    vals, ref, cells = wbinfo["VALUES"], wbinfo["REF"], wbinfo["BOQ_CELLS"]
    rows, wb = wbinfo["EXPORT_ROWS"], wbinfo["WB"]
    t = m["assembly"]["totals"]

    def cell(ref_str):
        sheet, coord = ref_str.split("!")
        sheet = sheet.strip("'")
        return vals[(sheet, coord)] if (sheet, coord) in vals else wb[sheet][coord].value

    def json_sum(trade=None, code=None):
        return round(sum(r["MEASURED_QUANTITY"] for r in rows
                         if (trade is None or r["TRADE"] == trade) and (code is None or r["ITEM_CODE"] == code)), 3)

    out = []

    def line(code, unit, json_v, wb_v, boq_v, frozen_v, included, note=""):
        vs = [v for v in (json_v, wb_v, boq_v, frozen_v) if v is not None]
        worst = round(max(vs) - min(vs), 6) if len(vs) > 1 else 0.0
        out.append({"ITEM_CODE": code, "UNIT": unit, "JSON": json_v, "WORKBOOK": wb_v, "BOQ": boq_v,
                    "FROZEN": frozen_v, "WORST_DIFFERENCE": worst, "TOLERANCE": CROSS_TOL_M2,
                    "AGREES": worst <= CROSS_TOL_M2, "BOQ_INCLUDED": included, "NOTE": note})

    line("AR-FL-01", "m2", json_sum("PORCELAIN_FLOOR"), round(cell(ref["FLOORS_BASE"]), 3),
         round(cell(cells["AR-FL-01"]), 3), t["INTERNAL_FLOOR_AREA_M2"], True,
         "internal, wet and kitchen components only")
    line("AR-FL-PENDING", "m2", json_sum("FLOOR_FINISH_UNCLASSIFIED"),
         round(cell(ref["FLOORS_PENDING"]), 3), round(cell(cells["AR-FL-PENDING"]), 3),
         m["floor_finish"]["PENDING_M2"], False, "stair, landing and unnamed - measured, finish unresolved")
    line("AR-CL-01", "m2", json_sum("CEILING"), round(cell(ref["CEILING"]), 3),
         round(cell(cells["AR-CL-01"]), 3), t["CEILING_M2"], False, "area final, finish type pending")
    line("AR-WP-01", "m2", json_sum("WALL_PORCELAIN"), round(cell(ref["PORCELAIN"]), 3),
         round(cell(cells["AR-WP-01"]), 3), t["WALL_PORCELAIN_NET_M2"], True, "")
    line("AR-WP-02", "m2", json_sum("TILE_PREPARATION"), round(cell(ref["PREP"]), 3),
         round(cell(cells["AR-WP-02"]), 3), t["TILE_PREPARATION_M2"], True, "")
    line("AR-BL-01", "m2", json_sum("BLOCKWORK_150"), round(cell(ref["BLOCK150"]), 3),
         round(cell(cells["AR-BL-01"]), 3), t["BLOCKWORK_150_NET_M2"], True, "")
    line("AR-BL-02", "m2", json_sum("BLOCKWORK_200"), round(cell(ref["BLOCK200"]), 3),
         round(cell(cells["AR-BL-02"]), 3), t["BLOCKWORK_200_NET_M2"], True, "")
    line("AR-PL-01", "m2", json_sum("INTERNAL_PLASTER"), round(cell(ref["INT_PLASTER"]), 3),
         round(cell(cells["AR-PL-01"]), 3), t["INTERNAL_PLASTER_NET_M2"], True, "")
    line("AR-PT-01", "m2", json_sum("INTERNAL_PAINT"), round(cell(ref["INT_PAINT"]), 3),
         round(cell(cells["AR-PT-01"]), 3), t["INTERNAL_PAINT_M2"], True, "")
    line("AR-PL-02", "m2", json_sum("EXTERNAL_PLASTER"), round(cell(ref["EXT_PLASTER"]), 3),
         round(cell(cells["AR-PL-02"]), 3), t["EXTERNAL_PLASTER_NET_M2"], True, "")
    line("AR-PT-02", "m2", json_sum("EXTERNAL_PAINT"), round(cell(ref["EXT_PAINT"]), 3),
         round(cell(cells["AR-PT-02"]), 3), t["EXTERNAL_PAINT_M2"], True, "")
    line("AR-IN-01", "m2", json_sum("WATERPROOFING"), round(cell(ref["INSULATION"]), 3),
         round(cell(cells["AR-IN-01"]), 3), t["WATERPROOFING_M2"], True, "")
    line("AR-RF-01", "m2", json_sum("PARAPET_BLOCKWORK"), round(cell(ref["PARAPET_BLOCK"]), 3),
         round(cell(cells["AR-RF-01"]), 3), t["PARAPET_BLOCKWORK_M2"], True, "")
    line("AR-DR-01", "m2", json_sum("DOORS"), round(cell(ref["DOORS"]), 3),
         round(cell(cells["AR-DR-01"]), 3), t["DOOR_AREA_M2"], True, "35 openings")
    line("AR-AL-01", "m2", json_sum("ALUMINIUM", "AR-AL-01"), round(cell(ref["WINDOWS_SUPPORTED"]), 4),
         round(cell(cells["AR-AL-01"]), 4), m["aluminium_split"]["SUPPORTED_M2"], True,
         "6 windows whose guide category has an authority")
    line("AR-AL-PENDING", "m2", json_sum("ALUMINIUM", "AR-AL-PENDING"),
         round(cell(ref["WINDOWS_PENDING"]), 4), round(cell(cells["AR-AL-PENDING"]), 4),
         m["aluminium_split"]["UNSUPPORTED_M2"], False,
         "AR-W-05: the guide category was used outside its band, so the height is ASK_THE_OWNER")
    line("AR-EX-01", "m2", json_sum("EXTERNAL_AREAS"), round(cell(ref["EXTERNAL"]), 3),
         round(cell(cells["AR-EX-01"]), 3), t["EXTERNAL_OPEN_AREA_M2"], False, "finish not stated")
    for code, key in (("AR-SK-01", "DRY_NAMED_INTERNAL_CANDIDATE"), ("AR-SK-02", "STAIR_OR_LANDING_PENDING"),
                      ("AR-SK-03", "UNNAMED_SPACE_PENDING")):
        line(code, "lm", json_sum("SKIRTING", code),
             round(cell({"AR-SK-01": ref["SK_DRY"], "AR-SK-02": ref["SK_STAIR"],
                         "AR-SK-03": ref["SK_UNNAMED"]}[code]), 3),
             round(cell(cells[code]), 3), None, False, "derived from frozen geometry, not part of the freeze")

    al = m["aluminium_split"]
    return {
        "ARTIFACT": "ALRASHED_QUANTITY_RECONCILIATION",
        "REVISION": REVISION,
        "GENERATED_AT_UTC": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "TOLERANCE_M2": CROSS_TOL_M2,
        "TOLERANCE_NOTE": "the artifacts publish dimensions at different rounding; the workbook's area sheet "
                          "carries raw dimensions and matches the frozen figures exactly, while a trade sheet "
                          "rebuilt from published millimetres can differ by a few mm2",
        "LINES": out,
        "ALL_AGREE": all(x["AGREES"] for x in out),
        "FLOOR_FINISH": m["floor_finish"],
        "COMPONENT_CENSUS": m["census"],
        "SKIRTING_BUCKETS": m["skirting_buckets"],
        "ALUMINIUM_SPLIT": {**al, "SUPPORTED_PLUS_UNSUPPORTED_M2": round(al["SUPPORTED_M2"]
                                                                         + al["UNSUPPORTED_M2"], 4),
                            "EQUALS_FROZEN_TOTAL": abs(al["SUPPORTED_M2"] + al["UNSUPPORTED_M2"]
                                                       - al["FROZEN_TOTAL_M2"]) < 1e-6},
        "HISTORICAL_OPENING_COUNTS": VA.HISTORICAL_COUNTS,
        "FROZEN": {"COMMIT": FROZEN_COMMIT, "DIGEST": FROZEN_DIGEST, "FILE_SHA256": frozen_file_sha256(),
                   "SHA256_MATCHES": frozen_file_sha256() == FROZEN_FILE_SHA256, "REWRITTEN": False},
    }


def finish():
    before = frozen_file_sha256()
    m = build()
    OUT.mkdir(parents=True, exist_ok=True)
    xl = OUT / "ALRASHED_DETAILED_QUANTITY_TAKEOFF.xlsx"
    wb = workbook(m, xl)
    rec_json = reconciliation(m, wb)
    (OUT / "ALRASHED_QUANTITY_RECONCILIATION.json").write_text(
        json.dumps(rec_json, indent=1, ensure_ascii=False, default=str), "utf-8")

    export = {
        "ARTIFACT": "ALRASHED_DETAILED_QUANTITY_EXPORT",
        "REVISION": REVISION,
        "PROJECT_ID": OI.PROJECT_ID, "PROJECT_NAME": PROJECT_NAME,
        "DRAWING_REVISION": DRAWING_REVISION, "SOURCE_FILES": SOURCE_FILES,
        "BASED_ON": {"ARTIFACT": "FULL_VILLA_BLIND_PRE_PRICING_TAKEOFF", "COMMIT": FROZEN_COMMIT,
                     "DIGEST": FROZEN_DIGEST, "FILE_SHA256": FROZEN_FILE_SHA256, "RESTATED": False},
        "SCHEMA_RULE": "DS-01 - MEASURED_QUANTITY, MEASURED_UNIT, WASTE_PERCENT, PROCUREMENT_QUANTITY, "
                       "UNIT_RATE and AMOUNT are separate fields; none is ever combined",
        "ENTITY_SEMANTICS": m["census"],
        "FLOOR_FINISH": m["floor_finish"],
        "SKIRTING_BUCKETS": m["skirting_buckets"],
        "ALUMINIUM_SPLIT": m["aluminium_split"],
        "APPROVAL_STATUS": "DRAFT",
        "RECORD_COUNT": len(wb["EXPORT_ROWS"]),
        "FIELDS_PER_RECORD": len(wb["EXPORT_ROWS"][0]),
        "RECORDS": wb["EXPORT_ROWS"],
        "QA": wb["QA"],
        "GENERATED_AT": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    }
    (OUT / "ALRASHED_DETAILED_QUANTITY_EXPORT.json").write_text(
        json.dumps(export, indent=1, ensure_ascii=False, default=str), "utf-8")

    comps = [c for f in m["floors"] for c in f["COMPONENTS"]]
    rec = {
        "ARTIFACT": "ALRASHED_DETAILED_QUANTITY_TAKEOFF",
        "REVISION": REVISION,
        "PROJECT_ID": OI.PROJECT_ID, "PROJECT_NAME": PROJECT_NAME,
        "WORKBOOK": {k: v for k, v in wb.items()
                     if k not in ("EXPORT_ROWS", "VALUES", "REF", "BOQ_CELLS", "WB")},
        "FLOORS": [{"FLOOR": f["FLOOR"], "COMPONENTS": len(f["COMPONENTS"]),
                    "RECTANGLES": sum(len(c["PARTS"]) for c in f["COMPONENTS"]),
                    "AREA_M2": f["FLOOR_TOTAL_M2"], "CLOSURE_RESIDUAL_M2": f["CLOSURE"]["RESIDUAL_M2"]}
                   for f in m["floors"]],
        "PROJECT_TOTAL_M2": m["project_total_m2"],
        "COMPONENT_CENSUS": m["census"],
        "FLOOR_FINISH": m["floor_finish"],
        "SKIRTING_BUCKETS": m["skirting_buckets"],
        "ALUMINIUM_SPLIT": m["aluminium_split"],
        "RECTANGLE_COUNT": sum(len(c["PARTS"]) for c in comps),
        "RECORD_COUNT": len(wb["EXPORT_ROWS"]),
        "TRADE_TOTALS": m["assembly"]["totals"],
        "FROZEN_UNCHANGED": {"COMMIT": FROZEN_COMMIT, "DIGEST": m["frozen"]["DIGEST"],
                             "FILE_SHA256_BEFORE": before, "FILE_SHA256_AFTER": frozen_file_sha256(),
                             "MATCHES": frozen_file_sha256() == FROZEN_FILE_SHA256 == before},
        "HISTORICAL_QUANTITY_USED": False,
        "GIT_HEAD": _git("rev-parse", "--short", "HEAD"),
    }
    rec["DIGEST"] = hashlib.sha256(json.dumps(
        {"F": rec["FLOORS"], "T": rec["TRADE_TOTALS"], "R": rec["RECORD_COUNT"], "C": rec["COMPONENT_CENSUS"]},
        sort_keys=True, default=str).encode()).hexdigest()[:16]
    (OUT / "ALRASHED_DETAILED_QUANTITY_TAKEOFF.json").write_text(
        json.dumps(rec, indent=1, ensure_ascii=False, default=str), "utf-8")
    assert rec["FROZEN_UNCHANGED"]["MATCHES"], "the frozen takeoff changed - refusing to publish"
    return rec, m, rec_json


if __name__ == "__main__":
    r, m, recon = finish()
    q = r["WORKBOOK"]["QA"]
    print(f"{r['ARTIFACT']} rev {r['REVISION']}  digest {r['DIGEST']}")
    print(f"  components {r['COMPONENT_CENSUS']['CONNECTED_COMPONENT_COUNT']} "
          f"= {r['COMPONENT_CENSUS']['NON_SLIVER_COMPONENT_COUNT']} non-sliver "
          f"+ {r['COMPONENT_CENSUS']['WALL_MATERIAL_OR_SLIVER_COUNT']} sliver")
    print(f"  floor finish  base {r['FLOOR_FINISH']['PORCELAIN_M2']}  pending {r['FLOOR_FINISH']['PENDING_M2']}")
    print(f"  records {r['RECORD_COUNT']}  formulas {r['WORKBOOK']['FORMULA_CELLS']}  "
          f"cached {r['WORKBOOK']['CACHED_VALUES_WRITTEN']}  "
          f"read back {r['WORKBOOK']['NUMERIC_RESULTS_READ_BACK']}/{r['WORKBOOK']['NUMERIC_RESULTS_EXPECTED']}")
    print(f"  QA {q['PASSED']}/{q['OF']}  {'ALL PASS' if q['ALL_PASS'] else 'FAILURES'}")
    for c in q["CHECKS"]:
        if not c["PASS"]:
            print("   FAIL", c["CHECK"], json.dumps(c["RESULT"], ensure_ascii=False, default=str)[:220])
    print(f"  reconciliation all agree: {recon['ALL_AGREE']}")
    print(f"  frozen unchanged: {r['FROZEN_UNCHANGED']['MATCHES']}")
