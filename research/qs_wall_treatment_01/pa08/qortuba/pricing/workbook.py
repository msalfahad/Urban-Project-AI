"""Write the pricing input matrix out as the trade-ordered workbook an estimator opens.

Trades are the sheets.  The room calculations follow as backup sheets, because that is the order a bill is priced in and the
order a reviewer checks it in.  The rate and amount columns are created and left empty: no rate was supplied for this project,
and a formula in an empty column is an invitation to fill it with a guess.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from research.qs_wall_treatment_01 import protocol as PR

OUT = Path(PR.OUT_DIR) / "pa08_qortuba_pricing"
QS = Path(PR.OUT_DIR) / "pa08_qortuba_qs01"
FILE = OUT / "QORTUBA_PRICING_INPUT_MATRIX.xlsx"

HEAD = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="E8EDF3")
BLOCKED = PatternFill("solid", fgColor="FFF4E5")
READY = PatternFill("solid", fgColor="EAF5EA")
THIN = Border(*[Side(style="thin", color="BFC9D4")] * 4)


def _auto(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header(ws, cols, widths, title=None):
    r = 1
    if title:
        ws.cell(1, 1, title).font = Font(bold=True, size=13)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        r = 3
    for i, c in enumerate(cols, 1):
        cell = ws.cell(r, i, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.freeze_panes = ws.cell(r + 1, 1)
    _auto(ws, widths)
    return r + 1


def build():
    mx = json.loads((OUT / "QORTUBA_PRICING_INPUT_MATRIX.json").read_text("utf-8"))
    rd = json.loads((OUT / "QORTUBA_PRICING_READINESS_SUMMARY.json").read_text("utf-8"))
    un = json.loads((OUT / "QORTUBA_MINIMUM_INPUTS_TO_UNLOCK.json").read_text("utf-8"))
    fr = json.loads((QS / "FREEZE_PA08_QORTUBA_ROOM_BY_ROOM_QS_01.json").read_text("utf-8"))
    wb = Workbook()

    # ---------------------------------------------------------------- cover
    ws = wb.active
    ws.title = "الغلاف"
    ws["A1"] = "قرطبة - مصفوفة مدخلات التسعير"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "QORTUBA PRICING INPUT MATRIX"
    ws["A2"].font = Font(size=12, color="555555")
    facts = [
        ("Workpaper", "PA08_QORTUBA_ROOM_BY_ROOM_QS_01"),
        ("Workpaper freeze digest", fr["DIGEST"]),
        ("Workpaper commit", fr["GIT_HEAD_AT_FREEZE"]),
        ("Storey", fr["STOREY"]),
        ("Trade rows", mx["COUNT"]),
        ("Rows with an established quantity", mx["CAN_PRICE_NOW"].get("YES", 0)),
        ("Rows blocked", mx["CAN_PRICE_NOW"].get("NO", 0)),
        ("Rates supplied", 0),
        ("Priced", "NO - no rate has been supplied, so no rate or amount cell is filled"),
        ("Rule", "MEASURED GEOMETRY, COMMERCIAL SCOPE, WASTE, PROCUREMENT, RATE and AMOUNT are separate layers"),
        ("Rule", "A geometric quantity is not a payable quantity and is not a purchase quantity"),
    ]
    for i, (k, v) in enumerate(facts, start=4):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    _auto(ws, [34, 110])

    # ---------------------------------------------------------------- readiness
    ws = wb.create_sheet("الجاهزية")
    cols = ["البند / Trade", "Trade (EN)", "الجاهزية", "بنود بكمية معتمدة", "إجمالي البنود", "السبب", "المطلوب لفتحها"]
    r = _header(ws, cols, [16, 26, 18, 20, 14, 78, 60], "Pricing readiness by trade")
    for x in rd["ROWS"]:
        ws.cell(r, 1, x["TRADE_AR"]); ws.cell(r, 2, x["TRADE"]); ws.cell(r, 3, x["READINESS"])
        ws.cell(r, 4, x["ITEMS_WITH_AN_ESTABLISHED_QUANTITY"]); ws.cell(r, 5, x["ITEMS"])
        ws.cell(r, 6, x["WHY"]); ws.cell(r, 7, "; ".join(x["UNLOCKS"]))
        fill = READY if x["READINESS"] == "READY_TO_PRICE" else BLOCKED if x["READINESS"] == "NOT_READY" else None
        for c in range(1, len(cols) + 1):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                ws.cell(r, c).fill = fill
        r += 1

    # ---------------------------------------------------------------- minimum inputs
    ws = wb.create_sheet("المطلوب")
    cols = ["المطلوب التالي / Next input", "النوع", "يفتح عدد بنود", "الترايدات", "البنود"]
    r = _header(ws, cols, [62, 18, 16, 40, 90], "The single next input that unlocks each blocked item")
    for x in un["ROWS"]:
        ws.cell(r, 1, x["NEXT_INPUT"]); ws.cell(r, 2, x["KIND"]); ws.cell(r, 3, x["UNBLOCKS_ITEMS"])
        ws.cell(r, 4, ", ".join(x["TRADES"])); ws.cell(r, 5, " | ".join(x["ITEMS"]))
        for c in range(1, len(cols) + 1):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
        r += 1

    # ---------------------------------------------------------------- one sheet per trade
    cols = ["البند", "الوصف", "الوحدة", "الكمية", "الهالك %", "كمية الشراء", "سعر الوحدة", "الإجمالي",
            "مصدر الكمية", "حالة الاعتماد", "ملاحظات"]
    widths = [34, 74, 10, 14, 12, 16, 14, 14, 44, 24, 60]
    for t in mx["TRADES"]:
        rows = [x for x in mx["ROWS"] if x["TRADE_CODE"] == t["CODE"]]
        if not rows:
            continue
        ws = wb.create_sheet(t["AR"][:31])
        r = _header(ws, cols, widths, f'{t["AR"]}  -  {t["EN"]}')
        for x in rows:
            ws.cell(r, 1, f'{x["ITEM"]} / {x["SUBITEM"]}')
            ws.cell(r, 2, x["FORMULA"])
            ws.cell(r, 3, x["UNIT"])
            ws.cell(r, 4, x["QUANTITY"])
            ws.cell(r, 9, x["SOURCE"])
            ws.cell(r, 10, x["STATUS"])
            ws.cell(r, 11, x["MISSING_INPUT"] or "")
            fill = READY if x["CAN_PRICE_NOW"] == "YES" else BLOCKED
            for c in range(1, len(cols) + 1):
                ws.cell(r, c).border = THIN
                ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
                ws.cell(r, c).fill = fill
            if x["QUANTITY"] is not None:
                ws.cell(r, 4).number_format = "0.000"
            r += 1
        ws.cell(r + 1, 1, "Empty columns are deliberate: waste, purchase quantity, rate and amount are not geometry.").font = \
            Font(italic=True, color="777777")

    # ---------------------------------------------------------------- backup: the room workpapers
    ws = wb.create_sheet("حساب الغرف")
    floors = json.loads((QS / "QORTUBA_QS01_FLOOR_CALCULATIONS.json").read_text("utf-8"))["ROWS"]
    sk = {x["ROOM_ID"]: x for x in json.loads((QS / "QORTUBA_QS01_SKIRTING_TAKEOFF.json").read_text("utf-8"))["ROWS"]}
    cols2 = ["الغرفة / Room", "رطب/جاف", "الشكل", "المعادلة", "المساحة م2", "محيط الجدار م", "خصم الأبواب م",
             "صافي النعلة م", "الحالة"]
    r = _header(ws, cols2, [26, 12, 26, 78, 14, 16, 16, 16, 22],
                "Calculation backup: the room-by-room workpaper the trade rows are cut from")
    for f in floors:
        s = sk.get(f["ROOM_ID"], {})
        ws.cell(r, 1, f["ROOM"]); ws.cell(r, 2, f["WET_OR_DRY"]); ws.cell(r, 3, f["SHAPE_TYPE"])
        ws.cell(r, 4, f["FORMULA"]); ws.cell(r, 5, f["METHOD_A_CAD_POLYGON_AREA_M2"])
        ws.cell(r, 6, s.get("GROSS_WALL_LINE_PERIMETER_LM")); ws.cell(r, 7, s.get("DOOR_WIDTH_DEDUCTION_LM"))
        ws.cell(r, 8, s.get("NET_SKIRTING_LM")); ws.cell(r, 9, f["STATUS"])
        for c in range(1, len(cols2) + 1):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(r, c).fill = SUB
        for c in (5, 6, 7, 8):
            ws.cell(r, c).number_format = "0.000"
        r += 1

    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(FILE)
    return FILE


if __name__ == "__main__":
    print("wrote", build())
