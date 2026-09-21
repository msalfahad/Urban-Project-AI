"""Write the Qortuba BOQ conversion workbook in the house's own column structure.

Fifteen columns, in Arabic, because that is the shape the Urban Projects bills already use.  Three of them carry quantities
and they are never merged: كمية القياس is what the drawing measured, قاعدة التحويل is what has to happen to it, and الكمية
النهائية is the payable figure.  The third is empty on every row but one, which is this phase's result rather than an
omission.

The room-by-room takeoff is demoted to a backup sheet at the end.  It is the workpaper behind the numbers, not the bill.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from research.qs_wall_treatment_01 import protocol as PR

OUT = Path(PR.OUT_DIR) / "pa08_qortuba_boq"
QS = Path(PR.OUT_DIR) / "pa08_qortuba_qs01"
FILE = OUT / "QORTUBA_BOQ_CONVERSION_WORKBOOK.xlsx"

NAVY = PatternFill("solid", fgColor="1F3A5F")
GREEN = PatternFill("solid", fgColor="D6EBD6")       # a final BOQ quantity
CREAM = PatternFill("solid", fgColor="FFF2CC")       # a measured input, not yet payable
PEACH = PatternFill("solid", fgColor="FBDBC8")       # a drawing or specification is missing
BLUE = PatternFill("solid", fgColor="DCE6F1")        # a rule decision is missing
GREY = PatternFill("solid", fgColor="E6E6E6")
SUB = PatternFill("solid", fgColor="F2F5F9")
THIN = Border(*[Side(style="thin", color="BFC9D4")] * 4)

STATUS_FILL = {"FINAL_QUANTITY_AVAILABLE": GREEN, "ONE_INPUT_REQUIRED": CREAM,
               "PROJECT_RULE_REQUIRED": BLUE, "SPEC_REQUIRED": PEACH,
               "DRAWING_REQUIRED": PEACH, "NOT_APPLICABLE": GREY}
STATUS_AR = {"FINAL_QUANTITY_AVAILABLE": "كمية نهائية جاهزة", "ONE_INPUT_REQUIRED": "ينقصه مُدخل واحد",
             "PROJECT_RULE_REQUIRED": "ينقصه قرار قاعدة", "SPEC_REQUIRED": "ينقصه مواصفة",
             "DRAWING_REQUIRED": "ينقصه مخطط", "NOT_APPLICABLE": "لا ينطبق"}

# §9: the fifteen columns, in the order the house bills use
COLS = ["البند", "الوصف", "وحدة التسعير", "كمية القياس", "وحدة القياس", "قاعدة التحويل", "الكمية النهائية",
        "% الهالك", "كمية الشراء", "سعر الوحدة", "الإجمالي", "مصدر القياس", "قاعدة القياس التجاري",
        "حالة الاعتماد", "ملاحظات"]
WIDTHS = [11, 34, 11, 13, 12, 46, 14, 9, 12, 11, 12, 30, 34, 20, 62]

LEVEL_AR = {"HISTORICAL_PRECEDENT": "سابقة مشروع سابق", "QORTUBA_PROJECT_RULE": "قاعدة مشروع قرطبة",
            "URBAN_STANDARD": "معيار أوربان"}


def _head(ws, cols, widths, title=None, sub=None):
    r = 1
    if title:
        ws.cell(1, 1, title).font = Font(bold=True, size=13, color="1F3A5F")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        if sub:
            ws.cell(2, 1, sub).font = Font(italic=True, size=9, color="666666")
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
        r = 4
    for i, c in enumerate(cols, 1):
        cell = ws.cell(r, i, c)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = NAVY
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[r].height = 30
    ws.freeze_panes = ws.cell(r + 1, 1)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return r + 1


def _put(ws, r, vals, fill=None, size=9):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = THIN
        cell.font = Font(size=size)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            cell.fill = fill
    return r + 1


def build():
    reg = json.loads((OUT / "URBAN_BOQ_RULE_REGISTRY.json").read_text("utf-8"))
    cm = json.loads((OUT / "QORTUBA_BOQ_CONVERSION_MAP.json").read_text("utf-8"))
    dec = json.loads((OUT / "QORTUBA_OWNER_DECISIONS_REQUIRED.json").read_text("utf-8"))
    rdy = json.loads((OUT / "QORTUBA_FINAL_BOQ_READINESS.json").read_text("utf-8"))
    saf = json.loads((OUT / "QORTUBA_RULE_SAFETY_REGISTER.json").read_text("utf-8"))
    cand = json.loads((OUT / "QORTUBA_CANDIDATE_BOQ_ITEMS.json").read_text("utf-8"))
    fz = json.loads((OUT / "FREEZE_URBAN_BOQ_RULE_REGISTRY.json").read_text("utf-8"))
    qfz = json.loads((QS / "FREEZE_PA08_QORTUBA_ROOM_BY_ROOM_QS_01.json").read_text("utf-8"))
    wb = Workbook()

    # ---------------------------------------------------------------- cover
    ws = wb.active
    ws.title = "الغلاف Cover"
    ws["A1"] = "قرطبة - خريطة التحويل إلى كميات جدول الكميات"
    ws["A1"].font = Font(bold=True, size=18, color="1F3A5F")
    ws["A2"] = "QORTUBA - BOQ CONVERSION MAP, built on the Urban Projects BOQ rule registry"
    ws["A2"].font = Font(size=11, color="555555")
    facts = [
        ("المشروع / Project", "قرطبة ق1 ش1 ج6  /  Qortuba, block 1, street 1, lane 6"),
        ("النطاق / Scope", qfz["STOREY"] + " - apartment only"),
        ("Measurement workpaper", "PA08_QORTUBA_ROOM_BY_ROOM_QS_01  " + qfz["DIGEST"][:16]),
        ("Rule registry freeze", fz["DIGEST"][:16]),
        ("Commit", fz["GIT_HEAD_AT_FREEZE"]),
        ("", ""),
        ("Rules in the registry", reg["COUNT"]),
        ("  at HISTORICAL_PRECEDENT", reg["BY_LEVEL"].get("HISTORICAL_PRECEDENT", 0)),
        ("  at QORTUBA_PROJECT_RULE", reg["BY_LEVEL"].get("QORTUBA_PROJECT_RULE", 0)),
        ("  at URBAN_STANDARD", reg["BY_LEVEL"].get("URBAN_STANDARD", 0)),
        ("", ""),
        ("BOQ items in the conversion map", cm["COUNT"]),
        ("Final BOQ quantities available now", cm["FINAL_QUANTITIES_WRITTEN"]),
        ("Owner decisions outstanding", dec["COUNT"]),
        ("Rates supplied", 0),
        ("", ""),
        ("THE RULE", "A precedent read off one previous project's bill is not a company standard. Nothing in this "
                     "registry has been promoted."),
        ("", "A rule belongs to the trade it was read from. The plaster half-deduction says nothing about ceramic, "
             "blockwork, paint or waterproofing."),
        ("", ""),
        ("THREE QUANTITY COLUMNS", "كمية القياس is what the drawing measured. قاعدة التحويل is what must happen to it. "
                                   "الكمية النهائية is the payable figure."),
        ("", "The third column is filled on one row only. That is the finding, not an omission."),
        ("", ""),
        ("NOT PRICED", "No rate has been supplied, so سعر الوحدة and الإجمالي are empty on every row."),
        ("NOT RECOMPUTED", "No Qortuba geometry was touched and no quantity was recomputed in this phase."),
        ("NOT BORROWED", "No historical quantity, rate, total or room dimension was used as a Qortuba input."),
    ]
    for i, (k, v) in enumerate(facts, start=4):
        ws.cell(i, 1, k).font = Font(bold=True, size=10)
        ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 120
    r = len(facts) + 6
    ws.cell(r, 1, "مفتاح الألوان / LEGEND").font = Font(bold=True)
    for j, (st, fill) in enumerate(STATUS_FILL.items()):
        ws.cell(r + 1 + j, 1, STATUS_AR[st]).fill = fill
        ws.cell(r + 1 + j, 1).border = THIN
        ws.cell(r + 1 + j, 2, st)

    # ---------------------------------------------------------------- the bill itself, §9 structure
    ws = wb.create_sheet("جدول الكميات BOQ")
    r = _head(ws, COLS, WIDTHS, "QORTUBA - BOQ CONVERSION MAP",
              "كمية القياس is a measurement. الكمية النهائية is a payable quantity. They are different columns because "
              "they are different things.")
    trade = None
    for x in cm["ROWS"]:
        if x["TRADE"] != trade:
            trade = x["TRADE"]
            ws.cell(r, 1, trade).font = Font(bold=True, size=11, color="1F3A5F")
            for c in range(1, len(COLS) + 1):
                ws.cell(r, c).fill = SUB
                ws.cell(r, c).border = THIN
            r += 1
        fin = x["FINAL_BOQ_QUANTITY"]
        vals = [x["MEASUREMENT_INPUT_ID"], x["BOQ_ITEM"], x["TARGET_PRICING_UNIT"],
                x["MEASURED_INPUT"], x["INPUT_UNIT"], x["CONVERSION_FORMULA"], fin.get("VALUE"),
                None, None, None, None,
                ("QORTUBA_QS01" if x["MEASURED_INPUT"] is not None else "لم يُقَس / not measured"),
                x["COMMERCIAL_RULE_REQUIRED"] or ("—" if not x["MISSING_INPUT"] else "لا قاعدة مطلوبة"),
                STATUS_AR[x["CURRENT_STATUS"]],
                " | ".join(z for z in (x["MISSING_INPUT"], x["NOTES"]) if z)]
        r = _put(ws, r, vals, STATUS_FILL[x["CURRENT_STATUS"]])
        for c in (4, 7, 9, 10, 11):
            ws.cell(r - 1, c).number_format = "0.000"

    # ---------------------------------------------------------------- the rule registry
    ws = wb.create_sheet("قواعد التسعير Rules")
    cols = ["الرقم", "البند / Trade", "البند / Item", "القاعدة / Rule", "وحدة التسعير", "أساس الكمية",
            "قاعدة الخصم", "قاعدة الإضافة", "الملف المصدر", "الورقة", "البند المصدر", "مستوى القاعدة",
            "ينطبق على قرطبة", "يتطلب موافقة المالك", "أمان التعميم", "ملاحظات"]
    r = _head(ws, cols, [8, 16, 30, 54, 11, 34, 30, 22, 30, 16, 26, 20, 14, 16, 26, 66],
              "URBAN PROJECTS BOQ RULE REGISTRY",
              "Every rule was read off ONE previous project. None has been promoted to an Urban standard.")
    for x in reg["ROWS"]:
        r = _put(ws, r, [x["RULE_ID"], x["TRADE"], x["ITEM"], x["RULE_DESCRIPTION"], x["PRICING_UNIT"],
                         x["QUANTITY_BASIS"], x["DEDUCTION_RULE"], x["ADDITION_RULE"],
                         x["SOURCE_WORKBOOK_TITLE_AR"] or "لا سابقة", x["SOURCE_SHEET"], x["SOURCE_ITEM"],
                         LEVEL_AR[x["RULE_LEVEL"]], "نعم" if x["APPLIES_TO_QORTUBA"] else "لا",
                         "نعم" if x["OWNER_APPROVAL_REQUIRED"] else "لا",
                         x["GENERALISATION_SAFETY"], x["NOTES"]],
                 GREEN if x["GENERALISATION_SAFETY"] == "SAFE_URBAN_STANDARD_CANDIDATE" else
                 PEACH if x["GENERALISATION_SAFETY"] == "CONTRACTOR_SPECIFIC" else CREAM)

    # ---------------------------------------------------------------- owner decisions
    ws = wb.create_sheet("قرارات المالك Decisions")
    cols = ["الرقم", "النوع", "السؤال / Question", "لماذا / Why it matters", "الخيارات / Options",
            "بنود تتأثر", "بنود لها قياس", "تصبح نهائية بهذا الجواب وحده", "إن لم يُجب"]
    r = _head(ws, cols, [8, 16, 58, 76, 46, 26, 14, 30, 46], "OWNER DECISIONS REQUIRED",
              "Ranked by how many items carrying a measured value this answer alone makes payable. "
              "A missing drawing is listed on the next sheet, not here: it is a document, not a decision.")
    for x in dec["ROWS"]:
        r = _put(ws, r, [x["DECISION_ID"], x["DECISION_KIND"], x["QUESTION"], x["WHY_IT_MATTERS"],
                         "\n".join(f"- {o}" for o in x["OPTIONS"]),
                         ", ".join(x["ROWS_AFFECTED"]), x["ROWS_CARRYING_A_MEASURED_VALUE"],
                         ", ".join(x["ROWS_THIS_ANSWER_ALONE_MAKES_FINAL"]) or "—",
                         x["IF_NOT_ANSWERED"]],
                 GREEN if x["ROWS_THIS_ANSWER_ALONE_MAKES_FINAL"] else CREAM)
    r += 2
    ws.cell(r, 1, "مستندات مطلوبة / DOCUMENTS REQUIRED - not decisions, simply not supplied").font = Font(bold=True)
    r += 1
    for d in dec["DOCUMENTS_NOT_DECISIONS"]:
        r = _put(ws, r, [None, "DOCUMENT", d["DOCUMENT"], None, None, ", ".join(d["ROWS_WAITING"]), d["COUNT"]], PEACH)

    # ---------------------------------------------------------------- readiness and safety
    ws = wb.create_sheet("الجاهزية Readiness")
    cols = ["البند / Trade", "عدد البنود", "كميات نهائية", "بنود متوقفة", "ما الذي يوقفها", "جاهز للتسعير"]
    r = _head(ws, cols, [22, 12, 14, 14, 96, 16], "FINAL BOQ READINESS BY TRADE",
              rdy["NO_SINGLE_READINESS_PERCENTAGE"])
    for x in rdy["ROWS"]:
        r = _put(ws, r, [x["TRADE"], x["BOQ_ITEMS"], x["FINAL_BOQ_QUANTITY_AVAILABLE"], x["BLOCKED"],
                         "\n".join(f"- {w}" for w in x["WHAT_BLOCKS_IT"]),
                         "نعم" if x["TRADE_READY_TO_PRICE"] else "لا"],
                 GREEN if x["TRADE_READY_TO_PRICE"] else CREAM)
    r += 2
    ws.cell(r, 1, "أمان التعميم / RULE SAFETY").font = Font(bold=True)
    r += 1
    for s in saf["ROWS"]:
        r = _put(ws, r, [s["SAFETY_CLASS"], s["COUNT"], None, None, s["MEANING"],
                         "نعم" if s["MAY_BE_APPLIED_TO_QORTUBA_WITHOUT_OWNER"] else "لا"],
                 GREEN if s["SAFETY_CLASS"] == "SAFE_URBAN_STANDARD_CANDIDATE" else PEACH)
    r += 2
    ws.cell(r, 1, "بند مقترح جديد / NEW CANDIDATE ITEM").font = Font(bold=True)
    r += 1
    for c in cand["ROWS"]:
        r = _put(ws, r, [c["CANDIDATE_ITEM_ID"], c["ITEM_AR"], c["PROPOSED_PRICING_UNIT_AR"],
                         c["QUANTITY_IF_APPROVED"]["VALUE"], c["WHY_NO_PRECEDENT"],
                         "معيار أوربان: لا" if not c["PROMOTED_TO_URBAN_STANDARD"] else "نعم"], CREAM)

    # ---------------------------------------------------------------- room backup, last
    ws = wb.create_sheet("حساب الغرف Backup")
    floors = json.loads((QS / "QORTUBA_QS01_FLOOR_CALCULATIONS.json").read_text("utf-8"))["ROWS"]
    sk = {x["ROOM_ID"]: x for x in json.loads((QS / "QORTUBA_QS01_SKIRTING_TAKEOFF.json").read_text("utf-8"))["ROWS"]}
    cols = ["الغرفة / Room", "رطب/جاف", "الشكل", "المعادلة / Formula", "المساحة م2", "محيط الجدار م",
            "خصم الأبواب م", "صافي النعلة م", "الحالة"]
    r = _head(ws, cols, [26, 12, 24, 76, 13, 15, 15, 15, 21], "CALCULATION BACKUP - ROOM BY ROOM TAKEOFF",
              "The workpaper behind the trade totals. It is not the bill, and nothing on this sheet was recomputed here.")
    for f in floors:
        s = sk.get(f["ROOM_ID"], {})
        r = _put(ws, r, [f["ROOM"], f["WET_OR_DRY"], f["SHAPE_TYPE"], f["FORMULA"],
                         f["METHOD_A_CAD_POLYGON_AREA_M2"], s.get("GROSS_WALL_LINE_PERIMETER_LM"),
                         s.get("DOOR_WIDTH_DEDUCTION_LM"), s.get("NET_SKIRTING_LM"), f["STATUS"]], SUB)
        for c in (5, 6, 7, 8):
            ws.cell(r - 1, c).number_format = "0.000"

    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(FILE)
    return FILE


if __name__ == "__main__":
    print("wrote", build())
