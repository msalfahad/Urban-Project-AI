"""ALSENAN - S3.1 multi-engine structural comparison (POST-FREEZE, comparison only).

Order (enforced):
  1. URBAN_PRODUCTION_FREEZE.json must exist and every hash in it must still match - production frozen first.
  2. The FREELANCER_QS_REFERENCE workbook is read by sha256 and its formulas are TRACED (lineage), not inferred
     from labels.
  3. EXTERNAL_ORACLE references (U-C4N, christiannp) enter only as recorded values from the owner's brief.
  4. Views are assembled with engine/source/comparison_scope; percentages only on matching bases.
  5. Rough steel = Urban concrete x URBAN_ROUGH_REBAR_PROFILE_V1 (sanity only) beside the actual BBS.

Nothing here writes to a production register. Freelancer / oracle values never change Urban quantities.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))

from engine.source import cad_oracle as CO  # noqa: E402
from engine.source import comparison_scope as CS  # noqa: E402
from engine.source import rough_rebar_sanity as RR  # noqa: E402
from engine.source import source_oracle_comparison as SOC  # noqa: E402
from engine.source import source_roles as SR  # noqa: E402
from engine.source import structural_population_discovery as SP  # noqa: E402

FREELANCER_WB = Path("/root/.claude/uploads/93607c01-16a4-590f-af7c-1c2701c3b240/d025d9f0-_________________1.xlsx")
FREELANCER_SHA = "ae608411ed5b9e64811a15c353adbdd017fa3e51abfdffcf0aa1454b85d1957c"
V3B = ROOT / "tests" / "alsenan" / "registers_v3b" / "BOQ_LINES_V3B.json"
S31 = ROOT / "research" / "alsenan_column_rebar_s3_1"
SLAB_REG = ROOT / "data" / "reports" / "URBAN_QTO_ALSENAN_PHASE_B2A1_GENERIC_QA_PATCH" / "registers" / \
    "SLAB_REGION_REGISTER.json"
PROFILE = ROOT / "engine" / "profiles" / "URBAN_ROUGH_REBAR_PROFILE_V1.json"


def sha(b):
    return hashlib.sha256(b).hexdigest()


def r3(x):
    return None if x is None else round(x, 3)


# ================================================================================================ 1. freeze gate
def check_freeze():
    fz = json.loads((HERE / "URBAN_PRODUCTION_FREEZE.json").read_text(encoding="utf-8"))
    for f, h in fz["inputs"].items():
        if sha((ROOT / f).read_bytes()) != h:
            raise SystemExit(f"production input {f} changed after the freeze record - comparison refused")
    return fz


# ================================================================================================ 2. lineage
SUMMARY_KEYS = [  # summary-sheet row label -> comparison category id (labels are matched, scope is TRACED)
    ("الخرسانه العاديه", "PLAIN_CONCRETE"), ("اجمالى القواعد", "FOUNDATIONS_RELATED"),
    ("اجمالي الشناجات والارضيه", "GROUND_BEAMS_AND_GROUND_SLAB"), ("اجمالى الحوائط والأعمدة", "WALLS_AND_COLUMNS"),
    ("الجسور", "BEAMS"), ("البلاطات", "SLABS"), ("الدرج+القبه", "STAIRS_AND_DOME"),
    ("حمام السباجة", "SWIMMING_POOL"), ("اجمالى الخرسانة المسلحة", "TOTAL_RC"), ("اجمالي كميات الحديد", "TOTAL_STEEL")]
REF = re.compile(r"(?:'([^']+)'|([^'!=+\-*/(),:\s]+))!\$?([A-Z]+)\$?(\d+)")
SUMR = re.compile(r"SUM\(\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)\)")
LOCAL = re.compile(r"(?<![!A-Z])\$?([A-Z]{1,2})\$?(\d+)")


def load_workbook_pair():
    from openpyxl import load_workbook
    raw = FREELANCER_WB.read_bytes()
    if sha(raw) != FREELANCER_SHA:
        raise SystemExit("freelancer workbook hash mismatch")
    return load_workbook(FREELANCER_WB), load_workbook(FREELANCER_WB, data_only=True)


def _label(ws, r):
    """Row label: column A of the row, or of the nearest labelled row above inside the same block."""
    for rr in range(r, 0, -1):
        v = ws.cell(row=rr, column=1).value
        if v is not None and str(v).strip():
            return str(v).strip(), rr
    return None, None


def _section(ws, r):
    for rr in range(r, 0, -1):
        for col in (1, 3):
            v = ws.cell(row=rr, column=col).value
            if v and str(v).strip().startswith("حصر"):
                return str(v).strip()
    return None


def trace(wf, wv, sheet, cell, seen=None):
    """Leaf rows feeding a cell: follow cross-sheet refs, SUM ranges and +/- of cells until a row-volume formula."""
    seen = seen if seen is not None else set()
    key = (sheet, cell)
    if key in seen:
        return []
    seen.add(key)
    ws = wf[sheet]
    f = ws[cell].value
    if not (isinstance(f, str) and f.startswith("=")):
        return [{"sheet": sheet, "cell": cell, "value": ws[cell].value, "formula": None, "kind": "TYPED_VALUE"}]
    if "#REF!" in f:
        return [{"sheet": sheet, "cell": cell, "formula": f, "kind": "BROKEN_REFERENCE"}]
    out = []
    refs = REF.findall(f)
    if refs:
        for a, b, c, r in refs:
            out += trace(wf, wv, a or b, f"{c}{r}", seen)
        return out
    m = SUMR.search(f)
    if m and f.strip().upper().startswith("=SUM("):
        c0, r0, c1, r1 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        for r in range(r0, r1 + 1):
            v = ws[f"{c0}{r}"].value
            if v is None:
                continue
            out += trace(wf, wv, sheet, f"{c0}{r}", seen)
        return out
    locs = LOCAL.findall(f[1:])
    row = int(re.search(r"\d+", cell).group())
    if all(int(r) == row for _, r in locs):           # a row formula (E x F etc.) -> a leaf
        lab, lab_row = _label(ws, row)
        return [{"sheet": sheet, "cell": cell, "row": row, "formula": f, "kind": "LEAF_ROW",
                 "label": lab, "label_row": lab_row, "section": _section(ws, row),
                 "B": ws[f"B{row}"].value, "C": ws[f"C{row}"].value, "D": ws[f"D{row}"].value,
                 "count": ws[f"F{row}"].value, "value": wv[sheet][cell].value}]
    for c, r in locs:
        out += trace(wf, wv, sheet, f"{c}{r}", seen)
    return out


def classify_leaf(sheet, section, label):
    """Physical kind of a freelancer leaf row, from its sheet, section header and label (comparison mapping)."""
    s, lab = section or "", label or ""
    if lab.startswith("العاديه"):
        return "PLAIN_CONCRETE"
    if re.match(r"^F[\d.]*|^F\.", lab) and "القواعد" in sheet:
        return "FOOTING"
    if lab.startswith("STB"):
        return "STRAP_BEAM"
    if "الشناج المحيط" in lab:
        return "PERIMETER_GROUND_BEAM"
    if "الارضيه الخرسانيه" in lab:
        return "GROUND_SLAB"
    if "شناج السور" in lab:
        return "BOUNDARY_GROUND_BEAM"
    if "شناجات" in lab:
        return "GROUND_BEAM"
    if "المصعد" in lab:
        return "LIFT_WALL"
    if "رقاب" in s:
        return "COLUMN_NECK"
    if "الأعمدة" in s:
        return {"الأرضي": "COLUMN_GF", "الأول": "COLUMN_1F", "السطح": "COLUMN_2F"}.get(s.split()[-1], "COLUMN")
    if "كمرات" in s:
        return {"الأرضي": "BEAM_GROSS_GF", "الأول": "BEAM_GROSS_1F", "السطح": "BEAM_GROSS_2F"}.get(s.split()[-1],
                                                                                                    "BEAM_GROSS")
    if "بلاطات" in s:
        return {"الأرضي": "SLAB_NET_GF", "الأول": "SLAB_NET_1F", "السطح": "SLAB_NET_2F"}.get(s.split()[-1], "SLAB")
    if "القبه" in lab:
        return "DOME"
    if "السلالم" in s:
        return "STAIR"
    if "السباحة" in s:
        return "POOL_ROOM_FLOOR" if "الغرفه" in lab else ("POOL_PUMP_WALL" if "المضخة" in lab else "POOL")
    return "UNCLASSIFIED"


def lineage():
    wf, wv = load_workbook_pair()
    summ = wf.worksheets[0]
    rows = {}
    for r in range(1, summ.max_row + 1):
        lab = summ.cell(row=r, column=1).value
        if lab:
            rows[str(lab).strip()] = r
    cats = []
    for lab, cid in SUMMARY_KEYS:
        r = rows[lab]
        qcol = "F" if cid == "TOTAL_RC" else ("H" if cid == "TOTAL_STEEL" else "E")
        cell = f"{qcol}{r}"
        f = summ[cell].value
        leaves = trace(wf, wv, summ.title, cell) if cid not in ("TOTAL_RC",) else []
        if cid == "TOTAL_RC":
            src = [{"sheet": summ.title, "cell": f"E{i}", "kind": "SUMMARY_ROW"} for i in range(17, 24)]
        else:
            src = leaves
        for x in src:
            if x.get("kind") == "LEAF_ROW":
                x["physical_kind"] = classify_leaf(x["sheet"], x["section"], x["label"])
        steel = summ[f"H{r}"].value
        subk = defaultdict(float)
        for x in src:
            if x.get("kind") == "LEAF_ROW":
                subk[x["physical_kind"]] += x["value"] or 0.0
        cats.append({"category_id": cid, "arabic_name": lab, "summary_cell": f"{summ.title}!{cell}",
                     "summary_value": wv[summ.title][cell].value, "formula": f,
                     "steel_cell": f"{summ.title}!H{r}" if steel is not None else None,
                     "steel_value_t": float(steel) if steel not in (None, "") else None,
                     "steel_is_formula": isinstance(steel, str) and str(steel).startswith("="),
                     "source_sheets": sorted({x["sheet"] for x in src}),
                     "source_cells": [f"{x['sheet']}!{x['cell']}" for x in src],
                     "subcomponents": [{k: x.get(k) for k in ("sheet", "cell", "row", "label", "section", "B", "C", "D",
                                                              "count", "formula", "value", "physical_kind", "kind")}
                                       for x in src],
                     "quantity_by_physical_kind": {k: r3(v) for k, v in sorted(subk.items())},
                     "quantity": wv[summ.title][cell].value, "unit": "t" if cid == "TOTAL_STEEL" else "m3"})
    notes = lineage_notes(cats, wf, wv)
    for c in cats:
        c["possible_overlap"] = notes["overlap"].get(c["category_id"], [])
        c["notes"] = notes["notes"].get(c["category_id"], [])
    return cats, notes


def lineage_notes(cats, wf, wv):
    by = {c["category_id"]: c for c in cats}
    ov, nt = defaultdict(list), defaultdict(list)
    lbl = {"FOUNDATIONS_RELATED": "footings", "GROUND_BEAMS_AND_GROUND_SLAB": "ground beams + ground slab",
           "WALLS_AND_COLUMNS": "walls + columns", "STAIRS_AND_DOME": "stairs + dome", "SWIMMING_POOL": "pool"}
    expect = {"FOUNDATIONS_RELATED": {"FOOTING"}, "GROUND_BEAMS_AND_GROUND_SLAB": {"GROUND_BEAM", "GROUND_SLAB"},
              "WALLS_AND_COLUMNS": {"COLUMN_GF", "COLUMN_1F", "COLUMN_2F", "COLUMN"},
              "STAIRS_AND_DOME": {"STAIR", "DOME"}, "SWIMMING_POOL": {"POOL"}}
    for cid, exp in expect.items():
        extra = {k: v for k, v in by[cid]["quantity_by_physical_kind"].items() if k not in exp}
        if extra:
            nt[cid].append(f"label reads '{lbl[cid]}' but the traced rows also include {extra} m3 "
                           "(scope wider than the label)")
    ov["FOUNDATIONS_RELATED"].append("PERIMETER_GROUND_BEAM sits in the foundations total; Urban classes it as "
                                     "ground beam (exterior GB)")
    ov["GROUND_BEAMS_AND_GROUND_SLAB"].append("BOUNDARY_GROUND_BEAM (boundary wall beam) counted here; Urban keeps "
                                              "boundary wall separate")
    ov["WALLS_AND_COLUMNS"].append("COLUMN_NECK (1.5 m necks) counted with columns; Urban treats necks as "
                                   "foundation-related")
    ov["BEAMS"].append("BEAM_GROSS includes the slab depth over each beam -> overlaps Urban SLAB full depth")
    ov["SLABS"].append("SLAB_NET areas exclude the beam footprints (allocated to BEAM_GROSS)")
    b = wf["كمرات"]
    if b["I82"].value and "I60" in str(b["I82"].value) and b["I59"].value:
        nt["BEAMS"].append(f"1F subtotal كمرات!I82 = {b['I82'].value} skips row 59 "
                           f"(B1 {wv['كمرات']['I59'].value:.3f} m3) that كمرات!I79 includes - formula inconsistency")
    pool = wf["حمام السباحة"]
    if "#REF!" in str(pool["D19"].value):
        shown = str(pool["D19"].value).lstrip("=").replace("#REF!", "<BROKEN_REF>")
        nt["SWIMMING_POOL"].append(f"stray cross-sheet total حمام السباحة!D19 ({shown}) contains a broken "
                                   "reference and evaluates to an error (not used by the summary)")
    st = by["TOTAL_STEEL"]
    if not st["steel_is_formula"]:
        nt["TOTAL_STEEL"].append("steel tonnages on the summary sheet are typed values, not formulas - no "
                                 "bar-by-bar lineage exists in the workbook")
    s1 = wf["بلاطات "]
    nt["SLABS"].append(f"1F slab deduction 'بلاطات '!F16 = {s1['F16'].value} is multiplied by G16 = "
                       f"{s1['G16'].value} - the deduction is zeroed")
    return {"overlap": ov, "notes": nt}


# ================================================================================================ 3. references
ORACLES = {  # recorded from the owner's S3.1 brief §16 (blind MCP runs) - comparison only
    "UC4N": {"source": "owner brief S3.1 §16 (U-C4N blind AutoCAD MCP run)", "built_up_m2": 590.362,
             "net_slab_m2": {"GF": 291.425, "1F": 171.235, "2F": 55.590}, "total_rc_m3": 337.812,
             "FOOTING": 66.579, "STRAP_BEAM": 6.122, "GROUND_BEAM": 36.005, "GROUND_SLAB": 32.183,
             "COLUMN": {"FOUNDATION": 6.150, "GF": 26.235, "1F": 8.888, "2F": 3.797},
             "BEAM_DOWNSTAND_TOTAL": 50.841, "SLAB_TOTAL": 95.689, "STAIR": 5.325,
             "net_rebar_t": 22.619, "rebar_state": "KNOWN_INCOMPLETE"},
    "CHRISTIANNP": {"source": "owner brief S3.1 §16 (christiannp blind MCP run)", "built_up_m2": 598.708,
                    "net_slab_m2": {"GF": 290.555, "1F": 179.325, "2F": 55.275}, "total_rc_m3": 333.173,
                    "FOOTING": 65.669, "STRAP_BEAM": 6.053, "GROUND_BEAM": 36.006, "GROUND_SLAB": 32.404,
                    "COLUMN": {"FOUNDATION": 6.150, "GF": 26.235, "1F": 8.888, "2F": 3.377},
                    "BEAM_DOWNSTAND": {"GF": 25.206, "1F": 18.562, "2F": 2.823},
                    "SLAB": {"GF": 58.111, "1F": 28.692, "2F": 9.950}, "STAIR": 5.048,
                    "net_rebar_t": 22.916, "rebar_state": "KNOWN_INCOMPLETE"}}


def urban_components():
    """Urban physical components from the frozen V3b BOQ lines (+ B2A1 slab register for areas, S3.1 for column
    steel). State: COMPUTED -> VERIFIED, PARTIAL -> LOWER_BOUND, owner method / review -> PROVISIONAL, BLOCKED."""
    lines = json.loads(V3B.read_text(encoding="utf-8"))["lines"]
    kind = {"C-FTG": ("FOOTING", "FOUNDATION_PAD"), "C-FTG-FF10": ("FOOTING", "FOUNDATION_PAD"),
            "C-STR": ("STRAP_BEAM", "FOUNDATION_STRAP_BEAM"), "C-NECK": ("COLUMN_NECK", "FOUNDATION_NECK"),
            "C-GB-INT": ("GROUND_BEAM", "GROUND_BEAM"), "C-GB-EXT": ("GROUND_BEAM_EXTERIOR", "GROUND_BEAM"),
            "C-GSLAB-ZONE-1": ("GROUND_SLAB", "GROUND_SLAB"), "C-GSLAB-ZONE-2": ("GROUND_SLAB", "GROUND_SLAB"),
            "C-COL-GF": ("COLUMN_GF", "COLUMN"), "C-COL-1F": ("COLUMN_1F", "COLUMN"),
            "C-COL-2F": ("COLUMN_2F", "COLUMN"), "C-COL-GF-RES": ("COLUMN_GF", "COLUMN"),
            "C-COL-1F-RES": ("COLUMN_1F", "COLUMN"), "C-COL-2F-RES": ("COLUMN_2F", "COLUMN"),
            "C-JNT-GF": ("JOINT_GF", "BEAM_COLUMN_JOINT"), "C-JNT-1F": ("JOINT_1F", "BEAM_COLUMN_JOINT"),
            "C-JNT-2F": ("JOINT_2F", "BEAM_COLUMN_JOINT"),
            "C-BEAM-GF": ("BEAM_DOWNSTAND_GF", "BEAM"), "C-BEAM-1F": ("BEAM_DOWNSTAND_1F", "BEAM"),
            "C-BEAM-2F": ("BEAM_DOWNSTAND_2F", "BEAM"), "C-BEAM-GF-RES": ("BEAM_DOWNSTAND_GF", "BEAM"),
            "C-BEAM-1F-RES": ("BEAM_DOWNSTAND_1F", "BEAM"),
            "C-SLAB-GF": ("SLAB_FULL_GF", "SOLID_SLAB"), "C-SLAB-1F": ("SLAB_FULL_1F", "SOLID_SLAB"),
            "C-SLAB-2F": ("SLAB_FULL_2F", "SOLID_SLAB"),
            "C-LINT-GF": ("LINTEL", "LINTEL"), "C-LINT-1F": ("LINTEL", "LINTEL"), "C-LINT-2F": ("LINTEL", "LINTEL"),
            "C-STAIR": ("STAIR", "STAIR"), "C-POOL": ("POOL", "POOL_WALL"), "C-SWALL": ("STRUCTURAL_WALL",
                                                                                       "STRUCTURAL_WALL"),
            "C-BWALL": ("BOUNDARY_GROUND_BEAM", "BOUNDARY_GROUND_BEAM"),
            "C-BLD-FULL": ("PLAIN_CONCRETE", "PLAIN_CONCRETE"), "C-POOL-BLD": ("PLAIN_CONCRETE", "PLAIN_CONCRETE")}
    out = []
    for ln in lines:
        if ln["trade"] != "CONCRETE":
            continue
        code = ln["code"]
        if code.startswith("C-DOME-SHELL"):
            k = ("DOME", "DOME")
        elif code.startswith("C-DOME-RING"):
            k = ("DOME_RING_BEAM", "RING_BEAM")
        else:
            k = kind.get(code)
        if k is None:
            continue
        st = {"COMPUTED": "VERIFIED", "PARTIAL": "LOWER_BOUND", "BLOCKED": "BLOCKED"}.get(ln["status"], "PROVISIONAL")
        if code == "C-BLD-FULL" or "owner" in (ln.get("authority") or "").lower():
            st = "PROVISIONAL" if ln["status"] != "BLOCKED" else st
        if code.startswith("C-GSLAB"):
            st = "LOWER_BOUND"          # ground-slab extent not drawn (S2 STR-GRO-001): computed zones are a minimum
        out.append({"component_id": f"URBAN:{code}", "kind": k[0], "element_class": k[1], "quantity": ln["qty"],
                    "unit": "m3", "state": st if ln["qty"] is not None else "BLOCKED", "source_line": code,
                    "v3b_status": ln["status"], "formula": ln.get("formula")})
    return out, lines


def urban_rebar(lines):
    s31 = json.loads((S31 / "COLUMN_REBAR_S3_1_SUMMARY.json").read_text(encoding="utf-8"))["headline"]
    cat = {"FOOTINGS": "FOUNDATIONS_RELATED", "FOOTING_F_F10": "FOUNDATIONS_RELATED",
           "COLUMN_STARTERS": "FOUNDATIONS_RELATED", "GROUND": "GROUND_BEAMS_AND_GROUND_SLAB",
           "GROUND_SLAB": "GROUND_BEAMS_AND_GROUND_SLAB", "GROUND_BEAM_EXT": "GROUND_BEAMS_AND_GROUND_SLAB",
           "BEAMS": "BEAMS", "BEAM_RESIDUE": "BEAMS", "BEAM_SIDE_BARS": "BEAMS", "SLAB": "SLABS",
           "DOME": "STAIRS_AND_DOME", "STAIRS": "STAIRS_AND_DOME", "POOL": "SWIMMING_POOL"}
    act = defaultdict(lambda: {"released_kg": 0.0, "provisional_kg": 0.0, "blocked_components": [],
                               "complete": True, "sources": []})
    for ln in lines:
        if ln["trade"] != "REBAR" or ln["group"] in ("PROCUREMENT", "COLUMNS", "LINTELS", "BOUNDARY_WALL"):
            continue
        c = cat.get(ln["group"])
        if c is None:
            continue
        a = act[c]
        a["sources"].append(ln["code"])
        if ln["qty"] is None or ln["status"] == "BLOCKED":
            a["blocked_components"].append(ln["code"])
            a["complete"] = False
        else:
            a["released_kg"] += ln["qty"]
            if ln["status"] != "COMPUTED":
                a["complete"] = False
    a = act["WALLS_AND_COLUMNS"]
    a["released_kg"] = s31["VERIFIED_KG"] + s31["LOWER_BOUND_KG"]
    a["provisional_kg"] = s31["PROVISIONAL_KG"]
    a["blocked_components"] = [f"S3.1 blocked modelled {s31['BLOCKED_MODELLED_KG']} kg",
                               f"S3.1 unquantified blocked parts {s31['UNQUANTIFIED_BLOCKED_PARTS']}"]
    a["complete"] = False
    a["sources"] = ["research/alsenan_column_rebar_s3_1 (S3.1 columns, D2/162)"]
    return {k: dict(v) for k, v in act.items()}, s31


# ================================================================================================ 4. views
def view_defs():
    U = CS.make_view("URBAN_PHYSICAL_VIEW", [
        {"category": "FOOTINGS", "measurement_basis": "PAD_LxWxH_SCHEDULE", "included_components": ["FOOTING"],
         "overlap_policy": "pads only"},
        {"category": "STRAP_BEAMS", "measurement_basis": "STRAP_BAND", "included_components": ["STRAP_BEAM"],
         "overlap_policy": "between footings"},
        {"category": "NECKS", "measurement_basis": "NECK_SECTION_x_HEIGHT", "included_components": ["COLUMN_NECK"],
         "overlap_policy": "footing top to ground beam"},
        {"category": "GROUND_BEAMS", "measurement_basis": "GB_SPAN_x_SECTION",
         "included_components": ["GROUND_BEAM", "GROUND_BEAM_EXTERIOR"], "overlap_policy": "interior + exterior"},
        {"category": "GROUND_SLAB", "measurement_basis": "SLAB_ZONE_AREA_x_T", "included_components": ["GROUND_SLAB"],
         "overlap_policy": "zones minus beam bands / columns"},
        {"category": "BOUNDARY", "measurement_basis": "BOUNDARY_WALL_BEAM", "excluded_components": [],
         "included_components": ["BOUNDARY_GROUND_BEAM"], "overlap_policy": "separate from building"},
        {"category": "COLUMNS_GF", "measurement_basis": "COLUMN_CLEAR_HEIGHT_JOINT_SEPARATE",
         "included_components": ["COLUMN_GF"], "overlap_policy": "joint separate"},
        {"category": "COLUMNS_1F", "measurement_basis": "COLUMN_CLEAR_HEIGHT_JOINT_SEPARATE",
         "included_components": ["COLUMN_1F"], "overlap_policy": "joint separate"},
        {"category": "COLUMNS_2F", "measurement_basis": "COLUMN_CLEAR_HEIGHT_JOINT_SEPARATE",
         "included_components": ["COLUMN_2F"], "overlap_policy": "joint separate"},
        {"category": "JOINTS", "measurement_basis": "COLUMN_SECTION_x_BEAM_DOWNSTAND",
         "included_components": ["JOINT_GF", "JOINT_1F", "JOINT_2F"], "overlap_policy": "own component"},
        {"category": "BEAMS_GF", "measurement_basis": "BEAM_DOWNSTAND_ONLY",
         "included_components": ["BEAM_DOWNSTAND_GF"], "overlap_policy": "slab over beam in SLABS"},
        {"category": "BEAMS_1F", "measurement_basis": "BEAM_DOWNSTAND_ONLY",
         "included_components": ["BEAM_DOWNSTAND_1F"], "overlap_policy": "slab over beam in SLABS"},
        {"category": "BEAMS_2F", "measurement_basis": "BEAM_DOWNSTAND_ONLY",
         "included_components": ["BEAM_DOWNSTAND_2F"], "overlap_policy": "slab over beam in SLABS"},
        {"category": "SLABS_GF", "measurement_basis": "SLAB_FULL_DEPTH_NET_PLATE",
         "included_components": ["SLAB_FULL_GF"], "overlap_policy": "full depth incl. over beams"},
        {"category": "SLABS_1F", "measurement_basis": "SLAB_FULL_DEPTH_NET_PLATE",
         "included_components": ["SLAB_FULL_1F"], "overlap_policy": "full depth incl. over beams"},
        {"category": "SLABS_2F", "measurement_basis": "SLAB_FULL_DEPTH_NET_PLATE",
         "included_components": ["SLAB_FULL_2F"], "overlap_policy": "full depth incl. over beams"},
        {"category": "LINTELS", "measurement_basis": "LINTEL_SCHEDULE", "included_components": ["LINTEL"],
         "overlap_policy": "separate"},
        {"category": "STAIRS", "measurement_basis": "STAIR_WAIST_STEPS_LANDINGS", "included_components": ["STAIR"],
         "overlap_policy": "separate"},
        {"category": "DOME", "measurement_basis": "DOME_SHELL_AREA_x_T", "included_components": ["DOME"],
         "overlap_policy": "ring separate"},
        {"category": "DOME_RING", "measurement_basis": "RING_BEAM", "included_components": ["DOME_RING_BEAM"],
         "overlap_policy": "separate"},
        {"category": "POOL", "measurement_basis": "POOL_WALLS_BASE", "included_components": ["POOL"],
         "overlap_policy": "pool only"},
        {"category": "STRUCTURAL_WALLS", "measurement_basis": "WALL", "included_components": ["STRUCTURAL_WALL"],
         "overlap_policy": "separate"},
        {"category": "PLAIN_CONCRETE", "measurement_basis": "BLINDING", "included_components": ["PLAIN_CONCRETE"],
         "overlap_policy": "not reinforced"}])
    F = CS.make_view("FREELANCER_QS_VIEW", [
        {"category": "FOOTINGS", "measurement_basis": "PAD_LxWxH_SCHEDULE", "included_components": ["FOOTING"],
         "overlap_policy": "pads only"},
        {"category": "STRAP_BEAMS", "measurement_basis": "STRAP_BAND", "included_components": ["STRAP_BEAM"],
         "overlap_policy": "between footings"},
        {"category": "NECKS", "measurement_basis": "NECK_SECTION_x_1.5M", "included_components": ["COLUMN_NECK"],
         "overlap_policy": "assumed 1.5 m necks"},
        {"category": "GROUND_BEAMS", "measurement_basis": "GB_RUN_LENGTH_x_30x60",
         "included_components": ["GROUND_BEAM", "PERIMETER_GROUND_BEAM"],
         "overlap_policy": "longitudinal + transverse runs + perimeter"},
        {"category": "GROUND_SLAB", "measurement_basis": "GROUND_SLAB_GROSS_AREA_x_T",
         "included_components": ["GROUND_SLAB"], "overlap_policy": "one area"},
        {"category": "BOUNDARY", "measurement_basis": "BOUNDARY_WALL_BEAM", "included_components":
            ["BOUNDARY_GROUND_BEAM"], "overlap_policy": "counted with ground beams in the summary"},
        {"category": "COLUMNS_GF", "measurement_basis": "COLUMN_BELOW_BEAM_SOFFIT_3.75",
         "included_components": ["COLUMN_GF"], "overlap_policy": "no joint"},
        {"category": "COLUMNS_1F", "measurement_basis": "COLUMN_BELOW_BEAM_SOFFIT_3.45",
         "included_components": ["COLUMN_1F"], "overlap_policy": "no joint"},
        {"category": "COLUMNS_2F", "measurement_basis": "COLUMN_BELOW_BEAM_SOFFIT_3.45",
         "included_components": ["COLUMN_2F"], "overlap_policy": "no joint"},
        {"category": "BEAMS_GF", "measurement_basis": "BEAM_GROSS_DEPTH", "included_components": ["BEAM_GROSS_GF"],
         "overlap_policy": "slab over beam in BEAMS"},
        {"category": "BEAMS_1F", "measurement_basis": "BEAM_GROSS_DEPTH", "included_components": ["BEAM_GROSS_1F"],
         "overlap_policy": "slab over beam in BEAMS"},
        {"category": "BEAMS_2F", "measurement_basis": "BEAM_GROSS_DEPTH", "included_components": ["BEAM_GROSS_2F"],
         "overlap_policy": "slab over beam in BEAMS"},
        {"category": "SLABS_GF", "measurement_basis": "SLAB_NET_OF_BEAMS", "included_components": ["SLAB_NET_GF"],
         "overlap_policy": "between beams"},
        {"category": "SLABS_1F", "measurement_basis": "SLAB_NET_OF_BEAMS", "included_components": ["SLAB_NET_1F"],
         "overlap_policy": "between beams"},
        {"category": "SLABS_2F", "measurement_basis": "SLAB_NET_OF_BEAMS", "included_components": ["SLAB_NET_2F"],
         "overlap_policy": "between beams"},
        {"category": "STAIRS", "measurement_basis": "STAIR_SLAB_x_0.28", "included_components": ["STAIR"],
         "overlap_policy": "waist + steps as one 0.28 slab"},
        {"category": "DOME", "measurement_basis": "DOME_AREA_x_T", "included_components": ["DOME"],
         "overlap_policy": "3 domes"},
        {"category": "POOL", "measurement_basis": "POOL_WALLS_BASE_PLUS_PUMP_ROOM",
         "included_components": ["POOL", "POOL_PUMP_WALL", "POOL_ROOM_FLOOR"], "overlap_policy": "pool + pump room"},
        {"category": "STRUCTURAL_WALLS", "measurement_basis": "LIFT_WALLS_20", "included_components": ["LIFT_WALL"],
         "overlap_policy": "lift walls"},
        {"category": "PLAIN_CONCRETE", "measurement_basis": "FOOTPRINT_15x31.37x0.1",
         "included_components": ["PLAIN_CONCRETE"], "overlap_policy": "one footprint"}])
    return U, F


GROUPS = [  # normalisation groups: unions comparable although each side allocates differently
    {"group_id": "FRAME_GF", "categories": ["COLUMNS_GF", "BEAMS_GF", "SLABS_GF"], "plus_urban": ["JOINTS_GF"],
     "why": "beam gross vs downstand and slab net vs full depth reallocate the same concrete within a floor"},
    {"group_id": "FRAME_1F", "categories": ["COLUMNS_1F", "BEAMS_1F", "SLABS_1F"], "plus_urban": ["JOINTS_1F"],
     "why": "same"},
    {"group_id": "FRAME_2F", "categories": ["COLUMNS_2F", "BEAMS_2F", "SLABS_2F"], "plus_urban": ["JOINTS_2F"],
     "why": "same"},
    {"group_id": "GROUND_ZONE", "categories": ["STRAP_BEAMS", "GROUND_BEAMS", "GROUND_SLAB"],
     "why": "straps / perimeter beam / ground beams / ground slab are split differently across the QS categories"},
]


def pct(a, b):
    return None if (a is None or b is None or not b) else 100.0 * (a - b) / b


def build():
    fz = check_freeze()
    cats, notes = lineage()
    lin = {c["category_id"]: c for c in cats}
    ucomp, lines = urban_components()
    U, Fv = view_defs()
    ures = CS.assemble(U, [{k: c[k] for k in ("component_id", "kind", "quantity", "unit", "state")} for c in ucomp])
    fcomp = []
    for c in cats:
        if c["category_id"] in ("TOTAL_RC", "TOTAL_STEEL"):
            continue
        for x in c["subcomponents"]:
            if x.get("kind") == "LEAF_ROW":
                fcomp.append({"component_id": f"FQS:{x['sheet']}!{x['cell']}", "kind": x["physical_kind"],
                              "quantity": x["value"], "unit": "m3", "state": "REFERENCE"})
    fres = CS.assemble(Fv, fcomp)
    joints = {f: next((c["quantity"] for c in ucomp if c["kind"] == f"JOINT_{f}"), None) for f in ("GF", "1F", "2F")}
    slabreg = json.loads(SLAB_REG.read_text(encoding="utf-8"))["sheets"]
    # beam plan area per floor in the freelancer rows (width x length), for slab-area normalisation
    beam_plan = defaultdict(float)
    for x in lin["BEAMS"]["subcomponents"]:
        if x.get("kind") == "LEAF_ROW" and x["B"] is not None and x["D"] is not None:
            beam_plan[x["physical_kind"].split("_")[-1]] += float(x["B"]) * float(x["D"]) * float(x.get("count") or 1)
    O = ORACLES
    rows = []

    def row(section, cat, sub, basis, fl, uc, cn, ur, ustate, comp, cause, review=True, refs=(), norm=None):
        diff = {}
        for k, v in (("FREELANCER", fl), ("UC4N", uc), ("CHRISTIANNP", cn)):
            st = comp.get(k, CS.NOT_COMPARABLE)
            cv = CS.compare_values(ur, v, st)
            diff[k] = {"status": st, "difference": r3(cv["difference"]), "difference_percent": r3(
                cv["difference_percent"])}
        rows.append({"section": section, "category": cat, "subcategory": sub, "measurement_basis": basis,
                     "FREELANCER_QS": r3(fl), "UC4N": r3(uc), "CHRISTIANNP": r3(cn), "URBAN": r3(ur),
                     "URBAN_STATUS": ustate, "normalized": norm or {}, "vs": diff,
                     "likely_cause": cause, "review_required": review, "source_references": list(refs)})

    def uq(cat):
        c = ures["categories"][cat]
        valued = [x for x in c["components"] if x not in c["components_without_value"]]
        return c["quantity"] if valued else None          # all components blocked -> no value, never 0

    def ust(cat):
        c = ures["categories"][cat]
        if c["components_without_value"]:
            return "LOWER_BOUND (blocked components: " + ", ".join(c["components_without_value"]) + ")"
        st = set(c["by_state"])
        return "/".join(sorted(st)) or "NO_VALUE"

    def fq(cat):
        return fres["categories"][cat]["quantity"]
    D, N = CS.DIRECT, CS.NORMALIZED
    gross = {k: v["gross_outline_area_m2"] for k, v in slabreg.items()}
    net = {k: v["net_plate_area_m2"] for k, v in slabreg.items()}
    row("A", "GROSS_STRUCTURAL_FOOTPRINT", "built-up (sum of storeys)", "GROSS_SLAB_OUTLINE_SUM", None,
        O["UC4N"]["built_up_m2"], O["CHRISTIANNP"]["built_up_m2"], sum(gross.values()), "VERIFIED (B2A.1)",
        {"UC4N": D, "CHRISTIANNP": D},
        "oracle built-up taken as gross slab outline (declared basis); freelancer workbook has no built-up figure",
        refs=["B2A1 SLAB_REGION_REGISTER gross_outline_area_m2"])
    for f in ("GF", "1F", "2F"):
        fl_area = sum(float(x["B"]) for x in lin["SLABS"]["subcomponents"] if x.get("kind") == "LEAF_ROW" and
                      x["physical_kind"] == f"SLAB_NET_{f}")
        norm_fl = fl_area + beam_plan.get(f, 0.0)
        row("B", "NET_SLAB_AREA", f, "NET_PLATE_AREA (openings out)", fl_area, O["UC4N"]["net_slab_m2"][f],
            O["CHRISTIANNP"]["net_slab_m2"][f], net[f], "VERIFIED (B2A.1)",
            {"UC4N": D, "CHRISTIANNP": D, "FREELANCER": CS.NOT_COMPARABLE},
            "freelancer slab area is between beams (beam footprint allocated to gross beams); normalised = slab + "
            "beam plan area; Urban subtracts stair / void openings - check oracle opening treatment",
            refs=[f"B2A1 net_plate_area_m2 {f}", f"freelancer 'بلاطات ' {f}"],
            norm={"FREELANCER_SLAB_PLUS_BEAM_PLAN_m2": r3(norm_fl),
                  "status": N, "difference_vs_urban": r3(net[f] - norm_fl), "difference_percent": r3(pct(net[f],
                                                                                                      norm_fl))})
    row("C", "FOUNDATIONS", "isolated / combined footings", "PAD_LxWxH_SCHEDULE", fq("FOOTINGS"), O["UC4N"]["FOOTING"],
        O["CHRISTIANNP"]["FOOTING"], uq("FOOTINGS"), ust("FOOTINGS"), {"FREELANCER": D, "UC4N": D, "CHRISTIANNP": D},
        "Urban is a lower bound: F/F10 conflict outline and one more footing blocked; freelancer resolves F/F10 "
        "(F x4 + F10 x1)", refs=["V3b C-FTG", "freelancer القواعد!I13:I28"])
    row("D", "STRAP_BEAMS", "foundation straps", "STRAP_BAND", fq("STRAP_BEAMS"), O["UC4N"]["STRAP_BEAM"],
        O["CHRISTIANNP"]["STRAP_BEAM"], uq("STRAP_BEAMS"), ust("STRAP_BEAMS"),
        {"FREELANCER": D, "UC4N": D, "CHRISTIANNP": D},
        "Urban = freelancer; both oracles ~1.8x - likely strap length basis (oracle may measure into the footings) "
        "or extra strap population - verify with the oracle handles", refs=["V3b C-STR", "freelancer STB1-3"])
    row("E", "GROUND_BEAMS", "interior + exterior (+ perimeter in QS)", "GB_SPAN_x_SECTION", fq("GROUND_BEAMS"),
        O["UC4N"]["GROUND_BEAM"], O["CHRISTIANNP"]["GROUND_BEAM"], uq("GROUND_BEAMS"), ust("GROUND_BEAMS"),
        {"FREELANCER": CS.NOT_COMPARABLE, "UC4N": D, "CHRISTIANNP": D},
        "Urban exterior ground beams BLOCKED (depth 'follow arch'); freelancer 30x60 on all runs incl. a perimeter "
        "beam filed under foundations", refs=["V3b C-GB-INT / C-GB-EXT", "freelancer الشناجات + الشناج المحيط"])
    row("F", "GROUND_SLAB", "ground slab", "AREA_x_T", fq("GROUND_SLAB"), O["UC4N"]["GROUND_SLAB"],
        O["CHRISTIANNP"]["GROUND_SLAB"], uq("GROUND_SLAB"), ust("GROUND_SLAB"),
        {"FREELANCER": D, "UC4N": D, "CHRISTIANNP": D},
        "Urban computed only two closed zones (115 m2 gross / T=10); the others take ~295-320 m2 - ground slab "
        "extent is the largest real gap (S2 STR-GRO-001)", refs=["V3b C-GSLAB-ZONE-1/2", "freelancer الشناجات!B11"])
    for f, key in (("FOUNDATION", "NECKS"), ("GF", "COLUMNS_GF"), ("1F", "COLUMNS_1F"), ("2F", "COLUMNS_2F")):
        row("G", "COLUMNS_NECKS", f, "COLUMN volume (height basis differs per engine)", fq(key),
            O["UC4N"]["COLUMN"][f], O["CHRISTIANNP"]["COLUMN"][f], uq(key), ust(key),
            {"FREELANCER": D if f != "FOUNDATION" else CS.NOT_COMPARABLE, "UC4N": CS.NOT_COMPARABLE,
             "CHRISTIANNP": CS.NOT_COMPARABLE},
            "freelancer: section x clear height below beams (3.75 / 3.45), necks 1.5 m assumed; Urban: clear height, "
            "joint separate, some occurrences blocked; oracle height basis not declared - not comparable",
            refs=[f"V3b C-COL-{f}" if f != "FOUNDATION" else "V3b C-NECK", "freelancer 'الحوائط + الأعمدة'"])
    gtot = {}
    for g in GROUPS[:3]:
        f = g["group_id"].split("_")[1]
        u = sum((uq(c) or 0) for c in g["categories"]) + (joints[f] or 0)      # blocked parts excluded (lower bound)
        fl = sum((fq(c) or 0) for c in g["categories"])
        cn = (O["CHRISTIANNP"]["COLUMN"][f] + O["CHRISTIANNP"]["BEAM_DOWNSTAND"][f] + O["CHRISTIANNP"]["SLAB"][f])
        gtot[f] = (u, fl, cn)
        row("H", "FRAME_NORMALISED", f, "COLUMNS + JOINTS + BEAMS + SLABS of the floor (allocation-free)", fl, None, cn,
            u, "LOWER_BOUND (columns / beams partial)", {"FREELANCER": N, "CHRISTIANNP": CS.NOT_COMPARABLE},
            "union removes the beam gross / downstand / slab allocation difference; oracle columns height basis "
            "undeclared so the oracle union stays not comparable",
            refs=["comparison_scope normalisation group " + g["group_id"]])
    for f in ("GF", "1F", "2F"):
        row("H", "BEAMS", f, "per engine (see basis)", fq(f"BEAMS_{f}"), None, O["CHRISTIANNP"]["BEAM_DOWNSTAND"][f],
            uq(f"BEAMS_{f}"), ust(f"BEAMS_{f}"), {"FREELANCER": CS.NOT_COMPARABLE, "CHRISTIANNP": D},
            "freelancer gross depth vs Urban downstand: NOT_COMPARABLE alone - see FRAME_NORMALISED; oracle "
            "downstand is the same convention as Urban", refs=[f"V3b C-BEAM-{f}", "freelancer كمرات"])
        row("I", "SLABS", f, "per engine (see basis)", fq(f"SLABS_{f}"), None, O["CHRISTIANNP"]["SLAB"][f],
            uq(f"SLABS_{f}"), ust(f"SLABS_{f}"), {"FREELANCER": CS.NOT_COMPARABLE, "CHRISTIANNP": D},
            "freelancer net of beams vs Urban full depth on the net plate; oracle slab / oracle net area implies "
            f"{O['CHRISTIANNP']['SLAB'][f] / O['CHRISTIANNP']['net_slab_m2'][f]:.3f} m average thickness vs printed "
            f"{slabreg[f]['sheet_thickness_tags_cm'][0] / 100:.2f} m - oracle slab basis to verify",
            refs=[f"V3b C-SLAB-{f}", "B2A1 SLAB_REGION_REGISTER"])
    ub = sum(uq(f"BEAMS_{f}") or 0 for f in ("GF", "1F", "2F"))
    us = sum(uq(f"SLABS_{f}") or 0 for f in ("GF", "1F", "2F"))
    row("H", "BEAMS", "all floors", "BEAM_DOWNSTAND_ONLY", sum(fq(f"BEAMS_{f}") for f in ("GF", "1F", "2F")),
        O["UC4N"]["BEAM_DOWNSTAND_TOTAL"], sum(O["CHRISTIANNP"]["BEAM_DOWNSTAND"].values()), ub,
        "LOWER_BOUND (residue blocked)", {"FREELANCER": CS.NOT_COMPARABLE, "UC4N": D, "CHRISTIANNP": D},
        "same convention as both oracles; freelancer gross depth only comparable inside FRAME_NORMALISED")
    row("I", "SLABS", "all floors", "SLAB_FULL_DEPTH", sum(fq(f"SLABS_{f}") for f in ("GF", "1F", "2F")),
        O["UC4N"]["SLAB_TOTAL"], sum(O["CHRISTIANNP"]["SLAB"].values()), us, "VERIFIED",
        {"FREELANCER": CS.NOT_COMPARABLE, "UC4N": D, "CHRISTIANNP": D},
        "both oracles ~15 m3 above Urban on the same declared convention: implied thickness / area basis differs - "
        "real disagreement to review")
    row("J", "STAIRS", "all stairs", "STAIR", fq("STAIRS"), O["UC4N"]["STAIR"], O["CHRISTIANNP"]["STAIR"],
        uq("STAIRS"), ust("STAIRS"), {"FREELANCER": CS.NOT_COMPARABLE, "UC4N": CS.NOT_COMPARABLE,
                                      "CHRISTIANNP": CS.NOT_COMPARABLE},
        "Urban stairs BLOCKED (riser / waist inputs); freelancer 0.28 m slab-equivalent; oracle basis undeclared")
    row("K", "DOME", "dome shells", "DOME_SHELL", fq("DOME"), None, None, uq("DOME"), ust("DOME"),
        {"FREELANCER": D}, "freelancer 3 x (1.0 x 22.6 x 0.1); Urban shell 2 pi R h t per dome; rings blocked",
        refs=["V3b C-DOME-SHELL-*", "freelancer السلم!B13"])
    row("L", "POOL", "pool", "POOL", fq("POOL"), None, None, uq("POOL"), ust("POOL"), {"FREELANCER": CS.NOT_COMPARABLE},
        "freelancer adds pump-room walls + room floor and a 10.8 m2 x 0.4 base; Urban walls 10.10 x 0.20 x 1.15 + "
        "base 3.50 x 1.95 x 0.40 - scope and depth differ", refs=["V3b C-POOL", "freelancer حمام السباحة"])
    row("M", "SPECIAL_RC", "structural / lift walls", "WALL", fq("STRUCTURAL_WALLS"), None, None,
        uq("STRUCTURAL_WALLS"), ust("STRUCTURAL_WALLS"), {"FREELANCER": CS.NOT_COMPARABLE},
        "freelancer counts 20 cm lift walls (8 m x 1.5 m); Urban structural walls BLOCKED - population to confirm")
    row("M", "SPECIAL_RC", "boundary wall beam", "BOUNDARY", fq("BOUNDARY"), None, None, uq("BOUNDARY"),
        ust("BOUNDARY"), {"FREELANCER": CS.NOT_COMPARABLE}, "Urban boundary wall BLOCKED; freelancer 42 m x 30x60")
    u_tot = sum(c["quantity"] for c in ucomp if c["quantity"] is not None and c["kind"] != "PLAIN_CONCRETE")
    row("N", "TOTAL_CONCRETE", "reinforced concrete", "ALL RC (bases differ)", lin["TOTAL_RC"]["summary_value"],
        O["UC4N"]["total_rc_m3"], O["CHRISTIANNP"]["total_rc_m3"], u_tot, "LOWER_BOUND (blocked populations)",
        {}, "totals combine different scopes and blocked Urban populations - shown, never compared as %",
        refs=["freelancer ورقة1!F24"])
    # ---------------------------------------------------------------- rough rebar
    profile = json.loads(PROFILE.read_text(encoding="utf-8"))
    act, s31 = urban_rebar(lines)
    occs = [{"occurrence_id": c["component_id"], "element_class": c["element_class"], "concrete_m3": c["quantity"],
             "concrete_state": c["state"] if c["quantity"] is not None else "BLOCKED"} for c in ucomp]
    rough = RR.rough_check(occs, profile, act)
    fl_steel = {c["category_id"]: c.get("steel_value_t") for c in cats}
    rebar_rows = []
    for cat, r in rough["categories"].items():
        a = act.get(cat, {})
        rebar_rows.append({"category": cat, "urban_released_concrete_m3": r3(r["released_concrete_m3"]),
                           "urban_modelled_concrete_m3": r3(r["modelled_concrete_m3"]),
                           "urban_blocked_concrete_m3_known": r3(r["blocked_concrete_m3"]),
                           "rough_ratio_kg_per_m3": r["ratio_kg_per_m3"],
                           "rough_reference_kg_released": r3(r["rough_reference_kg_released_basis"]),
                           "rough_reference_kg_modelled": r3(r["rough_reference_kg_modelled_basis"]),
                           "actual_released_kg": r3(a.get("released_kg")),
                           "actual_provisional_kg": r3(a.get("provisional_kg")),
                           "actual_blocked_components": a.get("blocked_components", []),
                           "actual_complete": a.get("complete", False),
                           "comparison_state": r.get("comparison_state", "NO_ACTUAL"),
                           "freelancer_steel_t": fl_steel.get(cat),
                           "freelancer_concrete_m3": lin.get(cat, {}).get("summary_value"),
                           "note": "rough = Urban concrete x owner ratio; the ratio was calibrated on the freelancer "
                                   "convention (beam gross depth), so on Urban physical volumes BEAMS read low and "
                                   "SLABS high - compare BEAMS+SLABS together"})
    oracle_rebar = [{"oracle": k, "net_rebar_t": v["net_rebar_t"], "state": v["rebar_state"],
                     "answer": CO.answer("COUNT_TAGS", v["net_rebar_t"], state=CO.INCOMPLETE, oracle=k)["state"]}
                    for k, v in O.items()]
    # ---------------------------------------------------------------- oracle register
    soc_rows = []
    facts = [("FOOTINGS_CONCRETE_M3", uq("FOOTINGS"), "FOOTING", True, ust("FOOTINGS")),
             ("STRAP_BEAMS_CONCRETE_M3", uq("STRAP_BEAMS"), "STRAP_BEAM", True, ust("STRAP_BEAMS")),
             ("GROUND_BEAMS_CONCRETE_M3", uq("GROUND_BEAMS"), "GROUND_BEAM", True, ust("GROUND_BEAMS")),
             ("GROUND_SLAB_CONCRETE_M3", uq("GROUND_SLAB"), "GROUND_SLAB", True, ust("GROUND_SLAB")),
             ("BEAM_DOWNSTAND_TOTAL_M3", ub, None, True, "LOWER_BOUND"),
             ("SLAB_TOTAL_M3", us, None, True, "VERIFIED"),
             ("BUILT_UP_GROSS_M2", sum(gross.values()), "built_up_m2", True, "VERIFIED")]
    for f in ("GF", "1F", "2F"):
        facts.append((f"NET_SLAB_AREA_{f}_M2", net[f], ("net_slab_m2", f), True, "VERIFIED"))
        facts.append((f"COLUMNS_{f}_M3", uq(f"COLUMNS_{f}"), ("COLUMN", f), False, ust(f"COLUMNS_{f}")))
    for name, o in O.items():
        for fact, uval, okey, bmatch, ustate in facts:
            if fact == "BEAM_DOWNSTAND_TOTAL_M3":
                ov = o.get("BEAM_DOWNSTAND_TOTAL") or sum((o.get("BEAM_DOWNSTAND") or {}).values()) or None
            elif fact == "SLAB_TOTAL_M3":
                ov = o.get("SLAB_TOTAL") or sum((o.get("SLAB") or {}).values()) or None
            elif isinstance(okey, tuple):
                ov = o[okey[0]][okey[1]]
            else:
                ov = o.get(okey)
            ans = CO.answer("BOUNDING_BOX", ov, refs=[o["source"]], oracle=name)
            soc_rows.append(SOC.compare({"element_id": "ALSENAN", "fact_type": fact, "value": uval,
                                         "source_refs": ["V3b BOQ lines / B2A1 slab register"], "status": ustate},
                                        ans, basis_match=bmatch, rel_tol=0.02, close_rel_tol=0.05))
        ans = CO.answer("COUNT_TAGS", o["net_rebar_t"], state=CO.OK, entity_limit=1, returned=1, oracle=name)
        soc_rows.append(SOC.compare({"element_id": "ALSENAN", "fact_type": "NET_REBAR_T", "value": None,
                                     "status": "BBS_INCOMPLETE"}, ans, basis_match=False))
    soc = SOC.register(soc_rows)
    # ---------------------------------------------------------------- populations
    pops = population_evidence(ucomp)
    pop = SP.discover(pops)
    # ---------------------------------------------------------------- assumptions
    assumptions = assumption_register(notes)
    return {"freeze": fz, "lineage": cats, "notes": notes, "urban_components": ucomp, "urban_view": ures,
            "freelancer_view": fres, "rows": rows, "rough": rough, "rebar_rows": rebar_rows,
            "oracle_rebar": oracle_rebar, "soc": soc, "populations": pop, "assumptions": assumptions,
            "profile": profile, "s31_headline": s31, "beam_plan_area": dict(beam_plan), "frame": gtot}


def population_evidence(ucomp):
    s1 = ROOT / "research" / "alsenan_structural_census_s1"
    spc = json.loads((s1 / "SPECIAL_STRUCTURAL_OCCURRENCE_REGISTER.json").read_text(encoding="utf-8"))["rows"]
    kinds = {r["kind"] for r in spc}
    P, NP, B, U = SP.PRESENT, SP.NOT_PRESENT, SP.BLOCKED, SP.UNKNOWN
    fp = "S1 FOOTING_OCCURRENCE_REGISTER (foundation plan census, all outlines)"
    ev = {
        "RAFT": {"state": NP, "refs": [fp], "why": "all foundation outlines are isolated / combined pads"},
        "ISOLATED_FOOTING": {"state": P, "refs": [fp], "quantity_state": "LOWER_BOUND (F/F10 blocked)"},
        "COMBINED_FOOTING": {"state": P, "refs": [fp + " combined_footings_with_several_columns"]},
        "STRIP_FOOTING": {"state": NP, "refs": [fp]},
        "PILE_CAP": {"state": NP, "refs": [fp, "ST7757 p.9 note 1 (shallow foundations)"]},
        "NECK_PEDESTAL": {"state": B, "refs": ["V3b C-NECK"], "why": "founding / ground-beam level not printed"},
        "STRAP_BEAM": {"state": P, "refs": ["V3b C-STR"]},
        "GROUND_BEAM": {"state": P, "refs": ["V3b C-GB-INT"], "quantity_state": "exterior GB BLOCKED"},
        "GROUND_SLAB": {"state": P, "refs": ["V3b C-GSLAB-ZONE-1/2"], "quantity_state": "LOWER_BOUND (extent)"},
        "COLUMN": {"state": P, "refs": ["S1 COLUMN_OCCURRENCE_REGISTER (95)"]},
        "STRUCTURAL_WALL": {"state": B, "refs": ["V3b C-SWALL"], "why": "lift walls appear only in the freelancer "
                                                                      "reference - source to confirm",
                            "flags": ["FREELANCER_HINT_LIFT_WALLS"]},
        "PLANTED_COLUMN": {"state": P if "PLANTED_COLUMN" in kinds else U, "refs": ["S1 SPECIAL register"]},
        "TURNED_COLUMN": {"state": P if "TURN_COLUMN" in kinds else U, "refs": ["S1 SPECIAL register"]},
        "SIMPLE_BEAM": {"state": P, "refs": ["S1 BEAM_OCCURRENCE_REGISTER"]},
        "CONTINUOUS_BEAM": {"state": P, "refs": ["S1 CONTINUOUS_BEAM_OCCURRENCE_REGISTER"]},
        "RING_BEAM": {"state": B, "refs": ["V3b C-DOME-RING-*"], "why": "dome ring detail N.I.S."},
        "LINTEL": {"state": P, "refs": ["V3b C-LINT-*"]},
        "BEAM_OPENING": {"state": U, "refs": [], "why": "service widenings not located on plans (S2 STR-BEA-007)"},
        "SOLID_SLAB": {"state": P, "refs": ["S1 SLAB_PANEL_REGISTER (70 panels)"]},
        "FLAT_SLAB": {"state": NP, "refs": ["S1 SLAB_PANEL_REGISTER: all panels beam-supported"]},
        "RIBBED_HORDI_SLAB": {"state": NP, "refs": ["S1 SLAB_PANEL_REGISTER: solid panels only"]},
        "CANTILEVER": {"state": P, "refs": ["S1 BEAM_OCCURRENCE_REGISTER family CANTILEVER"]},
        "STAIR": {"state": B, "refs": ["V3b C-STAIR"], "why": "riser / waist inputs incomplete"},
        "LANDING": {"state": B, "refs": ["V3b C-STAIR"]},
        "LIFT_PIT": {"state": B, "refs": ["S2 STR-LIF-001"], "why": "pit dimensions not printed"},
        "LIFT_WALL": {"state": B, "refs": ["V3b C-SWALL"], "flags": ["FREELANCER_HINT_LIFT_WALLS"]},
        "LIFT_TIE_BEAM": {"state": B, "refs": ["S1 STRUCTURAL_LEVEL_REGISTER lift_tie_beam_rule_check (GF)"],
                          "why": "required by p.8 note 19, not drawn"},
        "POOL": {"state": P, "refs": ["V3b C-POOL"]},
        "WATER_TANK": {"state": U, "refs": []},
        "DOME": {"state": P, "refs": ["V3b C-DOME-SHELL-*"]},
        "BOUNDARY_WALL": {"state": B, "refs": ["V3b C-BWALL"]},
        "BOUNDARY_RC_COLUMN": {"state": U, "refs": []},
        "BOUNDARY_GROUND_BEAM": {"state": B, "refs": ["S1 BEAM_DEFINITION_REGISTER B.W (no plan occurrence)"],
                                 "flags": ["FREELANCER_HINT_BOUNDARY_BEAM_42M"]},
        "PARAPET": {"state": B, "refs": ["S2 STR-PAR-001"], "why": "parapet geometry missing"},
        "PARAPET_RC_STIFFENER_COLUMN": {"state": U, "refs": [], "flags": ["PARAPET_STIFFENER_SPACING_REQUIRED"],
                                        "why": "not printed; Urban 4 m candidate is a provisional scenario only"},
        "PARAPET_TOP_RING_BEAM": {"state": U, "refs": [], "flags": ["PARAPET_TOP_RING_BEAM_DETAIL_REQUIRED"]},
        "EQUIPMENT_BASE": {"state": U, "refs": []},
        "PLANTER_STRUCTURAL_ELEMENT": {"state": U, "refs": []},
        "OTHER_SPECIAL_RC_DETAIL": {"state": U, "refs": ["ST7757 p.15 casement beam (candidate)"]},
    }
    return ev


def assumption_register(notes):
    a = [
        ("SOURCE_ROLES", "Freelancer workbook = FREELANCER_QS_REFERENCE; U-C4N / christiannp = EXTERNAL_ORACLE; "
                         "neither creates or resolves Urban quantities", SR.FREELANCER_QS_REFERENCE),
        ("ORACLE_VALUES", "oracle values are the owner's recorded figures (S3.1 brief §16); no oracle run or handle "
                          "list is held in the repository", SR.EXTERNAL_ORACLE),
        ("ORACLE_BASIS", "oracle 'beam downstands' and 'slabs' are taken on the Urban convention (downstand only / "
                         "full slab depth) as stated in the brief; oracle column height basis is undeclared",
         SR.EXTERNAL_ORACLE),
        ("URBAN_CONCRETE", "Urban concrete = frozen V3b candidate BOQ lines (not FINAL); COMPUTED -> VERIFIED, "
                           "PARTIAL -> LOWER_BOUND, owner method -> PROVISIONAL; ground slab zones -> LOWER_BOUND",
         SR.PRODUCTION_SOURCE),
        ("URBAN_COLUMN_STEEL", "column steel = S3.1 (D2/162, transitions blocked); other categories = V3b net design "
                               "weights (D2/162)", SR.PRODUCTION_SOURCE),
        ("ROUGH_MAPPING", "BEAM_COLUMN_JOINT -> WALLS_AND_COLUMNS (column bars and ties run through the joint); "
                          "LINTEL not mapped -> ROUGH_RATIO_NOT_CONFIGURED; plain concrete excluded",
         SR.URBAN_OWNER_RULE),
        ("ROUGH_CALIBRATION", "the owner ratios were calibrated on the freelancer allocation (beam gross depth, slab "
                              "between beams); on Urban physical volumes BEAMS read low and SLABS high",
         SR.URBAN_OWNER_RULE),
        ("NORMALISATION", "FRAME_<floor> = columns + joints + beams + slabs is allocation-free; BEAMS or SLABS alone "
                          "are NOT_COMPARABLE across gross / downstand conventions", "COMPARISON"),
        ("NET_SLAB_NORMALISATION", "freelancer slab area + freelancer beam plan area (width x length) approximates a "
                                   "net plate area", "COMPARISON"),
        ("PARAPET_CANDIDATE", "4 m stiffener spacing is URBAN_STANDARD_CANDIDATE - scenario only, never verified",
         SR.URBAN_OWNER_RULE)]
    return [{"assumption_id": k, "statement": s, "role": r} for k, s, r in a]


# ================================================================================================ outputs
def dumps(o):
    return json.dumps(o, indent=1, sort_keys=True, ensure_ascii=False, default=str) + "\n"


def files_from(R):
    lin = R["lineage"]
    return {
        "FREELANCER_CATEGORY_LINEAGE.json": {"workbook_sha256": FREELANCER_SHA, "role": SR.FREELANCER_QS_REFERENCE,
                                             "categories": lin, "lineage_notes": R["notes"]},
        "FREELANCER_STRUCTURAL_SCOPE_MAP.json": {
            "role": SR.FREELANCER_QS_REFERENCE, "use": "COMPARISON_ONLY",
            "categories": [{"category_id": c["category_id"], "arabic_name": c["arabic_name"],
                            "summary_value": c["summary_value"], "unit": c["unit"],
                            "physical_scope": c["quantity_by_physical_kind"], "possible_overlap": c["possible_overlap"],
                            "notes": c["notes"]} for c in lin],
            "freelancer_view": R["freelancer_view"]},
        "URBAN_ROUGH_REBAR_PROFILE_V1.json": {
            "profile": R["profile"], "engine_profile_path": "engine/profiles/URBAN_ROUGH_REBAR_PROFILE_V1.json",
            "calibration_sample": {"role": SR.FREELANCER_QS_REFERENCE, "workbook_sha256": FREELANCER_SHA,
                                   "rows": [{"category": c["category_id"], "concrete_m3": c["summary_value"],
                                             "steel_t": c["steel_value_t"],
                                             "kg_per_m3": None if not c["steel_value_t"] or not c["summary_value"] else
                                             round(c["steel_value_t"] * 1000 / c["summary_value"], 1)}
                                            for c in lin if c["category_id"] not in ("PLAIN_CONCRETE", "TOTAL_RC",
                                                                                     "TOTAL_STEEL")],
                                   "rule": "calibration values live only here and in the profile provenance - "
                                           "never in BBS tests or production"}},
        "MULTI_ENGINE_CONCRETE_COMPARISON.json": {"rows": [r for r in R["rows"] if r["section"] not in ("A", "B")],
                                                  "urban_view": R["urban_view"], "frame_groups": GROUPS},
        "MULTI_ENGINE_AREA_COMPARISON.json": {"rows": [r for r in R["rows"] if r["section"] in ("A", "B")],
                                              "freelancer_beam_plan_area_m2": {k: r3(v) for k, v in
                                                                               R["beam_plan_area"].items()}},
        "MULTI_ENGINE_REBAR_COMPARISON.json": {"rough": R["rough"], "rows": R["rebar_rows"],
                                               "oracle_net_rebar": R["oracle_rebar"],
                                               "s3_1_column_headline": R["s31_headline"], "footer": RR.FOOTER},
        "SOURCE_ORACLE_COMPARISON_REGISTER.json": R["soc"],
        "STRUCTURAL_POPULATION_REGISTER.json": R["populations"],
        "ASSUMPTION_AND_BASIS_REGISTER.json": {"rows": R["assumptions"]},
    }


def main(write=True):
    R = build()
    files = files_from(R)
    blobs = {k: dumps(v).encode() for k, v in files.items()}
    if write:
        for k, b in blobs.items():
            (HERE / k).write_bytes(b)
        idx = {"round": "S3.1", "use": "COMPARISON_ONLY (post-freeze)", "production_freeze": "URBAN_PRODUCTION_FREEZE.json",
               "production_freeze_sha256": sha((HERE / "URBAN_PRODUCTION_FREEZE.json").read_bytes()),
               "freelancer_workbook_sha256": FREELANCER_SHA,
               "outputs": {k: sha(b) for k, b in sorted(blobs.items())},
               "production_registers_modified": False}
        (HERE / "INDEX.json").write_text(dumps(idx), encoding="utf-8")
    return R, files, blobs


if __name__ == "__main__":
    R, files, b1 = main()
    if "--twice" in sys.argv:
        _, _, b2 = main(write=False)
        print("built twice identical:", all(b1[k] == b2[k] for k in b1))
    for r in R["rows"]:
        print(r["section"], r["category"], r["subcategory"], "| FL", r["FREELANCER_QS"], "UC", r["UC4N"], "CN",
              r["CHRISTIANNP"], "UR", r["URBAN"], "|", {k: (v["status"][:6], v["difference_percent"]) for k, v in
                                                       r["vs"].items()}, r["normalized"] or "")
