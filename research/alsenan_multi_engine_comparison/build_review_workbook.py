"""ALSENAN_MULTI_ENGINE_STRUCTURAL_COMPARISON.xlsx - review workbook (comparison only, post-freeze).

Reads the comparison registers (sha256 checked against INDEX.json). Differences / percentages / rough steel are
live formulas; a NOT_COMPARABLE row never gets a percentage (the formula returns blank).
"""
import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "review" / "ALSENAN_MULTI_ENGINE_STRUCTURAL_COMPARISON.xlsx"


def load():
    idx = json.loads((HERE / "INDEX.json").read_text(encoding="utf-8"))
    R = {}
    for k, h in idx["outputs"].items():
        b = (HERE / k).read_bytes()
        if hashlib.sha256(b).hexdigest() != h:
            raise SystemExit(f"{k} does not match INDEX")
        R[k.replace(".json", "")] = json.loads(b)
    return idx, R


def build(idx, R):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    wb = Workbook()
    hf, hfill = Font(name="Arial", bold=True, color="FFFFFF", size=10), PatternFill("solid", fgColor="1F3864")
    body, bold = Font(name="Arial", size=9), Font(name="Arial", size=10, bold=True)

    def sheet(name, headers, rows, widths=None):
        ws = wb.create_sheet(name)
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font, c.fill = hf, hfill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, r in enumerate(rows, 2):
            for j, v in enumerate(r, 1):
                c = ws.cell(row=i, column=j, value=v)
                if isinstance(v, str) and v.startswith("="):
                    c.data_type = "s"          # a quoted source formula is text, never a live formula here
                c.font = body
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{L(len(headers))}{max(2, len(rows) + 1)}"
        for j, h in enumerate(headers, 1):
            ws.column_dimensions[L(j)].width = (widths or {}).get(h, max(11, min(40, len(h) + 3)))
        ws.row_dimensions[1].height = 45
        return ws

    # ---------------------------------------------------------------- 00_Summary
    ws = wb.active
    ws.title = "00_Summary"
    lines = [
        ("ALSENAN - S3.1 MULTI-ENGINE STRUCTURAL COMPARISON (comparison only)", True),
        ("Production was frozen first (URBAN_PRODUCTION_FREEZE.json, sha256 "
         f"{idx['production_freeze_sha256'][:16]}...). Nothing in this workbook changes an Urban quantity.", False),
        ("Source roles: Urban = PRODUCTION_SOURCE | U-C4N, christiannp = EXTERNAL_ORACLE (values recorded from the "
         "owner's brief) | freelancer workbook = FREELANCER_QS_REFERENCE (sha256 "
         f"{idx['freelancer_workbook_sha256'][:16]}...) | rough ratios = URBAN_OWNER_RULE, sanity only.", False),
        ("Difference = Urban - reference; % = difference / reference. A percentage appears only when the basis "
         "matches (DIRECTLY_COMPARABLE) or after a declared normalisation (COMPARABLE_AFTER_NORMALIZATION).", False),
        ("Rough kg/m3 values are Urban estimating/sanity heuristics and are not a substitute for bar-by-bar "
         "reinforcement takeoff.", True),
        ("", False), ("Sheets: 01_Comparison (sections A-N) | 02_Rebar (section O: actual vs rough) | "
                      "03_Freelancer_Lineage | 04_Scope_Map | 05_Oracle_Register | 06_Populations | 07_Assumptions",
                      False)]
    for i, (t, b) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = bold if b else body
    ws.column_dimensions["A"].width = 160
    # ---------------------------------------------------------------- 01_Comparison
    H = ["SECTION", "CATEGORY", "SUBCATEGORY", "MEASUREMENT BASIS", "FREELANCER_QS", "UC4N", "CHRISTIANNP", "URBAN",
         "URBAN STATUS", "NORMALIZED FREELANCER", "STATUS vs FREELANCER", "DIFFERENCE vs FREELANCER",
         "DIFFERENCE % vs FREELANCER", "STATUS vs UC4N", "DIFFERENCE vs UC4N", "DIFFERENCE % vs UC4N",
         "STATUS vs CHRISTIANNP", "DIFFERENCE vs CHRISTIANNP", "DIFFERENCE % vs CHRISTIANNP", "LIKELY CAUSE",
         "REVIEW REQUIRED", "SOURCE REFERENCES"]
    rows = R["MULTI_ENGINE_AREA_COMPARISON"]["rows"] + R["MULTI_ENGINE_CONCRETE_COMPARISON"]["rows"]
    rows.sort(key=lambda r: r["section"])
    data = []
    for r in rows:
        nf = r["normalized"].get("FREELANCER_SLAB_PLUS_BEAM_PLAN_m2")
        stf = r["normalized"].get("status") if nf is not None else r["vs"]["FREELANCER"]["status"]
        data.append([r["section"], r["category"], r["subcategory"], r["measurement_basis"], r["FREELANCER_QS"],
                     r["UC4N"], r["CHRISTIANNP"], r["URBAN"], r["URBAN_STATUS"], nf, stf, None, None,
                     r["vs"]["UC4N"]["status"], None, None, r["vs"]["CHRISTIANNP"]["status"], None, None,
                     r["likely_cause"], "YES" if r["review_required"] else "NO", "; ".join(r["source_references"])])
    ws = sheet("01_Comparison", H, data, widths={"LIKELY CAUSE": 70, "MEASUREMENT BASIS": 34, "URBAN STATUS": 30,
                                                 "SOURCE REFERENCES": 40})
    for i in range(2, len(data) + 2):
        ref = f'IF(J{i}<>"",J{i},E{i})'
        ws[f"L{i}"] = f'=IF(OR(K{i}="NOT_COMPARABLE",H{i}="",E{i}=""),"",H{i}-{ref})'
        ws[f"M{i}"] = f'=IF(OR(L{i}="",{ref}=0),"",L{i}/{ref})'
        ws[f"O{i}"] = f'=IF(OR(N{i}="NOT_COMPARABLE",H{i}="",F{i}=""),"",H{i}-F{i})'
        ws[f"P{i}"] = f'=IF(OR(O{i}="",F{i}=0),"",O{i}/F{i})'
        ws[f"R{i}"] = f'=IF(OR(Q{i}="NOT_COMPARABLE",H{i}="",G{i}=""),"",H{i}-G{i})'
        ws[f"S{i}"] = f'=IF(OR(R{i}="",G{i}=0),"",R{i}/G{i})'
        for c in "MPS":
            ws[f"{c}{i}"].number_format = "0.0%"
    # ---------------------------------------------------------------- 02_Rebar
    rr = R["MULTI_ENGINE_REBAR_COMPARISON"]
    H = ["CATEGORY", "URBAN CONCRETE M3 (released)", "URBAN CONCRETE M3 (modelled)", "ROUGH KG/M3 (owner)",
         "ROUGH REBAR KG (modelled) = concrete x ratio", "ACTUAL RELEASED KG", "ACTUAL PROVISIONAL KG",
         "ACTUAL PROJECTED KG", "BBS COMPLETE", "VARIANCE KG (only if BBS complete)", "ACTUAL BLOCKED COMPONENTS",
         "FREELANCER CONCRETE M3", "FREELANCER STEEL T", "FREELANCER KG/M3 (implied)", "STATE", "NOTE"]
    data = [[x["category"], x["urban_released_concrete_m3"], x["urban_modelled_concrete_m3"],
             x["rough_ratio_kg_per_m3"], None, x["actual_released_kg"], x["actual_provisional_kg"], None,
             "YES" if x["actual_complete"] else "NO", None, "; ".join(x["actual_blocked_components"]),
             x["freelancer_concrete_m3"], x["freelancer_steel_t"], None, x["comparison_state"], x["note"]]
            for x in sorted(rr["rows"], key=lambda x: x["category"])]
    ws = sheet("02_Rebar", H, data, widths={"NOTE": 60, "ACTUAL BLOCKED COMPONENTS": 50})
    n = len(data)
    for i in range(2, n + 2):
        ws[f"E{i}"] = f'=IF(D{i}="","",C{i}*D{i})'
        ws[f"H{i}"] = f"=F{i}+G{i}"
        ws[f"J{i}"] = f'=IF(I{i}="YES",F{i}-E{i},"")'
        ws[f"N{i}"] = f'=IF(OR(L{i}="",M{i}=""),"",M{i}*1000/L{i})'
    t = n + 2
    ws[f"A{t}"] = "TOTAL (informational)"
    for c in "BCEFGHLM":
        ws[f"{c}{t}"] = f"=SUM({c}2:{c}{n + 1})"
    ws[f"A{t + 2}"] = rr["footer"]
    ws[f"A{t + 3}"] = ("Whole-project kg/m3 is an informational KPI only - rough estimation is category-specific. "
                       "A rough-vs-actual gap is a variance, never 'missing steel'.")
    ws[f"A{t + 4}"] = "Oracle net drawing rebar (known incomplete -> ORACLE_UNAVAILABLE for comparison): " + \
        "; ".join(f"{o['oracle']} {o['net_rebar_t']} t" for o in rr["oracle_net_rebar"])
    # ---------------------------------------------------------------- 03_Freelancer_Lineage
    lin = R["FREELANCER_CATEGORY_LINEAGE"]["categories"]
    H = ["CATEGORY", "ARABIC NAME", "SUMMARY CELL", "SUMMARY VALUE", "FORMULA", "SOURCE SHEET", "SOURCE CELL",
         "ROW LABEL", "SECTION", "B", "C", "D", "COUNT", "ROW FORMULA", "ROW VALUE", "PHYSICAL KIND", "LEAF KIND"]
    data, firsts = [], []
    for c in lin:
        firsts.append((len(data) + 2, len(c["subcomponents"]), c))
        for x in c["subcomponents"]:
            data.append([c["category_id"], c["arabic_name"], c["summary_cell"], c["summary_value"], c["formula"],
                         x.get("sheet"), x.get("cell"), x.get("label"), x.get("section"), x.get("B"), x.get("C"),
                         x.get("D"), x.get("count"), x.get("formula"), x.get("value"), x.get("physical_kind"),
                         x.get("kind")])
    ws = sheet("03_Freelancer_Lineage", H, data, widths={"SECTION": 34, "ROW LABEL": 26, "FORMULA": 30})
    chk = ws.max_column + 2
    ws.cell(row=1, column=chk, value="CATEGORY").font = bold
    ws.cell(row=1, column=chk + 1, value="SUM OF TRACED ROWS (formula)").font = bold
    ws.cell(row=1, column=chk + 2, value="SUMMARY VALUE").font = bold
    ws.cell(row=1, column=chk + 3, value="CHECK").font = bold
    for k, (r0, cnt, c) in enumerate([f for f in firsts if f[2]["unit"] == "m3" and
                                       f[2]["category_id"] != "TOTAL_RC"], 2):
        ws.cell(row=k, column=chk, value=c["category_id"])
        ws.cell(row=k, column=chk + 1, value=f"=SUM(O{r0}:O{r0 + cnt - 1})")
        ws.cell(row=k, column=chk + 2, value=c["summary_value"])
        a, b = L(chk + 1), L(chk + 2)
        ws.cell(row=k, column=chk + 3, value=f'=IF(ABS({a}{k}-{b}{k})<0.0005,"OK","MISMATCH")')
    # ---------------------------------------------------------------- 04_Scope_Map
    sm = R["FREELANCER_STRUCTURAL_SCOPE_MAP"]["categories"]
    H = ["CATEGORY", "ARABIC NAME", "SUMMARY VALUE", "PHYSICAL SCOPE (traced m3 by kind)", "POSSIBLE OVERLAP",
         "NOTES (scope vs label)"]
    sheet("04_Scope_Map", H, [[c["category_id"], c["arabic_name"], c["summary_value"],
                               json.dumps(c["physical_scope"], ensure_ascii=False),
                               " | ".join(c["possible_overlap"]), " | ".join(c["notes"])] for c in sm],
          widths={"PHYSICAL SCOPE (traced m3 by kind)": 60, "POSSIBLE OVERLAP": 60, "NOTES (scope vs label)": 80})
    # ---------------------------------------------------------------- 05_Oracle_Register
    soc = R["SOURCE_ORACLE_COMPARISON_REGISTER"]
    H = ["comparison_id", "fact_type", "urban_value", "urban_status", "oracle_name", "oracle_value", "oracle_status",
         "difference", "difference_percent", "basis_match", "result", "note"]
    sheet("05_Oracle_Register", H, [[r.get(k) if not isinstance(r.get(k), list) else "; ".join(map(str, r[k]))
                                     for k in H] for r in soc["rows"]])
    # ---------------------------------------------------------------- 06_Populations
    pop = R["STRUCTURAL_POPULATION_REGISTER"]
    H = ["GROUP", "POPULATION", "STATE", "QUANTITY STATE", "REFS", "WHY", "FLAGS"]
    sheet("06_Populations", H, [[r["group"], r["population"], r["state"], r.get("quantity_state"),
                                 "; ".join(r["refs"]), r.get("why"), "; ".join(r["flags"])] for r in pop["rows"]],
          widths={"REFS": 50, "WHY": 60})
    # ---------------------------------------------------------------- 07_Assumptions
    sheet("07_Assumptions", ["ASSUMPTION", "STATEMENT", "ROLE"],
          [[a["assumption_id"], a["statement"], a["role"]] for a in R["ASSUMPTION_AND_BASIS_REGISTER"]["rows"]],
          widths={"STATEMENT": 120})
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)


def main():
    idx, R = load()
    build(idx, R)
    rc = os.environ.get("RECALC_SCRIPT")
    if rc and Path(rc).exists():
        out = subprocess.run([sys.executable, rc, str(OUT), "180"], capture_output=True, text=True)
        res = json.loads(out.stdout[out.stdout.index("{"):])
        if res.get("status") != "success" or res.get("total_errors"):
            raise SystemExit(f"recalc failed: {res}")
        print("recalculated", res["total_formulas"], "formulas, 0 errors")
    man = {"workbook": OUT.name, "sha256": hashlib.sha256(OUT.read_bytes()).hexdigest(),
           "registers_consumed": idx["outputs"]}
    (OUT.parent / "REVIEW_MANIFEST.json").write_text(json.dumps(man, indent=1, sort_keys=True) + "\n")


if __name__ == "__main__":
    main()
