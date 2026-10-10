"""Alsenan Phase B1 benchmark readers: hash manifest, raw extraction and the normalised benchmark model.

Reads the pinned freelancer workbooks (xlsx with live formulas, BIFF xls with values only) and the web-app PDF text
layer. Nothing here is imported by any engine module and nothing here feeds an engine: it is a lab reader for the
comparison only (Phase B1 firewall: research/external_engine_lab only, origin = BENCHMARK on every record).

Raw extraction keeps every non-empty row with its cell references, original Arabic text, values and (xlsx) formulas.
The normalised model maps rows to a comparison taxonomy by declarative rules (file family + sheet + section + label);
no benchmark quantity is changed, original description and original unit are kept.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

ORIGIN = "BENCHMARK"
POLICY_ID = "ALSENAN_B1_BENCHMARK_READER_V1"

# the seven Phase B1 attachments, pinned by hash (no quantity is read before the hash matches)
FILES = {
    "FIN": {"name": "d901e652-____________________________1.xlsx",
            "sha256": "de448191fb9d0ababd33298289047c610366aba23912499c280146100be1151a",
            "family": "FINISHES", "kind": "FREELANCER_MANUAL_QS",
            "title_en": "floors, waterproofing, gypsum decor and railing"},
    "BLK": {"name": "5aab715f-_______.xlsx",
            "sha256": "873308bb6a3c99fd49332d6db881c41f51c94412a1caed986fe5d4a25179175d",
            "family": "BLOCKWORK", "kind": "FREELANCER_MANUAL_QS", "title_en": "blockwork"},
    "PNT": {"name": "e1c1f202-_____.xls",
            "sha256": "d1b459e531d78767041fe43a1ffa2f4560ad99939793044bfcbd1dbd13fa68f1",
            "family": "PAINT", "kind": "FREELANCER_MANUAL_QS", "title_en": "paint"},
    "RC": {"name": "d025d9f0-_________________1.xlsx",
           "sha256": "ae608411ed5b9e64811a15c353adbdd017fa3e51abfdffcf0aa1454b85d1957c",
           "family": "CONCRETE", "kind": "FREELANCER_MANUAL_QS", "title_en": "reinforced concrete"},
    "ALU": {"name": "64f4b7fc-__________.xlsx",
            "sha256": "7826608fc716ca0cdcd587aaf5ff9dd003d6968d6f77bca3289819844b2703e4",
            "family": "ALUMINIUM", "kind": "FREELANCER_MANUAL_QS", "title_en": "aluminium"},
    "PLS": {"name": "563d9805-______________________.xls",
            "sha256": "bf6e3123b988534dafcc488978eb8371e4f07a5f52857d01eb472db1aed22121",
            "family": "PLASTER", "kind": "FREELANCER_MANUAL_QS", "title_en": "internal and external plaster"},
    "WEB": {"name": "01e19893-Alsenan_Chalet_by_category.pdf",
            "sha256": "e92938e9b8df6f336a3b74ccf7962d1245f2e46e6d3645a71e95be6800ebc3bb",
            "family": "WEB_APP", "kind": "WEB_APP_COMMERCIAL_REPORT", "title_en": "Alsenan Chalet - Cost By Category"},
}


def sha256(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def verify(bench_dir) -> dict:
    """hash every pinned file; any mismatch stops the phase (fail closed)"""
    out = {}
    for k, f in FILES.items():
        p = Path(bench_dir) / f["name"]
        h = sha256(p) if p.exists() else None
        out[k] = {"present": p.exists(), "sha256": h, "match": h == f["sha256"], "bytes": p.stat().st_size if h else None}
    bad = [k for k, v in out.items() if not v["match"]]
    if bad:
        raise RuntimeError(f"benchmark hash mismatch / missing: {bad}")
    return out


def _v(x):
    if isinstance(x, (datetime, date)):
        return x.isoformat()
    if isinstance(x, float) and x.is_integer() and abs(x) < 1e15:
        return x
    return x


# ------------------------------------------------------------------ raw extraction
def raw_xlsx(path, key) -> tuple[list, list]:
    import openpyxl
    vals = openpyxl.load_workbook(path, data_only=True)
    forms = openpyxl.load_workbook(path, data_only=False)
    rows, sheets = [], []
    for i, ws in enumerate(vals.worksheets):
        wf = forms.worksheets[i]
        sheets.append({"sheet": ws.title, "max_row": ws.max_row, "max_column": ws.max_column})
        for r in ws.iter_rows():
            cells = {}
            for c in r:
                if c.value is None:
                    continue
                f = wf[c.coordinate].value
                cells[c.coordinate] = {"v": _v(c.value),
                                       "f": f if isinstance(f, str) and f.startswith("=") else None}
            if cells:
                rows.append({"file": key, "sheet": ws.title, "row": r[0].row, "cells": cells})
    return rows, sheets


def raw_xls(path, key) -> tuple[list, list]:
    import xlrd
    b = xlrd.open_workbook(path)
    rows, sheets = [], []
    for s in b.sheets():
        sheets.append({"sheet": s.name, "max_row": s.nrows, "max_column": s.ncols})
        for r in range(s.nrows):
            cells = {}
            for c in range(s.ncols):
                v = s.cell_value(r, c)
                if v in ("", None):
                    continue
                cells[f"{xlrd.colname(c)}{r + 1}"] = {"v": _v(v), "f": None}
            if cells:
                rows.append({"file": key, "sheet": s.name, "row": r + 1, "cells": cells})
    return rows, sheets


def col_of(ref) -> str:
    return re.match(r"[A-Z]+", ref).group(0)


def cell(row, col):
    """value of column `col` in a raw row (None if empty)"""
    for k, c in row["cells"].items():
        if col_of(k) == col:
            return c["v"]
    return None


def formula(row, col):
    for k, c in row["cells"].items():
        if col_of(k) == col:
            return c["f"]
    return None


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def txt(v):
    return v.strip() if isinstance(v, str) and v.strip() else None


# ------------------------------------------------------------------ web-app PDF
WEB_ROW = re.compile(r"^(?P<desc>.*?)\s*(?P<qty>[0-9][0-9,]*(?:\.[0-9]+)?)\s*(?P<unit>qty|item|m³|ton|KWD)\s+"
                     r"(?P<rate>[0-9,]+(?:\.[0-9]+)?)\s*KWD\s+(?P<total>[0-9,]+(?:\.[0-9]+)?)\s*KWD$")
WEB_CAT = re.compile(r"^CategoryTotal:(?P<name>.*?)(?P<amount>[0-9,]+(?:\.[0-9]+)?)KWD$")
WEB_TYPE_TOTAL = re.compile(r"^Type Total:\s*(?P<name>.*?)\s*(?P<amount>[0-9,]+(?:\.[0-9]+)?)KWD$")
WEB_GRAND = re.compile(r"^GrandTotal\s*(?P<amount>[0-9,]+(?:\.[0-9]+)?)KWD$")
WEB_SKIP = re.compile(r"^(AlsenanChalet|Cost By Category|Page [0-9]+|[0-9]{2}/[0-9]{2}/[0-9]{4}|Description Contractor.*)$")


def nfkc(s):
    return unicodedata.normalize("NFKC", s) if s else s


def _f(s):
    return float(s.replace(",", ""))


def web_rows(path) -> dict:
    """parse the web-app text layer: category / type / rows with qty, unit, rate, total; check every total"""
    import pypdf
    r = pypdf.PdfReader(str(path))
    lines = []
    for i, p in enumerate(r.pages):
        for ln in (p.extract_text() or "").splitlines():
            ln = ln.strip()
            if ln:
                lines.append((i + 1, ln))
    rows, cats, types, pending, ctype, prev = [], [], [], [], None, None
    grand = None
    for page, ln in lines:
        if WEB_SKIP.match(ln):
            if ln.startswith("Description") and prev is not None:
                ctype = prev
            prev = None
            continue
        m = WEB_ROW.match(ln)
        if m:
            pending.append({"page": page, "type_raw": ctype, "desc_raw": m["desc"].strip(), "desc_nfkc": nfkc(m["desc"].strip()),
                            "qty": _f(m["qty"]), "unit": m["unit"], "rate_kwd": _f(m["rate"]),
                            "total_kwd": _f(m["total"]), "line": ln})
            prev = None
            continue
        m = WEB_TYPE_TOTAL.match(ln)
        if m:
            types.append({"type_raw": m["name"], "type_nfkc": nfkc(m["name"]), "total_kwd": _f(m["amount"]),
                          "rows": len([x for x in pending if x["type_raw"] == ctype])})
            prev = None
            continue
        m = WEB_CAT.match(ln)
        if m:
            cid = len(cats)
            for x in pending:
                x["category_index"] = cid
                x["category_raw"] = m["name"]
                x["category_nfkc"] = nfkc(m["name"])
            cats.append({"index": cid, "category_raw": m["name"], "category_nfkc": nfkc(m["name"]),
                         "total_kwd": _f(m["amount"]), "rows": len(pending),
                         "rows_total_kwd": round(sum(x["total_kwd"] for x in pending), 2)})
            rows.extend(pending)
            pending, prev = [], None
            continue
        m = WEB_GRAND.match(ln)
        if m:
            grand = _f(m["amount"])
            continue
        prev = ln
    for i, x in enumerate(rows):
        x["id"] = f"WEB-{i + 1:03d}"
        x["arith_ok"] = abs(x["qty"] * x["rate_kwd"] - x["total_kwd"]) < 0.01 if x["unit"] != "KWD" else None
    checks = {"categories_sum_to_rows": all(abs(c["total_kwd"] - c["rows_total_kwd"]) < 0.01 for c in cats),
              "grand_total_kwd": grand,
              "grand_equals_categories": grand is not None and abs(sum(c["total_kwd"] for c in cats) - grand) < 0.01,
              "unparsed_rows": len(pending), "pages": len(r.pages), "lines": len(lines)}
    return {"rows": rows, "categories": cats, "checks": checks}


# ------------------------------------------------------------------ normalisation
FLOOR_WORDS = (("السرداب", "FOUNDATION"), ("رقاب", "FOUNDATION"), ("القواعد", "FOUNDATION"), ("الأرض", "GF"),
               ("الارض", "GF"), ("الأول", "1F"), ("الاول", "1F"), ("دور اول", "1F"), ("السطح", "ROOF"))
TOTAL_WORDS = ("إجمالي", "الإجمالي", "اجمالي", "اجمالى", "Total", "ما قبلة", "ما بعدة", "الإجمالي الكلي",
               "إجمالي الكلي", "اجــــــــــــــمالى")


HEADER_NOISE = ("بيان البند", "العرض م", "الحجم العنصر", "Page", "الأبعاد", "طولي", "سم  متر", "البيان", "ملاحظات",
                "الحجم صافي")


def update_context(c, t, fam=None):
    c["section"], c["label"] = t.strip(), None
    fl = floor_of(t)
    if fl:
        c["floor"] = fl
        c["floor_src"] = "SUB" if "الطابق" in t else "BLOCK"
    elif fam == "BLOCKWORK" and "خصم" in t and c.get("floor_src") == "SUB":
        c["floor"] = None                                   # a whole-block deduction list: no floor stated
    if "خصم" in t:
        c["mode"] = "DEDUCTION"
    elif any(w in t for w in ("سمك", "الخارج", "مبانى", "مباني", "الطابوق")):
        c["mode"] = None
    if "سمك 15" in t:
        c["thick"] = "150"
    elif "سمك 20" in t:
        c["thick"] = "200"
    elif "الخارج" in t:
        c["thick"] = "EXT"
    if any(w in t for w in ("المطابخ", "عازل")):
        c["wet"] = True
    elif any(w in t for w in ("ارضيات", "ديكور  ")):
        c["wet"] = False


def floor_of(text):
    for w, fl in FLOOR_WORDS:
        if text and w in text:
            return fl
    return None


def is_total_label(t):
    return bool(t) and any(t.startswith(w) or t == w for w in TOTAL_WORDS)


GLOSSARY = {
    # concrete
    "العاديه": "plain concrete (blinding)", "الشناج المحيط": "perimeter strap / edge beam",
    "الارضيه الخرسانيه": "concrete ground slab", "شناجات طوليه": "ground beams, longitudinal",
    "شناجات عرضيه": "ground beams, transverse", "شناج السور": "fence ground beam",
    "حوائط 20 سم المصعد": "elevator walls 20 cm", "S.S  16 CM": "solid slab 16 cm", "S.S 16 CM": "solid slab 16 cm",
    "بلاطة + الدرج سلم المدخل": "entrance stair slab + steps", "بلاطة + الدرج سلم 2": "stair 2 slab + steps",
    "القبه": "dome", "الدرج + البلاطة": "main stair steps + slab", "البسطة": "stair landing",
    "الأرضية": "pool floor slab", "حوائط حمام السباحة": "pool walls", "حوائط المضخة": "pump-room walls",
    "درج الحمام ": "pool steps", "ارضيه الغرفه": "pump-room floor",
    # finishes
    "استقبال+طعام": "reception + dining", "غرفه نوم رئيسيه": "master bedroom", "مدخل": "entrance hall",
    "ديوانيه": "diwaniya", "سايق ": "driver's room", "معيشه": "living", "غرفه نوم 1": "bedroom 1",
    "غرفه نوم 2": "bedroom 2", "غرفه نوم 3": "bedroom 3", "غرفه نوم 4": "bedroom 4", "غرفه خادمه": "maid's room",
    "طرقه": "corridor", "مغسله الديوانيه": "diwaniya washroom", "حمامات ": "bathrooms", "مطبخ": "kitchen",
    "حمام": "bathroom", "مغسله": "washroom", "غسيل": "laundry", "الملاحق والاسطح": "annexes and roofs",
    "درج": "stair treads", "بسطه": "landing", "نعلات": "skirting", "تواشيح": "nosings",
    # aluminium / blockwork
    "باب خارجي": "external door", "شباك حمام": "bathroom window", "شباك المسبح": "pool-facing window",
    "شباك": "window", "باب": "door", "البروزات ": "projections", "السور": "boundary fence wall", "الدروه": "parapet",
    # paint / plaster rooms
    "الديوانية": "diwaniya", "غرفة السائق": "driver's room", "غرفة ": "room", "صالة": "hall", "حائط درج": "stair wall",
    "مغسلة     ": "washroom", "غرفة رئيسية  1  دور اول": "main room 1, first floor", "غرفة رئيسية 2  ": "main room 2",
    "غرفة رئسية 3": "main room 3", "غرفة رئيسية 4": "main room 4", "معيشة": "living", "ممر مغاسل": "washroom corridor",
    " درج الخدمات   لشواحط الدرج": "service stair (stair soffits)", "درج الصالة": "hall stair",
    "غرفة خدامة    السطح": "maid's room, roof", "ممر": "corridor", "حوائط     بيت الدرج": "stairwell walls",
    "تابع": "continued", "شرشوب أبواب": "door lintel bands", "شباك    فراغات خارجية": "window, external openings",
    "شباك    ": "window", "ش": "window", "لشباك حمامات": "bathroom windows", "باب   فراغات داخلية": "door, internal openings",
    "الغرف والصالات": "rooms and halls", "الحمامات": "bathrooms", "كوف لايت ": "cove light",
    "كوف لايت حمامات": "cove light, bathrooms", "للحمامات والمطابخ": "bathrooms and kitchens",
    "  واجهة مرتفعة  واجهة غربية جنوبية": "facade high, south-west", "واجهة متوسطة": "facade middle",
    "واجهة مرتفعة": "facade high", "واجهة ملحق": "annex facade", "واجهة    واجهة جنوبية شرقية": "facade south-east",
    "منخفضة": "low", "لملحق": "annex", "اطار": "frame band", "واجهة     شمالية شرقية": "facade north-east",
    "مرتفعة": "high", "للملحق      راسية": "annex, vertical", "واجهة ملحق   واجهة شمالية غربية": "annex facade, north-west",
    "واجهة اعلى حمام السباحة": "facade above the pool", "قبة من الداخل والخارج": "dome inside and outside",
    "سقف اعلى حمام لبسباحة": "ceiling above the pool", "واجهة": "facade", "واجهة    واجهة السطح": "roof facade",
    "دروة": "parapet", "دروة سطح الملحق": "annex-roof parapet", "دروة    سطح الملاحق": "annex-roof parapets",
    "دروة    ": "parapet", "سور": "boundary wall", "كتف": "pier", "شرشوب م2": "lintel band m2", "للمساح": "for plaster",
    "للطرطشة": "for spatter coat",
}


GLOSSARY.update({
    "(subtotal)": "subtotal (unlabelled)", "Total": "total", "إجمالي": "total", "الإجمالي": "total", "اجمالي": "total",
    "إجمالي الكلي": "grand total", "الإجمالي الكلي": "grand total", "اجــــــــــــــمالى": "total", "ما قبلة": "brought forward",
    "ما بعدة": "carried forward", "الخرسانه العاديه": "plain concrete", "اجمالى القواعد": "total footings",
    "اجمالي الشناجات والارضيه": "total ground beams and ground slab", "اجمالى الحوائط والأعمدة": "total walls and columns",
    "الجسور": "beams", "البلاطات": "slabs", "الدرج+القبه": "stairs + dome", "حمام السباجة": "swimming pool",
    "اجمالى الخرسانة المسلحة": "total reinforced concrete", "اجمالي كميات الحديد": "total rebar tonnage",
    "اجمالى المباني": "total blockwork", "مبانى الخارجي م2": "external blockwork m2", "مبانى سمك 15سم  م2": "blockwork 15 cm m2",
    "مبانى سمك 20 سم م2": "blockwork 20 cm m2", "الإجمالي  الخارجي": "total external", "الإجمالي الخصم": "total deduction",
    "الإجمالي بعد الخصم": "total after deduction", "اجمالي خصم فتحات": "total opening deduction",
    "اجمالي الحوائط الخارجيه الكلي": "total external walls", "اجمالي الأبواب م2": "total doors m2",
    "اجمالي الشبابيك م2": "total windows m2", "اجمالي كامل": "grand total", "اجمالي الكميات": "total quantities",
    "اجمالي الارضيات": "total floors", "اجمالي النعلات": "total skirting", "اجمالي الحوائط": "total walls",
    "ارضيات حمام السباحه": "pool floors", "حوائط حمام السباحه": "pool walls", "ارضيات الحوش": "yard floor",
    "نعلات الحوش": "yard skirting", "اجمالي الدرج والبسطات": "total stairs and landings", "اجمالي التواشيح": "total nosings",
    "عازل ارضيات حمامات": "wet-room floor waterproofing", "نعلات العازل للحمامات": "wet-room waterproofing upturns",
    "عازل اسطح و ملاحق": "roof and annex waterproofing", "نعلات العازل الاسطح": "roof waterproofing upturns",
    "ديكور الصالات والغرف": "gypsum decor halls and rooms", "كرانيش الصالات والغرف": "cornices halls and rooms",
    "ديكور الحمامات والمطابح": "gypsum decor bathrooms and kitchens", "كرانيش الحمامات والمطابخ": "cornices bathrooms and kitchens",
    "دربزين داخلي": "internal railing", "اجمالي مسطحات الدرج": "total stair areas", "الارضيات": "floors",
    "النعلات": "skirting", "الحوائط": "walls", "الصالات والغرف": "halls and rooms", "كرانيش الصالات": "cornices halls",
    "الحمامات والمطابخ": "bathrooms and kitchens", "كرانيش الحمامات": "cornices bathrooms", "yard": "yard",
    "صبغ الحوائط م2": "wall paint m2", "صبغ الديكور مع الكوف لايت م2": "decor paint incl. cove light m2", "مساح م2": "plaster m2",
    "زوايا ونهايات م.ط": "corners and ends m (2 m = 1 m)", "طرطشة حمامات ومطابخ م2": "spatter coat bathrooms and kitchens m2",
    "طرطشة تحت نعلة م2": "spatter coat under skirting m2", "مساح خارجى م2": "external plaster m2",
    "زوايا وفواصل ومبروم م.ط": "external corners, joints and beads m (2 m = 1 m)", "لابواب مصاعد": "for lift doors",
    "لشواحط الدرج": "for stair soffits", "حمامات": "bathrooms", "سايق": "driver's room", "غرفة": "room",
    "مغسلة": "washroom", "درج الحمام": "pool steps", "غرفة رئيسية 2": "main room 2",
    "واجهة مرتفعة  واجهة غربية جنوبية": "facade high, south-west", "البروزات": "projections",
})


def en(label):
    if label is None:
        return "(unlabelled row - continuation of the section)"
    t = label.strip()
    if re.fullmatch(r"(F[\d.]*[FN]?|STB\d+|[BG]?CB\d+|B\d+|C\d*|CN|PC\d*|S\.S\s+\d+ CM)", t):
        return f"type mark {t}"
    return GLOSSARY.get(label) or GLOSSARY.get(t) or "(no glossary entry - see original)"


def _footing_type(label):
    t = (label or "").strip().upper().replace(".", "")
    return t if re.fullmatch(r"F\d*|FF|FN", t) else None


def _strap_type(label):
    m = re.fullmatch(r"STB(\d+)", (label or "").strip().upper())
    return f"SB{m.group(1)}" if m else None


def _column_type(label):
    t = (label or "").strip().upper()
    return t if re.fullmatch(r"C\d*|CN|PC\d*", t) else None


def _beam_type(label):
    t = (label or "").strip().upper()
    m = re.fullmatch(r"[BG]?CB(\d+)", t)
    if m:
        return f"CB{m.group(1)}"
    return t if re.fullmatch(r"B\d+|BW|CA", t) else None


def _row_numbers(row, cols):
    return {c: num(cell(row, c)) for c in cols}


def normalise(raw: list) -> list:
    """one normalised record per raw row that carries a quantity, with section / floor context and taxonomy.
    Header rows set the context; total rows are kept and marked (never mixed with details)."""
    out = []
    ctx = {}
    for r in raw:
        key, sheet = r["file"], r["sheet"]
        sk = (key, sheet)
        c = ctx.setdefault(sk, {"section": None, "floor": None, "label": None})
        fam = FILES[key]["family"]
        texts = [(k, v["v"]) for k, v in r["cells"].items() if isinstance(v["v"], str) and v["v"].strip()]
        nums = [(k, v["v"]) for k, v in r["cells"].items() if num(v["v"]) is not None]
        rec = None
        if fam == "CONCRETE":
            rec = _norm_concrete(r, c)
        elif fam in ("BLOCKWORK", "ALUMINIUM"):
            rec = _norm_blk_alu(r, c, fam)
        elif fam == "FINISHES":
            rec = _norm_finishes(r, c)
        elif fam in ("PAINT", "PLASTER"):
            rec = _norm_paint_plaster(r, c, fam)
        if rec is None:
            t = " ".join(str(v) for _, v in texts)
            if texts and not nums and not any(w in t for w in HEADER_NOISE):
                update_context(c, t, fam)                   # a header line: new section context
            continue
        rec.update({"file": key, "sheet": sheet, "row": r["row"], "origin": ORIGIN,
                    "section": rec.get("section", c["section"]),
                    "floor": rec.get("floor") or c["floor"],
                    "cells": sorted(r["cells"])})
        rec["id"] = f"{key}:{sheet}:{r['row']}"
        rec["label_en"] = en(rec.get("label_ar"))
        out.append(rec)
    return out


def _qty_rec(label, qty, unit, kind, trade, sub, **kw):
    return dict(label_ar=label, qty=qty, unit=unit, row_kind=kind, trade=trade, subtrade=sub, **kw)


def _norm_concrete(r, c):
    sheet = r["sheet"]
    a = txt(cell(r, "A"))
    if sheet == "ورقة1":                                    # cover: per-group totals and rebar tonnage
        if a and num(cell(r, "E")) is not None:
            rec = _qty_rec(a, num(cell(r, "E")), "m3", "COVER_TOTAL", *_rc_cover_trade(a))
            rec["formula"] = formula(r, "E")
            rec["rebar_t"] = num(cell(r, "H"))
            rec["rebar_formula"] = formula(r, "H")
            return rec
        if a and num(cell(r, "F")) is not None:
            return _qty_rec(a, num(cell(r, "F")), "m3", "COVER_TOTAL", "OTHER_CONCRETE", "ALL_REINFORCED",
                            formula=formula(r, "F"))
        if a and num(cell(r, "H")) is not None:
            return _qty_rec(a, num(cell(r, "H")), "t", "COVER_TOTAL", "REBAR", "ALL", formula=formula(r, "H"))
        return None
    netcol = {"السلم": "G"}.get(sheet, "I")
    cntcol = {"بلاطات ": "H"}.get(sheet, "F")
    net = num(cell(r, netcol))
    if a and is_total_label(a) or (a is None and net is not None and not any(
            num(cell(r, x)) is not None for x in "BCDE")):
        if net is None and num(cell(r, "B")) is not None:
            net = num(cell(r, "B"))
        if net is None:
            return None
        return _qty_rec(a or "(subtotal)", net, "m3", "SUBTOTAL", None, None, formula=formula(r, netcol) or formula(r, "B"))
    if net is None:
        return None
    if a:
        c["label"] = a
    label = a or c["label"]
    dims = _row_numbers(r, "BCD")
    cnt = num(cell(r, cntcol))
    rec = _qty_rec(label, net, "m3", "DETAIL", *_rc_trade(sheet, label, c["section"]), count=cnt if cnt is not None else 1.0,
                   dims={"B": dims["B"], "C": dims["C"], "D": dims["D"]}, unit_volume=num(cell(r, "E")),
                   formula=formula(r, netcol), unit_formula=formula(r, "E"), continuation=a is None)
    if sheet == "بلاطات ":
        rec["deduction"] = {"each_m3": num(cell(r, "F")), "count": num(cell(r, "G")), "formula": formula(r, "F")}
    return rec


def _rc_cover_trade(a):
    if "العاديه" in a:
        return "BLINDING", "PLAIN"
    for w, t in (("القواعد", ("FOOTING", "FOOTINGS+STRAPS+PERIMETER")), ("الشناجات", ("GROUND_BEAM", "GROUND_SLAB+BEAMS")),
                 ("الأعمدة", ("COLUMN", "NECKS+COLUMNS+WALLS")), ("الجسور", ("BEAM", "ALL")),
                 ("البلاطات", ("SLAB", "ALL")), ("الدرج", ("STAIR_CONCRETE", "STAIRS+DOMES")),
                 ("حمام", ("OTHER_CONCRETE", "POOL"))):
        if w in a:
            return t
    return "OTHER_CONCRETE", None


def _rc_trade(sheet, label, section):
    if sheet == "القواعد":
        if label == "العاديه":
            return "BLINDING", "PLAIN"
        if _footing_type(label):
            return "FOOTING", _footing_type(label)
        if _strap_type(label):
            return "STRAP_BEAM", _strap_type(label)
        return "GROUND_BEAM", "PERIMETER_STRAP"
    if sheet == "الشناجات":
        return ("OTHER_CONCRETE", "GROUND_SLAB") if "الارضيه" in label else ("GROUND_BEAM", label)
    if sheet == "الحوائط + الأعمدة":
        if "حوائط" in label:
            return "OTHER_CONCRETE", "ELEVATOR_WALL"
        sub = "NECK" if section and "رقاب" in section else "COLUMN"
        return "COLUMN", f"{sub}:{_column_type(label)}"
    if sheet == "كمرات":
        return "BEAM", _beam_type(label)
    if sheet == "بلاطات ":
        return "SLAB", "SOLID"
    if sheet == "السلم":
        return ("OTHER_CONCRETE", "DOME") if "القبه" in label else ("STAIR_CONCRETE", label)
    if sheet == "حمام السباحة":
        return "OTHER_CONCRETE", "POOL"
    return "OTHER_CONCRETE", None


def _norm_blk_alu(r, c, fam):
    sheet = r["sheet"]
    a = txt(cell(r, "A"))
    if sheet == "الغلاف":
        q = num(cell(r, "E")) if num(cell(r, "E")) is not None else num(cell(r, "F"))
        if a and q is not None:
            return _qty_rec(a, q, "m2", "COVER_TOTAL", "BLOCKWORK" if fam == "BLOCKWORK" else "ALUMINIUM",
                            _cover_sub(a), formula=formula(r, "E") or formula(r, "F"))
        return None
    tot = num(cell(r, "H"))
    if tot is None:
        return None
    b, cl, d = num(cell(r, "B")), num(cell(r, "C")), num(cell(r, "D"))
    if cl is None and d is None:
        return _qty_rec(a or "(subtotal)", tot, "m2", "SUBTOTAL", None, None, formula=formula(r, "H"))
    sec = c["section"] or ""
    if fam == "BLOCKWORK":
        th = c.get("thick")
        sub = {"EXT": "EXTERNAL", "150": "INTERNAL_150", "200": "INTERNAL_200"}.get(th, "UNKNOWN")
        if c.get("mode") == "DEDUCTION":
            sub += ":DEDUCTION"
        elif a:
            sub += {"السور": ":FENCE", "الدروه": ":PARAPET", "البروزات ": ":PROJECTION"}.get(a, "")
        trade = "BLOCKWORK"
    else:
        # no label inheritance: an unlabelled row is a window by the cover arithmetic (doors = the rows labelled باب)
        sub = "DOOR" if a and "باب" in a else "WINDOW"
        trade = "ALUMINIUM"
    return _qty_rec(a, tot, "m2", "DETAIL", trade, sub, count=b if b is not None else 1.0,
                    dims={"length_m": cl, "height_or_width_m": d}, formula=formula(r, "H"),
                    label_inherited=False, subtrade_basis=None if a else "COVER_ARITHMETIC (doors total = labelled باب rows)")


def _cover_sub(a):
    for w, s in (("الخارجي", "EXTERNAL"), ("15", "INTERNAL_150"), ("20", "INTERNAL_200"), ("الأبواب", "DOOR"),
                 ("الشبابيك", "WINDOW")):
        if w in a:
            return s
    return "ALL"


def _norm_finishes(r, c):
    sheet = r["sheet"]
    a = txt(cell(r, "A"))
    if sheet == "الغلاف":
        q = num(cell(r, "E"))
        if a and q is not None:
            return _qty_rec(a, q, txt(cell(r, "F")) or "", "COVER_TOTAL", *_fin_cover_trade(a, txt(cell(r, "G"))),
                            formula=formula(r, "E"), source_formula=formula(r, "B"))
        return None
    C, D, F, G, H, B = (num(cell(r, x)) for x in "CDFGHB")
    if C is None and D is None and G is None and H is None:
        return None                                         # a header / caption line
    fH, fD, fG = formula(r, "H") or "", formula(r, "D") or "", formula(r, "G") or ""
    summed = any(f.startswith("=SUM") for f in (fH, fD, fG)) or (fH.count("+") and "*" not in fH)
    if (a and is_total_label(a)) or summed or (D is None and C is None and B is None and sheet != "رخام+حوش"):
        val = H if H is not None else (D if D is not None else (G if G is not None else C))
        return _qty_rec(a or "(subtotal)", val, "", "SUBTOTAL", None, None, formula=fH or fD or fG or None,
                        perimeter_sum=C if fD.startswith("=SUM") else None)
    if a:
        c["label"] = a
    label = a or c["label"]
    sec = c["section"] or ""
    if sheet == "رخام+حوش":
        if "الحوش" in sec:
            return _qty_rec(label or "yard", D, "m2", "DETAIL", "FLOOR", "EXTERNAL_YARD", perimeter_m=C)
        if label == "نعلات":
            return _qty_rec(label, G, "m", "DETAIL", "SKIRTING", "MARBLE_STAIR", stair=sec)
        if label == "تواشيح":
            return _qty_rec(label, G, "nr", "DETAIL", "MARBLE", "NOSING", stair=sec)
        return _qty_rec(label, H, "m2", "DETAIL", "MARBLE", "STAIR" if label == "درج" else "LANDING", stair=sec,
                        count=B, dims={"length_m": C, "area_m2": D})
    wet = bool(c.get("wet"))
    trade = {"سيراميك": "FLOOR", "ديكور": "CEILING", "العازل": "WATERPROOFING"}[sheet]
    room_class = "WET" if wet else "DRY"
    rec = _qty_rec(label, D, "m2", "DETAIL", trade, f"{room_class}_ROOM", perimeter_m=C, height_m=F,
                   wall_m2=G, area_formula=formula(r, "H"), wall_formula=formula(r, "G"), continuation=a is None,
                   room_class=room_class)
    if sheet == "العازل" and label == "الملاحق والاسطح":
        rec["subtrade"] = "ROOF"
    return rec


def _fin_cover_trade(a, g):
    g = g or ""
    if "حمام السباحه" in a:
        return "FLOOR" if "ارضيات" in a else "WALL_TILE", "POOL"
    if "الحوش" in a:
        return ("FLOOR" if "ارضيات" in a else "SKIRTING"), "EXTERNAL_YARD"
    if "رخام" in g:
        return ("MARBLE" if "الدرج" in a or "التواشيح" in a else "SKIRTING"), "STAIR"
    if "عازل" in g:
        return "WATERPROOFING", ("ROOF" if "اسطح" in a else "WET_ROOM") + (":UPTURN" if "نعلات" in a else ":FLOOR")
    if "ديكور" in g:
        return "CEILING", ("CORNICE" if "كرانيش" in a else "GYPSUM") + (":WET" if "الحمامات" in a else ":DRY")
    if "سيراميك" in g:
        return {"اجمالي الارضيات": ("FLOOR", "ALL"), "اجمالي النعلات": ("SKIRTING", "ALL"),
                "اجمالي الحوائط": ("WALL_TILE", "WET_ROOMS")}.get(a, ("FLOOR", None))
    if "دربزين" in a:
        return "RAILING", "INTERNAL"
    return "OTHER", None


def _norm_paint_plaster(r, c, fam):
    sheet = r["sheet"]
    if sheet == "الفاتورة":
        a = txt(cell(r, "A"))
        q = num(cell(r, "F"))
        if a and q is not None:
            return _qty_rec(a, q, "", "COVER_TOTAL", "PAINT" if fam == "PAINT" else "PLASTER", _pp_cover_sub(a),
                            gross=num(cell(r, "B")), deduction=cell(r, "D"))
        return None
    if sheet != "البيــــــــــــان":
        return None
    d = txt(cell(r, "D"))
    E, F, G, H, I, J = (num(cell(r, x)) for x in "EFGHIJ")
    if I is None and J is None:
        if d and fam == "PLASTER" and num(cell(r, "B")) is not None:
            return _qty_rec(d, num(cell(r, "B")), "m", "DETAIL", "PLASTER", "CORNERS_NOTE")
        return None
    measured = any(v is not None for v in (F, G, H))
    if not measured:                                        # a running / grand subtotal (J only, or I only)
        return _qty_rec(d or "(subtotal)", J if J is not None else I, "m2", "SUBTOTAL", "PAINT" if fam == "PAINT"
                        else "PLASTER", None, ends_m=num(cell(r, "A")) if fam == "PLASTER" else None,
                        corners_m=num(cell(r, "B")) if fam == "PLASTER" else None)
    if d:
        c["label"] = d
    label = d or c["label"]
    sec = c["section"] or ""
    trade = "PAINT" if fam == "PAINT" else "PLASTER"
    sub = ("DECOR" if "الديكور" in sec else "EXTERNAL" if "خارجى" in sec else
           "SPATTER" if "طرطشة" in sec else "INTERNAL_WALL")
    if label and (label.strip() in ("ش", "شباك", "باب") or label.startswith(("شباك", "باب", "لشباك"))):
        sub = "OPENING_DEDUCTION"
    if label and label.strip() in ("للمساح", "للطرطشة"):
        sub = "OPENING_DEDUCTION_SUMMARY"
    rec = _qty_rec(label, I if I is not None else J, "m2", "DETAIL", trade, sub,
                   count=E, height_m=F, width_m=G, length_m=H, subtotal=J,
                   label_inherited=d is None)
    if fam == "PLASTER":
        rec["ends_m"] = num(cell(r, "A"))
        rec["corners_m"] = num(cell(r, "B"))
    return rec


def _pp_cover_sub(a):
    for w, s in (("الحوائط", "INTERNAL_WALL"), ("الديكور", "DECOR"), ("مساح خارجى", "EXTERNAL"), ("مساح م2", "INTERNAL_WALL"),
                 ("زوايا", "CORNERS"), ("طرطشة تحت", "SPATTER_UNDER_SKIRTING"), ("طرطشة", "SPATTER_WET"),
                 ("فرفيس", "LUMP")):
        if w in a:
            return s
    return None


def policy_record() -> dict:
    rec = {"id": POLICY_ID, "origin": ORIGIN, "files": {k: v["sha256"] for k, v in FILES.items()},
           "rules": ["hash before read", "every non-empty row kept raw (cell refs, Arabic, values, formulas)",
                     "header rows set section / floor context", "total rows kept and marked, never mixed with details",
                     "no benchmark value changed; original description and unit kept"]}
    rec["digest"] = hashlib.sha256(repr(sorted(rec.items())).encode()).hexdigest()
    return rec
