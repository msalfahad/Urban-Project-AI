"""Alsenan Phase B1 - benchmark reveal + forensic comparison ONLY (EXTRACT -> MAP -> COMPARE -> CLASSIFY -> FREEZE).

Reads the frozen Phase A3 registers (tests/alsenan/registers_a3, commit 9aa2741) and the hash-pinned benchmark files;
writes the B1 comparison registers. It imports no engine quantity module, changes no engine, and no value it reads
from a benchmark file can reach an engine input (the A3 registers are only read).

    python3 research/external_engine_lab/alsenan_phase_b1.py <benchmark_dir> <register_dir> [code_commit]

Comparison discipline:
- comparability gate first; only EXACT_COMPARABLE pairs get a percentage;
- identity between a benchmark row and an Urban object is never taken from the value being compared (a value
  coincidence is reported as LOW-confidence diagnostic, not as a match);
- every difference gets one primary class (and secondary classes); a difference is investigated, never assumed to be
  an Urban error;
- tolerances are the ones declared in the committed recommendation (alsenan_phase_b1_recommendation.json).
"""

from __future__ import annotations

import hashlib
import json
import re
import statistics
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).parent))
import alsenan_b1_benchmark as BB                                                 # noqa: E402

A3_DIR = ROOT / "tests/alsenan/registers_a3"
A3_COMMIT, A3_CODE_COMMIT = "9aa2741", "7ac5d9e"
A3_FREEZE = "ALSENAN_P7757_ST7757_PHASE_A3_FREEZE"
A3_FREEZE_SHA = "de01b42d1478628e618a508f09316c4f9a879d4883d451224653866e18d5c849"
REC_PATH = ROOT / "research/external_engine_lab/alsenan_phase_b1_recommendation.json"
RECOMMENDATION_COMMIT = "11d661c"
FREEZE = "ALSENAN_PHASE_B1_COMPARISON_FREEZE"
XLSX_NAME = "URBAN_QTO_ALSENAN_PHASE_B1_COMPARISON.xlsx"
BANNER = ("PHASE B1 - BENCHMARK REVEAL + FORENSIC COMPARISON ONLY - ALSENAN P7757 + ST7757 - NO ENGINE CHANGE - "
          "NO PRICING EVALUATION - BENCHMARK VALUES NEVER ENTER AN ENGINE - NO PRODUCTION MIGRATION")

CLASSES = ("MATCH_WITHIN_TOLERANCE", "SCOPE_MISMATCH", "SOURCE_GAP", "MANUAL_ASSUMPTION_NOT_IN_SOURCE",
           "MANUAL_METHOD_DIFFERENCE", "SOURCE_DECODER_DEFECT", "UNIT_FRAME_DEFECT", "GEOMETRY_DEFECT",
           "TOPOLOGY_DEFECT", "ROOM_SEMANTIC_DEFECT", "OPENING_DEFECT", "HEIGHT_EVIDENCE_GAP", "TRADE_RULE_GAP",
           "STRUCTURAL_TYPE_BINDING_DEFECT", "STRUCTURAL_LENGTH_DEFECT", "STRUCTURAL_VOLUME_METHOD_DIFFERENCE",
           "DUPLICATE_OR_OMITTED_ITEM", "REPORTING_DEFECT", "BENCHMARK_POSSIBLE_ERROR", "UNRESOLVED_REQUIRES_REVIEW")
COMPARABILITY = ("EXACT_COMPARABLE", "PARTIAL_SCOPE_COMPARABLE", "UNIT_CONVERSION_REQUIRED", "SCOPE_MISMATCH",
                 "URBAN_BLOCKED", "BENCHMARK_ONLY", "URBAN_ONLY", "NOT_COMPARABLE")
TRADES = ("FLOOR", "CEILING", "SKIRTING", "BLOCKWORK", "PLASTER", "PAINT", "TILE", "WATERPROOFING", "ALUMINIUM",
          "DOORS", "WINDOWS", "FOOTING CONCRETE", "STRUCTURAL CONCRETE", "REBAR", "MARBLE / STAIRS", "OTHER")
COMPLETE, BLOCKED = "COMPUTED_SHADOW_COMPLETE", "BLOCKED"
STOREY_SHEET = {"GF": "GF_ROOF_SLAB", "1F": "1F_ROOF_SLAB", "ROOF": "2F_ROOF_SLAB"}
STOREY_BAND = {"FOUNDATION": "FOUNDATION", "GF": "GROUND FLOOR", "1F": "1ST FLOOR", "ROOF": "2ND FLOOR & TOP"}


def jl(p):
    return json.loads(Path(p).read_text())


def digest(o) -> str:
    return hashlib.sha256(json.dumps(o, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def r6(x):
    return None if x is None else round(float(x), 6)


def tolerances() -> dict:
    return jl(REC_PATH)["tolerances_declared"]


# ------------------------------------------------------------------ difference metrics
def pct_diff(u, b):
    """(Urban - Benchmark) / Benchmark - descriptive only"""
    if u is None or b is None or b == 0:
        return None
    return (u - b) / b


def band(kind, u, b, tol):
    if u is None or b is None:
        return None
    if kind == "count":
        return "EXACT" if u == b else "DIFFERENT"
    d, p = abs(u - b), abs(pct_diff(u, b) or 0.0) * 100
    if d < 1e-9:
        return "EXACT"
    if kind == "linear":
        if d <= max(0.05, 0.01 * abs(b)):
            return "CLOSE"
        return "INVESTIGATE" if d <= max(0.10, 0.02 * abs(b)) else "MATERIAL"
    a = tol["area_volume"]
    if p <= a["close_match_pct"]:
        return "CLOSE"
    return "INVESTIGATE" if p <= a["investigate_pct"] else "MATERIAL"


class Book:
    """collects comparison rows; enforces the comparability gate and the class vocabulary"""

    def __init__(self, tol):
        self.tol, self.rows = tol, []

    def add(self, family, trade, item, *, unit, urban, bench, comparability, confidence, kind="area",
            primary=None, secondary=(), floor=None, urban_ref=None, urban_status=None, bench_refs=(),
            method=None, fix=None, note=None, extra=None):
        assert comparability in COMPARABILITY, comparability
        assert trade in TRADES, trade
        row = {"id": f"CMP-{len(self.rows) + 1:03d}", "family": family, "trade": trade, "item": item,
               "floor": floor, "unit": unit, "urban_qty": r6(urban) if isinstance(urban, (int, float)) else urban,
               "urban_ref": urban_ref, "urban_status": urban_status,
               "bench_qty": r6(bench) if isinstance(bench, (int, float)) else bench, "bench_refs": list(bench_refs),
               "comparability": comparability, "confidence": confidence, "metric_kind": kind,
               "abs_diff": None, "pct_diff": None, "band": None, "method": method, "proposed_fix": fix, "note": note}
        if comparability == "EXACT_COMPARABLE" and isinstance(urban, (int, float)) and isinstance(bench, (int, float)):
            row["abs_diff"] = r6(urban - bench)
            p = pct_diff(urban, bench)
            row["pct_diff"] = None if p is None else round(p * 100, 4)
            row["band"] = band(kind, urban, bench, self.tol)
            if primary is None:
                primary = ("MATCH_WITHIN_TOLERANCE" if row["band"] in ("EXACT", "CLOSE")
                           else "UNRESOLVED_REQUIRES_REVIEW")
        if primary is None:
            primary = {"URBAN_BLOCKED": "SOURCE_GAP", "BENCHMARK_ONLY": "SCOPE_MISMATCH", "URBAN_ONLY": "SCOPE_MISMATCH",
                       "SCOPE_MISMATCH": "SCOPE_MISMATCH"}.get(comparability, "UNRESOLVED_REQUIRES_REVIEW")
        assert primary in CLASSES and all(s in CLASSES for s in secondary), (primary, secondary)
        row["primary_class"], row["secondary_classes"] = primary, list(secondary)
        if extra:
            row.update(extra)
        self.rows.append(row)
        return row


def _man(N, **kw):
    return [n for n in N if all(n.get(k) == v for k, v in kw.items())]


def _ids(rs):
    return [r["id"] for r in rs]


# ------------------------------------------------------------------ structure
def footings(bk, R3, N):
    fr = R3["FOOTING_REGISTER"]
    summ = {s["type"]: s for s in fr["type_summary"]}
    lib = R3["STRUCT_TYPE_LIBRARY"]["footings"]
    occ = fr["occurrences"]
    man = defaultdict(list)
    for n in _man(N, trade="FOOTING", row_kind="DETAIL"):
        man[n["subtrade"]].append(n)
    rows, l4l_u, l4l_b = [], 0.0, 0.0
    for t in sorted(lib, key=lambda x: (len(x), x)):
        s, ms = summ.get(t), man.get(t, [])
        uc = s["count_tagged"] if s else 0
        mc = sum(m["count"] for m in ms) if ms else 0
        u_occ = [o for o in occ if o["type"] == t]
        conflict = [o for o in u_occ if o["geometry_state"].startswith("SOURCE_CONFLICT")]
        confirmed = [o for o in u_occ if o["status"] == COMPLETE]
        if uc == mc:
            prim, sec = "MATCH_WITHIN_TOLERANCE", []
            why = "same count" if uc else "both sides: type not tagged / not measured"
        elif uc > mc and conflict and uc - len(conflict) == mc:
            prim, sec = "SOURCE_GAP", ["MANUAL_ASSUMPTION_NOT_IN_SOURCE"]
            why = (f"the {len(conflict)} Urban tag(s) the manual does not count lie in a SOURCE_CONFLICT outline "
                   "(blocked in Urban, OQ3-S1); the manual resolved the conflict by an assumption")
        elif uc > mc and len(confirmed) == uc and all(o["geometry_state"] == "GEOMETRY_CONFIRMED" for o in u_occ):
            prim, sec = "BENCHMARK_POSSIBLE_ERROR", ["DUPLICATE_OR_OMITTED_ITEM"]
            why = (f"every Urban tag of {t} sits in its own outline drawn whole at the scheduled size "
                   f"({uc} separate outlines); the manual counts {mc}")
        else:
            prim, sec, why = "UNRESOLVED_REQUIRES_REVIEW", [], "count difference without a source explanation"
        dims_b = sorted([d for d in (ms[0]["dims"]["B"], ms[0]["dims"]["C"]) if d is not None]) if ms else None
        dims_s = sorted([lib[t]["L_cm"] / 100, lib[t]["W_cm"] / 100])
        dims_ok = (dims_b == dims_s and abs(ms[0]["dims"]["D"] - lib[t]["H_cm"] / 100) < 1e-9) if ms else None
        c = bk.add("FOOTING_COUNT", "FOOTING CONCRETE", f"{t} count", unit="nr", urban=uc, bench=mc,
                   comparability="EXACT_COMPARABLE", confidence="HIGH", kind="count", primary=prim, secondary=sec,
                   floor="FOUNDATION", urban_ref=f"FOOTING_REGISTER.type_summary[{t}].count_tagged",
                   urban_status=s["status"] if s else None, bench_refs=_ids(ms), note=why,
                   extra={"urban_computed": s["count_computed"] if s else 0, "urban_conflict": len(conflict),
                          "schedule_dims_m": [lib[t]["L_cm"] / 100, lib[t]["W_cm"] / 100, lib[t]["H_cm"] / 100],
                          "manual_dims_m": [ms[0]["dims"][k] for k in "BCD"] if ms else None,
                          "manual_dims_equal_schedule": dims_ok})
        uv = s["m3_total_computed"] if s and s["count_computed"] else None
        mv = sum(m["qty"] for m in ms) if ms else None
        if uc == 0 and mc == 0:
            comp, uv, mv = "EXACT_COMPARABLE", 0.0, 0.0             # not tagged on the plan, not in the manual
        elif s and s["status"] == BLOCKED or (s and not s["count_computed"]):
            comp = "URBAN_BLOCKED" if mv is not None else "NOT_COMPARABLE"
        elif s and s["status"] == COMPLETE:
            comp = "EXACT_COMPARABLE" if mv is not None else "URBAN_ONLY"
        elif s:
            comp = "PARTIAL_SCOPE_COMPARABLE"
        else:
            comp = "BENCHMARK_ONLY" if mv else "EXACT_COMPARABLE"
            if not mv:
                uv, mv = 0.0, 0.0                               # not tagged on the plan, not in the manual
        vrow = bk.add("FOOTING_VOLUME", "FOOTING CONCRETE", f"{t} concrete", unit="m3", urban=uv, bench=mv,
                      comparability=comp, confidence="HIGH" if comp == "EXACT_COMPARABLE" else "MEDIUM",
                      kind="volume", primary=(None if prim == "MATCH_WITHIN_TOLERANCE" and comp == "EXACT_COMPARABLE"
                                              else prim if comp in ("EXACT_COMPARABLE", "PARTIAL_SCOPE_COMPARABLE") else None),
                      secondary=sec if prim != "MATCH_WITHIN_TOLERANCE" else [], floor="FOUNDATION",
                      urban_ref=f"FOOTING_REGISTER.type_summary[{t}].m3_total_computed",
                      urban_status=s["status"] if s else None, bench_refs=_ids(ms),
                      note=None if comp == "EXACT_COMPARABLE" else
                      ("Urban computed subset only" if comp == "PARTIAL_SCOPE_COMPARABLE" else why),
                      extra={"urban_each_m3": s["m3_each"] if s else None,
                             "manual_each_m3": r6(ms[0]["unit_volume"]) if ms else None})
        if s and s["status"] == COMPLETE and uc == mc and mv is not None:
            l4l_u += uv
            l4l_b += mv
        rows.append({"type": t, "urban_tags": uc, "urban_computed": s["count_computed"] if s else 0,
                     "urban_conflict": len(conflict), "manual_count": mc, "difference": uc - mc,
                     "source_status": s["status"] if s else "NOT_TAGGED_ON_PLAN", "count_class": prim,
                     "urban_m3": r6(uv), "manual_m3": r6(mv), "schedule_each_m3": s["m3_each"] if s else None,
                     "manual_each_m3": r6(ms[0]["unit_volume"]) if ms else None, "manual_dims_equal_schedule": dims_ok,
                     "count_row": c["id"], "volume_row": vrow["id"], "manual_rows": _ids(ms)})
    # the F occurrences outside the combined outline are a like-for-like pair too (4 computed vs 4 manual)
    fsum = summ["F"]
    fman = sum(m["qty"] for m in man.get("F", []))
    fmc = sum(m["count"] for m in man.get("F", []))
    total_u = sum(s["m3_total_computed"] for s in summ.values())
    total_b = sum(m["qty"] for ms in man.values() for m in ms)
    ll = bk.add("FOOTING_TOTAL", "FOOTING CONCRETE", "footing concrete, like-for-like (types with equal counts, "
                "Urban complete)", unit="m3", urban=l4l_u, bench=l4l_b, comparability="EXACT_COMPARABLE",
                confidence="HIGH", kind="volume", floor="FOUNDATION",
                urban_ref="FOOTING_REGISTER.type_summary (COMPLETE types with equal counts)",
                note="excludes F (Urban 4 of 5 computed), F3 (count differs), F10 (Urban blocked)")
    llf = bk.add("FOOTING_TOTAL", "FOOTING CONCRETE", "footing concrete, like-for-like incl. the 4 computed F",
                 unit="m3", urban=l4l_u + fsum["m3_total_computed"], bench=l4l_b + fman,
                 comparability="EXACT_COMPARABLE" if fsum["count_computed"] == fmc else "PARTIAL_SCOPE_COMPARABLE",
                 confidence="MEDIUM", kind="volume", floor="FOUNDATION",
                 note="the 4 computed F vs the manual 4 F: identity of which four is not stated by the manual")
    tot = bk.add("FOOTING_TOTAL", "FOOTING CONCRETE", "footing concrete, all types as reported", unit="m3",
                 urban=total_u, bench=total_b, comparability="PARTIAL_SCOPE_COMPARABLE", confidence="MEDIUM",
                 kind="volume", floor="FOUNDATION", primary="SCOPE_MISMATCH",
                 secondary=["BENCHMARK_POSSIBLE_ERROR", "SOURCE_GAP"],
                 note="Urban 25 of 27 tags; the manual counts F3 once and F10 at the schedule size",
                 extra={"descriptive_difference_m3": r6(total_u - total_b)})
    return {"types": rows, "like_for_like": [ll["id"], llf["id"]], "all_reported": tot["id"],
            "urban_total_m3": r6(total_u), "manual_total_m3": r6(total_b),
            "like_for_like_m3": {"urban": r6(l4l_u), "manual": r6(l4l_b)}}


def f_f10(R3, N):
    fr = R3["FOOTING_REGISTER"]
    fx = fr["forensic"][0]
    lib = R3["STRUCT_TYPE_LIBRARY"]["footings"]
    occ = fr["occurrences"]
    f_out = [o for o in occ if o["type"] == "F" and o["geometry_state"] != "SOURCE_CONFLICT_COMBINED_OUTLINE"]
    mf = _man(N, trade="FOOTING", subtrade="F", row_kind="DETAIL")
    m10 = _man(N, trade="FOOTING", subtrade="F10", row_kind="DETAIL")
    F, F10 = lib["F"], lib["F10"]
    vF = F["L_cm"] * F["W_cm"] * F["H_cm"] / 1e6
    v10 = F10["L_cm"] * F10["W_cm"] * F10["H_cm"] / 1e6
    L, W = fx["drawn_mm"][0] / 1000, fx["drawn_mm"][1] / 1000
    overlap_len = F10["L_cm"] / 100 + F["L_cm"] / 100 - L
    approaches = {
        "A_SEPARATE_SCHEDULED": {"m3": r6(vF + v10), "basis": "F10 + F, each at its schedule size"},
        "B_COMBINED_ENVELOPE_AT_F10_DEPTH": {"m3": r6(L * W * F10["H_cm"] / 100),
                                             "basis": f"drawn envelope {L} x {W} x F10 depth"},
        "B_COMBINED_ENVELOPE_AT_F_DEPTH": {"m3": r6(L * W * F["H_cm"] / 100), "basis": f"drawn envelope {L} x {W} x F depth"},
        "C_OVERLAP_DEDUCTED": {"m3": r6(vF + v10 - overlap_len * min(F["W_cm"], F10["W_cm"]) / 100 * min(F["H_cm"], F10["H_cm"]) / 100),
                               "basis": f"A minus the plan overlap {r6(overlap_len)} m x the narrower width x the "
                                        "shallower depth (geometry assumed: F inside the east end)"},
        "E_F_ABSORBED_INTO_F10": {"m3": r6(v10), "basis": "F10 at its schedule size; the F mark not counted"},
    }
    man_f10 = sum(m["qty"] for m in m10)
    manual = {"F_rows": _ids(mf), "F_count": sum(m["count"] for m in mf), "F10_rows": _ids(m10),
              "F10_count": sum(m["count"] for m in m10), "F10_m3": r6(man_f10),
              "F10_dims_m": [m10[0]["dims"][k] for k in "BCD"] if m10 else None,
              "location_stated": False}
    matches = [k for k, v in approaches.items() if abs(v["m3"] - man_f10) < 1e-6]
    f_counts_match = manual["F_count"] == len(f_out)
    method = ("E_F_ABSORBED_INTO_F10" if "E_F_ABSORBED_INTO_F10" in matches and f_counts_match else "UNRESOLVED")
    return {"SCHEMA": "URBAN_ALSENAN_B1_F_F10_FORENSIC_V1",
            "source": {"outline": fx["outline"], "drawn_mm": fx["drawn_mm"],
                       "marks_inside": [{"type": m["type"], "key": m["key"]} for m in fx["marks_inside"]],
                       "explicit_local_dimensions": fx["explicit_local_dimensions"],
                       "scheduled_depths_cm": {"F": F["H_cm"], "F10": F10["H_cm"]},
                       "F_marks_outside_the_outline": len(f_out)},
            "urban_a3": {"state": "BLOCKED (SOURCE_CONFLICT_COMBINED_OUTLINE)", "owner_question": "OQ3-S1"},
            "manual": manual, "approaches": approaches, "approaches_matching_manual_m3": matches,
            "manual_method": method,
            "manual_method_confidence": "MEDIUM" if method != "UNRESOLVED" else "LOW",
            "reasoning": [f"the manual F count ({manual['F_count']}) equals the Urban F tags OUTSIDE the combined outline "
                          f"({len(f_out)})" if f_counts_match else "the manual F count does not equal the F tags outside",
                          f"the manual F10 volume ({r6(man_f10)} m3) equals approach(es) {matches}",
                          "the manual gives no location, so which F it dropped is inferred, not read"],
            "against_source": ["two independent type marks exist inside one outline", f"one {L} x {W} m envelope is drawn",
                               "no explicit local dimension binds either mark", "the scheduled depths differ (30 / 50 cm)",
                               "the manual value equals neither the drawn envelope nor two scheduled footings"],
            "classification": {"primary": "MANUAL_ASSUMPTION_NOT_IN_SOURCE", "secondary": ["SOURCE_GAP"]},
            "urban_change": "NONE - Phase A3 stays blocked; OQ3-S1 carried forward with this evidence"}


def straps(bk, R3, N):
    sr = {r["type"]: r for r in R3["STRAP_BEAM_REGISTER"]["rows"]}
    out = []
    for t in sorted(sr):
        u = sr[t]
        ms = _man(N, trade="STRAP_BEAM", subtrade=t, row_kind="DETAIL")
        if not ms:
            continue
        m = ms[0]
        d = [m["dims"][k] for k in "BCD"]
        sec = sorted([u["B_cm"] / 100, u["D_cm"] / 100])
        rest = list(d)
        sec_ok = True
        for v in sec:
            hit = next((x for x in rest if abs(x - v) < 1e-9), None)
            if hit is None:
                sec_ok = False
                break
            rest.remove(hit)
        mlen = rest[0] if sec_ok and len(rest) == 1 else None
        dep = u.get("support_outline_conflicts") or []
        lr = bk.add("STRAP", "FOOTING CONCRETE", f"{t} clear length", unit="m", urban=u["length_m"], bench=mlen,
                    comparability="EXACT_COMPARABLE" if mlen is not None else "NOT_COMPARABLE", confidence="HIGH",
                    kind="linear", floor="FOUNDATION", urban_ref=f"STRAP_BEAM_REGISTER.rows[{t}].length_m",
                    urban_status=u["status"], bench_refs=[m["id"]],
                    method="manual length = the third dimension after the scheduled B and D are matched",
                    note=("one end on the combined outline H6939: the manual measured to the same drawn face"
                          if dep else None))
        vr = bk.add("STRAP", "FOOTING CONCRETE", f"{t} concrete", unit="m3", urban=u["volume_m3"], bench=m["qty"],
                    comparability="EXACT_COMPARABLE", confidence="HIGH", kind="volume", floor="FOUNDATION",
                    urban_ref=f"STRAP_BEAM_REGISTER.rows[{t}].volume_m3", urban_status=u["status"], bench_refs=[m["id"]])
        out.append({"type": t, "B_cm": u["B_cm"], "D_cm": u["D_cm"], "manual_dims_m": d, "section_equal": sec_ok,
                    "urban_length_m": u["length_m"], "manual_length_m": mlen, "urban_m3": u["volume_m3"],
                    "manual_m3": m["qty"], "support_outline_conflicts": dep, "length_row": lr["id"], "volume_row": vr["id"],
                    "manual_method": "length x B x D, length rounded to 0.05 m"})
    return out


def foundation(bk, R3, N, ft, st):
    cov = {n["subtrade"]: n for n in _man(N, file="RC", row_kind="COVER_TOTAL")}
    blind = _man(N, trade="BLINDING", row_kind="DETAIL")
    perim = _man(N, trade="GROUND_BEAM", subtrade="PERIMETER_STRAP", row_kind="DETAIL")
    gslab = _man(N, trade="OTHER_CONCRETE", subtrade="GROUND_SLAB", row_kind="DETAIL")
    gbeam = [n for n in _man(N, trade="GROUND_BEAM", row_kind="DETAIL") if n["subtrade"] != "PERIMETER_STRAP"]
    necks = [n for n in _man(N, trade="COLUMN", row_kind="DETAIL") if n["subtrade"].startswith("NECK")]
    walls = _man(N, trade="OTHER_CONCRETE", subtrade="ELEVATOR_WALL", row_kind="DETAIL")
    straps_m = sum(n["qty"] for n in _man(N, trade="STRAP_BEAM", row_kind="DETAIL"))
    straps_u = sum(s["urban_m3"] for s in st)
    ct = R3["CONCRETE_REGISTER"]["totals"]
    comp = []
    comp.append(bk.add("FOUNDATION", "STRUCTURAL CONCRETE", "strap beams SB1-SB3", unit="m3", urban=straps_u,
                       bench=straps_m, comparability="EXACT_COMPARABLE", confidence="HIGH", kind="volume",
                       floor="FOUNDATION", urban_ref="CONCRETE_REGISTER.totals.straps_m3"))
    for label, rows, urban_item in (("blinding (plain concrete)", blind, None),
                                    ("perimeter strap / edge beam", perim, None),
                                    ("concrete ground slab", gslab, None),
                                    ("elevator walls (neck level)", walls, None)):
        comp.append(bk.add("FOUNDATION", "STRUCTURAL CONCRETE", label, unit="m3", urban=None,
                           bench=sum(n["qty"] for n in rows), comparability="BENCHMARK_ONLY", confidence="HIGH",
                           floor="FOUNDATION", bench_refs=_ids(rows), primary="SCOPE_MISMATCH",
                           secondary=["TRADE_RULE_GAP"] if "blinding" in label or "slab" in label else [],
                           note="no Urban item in Phase A3"))
    comp.append(bk.add("FOUNDATION", "STRUCTURAL CONCRETE", "ground beams (longitudinal / transverse / fence)",
                       unit="m3", urban=None, bench=sum(n["qty"] for n in gbeam), comparability="URBAN_BLOCKED",
                       confidence="MEDIUM", floor="FOUNDATION", bench_refs=_ids(gbeam),
                       urban_ref="CONCRETE_REGISTER S3-GB-VOL (BLOCKED_HEIGHT)", primary="SOURCE_GAP",
                       secondary=["MANUAL_ASSUMPTION_NOT_IN_SOURCE"],
                       note="manual lengths are round totals (" + ", ".join(f"{n['dims']['C']:g}" for n in gbeam) + " m)"))
    comp.append(bk.add("FOUNDATION", "STRUCTURAL CONCRETE", "column necks", unit="m3", urban=None,
                       bench=sum(n["qty"] for n in necks), comparability="URBAN_BLOCKED", confidence="HIGH",
                       floor="FOUNDATION", bench_refs=_ids(necks), urban_ref="CONCRETE_REGISTER S3-COL-VOL",
                       primary="HEIGHT_EVIDENCE_GAP", secondary=["MANUAL_ASSUMPTION_NOT_IN_SOURCE"],
                       note="manual neck height 1.5 m (no founding level in the source)"))
    footing_cover = next((n for n in _man(N, file="RC", row_kind="COVER_TOTAL") if n["trade"] == "FOOTING"), None)
    return {"decomposition": [c["id"] for c in comp],
            "manual_footing_sheet_total_m3": r6(footing_cover["qty"]) if footing_cover else None,
            "manual_footing_sheet_includes": ["footings", "strap beams", "perimeter strap"],
            "urban_deterministic_m3": ct["deterministic_total_m3"],
            "manual_like_for_like_m3": r6(ft["manual_total_m3"] + straps_m),
            "note": "Urban footings + straps is compared only against manual footings + straps; blinding, perimeter "
                    "strap, ground slab, ground beams, necks and elevator walls are separate scopes"}


def columns(bk, R3, N, VE):
    col = R3["COLUMN_REGISTER"]
    lib = col["library"]
    marks = col["marks"]
    printed = col["printed_size_labels"]["printed"]
    rows = []
    man = defaultdict(lambda: defaultdict(list))
    for n in _man(N, trade="COLUMN", row_kind="DETAIL"):
        kind, t = n["subtrade"].split(":")
        man["FOUNDATION" if kind == "NECK" else n["floor"]][t].append(n)
    fsize = {t: f"{int(v['FOUNDATION']['B_cm'])}x{int(v['FOUNDATION']['D_cm'])}" for t, v in lib.items()
             if v["FOUNDATION"]["B_cm"]}
    size_types = defaultdict(list)
    for t, s in fsize.items():
        size_types[s].append(t)
    for t in sorted(set(lib) | set(man["FOUNDATION"]), key=lambda x: (x[:1], len(x), x)):
        uc = marks.get(t, 0)
        ms = man["FOUNDATION"].get(t, [])
        mc = sum(m["count"] for m in ms)
        if t not in lib:
            bk.add("COLUMN_COUNT", "STRUCTURAL CONCRETE", f"{t} neck count", unit="nr", urban=None, bench=mc,
                   comparability="BENCHMARK_ONLY", confidence="HIGH", floor="FOUNDATION", bench_refs=_ids(ms),
                   note="type not in the column schedule")
            continue
        corr = None
        if t in fsize and len(size_types[fsize[t]]) == 1:
            corr = printed.get(fsize[t], 0)
        if uc == mc:
            prim, why = "MATCH_WITHIN_TOLERANCE", "same count"
        elif corr is not None and corr == uc:
            prim, why = "BENCHMARK_POSSIBLE_ERROR", (f"{fsize[t]} is unique to {t} in the foundation band and the plan "
                                                    f"prints it {corr} time(s) = the Urban mark count")
        else:
            prim, why = "UNRESOLVED_REQUIRES_REVIEW", "count differs; no independent printed-size corroboration"
        cr = bk.add("COLUMN_COUNT", "STRUCTURAL CONCRETE", f"{t} neck count", unit="nr", urban=uc, bench=mc,
                    comparability="EXACT_COMPARABLE", confidence="HIGH", kind="count", primary=prim,
                    floor="FOUNDATION", urban_ref=f"COLUMN_REGISTER.marks[{t}]", bench_refs=_ids(ms), note=why,
                    extra={"printed_label_corroboration": corr})
        rows.append({"type": t, "urban_marks": uc, "manual_necks": mc, "class": prim, "row": cr["id"],
                     "printed_corroboration": corr})
    sections = []
    for storey, band_name in STOREY_BAND.items():
        for t, ms in sorted(man[storey].items()):
            m = ms[0]
            sched = lib.get(t, {}).get(band_name, {})
            mb = sorted([m["dims"]["B"], m["dims"]["C"]])
            if not sched or sched.get("B_cm") is None:
                state = "TYPE_NOT_IN_BAND" if t in lib else "TYPE_NOT_IN_SCHEDULE"
                sections.append({"storey": storey, "type": t, "manual_m": mb, "schedule_cm": None, "state": state,
                                 "manual_height_m": m["dims"]["D"], "manual_count": sum(x["count"] for x in ms)})
                continue
            ss = sorted([sched["B_cm"] / 100, sched["D_cm"] / 100])
            ok = all(abs(a - b) < 1e-9 for a, b in zip(mb, ss))
            fs = fsize.get(t)
            label_support = storey == "FOUNDATION" and fs and len(size_types[fs]) == 1 and printed.get(fs, 0) > 0
            if not ok:
                bk.add("COLUMN_SECTION", "STRUCTURAL CONCRETE", f"{t} section ({storey})", unit="m",
                       urban=f"{ss[0]:g} x {ss[1]:g}", bench=f"{mb[0]:g} x {mb[1]:g}", comparability="EXACT_COMPARABLE",
                       confidence="HIGH" if label_support else "MEDIUM", kind="section",
                       primary="BENCHMARK_POSSIBLE_ERROR", floor=storey,
                       urban_ref=f"COLUMN_REGISTER.library[{t}][{band_name}]", bench_refs=_ids(ms),
                       note="schedule section is the authority" + (f"; the plan prints {fs}" if label_support else ""))
            sections.append({"storey": storey, "type": t, "manual_m": mb, "schedule_m": ss, "equal": ok,
                             "manual_height_m": m["dims"]["D"], "manual_count": sum(x["count"] for x in ms)})
    f2f = {x["from"]: x["floor_to_floor_m"] for x in VE["floor_to_floor"]}
    heights = {}
    for storey, f in (("GF", "GF"), ("1F", "1F"), ("ROOF", "2F")):
        hs = sorted({m["dims"]["D"] for ms in man[storey].values() for m in ms})
        heights[storey] = {"manual_heights_m": hs, "floor_to_floor_m": f2f.get(f),
                           "floor_to_floor_minus_manual_m": [r6(f2f[f] - h) for h in hs] if f in f2f else None}
    neck_h = sorted({m["dims"]["D"] for ms in man["FOUNDATION"].values() for m in ms})
    method = {"neck_height_m": neck_h, "storeys": heights,
              "inferred": "column height = floor-to-floor minus 0.75 m (the depth of the 75 cm main beams), i.e. to "
                          "the beam soffit; neck height 1.5 m is not in the source",
              "classes": {"necks": "MANUAL_ASSUMPTION_NOT_IN_SOURCE", "storeys": "MANUAL_METHOD_DIFFERENCE"}}
    storey_counts = []
    for storey, band_name in (("GF", "GROUND FLOOR"), ("1F", "1ST FLOOR"), ("ROOF", "2ND FLOOR & TOP")):
        mc = sum(m["count"] for ms in man[storey].values() for m in ms)
        corr = col["storeys"][band_name]["corroboration"] or {}
        r = bk.add("COLUMN_STOREY", "STRUCTURAL CONCRETE", f"columns drawn / counted ({storey})", unit="nr",
                   urban=corr.get("drawn_count"), bench=mc, comparability="PARTIAL_SCOPE_COMPARABLE",
                   confidence="LOW", floor=storey, primary="SCOPE_MISMATCH",
                   urban_ref=f"COLUMN_REGISTER.storeys[{band_name}].corroboration.drawn_count",
                   note="Urban: untyped outlines on the storey slab plan; manual: typed occurrences incl. planted "
                        "columns PC")
        storey_counts.append({"storey": storey, "urban_drawn_outlines": corr.get("drawn_count"),
                              "urban_schedule_expectation": corr.get("expected_count"), "manual_count": mc,
                              "row": r["id"]})
    vols = {s: r6(sum(m["qty"] for ms in man[s].values() for m in ms)) for s in ("FOUNDATION", "GF", "1F", "ROOF")}
    bk.add("COLUMN_VOLUME", "STRUCTURAL CONCRETE", "column concrete (necks + storeys)", unit="m3", urban=None,
           bench=sum(v for v in vols.values() if v), comparability="URBAN_BLOCKED", confidence="MEDIUM",
           primary="HEIGHT_EVIDENCE_GAP", secondary=["MANUAL_METHOD_DIFFERENCE"],
           urban_ref="CONCRETE_REGISTER S3-COL-VOL (BLOCKED_HEIGHT)", note="DIAGNOSTIC ONLY - never fed to Urban")
    return {"neck_counts": rows, "sections": sections, "height_method": method, "storey_counts": storey_counts,
            "manual_volumes_m3": vols}


def beams(bk, R3, N):
    lib_s = R3["STRUCT_TYPE_LIBRARY"]["simple_beams"]
    lib_c = R3["STRUCT_TYPE_LIBRARY"]["continuous_beams"]
    br = {r["type"]: r for r in R3["BEAM_REGISTER"]["rows"]}
    man = [n for n in _man(N, trade="BEAM", row_kind="DETAIL") if n["subtrade"]]
    cb = []
    for n in man:
        t = n["subtrade"]
        if not t.startswith("CB"):
            continue
        u = br.get(t)
        s = lib_c.get(t)
        ml = n["dims"]["D"]
        sec_ok = s is not None and sorted([n["dims"]["B"], n["dims"]["C"]]) == sorted([s["B_cm"] / 100, s["H_cm"] / 100])
        r = bk.add("BEAM_CB", "STRUCTURAL CONCRETE", f"{t} length ({n['label_ar']})", unit="m",
                   urban=u["length_m"] if u else None, bench=ml, comparability="EXACT_COMPARABLE" if u else "BENCHMARK_ONLY",
                   confidence="MEDIUM", kind="linear", floor=n["floor"],
                   urban_ref=f"BEAM_REGISTER.rows[{t}].length_m (sum of schedule spans, centre line)",
                   urban_status=u["status"] if u else None, bench_refs=[n["id"]],
                   method="Urban: printed schedule spans (centre line); manual: one length per beam read from the plan",
                   primary=None, extra={"section_equal": sec_ok, "schedule_spans_m": s["spans_m"] if s else None})
        if r["band"] not in (None, "EXACT", "CLOSE"):
            r["primary_class"], r["secondary_classes"] = "MANUAL_METHOD_DIFFERENCE", ["STRUCTURAL_LENGTH_DEFECT"]
            r["note"] = "basis differs (schedule span vs drawn length); not an Urban error by itself"
        cb.append({"type": t, "manual_label": n["label_ar"], "floor": n["floor"], "manual_length_m": ml,
                   "urban_schedule_length_m": u["length_m"] if u else None, "pct": r["pct_diff"],
                   "section_equal": sec_ok, "row": r["id"],
                   "urban_marks_per_sheet": u["marks_per_sheet"] if u else None})
    simple = defaultdict(lambda: {"rows": 0, "occurrences": 0.0, "length_m": 0.0, "m3": 0.0, "ids": [], "sections": set()})
    for n in man:
        t = n["subtrade"]
        if t.startswith("CB"):
            continue
        k = (n["floor"], t)
        simple[k]["rows"] += 1
        simple[k]["occurrences"] += n["count"]
        simple[k]["length_m"] += n["count"] * n["dims"]["D"]
        simple[k]["m3"] += n["qty"]
        simple[k]["ids"].append(n["id"])
        simple[k]["sections"].add((n["dims"]["B"], n["dims"]["C"]))
    srows = []
    for (fl, t), v in sorted(simple.items()):
        s = lib_s.get(t)
        u = br.get(t)
        sec = sorted(v["sections"])
        sec_ok = s is not None and all(abs(b - s["B_cm"] / 100) < 1e-9 and abs(d - s["D_cm"] / 100) < 1e-9 for b, d in sec)
        um = (u["marks_per_sheet"].get(STOREY_SHEET.get(fl), 0) if u else None)
        if not sec_ok:
            bk.add("BEAM_SECTION", "STRUCTURAL CONCRETE", f"{t} section ({fl})", unit="m",
                   urban=f"{s['B_cm'] / 100:g} x {s['D_cm'] / 100:g}" if s else None,
                   bench=" / ".join(f"{b:g} x {d:g}" for b, d in sec), comparability="EXACT_COMPARABLE" if s else "BENCHMARK_ONLY",
                   confidence="HIGH", kind="section", primary="BENCHMARK_POSSIBLE_ERROR", floor=fl, bench_refs=v["ids"])
        srows.append({"floor": fl, "type": t, "manual_rows": v["rows"], "manual_occurrences": v["occurrences"],
                      "manual_length_m": r6(v["length_m"]), "manual_m3": r6(v["m3"]), "section_equal_schedule": sec_ok,
                      "urban_marks_on_sheet": um, "urban_length_m": None, "ids": v["ids"]})
    measured = [(k, b) for k, v in R3["BEAM_REGISTER"]["sheets"].items() for b in v["simple_bands"]
                if b["state"].startswith("MEASURED")]
    for sheet, b in measured:
        fl = next(f for f, s in STOREY_SHEET.items() if s == sheet)
        v = simple.get((fl, b["type"]))
        if not v:
            continue
        bk.add("BEAM_SIMPLE", "STRUCTURAL CONCRETE", f"{b['type']} length ({fl}, the one band Urban measured)",
               unit="m", urban=b["length_m"], bench=r6(v["length_m"]),
               comparability="PARTIAL_SCOPE_COMPARABLE", confidence="LOW", floor=fl,
               primary="UNRESOLVED_REQUIRES_REVIEW", secondary=["STRUCTURAL_LENGTH_DEFECT", "MANUAL_METHOD_DIFFERENCE"],
               urban_ref=f"BEAM_REGISTER.sheets[{sheet}].simple_bands", bench_refs=v["ids"],
               note="Urban measured the clear span between support outlines of one band; the manual length may be "
                    "gross (support to support centre or over the supports) - basis not stated")
    vol = defaultdict(float)
    for n in man:
        vol[n["floor"]] += n["qty"]
    bk.add("BEAM_VOLUME", "STRUCTURAL CONCRETE", "beam concrete (all floors)", unit="m3", urban=None,
           bench=sum(vol.values()), comparability="URBAN_BLOCKED", confidence="MEDIUM",
           primary="STRUCTURAL_VOLUME_METHOD_DIFFERENCE", secondary=["SOURCE_GAP"],
           urban_ref="CONCRETE_REGISTER S3-BM-VOL", note="manual: length x B x full D (no slab deduction); Urban "
                                                         "blocks clear lengths - DIAGNOSTIC ONLY")
    return {"continuous": cb, "simple": srows, "manual_m3_by_floor": {k: r6(v) for k, v in sorted(vol.items())},
            "method": {"length_basis": "one length per occurrence row read from the plan (continuous beams within "
                                       "-6.4 % .. +12.1 % of the schedule spans), not the printed spans",
                       "depth": "full schedule D (slab thickness not deducted)",
                       "counts": "a manual row is a length run, not a plan mark - counts are not comparable"}}


def slabs(bk, R3, N):
    out = []
    for fl, sheet in STOREY_SHEET.items():
        u = R3["SLAB_REGISTER"]["slabs"].get(sheet, {})
        ms = [n for n in _man(N, trade="SLAB", row_kind="DETAIL") if n["floor"] == fl]
        ut = sorted(int(k) for k in (u.get("thickness_marks_cm") or {}))
        mt = sorted({int(round(n["dims"]["D"] * 100)) for n in ms})
        comp = "EXACT_COMPARABLE" if ut and mt else "NOT_COMPARABLE"
        prim = "MATCH_WITHIN_TOLERANCE" if ut == mt else ("SCOPE_MISMATCH" if set(ut) <= set(mt) else "UNRESOLVED_REQUIRES_REVIEW")
        bk.add("SLAB_THICKNESS", "STRUCTURAL CONCRETE", f"slab thickness set ({fl})", unit="cm",
               urban=", ".join(map(str, ut)), bench=", ".join(map(str, mt)), comparability=comp, confidence="HIGH",
               kind="set", primary=prim, floor=fl, urban_ref=f"SLAB_REGISTER.slabs[{sheet}].thickness_marks_cm",
               bench_refs=_ids(ms))
        area = sum(n["dims"]["B"] for n in ms)
        bk.add("SLAB_VOLUME", "STRUCTURAL CONCRETE", f"slab concrete ({fl})", unit="m3", urban=None,
               bench=sum(n["qty"] for n in ms), comparability="URBAN_BLOCKED", confidence="MEDIUM",
               primary="SOURCE_GAP", secondary=["TRADE_RULE_GAP"], floor=fl, bench_refs=_ids(ms),
               urban_ref=f"SLAB_REGISTER.slabs[{sheet}] (SLAB_OUTLINE_NOT_ESTABLISHED)",
               note=f"manual area {r6(area)} m2 as a single number per part; no void or beam deduction applied")
        out.append({"floor": fl, "urban_thickness_cm": ut, "manual_thickness_cm": mt, "manual_area_m2": r6(area),
                    "manual_m3": r6(sum(n["qty"] for n in ms)),
                    "manual_deductions": [n.get("deduction") for n in ms if n.get("deduction")],
                    "urban_void_labels": u.get("void_labels")})
    return out


def rebar(bk, N):
    cov = _man(N, file="RC", row_kind="COVER_TOTAL")
    rows = []
    for n in cov:
        if n.get("rebar_t") is None or n["trade"] in ("REBAR",):
            continue
        rows.append({"group": n["label_ar"], "group_en": n["label_en"], "concrete_m3": r6(n["qty"]),
                     "rebar_t": n["rebar_t"], "kg_per_m3": r6(n["rebar_t"] * 1000 / n["qty"]),
                     "rebar_cell_formula": n.get("rebar_formula")})
    tot = next((n for n in cov if n["trade"] == "REBAR"), None)
    bk.add("REBAR", "REBAR", "rebar tonnage (all groups)", unit="t", urban=None, bench=tot["qty"] if tot else None,
           comparability="URBAN_BLOCKED", confidence="HIGH", primary="MANUAL_ASSUMPTION_NOT_IN_SOURCE",
           secondary=["SOURCE_GAP"], urban_ref="REBAR_EVIDENCE_REGISTER (weight BLOCKED; kg/m3 refused)",
           bench_refs=[tot["id"]] if tot else [], note="typed constants per group, no bar schedule: an allowance")
    typed = all(r["rebar_cell_formula"] is None for r in rows)
    return {"groups": rows, "total_t": tot["qty"] if tot else None, "all_typed_constants": typed,
            "kg_per_m3_range": [min(r["kg_per_m3"] for r in rows), max(r["kg_per_m3"] for r in rows)] if rows else None,
            "method_class": "MANUAL_ESTIMATE_METHOD", "deterministic": False,
            "rule": "never gold for Urban's deterministic rebar engine"}


def other_structure(bk, N):
    for label, rows in (("stairs + domes", _man(N, trade="STAIR_CONCRETE", row_kind="DETAIL") +
                         _man(N, trade="OTHER_CONCRETE", subtrade="DOME", row_kind="DETAIL")),
                        ("swimming pool", _man(N, trade="OTHER_CONCRETE", subtrade="POOL", row_kind="DETAIL"))):
        bk.add("OTHER_STRUCTURE", "STRUCTURAL CONCRETE", label, unit="m3", urban=None, bench=sum(n["qty"] for n in rows),
               comparability="URBAN_BLOCKED" if "stairs" in label else "BENCHMARK_ONLY", confidence="MEDIUM",
               primary="SOURCE_GAP" if "stairs" in label else "SCOPE_MISMATCH", bench_refs=_ids(rows),
               urban_ref="CONCRETE_REGISTER S3-STR-VOL" if "stairs" in label else None)


# ------------------------------------------------------------------ architecture
def _not_detected_cause(R3, fl, lab, klass):
    """why Urban has no label-identified room for a manual row (evidence from the A3 room register, never a value)"""
    ufl = {"ROOF": "2F"}.get(fl, fl)
    blocked = [r for r in R3["ROOM_REGISTER"]["rooms"] if r["floor"] == ufl and r["status"] != COMPLETE]
    issues = sorted({i for r in blocked for i in (r["blocker"] or "").split("; ") if i})
    if fl == "GF" and lab in ("استقبال+طعام", "مدخل"):
        return {"primary": "SOURCE_GAP", "secondary": ["TOPOLOGY_DEFECT"],
                "why": "the GF zone is open to the outside at the west entrance (no door / screen drawn) - OQ3-A1"}
    if not [r for r in R3["ROOM_REGISTER"]["rooms"] if r["floor"] == ufl]:
        return {"primary": "ROOM_SEMANTIC_DEFECT", "secondary": ["UNRESOLVED_REQUIRES_REVIEW"],
                "why": f"no labelled site found on {ufl} in Phase A3"}
    if klass == "DRY" and any(i in issues for i in ("MULTIPLE_SEMANTIC_LABELS", "NEAR_MISS_BOUNDARY_GAP",
                                                      "OPENING_CLOSURE_UNRESOLVED")):
        return {"primary": "TOPOLOGY_DEFECT", "secondary": ["ROOM_SEMANTIC_DEFECT"],
                "why": f"Urban sites on {ufl} merge several labelled rooms ({', '.join(issues)})"}
    return None


ROOM_LABEL_MAP = {  # identity by LABEL on the same floor (never by area)
    ("GF", "مطبخ"): "KITCHEN", ("GF", "غرفه نوم رئيسيه"): "MASTER BED ROOM", ("GF", "ديوانيه"): "DEWANEYA / Wash",
}


def rooms(bk, R3, N):
    urooms = R3["ROOM_REGISTER"]["rooms"]
    cand = {(c["floor"], c["room"]): c for c in R3["ROOM_REGISTER"]["candidate_areas_review"]}
    man = [n for n in _man(N, file="FIN", sheet="سيراميك", trade="FLOOR", row_kind="DETAIL")]
    labels_per_floor = Counter((n["floor"], (n["label_ar"] or "").strip()) for n in man)
    out = []
    for n in man:
        fl, lab = n["floor"], (n["label_ar"] or "").strip()
        ul = ROOM_LABEL_MAP.get((fl, lab))
        unique_manual = labels_per_floor[(fl, lab)] == 1
        hits = [r for r in urooms if r["floor"] == fl and r["room"] == ul] if ul else []
        bucket, urban, uref, conf, comp, prim, sec = "URBAN_NOT_DETECTED", None, None, "LOW", "NOT_COMPARABLE", None, []
        cause = _not_detected_cause(R3, fl, lab, n["room_class"])
        if cause:
            prim, sec = cause["primary"], cause["secondary"]
        if ul and unique_manual and len(hits) == 1:
            u = hits[0]
            if u["status"] == COMPLETE:
                bucket, urban, uref = "URBAN_CERTIFIED", u["floor_area_m2"], f"ROOM_REGISTER.rooms[{fl} {ul}].floor_area_m2"
                conf, comp = "HIGH", "EXACT_COMPARABLE"
            else:
                c = cand.get((fl, ul))
                bucket, urban = "URBAN_REVIEW_REQUIRED", c["candidate_area_m2"] if c else None
                uref = f"ROOM_REGISTER.candidate_areas_review[{fl} {ul}] (BLOCKED: {u['blocker']})"
                conf, comp, prim = "MEDIUM", "URBAN_BLOCKED", "TOPOLOGY_DEFECT" if "TOPOLOGY" in (u["blocker"] or "") \
                    else "ROOM_SEMANTIC_DEFECT"
                sec = ["UNRESOLVED_REQUIRES_REVIEW"]
        r = bk.add("ROOM_FLOOR", "FLOOR", f"{n['label_en']} floor area", unit="m2", urban=urban, bench=n["qty"],
                   comparability=comp, confidence=conf, primary=prim, secondary=sec, floor=fl, urban_ref=uref,
                   bench_refs=[n["id"]], note=None if comp == "EXACT_COMPARABLE" else
                   ("diagnostic only - Urban candidate is not certified" if bucket == "URBAN_REVIEW_REQUIRED" else
                    "no label-identified Urban room (identity is never taken from the area)"),
                   extra={"bucket": bucket, "room_class": n["room_class"],
                          "cause": None if bucket != "URBAN_NOT_DETECTED" else (cause or {}).get("why"),
                          "diagnostic_difference_m2": r6(urban - n["qty"]) if bucket == "URBAN_REVIEW_REQUIRED" and urban else None})
        if comp == "EXACT_COMPARABLE":
            u = hits[0]
            bk.add("ROOM_PERIMETER", "SKIRTING", f"{n['label_en']} perimeter", unit="m", urban=u["perimeter_m"],
                   bench=n["perimeter_m"], comparability="EXACT_COMPARABLE", confidence="HIGH", kind="linear",
                   floor=fl, urban_ref=f"ROOM_REGISTER.rooms[{fl} {ul}].perimeter_m", bench_refs=[n["id"]])
        out.append({"floor": fl, "manual_room": lab, "manual_room_en": n["label_en"], "class": n["room_class"],
                    "manual_area_m2": n["qty"], "manual_perimeter_m": n["perimeter_m"], "bucket": bucket,
                    "urban_room": ul if hits else None, "urban_area_m2": r6(urban), "row": r["id"]})
    # Urban certified rooms without a label-identified manual row: value coincidences are reported, never matched
    claimed = {o["urban_room"] for o in out if o["bucket"] == "URBAN_CERTIFIED"}
    coincid = []
    for u in urooms:
        if u["status"] != COMPLETE or u["room"] in claimed:
            continue
        cands = [n for n in man if n["floor"] == u["floor"] and n["room_class"] == "WET"] if u["room"] in ("W.C", "BATH") \
            else [n for n in man if n["floor"] == u["floor"]]
        near = min(cands, key=lambda n: abs(n["qty"] - u["floor_area_m2"])) if cands else None
        coincid.append({"floor": u["floor"], "urban_room": u["room"], "urban_area_m2": u["floor_area_m2"],
                        "urban_perimeter_m": u["perimeter_m"], "nearest_manual_row": near["id"] if near else None,
                        "nearest_manual_label": near["label_ar"] if near else None,
                        "nearest_manual_area_m2": near["qty"] if near else None,
                        "nearest_manual_perimeter_m": near["perimeter_m"] if near else None,
                        "area_pct": r6(pct_diff(u["floor_area_m2"], near["qty"]) * 100) if near else None,
                        "use": "VALUE_COINCIDENCE - LOW confidence; not a match, not in any metric"})
        bk.add("ROOM_FLOOR", "FLOOR", f"{u['floor']} {u['room']} {u['floor_area_m2']} m2 (no label identity)", unit="m2",
               urban=u["floor_area_m2"], bench=near["qty"] if near else None, comparability="NOT_COMPARABLE",
               confidence="LOW", primary="UNRESOLVED_REQUIRES_REVIEW", floor=u["floor"],
               urban_ref=f"ROOM_REGISTER.rooms[{u['floor']} {u['room']}]", bench_refs=[near["id"]] if near else [],
               note="manual rows carry no room identity that matches this Urban room; nearest manual value shown")
    floors = {}
    for fl in ("GF", "1F", "ROOF"):
        mt = sum(n["qty"] for n in man if n["floor"] == fl)
        ufl = {"GF": "GF", "1F": "1F", "ROOF": "2F"}[fl]
        row = next((r for r in R3["ARCH_BOQ"]["rows"] if r["item"] == f"A3-FLR-{ufl}"), None)
        uq = row["qty"] if row else None
        bk.add("FLOOR_TOTAL", "FLOOR", f"floor area ({fl})", unit="m2", urban=uq, bench=mt,
               comparability="PARTIAL_SCOPE_COMPARABLE" if uq else "URBAN_BLOCKED", confidence="HIGH",
               primary="SCOPE_MISMATCH" if uq else "ROOM_SEMANTIC_DEFECT", secondary=["TOPOLOGY_DEFECT"] if uq else [],
               floor=fl,
               urban_ref=f"ARCH_BOQ.{row['item']} ({row['status']})" if row else None,
               note="Urban = certified rooms only; manual = all rooms of the floor",
               extra={"coverage_ratio": r6(uq / mt) if uq and mt else None})
        floors[fl] = {"manual_m2": r6(mt), "urban_certified_m2": uq, "coverage": r6(uq / mt) if uq and mt else None}
    return {"rooms": out, "value_coincidences": coincid, "floors": floors,
            "buckets": dict(Counter(o["bucket"] for o in out))}


def ceilings(bk, R3, N):
    man = _man(N, file="FIN", sheet="ديكور", trade="CEILING", row_kind="DETAIL")
    out = []
    for n in man:
        ul = ROOM_LABEL_MAP.get((n["floor"], (n["label_ar"] or "").strip()))
        u = next((r for r in R3["ROOM_REGISTER"]["rooms"] if r["floor"] == n["floor"] and r["room"] == ul
                  and r["status"] == COMPLETE), None) if ul else None
        if u:
            r = bk.add("CEILING", "CEILING", f"{n['label_en']} ceiling (base area)", unit="m2",
                       urban=u["base_ceiling_area_m2"], bench=n["qty"], comparability="EXACT_COMPARABLE",
                       confidence="HIGH", floor=n["floor"], urban_ref=f"ROOM_REGISTER.rooms[{n['floor']} {ul}].base_ceiling_area_m2",
                       bench_refs=[n["id"]], method="manual gypsum decor area = floor area")
            out.append(r["id"])
    tot = sum(n["qty"] for n in man)
    bk.add("CEILING", "CEILING", "gypsum decor, all rooms", unit="m2", urban=None, bench=tot,
           comparability="PARTIAL_SCOPE_COMPARABLE", confidence="MEDIUM", primary="SCOPE_MISMATCH",
           note="Urban base-ceiling exists only for certified rooms; decor material not in the source")
    cor = _man(N, file="FIN", sheet="الغلاف", trade="CEILING")
    for c in cor:
        if "CORNICE" in (c["subtrade"] or ""):
            bk.add("CEILING", "CEILING", c["label_en"], unit="m", urban=None, bench=c["qty"], comparability="BENCHMARK_ONLY",
                   confidence="HIGH", primary="SCOPE_MISMATCH", bench_refs=[c["id"]],
                   note="cornice length = room perimeter in the manual")
    return out


def wall_finishes(bk, R3, N):
    tile = _man(N, file="FIN", sheet="سيراميك", trade="FLOOR", row_kind="DETAIL", room_class="WET")
    hs = sorted({n["height_m"] for n in tile if n["height_m"]})
    bk.add("WALL_TILE", "TILE", "wet-room wall tile", unit="m2", urban=None,
           bench=sum(n["wall_m2"] or 0 for n in tile), comparability="URBAN_BLOCKED", confidence="HIGH",
           primary="HEIGHT_EVIDENCE_GAP", secondary=["TRADE_RULE_GAP"],
           urban_ref="ARCH_BOQ A3-WFN-* (BLOCKED)", bench_refs=_ids(tile),
           method=f"manual: room perimeter x {hs} m, no opening deduction")
    k = next((n for n in tile if n["floor"] == "GF" and (n["label_ar"] or "").strip() == "مطبخ"), None)
    u = next((r for r in R3["ROOM_REGISTER"]["rooms"] if r["floor"] == "GF" and r["room"] == "KITCHEN"), None)
    if k and u:
        bk.add("WALL_TILE", "TILE", "kitchen wall run (manual perimeter vs Urban wall-face length)", unit="m",
               urban=u["wall_face_length_m"], bench=k["perimeter_m"], comparability="SCOPE_MISMATCH",
               confidence="MEDIUM", primary="MANUAL_METHOD_DIFFERENCE", floor="GF",
               urban_ref="ROOM_REGISTER.rooms[GF KITCHEN].wall_face_length_m", bench_refs=[k["id"]],
               note="Urban wall-face length excludes the opening widths; the manual tiles the full perimeter")
    paint = [n for n in _man(N, file="PNT", trade="PAINT", row_kind="DETAIL") if n["subtrade"] == "INTERNAL_WALL"]
    plast = [n for n in _man(N, file="PLS", trade="PLASTER", row_kind="DETAIL") if n["subtrade"] == "INTERNAL_WALL"]
    pc = {n["subtrade"]: n for n in _man(N, file="PNT", row_kind="COVER_TOTAL")}
    lc = {n["subtrade"]: n for n in _man(N, file="PLS", row_kind="COVER_TOTAL")}
    for trade, rows, cov, key in (("PAINT", paint, pc, "INTERNAL_WALL"), ("PLASTER", plast, lc, "INTERNAL_WALL")):
        c = cov.get(key)
        bk.add("WALL_FINISH", trade, f"internal {trade.lower()} (walls)", unit="m2", urban=None,
               bench=c["qty"] if c else None, comparability="URBAN_BLOCKED", confidence="HIGH",
               primary="HEIGHT_EVIDENCE_GAP", secondary=["TRADE_RULE_GAP"],
               urban_ref="ARCH_BOQ A3-WFA-* (BLOCKED_HEIGHT)", bench_refs=[c["id"]] if c else [],
               method="manual: perimeter x height per room, openings deducted at half")
    ext = lc.get("EXTERNAL")
    if ext:
        bk.add("WALL_FINISH", "PLASTER", "external plaster", unit="m2", urban=None, bench=ext["qty"],
               comparability="URBAN_BLOCKED", confidence="HIGH", primary="HEIGHT_EVIDENCE_GAP",
               urban_ref="ARCH_BOQ A3-EXT-FAC (ELEVATIONS_SW_NE_RASTER_ONLY)", bench_refs=[ext["id"]])
    heights = sorted({(n["height_m"]) for n in paint if n["height_m"]})
    method = {"paint_heights_m": heights,
              "rooms": [{"label": n["label_ar"], "label_en": n["label_en"], "height_m": n["height_m"],
                         "length_m": n["length_m"], "m2": n["qty"]} for n in paint],
              "opening_deduction": "openings listed once (paint 116.495 m2) and deducted at half (58.2475)",
              "plaster_vs_paint": "plaster gross excludes the stair-soffit rows that paint includes",
              "inferred_height_rule": "GF 3.6 m (= 4.50 - 0.90), 1F / roof 3.4 m (= 4.20 - 0.80), living 3.5 m, "
                                      "stairwell 12.9 m - a clear height under a false ceiling, not in the source"}
    return {"paint_plaster_method": method}


def waterproofing(bk, N):
    for c in _man(N, file="FIN", sheet="الغلاف", trade="WATERPROOFING"):
        roof = c["subtrade"].startswith("ROOF")
        bk.add("WATERPROOFING", "WATERPROOFING", c["label_en"], unit=c["unit"] or "m2", urban=None, bench=c["qty"],
               comparability="BENCHMARK_ONLY" if roof else "URBAN_BLOCKED", confidence="HIGH",
               primary="SCOPE_MISMATCH" if roof else "SOURCE_GAP", secondary=[] if roof else ["TRADE_RULE_GAP"],
               urban_ref=None if roof else "ARCH_BOQ A3-WPM-* (BLOCKED_MATERIAL)", bench_refs=[c["id"]],
               method=("roof: one total area and perimeter" if roof else
                       "wet rooms: floor areas + perimeter x 1.0 m upturn (recorded as m.t although computed as m2)"))


def blockwork(bk, R3, N):
    man = [n for n in _man(N, file="BLK", trade="BLOCKWORK", row_kind="DETAIL")]
    ub = {r["item"]: r for r in R3["BLOCKWORK_REGISTER"]["rows"]}
    out = []
    for fl, ufl in (("GF", "GF"), ("1F", "1F"), ("ROOF", "2F")):
        rows = [n for n in man if n["floor"] == fl and ":" not in n["subtrade"]]
        by = defaultdict(float)
        for n in rows:
            by[n["subtrade"]] += (n["count"] or 1) * (n["dims"]["length_m"] or 0)
        u150 = (ub.get(f"A3-BLK-{ufl}-T150") or {}).get("qty")
        u200 = (ub.get(f"A3-BLK-{ufl}-T200") or {}).get("qty")
        mlen = sum(by.values())
        ulen = (u150 or 0) + (u200 or 0)
        bk.add("BLOCKWORK", "BLOCKWORK", f"wall length ({fl}, all classes, building walls)", unit="m", urban=ulen,
               bench=mlen, comparability="PARTIAL_SCOPE_COMPARABLE", confidence="MEDIUM", primary="SCOPE_MISMATCH",
               secondary=["GEOMETRY_DEFECT"], floor=fl, urban_ref=f"BLOCKWORK_REGISTER A3-BLK-{ufl}-T150/T200 "
               "(established bands only)", bench_refs=_ids(rows),
               note="Urban = established bands only (ambiguous bands excluded); manual = gross wall runs; thickness "
                    "classes differ (manual: external / 15 / 20 cm, Urban: measured 150 / 200 mm)",
               extra={"coverage_ratio": r6(ulen / mlen) if mlen else None})
        heights = sorted({n["dims"]["height_or_width_m"] for n in rows})
        out.append({"floor": fl, "manual_length_m": {k: r6(v) for k, v in sorted(by.items())},
                    "manual_total_m": r6(mlen), "urban_T150_m": u150, "urban_T200_m": u200, "urban_total_m": r6(ulen),
                    "coverage": r6(ulen / mlen) if mlen else None, "manual_heights_m": heights})
    cov = {n["subtrade"]: n for n in _man(N, file="BLK", row_kind="COVER_TOTAL")}
    for k in ("EXTERNAL", "INTERNAL_150", "INTERNAL_200"):
        c = cov.get(k)
        if c:
            bk.add("BLOCKWORK", "BLOCKWORK", f"blockwork area {k.lower()}", unit="m2", urban=None, bench=c["qty"],
                   comparability="URBAN_BLOCKED", confidence="HIGH", primary="HEIGHT_EVIDENCE_GAP",
                   urban_ref="ARCH_BOQ A3-BLA-* (BLOCKED_HEIGHT)", bench_refs=[c["id"]],
                   note="net of openings; includes fence / parapet / projection rows" if k != "INTERNAL_200" else "net of openings")
    extra = [{"label": n["label_ar"], "label_en": n["label_en"], "subtrade": n["subtrade"], "m2": n["qty"]}
             for n in man if any(s in n["subtrade"] for s in ("FENCE", "PARAPET", "PROJECTION"))]
    return {"floors": out, "non_building_rows": extra}


def aluminium(bk, R3, N, tol):
    g = R3["ALUMINIUM_GLAZING_REGISTER"]
    win = R3["OPENING_REGISTER"]["windows"]
    alu = _man(N, file="ALU", trade="ALUMINIUM", row_kind="DETAIL")
    pool = {fl: [n for n in alu if n["floor"] == fl and (n["label_ar"] or "").strip() == "شباك المسبح"]
            for fl in ("GF", "1F", "ROOF")}
    sal = g["salon"]
    smatch = [n for n in pool["GF"] if abs(n["dims"]["length_m"] - sal["width_m"]) <= 1e-9]
    out = {"salon": None, "master_bedroom_curved": None, "first_floor_curved": None, "windows": [], "doors": None}
    special = set()
    if len(smatch) == 1:
        m = smatch[0]
        mh = m["dims"]["height_or_width_m"]
        w = bk.add("ALUMINIUM", "ALUMINIUM", "salon sea-view opening width", unit="m", urban=sal["width_m"],
                   bench=m["dims"]["length_m"], comparability="EXACT_COMPARABLE", confidence="HIGH", kind="linear",
                   floor="GF", urban_ref="ALUMINIUM_GLAZING_REGISTER.salon.width_m (SOURCE)", bench_refs=[m["id"]],
                   method="identity: the GF pool-facing window whose width equals the source width (the Urban "
                          "binding is the printed dimension + CAD gap, not this row)")
        h = bk.add("ALUMINIUM", "ALUMINIUM", "salon opening height", unit="m", urban=sal["height_m"], bench=mh,
                   comparability="EXACT_COMPARABLE", confidence="HIGH", kind="linear", floor="GF",
                   primary="HEIGHT_EVIDENCE_GAP", secondary=["MANUAL_ASSUMPTION_NOT_IN_SOURCE"],
                   urban_ref="ALUMINIUM_GLAZING_REGISTER.salon.height_m (PROJECT_OWNER_DERIVED_DIMENSION)",
                   bench_refs=[m["id"]], note="no source height exists; owner-derived 3.65 m vs manual 4.30 m")
        a = bk.add("ALUMINIUM", "ALUMINIUM", "salon aluminium + glass area", unit="m2", urban=sal["area_m2"],
                   bench=m["qty"], comparability="EXACT_COMPARABLE", confidence="HIGH", floor="GF",
                   primary="HEIGHT_EVIDENCE_GAP", secondary=["MANUAL_ASSUMPTION_NOT_IN_SOURCE"],
                   urban_ref="ALUMINIUM_GLAZING_REGISTER.salon.area_m2", bench_refs=[m["id"]],
                   note=f"the whole difference is the height: {sal['width_m']} x ({sal['height_m']} - {mh}) = "
                        f"{r6(sal['width_m'] * (sal['height_m'] - mh))} m2; Urban keeps the owner value")
        special.add(m["id"])
        out["salon"] = {"rows": [w["id"], h["id"], a["id"]], "manual_width_m": m["dims"]["length_m"],
                        "manual_height_m": mh, "manual_area_m2": m["qty"], "urban_area_m2": sal["area_m2"],
                        "attributable_to_height_m2": r6(sal["width_m"] * (sal["height_m"] - mh)),
                        "width_attributable_m2": 0.0}
    # master-bedroom curved: by elimination (the other GF pool-facing window; MB glazing faces the pool - owner fact)
    others = [n for n in pool["GF"] if n not in smatch]
    item = next(c for c in g["curved"]["items"] if c["id"] == g["curved"]["master_bedroom"]["item"])
    if len(others) == 1:
        m = others[0]
        arcs = item["arc_lengths_mm"]
        bases = {"inner_arc": min(arcs) / 1000, "mean_arc (Urban)": sum(arcs) / len(arcs) / 1000,
                 "outer_arc": max(arcs) / 1000, "chord": item["chord_mm"] / 1000}
        ml = m["dims"]["length_m"]
        diffs = {k: r6(pct_diff(v, ml) * 100) for k, v in bases.items()}
        close = [k for k, v in bases.items() if band("linear", v, ml, tol) in ("EXACT", "CLOSE")]
        r = bk.add("ALUMINIUM", "ALUMINIUM", "master-bedroom curved glazing length", unit="m",
                   urban=g["curved"]["master_bedroom"]["developed_length_m"], bench=ml,
                   comparability="EXACT_COMPARABLE", confidence="MEDIUM", kind="linear", floor="GF",
                   primary=None, urban_ref="ALUMINIUM_GLAZING_REGISTER.curved.master_bedroom (mean of 4 arcs)",
                   bench_refs=[m["id"]], method="identity by elimination: the GF pool-facing window that is not the "
                                                "salon (owner fact: the MB curved glazing faces the pool)")
        if r["band"] not in ("EXACT", "CLOSE"):
            r["primary_class"], r["secondary_classes"] = "MANUAL_METHOD_DIFFERENCE", ["UNRESOLVED_REQUIRES_REVIEW"]
        bk.add("ALUMINIUM", "ALUMINIUM", "master-bedroom curved glazing area", unit="m2", urban=None, bench=m["qty"],
               comparability="URBAN_BLOCKED", confidence="MEDIUM", primary="HEIGHT_EVIDENCE_GAP",
               secondary=["MANUAL_ASSUMPTION_NOT_IN_SOURCE"], floor="GF", bench_refs=[m["id"]],
               urban_ref="ARCH_BOQ A3-CGL-MB-A (BLOCKED_HEIGHT)",
               note=f"manual height {m['dims']['height_or_width_m']} m (not in the source)")
        special.add(m["id"])
        out["master_bedroom_curved"] = {"row": r["id"], "manual_length_m": ml,
                                        "manual_height_m": m["dims"]["height_or_width_m"], "urban_bases_m": {k: r6(v) for k, v in bases.items()},
                                        "pct_vs_manual": diffs, "bases_within_tolerance": close,
                                        "closest_basis": min(diffs, key=lambda k: abs(diffs[k])),
                                        "classification": "CLOSEST_TO_INNER_ARC_OUTSIDE_TOLERANCE" if not close else
                                        f"WITHIN_TOLERANCE: {close}",
                                        "identity_confidence": "MEDIUM"}
    c1 = next((c for c in g["curved"]["items"] if c["floor"] == "1F"), None)
    if c1 and pool["1F"]:
        straight = sorted(w["width_mm"] / 1000 for w in win if w["floor"] == "1F")
        rest = [n for n in pool["1F"] if not any(abs(n["dims"]["length_m"] - s) <= 0.01 for s in straight)]
        if len(rest) == 1:
            m = rest[0]
            arcs = c1["arc_lengths_mm"]
            bk.add("ALUMINIUM", "ALUMINIUM", "1F curved glazing length", unit="m", urban=c1["developed_length_mm"] / 1000,
                   bench=m["dims"]["length_m"], comparability="EXACT_COMPARABLE", confidence="LOW", kind="linear",
                   floor="1F", urban_ref="ALUMINIUM_GLAZING_REGISTER.curved 1F-CG01", bench_refs=[m["id"]],
                   method="identity by elimination against the Urban straight-window widths (value-based - LOW)")
            special.add(m["id"])
            out["first_floor_curved"] = {"manual_length_m": m["dims"]["length_m"], "urban_mean_m": c1["developed_length_mm"] / 1000,
                                         "urban_arcs_m": [a / 1000 for a in arcs], "identity_confidence": "LOW"}
    # window widths per floor (multiset matching within the declared tolerance)
    wtol = tolerances()["width_match_for_openings_mm"] / 1000
    for fl, ufl in (("GF", "GF"), ("1F", "1F"), ("ROOF", "2F")):
        mw = []
        for n in alu:
            if n["floor"] == fl and n["subtrade"] == "WINDOW" and n["id"] not in special:
                mw += [n["dims"]["length_m"]] * int(n["count"] or 1)
        uw = sorted(w["width_mm"] / 1000 for w in win if w["floor"] == ufl and w["width_mm"] < 5000)
        left_u, left_m, pairs = list(uw), sorted(mw), 0
        for v in sorted(mw):
            hit = next((x for x in left_u if abs(x - v) <= wtol), None)
            if hit is not None:
                left_u.remove(hit)
                left_m.remove(v)
                pairs += 1
        doors = [n["dims"]["length_m"] for n in alu if n["floor"] == fl and n["subtrade"] == "DOOR"
                 for _ in range(int(n["count"] or 1))]
        as_doors = []
        for v in list(left_u):
            if any(abs(v - d) <= wtol for d in doors) and len(as_doors) < len(doors):
                as_doors.append(v)
        r = bk.add("WINDOWS", "WINDOWS", f"straight window count ({fl})", unit="nr", urban=len(uw), bench=len(mw),
                   comparability="EXACT_COMPARABLE", confidence="MEDIUM", kind="count", floor=fl,
                   primary=(None if len(uw) == len(mw) and not left_u and not left_m else
                            "OPENING_DEFECT" if as_doors else "UNRESOLVED_REQUIRES_REVIEW"),
                   secondary=[] if len(uw) == len(mw) and not left_u and not left_m else ["BENCHMARK_POSSIBLE_ERROR"],
                   urban_ref=f"OPENING_REGISTER.windows[{ufl}] (salon and curved excluded)",
                   note=f"{pairs} widths pair within {wtol * 1000:g} mm; Urban-only widths {left_u}; manual-only "
                        f"widths {left_m}" + (f"; {len(as_doors)} Urban-only width(s) equal manual aluminium DOOR widths "
                                              f"{doors}" if as_doors else ""))
        out["windows"].append({"floor": fl, "urban_widths_m": uw, "manual_widths_m": sorted(mw), "paired": pairs,
                               "urban_only": left_u, "manual_only": left_m, "urban_only_equal_to_manual_door_widths": as_doors,
                               "row": r["id"]})
    dm = sum(n["count"] or 1 for n in alu if n["subtrade"] == "DOOR")
    ud = sum((r["qty"] or 0) for r in R3["ARCH_BOQ"]["rows"] if r["item"].startswith("A3-DOR-"))
    bk.add("DOORS", "DOORS", "door count (Urban door symbols vs manual aluminium doors)", unit="nr", urban=ud, bench=dm,
           comparability="SCOPE_MISMATCH", confidence="LOW", primary="SCOPE_MISMATCH",
           urban_ref="ARCH_BOQ A3-DOR-* (all door symbols, material unknown)",
           note="the manual lists aluminium doors only; no door schedule exists; wood doors appear only in the web-app")
    out["doors"] = {"urban_door_symbols": ud, "manual_aluminium_doors": dm}
    tot = {n["subtrade"]: n["qty"] for n in _man(N, file="ALU", row_kind="COVER_TOTAL")}
    bk.add("ALUMINIUM", "ALUMINIUM", "aluminium windows + doors area (all)", unit="m2", urban=None, bench=tot.get("ALL"),
           comparability="URBAN_BLOCKED", confidence="HIGH", primary="HEIGHT_EVIDENCE_GAP",
           urban_ref="ARCH_BOQ A3-WNA-* (WINDOW_HEIGHT_NOT_PRINTED)", note="manual window heights 2.0 m / bath 0.8 m / "
                                                                           "pool 4.3 (GF) and 3.0 (1F): not in the source")
    return out


def reception(bk, R3, N):
    pnt = _man(N, file="PNT", trade="PAINT", row_kind="DETAIL")
    tall = [n for n in pnt if (n["height_m"] or 0) > 4.5]
    rd = next((n for n in _man(N, file="FIN", sheet="سيراميك", row_kind="DETAIL") if (n["label_ar"] or "").strip() == "استقبال+طعام"), None)
    bk.add("RECEPTION", "PAINT", "reception double-height walls", unit="m2", urban=None, bench=None,
           comparability="NOT_COMPARABLE", confidence="LOW", primary="HEIGHT_EVIDENCE_GAP",
           urban_ref="DOUBLE_HEIGHT_ZONE_REGISTER (zone corroborated; wall heights BLOCKED)",
           note="no manual row carries a double-height reception wall; the only wall above 4.5 m is the stairwell "
                f"({', '.join(n['label_en'] + ' ' + str(n['height_m']) + ' m' for n in tall)})")
    return {"manual_zone_row": rd["id"] if rd else None, "manual_zone_area_m2": rd["qty"] if rd else None,
            "manual_zone_perimeter_m": rd["perimeter_m"] if rd else None,
            "manual_treatment": "ONE closed room 'reception + dining' (saloon included, no separate row) with its own "
                                "perimeter; the entrance hall is a separate room (10.46 m2) - the manual assumes the "
                                "zone is closed at the entrance",
            "manual_tall_walls": [{"label": n["label_ar"], "height_m": n["height_m"], "m2": n["qty"]} for n in tall],
            "urban": "zone open at the west entrance (OQ3-A1); double height corroborated, extent blocked"}


def master_bedroom(R3, N):
    m = next((n for n in _man(N, file="FIN", sheet="سيراميك", row_kind="DETAIL", floor="GF")
              if (n["label_ar"] or "").strip() == "غرفه نوم رئيسيه"), None)
    c = next((x for x in R3["ROOM_REGISTER"]["candidate_areas_review"] if x["floor"] == "GF" and x["room"] == "MASTER BED ROOM"), None)
    box = 0.50 * 1.46
    out = {"manual_area_m2": m["qty"] if m else None, "manual_perimeter_m": m["perimeter_m"] if m else None,
           "urban_candidate_m2": c["candidate_area_m2"] if c else None, "box_m2": box}
    if m and c:
        d = m["qty"] - c["candidate_area_m2"]
        out["manual_minus_urban_m2"] = r6(d)
        out["likely_treatment"] = "NO_DEDUCTION" if d >= -0.05 else ("DEDUCTED" if abs(d + box) < 0.1 else "UNKNOWN")
        out["reasoning"] = ("the manual area is not smaller than the Urban candidate site, so the 0.73 m2 box is not "
                            "deducted in the manual (consistent with WARDROBE / JOINERY with floor finish under it)"
                            if d >= -0.05 else "the manual area is smaller - possible deduction")
        out["classification"] = "WARDROBE/JOINERY (probable, MEDIUM) - Urban topology unchanged, OQ3-A2 carried"
    return out


def marble(bk, N):
    stairs = [n for n in _man(N, file="FIN", sheet="رخام+حوش", trade="MARBLE", row_kind="DETAIL") if n["subtrade"] in ("STAIR", "LANDING")]
    cov = next((n for n in _man(N, file="FIN", sheet="الغلاف", trade="MARBLE") if "الدرج" in n["label_ar"]), None)
    bk.add("MARBLE", "MARBLE / STAIRS", "marble stairs + landings", unit="m2", urban=None, bench=sum(n["qty"] for n in stairs),
           comparability="URBAN_BLOCKED", confidence="HIGH", primary="SOURCE_GAP", secondary=["BENCHMARK_POSSIBLE_ERROR"],
           urban_ref="ARCH_BOQ A3-STR-* (RISER_COUNT_PER_FLIGHT_NOT_PROVED)", bench_refs=_ids(stairs),
           note=f"stairs + landings rows = {r6(sum(n['qty'] for n in stairs))} m2; the manual total "
                f"{cov['qty'] if cov else None} also contains the yard floor (see MANUAL_BOQ_QA)")
    for lab, tr, sub in (("yard floor", "FLOOR", "EXTERNAL_YARD"), ("pool floor", "FLOOR", "POOL"),
                         ("pool walls", "TILE", "POOL"), ("internal railing", "OTHER", "INTERNAL")):
        c = next((n for n in _man(N, file="FIN", sheet="الغلاف") if n["subtrade"] == sub and
                  (n["trade"] == tr or (tr == "TILE" and n["trade"] == "WALL_TILE") or (tr == "OTHER" and n["trade"] == "RAILING"))), None)
        if c:
            bk.add("OTHER_ARCH", tr, lab, unit=c["unit"] or "", urban=None, bench=c["qty"], comparability="BENCHMARK_ONLY",
                   confidence="HIGH", primary="SCOPE_MISMATCH", bench_refs=[c["id"]], note="no Urban item in Phase A3")
    jny = [r for r in []]
    return {"stairs_rows_m2": r6(sum(n["qty"] for n in stairs)), "cover_total_m2": cov["qty"] if cov else None}


def urban_only(bk, R3):
    for r in R3["ARCH_BOQ"]["rows"]:
        if r["item"].startswith("A3-JNY") and r["qty"]:
            bk.add("URBAN_ONLY", "OTHER", f"counter-run joinery front ({r['floor']})", unit="m", urban=r["qty"],
                   bench=None, comparability="URBAN_ONLY", confidence="HIGH", primary="SCOPE_MISMATCH",
                   floor=r["floor"], urban_ref=f"ARCH_BOQ.{r['item']}", note="no joinery item in the manual BOQ")


# ------------------------------------------------------------------ manual BOQ QA
def _refs(f):
    return re.findall(r"(?:'?([^'!=]+)'?!)?\$?([A-Z]+)\$?(\d+)", f or "")


def manual_qa(raw, N) -> dict:
    """generic arithmetic / formula / duplicate checks over the raw rows (formulas in xlsx, values in xls)"""
    cells = {}
    for r in raw:
        for k, c in r["cells"].items():
            cells[(r["file"], r["sheet"], k)] = c
    findings = []

    def val(fk, sh, ref):
        c = cells.get((fk, sh, ref))
        return BB.num(c["v"]) if c else 0.0

    for (fk, sh, ref), c in sorted(cells.items()):
        f = c["f"]
        if isinstance(c["v"], str) and "#REF!" in c["v"] or (f and "#REF!" in f):
            findings.append({"check": "BROKEN_REFERENCE", "file": fk, "sheet": sh, "cell": ref, "formula": f,
                             "value": c["v"], "impact": "display cell only (not carried to any total)",
                             "class": "REPORTING_DEFECT"})
            continue
        if not f:
            continue
        row = int(re.search(r"\d+", ref).group(0))
        m = re.fullmatch(r"=\$?([A-Z]+)\$?(\d+)(\*\$?([A-Z]+)\$?(\d+))+", f.replace(" ", ""))
        if m:
            refs = re.findall(r"([A-Z]+)(\d+)", f)
            other = [f"{a}{b}" for a, b in refs if int(b) != row]
            if other:
                own = [f"{a}{row}" for a, b in refs if int(b) != row]
                same = all(abs(val(fk, sh, o) - val(fk, sh, w)) < 1e-12 for o, w in zip(other, own))
                findings.append({"check": "PRODUCT_REFERENCES_ANOTHER_ROW", "file": fk, "sheet": sh, "cell": ref,
                                 "formula": f, "other_row_refs": other,
                                 "impact": "none (the referenced value equals the own-row value)" if same else "VALUE",
                                 "class": "REPORTING_DEFECT" if same else "BENCHMARK_POSSIBLE_ERROR"})
            prod = 1.0
            for a, b in refs:
                prod *= val(fk, sh, f"{a}{b}")
            if BB.num(c["v"]) is not None and abs(prod - c["v"]) > 1e-6:
                findings.append({"check": "PRODUCT_CACHE_MISMATCH", "file": fk, "sheet": sh, "cell": ref, "formula": f,
                                 "cached": c["v"], "recomputed": prod, "class": "BENCHMARK_POSSIBLE_ERROR"})
        m = re.fullmatch(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", f.replace(" ", ""))
        if m:
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            s = sum(val(fk, sh, f"{col}{rr}") or 0.0 for rr in range(r1, r2 + 1)
                    for col in ([c1] if c1 == c2 else [c1, c2]))
            if BB.num(c["v"]) is not None and abs(s - c["v"]) > 1e-6:
                findings.append({"check": "SUM_CACHE_MISMATCH", "file": fk, "sheet": sh, "cell": ref, "formula": f,
                                 "cached": c["v"], "recomputed": s, "class": "BENCHMARK_POSSIBLE_ERROR"})
            if c1 != c2:
                extra = [f"{c2}{rr}" for rr in range(r1, r2 + 1) if (fk, sh, f"{c2}{rr}") in cells]
                findings.append({"check": "SUM_RANGE_SPANS_EXTRA_COLUMN", "file": fk, "sheet": sh, "cell": ref,
                                 "formula": f, "extra_column_cells_with_values": extra,
                                 "impact": "none (no values in the extra column)" if not extra else "VALUE",
                                 "class": "REPORTING_DEFECT"})
            # a SUM whose range starts after the first detail row of its block skips a row
            block_first = None
            for rr in range(r1 - 1, max(r1 - 30, 0), -1):
                v = cells.get((fk, sh, f"{c1}{rr}"))
                if v is None:
                    lab = [cells.get((fk, sh, f"A{rr}"))]
                    if lab[0] and isinstance(lab[0]["v"], str) and BB.is_total_label(lab[0]["v"].strip()):
                        break
                    if any((fk, sh, f"{x}{rr}") in cells for x in "BCDEFGHI"):
                        continue
                    break
                if BB.num(v["v"]) is not None and not (v["f"] or "").startswith("=SUM"):
                    lab = cells.get((fk, sh, f"A{rr}"))
                    if lab and isinstance(lab["v"], str) and BB.is_total_label(lab["v"].strip()):
                        break
                    block_first = rr
                else:
                    break
            if block_first is not None:
                skipped = [f"{c1}{rr}" for rr in range(block_first, r1)]
                carried = [k for (a, b, k), x in cells.items() if a == fk and b == sh and x["f"] and
                           re.fullmatch(rf"=SUM\({c1}{block_first}:{c1}\d+\)", x["f"])]
                findings.append({"check": "SUM_SKIPS_LEADING_ROWS", "file": fk, "sheet": sh, "cell": ref, "formula": f,
                                 "skipped_cells": skipped,
                                 "skipped_value": r6(sum(val(fk, sh, x) or 0 for x in skipped)),
                                 "same_block_sums_including_them": carried, "class": "BENCHMARK_POSSIBLE_ERROR"})
    # a reported total that also contains another separately reported total (double count)
    cover = [r for r in raw if r["sheet"] == "الغلاف"]
    cover_refs = defaultdict(list)
    for r in cover:
        for k, c in r["cells"].items():
            for sh, col, rr in _refs(c["f"]):
                if sh:
                    cover_refs[(r["file"], sh.strip("'"), f"{col}{rr}")].append((k, c["f"]))
    for (fk, sh, ref), users in sorted(cover_refs.items()):
        c = cells.get((fk, sh, ref))
        m = re.fullmatch(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", (c or {}).get("f") or "")
        if not m:
            continue
        inside = [(k2, u2) for (fk2, sh2, k2), u2 in cover_refs.items() if fk2 == fk and sh2 == sh and k2 != ref and
                  BB.col_of(k2) == m.group(1) and int(m.group(2)) <= int(re.search(r"\d+", k2).group(0)) <= int(m.group(4))]
        for k2, u2 in inside:
            findings.append({"check": "TOTAL_CONTAINS_ANOTHER_REPORTED_TOTAL", "file": fk, "sheet": sh,
                             "cell": ref, "formula": c["f"], "contains": k2, "contained_value": val(fk, sh, k2),
                             "both_reported_on_cover": [x[0] for x in users] + [x[0] for x in u2],
                             "class": "BENCHMARK_POSSIBLE_ERROR", "secondary": "DUPLICATE_OR_OMITTED_ITEM"})
    # xls rows: I = count x height x width x length (present factors)
    for n in N:
        if n["file"] in ("PNT", "PLS") and n["row_kind"] == "DETAIL" and n.get("length_m") is not None:
            fac = [x for x in (n.get("count"), n.get("height_m"), n.get("width_m"), n.get("length_m")) if x]
            p = 1.0
            for x in fac:
                p *= x
            if n["qty"] is not None and abs(p - n["qty"]) > 1e-6:
                findings.append({"check": "XLS_ROW_PRODUCT", "file": n["file"], "sheet": n["sheet"], "row": n["row"],
                                 "label": n["label_ar"], "value": n["qty"], "product_of_factors": r6(p),
                                 "class": "BENCHMARK_POSSIBLE_ERROR"})
    # xls: a grand subtotal (J only) equals the J subtotals above it since the previous grand subtotal
    checked = 0
    for fk in ("PNT", "PLS"):
        rows = [n for n in N if n["file"] == fk and n["sheet"] == "البيــــــــــــان"]
        run = []
        for n in rows:
            if n["row_kind"] == "DETAIL" and n.get("subtotal") is not None:
                run.append(n["subtotal"])
            elif n["row_kind"] == "DETAIL" and n.get("subtotal") is None and n["qty"] is not None and not n.get("length_m") is None \
                    and n.get("label_inherited") is False and False:
                pass
            elif n["row_kind"] == "SUBTOTAL":
                if run:
                    checked += 1
                    if abs(sum(run) - n["qty"]) > 1e-6:
                        findings.append({"check": "XLS_SUBTOTAL_SUM", "file": fk, "sheet": n["sheet"], "row": n["row"],
                                         "value": n["qty"], "sum_of_subtotals": r6(sum(run)),
                                         "class": "BENCHMARK_POSSIBLE_ERROR"})
                run = []
    findings.append({"check": "XLS_SUBTOTALS_CHECKED", "groups": checked, "class": "INFO"})
    # typed constants on covers (no link to a detail sheet)
    typed = []
    for n in N:
        if n["row_kind"] == "COVER_TOTAL" and n["file"] in ("FIN", "ALU", "RC"):
            raw_row = next(r for r in raw if r["file"] == n["file"] and r["sheet"] == n["sheet"] and r["row"] == n["row"])
            b = next((c for k, c in raw_row["cells"].items() if BB.col_of(k) == "B"), None)
            if b and BB.num(b["v"]) is not None and not b["f"]:
                typed.append({"file": n["file"], "row": n["row"], "label": n["label_ar"], "value": b["v"]})
            h = n.get("rebar_t")
            if h is not None and not n.get("rebar_formula"):
                typed.append({"file": n["file"], "row": n["row"], "label": n["label_ar"] + " (rebar t)", "value": h})
    findings.append({"check": "TYPED_CONSTANTS_ON_COVERS", "items": typed, "class": "REPORTING_DEFECT",
                     "impact": "values not traceable to a detail row (rebar tonnage, pool, waterproofing, railing)"})
    # unit ambiguity
    findings.append({"check": "UNIT_AMBIGUITY", "file": "FIN", "label": "نعلات العازل للحمامات",
                     "detail": "upturn computed as perimeter x 1.0 m (m2) but reported in m.t (numerically equal)",
                     "class": "REPORTING_DEFECT"})
    return {"SCHEMA": "URBAN_ALSENAN_B1_MANUAL_BOQ_QA_V1", "findings": findings,
            "counts": dict(Counter(f["check"] for f in findings)),
            "rule": "a benchmark defect is recorded as BENCHMARK_POSSIBLE_ERROR; Urban is never forced to match"}


# ------------------------------------------------------------------ web-app mapping
WEB_MAP = {  # web row -> (English interpretation, freelancer item(s) or a non-quantity class)
    "WEB-001": ("formwork timber (lump)", "LUMP_SUM"), "WEB-002": ("props / jacks (lump)", "LUMP_SUM"),
    "WEB-003": ("structure labour (lump)", "LUMP_SUM"), "WEB-004": ("ceramic tiles (material)", "PROCUREMENT"),
    "WEB-005": ("cement", "PROCUREMENT"), "WEB-006": ("sand, 19 m3 loads", "PROCUREMENT"),
    "WEB-007": ("tile adhesive", "PROCUREMENT"), "WEB-008": ("grout", "PROCUREMENT"),
    "WEB-009": ("tile spacers", "PROCUREMENT"), "WEB-010": ("marble stairs", "FIN:الغلاف:24"),
    "WEB-011": ("marble skirting", "FIN:الغلاف:25"), "WEB-012": ("stair nosings", "FIN:الغلاف:26"),
    "WEB-013": ("tile installation per m2", "AGGREGATE_NOT_REPRODUCIBLE"),
    "WEB-014": ("skirting installation", "FIN:الغلاف:16+FIN:الغلاف:25"), "WEB-015": ("nosings labour", "FIN:الغلاف:26"),
    "WEB-016": ("yard floors labour", "FIN:الغلاف:22"), "WEB-017": ("yard skirting labour", "FIN:الغلاف:23"),
    "WEB-018": ("showers", "NO_FREELANCER_ITEM"), "WEB-019": ("floor drains installation", "NO_FREELANCER_ITEM"),
    "WEB-020": ("45-degree corners", "NO_FREELANCER_ITEM"), "WEB-021": ("stair installation", "FIN:الغلاف:24"),
    "WEB-022": ("stair skirting labour", "FIN:الغلاف:25"), "WEB-023": ("scaffold (lump)", "LUMP_SUM"),
    "WEB-024": ("internal plaster", "PLS:الفاتورة:16"), "WEB-025": ("plaster corners (gross metres)", "PLS:الفاتورة:17#gross"),
    "WEB-026": ("spatter coat bathrooms", "PLS:الفاتورة:18"), "WEB-027": ("spatter coat under skirting", "PLS:الفاتورة:19"),
    "WEB-028": ("facade lump item (freelancer row has no quantity)", "LUMP_SUM"),
    "WEB-029": ("external plaster", "PLS:الفاتورة:25"),
    "WEB-030": ("external corners / joints / beads (gross metres)", "PLS:الفاتورة:26#gross"),
    "WEB-031": ("sigma external finish (rate 0)", "PLS:الفاتورة:25"), "WEB-032": ("internal mesh", "PROCUREMENT"),
    "WEB-033": ("external mesh", "PROCUREMENT"), "WEB-034": ("external plaster corners (material)", "PROCUREMENT"),
    "WEB-035": ("additional materials", "LUMP_SUM"), "WEB-036": ("cement external", "PROCUREMENT"),
    "WEB-037": ("cement internal", "PROCUREMENT"), "WEB-038": ("washed sand loads", "PROCUREMENT"),
    "WEB-039": ("corner beads", "PROCUREMENT"), "WEB-040": ("sigma material", "PROCUREMENT"),
    "WEB-041": ("zero aggregate", "PROCUREMENT"), "WEB-042": ("bonding agent", "PROCUREMENT"),
    "WEB-043": ("PVC beads", "PROCUREMENT"), "WEB-044": ("concrete pump 10 m trips", "LOGISTICS"),
    "WEB-045": ("concrete pump 30 m trips", "LOGISTICS"), "WEB-046": ("block transport", "LOGISTICS"),
    "WEB-047": ("porcelain transport", "LOGISTICS"), "WEB-048": ("sand transport", "LOGISTICS"),
    "WEB-049": ("steel transport", "LOGISTICS"), "WEB-050": ("cleaning", "LOGISTICS"),
    "WEB-051": ("ready-mix concrete", "RC:ورقة1:24"), "WEB-052": ("plain concrete", "RC:ورقة1:15"),
    "WEB-053": ("air conditioning (subcontract)", "LUMP_SUM"), "WEB-054": ("Kuwaiti rebar", "RC:ورقة1:25"),
    "WEB-055": ("sanitary materials", "LUMP_SUM"), "WEB-056": ("Arabic boiler system", "LUMP_SUM"),
    "WEB-057": ("water tanks", "NO_FREELANCER_ITEM"), "WEB-058": ("pump pit", "NO_FREELANCER_ITEM"),
    "WEB-059": ("pit covers", "NO_FREELANCER_ITEM"), "WEB-060": ("sanitary labour", "LUMP_SUM"),
    "WEB-061": ("electrical materials", "LUMP_SUM"), "WEB-062": ("electrical labour", "LUMP_SUM"),
    "WEB-063": ("aluminium (lump)", "LUMP_SUM_FOR:ALU:الغلاف:20"), "WEB-064": ("blocks, 22 thousand (pieces)", "PROCUREMENT"),
    "WEB-065": ("white blocks", "PROCUREMENT"), "WEB-066": ("mortar", "PROCUREMENT"),
    "WEB-067": ("external door / window lintels", "LUMP_SUM"), "WEB-068": ("ties", "PROCUREMENT"),
    "WEB-069": ("ties + shots", "PROCUREMENT"), "WEB-070": ("white cement", "PROCUREMENT"),
    "WEB-071": ("nail-gun shots", "PROCUREMENT"), "WEB-072": ("pool (subcontract)", "LUMP_SUM"),
    "WEB-073": ("gypsum decor (subcontract)", "LUMP_SUM_FOR:FIN:الغلاف:34+FIN:الغلاف:36"),
    "WEB-074": ("internal paint labour", "LUMP_SUM_FOR:PNT:الفاتورة:22"), "WEB-075": ("paint materials", "LUMP_SUM_FOR:PNT:الفاتورة:22"),
    "WEB-076": ("wood doors", "NO_FREELANCER_ITEM"), "WEB-077": ("aluminium doors", "ALU:الكميات:14+ALU:الكميات:51"),
    "WEB-078": ("cement (waterproofing)", "PROCUREMENT"), "WEB-079": ("membrane rolls", "PROCUREMENT"),
    "WEB-080": ("bitumen", "PROCUREMENT"), "WEB-081": ("special materials", "PROCUREMENT"),
    "WEB-082": ("waterproofing floors and roofs", "FIN:العازل:50"), "WEB-083": ("waterproofing upturns", "FIN:العازل:51"),
    "WEB-084": ("railing", "FIN:الغلاف:39"), "WEB-085": ("roof steel ladder", "NO_FREELANCER_ITEM"),
    "WEB-086": ("ventilation (subcontract)", "LUMP_SUM"), "WEB-087": ("scaffold labour", "LUMP_SUM"),
    "WEB-088": ("extra labourers", "LABOUR_DAYS"),
}


def web_mapping(web, N):
    by_id = {n["id"]: n for n in N}
    rows = []
    for r in web["rows"]:
        interp, tgt = WEB_MAP.get(r["id"], ("(unmapped)", "UNMAPPED"))
        lump = tgt.startswith("LUMP_SUM_FOR:")
        srcs = [x for x in (tgt.split(":", 1)[1] if lump else tgt).split("+") if x.split("#")[0].count(":") == 2]
        recs = [(by_id[x.split("#")[0]], x.split("#")[1] if "#" in x else None) for x in srcs if x.split("#")[0] in by_id]
        fq = None
        if recs and not lump:
            fq = r6(sum((x["count"] if x["trade"] == "ALUMINIUM" and x["subtrade"] == "DOOR" else
                         x["gross"] if f == "gross" else x["qty"]) for x, f in recs))
        rel = r6((r["qty"] - fq) / fq * 100) if fq else None
        rows.append({"id": r["id"], "page": r["page"], "category_raw": r["category_raw"], "category_nfkc": r["category_nfkc"],
                     "type_raw": r["type_raw"], "description_raw": r["desc_raw"], "description_nfkc": r["desc_nfkc"],
                     "interpretation_en": interp, "qty": r["qty"], "unit": r["unit"], "rate_kwd": r["rate_kwd"],
                     "total_kwd": r["total_kwd"], "arith_ok": r["arith_ok"],
                     "source_freelancer_items": [x["id"] for x, _ in recs] if recs else [],
                     "mapping_class": ("FREELANCER_QTY" if fq is not None and abs(rel or 0) <= 10 else
                                       "FREELANCER_ITEM_DIFFERENT_QTY" if fq is not None else
                                       "LUMP_SUM_OVER_FREELANCER_ITEM" if lump else tgt.split(":")[0]),
                     "freelancer_qty": fq, "web_vs_freelancer_pct": rel})
    mapped = [x for x in rows if x["mapping_class"] == "FREELANCER_QTY"]
    return {"SCHEMA": "URBAN_ALSENAN_B1_WEB_APP_ROWS_V1", "rows": rows, "categories": web["categories"],
            "checks": web["checks"], "mapping_counts": dict(Counter(x["mapping_class"] for x in rows)),
            "arithmetic_inconsistent_rows": [x["id"] for x in rows if x["arith_ok"] is False],
            "rounding_up_range_pct": [min(x["web_vs_freelancer_pct"] for x in mapped if x["web_vs_freelancer_pct"] is not None),
                                      max(x["web_vs_freelancer_pct"] for x in mapped if x["web_vs_freelancer_pct"] is not None)]
            if mapped else None,
            "dependency": "DERIVED_FROM_FREELANCER_BOQ", "pricing": "recorded, NOT evaluated in Phase B1"}


# ------------------------------------------------------------------ backlog / owner questions
def backlog(D):
    items = [
        {"id": "FIX-B1-01", "priority": "P1", "title": "Opening type from plan symbol alone (window vs glazed door)",
         "evidence": "GF: Urban-only straight widths equal the manual aluminium external-door width; window counts "
                     "are published COMPLETE",
         "proposed_fix": "publish the glazed-opening COUNT as complete but its window / door TYPE as BLOCKED unless a "
                         "sill / head / door-swing / elevation fact proves it (generic OPENING_TYPE_AUTHORITY)",
         "materiality": "MEDIUM", "generic": "HIGH", "silent_error_risk": "HIGH", "ease": "MEDIUM", "trades": ["WINDOWS", "DOORS", "ALUMINIUM", "PAINT", "PLASTER", "BLOCKWORK"]},
        {"id": "FIX-B1-02", "priority": "P1", "title": "Column / wall clear height method as a declared, owner-approved rule",
         "evidence": "manual heights = floor-to-floor - 0.75 m (beam soffit) for columns, 3.6 / 3.4 m for walls",
         "proposed_fix": "a versioned HEIGHT_METHOD fact (e.g. column = storey height - main beam depth) applied only "
                         "when the owner approves it; never inferred from the benchmark",
         "materiality": "HIGH", "generic": "HIGH", "silent_error_risk": "LOW", "ease": "HIGH",
         "trades": ["STRUCTURAL CONCRETE", "BLOCKWORK", "PLASTER", "PAINT", "TILE"]},
        {"id": "FIX-B1-03", "priority": "P1", "title": "Simple-beam binding: mark beside its band",
         "evidence": "Urban measured 1 simple band (B24); the manual measures every beam",
         "proposed_fix": "bind a mark to the parallel band whose breadth equals the scheduled B and whose centre line "
                         "is within one band width, ambiguity-guarded (refuse when two bands qualify)",
         "materiality": "HIGH", "generic": "HIGH", "silent_error_risk": "MEDIUM", "ease": "MEDIUM", "trades": ["STRUCTURAL CONCRETE"]},
        {"id": "FIX-B1-04", "priority": "P1", "title": "Slab-edge outline + void deduction",
         "evidence": "manual slab areas 207 / 120 / 45.19 m2; Urban SLAB_OUTLINE_NOT_ESTABLISHED",
         "proposed_fix": "slab-edge role inference on the structural plans (closed outline on the slab layer, void "
                         "labels and opening hatches subtracted), beam/slab overlap rule declared",
         "materiality": "HIGH", "generic": "HIGH", "silent_error_risk": "LOW", "ease": "MEDIUM", "trades": ["STRUCTURAL CONCRETE"]},
        {"id": "FIX-B1-05", "priority": "P2", "title": "CURVED_OPENING_MEASUREMENT_POLICY",
         "evidence": "GF manual 6.57 m is below every Urban arc (closest: inner 6.658 m); 1F manual 3.15 m equals the mean",
         "proposed_fix": "publish all bases (inner / centre / outer / chord) and let a declared policy pick the "
                         "commercial basis; do not choose from this benchmark",
         "materiality": "LOW", "generic": "MEDIUM", "silent_error_risk": "MEDIUM", "ease": "HIGH", "trades": ["ALUMINIUM"]},
        {"id": "FIX-B1-06", "priority": "P2", "title": "Open-zone closure at the GF entrance (reception / dining / saloon)",
         "evidence": "manual treats the zone as one closed room (137.5 m2) plus a separate entrance room (10.46 m2)",
         "proposed_fix": "closure only from an owner answer (OQ3-A1) or a proven door / screen symbol; then certify",
         "materiality": "HIGH", "generic": "MEDIUM", "silent_error_risk": "LOW", "ease": "HIGH", "trades": ["FLOOR", "CEILING", "SKIRTING", "PAINT", "PLASTER"]},
        {"id": "FIX-B1-07", "priority": "P2", "title": "Room certification for multi-label and split-label sites",
         "evidence": "dewaneya, GF master bedroom, 1F laundry / bedrooms are measured by the manual; Urban blocks them",
         "proposed_fix": "separate label sites by proven single-line partitions / joinery; certify sites whose only "
                         "issue is an inferred wardrobe box (OQ3-A2)",
         "materiality": "HIGH", "generic": "HIGH", "silent_error_risk": "LOW", "ease": "LOW", "trades": ["FLOOR", "CEILING", "SKIRTING"]},
        {"id": "FIX-B1-08", "priority": "P3", "title": "Continuous-beam length per floor and per occurrence",
         "evidence": "the manual reports CBs per floor with drawn lengths; Urban reports one schedule row",
         "proposed_fix": "report CB occurrences per slab sheet and both length bases (schedule span / measured)",
         "materiality": "LOW", "generic": "MEDIUM", "silent_error_risk": "LOW", "ease": "HIGH", "trades": ["STRUCTURAL CONCRETE"]},
        {"id": "FIX-B1-09", "priority": "P3", "title": "Blinding / ground slab / ground beams / perimeter strap items",
         "evidence": "manual carries them; Urban has no item",
         "proposed_fix": "add the items with source-based extents (footprint offset rule declared) or BLOCKED rows",
         "materiality": "HIGH", "generic": "HIGH", "silent_error_risk": "MEDIUM (missing item = silent omission)",
         "ease": "MEDIUM", "trades": ["STRUCTURAL CONCRETE"]},
        {"id": "FIX-B1-10", "priority": "P4", "title": "Rebar tonnage", "evidence": "manual = typed allowances (75-200 kg/m3)",
         "proposed_fix": "SOURCE_IMPOSSIBLE without bar schedules (cutting lengths, laps, hooks); keep fail-closed",
         "materiality": "HIGH", "generic": "-", "silent_error_risk": "LOW", "ease": "-", "trades": ["REBAR"]},
        {"id": "FIX-B1-11", "priority": "P4", "title": "Finish heights / neck height / ground-beam extents",
         "evidence": "manual heights and round ground-beam lengths are not in the source",
         "proposed_fix": "SOURCE_IMPOSSIBLE without levels / false-ceiling heights: owner method facts only",
         "materiality": "HIGH", "generic": "-", "silent_error_risk": "LOW", "ease": "-", "trades": ["PAINT", "PLASTER", "TILE", "BLOCKWORK", "STRUCTURAL CONCRETE"]},
    ]
    return {"SCHEMA": "URBAN_ALSENAN_B1_PROPOSED_FIX_BACKLOG_V1", "items": items,
            "ranking": "A materiality, B generic usefulness, C silent-error risk, D ease, E trades affected",
            "source_impossible": [i["id"] for i in items if i["priority"] == "P4"],
            "implemented": "NONE - Phase B1 changes no engine"}


OWNER_QUESTIONS = [
    {"id": "OQ-B1-1", "group": "ALUMINIUM",
     "question": "Salon sea-view opening: your derived height is 3.65 m; the freelancer measured both ground-floor "
                 "pool windows at 4.30 m. Which height applies to the salon aluminium (23.10 m2 at 3.65 m, 27.22 m2 at "
                 "4.30 m)?", "impact": "4.11 m2 of salon aluminium",
     "evidence": "ARCHITECTURAL_COMPARISON.aluminium.salon"},
    {"id": "OQ-B1-2", "group": "OPENINGS",
     "question": "Ground floor: the drawing shows five 1.00 m glazed openings in external walls; the freelancer lists "
                 "four 1.00 x 2.40 m aluminium doors and one 1.00 m window. Are four of them glazed aluminium doors?",
     "impact": "window vs door count and areas", "evidence": "ARCHITECTURAL_COMPARISON.aluminium.windows[GF]"},
    {"id": "OQ-B1-3", "group": "ALUMINIUM",
     "question": "For curved aluminium glazing, on which line do you price the length - inside face, centre line or "
                 "outside face of the frame?", "impact": "master-bedroom curved glazing basis (6.66 / 6.72 / 6.79 m)",
     "evidence": "ARCHITECTURAL_COMPARISON.aluminium.master_bedroom_curved"},
]
CARRIED = [
    {"id": "OQ3-S1", "new_evidence": "the freelancer counts F10 once at its schedule size (1.96 m3) and does not count "
                                     "the F inside the same outline (F count 4 = F tags outside the outline)"},
    {"id": "OQ3-A1", "new_evidence": "the freelancer treats reception + dining (+ saloon) as one closed room of "
                                     "137.5 m2 and the entrance hall as a separate 10.46 m2 room"},
    {"id": "OQ3-A2", "new_evidence": "the freelancer master-bedroom area 16.66 m2 is not smaller than Urban's 16.55 m2 "
                                     "candidate: the 0.50 x 1.46 m box is not deducted (wardrobe-like)"},
]


# ------------------------------------------------------------------ registers
def metrics(rows):
    out = {}
    for t in TRADES:
        rs = [r for r in rows if r["trade"] == t]
        ex = [r for r in rs if r["comparability"] == "EXACT_COMPARABLE" and r["pct_diff"] is not None
              and r["confidence"] in ("HIGH", "MEDIUM") and r["metric_kind"] in ("area", "volume", "linear")]
        hi = [r for r in ex if r["confidence"] == "HIGH"]
        exc = [r for r in rs if r["comparability"] == "EXACT_COMPARABLE" and r["metric_kind"] == "count"]
        out[t] = {"rows": len(rs), "exact_comparable": sum(r["comparability"] == "EXACT_COMPARABLE" for r in rs),
                  "quantity_rows_high": len(hi), "quantity_rows_medium": len(ex) - len(hi),
                  "within_tolerance_high": sum(r["band"] in ("EXACT", "CLOSE") for r in hi),
                  "max_abs_pct_high": max((abs(r["pct_diff"]) for r in hi), default=None),
                  "median_abs_pct_high": statistics.median([abs(r["pct_diff"]) for r in hi]) if hi else None,
                  "count_rows": len(exc), "count_rows_equal": sum(r["band"] == "EXACT" for r in exc)}
    return {"by_trade": out, "global_accuracy": "NOT COMPUTED (by rule)",
            "rule": "only EXACT_COMPARABLE rows with HIGH confidence enter the high metrics; MEDIUM shown apart; "
                    "LOW never assesses engine accuracy"}


def build(bench_dir, commit=None) -> dict:
    tol = tolerances()
    hashes = BB.verify(bench_dir)
    R3 = {p.stem: jl(p) for p in sorted(A3_DIR.glob("*.json"))}
    a3_sha = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(A3_DIR.glob("*.json"))}
    if a3_sha.get(f"{A3_FREEZE}.json") != A3_FREEZE_SHA:
        raise RuntimeError("Phase A3 freeze changed - B1 compares only against the frozen A3")
    raw, sheets = [], {}
    for k, f in BB.FILES.items():
        p = Path(bench_dir) / f["name"]
        if f["name"].endswith(".xlsx"):
            rr, sh = BB.raw_xlsx(p, k)
        elif f["name"].endswith(".xls"):
            rr, sh = BB.raw_xls(p, k)
        else:
            continue
        raw += rr
        sheets[k] = sh
    N = BB.normalise(raw)
    web = BB.web_rows(Path(bench_dir) / BB.FILES["WEB"]["name"])
    bk = Book(tol)
    D = {}
    D["footings"] = footings(bk, R3, N)
    D["f_f10"] = f_f10(R3, N)
    D["straps"] = straps(bk, R3, N)
    D["foundation"] = foundation(bk, R3, N, D["footings"], D["straps"])
    D["columns"] = columns(bk, R3, N, R3["PDF_VERTICAL_EVIDENCE_REGISTER"])
    D["beams"] = beams(bk, R3, N)
    D["slabs"] = slabs(bk, R3, N)
    D["rebar"] = rebar(bk, N)
    other_structure(bk, N)
    D["rooms"] = rooms(bk, R3, N)
    D["ceilings"] = ceilings(bk, R3, N)
    D["walls"] = wall_finishes(bk, R3, N)
    waterproofing(bk, N)
    D["blockwork"] = blockwork(bk, R3, N)
    D["aluminium"] = aluminium(bk, R3, N, tol)
    D["reception"] = reception(bk, R3, N)
    D["master_bedroom"] = master_bedroom(R3, N)
    D["marble"] = marble(bk, N)
    urban_only(bk, R3)
    qa = manual_qa(raw, N)
    wm = web_mapping(web, N)
    return {"tol": tol, "hashes": hashes, "sheets": sheets, "raw": raw, "N": N, "web": wm, "rows": bk.rows, "D": D,
            "qa": qa, "a3_sha": a3_sha, "R3": R3, "commit": commit or _git("rev-parse", "--short", "HEAD")}


def _git(*a):
    try:
        return subprocess.run(["git", *a], capture_output=True, text=True, cwd=ROOT, check=True).stdout.strip()
    except Exception:                                                                # pragma: no cover
        return None


def registers(ctx) -> dict:
    rows, D, N = ctx["rows"], ctx["D"], ctx["N"]
    regs = {}
    regs["BENCHMARK_SOURCE_MANIFEST"] = {
        "SCHEMA": "URBAN_ALSENAN_B1_BENCHMARK_SOURCE_MANIFEST_V1", "reader_policy": BB.policy_record(),
        "files": [{"key": k, "filename": f["name"], "type": Path(f["name"]).suffix.lstrip("."), "size_bytes": ctx["hashes"][k]["bytes"],
                   "sha256": f["sha256"], "hash_verified": ctx["hashes"][k]["match"], "kind": f["kind"],
                   "family": f["family"], "title_en": f["title_en"], "language": "Arabic (labels), English (codes)",
                   "discipline": "STRUCTURAL" if f["family"] == "CONCRETE" else ("COMMERCIAL" if k == "WEB" else "ARCHITECTURAL"),
                   "sheets_or_pages": ctx["sheets"].get(k) or {"pages": ctx["web"]["checks"]["pages"]},
                   "formulas": "LIVE (xlsx)" if f["name"].endswith(".xlsx") else ("VALUES_ONLY (BIFF xls)" if f["name"].endswith(".xls") else "TEXT LAYER (pdf)")}
                  for k, f in BB.FILES.items()],
        "immutability": "files are read from a read-only pinned copy; every run re-hashes and fails closed on mismatch; "
                        "the binaries are not committed",
        "not_used": ["the older web-app PDF 737b63b5 (different hash)", "any other upload"]}
    regs["BENCHMARK_DEPENDENCY"] = {
        "SCHEMA": "URBAN_ALSENAN_B1_BENCHMARK_DEPENDENCY_V1", "WEB_APP_DEPENDENCY": "DERIVED_FROM_FREELANCER_BOQ",
        "basis": ["owner statement", f"{ctx['web']['mapping_counts'].get('FREELANCER_QTY', 0)} web rows carry a freelancer "
                  f"quantity rounded up by {ctx['web']['rounding_up_range_pct']} %"],
        "votes": "a web-app quantity is never a second confirmation of a freelancer quantity",
        "freelancer_files_relationship": "six workbooks by the same measurer (covers NO:1, 2, 3, 3, 5, 6; NO:4 not supplied); "
                                         "covers summarise their own detail sheets by cell reference",
        "use_of_web_app": ["category mapping", "commercial presentation", "cost structure", "current Urban Projects workflow"]}
    regs["FREELANCER_RAW_ROWS"] = {"SCHEMA": "URBAN_ALSENAN_B1_FREELANCER_RAW_ROWS_V1", "origin": BB.ORIGIN,
                                   "rows": ctx["raw"], "count": len(ctx["raw"]),
                                   "per_file": dict(Counter(r["file"] for r in ctx["raw"]))}
    regs["WEB_APP_ROWS"] = ctx["web"]
    regs["BENCHMARK_NORMALISED"] = {"SCHEMA": "URBAN_ALSENAN_B1_BENCHMARK_NORMALISED_V1", "origin": BB.ORIGIN,
                                    "records": N, "count": len(N),
                                    "by_kind": dict(Counter(n["row_kind"] for n in N)),
                                    "quantity_detail_rows": sum(1 for n in N if n["row_kind"] == "DETAIL" and n["qty"] is not None),
                                    "rule": "original Arabic description and original unit kept; no value changed"}
    regs["COMPARABILITY_MATRIX"] = {"SCHEMA": "URBAN_ALSENAN_B1_COMPARABILITY_MATRIX_V1", "tolerances": ctx["tol"],
                                    "rows": rows, "counts": dict(Counter(r["comparability"] for r in rows)),
                                    "confidence": dict(Counter(r["confidence"] for r in rows)),
                                    "metrics": metrics(rows)}
    regs["FOOTING_COMPARISON"] = {"SCHEMA": "URBAN_ALSENAN_B1_FOOTING_COMPARISON_V1", **D["footings"],
                                  "straps": D["straps"], "foundation": D["foundation"]}
    regs["F_F10_FORENSIC_COMPARISON"] = D["f_f10"]
    regs["STRUCTURAL_COMPARISON"] = {"SCHEMA": "URBAN_ALSENAN_B1_STRUCTURAL_COMPARISON_V1", "columns": D["columns"],
                                     "beams": D["beams"], "slabs": D["slabs"], "rebar": D["rebar"],
                                     "rows": [r["id"] for r in rows if r["family"] not in
                                              ("ROOM_FLOOR", "ROOM_PERIMETER", "FLOOR_TOTAL", "CEILING", "WALL_TILE",
                                               "WALL_FINISH", "WATERPROOFING", "BLOCKWORK", "ALUMINIUM", "WINDOWS",
                                               "DOORS", "RECEPTION", "MARBLE", "OTHER_ARCH", "URBAN_ONLY")]}
    regs["ARCHITECTURAL_COMPARISON"] = {"SCHEMA": "URBAN_ALSENAN_B1_ARCHITECTURAL_COMPARISON_V1", "rooms": D["rooms"],
                                        "ceilings": D["ceilings"], "walls": D["walls"], "blockwork": D["blockwork"],
                                        "aluminium": D["aluminium"], "reception": D["reception"],
                                        "master_bedroom": D["master_bedroom"], "marble": D["marble"]}
    regs["MANUAL_BOQ_QA"] = ctx["qa"]
    diffs = [r for r in rows if r["primary_class"] != "MATCH_WITHIN_TOLERANCE"]
    regs["DIFFERENCE_REGISTER"] = {"SCHEMA": "URBAN_ALSENAN_B1_DIFFERENCE_REGISTER_V1",
                                   "differences": [{k: r[k] for k in ("id", "family", "trade", "item", "floor", "unit",
                                                                       "urban_qty", "bench_qty", "comparability",
                                                                       "abs_diff", "pct_diff", "band", "primary_class",
                                                                       "secondary_classes", "confidence", "note")}
                                                   for r in diffs],
                                   "primary_counts": dict(Counter(r["primary_class"] for r in rows)),
                                   "difference_primary_counts": dict(Counter(r["primary_class"] for r in diffs)),
                                   "qa_findings": [f["check"] for f in ctx["qa"]["findings"]],
                                   "rule": "classified, NOT fixed; a difference is not automatically an Urban error"}
    regs["PROPOSED_FIX_BACKLOG"] = backlog(D)
    regs["OWNER_QUESTION_REGISTER_B1"] = {"SCHEMA": "URBAN_ALSENAN_B1_OWNER_QUESTIONS_V1", "new": OWNER_QUESTIONS,
                                          "carried_forward": CARRIED,
                                          "never_asked": ["expected quantities", "files (no extra DWG / schedules / levels)"]}
    regs[FREEZE] = freeze(ctx, regs)
    regs["TEST_RESULTS"] = {"SCHEMA": "URBAN_ALSENAN_B1_TEST_RESULTS_V1", "state": "RECORDED_BY_THE_PACKAGE_STEP",
                            "why": "the one full suite runs from the final commit"}
    return regs


def freeze(ctx, regs) -> dict:
    eng_a3, eng_now = _git("rev-parse", f"{A3_CODE_COMMIT}:engine"), _git("rev-parse", "HEAD:engine")
    rec = {"SCHEMA": "URBAN_ALSENAN_PHASE_B1_COMPARISON_FREEZE_V1", "phase": "PHASE_B1",
           "nature": "benchmark reveal + forensic comparison ONLY - anti-calibration boundary for Phase B2",
           "code_commit": ctx["commit"], "recommendation_commit": RECOMMENDATION_COMMIT,
           "phase_a3": {"registers_commit": A3_COMMIT, "code_commit": A3_CODE_COMMIT, "freeze_file": f"{A3_FREEZE}.json",
                        "freeze_sha256": ctx["a3_sha"][f"{A3_FREEZE}.json"],
                        "registers_digest": digest(ctx["a3_sha"])},
           "engine_tree": {"at_a3_code_commit": eng_a3, "at_b1_build": eng_now, "unchanged": eng_a3 == eng_now},
           "benchmark_sha256": {k: f["sha256"] for k, f in BB.FILES.items()},
           "extracted_rows_digest": digest(ctx["raw"]), "normalised_digest": digest(ctx["N"]),
           "web_rows_digest": digest(ctx["web"]["rows"]),
           "mapping_digest": digest([(r["id"], r["urban_ref"], r["bench_refs"], r["comparability"]) for r in ctx["rows"]]),
           "comparison_digest": digest(ctx["rows"]),
           "classification_digest": digest([(r["id"], r["primary_class"], r["secondary_classes"]) for r in ctx["rows"]]),
           "manual_qa_digest": digest(ctx["qa"]), "owner_questions_digest": digest([OWNER_QUESTIONS, CARRIED]),
           "backlog_digest": digest(regs["PROPOSED_FIX_BACKLOG"]["items"]),
           "tolerances": ctx["tol"], "counts": {"raw_rows": len(ctx["raw"]), "normalised": len(ctx["N"]),
                                                "comparison_rows": len(ctx["rows"]),
                                                "comparability": dict(Counter(r["comparability"] for r in ctx["rows"]))},
           "BENCHMARK_OPENED": True, "ENGINE_CHANGED": eng_a3 != eng_now, "QORTUBA_CHANGED": False,
           "PHASE_A3_FREEZE_PRESERVED": ctx["a3_sha"][f"{A3_FREEZE}.json"] == A3_FREEZE_SHA,
           "rule": "Phase B2 may change generic engines only against PROPOSED_FIX_BACKLOG items; any A3 quantity that "
                   "moves toward a benchmark value must cite a generic defect, never a benchmark row"}
    return rec


# ------------------------------------------------------------------ workbook (deterministic, no formulas)
def _cells(v):
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    return v


def workbook(regs) -> dict:
    cm = regs["COMPARABILITY_MATRIX"]["rows"]
    pick = lambda *fam: [r for r in cm if r["family"] in fam]                       # noqa: E731
    H = ["ID", "FAMILY", "TRADE", "ITEM", "FLOOR", "UNIT", "URBAN", "BENCHMARK", "COMPARABILITY", "ABS DIFF",
         "PCT DIFF", "BAND", "PRIMARY CLASS", "SECONDARY", "CONFIDENCE", "URBAN REF", "BENCHMARK REFS", "NOTE"]
    line = lambda r: [r["id"], r["family"], r["trade"], r["item"], r["floor"], r["unit"], r["urban_qty"], r["bench_qty"],  # noqa: E731
                      r["comparability"], r["abs_diff"], r["pct_diff"], r["band"], r["primary_class"],
                      ", ".join(r["secondary_classes"]), r["confidence"], r["urban_ref"], ", ".join(r["bench_refs"]),
                      r["note"] or r.get("method") or ""]
    m = {}
    m["00_READ_ME"] = (["TOPIC", "TEXT"], [
        ["WHAT", "Phase B1: freelancer BOQ + web-app report extracted, mapped and compared with the frozen Phase A3."],
        ["RULE", "Only EXACT_COMPARABLE rows carry a percentage; a difference is classified, not fixed."],
        ["RULE", "No engine was changed; no benchmark value enters an engine; no pricing is evaluated."],
        ["A3", f"{A3_COMMIT} / freeze sha256 {A3_FREEZE_SHA}"]])
    m["01_BENCHMARK_FILES"] = (["KEY", "FILE", "SHA256", "KIND", "FAMILY", "TITLE", "FORMULAS"],
                               [[f["key"], f["filename"], f["sha256"], f["kind"], f["family"], f["title_en"], f["formulas"]]
                                for f in regs["BENCHMARK_SOURCE_MANIFEST"]["files"]])
    mt = regs["COMPARABILITY_MATRIX"]["metrics"]["by_trade"]
    m["02_COMPARISON_SUMMARY"] = (["TRADE", "ROWS", "EXACT COMPARABLE", "QTY ROWS HIGH", "WITHIN TOL HIGH",
                                   "MAX |PCT| HIGH", "MEDIAN |PCT| HIGH", "COUNT ROWS", "COUNT ROWS EQUAL"],
                                  [[t, v["rows"], v["exact_comparable"], v["quantity_rows_high"], v["within_tolerance_high"],
                                    v["max_abs_pct_high"], v["median_abs_pct_high"], v["count_rows"], v["count_rows_equal"]]
                                   for t, v in mt.items() if v["rows"]])
    for name, fams in (("03_FOOTINGS", ("FOOTING_COUNT", "FOOTING_VOLUME", "FOOTING_TOTAL", "STRAP")),
                       ("04_CONCRETE", ("FOUNDATION", "OTHER_STRUCTURE")), ("05_COLUMNS", ("COLUMN_COUNT", "COLUMN_SECTION", "COLUMN_STOREY", "COLUMN_VOLUME")),
                       ("06_BEAMS", ("BEAM_CB", "BEAM_SECTION", "BEAM_SIMPLE", "BEAM_VOLUME")), ("07_SLABS", ("SLAB_THICKNESS", "SLAB_VOLUME")),
                       ("08_REBAR", ("REBAR",)), ("09_ROOM_AREAS", ("ROOM_FLOOR", "ROOM_PERIMETER", "FLOOR_TOTAL", "CEILING")),
                       ("10_BLOCKWORK", ("BLOCKWORK",)), ("11_PLASTER_PAINT", ("WALL_FINISH", "RECEPTION")),
                       ("12_TILES", ("WALL_TILE", "MARBLE", "OTHER_ARCH")), ("13_WATERPROOFING", ("WATERPROOFING",)),
                       ("14_ALUMINIUM", ("ALUMINIUM",)), ("15_DOORS_WINDOWS", ("WINDOWS", "DOORS"))):
        m[name] = (H, [line(r) for r in pick(*fams)])
    m["16_WEB_APP_MAPPING"] = (["ID", "PAGE", "CATEGORY (NFKC)", "DESCRIPTION (NFKC)", "INTERPRETATION", "QTY", "UNIT",
                                "RATE KWD", "TOTAL KWD", "MAPPING", "FREELANCER ITEMS", "FREELANCER QTY", "WEB vs FREELANCER %"],
                               [[w["id"], w["page"], w["category_nfkc"], w["description_nfkc"], w["interpretation_en"], w["qty"],
                                 w["unit"], w["rate_kwd"], w["total_kwd"], w["mapping_class"], ", ".join(w["source_freelancer_items"]),
                                 w["freelancer_qty"], w["web_vs_freelancer_pct"]] for w in regs["WEB_APP_ROWS"]["rows"]])
    m["17_MANUAL_BOQ_QA"] = (["CHECK", "FILE", "SHEET", "CELL / ROW", "CLASS", "DETAIL"],
                             [[f["check"], f.get("file"), f.get("sheet"), f.get("cell") or f.get("row"), f["class"],
                               {k: v for k, v in f.items() if k not in ("check", "file", "sheet", "cell", "row", "class")}]
                              for f in regs["MANUAL_BOQ_QA"]["findings"]])
    m["18_DIFFERENCES"] = (H, [line(r) for r in cm if r["primary_class"] != "MATCH_WITHIN_TOLERANCE"])
    m["19_FIX_BACKLOG"] = (["ID", "PRIORITY", "TITLE", "EVIDENCE", "PROPOSED FIX", "MATERIALITY", "GENERIC", "SILENT RISK", "EASE", "TRADES"],
                           [[i["id"], i["priority"], i["title"], i["evidence"], i["proposed_fix"], i["materiality"],
                             i["generic"], i["silent_error_risk"], i["ease"], i["trades"]] for i in regs["PROPOSED_FIX_BACKLOG"]["items"]])
    return {k: {"header": h, "rows": [[_cells(c) for c in r] for r in rs]} for k, (h, rs) in m.items()}


def write_xlsx(model, path) -> dict:
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    sys.path.insert(0, str(ROOT))
    from engine import boq_rc1_xlsx as BX                                        # used read-only (zip repack)
    wb = Workbook()
    wb.remove(wb.active)
    for name, s in model.items():
        ws = wb.create_sheet(name)
        ws.append([BANNER])
        ws.append(["COMPARISON VIEW - NOT ADDITIVE - no formulas; every value copied from a B1 register"])
        ws.append(s["header"])
        for r in s["rows"]:
            ws.append(r)
        ws["A1"].font = Font(bold=True, color="C00000")
        for c in ws[3]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        ws.freeze_panes = "B4"
    wb.properties.creator = "Urban QTO (Phase B1 comparison)"
    wb.properties.created = wb.properties.modified = datetime(2026, 10, 3)
    wb.save(str(path))
    BX._repack(path)
    data = Path(path).read_bytes()
    return {"sheets": list(model), "file_sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data),
            "content_digest": digest(model)}


def readback(path, model) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    diffs, formulas = [], 0
    for name, s in model.items():
        ws = wb[name]
        got = [list(r) for r in ws.iter_rows(min_row=3, values_only=True)]
        want = [s["header"]] + s["rows"]
        for i, (g, w) in enumerate(zip(got, want)):
            g = (g + [None] * len(w))[:len(w)]
            for j, (a, b) in enumerate(zip(g, w)):
                if isinstance(a, str) and a.startswith("="):
                    formulas += 1
                if not (a == b or (a in (None, "") and b in (None, "")) or
                        (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 1e-12)):
                    diffs.append([name, i, j, a, b])
        if len(got) != len(want):
            diffs.append([name, "row_count", len(got), len(want)])
    return {"state": "PASS" if not diffs and not formulas else "FAIL", "differences": diffs[:20], "formulas": formulas}


def main(bench_dir, regdir, commit=None):
    import tempfile
    ctx = build(bench_dir, commit)
    regs = registers(ctx)
    model = workbook(regs)
    with tempfile.TemporaryDirectory() as d:
        a = write_xlsx(model, Path(d) / XLSX_NAME)
        b = write_xlsx(model, Path(d) / ("2_" + XLSX_NAME))
        rb = readback(Path(d) / XLSX_NAME, model)
    regs[FREEZE]["xlsx"] = {"content_digest": a["content_digest"], "file_sha256": a["file_sha256"],
                            "rewrite_identical": a["file_sha256"] == b["file_sha256"], "readback": rb["state"]}
    regdir = Path(regdir)
    regdir.mkdir(parents=True, exist_ok=True)
    for n, o in regs.items():
        (regdir / f"{n}.json").write_text(json.dumps(o, indent=1, ensure_ascii=False, default=str) + "\n")
    cm = regs["COMPARABILITY_MATRIX"]
    print(json.dumps({"raw_rows": len(ctx["raw"]), "normalised": len(ctx["N"]), "comparisons": len(ctx["rows"]),
                      "comparability": cm["counts"], "xlsx": regs[FREEZE]["xlsx"],
                      "engine_unchanged": regs[FREEZE]["engine_tree"]["unchanged"],
                      "primary": regs["DIFFERENCE_REGISTER"]["primary_counts"]}, indent=1, ensure_ascii=False))
    return regs, ctx


if __name__ == "__main__":
    main(*sys.argv[1:4])
