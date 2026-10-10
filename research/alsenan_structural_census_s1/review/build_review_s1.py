"""Human-review deliverables for the frozen Alsenan structural census S1 (no rebar, no kg).

    python research/alsenan_structural_census_s1/review/build_review_s1.py

Writes, next to this script:
    ALSENAN_STRUCTURAL_CENSUS_REVIEW.xlsx         occurrence-level workbook (11 sheets)
    ALSENAN_STRUCTURAL_CENSUS_VISUAL_REVIEW.pdf   annotated plan sheets (vector)
    REVIEW_MANIFEST.json                          register hashes consumed + id cross-check + output hashes

Every label in the PDF and every row in the workbook comes from the frozen S1 register JSONs, whose sha256 are checked
against INDEX.json first. The DXF is read only for the drawing backdrop and for member geometry (line / outline
coordinates), after a rebuild of the census has been proved byte-identical to the frozen registers. Nothing in the
census is changed here: a label that looks untidy on the drawing stays untidy.
"""
import hashlib
import json
import math
import os
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
S1 = HERE.parent
ROOT = S1.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "research" / "external_engine_lab"))
sys.path.insert(0, str(S1))

XLSX = HERE / "ALSENAN_STRUCTURAL_CENSUS_REVIEW.xlsx"
PDF = HERE / "ALSENAN_STRUCTURAL_CENSUS_VISUAL_REVIEW.pdf"
MANIFEST = HERE / "REVIEW_MANIFEST.json"

FLAGS = ("SOURCE_CONFLICT", "BLOCKED", "AMBIGUOUS_MATCH", "MISSING_SCHEDULE_DEFINITION")
FLAG_COLOUR = {"SOURCE_CONFLICT": "#D62728", "BLOCKED": "#FF7F0E", "AMBIGUOUS_MATCH": "#9467BD",
               "MISSING_SCHEDULE_DEFINITION": "#17A2B8"}
FLAG_MARK = {"SOURCE_CONFLICT": "D", "BLOCKED": "s", "AMBIGUOUS_MATCH": "^", "MISSING_SCHEDULE_DEFINITION": "o"}
FLAG_SHORT = {"SOURCE_CONFLICT": "CONFLICT", "BLOCKED": "BLOCKED", "AMBIGUOUS_MATCH": "AMBIGUOUS",
              "MISSING_SCHEDULE_DEFINITION": "NO DEFINITION"}


# ============================================================================================ frozen registers
def load_frozen():
    idx = json.loads((S1 / "INDEX.json").read_text(encoding="utf-8"))
    regs = {}
    for name, v in idx["registers"].items():
        raw = (S1 / v["file"]).read_bytes()
        if hashlib.sha256(raw).hexdigest() != v["sha256"]:
            raise SystemExit(f"register {name} does not match the frozen INDEX hash")
        regs[name] = json.loads(raw)
    return idx, regs


def flags(r):
    """Review markers for one register row (shared by the workbook and the drawings)."""
    f = set()
    issues = " ".join(r.get("issues") or [])
    status = str(r.get("status") or "")
    if status == "SOURCE_CONFLICT" or r.get("type_authority") == "TAG_CONFLICT" or r.get("competing_tags") \
            or "CONFLICT" in issues or "DIFFERS" in issues or "SAME_SPAN" in issues or "OVERLAP" in issues:
        f.add("SOURCE_CONFLICT")
    if r.get("terminal_state") == "COUNTED_BLOCKED" or status == "BLOCKED" or "BLOCKED" in issues \
            or r.get("category_state") == "BLOCKED_CURVED_SPAN":
        f.add("BLOCKED")
    tb = r.get("tie_topology_band") or {}
    if r.get("binding") == "AMBIGUOUS" or "AMBIGUOUS" in issues or tb.get("state") in ("BOUNDARY_GAP", "OUT_OF_RANGE") \
            or r.get("category_state") == "AMBIGUOUS_LENGTH_BASIS" or r.get("_gap_event"):
        f.add("AMBIGUOUS_MATCH")
    if r.get("definition_state") == "NONE" or "TAG_WITHOUT_SCHEDULE_DEFINITION" in issues \
            or "MEMBER_WITH_NO_TAG" in issues or ("footing_id" in r and r.get("type") is None) \
            or status in ("NO_BAR_ANNOTATION", "ANNOTATION_ONE_DIRECTION_ONLY"):
        f.add("MISSING_SCHEDULE_DEFINITION")
    return [x for x in FLAGS if x in f]


def jtxt(v):
    if v is None:
        return None
    if isinstance(v, (str, int, float, bool)):
        return v
    return json.dumps(v, ensure_ascii=False, separators=(", ", ": "))


def bar(s):
    return (s or "").replace("%%C", "Ø").replace("%%c", "Ø") if isinstance(s, str) else s


def sup(s):
    if not s:
        return None
    k = s.get("support_kind") or s.get("kind")
    ref = s.get("ref")
    return f"{k}{(' ' + ref) if ref else ''}"


def bars_txt(b):
    if not b or not isinstance(b, dict) or b.get("count") is None:
        return None
    return f"{b['count']}Ø{b.get('dia_mm')}{'/m' if b.get('per_m') else ''}"


# ============================================================================================ row models
def model(regs):
    """One list of review rows per sheet. Each row carries its census id and its review flags."""
    occ = regs["COLUMN_OCCURRENCE_REGISTER"]["rows"]
    chains = {c["chain_id"]: c for c in regs["COLUMN_VERTICAL_CHAIN_REGISTER"]["rows"]}
    ties = {b["band_id"]: b for b in regs["COLUMN_TIE_TOPOLOGY_RULES"]["rows"]}
    M = {}
    cols = []
    for r in occ:
        tb = r["tie_topology_band"] or {}
        band = tb.get("band") if isinstance(tb.get("band"), str) else None
        ch = chains[r["chain_id"]]
        cols.append({
            "Column ID": r["column_id"], "Chain ID": r["chain_id"], "Grid X": r["grid"]["X"], "Grid Y": r["grid"]["Y"],
            "X (mm, sheet-local)": r["plan_centre_mm"][0], "Y (mm, sheet-local)": r["plan_centre_mm"][1],
            "Type": r["column_type"], "Type authority": r["type_authority"], "Floor": r["floor"],
            "Chain continues through": " > ".join(ch["continues_through"]),
            "From below": r["continues_from_below"], "Continues above": r["continues_above"],
            "Type below": r["type_below"], "Type above": r["type_above"], "Planted on": r["planted_on"],
            "Drawn section (cm)": "x".join(f"{v:g}" for v in r["drawn_section_cm"]),
            "Schedule B x D (cm)": "x".join(f"{v:g}" for v in r["schedule_section_cm"]) if r["schedule_section_cm"] else None,
            "Drawn vs schedule": r["drawn_vs_schedule"],
            "Main bars": bars_txt(r["longitudinal_bars"]) or jtxt(r["longitudinal_bars"]),
            "Tie rule": "Ø8 @ 6/m (P9-COL-TIES; per set vs per tie = BLOCKED_METHOD)",
            "Tie band": band or tb.get("state"), "Tie band state": tb.get("state"),
            "Closed ties / level (detail)": ties[band]["closed_ties_in_section"] if band in ties else None,
            "Bars in detail sketch": ties[band]["bars_drawn"] if band in ties else None,
            "Tmin check": (f"{r['min_thickness_check']['state']} (min {r['min_thickness_check']['t_min_cm']} cm)"
                           if r["min_thickness_check"] else None),
            "Plan sheets": ", ".join(r["plan_source"]["sheets"]), "DXF handles": ", ".join(r["plan_source"]["handles"]),
            "Status": r["status"], "Terminal state": r["terminal_state"], "_flags": flags(r), "_src": r})
    M["01_Columns"] = cols
    chr_ = []
    for c in regs["COLUMN_VERTICAL_CHAIN_REGISTER"]["rows"]:
        evn = [e["event"] for e in c["events"]]
        c2 = dict(c, _gap_event="GAP_IN_CHAIN" in evn)
        chr_.append({
            "Chain ID": c["chain_id"], "Grid X": c["grid"]["X"], "Grid Y": c["grid"]["Y"], "Type": c["column_type"],
            "Type authority": c["type_authority"], "Starts at": c["starts_at"], "Planted on": c["planted_on"],
            "Continues through": " > ".join(c["continues_through"]), "Terminates at": c["terminates_at"],
            "Type changes": jtxt(c["changes_type_at"]) if c["changes_type_at"] else None,
            "Tag CAP": (c["tags_by_plan"].get("CAP") or {}).get("tag"),
            "Tag FP": (c["tags_by_plan"].get("FP") or {}).get("tag"),
            "Tag GBP": (c["tags_by_plan"].get("GBP") or {}).get("tag"),
            "CAP size label": c["axis_plan_size_label"],
            "Members (sheet:handle)": ", ".join(f"{k}:{v}" for k, v in c["members_by_sheet"].items()),
            "Transition outlines": ", ".join(f"{k}:{v}" for k, v in c["transition_outlines"].items()) or None,
            "Labels": ", ".join(f"{lb['label']} ({lb['sheet']}:{lb['handle']})" for lb in c["labels"]) or None,
            "Events": ", ".join(evn) or None,
            "Status": "SOURCE_CONFLICT" if c["type_authority"] == "TAG_CONFLICT" else "COUNTED",
            "_flags": flags(dict(c2, status="SOURCE_CONFLICT" if c["type_authority"] == "TAG_CONFLICT" else None)),
            "_src": c})
    M["02_Column_Chains"] = chr_
    # footings
    fdefs = {d["footing_type"]: d for d in regs["FOOTING_DEFINITION_REGISTER"]["rows"]}
    col_by_chain = defaultdict(list)
    for r in occ:
        if r["floor"] == "FOUNDATION":
            col_by_chain[r["chain_id"]].append(r["column_id"])
    overl = set()
    for m in regs["FOOTING_OCCURRENCE_REGISTER"]["mismatches"]:
        if m["kind"] == "FOOTING_OUTLINES_OVERLAP":
            overl |= {f.split(":", 1)[1] for f in m["footings"]}
    fr = []
    for r in regs["FOOTING_OCCURRENCE_REGISTER"]["rows"]:
        d = fdefs.get(r["type"]) if r["type"] else None
        comps = d["components"] if d else []
        o = r["outline"] or {}
        r2 = dict(r, issues=list(r["issues"]) + (["FOOTING_OUTLINES_OVERLAP"] if o.get("handle") in overl else []))
        fr.append({
            "Footing ID": r["footing_id"], "Type": r["type"],
            "Candidate types": ", ".join(r["candidate_types"] or []) or None,
            "Competing tags": ", ".join(t["text"] for t in r["competing_tags"]) or None,
            "Tag": (r["tag"] or {}).get("text"), "Tag handle": (r["tag"] or {}).get("handle"),
            "Outline handle": o.get("handle"), "Outline geometry": o.get("geometry"),
            "Outline bbox (mm)": jtxt(o.get("bbox")),
            "Drawn L x W (cm)": f"{r['sizes']['drawn_L_cm']:g} x {r['sizes']['drawn_W_cm']:g}" if r["sizes"] else None,
            "Schedule L x W x D (cm)": f"{d['L_cm']:g} x {d['W_cm']:g} x {d['D_cm']:g}" if d else None,
            "Size match": r["sizes"]["match"] if r["sizes"] else None,
            "Supported columns": ", ".join(c for ch in r["supported_columns"] for c in col_by_chain.get(ch, [ch])),
            "Supported column types": ", ".join(r["supported_column_types"]),
            "Bars (schedule)": "; ".join(f"{c['component']} {c['count']}Ø{c['dia_mm']}"
                                         f"{'/m' if c['count_mode'] == 'BARS_PER_METRE' else ''}" for c in comps) or None,
            "Count mode": comps[0]["count_mode"] if comps else None,
            "BOXED raw": (d["boxed"].get("raw_boxed_value") if d else None) or None,
            "BOXED semantics": ("BLOCKED" if d and d["boxed"].get("raw_boxed_value") else None),
            "Issues": ", ".join(r2["issues"]) or None, "Terminal state": r["terminal_state"],
            "_flags": flags(r2), "_src": r})
    M["03_Footings"] = fr
    M["03_defs"] = regs["FOOTING_DEFINITION_REGISTER"]["rows"]
    M["03_mismatch"] = regs["FOOTING_OCCURRENCE_REGISTER"]["mismatches"]
    # beams
    bdefs = defaultdict(list)
    for d in regs["BEAM_DEFINITION_REGISTER"]["rows"]:
        bdefs[d["beam_type"]].append(d)

    def sched(t):
        ds = bdefs.get(t) or []
        if not ds:
            return None, None, None
        if len(ds) > 1:
            return " | ".join(f"{d['B_cm']:g}x{d['H_cm']:g}" for d in ds) + " (DUPLICATE KEY)", None, None
        d = ds[0]
        f = d.get("fields") or {}
        if d["element"] == "TYPICAL_DETAIL":
            return d["normalized_rule"], None, None
        if d["element"] == "CONTINUOUS_BEAM":
            b = f"spans {f.get('spans_m')} m; bottom " + ", ".join(bars_txt(x) or "-" for x in f.get("bottom_bars_per_span", [])) \
                + "; support top " + ", ".join(f"{k.replace('support_', '')} {bars_txt(v) or '-'}" for k, v in
                                                (f.get("support_top_bars") or {}).items()) \
                + "; stirrups/m " + ", ".join(bars_txt(x) or "-" for x in f.get("stirrups_per_span", []))
        else:
            b = f"bottom {bars_txt(f.get('bottom'))}; top {bars_txt(f.get('top'))}; stirrups/m {bars_txt(f.get('stirrups_per_m'))}"
        rem = d.get("REMARKS_side_bars")
        return f"{d['B_cm']:g} x {d['H_cm']:g}", b, (bar(rem.get("raw")) if rem else None)

    def brow(r):
        sd, sb, rem = sched(r.get("beam_type")) if r.get("beam_type") else (None, None, None)
        return {
            "Beam occurrence ID": r["beam_id"], "Floor": r["floor"], "Sheet": r.get("sheet"), "Family": r["family"],
            "Type": r.get("beam_type"), "Tag handle": r.get("tag_handle"), "Binding": r.get("binding"),
            "Member": r.get("member"), "Member candidates": ", ".join(r.get("member_candidates") or []) or None,
            "Start support": sup(r.get("start_support")), "End support": sup(r.get("end_support")),
            "Centreline length (m)": r.get("support_centreline_length_m"), "Clear length (m)": r.get("clear_length_m"),
            "Drawn length (m)": r.get("drawn_length_m"), "Drawn width (mm)": r.get("width_drawn_mm"),
            "Schedule B x H (cm)": sd, "Schedule bars": sb, "REMARKS (side bars)": rem,
            "Orientation (deg)": r.get("orientation_deg"), "Issues": ", ".join(r.get("issues") or []) or None,
            "Status": r.get("status"), "Terminal state": r["terminal_state"], "_flags": flags(r), "_src": r}
    bro = regs["BEAM_OCCURRENCE_REGISTER"]["rows"]
    M["04_Simple_Beams"] = [brow(r) for r in bro if r["family"] in ("SIMPLE", "STAIR", "CANTILEVER", "OTHER", "DOME_RING")]
    M["05_Continuous_Beams"] = [brow(r) for r in bro if r["family"] == "CB"]
    M["05_groups"] = regs["CONTINUOUS_BEAM_OCCURRENCE_REGISTER"]["rows"]
    gbr = []
    for r in bro:
        if r["family"] in ("GB", "EXTERIOR_GB", "STRAP", "LIFT_TIE", "LINTEL"):
            x = brow(r)
            x["GB category (c/c)"] = r.get("beam_type") if r["family"] in ("GB", "EXTERIOR_GB") else None
            x["GB category (clear)"] = r.get("category_clear_basis")
            x["Category state"] = r.get("category_state")
            x["Exterior state"] = r.get("exterior_state")
            gbr.append(x)
    M["06_Ground_Strap_Beams"] = gbr
    # slabs
    ann = defaultdict(list)
    for a in regs["SLAB_REBAR_SOURCE_REGISTER"]["rows"]:
        if a.get("panel"):
            ann[a["panel"]].append(a)
    sr = []
    for p in regs["SLAB_PANEL_REGISTER"]["rows"]:
        al = ann.get(p["panel_id"], [])
        edges = "; ".join(f"e{e['edge_index']} {e['length_m']}m {e['support']}"
                          f"{(' ' + e['support_ref']) if e.get('support_ref') else ''} {e['continuity']}"
                          for e in p["edges"] or [])
        sr.append({
            "Panel ID": p["panel_id"], "Floor": p["floor"], "Class": p["class"], "Area (m2)": p["area_m2"],
            "Perimeter (m)": p["perimeter_m"], "Lx (m)": p["Lx_m"], "Ly (m)": p["Ly_m"],
            "Rectangularity": p["rectangularity"], "Vertices": len(p["polygon_mm"]),
            "Thickness (mm)": p["effective_thickness_mm"], "Thickness authority": p["thickness_authority"],
            "Local T marks": ", ".join(f"{t['raw']} ({t['handle']})" for t in p["local_thickness_notes"]) or None,
            "Supports": jtxt(p["support_summary"]), "Continuity": jtxt(p["continuity_summary"]),
            "Edges (support / continuity)": edges or None,
            "Bound annotations": "; ".join(f"{bar(a['raw'])} {a['direction'] or ''} [{a['layer']}]".replace("  ", " ")
                                           for a in al) or None,
            "Annotation count": len(al), "Has TOP": any(a["layer"] == "TOP" for a in al),
            "Has T&B": any(a["TandB"] for a in al),
            "Polygon (mm, sheet-local)": jtxt([[round(x), round(y)] for x, y in p["polygon_mm"]]),
            "Issues": ", ".join(p["issues"]) or None, "Status": p["status"], "Terminal state": p["terminal_state"],
            "_flags": flags(p), "_src": p})
    M["07_Slabs"] = sr
    M["07_ann"] = regs["SLAB_REBAR_SOURCE_REGISTER"]["rows"]
    M["08_Special_Elements"] = [dict(_flags=flags(r), _src=r, **{
        "Special ID": r["special_id"], "Kind": r["kind"], "Floor": r.get("floor"),
        "Rule": r.get("rule"), "Rule status": r.get("rule_status"), "Status": r.get("status"),
        "Terminal state": r["terminal_state"],
        "Detail": jtxt({k: v for k, v in r.items() if k not in ("special_id", "kind", "floor", "rule", "rule_status",
                                                                  "status", "terminal_state", "source_id")})})
        for r in regs["SPECIAL_STRUCTURAL_OCCURRENCE_REGISTER"]["rows"]]
    M["08_dome"] = regs["SPECIAL_STRUCTURAL_OCCURRENCE_REGISTER"]["dome_detail_items"]
    M["09_Pool"] = [dict(_flags=flags(r), _src=r, **{
        "Pool item ID": r["pool_item_id"], "Component": r["component"], "Role": r["role"], "Raw text": bar(r["raw"]),
        "Count": (r["parsed"] or {}).get("count"), "Dia (mm)": (r["parsed"] or {}).get("dia_mm"),
        "Count mode": (r["parsed"] or {}).get("count_mode"), "Spacing (mm)": (r["parsed"] or {}).get("spacing_mm"),
        "DET position (mm)": jtxt(r["position_mm"]), "Mapping channel": r["mapping_channel"], "Status": r["status"],
        "Terminal state": r["terminal_state"]}) for r in regs["POOL_STRUCTURAL_REGISTER"]["rows"]]
    M["09_summary"] = regs["POOL_STRUCTURAL_REGISTER"]["pool"]
    q = [{"Rank": r["rank"], "Area": r["area"], "Item": r["item"], "Question": r["question"],
          "Evidence": jtxt(r["evidence"]), "Impact": r["impact"], "Mohammad answer": None, "_flags": [], "_src": r}
         for r in regs["STRUCTURAL_REVIEW_QUEUE"]["rows"]]
    for i, t in enumerate(regs["COLUMN_TIE_TOPOLOGY_RULES"]["open_questions"], 1):
        q.append({"Rank": f"T{i}", "Area": "COLUMN_TIES", "Item": "stirrup rule (p.9)", "Question": t,
                  "Evidence": "COLUMN_TIE_TOPOLOGY_RULES.json", "Impact": "column ties", "Mohammad answer": None,
                  "_flags": [], "_src": {}})
    M["10_Review_Questions"] = q
    return M


# ============================================================================================ workbook
def write_xlsx(idx, regs, M):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    F = "Arial"
    hfont = Font(name=F, bold=True, color="FFFFFF", size=10)
    hfill = PatternFill("solid", fgColor="1F3864")
    body = Font(name=F, size=9)
    bold = Font(name=F, size=10, bold=True)
    title = Font(name=F, size=14, bold=True)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ffill = {k: PatternFill("solid", fgColor=v.lstrip("#")) for k, v in FLAG_COLOUR.items()}
    yellow = PatternFill("solid", fgColor="FFFF00")
    wb = Workbook()
    wb.remove(wb.active)
    pos = {}  # sheet -> {header: column letter}, plus data row range

    def table(ws, rows, start_row=1, widths=None, autofilter=True, flagcol=True, fixed=None):
        heads = [h for h in (fixed or (rows[0].keys() if rows else [])) if not h.startswith("_")]
        if flagcol:
            heads = heads + ["Review flags"]
        for j, h in enumerate(heads, 1):
            c = ws.cell(start_row, j, h)
            c.font, c.fill, c.border = hfont, hfill, box
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, r in enumerate(rows, start_row + 1):
            for j, h in enumerate(heads, 1):
                v = (", ".join(r["_flags"]) or None) if h == "Review flags" else r.get(h)
                c = ws.cell(i, j, v)
                c.font, c.border = body, box
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            if flagcol and r.get("_flags"):
                ws.cell(i, len(heads)).fill = ffill[r["_flags"][0]]
                ws.cell(i, len(heads)).font = Font(name=F, size=9, bold=True, color="FFFFFF")
        for j, h in enumerate(heads, 1):
            w = max([len(str(h))] + [min(len(str(r.get(h) or "")), 60) for r in rows[:400]]) + 2
            ws.column_dimensions[L(j)].width = min(max(w, 8), 62)
        if autofilter and rows:
            ws.auto_filter.ref = f"A{start_row}:{L(len(heads))}{start_row + len(rows)}"
        return {h: L(j) for j, h in enumerate(heads, 1)}, start_row + 1, start_row + len(rows)

    order = ["01_Columns", "02_Column_Chains", "03_Footings", "04_Simple_Beams", "05_Continuous_Beams",
             "06_Ground_Strap_Beams", "07_Slabs", "08_Special_Elements", "09_Pool", "10_Review_Questions"]
    summary = wb.create_sheet("00_Summary")
    for name in order:
        ws = wb.create_sheet(name)
        cols, r0, r1 = table(ws, M[name])
        pos[name] = (cols, r0, r1)
        ws.freeze_panes = "B2"
        nxt = r1 + 3
        if name == "03_Footings":
            ws.cell(nxt, 1, "Footing schedule definitions (schedule rows are definitions, not occurrences)").font = bold
            drows = [{"Type": d["footing_type"], "Element": d["element"], "L x W x D (cm)": f"{d['L_cm']:g} x {d['W_cm']:g} x {d['D_cm']:g}",
                      "Components": "; ".join(f"{c['component']} {c['layer']} {c['count']}Ø{c['dia_mm']} ({c['count_mode']})"
                                              for c in d["components"]),
                      "BOXED raw": d["boxed"].get("raw_boxed_value") or None, "Plan occurrences": d["plan_occurrences"],
                      "State": d["state"], "Remarks": "; ".join(d["remarks"]) or None, "_flags": []}
                     for d in M["03_defs"]]
            table(ws, drows, nxt + 1, autofilter=False, flagcol=False)
            nxt += len(drows) + 4
            ws.cell(nxt, 1, "Footing / column mismatches").font = bold
            table(ws, [{"Kind": m["kind"], "Detail": jtxt({k: v for k, v in m.items() if k != "kind"}), "_flags": []}
                       for m in M["03_mismatch"]], nxt + 1, autofilter=False, flagcol=False)
        if name == "05_Continuous_Beams":
            ws.cell(nxt, 1, "CB groups vs schedule (never forced)").font = bold
            grows = [{"CB group ID": c["cb_id"], "Type": c["type"], "Floor": c["floor"],
                      "Schedule spans (m)": jtxt(c["schedule_spans_m"]),
                      "Plan spans c/c (m)": jtxt([s["cc_m"] for s in c["plan_spans"]]),
                      "Plan spans clear (m)": jtxt([s["clear_m"] for s in c["plan_spans"]]),
                      "Schedule B x H (cm)": f"{c['schedule_B_cm']:g} x {c['schedule_H_cm']:g}",
                      "Drawn width (mm)": jtxt(c["drawn_width_mm"]), "Match state": c["match_state"],
                      "Extension check": jtxt(c["extension_check"]), "Terminal state": c["terminal_state"],
                      "_flags": ["SOURCE_CONFLICT"] if c["match_state"] != "MATCH_CONFIRMED" else []} for c in M["05_groups"]]
            table(ws, grows, nxt + 1, autofilter=False)
        if name == "07_Slabs":
            ws.cell(nxt, 1, "Every slab bar text (one row per annotation; TOP / BOTTOM / T&B as printed)").font = bold
            arows = [{"Annotation ID": a["source_row_id"], "Panel ID": a["panel"], "Floor": a["floor"],
                      "Raw": bar(a["raw"]), "Owner": a["owner"], "Binding": a["binding"], "Direction": a["direction"],
                      "Layer": a["layer"], "T&B": a["TandB"], "Count": a["count"], "Dia (mm)": a["dia_mm"],
                      "Count mode": a["count_mode"], "Spacing (mm)": a.get("spacing_mm"),
                      "Support beam line": a.get("support_beam_line"), "Corner bars": a.get("corner_reinforcement"),
                      "Text handle": a["text_handle"], "Note": a.get("note"), "Terminal state": a["terminal_state"],
                      "_flags": flags(a)} for a in M["07_ann"]]
            table(ws, arows, nxt + 1, autofilter=False)
        if name == "08_Special_Elements":
            ws.cell(nxt, 1, "Dome detail bar texts (p.7)").font = bold
            table(ws, [{"Dome item ID": d["dome_item_id"], "Component": d["component"], "Role": d["role"],
                        "Raw": bar(d["raw"]), "Terminal state": d["terminal_state"], "_flags": []} for d in M["08_dome"]],
                  nxt + 1, autofilter=False, flagcol=False)
        if name == "09_Pool":
            ws.cell(nxt, 1, "Pool occurrence and geometry").font = bold
            for i, (k, v) in enumerate(M["09_summary"].items(), nxt + 1):
                ws.cell(i, 1, k).font = bold
                ws.cell(i, 2, jtxt(v)).font = body
        if name == "10_Review_Questions":
            a = cols["Mohammad answer"]
            for i in range(r0, r1 + 1):
                ws[f"{a}{i}"].fill = yellow
            ws.column_dimensions[a].width = 50
            ws.column_dimensions[cols["Question"]].width = 80

    # ---- 00 summary (formulas over the data sheets; register values beside them as a check)
    s = summary
    s["A1"] = "ALSENAN STRUCTURAL CENSUS S1 - HUMAN REVIEW WORKBOOK"
    s["A1"].font = title
    s["A2"] = "Occurrence census only. NO REBAR KG, NO BAR LENGTHS. Plan decides occurrence; schedule only defines."
    s["A3"] = (f"Source: {regs['STRUCTURAL_PROJECT_RULE_REGISTER']['source']['drawing']} sha256 "
               f"{regs['STRUCTURAL_PROJECT_RULE_REGISTER']['source']['dxf_sha256']}; frozen registers: INDEX.json "
               f"(built twice identical: {idx['built_twice_identical']}, frozen before benchmark: "
               f"{idx['frozen_before_benchmark']}). Each data sheet = rows of one frozen register; IDs match the visual "
               f"review PDF labels.")
    s["A4"] = "How to review: filter each sheet by 'Review flags'; write answers only in the yellow column of " \
              "10_Review_Questions."
    for c in ("A2", "A3", "A4"):
        s[c].font = body
    s["A6"] = "Review flag legend"
    s["A6"].font = bold
    for i, (k, v) in enumerate(FLAG_COLOUR.items(), 7):
        s.cell(i, 1, k).fill = ffill[k]
        s.cell(i, 1).font = Font(name=F, size=9, bold=True, color="FFFFFF")
        s.cell(i, 2, {"SOURCE_CONFLICT": "two printed sources disagree (tags, sizes, schedule keys, spans)",
                      "BLOCKED": "counted, but the method / geometry / rule is not available - no value assumed",
                      "AMBIGUOUS_MATCH": "binding or band not unique (tag between two members, L on a band limit, "
                                         "GB length basis, chain gap)",
                      "MISSING_SCHEDULE_DEFINITION": "occurrence with no schedule row / no tag / no bar text"}[k]).font = body
    row = 12
    s.cell(row, 1, "Counts (formulas over the data sheets)").font = bold
    s.cell(row, 2, "Formula").font = bold
    s.cell(row, 3, "Register value").font = bold
    s.cell(row, 4, "Register source").font = bold
    row += 1

    def cnt(label, sheet, crit=None, reg=None, regsrc=None):
        nonlocal row
        cols, r0, r1 = pos[sheet]
        if crit is None:
            f = f"=COUNTA('{sheet}'!A{r0}:A{r1})"
        else:
            parts = ",".join(f"'{sheet}'!{cols[h]}{r0}:{cols[h]}{r1},\"{v}\"" for h, v in crit)
            f = f"=COUNTIFS({parts})"
        s.cell(row, 1, label).font = body
        s.cell(row, 2, f).font = body
        if reg is not None:
            s.cell(row, 3, reg).font = Font(name=F, size=9, color="0000FF")
            s.cell(row, 4, regsrc).font = body
            s.cell(row, 5, f"=IF(B{row}=C{row},\"OK\",\"CHECK\")").font = body
        row += 1

    co = regs["COLUMN_OCCURRENCE_REGISTER"]["summary"]
    cnt("Column occurrences", "01_Columns", reg=co["occurrences"], regsrc="COLUMN_OCCURRENCE_REGISTER.summary.occurrences")
    cnt("Column positions (chains)", "02_Column_Chains", reg=co["physical_positions"],
        regsrc="COLUMN_OCCURRENCE_REGISTER.summary.physical_positions")
    cnt("Footing occurrences", "03_Footings", reg=len(regs["FOOTING_OCCURRENCE_REGISTER"]["rows"]),
        regsrc="FOOTING_OCCURRENCE_REGISTER rows")
    for sh in ("04_Simple_Beams", "05_Continuous_Beams", "06_Ground_Strap_Beams", "07_Slabs", "08_Special_Elements",
               "09_Pool"):
        cnt(f"{sh} rows", sh)
    row += 1
    s.cell(row, 1, "Flagged rows per sheet").font = bold
    for j, k in enumerate(FLAGS, 2):
        s.cell(row, j, k).font = bold
    row += 1
    for sh in order[:-1]:
        cols, r0, r1 = pos[sh]
        s.cell(row, 1, sh).font = body
        for j, k in enumerate(FLAGS, 2):
            c = cols["Review flags"]
            s.cell(row, j, f"=COUNTIF('{sh}'!{c}{r0}:{c}{r1},\"*{k}*\")").font = body
        row += 1
    # column type x floor matrix
    row += 1
    s.cell(row, 1, "Column type x floor (COUNTIFS over 01_Columns)").font = bold
    floors = ["FOUNDATION", "GF", "1F", "2F"]
    for j, f in enumerate(floors + ["Total", "Register total", "Check"], 2):
        s.cell(row, j, f).font = bold
    row += 1
    cols, r0, r1 = pos["01_Columns"]
    top = row
    for mr in regs["COLUMN_TYPE_FLOOR_MATRIX"]["rows"]:
        t = mr["column_type"]
        s.cell(row, 1, t).font = body
        for j, f in enumerate(floors, 2):
            s.cell(row, j, f"=COUNTIFS('01_Columns'!{cols['Type']}{r0}:{cols['Type']}{r1},\"{t}\","
                           f"'01_Columns'!{cols['Floor']}{r0}:{cols['Floor']}{r1},\"{f}\")").font = body
        s.cell(row, 6, f"=SUM(B{row}:E{row})").font = body
        s.cell(row, 7, mr["total"]).font = Font(name=F, size=9, color="0000FF")
        s.cell(row, 8, f"=IF(F{row}=G{row},\"OK\",\"CHECK\")").font = body
        row += 1
    s.cell(row, 1, "Total").font = bold
    for j in range(2, 8):
        cl = L(j)
        s.cell(row, j, f"=SUM({cl}{top}:{cl}{row - 1})").font = bold
    s.cell(row, 8, f"=IF(F{row}=G{row},\"OK\",\"CHECK\")").font = bold
    row += 2
    # beam type x floor
    s.cell(row, 1, "Beam occurrences by type x floor (COUNTIFS over sheets 04 + 05 + 06)").font = bold
    bfl = ["FOUNDATION", "GROUND_BEAMS", "GF_ROOF", "1F_ROOF", "2F_ROOF"]
    for j, f in enumerate(bfl + ["Total", "Register total", "Check"], 2):
        s.cell(row, j, f).font = bold
    row += 1
    top = row
    for mr in regs["BEAM_TYPE_FLOOR_MATRIX"]["rows"]:
        t = mr["beam_type"]
        if ":" in t:
            fam, _, typ = t.partition(":")
        elif t in ("OTHER", "DOME_RING"):  # matrix rows keyed by family (members with no tag)
            fam, typ = t, None
        else:
            fam, typ = None, t
        s.cell(row, 1, t).font = body
        for j, f in enumerate(bfl, 2):
            parts = []
            for sh in ("04_Simple_Beams", "05_Continuous_Beams", "06_Ground_Strap_Beams"):
                c, a0, a1 = pos[sh]
                if fam and typ is None:
                    crit = [(c["Family"], fam)]
                elif fam:
                    crit = [(c["Family"], fam)] + ([(c["Type"], typ)] if typ != "CURVED_UNCATEGORISED" else
                                                   [(c["Type"], "")])
                else:
                    crit = [(c["Type"], typ)]
                crit.append((c["Floor"], f))
                parts.append("COUNTIFS(" + ",".join(f"'{sh}'!{cc}{a0}:{cc}{a1},\"{v}\"" for cc, v in crit) + ")")
            s.cell(row, j, "=" + "+".join(parts)).font = body
        s.cell(row, 7, f"=SUM(B{row}:F{row})").font = body
        s.cell(row, 8, mr["total"]).font = Font(name=F, size=9, color="0000FF")
        s.cell(row, 9, f"=IF(G{row}=H{row},\"OK\",\"CHECK\")").font = body
        row += 1
    s.cell(row, 1, "Total").font = bold
    for j in range(2, 9):
        cl = L(j)
        s.cell(row, j, f"=SUM({cl}{top}:{cl}{row - 1})").font = bold
    s.cell(row, 9, f"=IF(G{row}=H{row},\"OK\",\"CHECK\")").font = bold
    s.column_dimensions["A"].width = 44
    for c in "BCDEFGHI":
        s.column_dimensions[c].width = 16
    s.cell(row + 2, 1, "Blue numbers are copied from the frozen register JSON (the formulas recount the rows); "
                       "every Check must read OK.").font = body
    wb.save(XLSX)
    return {k: v for k, v in pos.items()}


# ============================================================================================ drawings
def verified_build():
    import build_census_s1 as BC
    import alsenan_structural_s1 as S
    idx = json.loads((S1 / "INDEX.json").read_text())
    texts = BC.serialise(BC.build_all())
    for n, v in idx["registers"].items():
        if hashlib.sha256(texts[n].encode("utf-8")).hexdigest() != v["sha256"]:
            raise SystemExit(f"rebuild differs from frozen register {n} - drawings refused")
    B = S.build()
    return S, B


def write_pdf(regs, M):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.lines import Line2D
    from matplotlib.patches import Polygon as MPoly, Rectangle, Arc
    from shapely.geometry import Polygon as SPoly

    S, B = verified_build()
    src = B["src"]
    ents = src.entities()
    labelled = defaultdict(set)

    # one building extent for every page (all sheets share the local frame) -> same scale on every page
    xs, ys = [], []
    for sh in ("FP", "CAP", "GBP", "GFRS", "FFRS", "SFRS"):
        for e in ents[sh]:
            if e["type"] == "LWPOLYLINE" and e["layer"] in ("S-COL.BON", "S-FOOTINGS"):
                xs += [q[0] for q in e["pts"]]
                ys += [q[1] for q in e["pts"]]
        for l in B["beams"][sh]["lines"] if sh in B["beams"] else []:
            for t in (l["t0"], l["t1"]):
                xs.append(t * l["dir"][0] + l["offset"] * l["normal"][0])
                ys.append(t * l["dir"][1] + l["offset"] * l["normal"][1])
    EXT = (min(xs) - 1800, min(ys) - 1800, max(xs) + 2600, max(ys) + 1800)

    def frame(sh):
        return EXT

    def backdrop(ax, sh, grid=False):
        segs = []
        for e in ents[sh]:
            t = e["type"]
            if t == "LINE":
                segs.append((e["a"], e["b"]))
            elif t == "LWPOLYLINE":
                for sg in e.get("segs", []):
                    if sg[0] == "ARC":
                        _, cxy, r, a0, a1 = sg
                        ax.add_patch(Arc(cxy, 2 * r, 2 * r, theta1=math.degrees(a0), theta2=math.degrees(a1),
                                         lw=0.12, color="#B0B0B0", zorder=1))
                    else:
                        segs.append(sg)
            elif t in ("ARC", "CIRCLE"):
                ax.add_patch(Arc(e["c"], 2 * e["r"], 2 * e["r"], theta1=math.degrees(e["a0"]),
                                 theta2=math.degrees(e["a1"]), lw=0.12, color="#B0B0B0", zorder=1))
        from matplotlib.collections import LineCollection
        ax.add_collection(LineCollection(segs, linewidths=0.12, colors="#B0B0B0", zorder=1))
        if grid:
            for a in B["C"]["axes"]["X"]:
                ax.plot([a["pos"], a["pos"]], a["span"], lw=0.25, color="#8FB3D9", ls=(0, (8, 4)), zorder=1)
                ax.text(a["pos"], a["span"][1] + 250, a["axis_id"], fontsize=9, ha="center", color="#2E5E8E")
            for a in B["C"]["axes"]["Y"]:
                ax.plot(a["span"], [a["pos"], a["pos"]], lw=0.25, color="#8FB3D9", ls=(0, (8, 4)), zorder=1)
                ax.text(a["span"][0] - 250, a["pos"], a["axis_id"], fontsize=9, ha="right", va="center", color="#2E5E8E")

    def page(title, sub, sh, grid=False):
        fig = plt.figure(figsize=(46.8, 33.1))  # A0 landscape (vector - zoom freely)
        ax = fig.add_axes([0.02, 0.03, 0.96, 0.89])
        x0, y0, x1, y1 = frame(sh)
        ax.set_xlim(x0, x1)
        ax.set_ylim(y0, y1)
        ax.set_aspect("equal")
        ax.axis("off")
        backdrop(ax, sh, grid)
        fig.text(0.02, 0.975, title, fontsize=30, weight="bold", va="top")
        fig.text(0.02, 0.952, sub, fontsize=14, va="top", color="#333333")
        fig.text(0.02, 0.935, "S1 census - NO REBAR KG. Labels are the frozen census IDs (identical to the review "
                              "workbook). Coordinates: sheet-local mm.", fontsize=11, va="top", color="#555555")
        h = [Line2D([], [], marker=FLAG_MARK[k], ls="", markersize=16, markerfacecolor="none",
                    markeredgecolor=FLAG_COLOUR[k], markeredgewidth=2.5, label=k) for k in FLAGS]
        fig.legend(handles=h, loc="upper right", fontsize=13, ncol=2, frameon=True, bbox_to_anchor=(0.985, 0.985))
        return fig, ax

    def mark(ax, x, y, fl, size=26):
        for i, k in enumerate(fl):
            ax.plot([x], [y], marker=FLAG_MARK[k], markersize=size + 7 * i, markerfacecolor="none",
                    markeredgecolor=FLAG_COLOUR[k], markeredgewidth=1.6, zorder=6)

    def flagtxt(fl):
        return (" [" + ", ".join(FLAG_SHORT[k] for k in fl) + "]") if fl else ""

    pdf = PdfPages(PDF)
    # ---------------------------------------------------------------- cover
    fig = plt.figure(figsize=(46.8, 33.1))
    fig.text(0.05, 0.9, "ALSENAN STRUCTURAL CENSUS S1 - VISUAL REVIEW", fontsize=48, weight="bold")
    lines = ["Generated from the frozen S1 census registers (INDEX.json hashes verified; the census rebuild was proved "
             "byte-identical before drawing).",
             "Every label is a census ID that also appears in ALSENAN_STRUCTURAL_CENSUS_REVIEW.xlsx. Nothing was "
             "edited to tidy the drawings.", "NO REBAR KG and NO BAR LENGTHS are calculated in this round.", "",
             "Pages: 2-5 columns by storey | 6 footings | 7 ground beams + straps | 8-10 roof-slab beams | "
             "11-14 slab panels",
             "Beam labels: occurrence ID | type | cc = support centre-to-centre length, cl = clear length (m).",
             "Slab labels: panel ID | area m2 | thickness mm (authority).", "",
             "Review markers (outline drawn around the labelled object):"]
    for i, t in enumerate(lines):
        fig.text(0.05, 0.84 - i * 0.03, t, fontsize=20)
    for i, k in enumerate(FLAGS):
        y = 0.84 - (len(lines) + i) * 0.03
        fig.lines.append(Line2D([0.065], [y + 0.006], marker=FLAG_MARK[k], markersize=26, markerfacecolor="none",
                                markeredgecolor=FLAG_COLOUR[k], markeredgewidth=3, transform=fig.transFigure))
        fig.text(0.085, y, f"{k}", fontsize=20, color=FLAG_COLOUR[k], weight="bold")
    pdf.savefig(fig)
    plt.close(fig)
    # ---------------------------------------------------------------- columns by storey
    colpts = {f"{sh}:{e['handle']}": e["pts"] for sh in ents for e in ents[sh]
              if e["type"] == "LWPOLYLINE" and e["layer"] == "S-COL.BON"}
    sheet_for = {"FOUNDATION": "CAP", "GF": "GFRS", "1F": "FFRS", "2F": "SFRS"}
    for st, sh in sheet_for.items():
        rows = [r for r in M["01_Columns"] if r["Floor"] == st]
        nfl = sum(1 for r in rows if r["_flags"])
        fig, ax = page(f"COLUMNS - {st} storey ({len(rows)} occurrences, {nfl} flagged)",
                       f"Backdrop: {src.sheets[sh]['title']} (PDF p.{src.sheets[sh]['pdf_page']}). Label: column ID / "
                       f"type / schedule section / tie band.", sh, grid=(sh == "CAP"))
        for r in rows:
            o = r["_src"]
            pairs = list(zip(o["plan_source"]["sheets"], o["plan_source"]["handles"]))
            oid = next((f"{s}:{h}" for s, h in pairs if s == sh), None) or (f"{pairs[0][0]}:{pairs[0][1]}" if pairs else None)
            pts = colpts.get(oid)
            if pts:
                ax.add_patch(MPoly(pts, closed=True, fc="#1F3864", ec="#1F3864", lw=0.4, alpha=0.85, zorder=4))
            x, y = o["plan_centre_mm"]
            fl = r["_flags"]
            colr = FLAG_COLOUR[fl[0]] if fl else "#1F3864"
            ax.text(x + 320, y + 160, f"{r['Column ID']}\n{r['Type']}  sched {r['Schedule B x D (cm)'] or '-'}  "
                                      f"{r['Tie band'] or ''}{flagtxt(fl)}",
                    fontsize=6.5, color=colr, zorder=7, va="bottom",
                    bbox=dict(fc="white", ec=colr, lw=0.3, alpha=0.85, pad=0.6))
            if fl:
                mark(ax, x, y, fl)
            labelled["01_Columns"].add(r["Column ID"])
        pdf.savefig(fig)
        plt.close(fig)
    # ---------------------------------------------------------------- footings
    rows = M["03_Footings"]
    fig, ax = page(f"FOOTINGS ({len(rows)} occurrences, {sum(1 for r in rows if r['_flags'])} flagged)",
                   "Backdrop: FOUNDATION PLAN (PDF p.2). Label: footing ID / type / drawn size / supported column IDs.",
                   "FP")
    for r in rows:
        o = r["_src"]
        bb = (o["outline"] or {}).get("bbox")
        fl = r["_flags"]
        colr = FLAG_COLOUR[fl[0]] if fl else "#2C7A2C"
        if bb:
            ax.add_patch(Rectangle((bb[0], bb[1]), bb[2] - bb[0], bb[3] - bb[1], fc=colr, alpha=0.10, ec=colr, lw=1.0, zorder=3))
            x, y = bb[0] + 60, bb[3] - 60
        else:
            x, y = o["tag"]["position_mm"]
        ax.text(x, y, f"{r['Footing ID']}\ntype {r['Type'] or 'NONE (' + (r['Candidate types'] or '') + ')'}  "
                      f"{r['Drawn L x W (cm)'] or ''}\ncols: {r['Supported columns']}{flagtxt(fl)}",
                fontsize=6.0, color=colr, va="top", zorder=7, bbox=dict(fc="white", ec=colr, lw=0.3, alpha=0.85, pad=0.6))
        if fl and bb:
            mark(ax, (bb[0] + bb[2]) / 2, (bb[1] + bb[3]) / 2, fl, size=34)
        labelled["03_Footings"].add(r["Footing ID"])
    for r in M["01_Columns"]:
        if r["Floor"] == "FOUNDATION":
            x, y = r["_src"]["plan_centre_mm"]
            ax.plot([x], [y], marker="s", markersize=3, color="#1F3864", zorder=5)
    pdf.savefig(fig)
    plt.close(fig)

    # ---------------------------------------------------------------- beams
    lines_by = {sh: {l["line_id"]: l for l in B["beams"][sh]["lines"] + B["beams"][sh].get("fragments", [])}
                for sh in ("GBP", "GFRS", "FFRS", "SFRS")}
    arcs_by = {sh: {a["band_id"]: a for a in B["beams"][sh]["arc_bands"]} for sh in ("GBP", "GFRS", "FFRS", "SFRS")}
    fam_col = {"SIMPLE": "#1F77B4", "CB": "#2CA02C", "STAIR": "#8C564B", "CANTILEVER": "#E377C2", "OTHER": "#7F7F7F",
               "DOME_RING": "#BCBD22", "GB": "#1F77B4", "EXTERIOR_GB": "#2CA02C", "STRAP": "#8C564B"}

    def pt(l, t):
        return (t * l["dir"][0] + l["offset"] * l["normal"][0], t * l["dir"][1] + l["offset"] * l["normal"][1])

    def draw_beam(ax, sh, r, member, s0, s1, txt, fl, fam):
        colr = FLAG_COLOUR[fl[0]] if fl else fam_col.get(fam, "#1F77B4")
        if member and member.startswith("BA") and member in arcs_by[sh]:
            a = arcs_by[sh][member]
            ax.add_patch(Arc(a["centre"], 2 * a["r_mid_mm"], 2 * a["r_mid_mm"], theta1=a["start_deg"],
                             theta2=a["start_deg"] + a["sweep_deg"], lw=2.2, color=colr, alpha=0.75, zorder=4))
            th = math.radians(a["start_deg"] + a["sweep_deg"] / 2)
            mx, my = a["centre"][0] + a["r_mid_mm"] * math.cos(th), a["centre"][1] + a["r_mid_mm"] * math.sin(th)
            ang = 0
        elif member and member in lines_by[sh] and s0 is not None and s1 is not None:
            l = lines_by[sh][member]
            p0, p1 = pt(l, s0), pt(l, s1)
            ax.plot([p0[0], p1[0]], [p0[1], p1[1]], lw=2.2, color=colr, alpha=0.75, solid_capstyle="butt", zorder=4)
            ax.plot([p0[0], p1[0]], [p0[1], p1[1]], ls="", marker="|", markersize=5, color=colr, zorder=5)
            mx, my = (p0[0] + p1[0]) / 2, (p0[1] + p1[1]) / 2
            ang = math.degrees(math.atan2(p1[1] - p0[1], p1[0] - p0[0]))
            ang = ang - 180 if ang > 90 else (ang + 180 if ang <= -90 else ang)
        else:
            tp = r.get("tag_position_mm")
            if not tp:
                return False
            mx, my, ang = tp[0], tp[1], 0
            for c in r.get("member_candidates") or []:
                l = lines_by[sh].get(c)
                if l:
                    p0, p1 = pt(l, l["t0"]), pt(l, l["t1"])
                    ax.plot([p0[0], p1[0]], [p0[1], p1[1]], lw=1.2, ls=(0, (3, 2)), color=colr, zorder=4)
        ax.text(mx, my, txt, fontsize=4.6, rotation=ang, rotation_mode="anchor", ha="center", va="center",
                color=colr, zorder=7, bbox=dict(fc="white", ec=colr, lw=0.25, alpha=0.85, pad=0.5))
        if fl:
            mark(ax, mx, my, fl, size=20)
        return True

    def L_(v):
        return "-" if v is None else f"{v:.2f}"

    gbp_spans = {}
    for m, spans in B["beams"]["GBP"]["spans"].items():
        for spn in spans:
            gbp_spans[(m, round(spn["s0"]))] = spn
    strap_band = {s["band"]["band_id"]: s["band"] for s in B["straps"]}
    # ground beams + straps page
    rows = [r for r in M["06_Ground_Strap_Beams"] if r["Family"] in ("GB", "EXTERIOR_GB", "STRAP")]
    fig, ax = page(f"GROUND BEAMS + STRAP BEAMS ({len(rows)} occurrences, {sum(1 for r in rows if r['_flags'])} flagged)",
                   "Backdrop: GROUND BEAMS PLAN (PDF p.3); straps from the FOUNDATION PLAN (same frame). Label: ID / "
                   "category or type / cc / cl (m).", "GBP")
    for r in rows:
        o = r["_src"]
        fl = r["_flags"]
        if r["Family"] == "STRAP":
            b = strap_band[o["beam_id"].rsplit("-", 1)[1]]
            colr = FLAG_COLOUR[fl[0]] if fl else fam_col["STRAP"]
            ax.plot([b["centre_a"][0], b["centre_b"][0]], [b["centre_a"][1], b["centre_b"][1]], lw=4, color=colr,
                    alpha=0.6, zorder=4)
            mx, my = (b["centre_a"][0] + b["centre_b"][0]) / 2, (b["centre_a"][1] + b["centre_b"][1]) / 2
            ax.text(mx + 300, my, f"{r['Beam occurrence ID']}\n{r['Type']} drawn {r['Drawn width (mm)']:.0f} mm  "
                                  f"L {L_(r['Drawn length (m)'])}{flagtxt(fl)}", fontsize=6.0, color=colr, zorder=7,
                    bbox=dict(fc="white", ec=colr, lw=0.3, alpha=0.85, pad=0.5))
            if fl:
                mark(ax, mx, my, fl)
            labelled["06_Ground_Strap_Beams"].add(r["Beam occurrence ID"])
            continue
        gid = o["beam_id"].replace("BM-GROUND_BEAMS-", "")
        parts = gid.split("-")
        member = parts[1]
        s0 = s1 = None
        if len(parts) == 3:
            spn = gbp_spans.get((member, int(parts[2])))
            s0, s1 = (spn["s0"], spn["s1"]) if spn else (None, None)
        txt = f"{r['Beam occurrence ID']}  {r['GB category (c/c)'] or 'CURVED'}  cc {L_(r['Centreline length (m)'])} / " \
              f"cl {L_(r['Clear length (m)'])}{' EXT' if r['Family'] == 'EXTERIOR_GB' else ''}{flagtxt(fl)}"
        if draw_beam(ax, "GBP", o, member, s0, s1, txt, fl, r["Family"]):
            labelled["06_Ground_Strap_Beams"].add(r["Beam occurrence ID"])
    pdf.savefig(fig)
    plt.close(fig)
    for sh, fl_name in (("GFRS", "GF_ROOF"), ("FFRS", "1F_ROOF"), ("SFRS", "2F_ROOF")):
        rows = [(n, r) for n in ("04_Simple_Beams", "05_Continuous_Beams") for r in M[n] if r["Floor"] == fl_name]
        fig, ax = page(f"BEAMS - {src.sheets[sh]['title']} ({len(rows)} occurrences, "
                       f"{sum(1 for _, r in rows if r['_flags'])} flagged)",
                       f"Backdrop: PDF p.{src.sheets[sh]['pdf_page']}. Label: occurrence ID / type / cc / cl (m). "
                       "Dashed = candidate members of an ambiguous tag.", sh)
        for n, r in rows:
            o = r["_src"]
            ss = o.get("span_stations_mm") or [None, None]
            txt = f"{r['Beam occurrence ID']}  {r['Type'] or 'NO TAG'}  cc {L_(r['Centreline length (m)'])} / " \
                  f"cl {L_(r['Clear length (m)'])}{flagtxt(r['_flags'])}"
            if draw_beam(ax, sh, o, o.get("member"), ss[0], ss[1], txt, r["_flags"], r["Family"]):
                labelled[n].add(r["Beam occurrence ID"])
        pdf.savefig(fig)
        plt.close(fig)
    # ---------------------------------------------------------------- slabs
    cls_col = {"SLAB_PANEL": "#4C9BE8", "STAIR_FLIGHT_ZONE": "#C49C94", "STAIR_IN_VOID_ZONE": "#E7A0A0",
               "DOME_ZONE": "#DBDB8D", "OPEN_TO_BELOW": "#DDDDDD", "OUTSIDE_BUILDING_OR_COURT": "#F2F2F2"}
    for sh in ("GFRS", "FFRS", "SFRS", "GBP"):
        rows = [r for r in M["07_Slabs"] if r["_src"]["sheet"] == sh]
        fig, ax = page(f"SLAB PANELS - {src.sheets[sh]['title']} ({len(rows)} faces, "
                       f"{sum(1 for r in rows if r['_flags'])} flagged)",
                       f"Backdrop: PDF p.{src.sheets[sh]['pdf_page']}. Label: panel ID / area m2 / thickness mm "
                       "(authority) / bar texts.", sh)
        for r in rows:
            p = r["_src"]
            fl = r["_flags"]
            c = cls_col.get(p["class"], "#4C9BE8")
            big = p["class"] == "OUTSIDE_BUILDING_OR_COURT"
            ax.add_patch(MPoly(p["polygon_mm"], closed=True, fc=c, alpha=0.05 if big else 0.22,
                               ec=FLAG_COLOUR[fl[0]] if fl else "#1F3864", lw=1.4 if fl else 0.7, zorder=2))
            for hole in p.get("holes_mm") or []:
                ax.add_patch(MPoly(hole, closed=True, fc="white", ec="#1F3864", lw=0.5, zorder=2))
            poly = SPoly(p["polygon_mm"], holes=p.get("holes_mm") or None)
            if not poly.is_valid:
                poly = poly.buffer(0)
            rp = poly.representative_point()
            t = r["Thickness (mm)"]
            txt = f"{r['Panel ID']}\n{r['Area (m2)']:.2f} m2  t={t if t else '-'} ({r['Thickness authority']})"
            if p["class"] != "SLAB_PANEL":
                txt += f"\n{p['class']}"
            if r["Bound annotations"]:
                txt += "\n" + r["Bound annotations"].replace("; ", "\n")
            txt += flagtxt(fl)
            colr = FLAG_COLOUR[fl[0]] if fl else "#1F3864"
            ax.text(rp.x, rp.y, txt, fontsize=4.8 if not big else 8, ha="center", va="center", color=colr, zorder=7,
                    bbox=dict(fc="white", ec=colr, lw=0.25, alpha=0.8, pad=0.5))
            if fl:
                mark(ax, rp.x, rp.y, fl, size=30)
            labelled["07_Slabs"].add(r["Panel ID"])
        pdf.savefig(fig)
        plt.close(fig)
    d = pdf.infodict()
    d["Title"] = "Alsenan structural census S1 - visual review"
    d["Subject"] = "Occurrence census review; no rebar kg"
    pdf.close()
    return labelled


# ============================================================================================ main
def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
    idx, regs = load_frozen()
    M = model(regs)
    write_xlsx(idx, regs, M)
    recalc = os.environ.get("RECALC_SCRIPT")  # LibreOffice recalculation (caches formula values for viewers)
    recalc_state = "NOT_RUN (formulas recalculate when opened in Excel)"
    if recalc and Path(recalc).exists():
        out = subprocess.run([sys.executable, recalc, str(XLSX), "180"], capture_output=True, text=True)
        res = json.loads(out.stdout[out.stdout.index("{"):])
        if res.get("status") != "success" or res.get("total_errors"):
            raise SystemExit(f"workbook recalculation failed: {res}")
        recalc_state = f"RECALCULATED ({res['total_formulas']} formulas, 0 errors)"
    labelled = write_pdf(regs, M)
    xl_ids = {
        "01_Columns": {r["Column ID"] for r in M["01_Columns"]},
        "03_Footings": {r["Footing ID"] for r in M["03_Footings"]},
        "04_Simple_Beams": {r["Beam occurrence ID"] for r in M["04_Simple_Beams"]},
        "05_Continuous_Beams": {r["Beam occurrence ID"] for r in M["05_Continuous_Beams"]},
        "06_Ground_Strap_Beams": {r["Beam occurrence ID"] for r in M["06_Ground_Strap_Beams"]
                                  if r["Family"] in ("GB", "EXTERIOR_GB", "STRAP")},
        "07_Slabs": {r["Panel ID"] for r in M["07_Slabs"]}}
    cross = {k: {"workbook_ids": len(v), "drawing_labels": len(labelled.get(k, set())),
                 "missing_on_drawing": sorted(v - labelled.get(k, set())),
                 "label_not_in_workbook": sorted(labelled.get(k, set()) - v)} for k, v in xl_ids.items()}
    not_drawn = [r["Beam occurrence ID"] for r in M["06_Ground_Strap_Beams"] if r["Family"] in ("LIFT_TIE", "LINTEL")]
    man = {"round": "S1", "purpose": "human review of the frozen census before any reinforcement calculation",
           "no_rebar_kg": True, "registers_consumed": {k: v["sha256"] for k, v in idx["registers"].items()},
           "flags": list(FLAGS), "id_cross_check": cross,
           "not_drawable_by_design": {"rows": not_drawn, "why": "required by rule / population not drawn on any plan "
                                                                "(listed in 06_Ground_Strap_Beams)"},
           "outputs": {XLSX.name: sha(XLSX), PDF.name: sha(PDF)}, "workbook_formulas": recalc_state,
           "flag_counts": {k: dict(Counter(f for r in v for f in r["_flags"])) for k, v in M.items()
                           if isinstance(v, list) and v and isinstance(v[0], dict) and "_flags" in v[0]}}
    MANIFEST.write_text(json.dumps(man, indent=1, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    bad = {k: v for k, v in cross.items() if v["missing_on_drawing"] or v["label_not_in_workbook"]}
    print(json.dumps({k: (v["workbook_ids"], v["drawing_labels"]) for k, v in cross.items()}))
    if bad:
        raise SystemExit(f"workbook / drawing id mismatch: {json.dumps(bad)[:2000]}")


if __name__ == "__main__":
    main()
