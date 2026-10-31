"""T-394 Alsenan: CL-GF, CL-1F and CL-LIFT of T-385 r2 traced to count / section / height basis / ownership / scope
(research only: no quantity, line, state or released cell changes).

Inputs (committed unless marked):
  research/agents_t332/CANDIDATE_V3021_CONCRETE_LINES.json  column lines; asserted equal (ids, element ids, values,
      states) to the V3.0.29 binding research/agents_t387/BINDING_T387.csv of the reproduced candidate of record
  research/agents_t248/LEDGER_T248.csv     member inventory (T-217 re-extracted), datums (T-214), status, T-119 id
  research/agents_t119/COLUMNS_TO_CONTROLLING_MEMBER.csv  owner split D_ctrl per member (owner manual method)
  research/agents_t256/LEDGER_T256.csv     schedule band of the 15 CONFLICT members
  .work/private/t394/OWN_COLUMNS_T394.json  own outline read (t394_extract.py)                      [private]
  the concrete + rebar workbook, sheet 'walls + columns' rows 11, 36 .. 47, 52 .. 61                 [private]
Outputs: .work/private/t394/RECON_T394.csv (per line, with his row) and MARKS_T394.csv (per storey x mark terms);
the engine-side summary research/agents_t394/summary_t394.json (engine figures + cause terms in m3 per storey; no row,\ncount or dimension of the workbook).

    python3 -I research/agents_t394/t394_alsenan.py --freelancer <private>/concrete_rebar.xlsx [--write]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import t394_rules as R  # noqa: E402

RES = HERE.parent
PRIVATE = HERE.parents[2] / "private" / "t394"
SHEET = "الحوائط + الأعمدة"
HIS_ROWS = {"GF": range(36, 48), "1F": range(52, 62)}
LIFT_ROW = 11
HIS_MARK = {"PC1": "P.C 20x50", "PC2": "P.C 20x70"}          # his planted-column rows, matched by his printed section
OFFSET = (75968, 14448)                                        # candidate SH12 frame - route G (T-248) frame, mm
STOREY = {"GF": {"base": 0.95, "top": 5.45}, "1F": {"base": 5.45, "top": 9.65}}   # T-214 / T-248 datums
CB5 = {"GFRS:35F", "GFRS:365"}                                 # T-287 G1: base +0.25 (annex ground slab)
RX = re.compile(r"SH12:(\w+)@(-?\d+),(-?\d+)\|(\w+)")


def load_lines():
    d = json.loads((RES / "agents_t332" / "CANDIDATE_V3021_CONCRETE_LINES.json").read_text(encoding="utf-8"))
    lines = {ln["line_id"]: ln for ln in d["lines"] if ln["item_code"] == "STR.CONC.COLUMN"}
    b = {r["line_id"]: r for r in csv.DictReader((RES / "agents_t387" / "BINDING_T387.csv").open(encoding="utf-8"))
         if r["item_code"] == "STR.CONC.COLUMN"}
    assert set(b) == set(lines), "column line ids differ from the V3.0.29 binding"
    for k, r in b.items():
        ln = lines[k]
        assert r["element_id"] == ln["element_id"]
        if r["status"] == "VALUED":
            assert abs(float(r["valued_m3"]) - ln["measured_net"]) < 1e-6
        else:
            assert ln["measured_net"] in (None, 0) or ln["evidence_state"] == "NOT_APPLICABLE_WITH_EVIDENCE"
    return lines, {k: r["status"] for k, r in b.items()}


def load_members():
    t248 = [r for r in csv.DictReader((RES / "agents_t248" / "LEDGER_T248.csv").open(encoding="utf-8"))
            if r["storey"] in ("GF", "1F")]
    t119 = {r["id"]: r for r in csv.DictReader((RES / "agents_t119" / "COLUMNS_TO_CONTROLLING_MEMBER.csv")
                                                .open(encoding="utf-8"))}
    t256 = {r["member"]: r for r in csv.DictReader((RES / "agents_t256" / "LEDGER_T256.csv").open(encoding="utf-8"))}
    own = {m["member"]: m for m in json.loads((PRIVATE / "OWN_COLUMNS_T394.json").read_text(encoding="utf-8"))["members"]}
    return t248, t119, t256, own


def bind(lines, t248):
    """line -> member: route G lines by mark + storey + position (frame offset, 3 mm); planted lines by label."""
    out = {}
    for lid, ln in lines.items():
        if ln["floor"] not in ("GF", "1F"):
            continue
        eid = ln["element_id"]
        if eid.startswith("HL:PLANTED_COLUMN:PC-1F-"):
            mark = "P.C 20x70" if eid.split(":")[2] == "PC-1F-A" else "P.C 20x50"
            hit = [r for r in t248 if r["storey"] == "1F" and r["mark"] == mark]
        else:
            if ":COL:SH06:PC@" in eid:          # route G planted lines superseded by the HL records (T-220)
                out[lid] = "SUPERSEDED_BY_HL"
                continue
            m = RX.search(eid)
            mark = m.group(1)
            x, y = int(m.group(2)) - OFFSET[0], int(m.group(3)) - OFFSET[1]
            hit = []
            for r in t248:
                n = RX.search(r["route_g_line"] or "")
                if n and n.group(1) == mark and n.group(4) == ln["floor"] and r["storey"] == ln["floor"] \
                        and abs(int(n.group(2)) - x) <= 3 and abs(int(n.group(3)) - y) <= 3:
                    hit.append(r)
        assert len(hit) == 1, (lid, eid, len(hit))
        out[lid] = hit[0]["member"]
    return out


def his_rows(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    rows = {}
    for st, rng in HIS_ROWS.items():
        for r in rng:
            mark = str(ws.cell(r, 1).value).strip()
            mark = HIS_MARK.get(mark, mark)
            b, c, d, n = (ws.cell(r, k).value for k in (2, 3, 4, 6))
            rows[(st, mark)] = {"row": r, "n": int(n), "a": round(float(b) * float(c), 6), "h": float(d),
                                "base": STOREY[st]["base"], "m3": float(ws.cell(r, 9).value)}
    lift = {"row": LIFT_ROW, "len": float(ws.cell(LIFT_ROW, 2).value), "t": float(ws.cell(LIFT_ROW, 3).value),
            "h": float(ws.cell(LIFT_ROW, 4).value), "m3": float(ws.cell(LIFT_ROW, 9).value)}
    return rows, lift


def d_ctrl(t119_row):
    """D_ctrl from the owner split: settled -> number; an open controlling depth -> None (split OPEN)."""
    if t119_row is None or t119_row["hold_cause"] == "FRAMING_MEMBER_HELD_D_OPEN":
        return None
    try:
        return float(t119_row["D_ctrl"])
    except ValueError:
        return None


def build(freelancer):
    lines, st387 = load_lines()
    t248, t119, t256, own = load_members()
    binding = bind(lines, t248)
    by_member = {r["member"]: r for r in t248}
    his, lift = his_rows(freelancer)
    per_line, groups = [], defaultdict(list)
    for lid, member in sorted(binding.items(), key=lambda kv: (lines[kv[0]]["floor"], kv[1])):
        ln = lines[lid]
        st = ln["floor"]
        if ln["evidence_state"] == "NOT_APPLICABLE_WITH_EVIDENCE":
            per_line.append({"line_id": lid, "member": member, "storey": st, "state": "NOT_APPLICABLE",
                             "engine_m3": None})
            continue
        r = by_member[member]
        o = own[member] if member in own else None
        held = st387[lid] == "HELD"
        base = 0.25 if member in CB5 else STOREY[st]["base"]
        h = round(STOREY[st]["top"] - base, 6)
        m3 = ln["measured_net"]
        a = round(m3 / h, 6) if m3 else None
        if a is not None:
            assert abs(a * h - m3) < 1e-6
        sec = "CONFLICT" if "SECTION_CONFLICT" in r["causes"] else "DRAWING-CONFIRMED"
        a_drawn = round(o["area_m2"], 6) if o and o.get("found") else None
        t1 = t119.get(r["owner_t119_id"])
        row = {"line_id": lid, "member": member, "storey": st, "mark": r["mark"],
               "state": "HELD" if held else "VALUED", "evidence_state": ln["evidence_state"],
               "engine_m3": None if held else round(m3, 6), "engine_a": a, "engine_h": None if held else h,
               "base": base, "top": STOREY[st]["top"], "a_drawn_own": a_drawn,
               "schedule": t256[member]["schedule_cm"] if member in t256 else "= drawn",
               "section_status": sec, "t248_status": r["status"], "t119_id": r["owner_t119_id"],
               "d_ctrl": d_ctrl(t1), "owner_split_status": t1["status"] + (":" + t1["hold_cause"] if t1 and
                                                                           t1["hold_cause"] else "") if t1 else "",
               "owner_column_m3": None, "owner_joint_head_m3": None,
               "conflict_drawn_minus_engine_m3": None}
        if not held and row["d_ctrl"] is not None:
            row["owner_column_m3"] = round(a * (h - row["d_ctrl"]), 6)
            row["owner_joint_head_m3"] = round(a * row["d_ctrl"], 6)
        if not held and sec == "CONFLICT" and a_drawn is not None:
            row["conflict_drawn_minus_engine_m3"] = round((a_drawn - a) * h, 6)
        if not held:
            assert row["engine_a"] is not None and (sec == "CONFLICT" or abs(a_drawn - a) < 1e-4), (member, a, a_drawn)
        per_line.append(row)
        groups[(st, r["mark"])].append(row)
    marks = []
    fam = {}
    for st in ("GF", "1F"):
        keys = sorted({k[1] for k in groups if k[0] == st} | {k[1] for k in his if k[0] == st})
        tot = defaultdict(float)
        for mk in keys:
            mem = [{"id": m["member"], "state": m["state"], "a": m["engine_a"] or 0.0, "h": m["engine_h"] or 0.0,
                    "base": m["base"], "d_ctrl": m["d_ctrl"], "section_status": m["section_status"],
                    "m3": None if m["state"] == "HELD" else m["engine_m3"]} for m in groups.get((st, mk), [])]
            hr = his.get((st, mk))
            t = R.decompose_mark(mem, hr)
            marks.append({"storey": st, "mark": mk, "his_row": hr["row"] if hr else None, **t})
            for k in R.TERM_KEYS + ("engine_m3", "his_m3", "residual", "n_drawn", "n_his", "n_held", "n_valued"):
                tot[k] += t[k]
        fam[st] = {k: round(v, 6) for k, v in tot.items()}
        fam[st]["closure"] = R.closure([{"term": k, "m3": fam[st][k], "cause": k} for k in R.TERM_KEYS])
        fam[st]["conflict_exposure_m3"] = round(sum(m.get("conflict_drawn_minus_engine_m3") or 0 for m in per_line
                                                    if m["storey"] == st), 6)
        fam[st]["owner_split_open_members"] = sum(1 for m in per_line if m["storey"] == st and m["state"] == "VALUED"
                                                  and m.get("d_ctrl") is None)
    # CL-LIFT (R-T394-6): drawn ring 2.20^2 - 1.80^2 gross, 4 lift column pieces 0.50 m2 (T-204 / T-322), no candidate line
    gross, pieces = round(2.20 ** 2 - 1.80 ** 2, 6), 0.50
    ring = R.ring_scope(gross, pieces, round(lift["len"] * lift["t"], 6), lift["h"], net_h=None)
    lift_terms = {"engine_m3": 0.0, "his_m3": lift["m3"],
                  "DOUBLE_COLUMN_PIECES": -ring["double_on_gross_m3"],          # his ring through the 4 lift columns
                  "HELD_NET_RING": -round(ring["net_m2"] * lift["h"], 6)}       # net band x his height: pit depth HELD
    lift_terms["residual"] = round(lift_terms["engine_m3"] - lift_terms["his_m3"], 6)
    lift_terms["bookkeeping_remainder"] = round(lift_terms["residual"] - lift_terms["DOUBLE_COLUMN_PIECES"]
                                                - lift_terms["HELD_NET_RING"], 6)
    return per_line, marks, fam, {"net_m2": ring["net_m2"], "gross_m2": gross, "column_pieces_m2": pieces,
                                  "his_plan_equals_gross": ring["double_on_gross_m3"] is not None, **lift_terms}


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--freelancer", required=True)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args(argv)
    per_line, marks, fam, lift = build(a.freelancer)
    eng = {st: {"lines": sum(1 for m in per_line if m["storey"] == st),
                "valued": sum(1 for m in per_line if m["storey"] == st and m["state"] == "VALUED"),
                "held": sum(1 for m in per_line if m["storey"] == st and m["state"] == "HELD"),
                "not_applicable": sum(1 for m in per_line if m["storey"] == st and m["state"] == "NOT_APPLICABLE"),
                "engine_valued_m3": round(sum(m["engine_m3"] or 0 for m in per_line if m["storey"] == st), 6)}
           for st in ("GF", "1F")}
    public = {"schema": "T394_SUMMARY_V1", "candidate": "V3.0.29 column lines (= V3.0.21 committed file, asserted)",
              "engine": eng,
              "terms_engine_minus_freelancer_m3": {st: {k: fam[st][k] for k in R.TERM_KEYS + ("residual",)}
                                                   for st in fam},
              "closure": {st: fam[st]["closure"] for st in fam},
              "conflict_exposure_drawn_minus_engine_m3": {st: fam[st]["conflict_exposure_m3"] for st in fam},
              "owner_split_open_valued_members": {st: fam[st]["owner_split_open_members"] for st in fam},
              "count_drawn_vs_his": {st: [fam[st]["n_drawn"], fam[st]["n_his"]] for st in fam},
              "lift": {k: lift[k] for k in ("gross_m2", "column_pieces_m2", "net_m2", "his_plan_equals_gross",
                                            "DOUBLE_COLUMN_PIECES", "HELD_NET_RING", "residual",
                                            "bookkeeping_remainder")},
              "binding": {"lines_bound": len(per_line),
                          "members": sorted({m["member"] for m in per_line if m["member"] != "SUPERSEDED_BY_HL"})},
              "per_line_engine": [{k: m.get(k) for k in ("line_id", "member", "storey", "mark", "state", "engine_m3",
                                                         "engine_h", "base", "top", "section_status", "t119_id",
                                                         "d_ctrl")} for m in per_line]}
    print(json.dumps({k: public[k] for k in ("engine", "terms_engine_minus_freelancer_m3", "closure",
                                             "conflict_exposure_drawn_minus_engine_m3", "lift",
                                             "owner_split_open_valued_members", "count_drawn_vs_his")}, indent=1))
    if a.write:
        PRIVATE.mkdir(parents=True, exist_ok=True)
        for name, rows in (("RECON_T394.csv", per_line), ("MARKS_T394.csv", marks)):
            keys = list(dict.fromkeys(k for r in rows for k in r))
            with (PRIVATE / name).open("w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, keys)
                w.writeheader()
                w.writerows(rows)
        (PRIVATE / "summary_t394.json").write_text(json.dumps(public, indent=1, sort_keys=True) + "\n",
                                                   encoding="utf-8")
        # tracked copy: engine-side figures and cause terms only (no row, count or dimension of the workbook)
        pub = {k: v for k, v in public.items() if k != "count_drawn_vs_his"}
        (HERE / "summary_t394.json").write_text(json.dumps(pub, indent=1, sort_keys=True) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
