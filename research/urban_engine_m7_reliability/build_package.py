"""M7 package builder: the CSV registers and the workbook, from the committed records only (no engine run, no external
record). Deterministic: the same records give the same files.

    python3 -I research/urban_engine_m7_reliability/build_package.py [--check]

Writes 03_ENGINE_WEAKNESS_REGISTER.csv, 06_COMPLETENESS_RECONCILIATION.csv, 07_MEASUREMENT_RELEASE_REGISTER.csv,
10_BEFORE_AFTER_QUANTITIES.csv and M7_RELIABILITY_WORKBOOK.xlsx (formulas for every money figure: the unit rate is one
input cell). --check rebuilds the CSVs in memory and compares them byte for byte.
"""

from __future__ import annotations

import csv
import io
import json
import sys
from collections import defaultdict
from pathlib import Path

PKG = Path(__file__).resolve().parent
ROOT = PKG.parents[1]
RATE_KWD_PER_M3 = 26.5          # the mission's concrete unit rate: a reporting input only, never used by the engine

WEAKNESSES = [
    # id, weakness, found by, rule (V5 convention or intake), ST7757 effect, synthetic evidence, status
    ("M7-W1", "MTEXT direction stored as a vector (text_direction) read as rotation 0", "metamorphic harness: "
     "translate / rotate changed beam mark binding on cluttered villas (22 quantity groups per drawing)",
     "intake (unconditional): MTEXT.get_rotation()", "135 structural and 216 architectural texts (dimension figures) "
     "re-read; no frozen output changed", "rotations and mirror pass on 8/8 development villas", "FIXED"),
    ("M7-W2", "a mirrored TEXT's insert and rotation read in its OCS, not world coordinates", "metamorphic: mirror lost "
     "every sheet", "intake (unconditional): OCS -> WCS", "none (no mirrored text)", "mirror passes", "FIXED"),
    ("M7-W3", "geometry predicates and shoelace areas lose digits far from the origin", "metamorphic: translate by "
     "1.2e6 units changed slab areas (metre drawings) and flipped beam cuts at flush columns (ST7757)",
     "GEOMETRY_LOCAL_ORIGIN", "ids with coordinates are now drawing-local; quantities unchanged (< 0.002 kg)",
     "translate passes on 10/10 held-out villas", "FIXED"),
    ("M7-W4", "exact-zero overlap tests on flush faces; a convex clip against an arc band reads 0",
     "metamorphic on ST7757: rot180 sent a 6.8 m2 ground-slab cell to UNRESOLVED; translate un-cut 3 bands",
     "TOUCH_IS_NOT_OVERLAP", "1F beams unresolved +0.27 m3, GF beams +0.03 m3, 1F slabs -0.08 m3, GF beam bars +3.3 kg",
     "rot180 / mirror pass", "FIXED"),
    ("M7-W5", "a column-sized rectangle on an unrelated layer read as a column (cuts beams, splits cells)",
     "metamorphic: an added furniture layer changed quantities on 5 of 8 villas", "COLUMN_LAYER_EVIDENCE (+ post-freeze "
     "P2: set-aside outlines stay weak evidence)", "two outlines on a generic layer stop cutting GF beams / slabs "
     "(GF slabs +0.04 m3, GF beams unresolved +0.12 m3); still reported (MEDIUM)", "noise_layer 1/10 (V4) -> 10/10 "
     "(V5) held-out", "FIXED"),
    ("M7-W6", "lines of another layer between two treads break the stair going", "metamorphic noise layer",
     "STAIR_TREADS_ONE_LAYER", "none", "noise_layer passes", "FIXED"),
    ("M7-W7", "every polyline rectangle under 1.0 drawing unit dropped (all columns of a metre drawing)",
     "metamorphic units_m on ST7757 (50 groups)", "UNIT_AWARE_RECTANGLES", "none in mm", "units pass 10/10", "FIXED"),
    ("M7-W8", "crossing outlines (composite columns, exploded polylines) return minimal cells, not the drawn "
     "rectangles", "metamorphic explode on ST7757: 4 composite column outlines lost, GF beams +0.30 m3",
     "RECTANGLE_FROM_FOUR_SIDES", "none on the drawing as delivered", "explode passes 10/10 (V4 0/10)", "FIXED"),
    ("M7-W9", "a clipped outline exploded into lines is no longer an outline", "metamorphic explode on ST7757: one "
     "footing fell back to its schedule size (+0.042 m3)", "POLYGON_FROM_LINE_RING", "the ST7757 case remains (a "
     "third line meets a corner): +0.042 m3 on the exploded copy only", "explode passes on synthetic villas",
     "PARTIAL"),
    ("M7-W10", "continuous beam read REVERSED: end anchorage put on the interior support", "metamorphic rotate on "
     "ST7757 (a 5D18 group 73.2 / 66.2 kg)", "CB_ENDS_BY_SHARED_SUPPORT", "1F beam bars +1.45 kg",
     "unit fixture both directions equal", "FIXED"),
    ("M7-W11", "a drawn length on a printed class limit decided by float noise (ground-beam span class, lintel "
     "width band)", "metamorphic units / rotate on ST7757; ST7757 local frame dropped a 1.000 m lintel",
     "LENGTH_CLASS_AT_MILLIMETRE (+ post-freeze P1 for lintels)", "GB low reading stable; the GF lintel OP-002 kept "
     "(0.048 m3, 5.19 kg)", "fixture", "FIXED"),
    ("M7-W12", "a duplicated annotation crashes the run (duplicate member ids)", "metamorphic dup_text",
     "DUPLICATE_ANNOTATION_READ_ONCE", "none", "dup_text 0/10 (V4) -> 10/10", "FIXED"),
    ("M7-W13", "a sheet frame drawn as separate lines is no sheet", "metamorphic explode lost every sheet",
     "SHEET_FRAMES_FROM_LINES", "none", "explode passes", "FIXED"),
    ("M7-W14", "bays merged through a support gap and dropped as one opening (M6 W7 re-diagnosed)",
     "cluttered evaluation; M6 held-out REGION_EXCLUDED faces larger than the design's opening bay",
     "SLAB_JUNCTION_GAP_CLOSURE + SLAB_OPENING_EXTENT_GUARD", "ground slab +0.486 m3 provisional, -0.562 m3 "
     "unresolved (two cells split at a support gap); 1F slabs -0.003 m3", "slab SILENT rows 16 / 18 (V4) -> 0 (V5)",
     "FIXED"),
    ("M7-W15", "a column mark bound to the nearest outline (a pit, a footing) instead of its section",
     "cluttered evaluation (lift pit beside a footing label)", "COLUMN_OUTLINE_BY_SECTION",
     "3 MEMBER_OUTLINE_SUSPECT findings resolved on the foundation plan", "column EXACT 49 -> 55 rows", "FIXED"),
    ("M7-W16", "a column absent from a storey plan but shown above silently dropped", "cluttered evaluation",
     "COLUMN_PRESENCE_CONTRADICTION", "none", "column SILENT rows 31 / 33 (V4) -> 0", "FIXED"),
    ("M7-W17", "a stair flight drawn inside a slab panel counted as slab", "cluttered evaluation",
     "SLAB_STAIR_FOOTPRINT_UNRESOLVED", "none", "stair bays IN_RANGE / REPORTED", "FIXED"),
    ("M6-W5", "legend rows drawn close together tie the caption vote", "M6 held-out V2", "LEGEND_NEAREST_CAPTION",
     "none (no legend on the plans)", "seeds 233 / 234: 1 false finding each -> 0", "FIXED"),
    ("M6-W6", "a legend that is not in the DXF (another sheet, the PDF) cannot name a fill", "M6 ST7757", "(none)",
     "13 of 14 fills UNKNOWN with candidates", "-", "OPEN"),
    ("M7-W18", "item ids carry drawing coordinates: a unit change renames them", "metamorphic units", "(none; "
     "decisions are bound to the source sha256 anyway)", "-", "-", "OPEN"),
    ("M7-W19", "RTEXT entities are neither read by the intake nor transformable by the harness", "metamorphic notes "
     "(10 RTEXT on ST7757)", "(none)", "10 texts not read", "-", "OPEN"),
    ("M7-W20", "residual placement sensitivity on ST7757: mirror -0.002 m3 GF slab, rot180 / mirror -0.13 kg GF beam "
     "bars, explode drift <= 1e-4 relative (arc representation)", "metamorphic on ST7757 (final run)", "(none)",
     "below 0.01 % of any family", "-", "OPEN"),
    ("M7-W21", "a storey whose plan is missing is reported, but carries no volume range", "held-out shortfall "
     "decomposition (124 m3 slab, 54 m3 beams over 5 drawings)", "(none)", "-", "-", "OPEN"),
    ("M7-W22", "untagged beam spans: no beam row is EXACT; 31 m3 of 678 m3 beam truth sits inside unresolved ranges",
     "held-out quantity rows (BEAMS 83/83 REPORTED)", "(convention UNLABELLED_SPAN_ACROSS_SUPPORT keeps them "
     "unresolved)", "1F / GF beam unresolved exposure 7.5 / 5.3 m3", "-", "OPEN (by design)"),
    ("M7-W23", "independent checks do not cover arc bands, arc-bounded and recovered panels (58 of 300 measured "
     "concrete items on ST7757 NOT_COMPARABLE)", "per-item checker", "(none)", "58 items unchecked", "-", "OPEN"),
]


def _csv(rows, fields):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields, lineterminator="\n")
    w.writeheader()
    for r in rows:
        w.writerow({k: (f"{v:.4f}" if isinstance(v, float) else v) for k, v in r.items()})
    return buf.getvalue()


def _rows(path):
    return list(csv.DictReader((PKG / path).open(encoding="utf-8")))


def weakness_register():
    f = ["ID", "WEAKNESS", "FOUND_BY", "RULE", "ST7757_EFFECT", "SYNTHETIC_EVIDENCE", "STATUS"]
    return _csv([dict(zip(f, w)) for w in WEAKNESSES], f)


def completeness_reconciliation():
    H = json.loads((PKG / "heldout" / "HELDOUT_SUMMARY.json").read_text(encoding="utf-8"))["runs"]
    out = []
    for cfg, name in (("C_A", "held-out seeds 301-340, development vocabulary (M6 config A2)"),
                      ("C_B", "held-out seeds 301-340, held-out vocabulary, hard (M6 config B2)")):
        for tag in ("V4", "V5"):
            s = H[f"{cfg}_{tag}"]
            fs = s["family_states"]
            out.append({"SECTION": "SYNTHETIC_HELDOUT", "SET": name, "ENGINE": tag, "INJECTED": s["injected"],
                        "DETECTED": s["detected"], "FALSE_FINDINGS": s["false_findings"],
                        "FALSE_BY_KIND": json.dumps(s["false_by_kind"], sort_keys=True),
                        "DESIGN_VOID_FLAGS": s["design_opening_bay_flags"],
                        "FAMILY_RIGHT": fs.get("RIGHT", 0), "FAMILY_WRONG": fs.get("WRONG", 0),
                        "FAMILY_UNKNOWN": fs.get("UNKNOWN", 0), "FAMILY_NA": fs.get("NA", 0), "DETAIL": ""})
    S = json.loads((PKG / "st7757" / "SUMMARY.json").read_text(encoding="utf-8"))["completeness_findings"]
    for tag in ("V4", "V5"):
        sev = S[f"{tag}_by_severity"]
        out.append({"SECTION": "ST7757_FINDINGS", "SET": "ST7757 (post-discovery)", "ENGINE": tag, "INJECTED": "",
                    "DETECTED": S[tag], "FALSE_FINDINGS": "", "FALSE_BY_KIND": "", "DESIGN_VOID_FLAGS": "",
                    "FAMILY_RIGHT": "", "FAMILY_WRONG": "", "FAMILY_UNKNOWN": "", "FAMILY_NA": "",
                    "DETAIL": f"findings HIGH {sev['HIGH']}, MEDIUM {sev['MEDIUM']}, LOW {sev['LOW']}"})
    for r in _rows("st7757/FINDINGS_DELTA.csv"):
        out.append({"SECTION": "ST7757_FINDING_DELTA", "SET": r["IN"], "ENGINE": "", "INJECTED": "", "DETECTED": "",
                    "FALSE_FINDINGS": "", "FALSE_BY_KIND": r["KIND"], "DESIGN_VOID_FLAGS": "", "FAMILY_RIGHT": "",
                    "FAMILY_WRONG": "", "FAMILY_UNKNOWN": "", "FAMILY_NA": "",
                    "DETAIL": f"{r['SEVERITY']} {r['SHEET']} {r['FAMILY']}: {r['EVIDENCE'][:160]}"})
    return _csv(out, list(out[0]))


def release_register():
    agg = defaultdict(lambda: {"N": 0, "BASE": 0.0, "LOW": 0.0, "HIGH": 0.0, "EXPOSURE_LOW": 0.0, "EXPOSURE_HIGH": 0.0})
    for r in _rows("st7757/RELEASE_REGISTER.csv"):
        a = agg[(r["TRADE"], r["FAMILY"], r["STOREY"], r["STATE"])]
        a["N"] += 1
        for k in ("BASE", "LOW", "HIGH", "EXPOSURE_LOW", "EXPOSURE_HIGH"):
            a[k] += float(r[k]) if r[k] not in ("", "None") else 0.0
    rows = [{"PROJECT": "ST7757", "TRADE": k[0], "FAMILY": k[1], "STOREY": k[2], "STATE": k[3],
             "UNIT": "m3" if k[0] == "CONCRETE" else "kg", **v} for k, v in sorted(agg.items())]
    return _csv(rows, list(rows[0]))


def _basis():
    b = defaultdict(list)
    for r in _rows("st7757/ABLATION.csv"):
        if r["TRADE"]:
            b[(r["TRADE"], r["FAMILY"], r["STOREY"])].append(
                f"{r['CONVENTION']} ({float(r['DELTA_MEASURED_DUE_TO_IT']):+.4f} / unresolved "
                f"{float(r['DELTA_UNRESOLVED_HIGH_DUE_TO_IT']):+.4f})")
    return b


def before_after():
    basis = _basis()
    rows = []
    for r in _rows("st7757/BY_FAMILY.csv"):
        k = (r["TRADE"], r["FAMILY"], r["STOREY"])
        rows.append({"SOURCE": "ST7757 (post-discovery)", "TRADE": r["TRADE"], "FAMILY": r["FAMILY"],
                     "STOREY": r["STOREY"], "UNIT": r["UNIT"], "TRUTH": "", "V4_MEASURED": float(r["V4_MEASURED"]),
                     "V5_MEASURED": float(r["V5_MEASURED"]), "DELTA_MEASURED": float(r["DELTA_MEASURED"]),
                     "V4_UNRESOLVED_HIGH": float(r["V4_UNRESOLVED_HIGH"]),
                     "V5_UNRESOLVED_HIGH": float(r["V5_UNRESOLVED_HIGH"]),
                     "V4_SILENT_M3": "", "V5_SILENT_M3": "", "BASIS": "; ".join(basis.get(k, [])) or
                     ("no change" if abs(float(r["DELTA_MEASURED"])) < 1e-4 and
                      abs(float(r["DELTA_UNRESOLVED_HIGH"])) < 1e-4 else "below the attribution resolution")})
    for cfg, name in (("Q_A", "held-out 301-340 development vocabulary"), ("Q_B", "held-out 301-340 held-out vocabulary")):
        agg = {}
        for tag in ("V4", "V5"):
            d = json.loads((PKG / "heldout" / f"{cfg}_{tag}.json").read_text(encoding="utf-8"))
            for rec in d["drawings"]:
                for x in rec["rows"]:
                    a = agg.setdefault(x["family"], defaultdict(float))
                    if tag == "V4":
                        a["truth"] += x["truth_m3"]
                    a[f"{tag}_m"] += x["engine_m3"]
                    a[f"{tag}_u"] += x["unresolved_high_m3"]
                    if x["state"] == "SILENT":
                        a[f"{tag}_s"] += abs(x["error_m3"])
        for fam, a in sorted(agg.items()):
            rows.append({"SOURCE": f"synthetic {name}", "TRADE": "CONCRETE", "FAMILY": fam, "STOREY": "ALL (40 villas)",
                         "UNIT": "m3", "TRUTH": round(a["truth"], 3), "V4_MEASURED": a["V4_m"], "V5_MEASURED": a["V5_m"],
                         "DELTA_MEASURED": a["V5_m"] - a["V4_m"], "V4_UNRESOLVED_HIGH": a["V4_u"],
                         "V5_UNRESOLVED_HIGH": a["V5_u"], "V4_SILENT_M3": round(a["V4_s"], 3),
                         "V5_SILENT_M3": round(a["V5_s"], 3),
                         "BASIS": "frozen engine af2822a; truth from the generator's design record"})
    return _csv(rows, list(rows[0]))


def files():
    return {"03_ENGINE_WEAKNESS_REGISTER.csv": weakness_register(),
            "06_COMPLETENESS_RECONCILIATION.csv": completeness_reconciliation(),
            "07_MEASUREMENT_RELEASE_REGISTER.csv": release_register(),
            "10_BEFORE_AFTER_QUANTITIES.csv": before_after()}


# ------------------------------------------------------------------ workbook
def workbook(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    F = Font(name="Arial", size=10)
    B = Font(name="Arial", size=10, bold=True)
    IN = Font(name="Arial", size=10, color="0000FF")
    HEAD = PatternFill("solid", fgColor="D9E2EC")
    wb = Workbook()

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = B, HEAD
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = F
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    H = json.loads((PKG / "heldout" / "HELDOUT_SUMMARY.json").read_text(encoding="utf-8"))["runs"]
    T = json.loads((PKG / "st7757" / "RELEASE_TOTALS.json").read_text(encoding="utf-8"))["totals"]
    by = _rows("st7757/BY_FAMILY.csv")
    ws = wb.active
    ws.title = "Dashboard"
    ws.append(["URBAN ENGINE M7 - RELIABILITY DASHBOARD (every figure with its denominator; no single percentage)"])
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws.append(["Input: concrete unit rate (KWD per m3) - reporting input only, edit here", RATE_KWD_PER_M3])
    ws["B2"].font = IN
    ws["A2"].font = B
    lines = [
        ("Held-out quantity rows SILENT, V4 (M6 engine) - dev vocabulary", H["Q_A_V4"]["silent_rows"],
         "of 286 family / storey rows over 40 villas"),
        ("Held-out quantity rows SILENT, V5 (M7) - dev vocabulary", H["Q_A_V5"]["silent_rows"], "of 286 rows"),
        ("Held-out quantity rows SILENT, V4 - held-out vocabulary", H["Q_B_V4"]["silent_rows"], "of 286 rows"),
        ("Held-out quantity rows SILENT, V5 - held-out vocabulary", H["Q_B_V5"]["silent_rows"], "of 286 rows"),
        ("Completeness: injected elements detected, V4 / V5 (dev)", f"{H['C_A_V4']['detected']} / {H['C_A_V5']['detected']}",
         f"of {H['C_A_V5']['injected']}"),
        ("Completeness: false findings, V4 / V5 (dev)", f"{H['C_A_V4']['false_findings']} / "
         f"{H['C_A_V5']['false_findings']}", "over 40 drawings"),
        ("Metamorphic properties held, V4 / V5", f"{sum(v['pass'] for v in H['P_V4'].values())} / "
         f"{sum(v['pass'] for v in H['P_V5'].values())}", "of 100 (10 villas x 10 operations)"),
        ("ST7757 concrete measured (provisional + machine supported), m3",
         round(T["CONCRETE"]["provisional"]["base"] + T["CONCRETE"]["machine_supported"]["base"], 3),
         f"{T['CONCRETE']['provisional']['n'] + T['CONCRETE']['machine_supported']['n']} items"),
        ("ST7757 concrete RELEASED (engineer approved), m3", T["CONCRETE"]["released"]["base"],
         f"{T['CONCRETE']['released']['n']} items - nothing is released without an engineer"),
        ("ST7757 concrete unresolved exposure high, m3", T["CONCRETE"]["unresolved"]["exposure_high"],
         f"{T['CONCRETE']['unresolved']['n']} items ({T['CONCRETE']['unresolved']['n_without_exposure']} without a "
         "range)"),
    ]
    ws.append([])
    ws.append(["Metric", "Value", "Denominator / note"])
    for c in ws[4]:
        c.font, c.fill = B, HEAD
    for ln in lines:
        ws.append(list(ln))
    ws.append(["ST7757 measured concrete at the unit rate, KWD (planning, not a quotation)", "=B12*$B$2",
               "provisional - no item released"])
    ws.append(["ST7757 unresolved exposure high at the unit rate, KWD", "=B14*$B$2", "never added to the total"])
    ws.append(["Verdict", "NOT QUOTATION-READY", "0 items engineer-approved; second villa blocked; gate INCOMPLETE"])
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.font = F
    ws.column_dimensions["A"].width = 78
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 70

    rows = [[r["TRADE"], r["FAMILY"], r["STOREY"], r["UNIT"], float(r["V5_MEASURED"]), float(r["V5_UNRESOLVED_HIGH"]),
             int(r["V5_N_MEASURED"]), int(r["V5_N_UNRESOLVED"])] for r in by]
    sheet("Quantities_by_floor", ["Trade", "Family", "Storey", "Unit", "V5 measured", "V5 unresolved high",
                                  "Items measured", "Items unresolved"], rows, [16, 22, 12, 6, 14, 18, 14, 16])
    basis = _basis()
    rows = [[r["TRADE"], r["FAMILY"], r["STOREY"], float(r["V4_MEASURED"]), float(r["V5_MEASURED"]),
             None, float(r["V4_UNRESOLVED_HIGH"]), float(r["V5_UNRESOLVED_HIGH"]),
             "; ".join(basis.get((r["TRADE"], r["FAMILY"], r["STOREY"]), []))] for r in by]
    ws = sheet("Family_reconciliation", ["Trade", "Family", "Storey", "V4 measured", "V5 measured", "Delta",
                                         "V4 unresolved high", "V5 unresolved high", "Basis (V5 convention: its own "
                                         "effect)"], rows, [16, 22, 12, 14, 14, 12, 18, 18, 110])
    for i in range(2, len(rows) + 2):
        ws[f"F{i}"] = f"=E{i}-D{i}"
        ws[f"F{i}"].font = F
    rows = []
    for tr, t in T.items():
        rows.append([tr, t["unit"], t["items"], t["released"]["n"], t["released"]["base"], t["machine_supported"]["n"],
                     t["machine_supported"]["base"], t["provisional"]["n"], t["provisional"]["base"],
                     t["unresolved"]["n"], t["unresolved"]["exposure_low"], t["unresolved"]["exposure_high"]])
    sheet("Measured_vs_unresolved", ["Trade", "Unit", "Items", "Released n", "Released", "Machine supported n",
                                     "Machine supported", "Provisional n", "Provisional", "Unresolved n",
                                     "Exposure low", "Exposure high"], rows, [16, 6, 8, 11, 11, 18, 16, 13, 13, 13, 13, 13])
    rows = []
    for cfg in ("C_A", "C_B"):
        for tag in ("V4", "V5"):
            s = H[f"{cfg}_{tag}"]
            fs = s["family_states"]
            rows.append([cfg, tag, s["injected"], s["detected"], s["false_findings"], s["design_opening_bay_flags"],
                         fs.get("RIGHT", 0), fs.get("WRONG", 0), fs.get("UNKNOWN", 0)])
    sheet("Detection_classification", ["Set", "Engine", "Injected", "Detected", "False findings", "Design void flags",
                                       "Family right", "Family wrong", "Family unknown"], rows,
          [8, 8, 10, 10, 14, 18, 13, 13, 15])
    rows = []
    for cfg in ("Q_A", "Q_B", "K"):
        for tag in ("V4", "V5"):
            s = H[f"{cfg}_{tag}"]
            for fam, st in sorted(s["states_by_family"].items()):
                rows.append([cfg, tag, fam, st.get("EXACT", 0), st.get("IN_RANGE", 0), st.get("REPORTED", 0),
                             st.get("SILENT", 0)])
    for tag in ("V4", "V5"):
        for op, v in sorted(H[f"P_{tag}"].items()):
            rows.append(["P", tag, op, v["pass"], "", "", v["fail"]])
    sheet("Validation", ["Run", "Engine", "Family / operation", "Exact (or pass)", "In range", "Reported",
                         "Silent (or fail)"], rows, [8, 8, 20, 16, 10, 10, 16])
    ba = list(csv.DictReader(io.StringIO(before_after())))
    rows = [[r["SOURCE"], r["FAMILY"], float(r["TRUTH"]), float(r["V4_MEASURED"]), float(r["V5_MEASURED"]),
             float(r["V4_UNRESOLVED_HIGH"]), float(r["V5_UNRESOLVED_HIGH"]), float(r["V4_SILENT_M3"]),
             float(r["V5_SILENT_M3"])] for r in ba if r["TRUTH"]]
    sheet("Quantity_error", ["Set (40 synthetic villas)", "Family", "Design truth m3", "V4 measured", "V5 measured",
                             "V4 unresolved high", "V5 unresolved high", "V4 silent error m3", "V5 silent error m3"],
          rows, [44, 12, 15, 13, 13, 18, 18, 18, 18])
    ws = sheet("Financial_exposure", ["Item", "m3", "KWD at the Dashboard rate", "Note"], [
        ["ST7757 measured concrete (provisional)", None, None, "planning figure, not released"],
        ["ST7757 unresolved exposure low", T["CONCRETE"]["unresolved"]["exposure_low"], None, "not in the total"],
        ["ST7757 unresolved exposure high", T["CONCRETE"]["unresolved"]["exposure_high"], None, "not in the total"],
        ["Synthetic held-out: V4 silent under-count (40 villas, dev)", None, None, "what the M6 engine would have hidden"],
        ["Synthetic held-out: V5 silent under-count (40 villas, dev)", None, None, ""]], [58, 14, 26, 44])
    ws["B2"] = f"={T['CONCRETE']['provisional']['base'] + T['CONCRETE']['machine_supported']['base']:.6f}"
    v4s = sum(float(r["V4_SILENT_M3"]) for r in ba if r["SOURCE"].endswith("development vocabulary"))
    v5s = sum(float(r["V5_SILENT_M3"]) for r in ba if r["SOURCE"].endswith("development vocabulary"))
    ws["B5"], ws["B6"] = round(v4s, 3), round(v5s, 3)
    for i in range(2, 7):
        ws[f"C{i}"] = f"=B{i}*Dashboard!$B$2"
        for c in ("B", "C"):
            ws[f"{c}{i}"].font = F
    conv = json.loads((ROOT / "engine/profiles/URBAN_M7_GENERIC_CONVENTIONS_V5.json").read_text(encoding="utf-8"))
    v4 = {c["id"] for c in json.loads((ROOT / "engine/profiles/URBAN_M5_GENERIC_CONVENTIONS_V4.json")
                                      .read_text(encoding="utf-8"))["conventions"]}
    fz = json.loads((PKG / "00_ENGINE_FREEZE.json").read_text(encoding="utf-8"))
    rows = [[c["id"], "M7 (V5)" if c["id"] not in v4 else "earlier", c["provenance"], c["title"], c["rationale"]]
            for c in conv["conventions"]]
    rows += [["FREEZE", fz["commit"], "", "engine freeze before the held-out seeds", ""]]
    rows += [[k, v, "", "sha256", ""] for k, v in sorted(fz["files_sha256"].items())]
    ws = sheet("Traceability", ["Id / file", "Set / commit / sha256", "Provenance", "Title", "Rationale"], rows,
               [36, 22, 20, 80, 110])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=False)
    wb.save(path)


def main(argv):
    out = files()
    if "--check" in argv:
        same = {k: (PKG / k).read_text(encoding="utf-8") == v for k, v in out.items()}
        print(json.dumps(same, indent=1))
        return 0 if all(same.values()) else 1
    for k, v in out.items():
        (PKG / k).write_text(v, encoding="utf-8")
    workbook(PKG / "M7_RELIABILITY_WORKBOOK.xlsx")
    print(json.dumps({k: len(v.splitlines()) for k, v in out.items()}))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
