"""Human-review deliverables for the Alsenan S3 column reinforcement (quantity authority = the workbook).

    python research/alsenan_column_rebar_s3/review/build_review_s3.py

Writes next to this script:
    ALSENAN_COLUMN_REBAR_REVIEW.xlsx        10 sheets; kg columns are live formulas (count x length x unit mass)
                                            with the register value beside them as a check
    ALSENAN_COLUMN_REBAR_VISUAL_REVIEW.pdf  annotated column plans per storey + link-set sketches per section
    REVIEW_MANIFEST.json                    register hashes consumed, id cross-check, output hashes

Every row comes from the S3 registers (sha256 checked against the S3 INDEX first). The DXF is read only for the
drawing backdrop, after the S1 census rebuild has been proved byte-identical (same guard as the S1 review).
"""
import hashlib
import json
import os
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
S3 = HERE.parent
ROOT = S3.parents[1]
S1R = ROOT / "research" / "alsenan_structural_census_s1" / "review"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "research" / "external_engine_lab"))
sys.path.insert(0, str(S1R.parent))
sys.path.insert(0, str(S1R))

XLSX = HERE / "ALSENAN_COLUMN_REBAR_REVIEW.xlsx"
PDF = HERE / "ALSENAN_COLUMN_REBAR_VISUAL_REVIEW.pdf"
MANIFEST = HERE / "REVIEW_MANIFEST.json"
FLOORS = ["FOUNDATION", "GF", "1F", "2F"]
STATES = ["REBAR_COMPLETE", "REBAR_LOWER_BOUND", "REBAR_PROVISIONAL", "REBAR_BLOCKED", "NOT_REQUIRED"]
MARK = {"SOURCE_CONFLICT": ("D", "#D62728"), "BLOCKED_COMPONENT": ("s", "#FF7F0E"),
        "RESOLVED_BY_CLAIM": ("^", "#9467BD"), "CLEAR_ZONE_NOT_ESTABLISHED": ("o", "#17A2B8")}
TOPO_COLOUR = {1: "#1F3864", 2: "#2E8B57", 3: "#B8860B"}


def sha(b):
    return hashlib.sha256(b).hexdigest()


def load():
    idx = json.loads((S3 / "INDEX.json").read_text(encoding="utf-8"))
    R = {}
    for name, h in idx["outputs"].items():
        raw = (S3 / name).read_bytes()
        if sha(raw) != h:
            raise SystemExit(f"S3 register {name} does not match the S3 INDEX")
        R[name.replace(".json", "")] = json.loads(raw)
    return idx, R


def fid(flags_by_id, ids):
    return ", ".join(ids)


# ================================================================================================ workbook
def write_xlsx(idx, R):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.comments import Comment

    wb = Workbook()
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    body = Font(name="Arial", size=9)
    bold = Font(name="Arial", size=10, bold=True)
    title_font = Font(name="Arial", size=14, bold=True)
    input_font = Font(name="Arial", size=9, color="0000FF")

    def sheet(name, headers, rows, widths=None, notes=None):
        ws = wb.create_sheet(name)
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if notes and h in notes:
                c.comment = Comment(notes[h], "S3")
        for i, r in enumerate(rows, 2):
            for j, v in enumerate(r, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = body
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{L(len(headers))}{max(2, len(rows) + 1)}"
        for j, h in enumerate(headers, 1):
            ws.column_dimensions[L(j)].width = (widths or {}).get(h, max(10, min(34, len(h) + 4)))
        ws.row_dimensions[1].height = 42
        return ws

    rr = R["COLUMN_REBAR_REGISTER"]["rows"]
    ties = {t["occurrence_id"]: t for t in R["COLUMN_TIE_REGISTER"]["rows"]}
    links = defaultdict(list)
    for g in R["COLUMN_LINK_GEOMETRY_REGISTER"]["rows"]:
        links[g["occurrence_id"]].append(g)
    mains = defaultdict(list)
    for m in R["COLUMN_MAIN_BAR_REGISTER"]["rows"]:
        mains[m["occurrence_id"]].append(m)
    laps = defaultdict(list)
    for m in R["COLUMN_LAP_STARTER_REGISTER"]["rows"]:
        laps[m["occurrence_id"]].append(m)
    flags = R["COLUMN_ENGINEERING_FLAGS"]["flags"]
    S = R["COLUMN_REBAR_SUMMARY"]
    tie_rule = R["COLUMN_TIE_REGISTER"]["tie_rule"]

    # ---------------------------------------------------------------- 00_Summary (inputs + formulas)
    ws = wb.active
    ws.title = "00_Summary"
    ws["A1"] = "ALSENAN - S3 COLUMN REINFORCEMENT REVIEW"
    ws["A1"].font = title_font
    ws["A2"] = (f"Generic engine {idx['engine']} ({idx['engine_policy']}); project data from the frozen S1 census "
                f"and S2 flags (hashes in REVIEW_MANIFEST.json). Columns only - footings, beams and slabs not "
                f"calculated. This workbook is the quantity authority; the PDF is a visual check.")
    ws["A2"].font = body
    ws["A4"], ws["B4"] = "Steel density (kg/m3) - INPUT", 7850
    ws["A4"].font, ws["B4"].font = bold, input_font
    ws["C4"] = "Engineering constant used by every kg formula in this workbook (blue = input)."
    ws["A5"], ws["B5"] = "Tie rule", tie_rule["raw"]
    ws["A6"], ws["B6"] = "Tie rate per m - INPUT", tie_rule["rate_per_m"]
    ws["B6"].font = input_font
    ws["A7"], ws["B7"] = "Equivalent spacing (mm)", "=1000/B6"
    ws["A8"], ws["B8"] = "Rate semantics", "SETS_PER_M by project claim ALS-S3-CLAIM-001 (1L=6, 2L=12, 3L=18 links/m)"
    for a in ("A5", "A6", "A7", "A8"):
        ws[a].font = bold
    ws["A10"] = "Occurrence conservation"
    ws["A10"].font = bold
    ws["A11"], ws["B11"] = "Occurrences (rows in 01_Occurrences)", "=COUNTA('01_Occurrences'!A2:A500)"
    ws["A12"], ws["B12"] = "Occurrences in the S1 census (register)", S["occurrences_in"]
    ws["A13"], ws["B13"] = "Check", '=IF(B11=B12,"OK","MISMATCH")'
    row = 15
    ws.cell(row=row, column=1, value="Occurrence terminal state").font = bold
    ws.cell(row=row, column=2, value="Count (formula)").font = bold
    ws.cell(row=row, column=3, value="Register").font = bold
    ws.cell(row=row, column=4, value="Check").font = bold
    for k, st in enumerate(STATES, 1):
        r = row + k
        ws.cell(row=r, column=1, value=st)
        ws.cell(row=r, column=2, value=f"=COUNTIF('01_Occurrences'!$AF$2:$AF$500,A{r})")
        ws.cell(row=r, column=3, value=S["occurrence_states"].get(st, 0))
        ws.cell(row=r, column=4, value=f'=IF(B{r}=C{r},"OK","MISMATCH")')
    row += len(STATES) + 2
    ws.cell(row=row, column=1, value="kg by floor and release bucket (SUMIFS over 01_Occurrences)").font = bold
    heads = ["Floor", "Verified kg", "Lower-bound kg", "Provisional kg", "Blocked kg", "Total kg", "Main core kg",
             "Tie core kg", "Hook kg (provisional)", "Lap/starter/anchorage kg", "Extras kg"]
    cols = ["AA", "AB", "AC", "AD", "S", "X", "Y", "N", "Z", "AE"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=row + 1, column=j, value=h).font = bold
    for i, fl in enumerate(FLOORS + ["TOTAL"], 2):
        r = row + i
        ws.cell(row=r, column=1, value=fl)
        for j, c in enumerate(cols, 2):
            if fl == "TOTAL":
                ws.cell(row=r, column=j, value=f"=SUM({L(j)}{row + 2}:{L(j)}{row + 5})")
            else:
                ws.cell(row=r, column=j, value=f"=SUMIFS('01_Occurrences'!${c}$2:${c}$500,'01_Occurrences'!$C$2:$C$500,$A{r})")
    tot_row = row + 6
    row = tot_row + 2
    ws.cell(row=row, column=1, value="Check against the register totals").font = bold
    rk = S["release_kg"]
    for i, (lab, key, col) in enumerate([("Verified", "verified", "B"), ("Lower bound", "lower_bound", "C"),
                                         ("Provisional", "provisional", "D"), ("Blocked", "blocked", "E"),
                                         ("Total", "total", "F")], 1):
        r = row + i
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=2, value=f"={col}{tot_row}")
        ws.cell(row=r, column=3, value=rk[key])
        ws.cell(row=r, column=4, value=f'=IF(ABS(B{r}-C{r})<0.05,"OK","MISMATCH")')
    row += 8
    lines = [
        "What the buckets mean:",
        "VERIFIED - every input printed / established (core runs of storeys with printed floor-to-floor heights).",
        "LOWER_BOUND - proven minimum: foundation storey (founding level only bounded), tie perimeter over the clear "
        "zone with RATE_COUNT, parts shared by both candidates of a type conflict.",
        "PROVISIONAL - value depends on an open method or interpretation: laps / anchorage (40D vs 70D), starters "
        "(typical detail), internal links (bar positions), hooks (Urban method), the likely C7 at X12-Y02.",
        "BLOCKED - computed but not releasable (type-conflict difference at X04-Y01; ties where the clear zone is not "
        "established). Unquantified blocked parts (no value at all): " + str(S["unquantified_blocked_parts"]) + ".",
        "",
        "Front summary (de-duplicated): kg in parts touched by an open flag = "
        f"{S['de_duplicated_kg_in_parts_with_open_flags']} kg - each part counted once, never once per flag.",
        f"Open column flags: {S['open_column_flags']}; resolved by project claim: {', '.join(S['resolved_by_claim'])}; "
        f"superseded by a quantified S3 flag: {', '.join(S['superseded'])}.",
        "Zone sensitivity (clear vs full storey, RATE_COUNT, tie core): "
        f"{S['zone_clear_vs_full']['clear_rate_count_core_kg']} vs {S['zone_clear_vs_full']['full_rate_count_core_kg']} kg "
        f"(+{S['zone_clear_vs_full']['difference_core_kg']} kg) over {S['zone_clear_vs_full']['segments_compared']} "
        "segments - see 09_Sensitivity.",
        "Level-count sensitivity (RATE_COUNT vs SPACING_WITH_ENDS, base zone): "
        f"{S['rate_count_vs_spacing_with_ends']['quantification']['current_kg']} vs "
        f"{S['rate_count_vs_spacing_with_ends']['quantification']['alternative_kg']} kg.",
        "Type summary (07) is a roll-up of storey segments, never TYPE x count.",
    ]
    for i, t in enumerate(lines):
        ws.cell(row=row + i, column=1, value=t).font = bold if i == 0 else body
    ws.column_dimensions["A"].width = 46
    for c in "BCDEFGHIJK":
        ws.column_dimensions[c].width = 16

    # ---------------------------------------------------------------- 01_Occurrences
    H = ["Occurrence ID", "Chain", "Floor", "Plan tag", "Type (released)", "Candidates", "Schedule row", "Section B x D (mm)",
         "Main bars", "Main Ø", "Core length (mm)", "Lap length (mm)", "Tie Ø", "Hook provisional kg", "Tie sets/m",
         "Tie zone (base)", "Tie levels (RATE_COUNT)", "Links/level", "Total kg", "Link-1 path (mm)",
         "Link-2 path (mm)", "Link-3 path (mm)", "Tie band state", "Main kg", "Tie core kg", "Lap/starter/anch kg",
         "Verified kg", "Lower-bound kg", "Provisional kg", "Blocked kg", "Extras kg", "Occurrence state",
         "MAIN_BARS", "TIES", "LAP", "STARTER", "ANCHORAGE", "Flags", "Source refs"]
    # column letters used by 00_Summary: C floor, S total, X main, Y tie core, Z lap, AA..AD buckets, AE extras?
    rows = []
    for o in rr:
        t = ties[o["occurrence_id"]]
        lk = {g["link_id"]: g["core_path_mm"] for g in links[o["occurrence_id"]]}
        core = next((m for m in mains[o["occurrence_id"]] if m["length_kind"] == "CORE_VERTICAL_RUN"), {})
        lapr = next((m for m in laps[o["occurrence_id"]] if m["length_kind"] == "LAP_SPLICE"), {})
        k = o["kg"]
        rows.append([o["occurrence_id"], o["chain_id"], o["floor"], o["plan_tag"], o["released_candidate"],
                     " / ".join(o["candidates"]), o["schedule_row"], f"{o['section_mm'][0]} x {o['section_mm'][1]}",
                     core.get("count"), core.get("dia_mm"), core.get("length_per_piece_mm"),
                     lapr.get("length_per_piece_mm"), t["tie_diameter_mm"], k["tie_hooks_provisional"],
                     t["tie_sets_per_m"] or t["rate_per_m"], t["base_zone"],
                     (t["levels_rate_count"] or {}).get("levels"), t["links_per_level"], k["total"],
                     lk.get("LINK-1"), lk.get("LINK-2"), lk.get("LINK-3"), t["band_state"], k["main_core"],
                     k["tie_core"], k["lap_starter_anchorage"], k["verified"], k["lower_bound"], k["provisional"],
                     k["blocked"], k["extras"], o["occurrence_state"], o["components"]["MAIN_BARS"],
                     o["components"]["TIES"], o["components"]["LAP"], o["components"]["STARTER"],
                     o["components"]["ANCHORAGE"], ", ".join(o["open_flags"]),
                     f"{o['source_refs']['plan']['drawing']} {'/'.join(o['source_refs']['plan']['sheets'])} "
                     f"handles {','.join(o['source_refs']['plan']['handles'])}; schedule {o['schedule_row']}"])
    ws1 = sheet("01_Occurrences", H, rows, widths={"Occurrence ID": 40, "Chain": 30, "Flags": 40, "Source refs": 50})
    # the summary refers to fixed letters - assert the layout
    assert [H.index(x) for x in ("Floor", "Total kg", "Main kg", "Tie core kg", "Lap/starter/anch kg", "Verified kg",
                                 "Lower-bound kg", "Provisional kg", "Blocked kg", "Hook provisional kg",
                                 "Extras kg", "Occurrence state")] == [2, 18, 23, 24, 25, 26, 27, 28, 29, 13, 30, 31]

    # ---------------------------------------------------------------- 02_Main_Bars (formula kg)
    H2 = ["Part ID", "Occurrence ID", "Floor", "Type", "Section B x D (mm)", "Length kind", "Bars", "Ø (mm)",
          "Length per bar (mm)", "kg (formula)", "kg (register)", "Check", "Release state", "Continues below",
          "Continues above", "Storey interval basis", "Rule", "Flags", "Why"]
    rows = []
    for o in rr:
        for m in mains[o["occurrence_id"]]:
            rows.append([m["part_id"], m["occurrence_id"], m["floor"], m["type"],
                         f"{m['section_mm'][0]} x {m['section_mm'][1]}" if m.get("section_mm") else None,
                         m["length_kind"], m["count"], m["dia_mm"], m["length_per_piece_mm"], None, m["kg"], None,
                         m["release_state"], m.get("chain_continues_below"), m.get("chain_continues_above"),
                         (m.get("interval") or {}).get("basis"), m.get("rule_id"), ", ".join(m["flags"]), m["why"]])
    ws2 = sheet("02_Main_Bars", H2, rows, widths={"Part ID": 50, "Occurrence ID": 38, "Why": 60,
                                                  "Storey interval basis": 50})
    for i in range(2, len(rows) + 2):
        ws2[f"J{i}"] = f'=IF(OR(G{i}="",I{i}=""),"",G{i}*I{i}/1000*PI()/4*(H{i}/1000)^2*\'00_Summary\'!$B$4)'
        ws2[f"L{i}"] = f'=IF(J{i}="","no value",IF(ABS(J{i}-K{i})<0.01,"OK","MISMATCH"))'

    # ---------------------------------------------------------------- 03_Tie_Sets
    H3 = ["Occurrence ID", "Floor", "Type", "Section B x D (mm)", "Tie Ø", "Rate /m", "Semantics used", "Claim",
          "Tie sets/m", "Links/m", "Equivalent spacing (mm)", "Band", "Band state", "Band claim", "Topology",
          "Links/level", "Clear zone (mm)", "Clear zone basis", "Full zone (mm)", "Base zone",
          "Levels RATE_COUNT (clear)", "Levels SPACING_WITH_ENDS (clear)", "Levels RATE_COUNT (full)",
          "Levels SPACING_WITH_ENDS (full)", "Links placed (base, RATE_COUNT)", "Perimeter path (mm)",
          "Sum of link paths (mm)", "Tie core kg", "Hook kg (provisional)", "Part states"]
    rows = []
    for o in rr:
        t = ties[o["occurrence_id"]]
        zc, zf = t["zones"]["CLEAR_COLUMN_ZONE"], t["zones"]["FULL_STOREY"]
        rows.append([o["occurrence_id"], o["floor"], t["type"], f"{t['section_mm'][0]} x {t['section_mm'][1]}",
                     t["tie_diameter_mm"], t["rate_per_m"], t["per_metre_semantics_used"], t["semantics_claim"],
                     t["tie_sets_per_m"], t["links_per_m"], t["equivalent_spacing_mm"], t["band"], t["band_state"],
                     t["band_claim"], t["topology"], t["links_per_level"], zc["length_mm"], zc["basis"],
                     zf["length_mm"], t["base_zone"], (zc["levels"] or {}).get("RATE_COUNT"),
                     (zc["levels"] or {}).get("SPACING_WITH_ENDS"), (zf["levels"] or {}).get("RATE_COUNT"),
                     (zf["levels"] or {}).get("SPACING_WITH_ENDS"), (t["levels_rate_count"] or {}).get("links"),
                     t["perimeter_mm"], t["sum_link_paths_mm"], t["tie_core_kg"], t["hook_kg_provisional"],
                     "; ".join(f"{p['length_kind']}:{p['release_state']}" for p in t["parts"])])
    sheet("03_Tie_Sets", H3, rows, widths={"Occurrence ID": 38, "Clear zone basis": 50, "Part states": 60})

    # ---------------------------------------------------------------- 04_Link_Cutting_Lengths (formula kg per level)
    H4 = ["Occurrence ID", "Floor", "Type", "Section B x D (mm)", "Link", "Bars restrained (range per long face)",
          "Bars restrained", "Across (mm)", "Along (mm)", "Core path (mm) = 2 x (across + along)", "Check path",
          "Hook 1 (mm)", "Hook 2 (mm)", "Hook state", "Geometry state", "Link ranges", "Bar position method",
          "Cover rule", "Cover (mm)", "Cover source", "Tie Ø", "Core kg per level (formula)"]
    rows = []
    for o in rr:
        for g in links[o["occurrence_id"]]:
            rows.append([g["occurrence_id"], g["floor"], g["type"], f"{g['section_mm'][0]} x {g['section_mm'][1]}",
                         g["link_id"], str(g["bar_range_per_long_face"]), g["bars_restrained"], g["across_mm"],
                         g["along_mm"], g["core_path_mm"], None, g["hook_1_mm"], g["hook_2_mm"], g["hook_state"],
                         g["geometry_state"], g["link_ranges_state"], g["bar_position_method"], g["cover_rule_id"],
                         g["cover_mm"], g["cover_source_ref"], g["tie_dia_mm"], None])
    ws4 = sheet("04_Link_Cutting_Lengths", H4, rows, widths={"Occurrence ID": 38})
    for i in range(2, len(rows) + 2):
        ws4[f"K{i}"] = f'=IF(ABS(2*(H{i}+I{i})-J{i})<0.01,"OK","MISMATCH")'
        ws4[f"V{i}"] = f"=J{i}/1000*PI()/4*(U{i}/1000)^2*'00_Summary'!$B$4"

    # ---------------------------------------------------------------- 05_Laps_Starters
    H5 = ["Part ID", "Occurrence ID", "Floor", "Type", "Length kind", "Bars", "Ø (mm)", "Length per bar (mm)",
          "kg (formula)", "kg (register)", "Release state", "Rule", "Alternative (x Ø)", "kg at alternative",
          "Footing", "Flags", "Why"]
    rows = []
    for o in rr:
        for m in laps[o["occurrence_id"]]:
            rows.append([m["part_id"], m["occurrence_id"], m["floor"], m["type"], m["length_kind"], m["count"],
                         m["dia_mm"], m["length_per_piece_mm"], None, m["kg"], m["release_state"], m.get("rule_id"),
                         (m.get("alternative_D") or [None])[0], m.get("kg_at_alternative"),
                         json.dumps(m.get("footing")) if m.get("footing") else None, ", ".join(m["flags"]), m["why"]])
    ws5 = sheet("05_Laps_Starters", H5, rows, widths={"Part ID": 50, "Occurrence ID": 38, "Why": 60, "Footing": 50})
    for i in range(2, len(rows) + 2):
        ws5[f"I{i}"] = f'=IF(OR(F{i}="",H{i}=""),"",F{i}*H{i}/1000*PI()/4*(G{i}/1000)^2*\'00_Summary\'!$B$4)'

    # ---------------------------------------------------------------- 06_Floor_Summary
    H6 = ["Floor", "Occurrences", "Main core kg", "Tie core kg", "Hook kg (provisional)", "Lap/starter/anch kg",
          "Extras kg", "Verified kg", "Lower-bound kg", "Provisional kg", "Blocked kg", "Total kg", "Check total"]
    ws6 = sheet("06_Floor_Summary", H6, [[f] for f in FLOORS + ["TOTAL"]])
    src = {"Main core kg": "X", "Tie core kg": "Y", "Hook kg (provisional)": "N", "Lap/starter/anch kg": "Z",
           "Extras kg": "AE", "Verified kg": "AA", "Lower-bound kg": "AB", "Provisional kg": "AC",
           "Blocked kg": "AD", "Total kg": "S"}
    for i, fl in enumerate(FLOORS + ["TOTAL"], 2):
        for j, h in enumerate(H6[1:], 2):
            if fl == "TOTAL":
                ws6.cell(row=i, column=j, value=f"=SUM({L(j)}2:{L(j)}5)" if h != "Check total" else None)
                continue
            if h == "Occurrences":
                ws6.cell(row=i, column=j, value=f"=COUNTIF('01_Occurrences'!$C$2:$C$500,$A{i})")
            elif h == "Check total":
                ws6.cell(row=i, column=j, value=f'=IF(ABS(SUM(H{i}:K{i})-L{i})<0.05,"OK","MISMATCH")')
            else:
                ws6.cell(row=i, column=j, value=f"=SUMIFS('01_Occurrences'!${src[h]}$2:${src[h]}$500,"
                                                f"'01_Occurrences'!$C$2:$C$500,$A{i})")

    # ---------------------------------------------------------------- 07_Type_Summary
    types = sorted({o["released_candidate"] for o in rr}, key=lambda t: (len(t), t))
    H7 = ["Type (released)", "Segments", "Main core kg", "Tie core kg", "Hook kg (provisional)",
          "Lap/starter/anch kg", "Total kg", "Note"]
    ws7 = sheet("07_Type_Summary", H7, [[t] for t in types])
    for i, t in enumerate(types, 2):
        ws7.cell(row=i, column=2, value=f"=COUNTIF('01_Occurrences'!$E$2:$E$500,$A{i})")
        for j, c in zip(range(3, 8), ["X", "Y", "N", "Z", "S"]):
            ws7.cell(row=i, column=j, value=f"=SUMIFS('01_Occurrences'!${c}$2:${c}$500,'01_Occurrences'!$E$2:$E$500,$A{i})")
        ws7.cell(row=i, column=8, value="roll-up of storey segments (each with its own section, bars and storey)")

    # ---------------------------------------------------------------- 08_Engineering_Flags
    H8 = ["Flag ID", "S3 kind", "Status", "Release effect", "Issue type", "Elements", "Summary",
          "Quantity affected (kg)", "Current kg (basis used)", "Alternative kg", "Alternative / candidates",
          "Released kg in the flagged parts",
          "Current interpretation", "Question for the consultant", "Where to check", "Resolution", "History"]
    rows = []
    for f in flags:
        q = f.get("s3_quantification_kg") or {}
        alt = q.get("alternative_kg")
        cand = q.get("kg_by_candidate") or q.get("differing_kg_by_candidate")
        rows.append([f["flag_id"], f["s3_kind"], f["status"], f["release_effect"], f["issue_type"],
                     len(f["element_ids"]), f["issue_summary"], q.get("quantity_affected_kg"),
                     q.get("current_kg"), alt,
                     (q.get("alternative") or "") + (" " + json.dumps(cand, ensure_ascii=False) if cand else ""),
                     q.get("current_released_kg"),
                     f.get("current_interpretation"), f.get("question_for_engineer"), f.get("where_to_check"),
                     f.get("resolution") if not isinstance(f.get("resolution"), dict) else json.dumps(f["resolution"]),
                     " -> ".join(h["to"] for h in f["history"])])
    sheet("08_Engineering_Flags", H8, rows, widths={"Summary": 70, "Question for the consultant": 70,
                                                     "Where to check": 50, "Alternative / candidates": 40})

    # ---------------------------------------------------------------- 09_Sensitivity
    H9 = ["Occurrence ID", "Floor", "Type", "Clear zone (mm)", "Full zone (mm)", "Levels clear RATE", "Levels clear SWE",
          "Levels full RATE", "Levels full SWE", "Tie core kg clear RATE", "Tie core kg clear SWE",
          "Tie core kg full RATE", "Tie core kg full SWE", "Full - clear (RATE) kg", "SWE - RATE (clear) kg",
          "Hook kg clear RATE", "Hook kg full RATE", "Drawn section (mm)", "Ties kg at drawn section",
          "Lap/anch kg 40D", "Lap/anch kg 70D"]
    rows = []
    for x in R["COLUMN_SENSITIVITY_REGISTER"]["rows"]:
        cr, cs = x.get("CLEAR_COLUMN_ZONE|RATE_COUNT") or {}, x.get("CLEAR_COLUMN_ZONE|SPACING_WITH_ENDS") or {}
        fr, fs = x.get("FULL_STOREY|RATE_COUNT") or {}, x.get("FULL_STOREY|SPACING_WITH_ENDS") or {}
        rows.append([x["occurrence_id"], x["floor"], x["type"], cr.get("zone_mm"), fr.get("zone_mm"),
                     cr.get("levels"), cs.get("levels"), fr.get("levels"), fs.get("levels"), cr.get("core_kg"),
                     cs.get("core_kg"), fr.get("core_kg"), fs.get("core_kg"), None, None, cr.get("hook_kg"),
                     fr.get("hook_kg"), str(x["drawn_section_mm"]) if x["drawn_section_mm"] else None,
                     x["ties_kg_drawn_section"], x["lap_kg_current"], x["lap_kg_alternative"]])
    ws9 = sheet("09_Sensitivity", H9, rows, widths={"Occurrence ID": 38})
    for i in range(2, len(rows) + 2):
        ws9[f"N{i}"] = f'=IF(OR(J{i}="",L{i}=""),"",L{i}-J{i})'
        ws9[f"O{i}"] = f'=IF(OR(J{i}="",K{i}=""),"",K{i}-J{i})'
    n = len(rows) + 3
    ws9[f"A{n}"] = "TOTAL"
    for c in "JKLMNOPQSTU":
        ws9[f"{c}{n}"] = f"=SUM({c}2:{c}{n - 2})"
    wb.save(XLSX)


# ================================================================================================ drawings
def write_pdf(R):
    import math
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.collections import LineCollection
    from matplotlib.lines import Line2D
    from matplotlib.patches import Arc, Polygon as MPoly, Rectangle
    import build_review_s1 as RV1

    S, B = RV1.verified_build()
    src = B["src"]
    ents = src.entities()
    xs, ys = [], []
    for sh in ("FP", "CAP", "GBP", "GFRS", "FFRS", "SFRS"):
        for e in ents[sh]:
            if e["type"] == "LWPOLYLINE" and e["layer"] in ("S-COL.BON", "S-FOOTINGS"):
                xs += [q[0] for q in e["pts"]]
                ys += [q[1] for q in e["pts"]]
    EXT = (min(xs) - 1800, min(ys) - 1800, max(xs) + 2600, max(ys) + 1800)
    colpts = {f"{sh}:{e['handle']}": e["pts"] for sh in ents for e in ents[sh]
              if e["type"] == "LWPOLYLINE" and e["layer"] == "S-COL.BON"}

    def backdrop(ax, sh):
        segs = []
        for e in ents[sh]:
            if e["type"] == "LINE":
                segs.append((e["a"], e["b"]))
            elif e["type"] == "LWPOLYLINE":
                for sg in e.get("segs", []):
                    if sg[0] == "ARC":
                        _, cxy, r, a0, a1 = sg
                        ax.add_patch(Arc(cxy, 2 * r, 2 * r, theta1=math.degrees(a0), theta2=math.degrees(a1),
                                         lw=0.12, color="#B8B8B8", zorder=1))
                    else:
                        segs.append(sg)
            elif e["type"] in ("ARC", "CIRCLE"):
                ax.add_patch(Arc(e["c"], 2 * e["r"], 2 * e["r"], theta1=math.degrees(e["a0"]),
                                 theta2=math.degrees(e["a1"]), lw=0.12, color="#B8B8B8", zorder=1))
        ax.add_collection(LineCollection(segs, linewidths=0.12, colors="#B8B8B8", zorder=1))

    rr = R["COLUMN_REBAR_REGISTER"]["rows"]
    ties = {t["occurrence_id"]: t for t in R["COLUMN_TIE_REGISTER"]["rows"]}
    occ_src = {}
    s1 = json.loads((ROOT / "research" / "alsenan_structural_census_s1" / "COLUMN_OCCURRENCE_REGISTER.json")
                    .read_text(encoding="utf-8"))
    for o in s1["rows"]:
        occ_src[o["column_id"]] = o
    labelled = set()
    pdf = PdfPages(PDF)
    fig = plt.figure(figsize=(46.8, 33.1))
    fig.text(0.05, 0.9, "ALSENAN - S3 COLUMN REINFORCEMENT - VISUAL REVIEW", fontsize=46, weight="bold")
    S3s = R["COLUMN_REBAR_SUMMARY"]
    lines = ["The workbook ALSENAN_COLUMN_REBAR_REVIEW.xlsx is the quantity authority. This PDF lets a reviewer check "
             "WHERE each storey segment is, what it was read as and how its ties are arranged.",
             "Label per column: occurrence ID | released type | schedule section B x D mm | main bars | tie topology "
             "(1L / 2L / 3L) | total kg | open S3/S2 flag ids.",
             "Column fill colour = links per tie level: 1L navy, 2L green, 3L ochre.",
             f"Occurrences: {S3s['occurrences_in']} in -> {S3s['terminal_records_out']} terminal records. "
             f"Topology: {S3s['links_per_level_counts']}.",
             "Last pages: link sets drawn to scale from the computed geometry (cover, bar positions, link ranges) for "
             "every distinct section / bar / topology combination.", "", "Markers:"]
    for i, t in enumerate(lines):
        fig.text(0.05, 0.84 - i * 0.03, t, fontsize=19)
    for i, (k, (m, c)) in enumerate(MARK.items()):
        y = 0.84 - (len(lines) + i) * 0.03
        fig.lines.append(Line2D([0.065], [y + 0.006], marker=m, markersize=26, markerfacecolor="none",
                                markeredgecolor=c, markeredgewidth=3, transform=fig.transFigure))
        fig.text(0.085, y, k, fontsize=19, color=c, weight="bold")
    pdf.savefig(fig)
    plt.close(fig)
    sheet_for = {"FOUNDATION": "CAP", "GF": "GFRS", "1F": "FFRS", "2F": "SFRS"}
    fc = Counter(f for r in rr for f in r["open_flags"])
    wide = sorted(f for f, n in fc.items() if n >= 0.8 * len(rr))       # project-wide method flags: page subtitle
    for st, sh in sheet_for.items():
        rows = [r for r in rr if r["floor"] == st]
        fig = plt.figure(figsize=(46.8, 33.1))
        ax = fig.add_axes([0.02, 0.03, 0.96, 0.89])
        ax.set_xlim(EXT[0], EXT[2])
        ax.set_ylim(EXT[1], EXT[3])
        ax.set_aspect("equal")
        ax.axis("off")
        backdrop(ax, sh)
        kg = sum(r["kg"]["total"] or 0 for r in rows)
        fig.text(0.02, 0.975, f"COLUMNS {st} - {len(rows)} storey segments - {kg:,.1f} kg (all buckets)",
                 fontsize=30, weight="bold", va="top")
        fig.text(0.02, 0.952, f"Backdrop: {src.sheets[sh]['title']} (PDF p.{src.sheets[sh]['pdf_page']}). "
                              "Label: ID | type | B x D | bars | topology | kg | element flags. Flags on (almost) "
                              f"every column, not repeated in labels: {', '.join(wide)}.", fontsize=14, va="top")
        h = [Line2D([], [], marker=m, ls="", markersize=16, markerfacecolor="none", markeredgecolor=c,
                    markeredgewidth=2.5, label=k) for k, (m, c) in MARK.items()]
        h += [Rectangle((0, 0), 1, 1, fc=c, label=f"{n}L") for n, c in TOPO_COLOUR.items()]
        fig.legend(handles=h, loc="upper right", fontsize=13, ncol=4, frameon=True, bbox_to_anchor=(0.985, 0.985))
        for r in rows:
            o = occ_src[r["occurrence_id"]]
            t = ties[r["occurrence_id"]]
            pairs = list(zip(o["plan_source"]["sheets"], o["plan_source"]["handles"]))
            oid = next((f"{s}:{hh}" for s, hh in pairs if s == sh), None) or \
                (f"{pairs[0][0]}:{pairs[0][1]}" if pairs else None)
            nl = r["links_per_level"] or 0
            pts = colpts.get(oid)
            if pts:
                ax.add_patch(MPoly(pts, closed=True, fc=TOPO_COLOUR.get(nl, "#888888"),
                                   ec=TOPO_COLOUR.get(nl, "#888888"), lw=0.4, alpha=0.9, zorder=4))
            x, y = o["plan_centre_mm"]
            marks = []
            if len(r["candidates"]) > 1:
                marks.append("SOURCE_CONFLICT")
            if any(v == "BLOCKED" for v in r["components"].values()):
                marks.append("BLOCKED_COMPONENT")
            if r["tie_band_state"] == "RESOLVED_BY_CLAIM":
                marks.append("RESOLVED_BY_CLAIM")
            if t["base_zone"] != "CLEAR_COLUMN_ZONE":
                marks.append("CLEAR_ZONE_NOT_ESTABLISHED")
            for i, k in enumerate(marks):
                m, c = MARK[k]
                ax.plot([x], [y], marker=m, markersize=26 + 7 * i, markerfacecolor="none", markeredgecolor=c,
                        markeredgewidth=1.6, zorder=6)
            colr = MARK[marks[0]][1] if marks else "#1F3864"
            cand = r["released_candidate"] + (f" (likely; also {'/'.join(c for c in r['candidates'] if c != r['released_candidate'])})"
                                              if len(r["candidates"]) > 1 else "")
            own = [f for f in r["open_flags"] if f not in wide]
            fl = ", ".join(own[:6]) + (" ..." if len(own) > 6 else "")
            ax.text(x + 320, y + 160, f"{r['occurrence_id']}\n{cand}  {r['section_mm'][0]}x{r['section_mm'][1]}  "
                                      f"{r['main_bars']}  {nl}L  {r['kg']['total']:.1f} kg\n{fl}",
                    fontsize=6.2, color=colr, zorder=7, va="bottom",
                    bbox=dict(fc="white", ec=colr, lw=0.3, alpha=0.86, pad=0.6))
            labelled.add(r["occurrence_id"])
        pdf.savefig(fig)
        plt.close(fig)
    # ---------------------------------------------------------------- link-set sketches
    geo = defaultdict(list)
    for g in R["COLUMN_LINK_GEOMETRY_REGISTER"]["rows"]:
        geo[g["occurrence_id"]].append(g)
    combos = {}
    for r in rr:
        gs = geo[r["occurrence_id"]]
        if not gs:
            continue
        key = (tuple(r["section_mm"]), r["main_bars"], r["links_per_level"],
               tuple(tuple(g["bar_range_per_long_face"]) for g in gs))
        combos.setdefault(key, []).append(r["occurrence_id"])
    keys = sorted(combos, key=lambda k: (k[2], k[0][1], k[0][0], k[1]))
    per_page = 12
    for p0 in range(0, len(keys), per_page):
        fig = plt.figure(figsize=(46.8, 33.1))
        fig.text(0.02, 0.975, "LINK SETS - drawn to scale from the computed geometry (cover, tie Ø, bar positions, "
                              "link bar ranges). One sketch per distinct section / bars / topology.",
                 fontsize=24, weight="bold", va="top")
        for i, key in enumerate(keys[p0:p0 + per_page]):
            (Bm, Dm), bars, nl, ranges = key
            occ = combos[key][0]
            gs = geo[occ]
            ax = fig.add_axes([0.03 + (i % 4) * 0.24, 0.66 - (i // 4) * 0.31, 0.21, 0.25])
            L_, T_ = max(Bm, Dm), min(Bm, Dm)
            ax.add_patch(Rectangle((0, 0), L_, T_, fc="#EEEEEE", ec="black", lw=1.2))
            cov, td = gs[0]["cover_mm"], gs[0]["tie_dia_mm"]
            pos = gs[0]["bar_positions_mm"]
            db = 2 * (pos[0] - cov - td)
            for yy in (cov + td + db / 2, T_ - cov - td - db / 2):
                for xx in pos:
                    ax.add_patch(plt.Circle((xx, yy), db / 2, color="black", zorder=5))
            cols = ["#1F77B4", "#D62728", "#2CA02C", "#9467BD"]
            for j, g in enumerate(gs):
                a, b = g["bar_range_per_long_face"]
                x0 = pos[a] - db / 2 - td / 2
                x1 = pos[b] + db / 2 + td / 2
                off = (j - (len(gs) - 1) / 2) * 3.0
                ax.add_patch(Rectangle((x0 + off, cov + td / 2 + off), x1 - x0, T_ - 2 * cov - td, fill=False,
                                       ec=cols[j % 4], lw=2.0, zorder=4))
                ax.text(x0 + off, T_ + 12 + 26 * j, f"{g['link_id']} bars {a}-{b}: path {g['core_path_mm']:.0f} mm "
                                                    f"({g['geometry_state']})", fontsize=10, color=cols[j % 4])
            ax.set_xlim(-30, L_ + 30)
            ax.set_ylim(-60, T_ + 30 + 26 * len(gs))
            ax.set_aspect("equal")
            ax.axis("off")
            ax.set_title(f"{T_} x {L_} mm  {bars}  {nl}L  cover {cov}  Ø{td}\n{len(combos[key])} segment(s), e.g. "
                         f"{occ}", fontsize=11)
        pdf.savefig(fig)
        plt.close(fig)
    pdf.close()
    return labelled


def main():
    idx, R = load()
    write_xlsx(idx, R)
    recalc = os.environ.get("RECALC_SCRIPT")
    recalc_state = "NOT_RUN (formulas recalculate when opened in Excel)"
    if recalc and Path(recalc).exists():
        out = subprocess.run([sys.executable, recalc, str(XLSX), "240"], capture_output=True, text=True)
        res = json.loads(out.stdout[out.stdout.index("{"):])
        if res.get("status") != "success" or res.get("total_errors"):
            raise SystemExit(f"workbook recalculation failed: {res}")
        recalc_state = f"RECALCULATED ({res['total_formulas']} formulas, 0 errors)"
    labelled = write_pdf(R) if "--no-pdf" not in sys.argv else set()
    ids = {r["occurrence_id"] for r in R["COLUMN_REBAR_REGISTER"]["rows"]}
    man = {"round": "S3", "purpose": "human review of the column reinforcement before the next structural trade",
           "quantity_authority": XLSX.name, "s3_registers_consumed": idx["outputs"],
           "id_cross_check": {"workbook_ids": len(ids), "drawing_labels": len(labelled),
                              "missing_on_drawing": sorted(ids - labelled) if labelled else "PDF not built",
                              "label_not_in_workbook": sorted(labelled - ids)},
           "workbook_formulas": recalc_state,
           "outputs": {XLSX.name: sha(XLSX.read_bytes()), **({PDF.name: sha(PDF.read_bytes())} if labelled else {})}}
    MANIFEST.write_text(json.dumps(man, indent=1, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({k: man[k] for k in ("id_cross_check", "workbook_formulas")}, default=str)[:800])
    if labelled and (ids - labelled or labelled - ids):
        raise SystemExit("workbook / drawing id mismatch")


if __name__ == "__main__":
    main()
