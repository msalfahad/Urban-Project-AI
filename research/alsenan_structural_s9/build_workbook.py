"""S9 workbook view: one Excel file over the frozen S9 CSVs (concrete, reinforcement, completeness, exceptions).

    python3 -I research/alsenan_structural_s9/build_workbook.py

Reads only the frozen S9 outputs (00 - 15; the freeze manifest is verified first) and writes
S9_STRUCTURAL_BOQ.xlsx next to them. The workbook is a view: it is not frozen, adds no quantity and is not committed
(*.xlsx is git-ignored). Totals are spreadsheet formulas over the listed rows, with a check against the frozen
summary. The digest of every CSV it reads is listed on the Read Me sheet.
"""

from __future__ import annotations

import csv
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from engine.source import delta_release as DR  # noqa: E402

OUT = HERE / "S9_STRUCTURAL_BOQ.xlsx"
MANIFEST = HERE / "16_S9_FREEZE_MANIFEST.json"
FONT = "Arial"
HEAD = PatternFill("solid", start_color="1F3864")
NOTE = PatternFill("solid", start_color="F2F2F2")
TOTAL = PatternFill("solid", start_color="D9E1F2")
RELEASED = ("RELEASED_FROZEN_STAGE", "RELEASED_S9_DELTA")


def rows(name):
    with open(HERE / name, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def num(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except ValueError:
        return v


def style_header(ws, row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def table(ws, top, cols, data, numeric=(), widths=None):
    """cols: [(header, key)]; returns (first data row, last data row)."""
    for j, (h, _) in enumerate(cols, 1):
        ws.cell(row=top, column=j, value=h)
    style_header(ws, top, len(cols))
    r = top
    for d in data:
        r += 1
        for j, (h, k) in enumerate(cols, 1):
            v = d.get(k)
            v = num(v) if k in numeric else (v if v not in ("",) else None)
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name=FONT)
            if k in numeric and isinstance(v, float):
                c.number_format = "#,##0.000;-#,##0.000;-"
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    return top + 1, r


def total_row(ws, r, label_col, label, sum_cols, first, last):
    ws.cell(row=r, column=label_col, value=label).font = Font(name=FONT, bold=True)
    for col in sum_cols:
        L = get_column_letter(col)
        c = ws.cell(row=r, column=col, value=f"=SUM({L}{first}:{L}{last})")
        c.font = Font(name=FONT, bold=True)
        c.number_format = "#,##0.000"
    for col in range(1, max(sum_cols) + 1):
        ws.cell(row=r, column=col).fill = TOTAL


def note(ws, r, text, span=8):
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(name=FONT, italic=True)
    c.fill = NOTE
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    ws.row_dimensions[r].height = 30


def main():
    DR.verify_frozen(MANIFEST, ROOT)
    s = json.loads((HERE / "14_RELEASE_SUMMARY.json").read_text(encoding="utf-8"))
    used = ["02_CONCRETE_BOQ_RECONCILIATION.csv", "03_REINFORCEMENT_BOQ_RECONCILIATION.csv",
            "04_OWNERSHIP_AND_PRECEDENCE_REGISTER.csv", "05_MISSING_AND_BLOCKED_REGISTER.csv",
            "06_FLOOR_SUMMARY.csv", "07_RELEASED_STRUCTURAL_BOQ.csv", "09_ENGINEER_RFI_REGISTER.csv",
            "11_OVERLAP_AND_DOUBLE_COUNT_AUDIT.csv", "12_CONSERVATION_CHECKS.csv", "13_COVERAGE_BY_FAMILY.csv",
            "14_RELEASE_SUMMARY.json"]
    wb = Workbook()

    # ---- Read Me
    ws = wb.active
    ws.title = "Read Me"
    lines = [("S9 whole-building structural BOQ reconciliation (workbook view)", True),
             (f"Baseline {s['baseline']}; date {s['date']}; frozen manifest 16_S9_FREEZE_MANIFEST.json "
              f"(sha256 {hashlib.sha256(MANIFEST.read_bytes()).hexdigest()}).", False),
             ("THIS IS A PARTIAL STRUCTURAL BOQ, NOT A COMPLETE BUILDING ESTIMATE. The project gate stays INCOMPLETE.",
              True),
             ("Only rows in the RELEASED_FROZEN_STAGE / RELEASED_S9_DELTA lanes are in a released total. Conditional, "
              "indicative, blocked and conflict rows are shown with their reason and are never added; an unmeasured "
              "component is unknown, never zero.", False),
             ("The workbook is a view over the frozen CSVs: it is not frozen and adds no quantity. Totals are "
              "formulas over the listed rows; the 'Check' cells compare them with the frozen summary.", False),
             ("", False), ("Lanes", True),
             ("RELEASED_FROZEN_STAGE: carried unchanged from a frozen stage (S3.1 - S8.8, with S9-C01 for column ties)",
              False),
             ("RELEASED_S9_DELTA: measured by S9 under rule S9-RR1 (one auditable delta each)", False),
             ("CONDITIONAL_NOT_RELEASED: a value under a stated condition (point and range)", False),
             ("INDICATIVE_ONLY_NOT_RELEASED: an indicative value whose unknowns have no bound", False),
             ("BLOCKED_UNQUANTIFIED / SOURCE_CONFLICT: no value (missing or conflicting source)", False),
             ("EXCLUDED_OWNED_BY_OTHER_FAMILY / SUPERSEDED / NOT_APPLICABLE / NOT_IN_SOURCE: never in a total", False),
             ("", False), ("Source files (sha256)", True)]
    lines += [(f"{n}  {hashlib.sha256((HERE / n).read_bytes()).hexdigest()}", False) for n in used]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=b, size=12 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140

    # ---- Released BOQ
    ws = wb.create_sheet("Released BOQ")
    boq = rows("07_RELEASED_STRUCTURAL_BOQ.csv")
    a1 = [x for x in boq if x["SECTION"] == "A1 CONCRETE"]
    a2 = [x for x in boq if x["SECTION"] == "A2 REINFORCEMENT"]
    ws.cell(row=1, column=1, value="A1 Released concrete (m3)").font = Font(name=FONT, bold=True, size=12)
    f1, l1 = table(ws, 2, [("Family", "FAMILY"), ("Storey", "STOREY"), ("Lane", "LANE"), ("Components", "COMPONENTS"),
                           ("m3", "QUANTITY")], a1, numeric=("COMPONENTS", "QUANTITY"), widths=[22, 14, 26, 14, 16, 16])
    total_row(ws, l1 + 1, 1, "Total released concrete (partial)", [5], f1, l1)
    ws.cell(row=l1 + 2, column=1, value="Check: frozen summary").font = Font(name=FONT)
    ws.cell(row=l1 + 2, column=5, value=s["released_concrete_m3_total"]).font = Font(name=FONT, color="0000FF")
    ws.cell(row=l1 + 3, column=1, value="Check: difference (must be 0)").font = Font(name=FONT)
    ws.cell(row=l1 + 3, column=5, value=f"=ROUND(E{l1 + 1}-E{l1 + 2},6)").font = Font(name=FONT)
    top = l1 + 6
    ws.cell(row=top - 1, column=1, value="A2 Released reinforcement (kg)").font = Font(name=FONT, bold=True, size=12)
    f2, l2 = table(ws, top, [("Family", "FAMILY"), ("Diameter mm", "DIA_MM"), ("Lane", "LANE"),
                             ("Items", "COMPONENTS"), ("kg", "QUANTITY")], a2, numeric=("DIA_MM", "COMPONENTS", "QUANTITY"))
    ws.cell(row=top, column=6, value="t")
    style_header(ws, top, 6)
    for r in range(f2, l2 + 1):
        c = ws.cell(row=r, column=6, value=f"=E{r}/1000")
        c.font = Font(name=FONT)
        c.number_format = "#,##0.000"
    total_row(ws, l2 + 1, 1, "Total released reinforcement (partial)", [5, 6], f2, l2)
    ws.cell(row=l2 + 2, column=1, value="Check: frozen summary").font = Font(name=FONT)
    ws.cell(row=l2 + 2, column=5, value=s["released_reinforcement_kg_total"]).font = Font(name=FONT, color="0000FF")
    ws.cell(row=l2 + 3, column=1, value="Check: difference (must be 0)").font = Font(name=FONT)
    ws.cell(row=l2 + 3, column=5, value=f"=ROUND(E{l2 + 1}-E{l2 + 2},6)").font = Font(name=FONT)
    note(ws, l2 + 5, "Blue cells are the frozen summary figures (inputs); black cells are formulas. Partial BOQ: see "
                     "the Completeness and Exceptions sheets for what is not released.", span=6)
    ws.freeze_panes = None

    # ---- Concrete
    ws = wb.create_sheet("Concrete")
    conc = rows("02_CONCRETE_BOQ_RECONCILIATION.csv")
    ccols = [("Component", "COMPONENT_ID"), ("Family", "FAMILY"), ("Subfamily", "SUBFAMILY"), ("Storey", "STOREY"),
             ("Owner family", "OWNER_FAMILY"), ("Owner stage", "OWNER_STAGE"), ("Source", "SOURCE"),
             ("Dimensions", "DIMENSIONS"), ("Gross m3", "GROSS_M3"), ("Deductions m3", "DEDUCTIONS_M3"),
             ("Net m3", "NET_M3"), ("Released m3", "RELEASED_M3"), ("Conditional m3", "CONDITIONAL_M3"),
             ("Conditional range m3", "CONDITIONAL_RANGE_M3"), ("Lane", "LANE"), ("Evidence", "EVIDENCE_STATUS"),
             ("S2 census", "S2_CENSUS_STATE"), ("Unresolved reason", "UNRESOLVED_REASON")]
    f, l = table(ws, 1, ccols, conc, numeric=("GROSS_M3", "DEDUCTIONS_M3", "NET_M3", "RELEASED_M3", "CONDITIONAL_M3"),
                 widths=[30, 16, 26, 12, 16, 22, 30, 30, 11, 11, 11, 12, 12, 16, 28, 26, 16, 50])
    r = l + 2
    ws.cell(row=r, column=1, value="Released m3 (RELEASED_FROZEN_STAGE + RELEASED_S9_DELTA)").font = Font(
        name=FONT, bold=True)
    c = ws.cell(row=r, column=12, value=f'=SUMIFS(L{f}:L{l},O{f}:O{l},"RELEASED_FROZEN_STAGE")'
                                        f'+SUMIFS(L{f}:L{l},O{f}:O{l},"RELEASED_S9_DELTA")')
    c.font = Font(name=FONT, bold=True)
    c.number_format = "#,##0.000"
    ws.cell(row=r + 1, column=1, value="Conditional + indicative m3, every lane (never in a total; the Completeness sheet "
                                                   "gives the conditional point by family)").font = Font(name=FONT)
    c = ws.cell(row=r + 1, column=13, value=f"=SUM(M{f}:M{l})")
    c.font = Font(name=FONT)
    c.number_format = "#,##0.000"
    ws.auto_filter.ref = f"A1:R{l}"

    # ---- Reinforcement
    ws = wb.create_sheet("Reinforcement")
    bars = rows("03_REINFORCEMENT_BOQ_RECONCILIATION.csv")
    bcols = [("Item", "ITEM_ID"), ("Family", "FAMILY"), ("Owner stage", "OWNER_STAGE"), ("Component", "COMPONENT_ID"),
             ("Storey", "STOREY"), ("Bar role", "BAR_ROLE"), ("Dia mm", "DIA_MM"), ("Count", "COUNT"),
             ("Spacing / rate", "SPACING_OR_RATE"), ("Shape", "SHAPE"), ("Cut length m", "CUT_LENGTH_M"),
             ("Total length m", "TOTAL_LENGTH_M"), ("Laps", "LAPS"), ("Anchorage", "ANCHORAGE"),
             ("Hooks / bends", "HOOKS_BENDS"), ("kg/m", "UNIT_MASS_KG_M"), ("Original kg", "ORIGINAL_KG"),
             ("Correction kg", "CORRECTION_KG"), ("Corrections", "CORRECTIONS"), ("Authoritative kg", "AUTHORITATIVE_KG"),
             ("Not released kg", "NOT_RELEASED_KG"), ("Source authority", "SOURCE_AUTHORITY"),
             ("Stage state", "STAGE_STATE"), ("Lane", "LANE"), ("Note", "NOTE")]
    f, l = table(ws, 1, bcols, bars, numeric=("DIA_MM", "COUNT", "CUT_LENGTH_M", "TOTAL_LENGTH_M", "UNIT_MASS_KG_M",
                                              "ORIGINAL_KG", "CORRECTION_KG", "AUTHORITATIVE_KG", "NOT_RELEASED_KG"),
                 widths=[34, 14, 20, 28, 10, 30, 8, 9, 14, 30, 10, 11, 12, 12, 12, 9, 12, 12, 14, 14, 13, 30, 22, 26,
                         40])
    r = l + 2
    ws.cell(row=r, column=1, value="Released kg (sum of Authoritative kg)").font = Font(name=FONT, bold=True)
    c = ws.cell(row=r, column=20, value=f"=SUM(T{f}:T{l})")
    c.font = Font(name=FONT, bold=True)
    c.number_format = "#,##0.000"
    ws.auto_filter.ref = f"A1:Y{l}"

    # ---- Completeness
    ws = wb.create_sheet("Completeness")
    ws.cell(row=1, column=1, value="Coverage by family").font = Font(name=FONT, bold=True, size=12)
    cov = rows("13_COVERAGE_BY_FAMILY.csv")
    kcols = [("Family", "FAMILY"), ("Components", "COMPONENTS"), ("Released", "RELEASED_COMPONENTS"),
             ("Conditional (valued)", "CONDITIONAL_COMPONENTS"), ("No value", "UNQUANTIFIED_COMPONENTS"),
             ("Released m3", "RELEASED_M3"), ("Conditional point m3", "CONDITIONAL_POINT_M3"),
             ("Released kg", "RELEASED_KG"), ("Not released (modelled) kg", "NOT_RELEASED_MODELLED_KG")]
    f, l = table(ws, 2, kcols, cov, numeric=tuple(k for _, k in kcols[1:]), widths=[20, 12, 10, 14, 10, 13, 15, 13, 16])
    total_row(ws, l + 1, 1, "Total", list(range(2, 10)), f, l)
    c = ws.cell(row=l + 1, column=10, value=f"=IFERROR(C{l + 1}/B{l + 1},0)")
    c.number_format = "0.0%"
    c.font = Font(name=FONT, bold=True)
    ws.cell(row=2, column=10, value="Share released (components)")
    style_header(ws, 2, 10)
    for r in range(f, l + 1):
        c = ws.cell(row=r, column=10, value=f"=IFERROR(C{r}/B{r},0)")
        c.number_format = "0.0%"
        c.font = Font(name=FONT)
    note(ws, l + 2, "'No value' components are blocked or in conflict: their quantity is unknown, never zero. "
                    "Conditional values are never in a released total.", span=10)
    top = l + 5
    ws.cell(row=top - 1, column=1, value="Coverage by storey and family").font = Font(name=FONT, bold=True, size=12)
    fl = rows("06_FLOOR_SUMMARY.csv")
    fcols = [("Storey", "STOREY"), ("Family", "FAMILY"), ("Components", "COMPONENTS"), ("Released m3", "RELEASED_M3"),
             ("Conditional point m3", "CONDITIONAL_POINT_M3"), ("Indicative m3", "INDICATIVE_M3"),
             ("No value", "NOT_QUANTIFIED_COMPONENTS"), ("Released kg", "RELEASED_KG"),
             ("Not released (modelled) kg", "NOT_RELEASED_MODELLED_KG")]
    f2, l2 = table(ws, top, fcols, fl, numeric=("COMPONENTS", "RELEASED_M3", "CONDITIONAL_POINT_M3", "INDICATIVE_M3",
                                                "NOT_QUANTIFIED_COMPONENTS", "RELEASED_KG", "NOT_RELEASED_MODELLED_KG"))
    ws.freeze_panes = None

    # ---- Exceptions
    ws = wb.create_sheet("Exceptions")
    ws.cell(row=1, column=1, value="Missing and blocked register").font = Font(name=FONT, bold=True, size=12)
    miss = rows("05_MISSING_AND_BLOCKED_REGISTER.csv")
    mcols = [("Component", "COMPONENT_ID"), ("Family", "FAMILY"), ("Storey", "STOREY"), ("Categories", "CATEGORIES"),
             ("Concrete lane", "CONCRETE_LANE"), ("Conditional m3", "CONDITIONAL_M3"),
             ("Conditional range m3", "CONDITIONAL_RANGE_M3"), ("Released rebar kg", "RELEASED_REBAR_KG"),
             ("Reason", "REASON"), ("Owner stage", "OWNER_STAGE")]
    f, l = table(ws, 2, mcols, miss, numeric=("CONDITIONAL_M3", "RELEASED_REBAR_KG"),
                 widths=[32, 18, 12, 36, 26, 13, 16, 14, 60, 26])
    top = l + 3
    ws.cell(row=top - 1, column=1, value="Overlap and double-count audit").font = Font(name=FONT, bold=True, size=12)
    ov = rows("11_OVERLAP_AND_DOUBLE_COUNT_AUDIT.csv")
    table(ws, top, [("Check", "CHECK_ID"), ("Subject", "SUBJECT"), ("Families", "FAMILIES"), ("Rule", "RULE"),
                    ("Evidence", "EVIDENCE"), ("Result", "RESULT")], ov)
    top2 = top + len(ov) + 3
    ws.cell(row=top2 - 1, column=1, value="Conservation checks").font = Font(name=FONT, bold=True, size=12)
    table(ws, top2, [("Check", "CHECK_ID"), ("Check", "CHECK"), ("Result", "RESULT"), ("Detail", "DETAIL")],
          rows("12_CONSERVATION_CHECKS.csv"))
    ws.freeze_panes = "A3"

    # ---- RFIs / Precedence
    ws = wb.create_sheet("RFIs")
    table(ws, 1, [("RFI", "RFI_ID"), ("Priority", "PRIORITY"), ("Source stage", "SOURCE_STAGE"),
                  ("Source id", "SOURCE_ID"), ("To", "TO"), ("Family", "FAMILY"), ("Question", "QUESTION"),
                  ("Status", "STATUS")], rows("09_ENGINEER_RFI_REGISTER.csv"), widths=[10, 10, 12, 18, 16, 16, 100, 10])
    ws = wb.create_sheet("Precedence")
    table(ws, 1, [("Trade", "TRADE"), ("Family", "FAMILY"), ("Version", "VERSION"), ("Kind", "KIND"),
                  ("Quantity", "QUANTITY"), ("Unit", "UNIT"), ("Running value", "RUNNING_VALUE"),
                  ("Authoritative", "AUTHORITATIVE"), ("Superseded", "SUPERSEDED"), ("Why", "WHY")],
          rows("04_OWNERSHIP_AND_PRECEDENCE_REGISTER.csv"), numeric=("QUANTITY", "RUNNING_VALUE"),
          widths=[11, 18, 16, 14, 14, 6, 14, 13, 11, 80])
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if c.font is None or c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold if c.font else False)
    wb.properties.creator = "S9 build_workbook.py"
    import datetime
    wb.properties.created = wb.properties.modified = datetime.datetime(2026, 10, 10)   # the S9 date (stable)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
