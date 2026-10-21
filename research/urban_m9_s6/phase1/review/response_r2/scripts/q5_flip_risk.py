"""Q5 (scratch): how many numeric values differ between two output folders, and the expected number of rounding flips
if outputs were canonicalised by rounding to D decimals (a value differing by d lands on opposite sides of a rounding
boundary with probability ~ d / 10**-D). Also compares the workbook cell by cell. usage: q5_flip_risk.py DIR_A DIR_B"""
import json, sys
from pathlib import Path
A, B = Path(sys.argv[1]), Path(sys.argv[2])
diffs = []


def walk(a, b):
    if isinstance(a, float) and isinstance(b, float):
        if a != b:
            diffs.append(abs(a - b))
    elif isinstance(a, dict) and isinstance(b, dict):
        for k in a.keys() & b.keys():
            walk(a[k], b[k])
    elif isinstance(a, list) and isinstance(b, list):
        for x, y in zip(a, b):
            walk(x, y)


nvals = [0]


def count(a):
    if isinstance(a, float):
        nvals[0] += 1
    elif isinstance(a, dict):
        for v in a.values():
            count(v)
    elif isinstance(a, list):
        for v in a:
            count(v)


for f in sorted(A.glob("*.json")):
    ja, jb = json.loads(f.read_bytes()), json.loads((B / f.name).read_bytes())
    count(ja); walk(ja, jb)
negz = 0
print("float values", nvals[0], "differing", len(diffs), "max", max(diffs) if diffs else 0)
for D in (6, 9, 12):
    print(f"  expected rounding flips at {D} decimals: {sum(min(1.0, d / 10 ** -D) for d in diffs):.3g}")
import openpyxl
wa = openpyxl.load_workbook(A / "URBAN_ROOM_BOQ.xlsx", data_only=False)
wb = openpyxl.load_workbook(B / "URBAN_ROOM_BOQ.xlsx", data_only=False)
cells = cdiff = 0; mx = 0.0; other = []
for ws in wa.worksheets:
    wsb = wb[ws.title]
    for row in ws.iter_rows():
        for c in row:
            v, w = c.value, wsb[c.coordinate].value
            cells += 1
            if v != w:
                cdiff += 1
                if isinstance(v, float) and isinstance(w, float):
                    mx = max(mx, abs(v - w))
                else:
                    other.append((ws.title, c.coordinate, v, w))
print("workbook cells", cells, "differing", cdiff, "max float diff", mx, "non-float differences", other[:5])
