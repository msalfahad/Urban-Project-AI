"""T-401 development-villa run (research, read only): the lift LIFT-01 from the raft to the roof against the candidate of
record ARCH_TAKEOFF_V3.0.32. Writes ids, codes, states, m / m2 / m3 and DXF handles only (no coordinate, no name).

    .work/venv/bin/python -I research/agents_t401/t401_alsenan.py \
        --lines <run>/v3032/V3_0_32_BOQ_LINES.json --geometry <.work/private>/t401/geometry.json \
        [--freelancer <.work/private>/freelancer_concrete.xlsx] --out research/agents_t401

<geometry> = research/agents_t387/t387_geometry.py on 'ST7757 dxf.dxf' (route G's own plan records, private).
Own reads (this task, 'ST7757 dxf.dxf' sha 9f9d1179, ST7757.pdf sha 74da1523) are the VALUES below, as dimensions
relative to the ring's south-west outer corner (not coordinates):
  ring S-BW 10EB 2.20 x 2.20 / 10EA 1.80 x 1.80 (FOUNDATION), = 7C2 / 7BE (GROUND BEAMS); columns S-COL.BON 1096 SW C2
  0.20 x 0.50, 1099 NW C1 0.20 x 0.50, 1093 NE C2 0.20 x 0.50, 1090 SE C9 0.25 x 1.00 (0.05 proud of the ring);
  raft 1B2B 4.60 x 4.50 (ring offset 0.90 / 1.30 from its SW corner), depth 0.55 (FTB 2B1E, p.9);
  ST p.14: pit wall top flush with the slab-on-grade TOP (the F.F.L. line is drawn above it), pit depth dimension raft
  top -> wall top 'As Per Lift Manufactures recommendations', above it 'AS PER PLAN', 'Lift door';
  ST p.8 note 19: tie beams round the lift at 3.00 m where the storey exceeds 4.30 m; FFL GF +1.00, 1F +5.50,
  2F +9.70, roof +13.90 (T-206 chain, DRAWING-CONFIRMED); GF build-up F.F.L. -> slab top NOT PRINTED (T-206).
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import lift_structure_v1 as X  # noqa: E402

ROOT = HERE.parents[1]
OWN_T391 = ROOT / "research" / "agents_t391" / "OWN_GROUND_T391.json"

# ---- own reads (metres, relative to the ring's SW outer corner) ----
SPEC = {
    "unit": "m",
    "raft": {"rect": [-0.90, -1.30, 3.70, 3.20], "depth": 0.55},
    "ring_outer": [0.0, 0.0, 2.20, 2.20], "shaft": [0.20, 0.20, 2.00, 2.00],
    "columns": [[0.0, 0.0, 0.20, 0.50], [0.0, 1.70, 0.20, 2.20], [2.00, 1.70, 2.20, 2.20], [2.00, 0.0, 2.25, 1.00]],
    "column_marks": ["C2 (1096)", "C1 (1099)", "C2 (1093)", "C9 (1090)"],
    "wall_top_datum": X.WALL_TOP_SLAB,
    "pit_depth": {"value": None, "basis": None, "datum": X.WALL_TOP_SLAB},
    "build_up": None,
    "tie_rule": {"limit": 4.30, "level": 3.00}, "tie_section": None,
    "shaft_wall_trade": "BLOCKWORK",
}
# storeys: FFL chain T-206; ring beams = the route G bands on the ring band (geometry check below re-finds them)
STOREYS = [
    {"name": "GF", "sheet": "SH06", "ffl": 1.00, "next_ffl": 5.50, "shaft": "OPENING",
     "bands": [("B14", "BAND0022"), ("B1", "BAND0023.1"), ("B7", "BAND0025.1"), ("B1", "BAND0026")]},
    {"name": "1F", "sheet": "SH04", "ffl": 5.50, "next_ffl": 9.70, "shaft": "OPENING",
     "bands": [("B7", "BAND0006"), ("B1", "BAND0007.1"), ("B7", "BAND0011.1"), ("B1", "BAND0012")]},
    {"name": "2F", "sheet": "SH02", "ffl": 9.70, "next_ffl": 13.90, "shaft": "CAP",
     "bands": [("B2", "BAND0007.1"), ("B1", "BAND0008.2"), ("B1", "BAND0011.2"), ("B1", "BAND0012")]},
]
SHAFT_FACE = {"SH06": None, "SH04": None, "SH02": "SITE-29936e52e9c9bb32"}
GROUND_CELL = "SITE-c1089bdb4cfcdc4a"
RING_FACES_T391 = ("OWN-082", "OWN-083", "OWN-086", "OWN-094")   # the ring band pieces of T-391 (named 'BEAM' there)
FREELANCER_LIFT = {"row": "walls+columns sheet row 11", "plan_len_m": 8.00, "t_m": 0.20, "h_m": 1.50, "m3": 2.40}


def _eid(ln):
    return ln.get("element_id") or (ln.get("location") or {}).get("element_id") or ""


def _val(ln):
    v = ln.get("measured_net")
    return None if v is None else float(v)


def load_lines(p):
    d = json.loads(Path(p).read_text(encoding="utf-8"))
    return d["lines"] if isinstance(d, dict) and "lines" in d else d


def find(lines, pred):
    return [ln for ln in lines if pred(ln)]


def ring_zone(sheet_ring, pad=300.0):
    x0, y0, x1, y1 = sheet_ring
    return (x0 - pad, y0 - pad, x1 + pad, y1 + pad)


def sheet_rings(geo):
    """The ring outer box on each route G sheet, from the shaft opening / cap face (1.80 x 1.80) + 0.20."""
    out = {}
    for sid in ("SH02", "SH04", "SH06"):
        for f in geo["sheets"][sid]["slab"]["faces"]:
            if abs(f["area_m2"] - 3.24) < 0.01 and f["class"] in ("OPENING", "PANEL"):
                xs = [p[0] for p in f["polygon"]]
                ys = [p[1] for p in f["polygon"]]
                if abs(max(xs) - min(xs) - 1800) < 2 and abs(max(ys) - min(ys) - 1800) < 2:
                    if f["class"] == "PANEL" and sid != "SH02":
                        continue
                    out[sid] = (min(xs) - 200, min(ys) - 200, max(xs) + 200, max(ys) + 200)
    step = out["SH06"][0] - out["SH04"][0]
    out["SH08"] = (out["SH06"][0] + step, out["SH06"][1], out["SH06"][2] + step, out["SH06"][3])
    return out


def overlap_census(geo, lines, rings):
    from shapely.geometry import Polygon, box
    by_eid = {}
    for ln in lines:
        by_eid.setdefault(_eid(ln), []).append(ln)
    rows, faces_rows = [], []
    for sid, storey, slab_prefix in (("SH08", "GROUND", "U-C:GROUND_SLAB:GROUND:"), ("SH06", "GF", "P-C:SLABS:GF:"),
                                     ("SH04", "1F", "P-C:SLABS:1F:"), ("SH02", "2F", "P-C:SLABS:2F:")):
        sh = geo["sheets"][sid]
        R = rings[sid]
        outer = box(*R)
        inner = box(R[0] + 200, R[1] + 200, R[2] - 200, R[3] - 200)
        zone = box(*ring_zone(R))
        claims = []
        for b in sh.get("beams", []):
            P = Polygon(b["polygon"]).buffer(0)
            if P.intersection(zone).area > 1e2:
                eid = f"P-C:BEAMS:{storey}:{b['band_id']}" if storey != "GROUND" else f"P-C:GROUND_BEAMS:GROUND:{b['band_id']}"
                claims.append(("BEAM", eid, P))
        for cb in sh.get("column_boxes", []):
            P = box(*cb)
            if P.intersection(zone).area > 1e2:
                claims.append(("COLUMN", "COLUMN_BOX", P))
        for f in sh["slab"]["faces"]:
            if len(f["polygon"]) < 3:
                continue
            P = Polygon(f["polygon"]).buffer(0)
            if P.intersection(zone).area <= 1e2:
                continue
            eid = slab_prefix + f["face_id"]
            fam = {"COLUMN": "COLUMN_FACE", "BEAM": "BEAM_FACE", "OPENING": "OPENING"}.get(f["class"], "SLAB")
            if f["region"] == "JUNCTION_GAP":
                fam = "JUNCTION_GAP"
            claims.append((fam, eid, P))
            faces_rows.append({"level": storey, "face": f["face_id"], "class": f["class"], "region": f["region"],
                               "area_m2": round(f["area_m2"], 4),
                               "on_shaft_m2": round(P.intersection(inner).area / 1e6, 4),
                               "on_ring_band_m2": round((P.intersection(outer).area - P.intersection(inner).area) / 1e6, 4),
                               "lines": " ".join(ln["line_id"] for ln in by_eid.get(eid, [])) or "NO_LINE"})
        for i in range(len(claims)):
            for j in range(i + 1, len(claims)):
                fa, ea, Pa = claims[i]
                fb, eb, Pb = claims[j]
                if fa in ("COLUMN_FACE", "BEAM_FACE", "OPENING") or fb in ("COLUMN_FACE", "BEAM_FACE", "OPENING"):
                    continue            # faces that are not billed by the slab line (the member's own lines bill them)
                sh_m2 = Pa.intersection(Pb).intersection(zone).area / 1e6
                if sh_m2 <= 1e-4:
                    continue
                if fa == fb == "COLUMN":
                    continue
                la, lb = by_eid.get(ea, []), by_eid.get(eb, [])
                if "COLUMN" in (fa, fb):        # the column lines (column plan / ledger) bill every lift column
                    la = la or [{"line_id": "COLUMN_LINES", "measured_net": 1.0}]
                    lb = lb or [{"line_id": "COLUMN_LINES", "measured_net": 1.0}]
                va = any(_val(x) is not None for x in la)
                vb = any(_val(x) is not None for x in lb)
                kind = "DOUBLE" if (va and vb) else ("HELD_OVERLAP" if (la and lb) else "NOT_BILLED")
                rows.append({"level": storey, "a": f"{fa}:{ea}", "b": f"{fb}:{eb}",
                             "a_lines": " ".join(x["line_id"] for x in la) or "NO_LINE",
                             "b_lines": " ".join(x["line_id"] for x in lb) or "NO_LINE",
                             "shared_m2": round(sh_m2, 4), "kind": kind})
    return rows, faces_rows


def freelancer_rows(path):
    if not path:
        return None
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb:
        for row in ws.iter_rows():
            vals = [c.value for c in row]
            txt = " ".join(str(v) for v in vals if isinstance(v, str))
            if "المصعد" in txt:        # the word for 'the lift'
                nums = [v for v in vals if isinstance(v, (int, float))]
                out.append({"sheet_index": wb.sheetnames.index(ws.title), "row": row[0].row, "numbers": nums})
    return out


def run(lines_path, geo_path, freelancer=None):
    lines = load_lines(lines_path)
    geo = json.loads(Path(geo_path).read_text(encoding="utf-8"))
    own = json.loads(OWN_T391.read_text(encoding="utf-8"))
    rings = sheet_rings(geo)
    by_eid = {}
    for ln in lines:
        by_eid.setdefault(_eid(ln), []).append(ln)

    def one(pred, what):
        hit = [ln for ln in lines if pred(ln)]
        return hit

    ff = one(lambda ln: _eid(ln).startswith("P-C:FTG:SH10:FF@"), "raft")
    lift_lines = one(lambda ln: "LIFT" in (_eid(ln) + " " + (ln.get("item_code") or "")).upper()
                     and (ln.get("item_code") or "").startswith("STR.CONC"), "lift")
    dependent = one(lambda ln: "LIFT" in json.dumps({k: ln.get(k) for k in ("item_code", "location", "physical_element")}).upper()
                    and not (ln.get("item_code") or "").startswith("STR.CONC"), "dependent")
    # columns of the lift: route G members whose outline lies on the ring (foundation SH10 necks; column plan SH12)
    R10 = rings["SH08"]
    step = rings["SH06"][0] - rings["SH04"][0]
    R10 = (R10[0] + step, R10[1], R10[2] + step, R10[3])
    R12 = (R10[0] + step, R10[1], R10[2] + step, R10[3])
    col_members = {}
    for sid, R in (("SH10", R10), ("SH12", R12)):
        for c in geo["sheets"][sid]["columns"]:
            b = c.get("box")
            if b and not (b[2] < R[0] - 60 or b[0] > R[2] + 60 or b[3] < R[1] - 60 or b[1] > R[3] + 60):
                col_members.setdefault(sid, []).append((c["member_id"], c["mark"], b))
    unbound_marks = [c["member_id"] for c in geo["sheets"]["SH12"]["columns"]
                     if not c.get("box") and c["mark"] in ("C1", "C2", "C9")
                     and R12[0] - 1500 < c["mark_at"][0] < R12[2] + 1500 and R12[1] - 1500 < c["mark_at"][1] < R12[3] + 1500]
    col_rows = []
    for sid, ms in sorted(col_members.items()):
        for mid, mark, b in ms:
            ls = [ln for ln in lines if mid + ":" in _eid(ln) or _eid(ln).endswith(mid) or (":" + mid.split(":", 1)[1]) in _eid(ln)]
            col_rows.append({"sheet": sid, "member": mid, "mark": mark,
                             "plan_m2": round((b[2] - b[0]) * (b[3] - b[1]) / 1e6, 4),
                             "lines": [(ln["line_id"], ln.get("item_code"), _val(ln), ln.get("evidence_state")) for ln in ls]})
    for mid in unbound_marks:
        ls = [ln for ln in lines if (":" + mid.split(":", 1)[1]) in _eid(ln)]
        col_rows.append({"sheet": "SH12", "member": mid, "mark": mid.split(":")[2].split("@")[0], "plan_m2": None,
                         "lines": [(ln["line_id"], ln.get("item_code"), _val(ln), ln.get("evidence_state")) for ln in ls],
                         "note": "mark without a bound outline on the column plan (outline_binding NEAREST found none)"})

    bind = {"RAFT": ff[0]["line_id"] if ff else None, "PIT_WALL": lift_lines[0]["line_id"] if lift_lines else None}
    storeys = []
    for st in STOREYS:
        rb = []
        for k, (mark, band) in enumerate(st["bands"]):
            eid = f"P-C:BEAMS:{st['name']}:{band}"
            ls = by_eid.get(eid, [])
            rb.append({"mark": mark, "length": None, "depth": None, "band": band})
            if ls:
                bind[f"RING_BEAM:{st['name']}:{mark}:{k + 1}"] = ls[0]["line_id"]
        b_geo = {b["band_id"]: b for b in geo["sheets"][st["sheet"]]["beams"]}
        for r in rb:
            g = b_geo[r["band"]]
            r["length"], r["depth"], r["width"] = round(g["length_m"], 4), g.get("depth_m"), round(g["width_m"], 4)
            if g.get("mark") != r["mark"]:
                raise X.Refused("RING_BEAM_MARK_CHANGED", r["band"])
        if st["shaft"] == "CAP":
            ls = by_eid.get("P-C:SLABS:2F:" + SHAFT_FACE[st["sheet"]], [])
            if ls:
                bind[f"SHAFT_SLAB:{st['name']}"] = ls[0]["line_id"]
        storeys.append({"name": st["name"], "ffl": st["ffl"], "next_ffl": st["next_ffl"], "shaft": st["shaft"],
                        "ring_beams": rb})
    spec = dict(SPEC, storeys=storeys, bind=bind)
    led = X.build(spec)

    # ground level: the held cell and the T-391 ring pieces
    faces = {f["face"]: f for f in own["faces"]}
    cell = next(b for b in own["engine_binding"] if b["engine_cell"] == GROUND_CELL)
    ring_pieces = {k: faces[k]["area_m2"] for k in RING_FACES_T391}
    in_cell = [k for k in RING_FACES_T391 if k in cell["own_faces"]]
    cell_lines = by_eid.get("U-C:GROUND_SLAB:GROUND:" + GROUND_CELL, [])

    overlaps, face_rows = overlap_census(geo, lines, rings)

    # values of the bound lines
    line_by_id = {ln["line_id"]: ln for ln in lines}
    for c in led["components"]:
        ln = line_by_id.get(c["line_id"]) if c["line_id"] else None
        c["line_value"] = _val(ln) if ln else None
        c["line_state"] = (ln.get("evidence_state") if ln else "") or ""
        c["line_code"] = (ln.get("item_code") if ln else "") or ""
    ring_beam_m3 = {st["name"]: round(sum((c["line_value"] or 0.0) for c in led["components"]
                                          if c["id"].startswith(f"RING_BEAM:{st['name']}:")), 6) for st in STOREYS}
    hand_ring_beam_m3 = {st["name"]: round(sum(r["width"] * r["length"] * r["depth"] for r in s["ring_beams"]), 6)
                         for st, s in zip(STOREYS, storeys)}

    fl = freelancer_rows(freelancer)
    net = led["ring"]["net_m2"]
    gross = led["ring"]["gross_m2"]
    fr = FREELANCER_LIFT
    recon = {
        "freelancer_m3": fr["m3"], "freelancer_plan_m2": round(fr["plan_len_m"] * fr["t_m"], 6),
        "drawn_gross_ring_m2": gross, "plan_difference_m2": round(fr["plan_len_m"] * fr["t_m"] - gross, 6),
        "columns_in_his_ring_m3": round(-(gross - net) * fr["h_m"], 6),
        "net_ring_at_his_height_m3": round(-net * fr["h_m"], 6),
        "engine_V3_0_32_m3": None, "residual_m3": -fr["m3"],
        "closure_m3": round(-fr["m3"] - (-(gross - net) * fr["h_m"] - net * fr["h_m"]), 9),
    }
    summary = {
        "task": "T-401", "method": X.METHOD_ID, "candidate": "ARCH_TAKEOFF_V3.0.32",
        "ring": led["ring"], "wall_height_m": led["wall_height_m"], "wall_missing": led["wall_missing"],
        "rings_found_on_sheets": sorted(rings),
        "ring_beams_line_m3": ring_beam_m3, "ring_beams_hand_m3": hand_ring_beam_m3,
        "ground": {"cell_line_ids": [ln["line_id"] for ln in cell_lines],
                   "cell_values": [_val(ln) for ln in cell_lines],
                   "cell_gross_m2": cell["gross_m2"], "cell_net_m2": cell["net_m2"],
                   "shaft_face_m2": faces["OWN-020"]["area_m2"], "ring_pieces_m2": ring_pieces,
                   "ring_pieces_in_cell": in_cell,
                   "ring_pieces_in_cell_m2": round(sum(ring_pieces[k] for k in in_cell), 4),
                   "ring_net_from_pieces_m2": round(sum(ring_pieces.values()), 4)},
        "freelancer_lift_rows": fl, "freelancer_reconciliation": recon,
        "overlaps": {"DOUBLE": [o for o in overlaps if o["kind"] == "DOUBLE"],
                     "HELD_OVERLAP": [o for o in overlaps if o["kind"] == "HELD_OVERLAP"],
                     "NOT_BILLED": [o for o in overlaps if o["kind"] == "NOT_BILLED"]},
        "lift_lines_in_candidate": [ln["line_id"] for ln in lift_lines],
        "dependent_lift_lines": [{"line_id": ln["line_id"], "item_code": ln.get("item_code"), "unit": ln.get("unit"),
                                  "value": _val(ln), "state": ln.get("evidence_state"),
                                  "element": _eid(ln)} for ln in dependent],
        "raft_line": [{"line_id": ln["line_id"], "value": _val(ln), "state": ln.get("evidence_state")} for ln in ff],
    }
    return led, col_rows, overlaps, face_rows, summary


def write(out, led, col_rows, overlaps, face_rows, summary):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    with (out / "COMPONENTS_T401.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["component_id", "component", "owner_trade", "unit", "value", "status", "candidate_line",
                    "line_code", "line_value", "line_state", "missing", "note"])
        for c in led["components"]:
            w.writerow([c["id"], c["component"], c["owner_trade"], c["unit"], c["value"], c["status"],
                        c["line_id"] or "MISSING" if c["status"] == X.MISSING else (c["line_id"] or ""),
                        c["line_code"], c["line_value"], c["line_state"], "; ".join(c["missing"]), c["note"]])
    with (out / "COLUMNS_T401.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["sheet", "member", "mark", "plan_m2", "line_id", "item_code", "value_m3", "state", "note"])
        for r in col_rows:
            for lid, code, v, st in r["lines"] or [("NO_LINE", "", None, "")]:
                w.writerow([r["sheet"], r["member"].split("@")[0], r["mark"], r["plan_m2"], lid, code, v, st,
                            r.get("note", "")])
    with (out / "OVERLAP_T401.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["level", "a", "b", "a_lines", "b_lines", "shared_m2", "kind"])
        for o in overlaps:
            w.writerow([o["level"], o["a"], o["b"], o["a_lines"], o["b_lines"], o["shared_m2"], o["kind"]])
    with (out / "FACES_T401.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["level", "face", "class", "region", "area_m2", "on_shaft_m2", "on_ring_band_m2", "lines"])
        for f in face_rows:
            w.writerow([f[k] for k in ("level", "face", "class", "region", "area_m2", "on_shaft_m2",
                                       "on_ring_band_m2", "lines")])
    (out / "summary_t401.json").write_text(json.dumps(summary, indent=1, sort_keys=True, ensure_ascii=False) + "\n",
                                           encoding="utf-8")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--lines", required=True)
    ap.add_argument("--geometry", required=True)
    ap.add_argument("--freelancer")
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    res = run(a.lines, a.geometry, a.freelancer)
    write(a.out, *res)
    s = res[-1]
    print(json.dumps({k: s[k] for k in ("ring", "wall_height_m", "ring_beams_line_m3", "ring_beams_hand_m3", "ground",
                                        "freelancer_reconciliation", "lift_lines_in_candidate",
                                        "dependent_lift_lines", "raft_line")}, indent=1))
    print("overlaps:", {k: len(v) for k, v in s["overlaps"].items()})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
