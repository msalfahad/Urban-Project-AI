"""ALSENAN_COVERAGE_RECOVERY_REVIEW.xlsx - review workbook of the coverage-recovery round (client data, gitignored).

Reads the frozen registers of this folder (hash-checked against INDEX.json) and writes 12 sheets. Totals and
differences are live formulas; the recalc script (LibreOffice) is run when RECALC_SCRIPT is set.
"""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "review" / "ALSENAN_COVERAGE_RECOVERY_REVIEW.xlsx"

RECOMMENDATIONS = [
    ("P0", "Scenario-layer publication for every trade (verified / lower bound / best / low / high / unquantified)",
     "removes the systemic BLOCKED = 0 undercount", "low", "all", "M"),
    ("P0", "Column concrete from column_concrete_geometry (net of slab, joints included)",
     "95/95 occurrences quantified, rebar blockers decoupled", "low", "columns", "S"),
    ("P0", "Ground slab from ground_slab_recovery cells; label scope as a question", "+101.7 m2 recovered",
     "low-medium (low = official)", "ground slab", "S"),
    ("P0", "Send the four consultant questions (exterior GB depth, slab-note scope, founding level, GF T16 void)",
     "turns the largest provisional bands into source facts", "none", "GB / slab / columns / slabs", "S"),
    ("P1", "Measure ambiguous wall bands as CANDIDATE + wall_band_reconciliation", "explains the 200 mm gap",
     "medium", "blockwork", "M"),
    ("P1", "Physical wall faces before finish semantics", "plaster / paint scope visible", "low", "finishes", "M"),
    ("P1", "Beam binding ladder on the 8 unbound tags; confirm CB8 spans", "closes beam residue", "medium", "beams", "M"),
    ("P1", "Store band / cell coordinates in frozen registers", "per-segment URBAN_MISSED + Method C", "low",
     "walls / slabs", "S"),
    ("P2", "Method C (room adjacency); second footing route", "multi-route confidence", "low", "walls / footings", "M"),
    ("P2", "Live read-only CAD oracle re-runs", "fresher post-freeze checks", "low", "all", "M"),
]


def J(name):
    return json.loads((HERE / name).read_text(encoding="utf-8"))


def check():
    idx = J("INDEX.json")
    for k, h in idx["outputs"].items():
        if hashlib.sha256((HERE / k).read_bytes()).hexdigest() != h:
            raise SystemExit(f"{k} does not match INDEX.json")
    return idx


def build():
    idx = check()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    wb = Workbook()
    hf, fill = Font(name="Arial", bold=True, color="FFFFFF", size=10), PatternFill("solid", fgColor="1F3864")
    body = Font(name="Arial", size=9)

    def sheet(name, headers, rows, first=None):
        ws = first or wb.create_sheet(name)
        ws.title = name
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font, c.fill, c.alignment = hf, fill, Alignment(wrap_text=True, vertical="top")
        for i, r in enumerate(rows, 2):
            for j, v in enumerate(r, 1):
                if isinstance(v, (list, dict)):
                    v = json.dumps(v, ensure_ascii=False)
                c = ws.cell(row=i, column=j, value=v)
                if isinstance(v, str) and v.startswith("="):
                    c.data_type = "s" if not getattr(v, "_formula", False) else "f"
                c.font = body
        for j, h in enumerate(headers, 1):
            ws.column_dimensions[L(j)].width = max(11, min(48, len(h) + 4))
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 42
        return ws

    def formula(ws, cell, f):
        ws[cell] = f
        ws[cell].font = Font(name="Arial", size=9, bold=True)

    Q, D = J("QUANTITY_SCENARIOS.json"), J("QUANTITY_COVERAGE_DASHBOARD.json")
    # 00 dashboard
    hd = ["TRADE", "UNIT", "OFFICIAL BEFORE", "VERIFIED (OFFICIAL AFTER)", "LOWER BOUND", "BEST PROVISIONAL",
          "LOW SCENARIO", "HIGH SCENARIO", "UNQUANTIFIED", "COVERAGE %", "RELEASE", "TOP BLOCKER", "HOW TO FIX",
          "BEST - BEFORE"]
    rows = [[r["trade"], r["unit"], r["official_before"], r["verified"], r["lower_bound"], r["best_provisional"],
             r["low"], r["high"] if r["high"] is not None else "UNBOUNDED", r["unquantified"], r["coverage_pct"],
             r["release"], r["top_blocker"], r["how_to_fix"], None] for r in D["rows"]]
    ws = sheet("00_Dashboard", hd, rows, first=wb.active)
    for i in range(2, len(rows) + 2):
        formula(ws, f"N{i}", f"=F{i}-C{i}")
    n = len(rows) + 3
    ws[f"A{n}"] = ("Provisional / scenario values show the likely full scope; they are not procurement values. "
                   "Official = verified source-derived quantity only.")
    # 01 coverage
    C = J("COVERAGE_METRICS.json")
    sheet("01_Trade_Coverage", ["TRADE", "OCCURRENCES", "POPULATION %", "GEOMETRY %", "SEMANTIC %", "SOURCE AUTHORITY %",
                                "OFFICIAL % OF BEST"],
          [[t["trade"], t["occurrences"], t["population_pct"], t["geometry_pct"], t["semantic_pct"],
            t["source_authority_pct"], t["quantity_official_pct_of_best"]] for t in C["trades"]] +
          [[""], ["ANOMALY CHECK", "STATE", "COVERED", "EXPECTED", "RATIO", "WHAT"]] +
          [[a["check_id"], a["state"], a.get("covered", a.get("emitted")), a.get("expected", a.get("occurrences")),
            a.get("ratio"), a.get("what")] for a in C["anomalies"]["findings"]])
    # 02 ground beams
    gb = Q["ground_beams"]
    ws = sheet("02_Ground_Beams", ["SPAN", "KIND", "LENGTH m", "B m", "DEPTH m", "DEPTH LEVEL", "DEPTH AUTHORITY",
                                   "TERMINAL", "BEST m3", "LOW m3", "HIGH m3"],
               [[s["span"], s["kind"], s["length_m"], s["B_m"], s["depth_m"], s["depth_level"], s["depth_authority"],
                 s["terminal"], s["part"]["best"], s["part"]["low"], s["part"]["high"]] for s in gb["spans"]])
    k = len(gb["spans"]) + 2
    ws[f"A{k}"] = "TOTAL (spans; boundary beam line C-BWALL separate)"
    for col in ("C", "I", "J", "K"):
        formula(ws, f"{col}{k}", f"=SUM({col}2:{col}{k - 1})")
    # 03 ground slab
    gs = Q["ground_slab"]
    ws = sheet("03_Ground_Slab", ["CELL", "FOOTPRINT", "PART", "AREA m2", "EFF WIDTH mm", "LABELS", "ROLE", "WHY"],
               [[c["cell_id"], c["footprint_id"], c["part"], c["area_m2"], c["eff_width_mm"], c["labels"], c["role"],
                 c["why"]] for c in gs["cells"]])
    k = len(gs["cells"]) + 2
    ws[f"A{k}"] = "OFFICIAL (labelled / released)"
    formula(ws, f"D{k}", f'=SUMIF(G2:G{k - 1},"LABELLED_OFFICIAL",D2:D{k - 1})')
    ws[f"A{k + 1}"] = "CANDIDATE (unlabelled cells)"
    formula(ws, f"D{k + 1}", f'=SUMIF(G2:G{k - 1},"UNLABELLED_CANDIDATE",D2:D{k - 1})')
    ws[f"A{k + 2}"] = f"thickness: {gs['thickness']}"
    # 04 columns
    col = Q["columns"]
    sheet("04_Columns", ["FLOOR", "OCCURRENCES", "GROSS INTERVAL m3", "VERIFIED m3", "BEST m3", "LOW m3", "HIGH m3",
                         "TERMINAL STATES"],
          [[fl, v["occurrences"], v["gross_interval_m3"], v["column_plus_joint"]["VERIFIED_QUANTITY"],
            v["column_plus_joint"]["BEST_PROVISIONAL_QUANTITY"], v["column_plus_joint"]["LOW_SCENARIO"],
            v["column_plus_joint"]["HIGH_SCENARIO"], v["terminal_states"]] for fl, v in col["per_floor"].items()] +
          [[""], ["basis", col["basis"]]])
    # 05 beams
    bm = Q["beams"]
    sheet("05_Beams", ["OBJECT", "FLOOR", "TYPE", "WHY (V3a)", "LENGTH m", "TERMINAL", "BEST m3", "LOW m3", "HIGH m3",
                       "BINDING LEVEL"],
          [[o["occurrence_id"], o["floor"], o["type"], o["why_v3a"], o["length_m"], o["terminal"], o["part"]["best"],
            o["part"]["low"], o["part"]["high"], o["binding_level"]] for o in bm["residue_objects"]] +
          [[""], ["DUPLICATES REMOVED"]] + [[d["occurrence_id"], "", "", d["why"], "", "duplicate of " + d["duplicate_of"]]
                                            for d in bm["duplicates_removed"]])
    # 06 slabs
    sl = Q["slabs"]
    rows = []
    for fl, v in sl["sheets"].items():
        rows.append([fl, "GROSS", v["gross_m2"], "", ""])
        rows.append([fl, "PLATE", v["plate_m2"], "", ""])
        for o in v["openings"]:
            rows.append([fl, o["opening_id"], o["area_m2"], o["state"], o["evidence"]])
    sheet("06_Slabs_Openings", ["SHEET", "ITEM / OPENING ID", "AREA m2", "STATE", "EVIDENCE"], rows)
    # 07 walls
    wl = Q["walls"]
    sheet("07_Walls", ["FLOOR|WIDTH", "A ESTABLISHED m", "A AMBIGUOUS m", "A COLUMN OVERLAP m", "B PAIRED WALL m",
                       "B OPENING SPAN m", "B COLUMN OVERLAP m", "B DUPLICATE m", "AMBIGUOUS NOT MEASURED m",
                       "MISSED / UNRESOLVED m", "METHOD C"],
          [[k, v["method_a"]["established_m"], v["method_a"]["ambiguous_m"], v["method_a"]["column_overlap_m"],
            v["method_b"].get("PAIRED_WALL", 0), v["method_b"].get("OPENING_SPAN", 0),
            v["method_b"].get("COLUMN_OVERLAP_POLICY", 0), v["method_b"].get("DUPLICATE_FACE", 0),
            v["explain"]["URBAN_AMBIGUOUS_NOT_MEASURED"], v["explain"]["URBAN_MISSED_OR_UNRESOLVED"], "NOT_RUN"]
           for k, v in wl["reconciliation"].items()] +
          [[""], ["PHYSICAL WALL FACES", json.dumps(wl["physical_faces"]["physical_area"]),
                  json.dumps(wl["physical_faces"]["by_floor_role"])]])
    # 08 remediation
    R = J("REMEDIATION_ATTEMPTS.json")["records"]
    sheet("08_Blocked_Remediation", ["FLAG", "KIND", "ELEMENT", "UNKNOWN FACT", "CURRENT", "PROVISIONAL", "LOW", "HIGH",
                                     "RESULT", "SUCCESSFUL METHOD", "ATTEMPTS", "NEXT METHOD", "WHERE TO LOOK",
                                     "SEARCH TEXT", "CONSULTANT QUESTION"],
          [[r["flag_id"], r["kind"], r["element"], r["unknown_fact"], r["current_quantity"], r["provisional_quantity"],
            r["low_scenario"], r["high_scenario"], r["result"], r["successful_method"],
            " > ".join(f"{a['level']}:{a['result']}" for a in r["attempts"]), r["next_automated_method"],
            r["where_to_look"], r["search_text"], r["consultant_question"]] for r in R])
    # 09 post-freeze
    P = J("POST_FREEZE_MCP_COMPARISON.json")["difference_register"]
    ws = sheet("09_MCP_Post_Freeze", ["TRADE", "UNIT", "URBAN OFFICIAL", "URBAN BEST", "URBAN LOW", "URBAN HIGH",
                                      "U-C4N", "CHRISTIANNP", "FREELANCER", "BEST - U-C4N", "BEST - CHRISTIANNP",
                                      "BASIS", "REASON CLASS", "URBAN ISSUE?", "DONOR ISSUE?", "SOURCE REVIEW?", "NOTE"],
               [[r["trade"], r["unit"], r["urban_official"], r["urban_best_provisional"], r["urban_low"], r["urban_high"],
                 r["donors"].get("UC4N"), r["donors"].get("CHRISTIANNP"), r["donors"].get("FREELANCER"), None, None,
                 r["basis_match"], r["reason_class"], r["likely_urban_issue"], r["likely_donor_issue"],
                 r["needs_source_review"], r["note"]] for r in P])
    for i, r in enumerate(P, 2):
        if r["donors"].get("UC4N") is not None:
            formula(ws, f"J{i}", f"=D{i}-G{i}")
        if r["donors"].get("CHRISTIANNP") is not None:
            formula(ws, f"K{i}", f"=D{i}-H{i}")
    # 10 root causes
    RC = J("ROOT_CAUSE_REGISTER.json")["rows"]
    keys = [k for k in RC[0] if k != "difference"]
    sheet("10_Root_Causes", ["DIFFERENCE"] + [k.upper() for k in keys], [[r["difference"]] + [r[k] for k in keys] for r in RC])
    # 11 recommendations
    sheet("11_Recommendations", ["RANK", "RECOMMENDATION", "BENEFIT", "RISK", "TRADE", "EFFORT"],
          [list(r) for r in RECOMMENDATIONS])
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    rs = os.environ.get("RECALC_SCRIPT")
    if rs:
        res = subprocess.run([sys.executable, rs, str(OUT), "90"], capture_output=True, text=True)
        out = json.loads(res.stdout[res.stdout.index("{"):])
        if out.get("status") != "success":
            raise SystemExit(f"recalc failed: {out}")
        print("recalculated", out.get("total_formulas"), "formulas, 0 errors")
    man = {"workbook": OUT.name, "sha256": hashlib.sha256(OUT.read_bytes()).hexdigest(),
           "registers_consumed": idx["outputs"]}
    (OUT.parent / "REVIEW_MANIFEST.json").write_text(json.dumps(man, indent=1, sort_keys=True) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build()
