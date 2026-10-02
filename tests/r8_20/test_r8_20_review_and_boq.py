"""R8.20 synthetic: human closure review bound to the closure digest; blocker classification for BOQ status; the XLSX
view (no formula, every cell = its report row, statuses visible, tamper refused); controlled-export identity; the
physical-rationale fact authorises no domain."""

from __future__ import annotations

import datetime

from openpyxl import load_workbook

from engine.source import boq_evidence as BE, boq_report as BR, boq_xlsx as BX, closure_release as CR
from engine.source import closure_review as RV, owner_facts as OF, source_anchor as SA

REC = {"closure_id": "TC-1", "geometry": [0.0, 0.0, 0.0, 20.0], "source_evidence_ids": ["A|H1||SEGMENT|0"]}
DIG = RV.record_digest(REC)


def review(**kw):
    return dict({"review_id": "R-1", "closure": "TC-1", "packet_digest": DIG, "decision": RV.ACCEPT,
                 "reviewer_role": "PROJECT_OWNER"}, **kw)


# ------------------------------------------------------------------------------- human review
def test_an_owner_accept_bound_to_the_digest_satisfies_human_review_only():
    s = RV.review_state(review(), closure="TC-1", current_digest=DIG)
    assert s["state"] == RV.ACCEPTED and s["satisfies_human_review"] is True
    ev = {"POLICY_FROZEN": True, "CROSS_ROUTE_AGREEMENT": None, "OWNER_OR_SOURCE_CORROBORATION": True,
          "SOURCE_ANCHOR": False, "HUMAN_REVIEW": s["satisfies_human_review"]}
    e = CR.evaluate("AUTHORISED_FOR_SHADOW", ev)
    assert e["level"] == "AUTHORISED_FOR_SHADOW" and e["missing_for_reviewed"] == ["CROSS_ROUTE_AGREEMENT",
                                                                                    "SOURCE_ANCHOR"]
    assert e["releases_anything"] is False


def test_a_changed_closure_another_closure_or_a_machine_review_never_counts():
    moved = RV.record_digest(dict(REC, geometry=[0.0, 0.0, 0.0, 20.5]))
    assert RV.review_state(review(), closure="TC-1", current_digest=moved)["state"] == RV.STALE
    assert RV.review_state(review(), closure="TC-2", current_digest=DIG)["state"] == RV.WRONG_CLOSURE
    for role in ("ENGINE", "CLAUDE", None):
        assert RV.review_state(review(reviewer_role=role), closure="TC-1",
                               current_digest=DIG)["satisfies_human_review"] is None
    assert RV.review_state(review(decision=RV.REJECT), closure="TC-1", current_digest=DIG)["satisfies_human_review"] \
        is False
    assert RV.review_state(None, closure="TC-1", current_digest=DIG)["state"] == RV.ABSENT


def test_all_five_conditions_are_needed_for_the_reviewed_level_and_nothing_releases():
    full = {c: True for c in CR.CONDITIONS}
    assert CR.evaluate("AUTHORISED_FOR_SHADOW", full)["level"] == CR.REVIEWED
    assert "RELEASED" not in CR.LEVELS and "a release" in CR.policy_record()["never"]


# ------------------------------------------------------------------------------- BOQ evidence
def test_quantity_affecting_and_unknown_blockers_keep_a_row_out_of_complete():
    assert BE.classify("SOURCE_ANCHOR: DXF anchored") == BE.RELEASE_ONLY
    assert BE.classify("SHADOW_ONLY: no approval") == BE.RELEASE_ONLY
    assert BE.classify("OBJECT_FOOTPRINT_IMPLICIT: convention") == BE.QUANTITY_AFFECTING
    assert BE.classify("SOMETHING_NEW: x") == BE.QUANTITY_AFFECTING
    un = BE.unresolved(["SOURCE_ANCHOR: a", "OBJECT_FOOTPRINT_IMPLICIT: b"])
    assert [u["blocker"] for u in un] == ["OBJECT_FOOTPRINT_IMPLICIT: b"]
    assert BR.status({"complete": 11.685, "authorised_subtotal": 11.685, "unresolved": un})[0] == BR.SUBTOTAL
    assert BR.status({"complete": 11.685, "authorised_subtotal": 11.685,
                      "unresolved": BE.unresolved(["SOURCE_ANCHOR: a"])})[0] == BR.COMPLETE


# ------------------------------------------------------------------------------- XLSX
TRACE = {"run_id": "RUN", "source_revision": "REV", "trade_row": "T", "sites": ["S1"], "surface_ids": None,
         "authority_digest": "d", "rule_ids": ["R@v1"], "owner_facts": [], "release_state": "SHADOW"}
ITEMS = {k: {"item_code": f"IT-{k}", "trade": "PAINT", "description_ar": "دهانات", "description_en": "Paint"}
         for k in ("A", "B", "C")}
EVS = [dict(TRACE, row_id="A", unit="m2", complete=307.361518, authorised_subtotal=307.361518, unresolved=[],
            floor="F2", source="DXF", blockers=["SHADOW_ONLY: x"],
            per_room={"S1": {"zones": ["BED"], "qty": 72.328302}, "S2": {"zones": ["HALL"], "qty": 235.033216}}),
       dict(TRACE, row_id="B", unit="m2", complete=None, authorised_subtotal=11.685, unresolved=[{"blocker": "X"}],
            floor="F2", source="DXF", blockers=["OBJECT_FOOTPRINT_IMPLICIT: y"], per_room=None),
       dict(TRACE, row_id="C", unit="lm", complete=None, authorised_subtotal=None, unresolved=[{"blocker": "Z"}],
            floor="F2", source="DXF", blockers=["Z"], per_room=None)]


def rep():
    return BR.build(EVS, ITEMS, run={"run_id": "RUN", "source_revision": "REV"})


def test_the_workbook_is_a_view_every_quantity_is_its_report_row_and_validates(tmp_path):
    p = tmp_path / "boq.xlsx"
    w = BX.write(rep(), p, created=datetime.datetime(2026, 10, 2))
    v = BX.validate(p, rep())
    assert v["state"] == "PASS" and v["summary_lines"] == 3 and v["breakdown_lines"] == 2 and v["trace_lines"] == 5
    assert w["content_digest"] == v["content_digest"] and not v["formulas"]
    ws = load_workbook(str(p))["01_BOQ_SUMMARY"]
    assert ws["A1"].value == BX.BANNER
    qty = {ws.cell(i, 1).value: ws.cell(i, 7).value for i in range(3, ws.max_row + 1)}
    assert qty == {"IT-A": 307.361518, "IT-B": 11.685, "IT-C": None}
    st = {ws.cell(i, 1).value: ws.cell(i, 9).value for i in range(3, ws.max_row + 1)}
    assert st == {"IT-A": BR.COMPLETE, "IT-B": BR.SUBTOTAL, "IT-C": BR.BLOCKED}
    assert all(ws.cell(i, 13).value == "NO - SHADOW" for i in range(3, ws.max_row + 1))


def test_an_edited_cell_or_a_formula_in_the_workbook_is_refused(tmp_path):
    p = tmp_path / "boq.xlsx"
    BX.write(rep(), p, created=datetime.datetime(2026, 10, 2))
    wb = load_workbook(str(p))
    wb["01_BOQ_SUMMARY"].cell(3, 7).value = 307.4
    wb.save(str(p))
    assert BX.validate(p, rep())["state"] == "FAIL"
    wb = load_workbook(str(p))
    wb["01_BOQ_SUMMARY"].cell(3, 7).value = 307.361518
    wb["02_ROOM_BREAKDOWN"].cell(10, 8).value = "=SUM(H4:H5)"
    wb.save(str(p))
    v = BX.validate(p, rep())
    assert v["state"] == "FAIL" and v["formulas"]


def test_the_breakdown_sheet_is_marked_non_additive_and_never_sums():
    e = BX.expected(rep())
    assert all(r[0].startswith("BREAKDOWN OF IT-A (not additive") for r in e["02_ROOM_BREAKDOWN"][1:])
    assert sum(r[7] for r in e["02_ROOM_BREAKDOWN"][1:]) != 0                     # the values are there ...
    assert not [r for r in e["01_BOQ_SUMMARY"][1:] if r[5] != "ALL"]             # ... and never added in the view
    assert "a formula" in BX.policy_record()["never"]


# ------------------------------------------------------------------------------- source anchor
def ent(h, geom, layer="WALL", t="LINE", path=()):
    return {"handle": h, "instance_path": list(path), "layer": layer, "entity_type": t, "geometry": geom}


def test_export_identity_is_all_or_nothing():
    a = [ent("A1", [0.0, 0.0, 1.0, 0.0]), ent("A2", [1.0, 0.0, 1.0, 1.0], path=("I9",))]
    assert SA.compare(a, [dict(x) for x in a], tol=1e-9)["state"] == SA.IDENTICAL
    moved = [a[0], ent("A2", [1.0, 0.0, 1.0, 1.001], path=("I9",))]
    r = SA.compare(a, moved, tol=1e-6)
    assert r["state"] == SA.DIFFERENT and r["changed"][0]["fields"] == ["geometry"]
    assert SA.compare(a, a[:1], tol=1e-9)["only_in_a"] == [["A2", ["I9"]]]
    relayer = [a[0], dict(a[1], layer="FIXTURE")]
    assert SA.compare(a, relayer, tol=1e-9)["changed"][0]["fields"] == ["layer"]
    assert "a hash of a re-saved source" in SA.policy_record()["never"]


# ------------------------------------------------------------------------------- physical rationale
def test_a_physical_construction_rationale_fact_authorises_no_domain():
    f = OF.from_record({"fact_id": "X-RATIONALE", "version": 1, "kind": "PHYSICAL_CONSTRUCTION_RATIONALE",
                        "authority": ["PROJECT_OWNER"], "scope": {}, "allowed_domains": [],
                        "parts": [{"key": "A|H1||SEGMENT|0", "fingerprint": "f", "physical_reading": "R"}]})
    assert f.allowed_domains == () and not set(f.allowed_domains) & set(OF.DOMAINS)
