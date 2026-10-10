"""Reporting V2 - generic model, renderer and readback on a synthetic register set (no project data)."""

import ast
import json
from pathlib import Path

import pytest

from engine.reporting_v2 import layout as L
from engine.reporting_v2 import model as M
from engine.reporting_v2 import readback as RB
from engine.reporting_v2 import terms as T
from engine.reporting_v2 import xlsx as X

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def regs(tmp_path):
    p = tmp_path / "REG.json"
    p.write_text(json.dumps({"slab": {"GF": 10.1, "1F": 20.2}, "beam": 0.1, "beam2": 0.2, "blocked": None, "count": 3,
                             "rows": [{"a": 1.5}, {"a": 2.25}]}))
    return M.Registers({"REG": p})


def _model(regs, *, breakdown_cls="BREAKDOWN_ONLY", summary_override=None):
    lines = [M.line("CON", "STRUCTURAL_CONCRETE", "Concrete", "خرسانة", "m3",
                    [M.part("GF", regs.q("REG:/slab/GF"), "COMPUTED"), M.part("1F", regs.q("REG:/slab/1F"), "COMPUTED"),
                     M.part("FOUNDATION", M.blocked("BLOCKED_HEIGHT"), "BLOCKED_HEIGHT")]),
             M.line("CNT", "DOORS", "Doors", "", "nr", [M.part("GF", regs.q("REG:/count"), "COMPUTED")])]
    lines[0]["matrix"] = "CONCRETE"
    rows = [M.row(["GF slab", regs.q("REG:/slab/GF"), "COMPUTED"], cls=breakdown_cls, status="COMPUTED", tech="COMPUTED",
                  explains="CON@GF"),
            M.row(["blocked thing", M.blocked("BLOCKED_HEIGHT"), "BLOCKED"], cls="BREAKDOWN_ONLY", status="BLOCKED",
                  tech="BLOCKED_HEIGHT", explains="CON@FOUNDATION")]
    sec = M.section("S", ("SECTION", "قسم"), [M.col("i", "ITEM"), M.col("q", "QTY", "", "qty", "m3"), M.col("s", "STATUS", "", "status")], rows)
    sh = M.sheet("01_GROUND_FLOOR", ("GROUND FLOOR", "الدور الأرضي"), "BREAKDOWN", [sec], level="GF")
    proj = {"name_en": "SYNTHETIC", "name_ar": "", "phase": "TEST", "release_state": "SHADOW", "subtitle": "synthetic",
            "header": [("Project", "synthetic")]}
    return L.assemble(proj, lines, [sh], regs)


def test_declared_sum_is_exact_decimal_addition(regs):
    s = M.qsum([regs.q("REG:/beam"), regs.q("REG:/beam2")])
    assert s["q"] == 0.3 and s["sum"] == ["REG:/beam", "REG:/beam2"]
    assert M.qsum([regs.q("REG:/count"), regs.q("REG:/count")])["q"] == 6


def test_line_status_blocked_is_never_zero(regs):
    m = _model(regs)
    con = m["lines"][0]
    assert con["status"] == "PARTIAL" and con["qty"]["q"] == 30.3 and len(con["qty"]["sum"]) == 2
    only_blocked = M.line("X", "PAINT", "x", "", "m2", [M.part("GF", M.blocked("BLOCKED_MATERIAL"), "BLOCKED_MATERIAL")])
    assert only_blocked["status"] == "BLOCKED" and only_blocked["qty"]["q"] is None


def test_valid_model_passes(regs):
    v = M.validate(_model(regs), regs)
    assert v["state"] == "PASS", v["problems"]
    assert v["quantity_cells"] > 0 and v["declared_sums"] >= 1


def test_additive_row_outside_summary_fails(regs):
    v = M.validate(_model(regs, breakdown_cls="ADDITIVE"), regs)
    assert v["state"] == "FAIL" and any("outside the summary" in p for p in v["problems"])


def test_quantity_that_differs_from_its_register_fails(regs):
    m = _model(regs)
    m["sheets"][1]["sections"][0]["rows"][0]["cells"][1]["q"] = 10.2
    assert any("REG:/slab/GF" in p for p in M.validate(m, regs)["problems"])


def test_declared_sum_mismatch_fails(regs):
    m = _model(regs)
    m["lines"][0]["qty"]["q"] = 30.4
    assert M.validate(m, regs)["state"] == "FAIL"


def test_line_missing_from_summary_fails(regs):
    m = _model(regs)
    lines_sec = [s for s in m["sheets"][0]["sections"] if s["id"] == "BOQ_LINES"][0]
    lines_sec["rows"] = [r for r in lines_sec["rows"] if r.get("explains") != "CNT"]
    assert any("not on the summary" in p for p in M.validate(m, regs)["problems"])


def test_unknown_status_code_fails_closed():
    with pytest.raises(KeyError):
        M.display_status("SOMETHING_NEW")
    with pytest.raises(KeyError):
        T.explain("SOMETHING_NEW")
    assert M.display_status("BLOCKED_ANYTHING") == "BLOCKED" and T.explain("BLOCKED_UPPER_MEMBER_UNBOUND")["priority"] == "HIGH"


def test_trade_totals_only_add_one_group_and_unit(regs):
    m = _model(regs)
    totals = [s for s in m["sheets"][0]["sections"] if s["id"] == "TRADE_TOTALS"][0]["rows"]
    assert all(r["role"] == "TOTAL" for r in totals) and len(totals) == 2
    matrix = [s for s in m["sheets"][0]["sections"] if s["id"] == "FLOOR_TOTALS"][0]
    total_row = matrix["rows"][-1]
    assert total_row["role"] == "TOTAL" and M.is_f(total_row["cells"][2]) and total_row["cells"][2]["v"] == pytest.approx(30.3)
    assert all(M.is_f(r["cells"][3]) and r["cells"][3]["expect"] is not None for r in totals)


def _n_formulas(m):
    return sum(1 for *_, c in M.iter_cells(m) if M.is_f(c))


def test_render_readback_and_tamper(regs, tmp_path):
    m = _model(regs)
    x = X.render(m, tmp_path / "r.xlsx")
    rb = RB.validate(tmp_path / "r.xlsx", m, x["cell_map"])
    nf = _n_formulas(m)
    assert rb["state"] == "PASS" and nf > 0 and rb["formula_cells"] == nf and rb["validated_quantity_cells"] > 0
    assert rb["sheets_protected"] == len(m["sheets"])
    x2 = X.render(m, tmp_path / "r2.xlsx")
    assert x2["file_sha256"] == x["file_sha256"]                         # deterministic bytes
    from openpyxl import load_workbook
    from openpyxl.styles import Protection
    wb = load_workbook(tmp_path / "r.xlsx")
    sheet, coord = next((e[0], e[1]) for e in x["cell_map"] if e[2] == "qty")
    fsheet, fcoord = next((e[0], e[1]) for e in x["cell_map"] if e[2] == "formula")
    wb[sheet][coord] = 999.0                                              # an engine value edited
    wb[sheet]["Z99"] = "=SUM(1,2)"                                        # a formula the model does not hold
    wb[fsheet][fcoord] = "=1+1"                                           # a report total replaced
    wb[sheet]["Z98"].protection = Protection(locked=False)               # an unlocked cell that is not an input
    wb.save(tmp_path / "t.xlsx")
    bad = RB.validate(tmp_path / "t.xlsx", m, x["cell_map"])
    kinds = {d[0] for d in bad["differences"]}
    assert bad["state"] == "FAIL" and bad["formula_cells"] == nf + 1 and any(d[1] == coord for d in bad["differences"])
    assert {"unmapped_formula", "unlocked_non_input", fsheet} <= kinds


def test_blocked_quantity_is_written_as_text_not_zero(regs, tmp_path):
    m = _model(regs)
    x = X.render(m, tmp_path / "r.xlsx")
    texts = [e for e in x["cell_map"] if e[2] == "qty_text"]
    assert texts and all(e[3] == "BLOCKED" for e in texts if e[3] != "—" and e[3] != "")


def test_reporting_package_imports_no_qto_engine():
    """The presentation layer measures nothing: no import of engine.source or any other engine module."""
    for p in sorted((ROOT / "engine" / "reporting_v2").glob("*.py")):
        tree = ast.parse(p.read_text())
        for n in ast.walk(tree):
            names = [a.name for a in n.names] if isinstance(n, ast.Import) else \
                ([n.module or ""] if isinstance(n, ast.ImportFrom) and n.level == 0 else [])
            for name in names:
                assert not name.startswith("engine") or name.startswith("engine.reporting_v2"), (p.name, name)


def test_line_totals_and_engine_cells_are_never_formulas(regs):
    """Formulas only add frozen rows up: every engine quantity (and every BOQ line total) is a register value."""
    m = _model(regs)
    lines = [s for s in m["sheets"][0]["sections"] if s["id"] == "BOQ_LINES"][0]
    keys = [c["key"] for c in lines["columns"]]
    for r in lines["rows"]:
        if r["role"] == "ITEM":
            assert M.is_q(r["cells"][keys.index("total")])
    for s, sec, i, r, j, c in M.iter_cells(m):
        if M.is_f(c):
            assert r["cls"] != "ADDITIVE" or sec["columns"][j]["kind"] != "qty" or r["role"] != "ITEM"


def _recon_model(regs):
    lines = [M.line("CON", "STRUCTURAL_CONCRETE", "Concrete", "خرسانة", "m3",
                    [M.part("GF", regs.q("REG:/slab/GF"), "COMPUTED"), M.part("1F", regs.q("REG:/slab/1F"), "COMPUTED"),
                     M.part("FOUNDATION", M.blocked("BLOCKED_HEIGHT"), "BLOCKED_HEIGHT")]),
             M.line("CON2", "STRUCTURAL_CONCRETE", "Concrete 2", "", "m3", [M.part("GF", regs.q("REG:/beam"), "COMPUTED")])]
    proj = {"name_en": "SYNTHETIC", "name_ar": "", "phase": "TEST", "release_state": "SHADOW", "subtitle": "synthetic",
            "header": [("Project", "synthetic")]}
    rs = L.reconciliation_sheet(lines, ["FOUNDATION", "GF", "1F"])
    return L.assemble(proj, lines, [rs], regs, levels=["FOUNDATION", "GF", "1F"])


def _status(m, key):
    sec = [x for x in m["sheets"][1]["sections"] if x["id"] == "ITEMS"][0]
    keys = [c["key"] for c in sec["columns"]]
    r = [x for x in sec["rows"] if x.get("key") == key][0]
    return r, keys


@pytest.mark.parametrize("manual,expect", [(None, "NOT CHECKED"), (10.1, "MATCH"), (10.1 * 1.004, "MATCH"), (10.1 * 1.02, "CLOSE"),
                                           (10.1 * 1.2, "REVIEW"), (10.1 * 0.9, "REVIEW"), ("n/a", "NOT COMPARABLE")])
def test_reconciliation_status_rules(regs, manual, expect):
    m = _recon_model(regs)
    assert M.validate(m, regs)["state"] == "PASS"
    r, keys = _status(m, "I:CON:0")
    r["cells"][keys.index("manual")]["input"] = manual
    M.evaluate(m)
    assert r["cells"][keys.index("rstatus")]["v"] == expect
    pct = r["cells"][keys.index("pct")]["v"]
    assert (pct == "") == (not isinstance(manual, float))


def test_engine_blocked_row_is_never_a_mismatch(regs):
    m = _recon_model(regs)
    r, keys = _status(m, "I:CON:2")
    assert r["cells"][keys.index("urban")]["blocked"] == "BLOCKED_HEIGHT"
    for manual in (None, 0.0, 5.0, "text"):
        r["cells"][keys.index("manual")]["input"] = manual
        M.evaluate(m)
        assert r["cells"][keys.index("rstatus")]["v"] == "ENGINE BLOCKED"
        assert r["cells"][keys.index("pct")]["v"] == "" and r["cells"][keys.index("diff")]["v"] == ""


def test_reconciliation_subtotals_and_no_feedback(regs):
    m = _recon_model(regs)
    sec = [x for x in m["sheets"][1]["sections"] if x["id"] == "ITEMS"][0]
    keys = [c["key"] for c in sec["columns"]]
    sub = [r for r in sec["rows"] if r["role"] == "SUBTOTAL"]
    tot = [r for r in sec["rows"] if r["role"] == "TOTAL"]
    assert len(sub) == 1 and sub[0]["cells"][keys.index("urban")]["v"] == pytest.approx(10.2)       # GF: 10.1 + 0.1
    assert len(tot) == 1 and tot[0]["cells"][keys.index("urban")]["v"] == pytest.approx(30.4)       # blocked part excluded
    assert all(r["cls"] != "ADDITIVE" for r in sec["rows"])
    before = [c for *_, c in M.iter_cells(m) if M.is_q(c)]
    summary = [c["v"] for *_, c in M.iter_cells({"sheets": [m["sheets"][0]]}) if M.is_f(c)]
    for r in sec["rows"]:
        if r["role"] == "ITEM":
            r["cells"][keys.index("manual")]["input"] = 99.0
    M.evaluate(m)
    assert [c for *_, c in M.iter_cells(m) if M.is_q(c)] == before
    assert [c["v"] for *_, c in M.iter_cells({"sheets": [m["sheets"][0]]}) if M.is_f(c)] == summary
    assert sub[0]["cells"][keys.index("manual")]["v"] == pytest.approx(198.0)


def test_reconciliation_xlsx_inputs_unlocked_engine_locked(regs, tmp_path):
    m = _recon_model(regs)
    x = X.render(m, tmp_path / "r.xlsx")
    rb = RB.validate(tmp_path / "r.xlsx", m, x["cell_map"])
    n_in = sum(1 for *_, c in M.iter_cells(m) if M.is_in(c))
    assert rb["state"] == "PASS" and rb["input_cells"] == n_in > 0
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "r.xlsx")[L.RECON_NAME]
    assert ws.auto_filter.ref and ws.protection.sheet and ws.conditional_formatting
    text = {e[3] for e in x["cell_map"] if e[2] == "formula"}
    assert any(t.startswith('=IF(') and '"ENGINE BLOCKED"' in t for t in text)


@pytest.mark.skipif(RB.soffice() is None, reason="LibreOffice not installed")
def test_reconciliation_formulas_recalculate_like_the_model(regs, tmp_path):
    m = _recon_model(regs)
    x = X.render(m, tmp_path / "r.xlsx")
    assert RB.recalc(tmp_path / "r.xlsx", x["cell_map"])["state"] == "PASS"
    edits = [[L.RECON_NAME, "ITEMS", "I:CON:0", "manual", 10.3], [L.RECON_NAME, "ITEMS", "I:CON:1", "manual", 20.2],
             [L.RECON_NAME, "ITEMS", "I:CON:2", "manual", 4.0], [L.RECON_NAME, "PARAMS", "CLOSE", "v", 0.05]]
    w = RB.whatif(tmp_path / "r.xlsx", m, x["cell_map"], edits, tmp_path)
    assert w["state"] == "PASS" and w["formula_cells_changed_by_edits"] > 0
    r, keys = _status(w["model_after"], "I:CON:0")
    assert r["cells"][keys.index("rstatus")]["v"] == "CLOSE"                                       # 1.98 % within 5 %
