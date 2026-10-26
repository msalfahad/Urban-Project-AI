"""T-235 r2 + part (4): URBAN_CHECKER_VERDICT_V2 schema, validator, 'verdict add' helper, the --verdicts-v2 CLI flag,
and the T-236 trade completeness report (TRADE_COMPLETENESS.json, reference_line_rows) as the completeness input.

Hand calculation (written before running the tool), hand-made synthetic villa (not the development villa):
  CONCRETE STR.CONC.FOOTING m3, storey FDN: F1 1.50 x 1.00 x 0.80 = 1.200 ; F2 1.00 x 1.00 x 0.80 = 0.800 ;
  F3 2.50 x 1.00 x 0.80 = 2.000  -> line 4.000 m3, 3 elements.
  Independent inventory (printed schedule, basis PRINTED_SCHEDULE): FDN F1, F2, F3 -> found 3, missing 0, duplicated 0,
  held 0 -> COMPLETE_AGAINST_INVENTORY.
  Checker WHOLE re-measure 4.010 m3 (+0.25 %, inside 1 %) -> COUNTED -> F1 / F2 / F3 VERIFIED 1.200 / 0.800 / 2.000,
  line VERIFIED, verified subtotal 4.000; every other row keeps its state.
  Adverse: the schedule also lists F4 -> MISSING 1 -> F4 an UNMEASURED row, line NOT VERIFIED (completeness fails);
  a strap ST1 in the schedule with no line -> UNMEASURED row under CONCRETE.STRAP; F2 linked by two lines ->
  DUPLICATED -> not VERIFIED; a report that knows a line the export does not -> refused; a footing line in m2 as well
  -> the trade spans two BOQ lines -> its completeness is not this line's -> not VERIFIED.
"""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from engine.source import boq_line as BL
from engine.source import trade_completeness_v1 as TC
from tools import boq_export_v1 as X

PID, REV = "SYNTH-VILLA-9", "9" * 64
CVER = "ARCH_TAKEOFF_V3.0.TEST"
FOOT = ("CONCRETE", "STR.CONC.FOOTING", "m3")


def ln(code, eid, ev="MEASURED", q=None, unit="m3", floor="FDN", **kw):
    return BL.make_line(project_id=kw.pop("project_id", PID), revision_set_sha256=kw.pop("revision_set_sha256", REV),
                        item_code=code, unit=unit, evidence_state=ev, measured_net=q, element_id=eid,
                        formula="L*B*H", rule_id="R-SYN", rule_version="1", floor=floor, **kw)


def villa():
    return [ln("STR.CONC.FOOTING", "F1", q=1.2, inputs={"L": 1.5, "B": 1.0, "H": 0.8}),
            ln("STR.CONC.FOOTING", "F2", q=0.8, inputs={"L": 1.0, "B": 1.0, "H": 0.8}),
            ln("STR.CONC.FOOTING", "F3", q=2.0, inputs={"L": 2.5, "B": 1.0, "H": 0.8}),
            ln("STR.CONC.SLAB", "S1", ev="BOUNDED_EXPOSURE", bounded_range={"low": 5.0, "high": 6.0},
               rfi={"question": "slab depth 0.20 or 0.24"}, floor="GF"),
            ln("FIN.PAINT", "P1", ev="PROVISIONAL", q=40.0, unit="m2", floor="GF", inputs={"gross": 44.0})]


def links(lines, extra=None):
    out = {x["line_id"]: [TC.element_key(x["location"]["floor"], x["location"]["element_id"])] for x in lines}
    out.update(extra or {})
    return out


def report(lines, labels=("F1", "F2", "F3"), straps=(), extra_links=None, report_lines=None):
    inv = [{"source": "SCHED-FDN", "basis": "PRINTED_SCHEDULE", "trade": "CONCRETE.FOOTING",
            "elements": [{"floor": "FDN", "label": lab, "ref": f"schedule row {lab}"} for lab in labels]}]
    if straps:
        inv.append({"source": "SCHED-ST", "basis": "PRINTED_SCHEDULE", "trade": "CONCRETE.STRAP",
                    "elements": [{"floor": "FDN", "label": lab, "ref": "strap schedule"} for lab in straps]})
    rl = lines if report_lines is None else report_lines
    return TC.check(inventories=inv, lines=rl, line_links=links(rl, extra_links))


def by(rows, eid):
    return [r for r in rows if r["element_id"] == eid][0]


def summ(path):
    return {(s["trade"], s["item_code"], s["unit"]): s for s in X.read_csv(path / "BOQ_TRADE_SUMMARY.csv")}


def add(lines, ids, qty, scope="WHOLE", **kw):
    base = dict(candidate_version=CVER, boq_lines_sha256=X.lines_sha256(lines), line_ids=ids, scope=scope,
                checker_qty=qty, evidence=[{"sheet": "S-02 footing schedule", "printed_value": "F1 1.50x1.00x0.80"}],
                run_token="20261024T085003-1-1", routine="Verifier", verdict_commit="c0ffee12",
                time="2026-10-24T10:00:00Z")
    base.update(kw)
    return X.verdict_add(lines, **base)


def run(tmp_path, lines, name="o", **kw):
    m = X.export(lines, tmp_path / name, candidate_version=CVER, boq_lines_sha256=X.lines_sha256(lines), **kw)
    return m, X.read_csv(tmp_path / name / "BOQ.csv")


def foot_ids(lines):
    return [x["line_id"] for x in lines if x["item_code"] == "STR.CONC.FOOTING"]


# ---- the schema and its validator -----------------------------------------------------------------------------------

def test_schema_file_is_the_tool_schema():
    p = X.ROOT / "research" / "agents_t235" / "URBAN_CHECKER_VERDICT_V2.schema.json"
    assert json.loads(p.read_text(encoding="utf-8")) == X.schema_document()
    assert set(X.schema_document()["required"]) == set(X.V2_REQUIRED)


def test_helper_fills_the_binding_and_the_record_is_valid():
    lines = villa()
    rec = add(lines, foot_ids(lines), 4.01)
    assert X.validate_v2(rec) == []
    assert (rec["project_id"], rec["revision_set_sha256"], rec["candidate_version"], rec["unit"]) == (PID, REV, CVER, "m3")
    assert rec["element_scope"] == {"ids": ["F1", "F2", "F3"], "count": 3}
    assert rec["engine_qty"] == 4.0 and rec["checker_qty"] == 4.01 and rec["line_ids"] == sorted(foot_ids(lines))
    assert rec["boq_lines_sha256"] == X.lines_sha256(lines) and rec["verdict_id"].startswith("V2-")
    assert add(lines, list(reversed(foot_ids(lines))), 4.01) == rec          # deterministic, order-free


@pytest.mark.parametrize("over,needle", [
    ({"evidence": [{"printed_value": "1.2"}]}, "evidence"),                   # no sheet
    ({"evidence": [{"sheet": "S-02"}]}, "evidence"),                          # no printed value
    ({"checker": {"run_token": "t", "routine": "Worker 3", "verdict_commit": "c0ffee12"}}, "routine"),
    ({"time": "yesterday"}, "time"),
    ({"scope": "ALL"}, "scope"),
    ({"engine_qty": "4.0"}, "engine_qty"),
    ({"checker_qty": True}, "checker_qty"),
    ({"unit": "cubic"}, "unit"),
    ({"element_scope": {"ids": ["F1", "F2", "F3"], "count": 4}}, "element_scope"),
    ({"element_scope": {"ids": ["F1", "F1", "F3"], "count": 3}}, "element_scope"),
    ({"revision_set_sha256": "abc"}, "revision_set_sha256"),
    ({"schema": "URBAN_SCOREBOARD_CHECKER_VERDICTS_V1"}, "schema")])
def test_validator_refuses_a_broken_binding(over, needle):
    lines = villa()
    rec = add(lines, foot_ids(lines), 4.01) | over
    probs = X.validate_v2(rec)
    assert probs and any(needle in p for p in probs)
    assert X.bind_verdict(rec, (PID, REV, CVER, X.lines_sha256(lines)), lines[:3], ["F1", "F2", "F3"]) in (
        "BINDING_INCOMPLETE", "NOT_V2")


def test_validator_refuses_duplicated_line_ids_and_missing_fields():
    lines = villa()
    rec = add(lines, foot_ids(lines), 4.01)
    assert any("line_ids" in p for p in X.validate_v2(rec | {"line_ids": rec["line_ids"] + rec["line_ids"][:1]}))
    for k in X.V2_REQUIRED:
        broken = dict(rec)
        broken.pop(k)
        assert any(k in p for p in X.validate_v2(broken)), k


@pytest.mark.parametrize("ids,why", [(["nope"], "unknown"), ("S1", "not measured"), ("mixed", "unit")])
def test_helper_refuses_what_a_verdict_cannot_bind(ids, why):
    lines = villa()
    if ids == "S1":
        ids = [x["line_id"] for x in lines if x["location"]["element_id"] == "S1"]
    elif ids == "mixed":
        ids = [lines[0]["line_id"], lines[4]["line_id"]]
    with pytest.raises(X.ExportError, match=why):
        add(lines, ids, 1.0)


def test_helper_refuses_mixed_identity_and_a_non_checker():
    lines = villa() + [ln("STR.CONC.FOOTING", "F9", q=1.0, inputs={"L": 1}, revision_set_sha256="8" * 64)]
    with pytest.raises(X.ExportError, match="revision set"):
        add(lines, foot_ids(lines)[:1], 1.2)
    with pytest.raises(X.ExportError, match="routine"):
        add(villa(), foot_ids(villa()), 4.0, routine="Worker 1")


def test_verdict_set_appends_sorted_and_refuses_a_repeated_id(tmp_path):
    lines = villa()
    p = tmp_path / "V.json"
    a = add(lines, foot_ids(lines)[:1], 1.2, scope="COMPONENT")
    b = add(lines, foot_ids(lines), 4.01)
    X.verdict_set_append(p, b)
    X.verdict_set_append(p, a)
    doc = json.loads(p.read_text(encoding="utf-8"))
    assert doc["schema"] == X.V2_SET_SCHEMA and [v["verdict_id"] for v in doc["verdicts"]] == sorted(
        [a["verdict_id"], b["verdict_id"]])
    with pytest.raises(X.ExportError, match="already"):
        X.verdict_set_append(p, a)
    assert X.validate_v2_set(doc) == []
    assert X.validate_v2_set(doc | {"schema": "X"})


# ---- export with V2 verdicts + the T-236 report (part 4) ------------------------------------------------------------

def test_whole_verdict_plus_complete_report_verifies_the_line(tmp_path):
    lines = villa()
    rec = add(lines, foot_ids(lines), 4.01)
    m, rows = run(tmp_path, lines, v2_records=[rec], completeness_report=report(lines))
    assert m["row_counts"]["VERIFIED"] == 3 and m["row_counts"]["MEASURED UNVERIFIED"] == 1   # paint stays
    assert [by(rows, e)["qty_verified"] for e in ("F1", "F2", "F3")] == ["1.200000", "0.800000", "2.000000"]
    s = summ(tmp_path / "o")[FOOT]
    assert s["line_status"] == "VERIFIED" and s["verified_subtotal"] == "4.000000"
    assert s["completeness"].startswith("PASS")
    assert by(rows, "S1")["quantity_state"] == "HELD" and by(rows, "P1")["quantity_state"] == "MEASURED UNVERIFIED"
    assert {v["verdict_id"]: v["status"] for v in m["v2_verdicts"]} == {rec["verdict_id"]: "COUNTED (WHOLE)"}


def test_without_a_report_or_with_a_v1_record_nothing_is_verified(tmp_path):
    lines = villa()
    rec = add(lines, foot_ids(lines), 4.01)
    m, rows = run(tmp_path, lines, v2_records=[rec])
    assert m["row_counts"]["VERIFIED"] == 0 and "completeness" in by(rows, "F1")["verdict_status"]


def test_missing_inventory_element_is_an_unmeasured_row_and_blocks_verified(tmp_path):
    lines = villa()
    rec = add(lines, foot_ids(lines), 4.01)
    m, rows = run(tmp_path, lines, v2_records=[rec], completeness_report=report(lines, labels=("F1", "F2", "F3", "F4")))
    f4 = by(rows, "F4")
    assert (f4["quantity_state"], f4["net"], f4["item_code"], f4["floor"]) == ("UNMEASURED", "", "STR.CONC.FOOTING", "FDN")
    assert "MISSING" in f4["unmeasured_reason"] and m["row_counts"]["VERIFIED"] == 0
    s = summ(tmp_path / "o")[FOOT]
    assert s["line_status"] == "NOT VERIFIED" and "MISSING 1" in s["completeness"] and s["unmeasured_components"] == "F4"


def test_a_trade_with_inventory_and_no_line_is_shown_unmeasured(tmp_path):
    lines = villa()
    m, rows = run(tmp_path, lines, completeness_report=report(lines, straps=("ST1",)))
    st = by(rows, "ST1")
    assert (st["trade"], st["item_code"], st["quantity_state"], st["net"]) == ("CONCRETE.STRAP", "", "UNMEASURED", "")
    # principal trades with no line and no inventory are rows too (never absent), quantity blank
    neck = [r for r in rows if r["trade"] == "CONCRETE.NECK"]
    assert len(neck) == 1 and neck[0]["quantity_state"] == "UNMEASURED" and neck[0]["net"] == ""
    assert "no line" in neck[0]["unmeasured_reason"]


def test_duplicated_element_blocks_verified(tmp_path):
    lines = villa()
    f3 = [x["line_id"] for x in lines if x["location"]["element_id"] == "F3"][0]
    rep = report(lines, extra_links={f3: ["FDN:F3", "FDN:F2"]})
    rec = add(lines, foot_ids(lines), 4.01)
    m, _rows = run(tmp_path, lines, v2_records=[rec], completeness_report=rep)
    assert m["row_counts"]["VERIFIED"] == 0 and "DUPLICATED" in summ(tmp_path / "o")[FOOT]["completeness"]


def test_report_of_other_lines_is_refused(tmp_path):
    lines = villa()
    other = lines + [ln("STR.CONC.FOOTING", "F7", q=0.5, inputs={"L": 1})]
    with pytest.raises(X.ExportError, match="report"):
        run(tmp_path, lines, completeness_report=report(other, labels=("F1", "F2", "F3", "F7"), report_lines=other))


def test_trade_spanning_two_boq_lines_is_not_this_lines_completeness(tmp_path):
    lines = villa() + [ln("STR.CONC.FOOTING", "F1", q=3.0, unit="m2", inputs={"L": 1.5, "B": 2.0})]
    rec = add(lines, [x["line_id"] for x in lines if x["item_code"] == "STR.CONC.FOOTING" and x["unit"] == "m3"], 4.01)
    m, _rows = run(tmp_path, lines, v2_records=[rec], completeness_report=report(lines))
    assert m["row_counts"]["VERIFIED"] == 0 and "SPANS" in summ(tmp_path / "o")[FOOT]["completeness"]


def test_inventory_v1_and_report_together_refused(tmp_path):
    lines = villa()
    inv = {"schema": X.INVENTORY_SCHEMA, "project_id": PID, "revision_set_sha256": REV, "elements": []}
    with pytest.raises(X.ExportError, match="one inventory"):
        run(tmp_path, lines, inventory_doc=inv, completeness_report=report(lines))


def test_stale_record_is_listed_never_dropped(tmp_path):
    lines = villa()
    rec = add(lines, foot_ids(lines), 4.01, candidate_version="ARCH_TAKEOFF_V3.0.OLD")
    # T-294: a record bound to its own lines' content stays valid on another candidate version (tests/t294_*); a
    # record issued under the T-245 whole-file binding (no binding field) is still STALE there, never migrated
    rec = {k: v for k, v in rec.items() if k not in ("binding", "line_content_sha256")}
    m, rows = run(tmp_path, lines, v2_records=[rec], completeness_report=report(lines))
    assert m["row_counts"]["VERIFIED"] == 0
    assert {v["verdict_id"]: v["status"] for v in m["v2_verdicts"]}[rec["verdict_id"]].startswith("STALE")


# ---- the command line: verdict add, verdict validate, export --verdicts-v2 --completeness ---------------------------

def test_cli_end_to_end(tmp_path):
    lines = villa()
    lp = tmp_path / "V3_0_T_BOQ_LINES.json"
    lp.write_text(BL.dumps(lines), encoding="utf-8")
    vp = tmp_path / "VERDICTS_V2.json"
    lid = ",".join(foot_ids(lines))
    assert X.main(["verdict-add", "--lines", str(lp), "--candidate-version", CVER, "--line-ids", lid, "--scope", "WHOLE",
                   "--checker-qty", "4.01", "--evidence", "S-02 footing schedule=F1 1.50x1.00x0.80",
                   "--run-token", "tok-1", "--routine", "Verifier", "--verdict-commit", "c0ffee12",
                   "--time", "2026-10-24T10:00:00Z", "--out", str(vp)]) == 0
    assert X.main(["verdict-validate", str(vp)]) == 0
    tin = tmp_path / "tc_in.json"
    tin.write_text(json.dumps({"inventories": [{"source": "SCHED-FDN", "basis": "PRINTED_SCHEDULE",
                                                "trade": "CONCRETE.FOOTING",
                                                "elements": [{"floor": "FDN", "label": f} for f in ("F1", "F2", "F3")]}],
                               "lines_file": lp.name, "line_links": links(lines)}), encoding="utf-8")
    from tools import trade_completeness as TCR
    assert TCR.main([str(tin), "--out", str(tmp_path / "tc")]) == 0
    assert X.main(["--lines", str(lp), "--candidate-version", CVER, "--out", str(tmp_path / "cli"),
                   "--verdicts-v2", str(vp), "--completeness", str(tmp_path / "tc" / "TRADE_COMPLETENESS.json"),
                   "--code-commit", "abc"]) == 0
    m = json.loads((tmp_path / "cli" / "EXPORT_MANIFEST.json").read_text(encoding="utf-8"))
    assert m["row_counts"]["VERIFIED"] == 3
    assert m["inputs_sha256"]["verdicts_v2"] == X.sha256_bytes(vp.read_bytes())
    assert m["inputs_sha256"]["completeness"] == X.sha256_bytes((tmp_path / "tc" / "TRADE_COMPLETENESS.json").read_bytes())
    ws = load_workbook(tmp_path / "cli" / "BOQ.xlsx")["BOQ"]
    head = [c.value for c in ws[1]]
    xr = {r["element_id"]: r for r in (dict(zip(head, [c.value for c in row])) for row in ws.iter_rows(min_row=2))}
    assert xr["F2"]["quantity_state"] == "VERIFIED" and xr["F2"]["qty_verified"] == 0.8
    assert xr["F2"]["qty_measured_unverified"] is None


def test_cli_validate_flags_a_broken_file(tmp_path, capsys):
    p = tmp_path / "bad.json"
    p.write_text(json.dumps({"schema": X.V2_SET_SCHEMA, "verdicts": [{"schema": X.V2_VERDICT_SCHEMA,
                                                                      "verdict_id": "V2-x"}]}), encoding="utf-8")
    assert X.main(["verdict-validate", str(p)]) == 1
    assert "V2-x" in capsys.readouterr().out


def test_cli_export_with_v2_records_is_byte_identical_twice(tmp_path):
    lines = villa()
    lp = tmp_path / "L.json"
    lp.write_text(BL.dumps(lines), encoding="utf-8")
    vp = tmp_path / "V.json"
    X.verdict_set_append(vp, add(lines, foot_ids(lines), 4.01))
    rp = tmp_path / "R.json"
    rp.write_text(TC.dumps(report(lines)), encoding="utf-8")
    for o in ("a", "b"):
        X.main(["--lines", str(lp), "--candidate-version", CVER, "--out", str(tmp_path / o), "--verdicts-v2", str(vp),
                "--completeness", str(rp), "--code-commit", "abc"])
    for f in ("BOQ.csv", "BOQ.xlsx", "BOQ_TRADE_SUMMARY.csv", "EXPORT_MANIFEST.json"):
        assert (tmp_path / "a" / f).read_bytes() == (tmp_path / "b" / f).read_bytes(), f


def test_cli_verdict_add_by_item_code_with_the_checker_lock(tmp_path):
    lines = villa()
    lp = tmp_path / "L.json"
    lp.write_text(BL.dumps(lines), encoding="utf-8")
    lock = tmp_path / "checker.lock"
    lock.write_text(json.dumps({"agent": "Verifier", "token": "20261024T085003-464-21325"}), encoding="utf-8")
    vp = tmp_path / "V.json"
    args = ["verdict-add", "--lines", str(lp), "--candidate-version", CVER, "--item-code", "STR.CONC.FOOTING",
            "--scope", "WHOLE", "--checker-qty", "4.01", "--evidence", "S-02=F1 1.50x1.00x0.80",
            "--lock-file", str(lock), "--verdict-commit", "c0ffee12", "--time", "2026-10-24T10:00:00Z", "--out", str(vp)]
    assert X.main(args) == 0
    rec = json.loads(vp.read_text(encoding="utf-8"))["verdicts"][0]
    assert rec["line_ids"] == sorted(foot_ids(lines)) and rec["checker"]["run_token"] == "20261024T085003-464-21325"
    assert rec == add(lines, foot_ids(lines), 4.01, run_token="20261024T085003-464-21325",
                      evidence=[{"sheet": "S-02", "printed_value": "F1 1.50x1.00x0.80"}])
    lock.write_text(json.dumps({"agent": "Worker 3", "token": "t"}), encoding="utf-8")
    with pytest.raises(X.ExportError, match="routine"):
        X.main(args[:-1] + [str(tmp_path / "W.json")])
    assert not (tmp_path / "W.json").exists()
