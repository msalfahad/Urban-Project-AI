"""The detailed takeoff: readable arithmetic that still adds up, and classifications that admit what is unsettled.

Each audit finding gets a validator rather than an assertion, and each validator is tested twice: once against
the published artifacts, where it must find nothing, and once against a deliberately corrupted copy, where it
must find the defect.  A test that only ever sees good data proves nothing about what it would catch.
"""

from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from research.qs_wall_treatment_01 import protocol as PR
from research.qs_wall_treatment_01.pa09.alrashed import (detailed_takeoff as DT, validation_amendment as VA,
                                                          window_standard as WS, workbook_calc as WC)

OUT = Path(PR.OUT_DIR) / "pa09_alrashed"
XL = OUT / "ALRASHED_DETAILED_QUANTITY_TAKEOFF.xlsx"
FROZEN = OUT / "FULL_VILLA_BLIND_PRE_PRICING_TAKEOFF.json"


@pytest.fixture(scope="module")
def wb():
    return load_workbook(XL)


@pytest.fixture(scope="module")
def vals(wb):
    return WC.evaluate_all(wb)


@pytest.fixture(scope="module")
def export():
    return json.loads((OUT / "ALRASHED_DETAILED_QUANTITY_EXPORT.json").read_text("utf-8"))


@pytest.fixture(scope="module")
def recon():
    return json.loads((OUT / "ALRASHED_QUANTITY_RECONCILIATION.json").read_text("utf-8"))


@pytest.fixture(scope="module")
def rec():
    return json.loads((OUT / "ALRASHED_DETAILED_QUANTITY_TAKEOFF.json").read_text("utf-8"))


@pytest.fixture(scope="module")
def frozen():
    return json.loads(FROZEN.read_text("utf-8"))


def cell(wb, vals, sheet, coord):
    return vals[(sheet, coord)] if (sheet, coord) in vals else wb[sheet][coord].value


def _find(ws, col, text):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == text:
            return r
    raise AssertionError(f"{text!r} not found in {ws.title}")


# ------------------------------------------------------------------ the validators the audit asked for
def problems_ar_fl_01(export, wb, vals):
    """AR-FL-01 must be the same number in the JSON, the floors sheet and the BOQ."""
    json_total = round(sum(r["MEASURED_QUANTITY"] for r in export["RECORDS"]
                           if r["TRADE"] == "PORCELAIN_FLOOR"), 3)
    sheet_total = cell(wb, vals, "الأرضيات", f"G{_find(wb['الأرضيات'], 5, 'AR-FL-01 — إجمالي الأرضيات المعتمدة')}")
    boq_row = _find(wb["BOQ حسب البند"], 2, "AR-FL-01")
    boq_total = cell(wb, vals, "BOQ حسب البند", f"F{boq_row}")
    out = []
    if abs(json_total - sheet_total) > DT.CROSS_TOL_M2:
        out.append(f"JSON {json_total} vs floors sheet {sheet_total}")
    if abs(sheet_total - boq_total) > DT.CROSS_TOL_M2:
        out.append(f"floors sheet {sheet_total} vs BOQ {boq_total}")
    return out


def problems_pending_in_porcelain(records):
    return [r["COMPONENT_REF"] for r in records
            if r["TRADE"] == "PORCELAIN_FLOOR" and r["COMPONENT_ROLE"] in DT.PENDING_FLOOR_ROLES]


def problems_entity_semantics(export):
    """A connected component may not be published as a room, and the census must reconcile."""
    c, out = export["ENTITY_SEMANTICS"], []
    if c["CONNECTED_COMPONENT_COUNT"] != c["NON_SLIVER_COMPONENT_COUNT"] + c["WALL_MATERIAL_OR_SLIVER_COUNT"]:
        out.append("census does not reconcile")
    for role, n in DT.EXPECTED_ROLE_COUNTS.items():
        if c["COUNT_BY_ROLE"].get(role) != n:
            out.append(f"{role}: {c['COUNT_BY_ROLE'].get(role)} expected {n}")
    for r in export["RECORDS"]:
        if r["ENTITY_TYPE"] == "CONNECTED_COMPONENT" and r["COMPONENT_ROLE"] is None:
            out.append(f"{r['COMPONENT_REF']} has no component role")
        if "ROOM_REF" in r or "ROOM_ROLE" in r:
            out.append(f"{r['ITEM_CODE']} still calls a component a room")
    return out


def problems_counts(counts):
    """Row counts and physical-object counts are different fields and must never be conflated."""
    out = []
    rows, objs, mult = counts["SCHEDULE_ROW_COUNT"], counts["PHYSICAL_OBJECT_COUNT"], counts["MULTIPLICITY"]
    if sum(v for k, v in rows.items() if k != "TOTAL") != rows["TOTAL"]:
        out.append("row counts do not add to their total")
    if sum(v for k, v in objs.items() if k != "TOTAL") != objs["TOTAL"]:
        out.append("object counts do not add to their total")
    extra = {"WINDOWS": 0, "DOORS": 0, "SLIDING_DOORS": 0}
    for mr in mult:
        key = {"WINDOW": "WINDOWS", "DOOR": "DOORS"}[mr["KIND"]]
        extra[key] += mr["PHYSICAL_OBJECT_COUNT"] - mr["ROW_COUNT"]
    for k in extra:
        if rows[k] + extra[k] != objs[k]:
            out.append(f"{k}: {rows[k]} rows + {extra[k]} extra objects != {objs[k]}")
    if objs["TOTAL"] == rows["TOTAL"]:
        out.append("rows and objects are being reported as the same population")
    if counts["SEPARATE_LEAF_COUNT_BLOCK"]["IS_THE_SAME_POPULATION_AS_THE_L_M_SCHEDULE"]:
        out.append("the leaf block is being treated as the same population")
    return out


def problems_guide_band(records):
    """A guide category used outside its own area band needs an authority, or the height is not final."""
    out = []
    for r in records:
        if r["TRADE"] != "ALUMINIUM" or not r["AUTHORITY"]:
            continue
        cat = r["DIMENSION_INPUTS"]["GUIDE_CATEGORY"]
        band = WS.TABLE[cat]["AREA_BAND_M2"] if cat in WS.TABLE else None
        area = r["DIMENSION_INPUTS"]["COMPONENT_AREA_M2"]
        outside = band and area is not None and not (band[0] <= area <= band[1])
        by_area = r["DIMENSION_INPUTS"]["CATEGORY_BASIS"] != "EXPLICIT_LABEL_MAP"
        if outside and by_area and (r["AUTHORITY"]["SUPPORTED"] or r["BOQ_INCLUDED"]):
            out.append(f"{r['SUBITEM']} uses {cat} at {area} m2, outside {band}, and is still treated as final")
    return out


def problems_ceiling(records):
    return [r["SUBITEM"] for r in records
            if r["TRADE"] == "CEILING" and (r["STATUS"] == "FINAL_QUANTITY_AVAILABLE" or r["BOQ_INCLUDED"])]


HISTORICAL = {950.22, 57.0, 1375.0, 600.16, 1.82, 1.715}


def problems_historical(records):
    return [f"{r['ITEM_CODE']}={r['MEASURED_QUANTITY']}" for r in records
            if r["MEASURED_QUANTITY"] in HISTORICAL]


def problems_frozen(path, record):
    out = []
    if hashlib.sha256(Path(path).read_bytes()).hexdigest() != DT.FROZEN_FILE_SHA256:
        out.append("the frozen file's bytes changed")
    if record["DIGEST"] != DT.FROZEN_DIGEST:
        out.append("the frozen digest changed")
    if record["GIT_HEAD"] != DT.FROZEN_COMMIT:
        out.append("the frozen commit changed")
    return out


# ------------------------------------------------------------------ 1. floor-finish classification
def test_ar_fl_01_agrees_across_json_workbook_and_boq(export, wb, vals):
    assert problems_ar_fl_01(export, wb, vals) == []


def test_the_check_catches_a_json_that_drifts_from_the_workbook(export, wb, vals):
    broken = copy.deepcopy(export)
    broken["RECORDS"].append(dict(broken["RECORDS"][0], MEASURED_QUANTITY=50.0, TRADE="PORCELAIN_FLOOR"))
    assert problems_ar_fl_01(broken, wb, vals)


def test_no_pending_floor_area_is_classified_as_porcelain(export):
    assert problems_pending_in_porcelain(export["RECORDS"]) == []


def test_the_check_catches_a_stair_slipped_into_porcelain(export):
    broken = copy.deepcopy(export["RECORDS"])
    stair = next(r for r in broken if r["COMPONENT_ROLE"] == "STAIR_OR_LANDING")
    stair["TRADE"] = "PORCELAIN_FLOOR"
    assert problems_pending_in_porcelain(broken)


def test_the_pending_area_is_measured_kept_and_out_of_the_base_boq(export, recon):
    pend = [r for r in export["RECORDS"] if r["TRADE"] == "FLOOR_FINISH_UNCLASSIFIED"]
    assert len(pend) == 33
    assert all(r["STATUS"] == "FINISH_CLASSIFICATION_PENDING" and r["BOQ_INCLUDED"] is False for r in pend)
    assert all(r["ITEM_CODE"] == "AR-FL-PENDING" for r in pend)
    assert abs(export["FLOOR_FINISH"]["PENDING_M2"] - 196.378) < 0.01
    by_role = export["FLOOR_FINISH"]["PENDING_BY_ROLE"]
    assert by_role["STAIR_OR_LANDING"] == {"COMPONENTS": 9, "AREA_M2": 67.563}
    assert by_role["UNNAMED_ON_DRAWING"] == {"COMPONENTS": 24, "AREA_M2": 128.816}
    assert next(l for l in recon["LINES"] if l["ITEM_CODE"] == "AR-FL-PENDING")["AGREES"]


# ------------------------------------------------------------------ 2. entity semantics
def test_the_census_publishes_components_not_rooms(export):
    assert problems_entity_semantics(export) == []
    c = export["ENTITY_SEMANTICS"]
    assert (c["CONNECTED_COMPONENT_COUNT"], c["NON_SLIVER_COMPONENT_COUNT"],
            c["WALL_MATERIAL_OR_SLIVER_COUNT"]) == (135, 64, 71)


def test_the_check_catches_a_component_relabelled_as_a_room(export):
    broken = copy.deepcopy(export)
    broken["RECORDS"][0]["ROOM_REF"] = broken["RECORDS"][0]["COMPONENT_REF"]
    assert problems_entity_semantics(broken)
    miscounted = copy.deepcopy(export)
    miscounted["ENTITY_SEMANTICS"]["NON_SLIVER_COMPONENT_COUNT"] = 135
    assert problems_entity_semantics(miscounted)


def test_the_workbook_says_component_where_it_means_component(wb):
    assert wb["حصر المساحات"]["B1"].value == "رمز المكوّن"
    assert _find(wb["ملخص الحصر"], 1, "مكوّنات متصلة (الإجمالي)")
    assert _find(wb["ملخص الحصر"], 1, "منها مادة جدار / شظية")


# ------------------------------------------------------------------ 3. historical opening counts
def test_rows_and_physical_objects_are_separate_fields():
    assert problems_counts(VA.HISTORICAL_COUNTS) == []
    c = VA.HISTORICAL_COUNTS
    assert c["SCHEDULE_ROW_COUNT"] == {"WINDOWS": 24, "DOORS": 5, "SLIDING_DOORS": 2, "TOTAL": 31}
    assert c["PHYSICAL_OBJECT_COUNT"] == {"WINDOWS": 26, "DOORS": 6, "SLIDING_DOORS": 2, "TOTAL": 34}


def test_the_check_catches_a_row_count_reported_as_an_object_count():
    broken = copy.deepcopy(VA.HISTORICAL_COUNTS)
    broken["PHYSICAL_OBJECT_COUNT"] = dict(broken["SCHEDULE_ROW_COUNT"])
    assert problems_counts(broken)
    mixed = copy.deepcopy(VA.HISTORICAL_COUNTS)
    mixed["PHYSICAL_OBJECT_COUNT"]["WINDOWS"] = 24
    mixed["PHYSICAL_OBJECT_COUNT"]["TOTAL"] = 32
    assert problems_counts(mixed)


# ------------------------------------------------------------------ 4. skirting
def test_skirting_is_three_buckets_and_not_one_decision_ready_quantity(export):
    b = export["SKIRTING_BUCKETS"]
    assert b["IS_ONE_DECISION_READY_QUANTITY"] is False
    assert b["DRY_NAMED_INTERNAL_CANDIDATE"] == {"COMPONENTS": 16, "LENGTH_M": 381.535}
    assert b["STAIR_OR_LANDING_PENDING"] == {"COMPONENTS": 9, "LENGTH_M": 97.355}
    assert b["UNNAMED_SPACE_PENDING"] == {"COMPONENTS": 24, "LENGTH_M": 349.603}
    assert abs(b["TOTAL_LENGTH_M"] - 828.493) < 1e-6
    rows = [r for r in export["RECORDS"] if r["TRADE"] == "SKIRTING"]
    assert len(rows) == 49
    assert all(r["STATUS"] == "DERIVED_NOT_IN_FROZEN_TAKEOFF" and r["BOQ_INCLUDED"] is False for r in rows)
    assert all(r["VERIFIED_AGAINST_FROZEN"] is False for r in rows)


def test_every_skirting_line_sits_in_the_provisional_section_of_the_boq(wb):
    ws = wb["BOQ حسب البند"]
    for code in ("AR-SK-01", "AR-SK-02", "AR-SK-03", "AR-SK-04"):
        r = _find(ws, 2, code)
        assert ws.cell(r, 1).value == "PROVISIONAL"
        assert ws.cell(r, 14).value == "لا"


# ------------------------------------------------------------------ 5. ceilings
def test_no_unresolved_ceiling_finish_is_final(export):
    assert problems_ceiling(export["RECORDS"]) == []
    ceil = [r for r in export["RECORDS"] if r["TRADE"] == "CEILING"]
    assert len(ceil) == 62
    assert all(r["FINISH_STATUS"] == "FINISH_CLASSIFICATION_PENDING" for r in ceil)


def test_the_check_catches_a_ceiling_marked_final(export):
    broken = copy.deepcopy(export["RECORDS"])
    next(r for r in broken if r["TRADE"] == "CEILING")["STATUS"] = "FINAL_QUANTITY_AVAILABLE"
    assert problems_ceiling(broken)


def test_the_ceiling_area_is_still_measured_and_still_final_as_a_measurement(wb, vals, frozen):
    total = cell(wb, vals, "الأسقف", f"E{wb['الأسقف'].max_row}")
    assert abs(total - frozen["TOTALS"]["CEILING_M2"]) < DT.CROSS_TOL_M2
    assert wb["الأسقف"].cell(wb["الأسقف"].max_row, 6).value == "FINAL_QUANTITY_AVAILABLE"


# ------------------------------------------------------------------ 6. the window whose category has no authority
def test_no_guide_category_is_used_outside_its_band_without_authority(export):
    assert problems_guide_band(export["RECORDS"]) == []


def test_the_check_catches_a_band_breach_that_is_treated_as_final(export):
    broken = copy.deepcopy(export["RECORDS"])
    w = next(r for r in broken if r["SUBITEM"] == "AR-W-05")
    w["AUTHORITY"]["SUPPORTED"] = True
    w["BOQ_INCLUDED"] = True
    assert problems_guide_band(broken)


def test_ar_w_05_asks_the_owner_rather_than_publishing_a_guide_height(export):
    w = next(r for r in export["RECORDS"] if r["SUBITEM"] == "AR-W-05")
    assert w["STATUS"] == "OWNER_INPUT_REQUIRED" and w["BOQ_INCLUDED"] is False
    assert w["ITEM_CODE"] == "AR-AL-PENDING"
    assert w["AUTHORITY"]["AUTHORITY"] == "NONE" and w["AUTHORITY"]["SUPPORTED"] is False
    assert w["AUTHORITY"]["BAND_M2"] == [35, 60] or tuple(w["AUTHORITY"]["BAND_M2"]) == (35, 60)
    assert "ASK_THE_OWNER" in w["DIMENSION_INPUTS"]["HEIGHT_SOURCE"]


def test_the_master_bedroom_category_rests_on_the_label_not_the_area(export):
    w = next(r for r in export["RECORDS"] if r["SUBITEM"] == "AR-W-04")
    assert w["DIMENSION_INPUTS"]["CATEGORY_BASIS"] == "EXPLICIT_LABEL_MAP"
    assert w["AUTHORITY"]["AUTHORITY"] == "EXPLICIT_LABEL_ON_THE_DRAWING"
    assert "label" in w["AUTHORITY"]["WHY"]
    assert w["BOQ_INCLUDED"] is True


def test_the_two_aluminium_lines_add_back_to_the_frozen_total(export, frozen):
    s = export["ALUMINIUM_SPLIT"]
    assert abs(s["SUPPORTED_M2"] + s["UNSUPPORTED_M2"] - frozen["TOTALS"]["ALUMINIUM_M2"]) < 1e-6
    assert s["UNSUPPORTED_WINDOWS"] == ["AR-W-05"]


# ------------------------------------------------------------------ 7. rule provenance
def test_us18_is_claimed_only_where_it_applies(export):
    wet = {"WET_ROOM", "KITCHEN"}
    for r in export["RECORDS"]:
        if r["RULE_ID"] == "US-18" and r["COMPONENT_ROLE"] is not None:
            assert r["COMPONENT_ROLE"] in wet, r["ITEM_CODE"]
    dry = [r for r in export["RECORDS"]
           if r["TRADE"] == "PORCELAIN_FLOOR" and r["COMPONENT_ROLE"] == "INTERNAL_ROOM"]
    assert dry and all(r["RULE_ID"] is None and r["RULE_LEVEL"] is None for r in dry)


def test_rule_ids_match_between_the_workbook_and_the_json(wb, export):
    json_rules = {r["COMPONENT_REF"]: r["RULE_ID"] for r in export["RECORDS"]
                  if r["TRADE"] in ("PORCELAIN_FLOOR", "FLOOR_FINISH_UNCLASSIFIED")}
    ws, seen = wb["الأرضيات"], 0
    for r in range(2, ws.max_row + 1):
        ref, rule = ws.cell(r, 2).value, ws.cell(r, 9).value
        if ref in json_rules and rule:
            assert json_rules[ref] == (None if rule == "—" else rule), ref
            seen += 1
    assert seen == 62


# ------------------------------------------------------------------ 8 and 9. QA and precision
def test_every_quality_check_is_evaluated_evidence_and_passes(export):
    q = export["QA"]
    assert q["ALL_PASS"] and q["PASSED"] == q["OF"] == 22
    for c in q["CHECKS"]:
        assert c["INPUTS"] and c["RESULT"] is not None and c["METHOD"]
        assert c["COMMIT"] and c["EVALUATED_AT_UTC"]
        assert "TOLERANCE" in c


def test_component_areas_recompute_from_raw_dimensions_within_tolerance(export):
    worst = 0.0
    for r in export["RECORDS"]:
        if not r["COMPONENTS"]:
            continue
        s = sum(p["RAW_LENGTH_M"] * p["RAW_WIDTH_M"] for p in r["COMPONENTS"])
        worst = max(worst, abs(s - r["MEASURED_QUANTITY"]))
    assert worst <= 5e-5, worst
    for r in export["RECORDS"]:
        for p in r["COMPONENTS"] or []:
            assert abs(p["RAW_LENGTH_M"] * p["RAW_WIDTH_M"] - p["RAW_AREA_M2"]) <= DT.AREA_TOL_M2
            assert p["DISPLAY_LENGTH_M"] == round(p["RAW_LENGTH_M"], 4)


def test_the_workbook_arithmetic_is_the_frozen_arithmetic(wb, vals, frozen):
    ws = wb["حصر المساحات"]
    frozen_area = {r["ROOM_REF"]: r["AREA_M2"] for f in frozen["FLOORS"] for r in f["ROOMS"]}
    seen = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 5).value != "إجمالي المكوّن":
            continue
        ref = ws.cell(row, 2).value
        assert abs(vals[(ws.title, f"H{row}")] - frozen_area[ref]) < 5e-5, ref
        seen += 1
    assert seen == 135


# ------------------------------------------------------------------ 10. the workbook as delivered
def test_the_workbook_keeps_its_thirteen_right_to_left_sheets(wb):
    assert wb.sheetnames == DT.SHEETS and len(wb.sheetnames) == 13
    assert all(wb[n].sheet_view.rightToLeft for n in wb.sheetnames)


def test_every_formula_carries_a_cached_result_and_none_is_an_error(wb, vals):
    cached = load_workbook(XL, data_only=True)
    numeric = missing = 0
    for (s, c), v in vals.items():
        got = cached[s][c].value
        if isinstance(v, (int, float)):
            numeric += 1
            if not isinstance(got, (int, float)):
                missing += 1
        if isinstance(got, str):
            assert not got.startswith("#"), f"{s}!{c} holds {got}"
    assert numeric > 800 and missing == 0


def test_pricing_columns_stay_blank_and_separate(wb, vals, export):
    ws = wb["BOQ حسب البند"]
    for r in range(2, ws.max_row + 1):
        for col in (7, 9):                       # waste %, unit rate
            assert ws.cell(r, col).value is None
        for col in (8, 10):                      # procurement quantity, amount
            v = ws.cell(r, col).value
            assert v is None or str(v).startswith("=IF")
    for r in export["RECORDS"]:
        assert r["WASTE_PERCENT"] is None and r["PROCUREMENT_QUANTITY"] is None
        assert r["UNIT_RATE"] is None and r["AMOUNT"] is None
        assert r["APPROVAL_STATUS"] == "DRAFT"


def test_procurement_and_amount_stay_blank_until_a_waste_and_a_rate_exist(wb, vals):
    ws = wb["BOQ حسب البند"]
    row = _find(ws, 2, "AR-BL-01")
    assert vals[(ws.title, f"H{row}")] is WC.BLANK
    assert vals[(ws.title, f"J{row}")] is WC.BLANK


# ------------------------------------------------------------------ the frozen artifact
def test_no_historical_figure_reached_a_quantity_field(export):
    assert problems_historical(export["RECORDS"]) == []
    assert export["BASED_ON"]["FILE_SHA256"] == DT.FROZEN_FILE_SHA256


def test_the_check_catches_a_historical_figure_used_as_a_quantity(export):
    broken = copy.deepcopy(export["RECORDS"])
    broken[0]["MEASURED_QUANTITY"] = 950.22
    assert problems_historical(broken)


def test_the_frozen_takeoff_is_byte_for_byte_unchanged(frozen, rec):
    assert problems_frozen(FROZEN, frozen) == []
    assert rec["FROZEN_UNCHANGED"]["MATCHES"] is True
    assert rec["FROZEN_UNCHANGED"]["FILE_SHA256_AFTER"] == DT.FROZEN_FILE_SHA256


def test_the_check_catches_a_changed_frozen_file(tmp_path, frozen):
    p = tmp_path / "frozen.json"
    p.write_text(json.dumps(frozen) + " ", "utf-8")
    assert problems_frozen(p, dict(frozen, DIGEST="deadbeef"))


# ------------------------------------------------------------------ the reconciliation
def test_every_reconciliation_line_agrees_inside_its_tolerance(recon):
    assert recon["ALL_AGREE"]
    assert recon["FROZEN"]["SHA256_MATCHES"] and recon["FROZEN"]["REWRITTEN"] is False
    assert recon["ALUMINIUM_SPLIT"]["EQUALS_FROZEN_TOTAL"]
    for line in recon["LINES"]:
        assert line["WORST_DIFFERENCE"] <= line["TOLERANCE"], line["ITEM_CODE"]
