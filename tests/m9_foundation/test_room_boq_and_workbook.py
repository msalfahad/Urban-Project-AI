"""M9 S1 room BOQ and the room-by-room workbook on the synthetic room schedule: net lines per room, missing data
blocked with its question (never invented), material results traced to their lines, and the workbook formulas -
recalculated by LibreOffice - reconcile room = floor = project, read PARTIAL where a room is unresolved, and balance."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest

from engine.source import boq_line as BL
from engine.source import coverage_census as CS
from engine.source import room_boq as RB
from engine.source import material_result as WP

ROOT = Path(__file__).resolve().parents[2]
SCHED = ROOT / "research/urban_m9_s1/SYNTHETIC_ROOMS_S1.json"


@pytest.fixture(scope="module")
def built():
    return RB.build(SCHED)


def _by(built, room, code):
    return [ln for ln in built["lines"] if ln["location"]["room_id"] == room and ln["item_code"] == code]


def test_room_lines_and_their_states(built):
    assert len(built["lines"]) == 12
    assert _by(built, "GF-BATH-01", "FIN.FLOOR.WET_TILE")[0]["measured_net"] == pytest.approx(6.0)
    assert _by(built, "GF-LIV-01", "FIN.FLOOR.PORCELAIN")[0]["measured_net"] == pytest.approx(5.4 * 4.2)
    kit = _by(built, "GF-KIT-01", "FIN.FLOOR.PORCELAIN")[0]
    assert kit["measured_net"] == pytest.approx(4.0 * 2.5 + 2.0 * 1.0) and kit["evidence_state"] == "SOURCE_DERIVED"
    store = _by(built, "1F-STORE-01", "FIN.FLOOR.PORCELAIN")[0]
    assert store["evidence_state"] == "BLOCKED_RFI" and store["measured_net"] is None
    assert store["rfi"]["question"].startswith("ROOM_GEOMETRY_NOT_ESTABLISHED")
    wp = {ln["item_code"]: ln["measured_net"] for ln in built["lines"] if ln["location"]["room_id"] == "1F-BATH-02"
          and ln["item_code"].startswith("FIN.WP")}
    assert wp == pytest.approx({"FIN.WP.FLOOR": 7.36, "FIN.WP.DOORWAY": 0.18, "FIN.WP.UPTURN": 2.02})
    for ln in built["lines"]:
        BL.validate(ln)
        assert ln["approval_state"] == "MACHINE_ONLY" and ln["source"]["geometry_source"] == "SYNTHETIC_FIXTURE"
        assert ln["rule"]["rule_set"] == "1.5.0"
        if ln["item_code"] == "FIN.WP.UPTURN":
            assert ln["rule"]["rule_id"] == "UP-WP-002" and ln["rule"]["version"] == "2.0.0"


def test_material_results_trace_to_their_lines_and_never_change_the_net(built):
    lines = {ln["line_id"]: ln for ln in built["lines"]}
    for m in built["materials"]:
        ln = lines[m["line_id"]]
        if m["method"].startswith("TILE"):
            assert WP.trace(ln, m) == []
        elif m["state"] != "BLOCKED":
            for c in m["components"]:
                assert sum(lines[i]["measured_net"] for i in c["line_ids"]) == pytest.approx(c["net_m2"])
    kit = next(m for m in built["materials"] if m["location"].get("room_id") == "GF-KIT-01")
    assert kit["state"] == "BLOCKED" and kit["consumed_m2"] is None
    reasons = " ".join(u["reason"] for u in built["unresolved"])
    for word in ("ROOM_GEOMETRY_NOT_ESTABLISHED", "POLYGON_OPTIMIZER_NOT_IMPLEMENTED", "PACKAGING_NOT_ESTABLISHED",
                 "DOORWAY_UPTURN_DETAIL_REQUIRED", "CORNER_TREATMENT_NOT_SPECIFIED"):
        assert word in reasons, word


def test_comparisons_are_recommendations_only(built):
    rows = [c for c in built["comparisons"] if c["room_id"] == "GF-BATH-01"]
    assert {c["spec_id"] for c in rows} == {"TILE-120x60", "TILE-60x60", "TILE-160x80"}
    sel = [c for c in rows if c["role"].startswith("SELECTED")]
    assert len(sel) == 1 and sel[0]["spec_id"] == "TILE-120x60"
    mat = next(m for m in built["materials"] if m["location"].get("room_id") == "GF-BATH-01" and m["method"].startswith("TILE"))
    assert mat["material_spec"]["spec_id"] == "TILE-120x60"


def test_missing_data_never_invents_a_result():
    s = RB.load(SCHED)
    s["assumptions"].pop("joint_m")
    b = RB.build(s)
    tiles = [m for m in b["materials"] if m["method"].startswith("TILE")]
    assert all(m["state"] == "BLOCKED" for m in tiles)
    assert any("JOINT_WIDTH_NOT_ESTABLISHED" in (m["blocked_by"] or "") for m in tiles)
    chk = RB.check(SCHED)
    flags = {r["room_id"]: r["flags"] for r in chk["rooms"]}
    assert "ROOM_GEOMETRY_NOT_ESTABLISHED" in flags["1F-STORE-01"] and chk["note"].startswith("diagnostic only")


def _recalc(path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        pytest.skip("LibreOffice not installed: the formula values cannot be recalculated here")
    td = tempfile.mkdtemp()
    out = Path(td) / "out"
    out.mkdir()
    subprocess.run([exe, "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile", "--convert-to",
                    "xlsx", "--outdir", str(out), str(path)], capture_output=True, timeout=240)
    f = out / Path(path).name
    if not f.exists():
        pytest.fail("LibreOffice produced no file")
    return f


@pytest.fixture(scope="module")
def workbook(built, tmp_path_factory):
    from engine import boq_room_workbook as WB
    d = tmp_path_factory.mktemp("wb")
    rows = CS.census(built["lines"], rooms_run=True, material_results=built["materials"])
    meta = {"project_id": built["project_id"], "rule_set": "1.5.0", "rule_set_digest": built["rule_set_digest"]}
    from datetime import datetime
    a = WB.write(d / "a.xlsx", room_boq=built, census_rows=rows, meta=meta, created=datetime(2026, 1, 1))
    b = WB.write(d / "b.xlsx", room_boq=built, census_rows=rows, meta=meta, created=datetime(2026, 1, 1))
    return {"a": a, "b": b, "path": d / "a.xlsx", "WB": WB}


def test_workbook_sheets_headers_font_and_determinism(workbook):
    from openpyxl import load_workbook
    WB = workbook["WB"]
    assert workbook["a"]["sha256"] == workbook["b"]["sha256"]
    wb = load_workbook(workbook["path"])
    need = ["FLOOR PORCELAIN", "WALL PORCELAIN", "PLASTER & PAINT", "WATERPROOFING", "CEILINGS", "SKIRTING", "MARBLE",
            "STAIRS & RAILINGS", "DOORS & WINDOWS", "ALUMINIUM", "STRUCTURAL", "PROCUREMENT", "UNRESOLVED"]
    assert all(n in wb.sheetnames for n in need)
    ws = wb["FLOOR PORCELAIN"]
    head = [c.value for c in ws[3]]
    assert head[:11] == ["FLOOR", "ROOM (AR/EN)", "ITEM", "MATERIAL SIZE", "NET M²", "WASTE %", "WASTE M²", "MATERIAL M²",
                         "BOXES/ROLLS", "SOURCE", "STATUS"]
    assert head == WB.HEAD
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if c.value is not None:
                    assert c.font.name == "Arial", (w.title, c.coordinate)
                if isinstance(c.value, str):
                    low = c.value.lower()
                    assert "kwd" not in low and "unit rate" not in low and "kg/m3 ratio" not in low.replace(
                        "never a kg/m3 ratio", "").replace("no kg/m3 ratio", "")
    for name in ("WALL PORCELAIN", "CEILINGS", "MARBLE"):
        for row in wb[name].iter_rows(min_row=4):
            assert row[WB.C["NET M²"] - 1].value is None and row[WB.C["MATERIAL M²"] - 1].value is None


def test_workbook_formulas_reconcile_room_floor_project(workbook, built):
    from openpyxl import load_workbook
    WB = workbook["WB"]
    f = _recalc(workbook["path"])
    wb = load_workbook(f, data_only=True)
    errs = [(w.title, c.coordinate) for w in wb.worksheets for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("#") and c.value.endswith(("!", "?", "A", "0"))]
    assert not errs, errs[:5]
    exp = WB.expected_totals(built)
    for sheet, key in (("FLOOR PORCELAIN", "floor_finish"), ("WATERPROOFING", "waterproofing")):
        ws = wb[sheet]
        g = lambda r, h: ws.cell(r, WB.C[h]).value                                         # noqa: E731
        rows = range(4, ws.max_row + 1)
        for r in rows:
            kind = g(r, "ROW TYPE")
            if kind in ("DETAIL", "MATERIAL") and g(r, "BALANCE CHECK M²") is not None:
                assert abs(g(r, "BALANCE CHECK M²")) < 1e-6, (sheet, r)
        details = [r for r in rows if g(r, "ROW TYPE") == "DETAIL"]
        open_rooms = {g(r, "ROOM ID") for r in details if str(g(r, "STATUS")).startswith("BLOCKED_RFI")}
        room_of = {}
        for r in rows:
            if g(r, "ROW TYPE") == "DETAIL":
                room_of.setdefault(g(r, "ROOM ID"), g(r, "FLOOR"))
        last_room = None
        for r in rows:
            kind = g(r, "ROW TYPE")
            if kind in ("DETAIL", "MATERIAL"):
                last_room = g(r, "ROOM ID")
            if kind == "ROOM_SUBTOTAL":
                v = g(r, "NET M²")
                if last_room in open_rooms:
                    assert isinstance(v, str) and v.startswith("PARTIAL")
                else:
                    assert v == pytest.approx(exp[key][("ROOM", last_room)])
            if kind == "FLOOR_SUBTOTAL":
                fl = g(r, "FLOOR")
                v = g(r, "NET M²")
                if any(room_of[x] == fl for x in open_rooms):
                    assert isinstance(v, str) and v.startswith("PARTIAL") and f"{exp[key][('FLOOR', fl)]:.3f}" in v
                else:
                    assert v == pytest.approx(exp[key][("FLOOR", fl)])
                    assert v == pytest.approx(sum(exp[key][("ROOM", x)] for x, f2 in room_of.items() if f2 == fl))
            if kind == "PROJECT_TOTAL":
                v = g(r, "NET M²")
                if open_rooms:
                    assert isinstance(v, str) and f"{exp[key][('PROJECT', '')]:.3f}" in v
                else:
                    assert v == pytest.approx(exp[key][("PROJECT", "")])
    ws = wb["WATERPROOFING"]
    mats = [r for r in range(4, ws.max_row + 1) if ws.cell(r, WB.C["ROW TYPE"]).value == "MATERIAL"]
    rolls = sum(ws.cell(r, WB.C["BOXES/ROLLS"]).value for r in mats)
    assert rolls == sum(m["units"]["consumed_units"] for m in built["materials"] if m["method"].startswith("MEMBRANE"))
