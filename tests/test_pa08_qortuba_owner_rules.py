"""OWNER RULES V1: the rules that keep an owner's decision from quietly becoming a measurement.

An owner's rule changes what a number means, never where a line is.  These tests hold that separation: the geometry that
came out of the frozen takeoff is the geometry that goes into every quantity below, a temporary default is never allowed
to pass for a source dimension, and an opening whose type is still unknown never reaches a payable figure.
"""

from __future__ import annotations

import json
from pathlib import Path

from research.qs_wall_treatment_01 import protocol as PR
from research.qs_wall_treatment_01.pa08.qortuba.boq import (opening_source_search as OS, owner_rules as OR,
                                                          workbook_boq as WB)

OUT = Path(PR.OUT_DIR) / "pa08_qortuba_boq"
QS = Path(PR.OUT_DIR) / "pa08_qortuba_qs01"


def reg(name, folder=OUT):
    return json.loads((folder / f"{name}.json").read_text("utf-8"))


def rules():
    return reg("URBAN_OWNER_RULES_V1")


def quants():
    return reg("QORTUBA_RECALCULATED_QUANTITIES_V1")


def by_id():
    return {x["QUANTITY_ID"]: x for x in quants()["ROWS"]}


# ------------------------------------------------------------------ §A the ladder
def test_the_priority_ladder_is_stored_in_order_and_precedent_sits_outside_it():
    R = rules()
    assert [p["RANK"] for p in R["PRIORITY"]] == [1, 2, 3, 4, 5]
    assert R["PRIORITY"][0]["LEVEL"] == "PROJECT_DRAWING_OR_SPECIFICATION"
    assert R["PRIORITY"][-1]["LEVEL"] == "UNKNOWN_ASK_OWNER"
    assert "NO rung" in R["WHERE_HISTORICAL_PRECEDENT_SITS"]


def test_only_an_urban_standard_applies_by_itself():
    R = rules()
    for x in R["URBAN_STANDARDS"]:
        assert x["RULE_LEVEL"] == "URBAN_STANDARD" and x["APPLIES_AUTOMATICALLY_BELOW_DRAWINGS"] is True
    for x in R["QORTUBA_PROJECT_RULES"]:
        assert x["RULE_LEVEL"] == "QORTUBA_PROJECT_RULE" and x["APPLIES_TO_OTHER_PROJECTS"] is False
    for x in R["TEMPORARY_DEFAULTS"]:
        assert x["RULE_LEVEL"] == "APPROVED_TEMPORARY_DEFAULT" and x["IS_SOURCE_TRUTH"] is False
    # §T, carried in the precedent register: a candidate is not a standard
    for s in reg("QORTUBA_RULE_SAFETY_REGISTER")["ROWS"]:
        assert s["MAY_BE_APPLIED_TO_QORTUBA_WITHOUT_OWNER"] is False


def test_the_four_three_metre_heights_are_project_rules_and_not_company_policy():
    names = {x["PARAMETER"]: x for x in rules()["QORTUBA_PROJECT_RULES"]}
    for p in ("QORTUBA_WALL_TILE_HEIGHT", "QORTUBA_BLOCKWORK_HEIGHT", "QORTUBA_INTERNAL_PLASTER_HEIGHT",
              "QORTUBA_INTERNAL_PAINT_HEIGHT"):
        assert names[p]["VALUE"] == 3.00 and names[p]["RULE_LEVEL"] == "QORTUBA_PROJECT_RULE"
    titles = {x["TITLE"] for x in rules()["URBAN_STANDARDS"]}
    assert not any("HEIGHT" in t and "3" in t for t in titles), "a project height must not become an Urban standard"


def test_the_superseded_precedents_are_named_and_demoted_not_deleted():
    sup = {x["RULE_ID"]: x for x in rules()["SUPERSEDED_HISTORICAL_RULES"]}
    for rid in ("R-02", "R-07", "R-09", "R-13"):
        assert rid in sup, rid
        assert sup[rid]["NEW_FILE_STATUS"] == "CONTRACTOR_SPECIFIC_REFERENCE_ONLY"
        assert sup[rid]["SUPERSEDED_BY"]
    still = {r["RULE_ID"] for r in reg("URBAN_BOQ_RULE_REGISTRY")["ROWS"]}
    assert set(sup) <= still, "a superseded rule stays on file"


# ------------------------------------------------------------------ §G the default that must not spread
def test_no_default_height_reaches_a_window_a_sliding_door_or_an_unknown():
    o = reg("QORTUBA_OPENING_REGISTER")
    assert set(OR.NO_DEFAULT_FOR) <= set(o["TYPES"])
    from research.qs_wall_treatment_01.pa08.qortuba.boq import opening_source_search as OS
    owner = set(OS.OWNER_SUPPLIED_HEIGHTS)
    for x in o["ROWS"]:
        if x["HEIGHT_M"] is None:
            assert x["HEIGHT_STATE"] == "NOT_ESTABLISHED"
            continue
        if x["OPENING_ID"] in owner:
            continue  # an owner override is a real dimension, not a default, and may reach any type
        if (x["OPENING_ID"] in OS.OWNER_SUPPLIED_PASSAGE_SUBTYPES
                or x["OPENING_ID"] in OS.OWNER_CLOSED_AS_MEASUREMENT_OPENINGS):
            continue  # so is a passage or full-height opening the owner closed
        # otherwise a height exists only on an ordinary door, and only from the owner's default
        assert x["TYPE"] == "DOOR" and x["HEIGHT_M"] == 2.20, x["OPENING_ID"]
        assert x["HEIGHT_STATE"] == "TEMPORARY_OWNER_DEFAULT"
    for x in o["ROWS"]:
        if x["OPENING_ID"] in owner:
            continue
        assert not (x["TYPE"] in OR.NO_DEFAULT_FOR and x["HEIGHT_M"] is not None), x["OPENING_ID"]


def test_a_source_width_is_never_replaced_by_the_default_width():
    d = {x["PARAMETER"]: x for x in rules()["TEMPORARY_DEFAULTS"]}
    assert d["DEFAULT_DOOR_WIDTH"]["VALUE"] == 1.00
    for x in reg("QORTUBA_OPENING_REGISTER")["ROWS"]:
        assert x["WIDTH_SOURCE"], x["OPENING_ID"]
        assert "DEFAULT" not in x["WIDTH_SOURCE"]
    assert reg("QORTUBA_OPENING_REGISTER")["WIDTHS_ARE_SOURCE_ESTABLISHED"] is True


def test_every_quantity_that_consumed_the_default_says_so():
    for x in quants()["ROWS"]:
        if x["USES_TEMPORARY_DEFAULT"]:
            assert "TD-02" in (x["TEMPORARY_DEFAULT_WARNING"] or "")
            assert "TD-02" in x["RULE_ID"] or "TD-02" in x["PARAMETER_SOURCE"]
        else:
            assert x["TEMPORARY_DEFAULT_WARNING"] is None


def test_an_unknown_type_never_enters_a_final_quantity():
    unknown = {o["OPENING_ID"] for o in reg("QORTUBA_OPENING_REGISTER")["ROWS"] if o["TYPE"] == "UNKNOWN"}
    for x in quants()["ROWS"]:
        if x["STATUS"] == "FINAL_QUANTITY_AVAILABLE" and x["UNIT"] == "M2":
            named = {o.get("OPENING_ID") for o in x["RESIDUAL_OPENINGS"]}
            assert not (named & unknown), x["QUANTITY_ID"]


# ------------------------------------------------------------------ §E the unit rule
def test_skirting_deducts_opening_widths_in_metres_and_nothing_else():
    for x in reg("QORTUBA_ROOM_FINISH_AND_SKIRTING_RECALC")["ROWS"]:
        if x["FINISH_CLASS"] == "CERAMIC_SERVICE_ROOM":
            assert x["SKIRTING_LM"] == 0.0, x["ROOM"]
            continue
        want = round(x["GROSS_WALL_LINE_LM"] - x["DOOR_OPENING_LM"] - x["GLAZED_OPENING_LM"]
                     - x["OPEN_PASSAGE_LM"], 3)
        assert abs(x["SKIRTING_LM"] - want) < 1e-9, (x["ROOM"], x["SKIRTING_LM"], want)
        # a column face is not an opening, so it stays on the path.  An open passage is the one thing the frozen
        # boundary traced as continuous WALL FACE that is not wall at all, so it is the only permitted shortfall.
        assert x["SKIRTING_LM"] >= x["WALL_FACE_LM"] - x["OPEN_PASSAGE_LM"] - 1e-9


def test_no_area_was_ever_subtracted_from_a_linear_quantity():
    for x in quants()["ROWS"]:
        if x["UNIT"] == "LM":
            assert "m2" not in x["FORMULA"], x["QUANTITY_ID"]


def test_the_hidden_profile_equals_the_hidden_skirting_exactly():
    q = by_id()
    assert q["Q-01"]["MEASURED_NET_QUANTITY"] == q["Q-02"]["MEASURED_NET_QUANTITY"]
    assert q["Q-01"]["UNIT"] == q["Q-02"]["UNIT"] == "LM"
    assert q["Q-01"]["TRADE"] != q["Q-02"]["TRADE"], "two BOQ rows, never merged"
    rr = {x["RULE_ID"]: x for x in rules()["COMMERCIAL_RATE_RULES"]}
    assert rr["US-08"]["FACTOR"] == 1.0 and rr["US-08"]["AMOUNTS_MERGED"] is False
    assert rr["US-09"]["FACTOR"] == 0.50 and rr["US-09"]["CHANGES_QUANTITY"] is False


def test_a_fully_ceramic_room_carries_no_skirting_and_no_profile():
    rooms = reg("QORTUBA_ROOM_FINISH_AND_SKIRTING_RECALC")
    assert "PAINTRY" in rooms["CERAMIC_SERVICE_ROOMS"] and rooms["CERAMIC_SERVICE_ROOMS"].count("BATH") == 3
    for x in rooms["ROWS"]:
        if x["FINISH_CLASS"] == "CERAMIC_SERVICE_ROOM":
            assert x["SKIRTING_LM"] == 0.0 and x["HIDDEN_PROFILE_LM"] == 0.0, x["ROOM"]


# ------------------------------------------------------------------ §N the membrane
def test_the_upturn_runs_through_the_doorway():
    q = by_id()["Q-04"]
    sk = reg("QORTUBA_QS01_SKIRTING_TAKEOFF", QS)["ROWS"]
    gross = round(sum(x["GROSS_WALL_LINE_PERIMETER_LM"] for x in sk if x["WET_OR_DRY"] == "WET"), 3)
    assert abs(q["MEASURED_NET_QUANTITY"] - gross) < 1e-9
    assert q["MEASURED_NET_QUANTITY"] > q["SUPERSEDES"]["PREVIOUS"], "the doorway must be restored, not deducted"
    doors = round(sum(x["DOOR_WIDTH_DEDUCTION_LM"] for x in sk if x["WET_OR_DRY"] == "WET"), 3)
    assert abs(q["MEASURED_NET_QUANTITY"] - q["SUPERSEDES"]["PREVIOUS"] - doors) < 1e-9


def test_the_upturn_area_is_reference_only_and_never_the_boq_quantity():
    q = by_id()
    a, lm = q["Q-04R"], q["Q-04"]
    assert abs(a["MEASURED_NET_QUANTITY"] - round(lm["MEASURED_NET_QUANTITY"] * OR.UPTURN_HEIGHT, 4)) < 1e-9
    assert "REFERENCE ONLY" in a["NOTE"] and a["UNIT"] == "M2" and lm["UNIT"] == "LM"


# ------------------------------------------------------------------ §F/§I/§J/§K the full deduction and the reveals
def test_every_opening_deduction_is_the_full_area_and_never_a_half():
    for w in reg("QORTUBA_BLOCKWORK_RECALC")["ROWS"]:
        if not w["BLOCKWORK_CONFIRMED"]:
            assert w["OPENING_DEDUCTION_M2"] is None, w["WALL_ID"]
            continue
        want = sum(o["W"] * o["H"] for o in w["OPENINGS_DEDUCTED"])
        assert abs(w["OPENING_DEDUCTION_M2"] - round(want, 4)) < 1e-9, w["WALL_ID"]
    for x in quants()["ROWS"]:
        # a half-opening convention would show as an explicit halving, not as any digits that happen to read 0.5
        f = x["FORMULA"].lower()
        assert "half" not in f and "x 0.5" not in f and "/ 2" not in f, x["QUANTITY_ID"]


def test_blockwork_adds_no_reveal_back():
    for x in quants()["ROWS"]:
        if x["BOQ_ITEM"].startswith("BLOCKWORK"):
            assert "reveal" not in x["FORMULA"].lower(), x["QUANTITY_ID"]


def test_the_plaster_reveal_is_three_sided_at_a_quarter_metre():
    assert OR.REVEAL_DEPTH == 0.25
    q = by_id()["Q-08"]
    # three sides where a head exists, two where it does not: the formula names the rule rather than one height
    assert "0.25 x (2 x height" in q["FORMULA"] and "width where a head exists" in q["FORMULA"]
    assert "sill" not in q["FORMULA"].lower()
    assert "US-07" in q["RULE_ID"]


def test_plaster_and_paint_cover_the_same_dry_faces():
    q = by_id()
    assert q["Q-08"]["MEASURED_NET_QUANTITY"] == q["Q-09"]["MEASURED_NET_QUANTITY"]
    assert q["Q-08"]["TRADE"] != q["Q-09"]["TRADE"]
    assert "by construction" in q["Q-08"]["NOTE"]


def test_a_ceramic_face_leaves_plaster_and_paint_and_joins_tile_preparation():
    q = by_id()
    rooms = {r["ROOM_NAME"] for r in reg("QORTUBA_ROOM_FINISH_AND_SKIRTING_RECALC")["ROWS"]
             if r["FINISH_CLASS"] == "CERAMIC_SERVICE_ROOM"}
    assert not (set(q["Q-08"]["ROOMS"]) & rooms) and not (set(q["Q-09"]["ROOMS"]) & rooms)
    assert set(q["Q-10"]["ROOMS"]) == rooms and "US-03" in q["Q-10"]["RULE_ID"]


# ------------------------------------------------------------------ §U partial means partial
def test_a_partial_quantity_carries_a_value_and_names_what_is_missing():
    for x in quants()["ROWS"]:
        if x["STATUS"] == "PARTIALLY_CALCULATED":
            # partial means a value plus a named reservation: either an opening with no height, or a temporary
            # default still standing in for one
            assert x["MEASURED_NET_QUANTITY"] is not None, x["QUANTITY_ID"]
            assert x["RESIDUAL_OPENINGS"] or x["USES_TEMPORARY_DEFAULT"], x["QUANTITY_ID"]
        if x["STATUS"] == "FINAL_QUANTITY_AVAILABLE":
            assert x["RESIDUAL_OPENINGS"] == [], x["QUANTITY_ID"]
            assert x["MEASURED_NET_QUANTITY"] is not None


def test_a_missing_height_blocks_only_the_rows_that_depend_on_it():
    q = quants()
    finals = {x["QUANTITY_ID"] for x in q["ROWS"] if x["STATUS"] == "FINAL_QUANTITY_AVAILABLE"}
    assert {"Q-03", "Q-04", "Q-11", "Q-12"} <= finals, "floor and perimeter items need no height"
    by = by_id()
    # the skirting pair is fixed by QP-09 and was never held by a height
    for qid in ("Q-01", "Q-02"):
        assert by[qid]["STATUS"] == "FINAL_QUANTITY_AVAILABLE"
        assert by[qid]["MEASURED_NET_QUANTITY"] == 76.389
        assert by[qid]["USES_TEMPORARY_DEFAULT"] is False
        assert by[qid]["RESIDUAL_OPENINGS"] == [], "no unheighted opening may hold up a linear quantity"
    assert q["BY_STATUS"]["FINAL_QUANTITY_AVAILABLE"] > 0 and q["BY_STATUS"]["PARTIALLY_CALCULATED"] > 0


def test_every_window_now_has_a_host_room_so_none_is_smeared_everywhere():
    """The audit joined each glazed element to its host band, so no window blocks a room it does not stand in."""
    op = reg("QORTUBA_OPENING_REGISTER")
    assert op["WITHOUT_A_HOST_ROOM"] == []
    glz = {g["BLUE_ELEMENT_ID"]: g for g in reg("QORTUBA_GLAZED_ELEMENT_HOSTS")["ROWS"]}
    q = by_id()
    for qid in ("Q-05", "Q-06", "Q-08", "Q-09", "Q-10"):
        rooms = set(q[qid]["ROOMS"])
        for o in q[qid]["RESIDUAL_OPENINGS"]:
            oid = o.get("OPENING_ID")
            if oid in glz and glz[oid]["ROOM"]:
                # a room answers to its full label and to its canonical name
                aliases = rooms | {r.split(" /")[0] for r in rooms}
                assert glz[oid]["ROOM"].split(" /")[0] in aliases, (qid, oid)
    # a glazed element can only move a blockwork group whose thickness matches its host wall
    for qid, x in q.items():
        if not qid.startswith("Q-07-"):
            continue
        t = int(qid.split("-")[-1])
        for o in x["RESIDUAL_OPENINGS"]:
            if o.get("OPENING_ID") in glz:
                assert int(glz[o["OPENING_ID"]]["HOST_WALL_THICKNESS_MM"]) == t, (qid, o["OPENING_ID"])


def test_an_opening_outside_the_apartment_does_not_smear_across_apartment_rows():  # noqa: D401
    out = {o["OPENING_ID"] for o in reg("QORTUBA_OPENING_REGISTER")["ROWS"]
           if not o["ROOMS"] and o["HOST_WALL_ID"]}
    assert out, "the fixture needs at least one site on a wall touching no apartment room"
    for x in quants()["ROWS"]:
        assert not (out & {o.get("OPENING_ID") for o in x["RESIDUAL_OPENINGS"]}), x["QUANTITY_ID"]


# ------------------------------------------------------------------ §V waste stays out
def test_no_waste_and_no_procurement_quantity_exists_anywhere():
    for x in quants()["ROWS"]:
        assert x["WASTE_FACTOR"] is None and x["PROCUREMENT_QUANTITY"] is None, x["QUANTITY_ID"]
    assert quants()["WASTE_APPLIED_ANYWHERE"] is False and quants()["RATES_SUPPLIED"] == 0


# ------------------------------------------------------------------ §Q geometry is not re-run
def test_a_rule_change_recalculates_quantities_and_never_geometry():
    R = rules()
    for stage in ("wall detection", "room detection", "opening detection", "floor polygons"):
        assert stage in R["STAGES_NEVER_RERUN_BY_A_RULE_CHANGE"]
    dep = R["DEPENDENCY_GRAPH"]
    assert dep["QORTUBA_INTERNAL_PLASTER_HEIGHT"] == ["INTERNAL_PLASTER", "TILE_PREPARATION"]
    assert "BLOCKWORK_BY_THICKNESS" not in dep["QORTUBA_INTERNAL_PLASTER_HEIGHT"]
    assert all(not any(s in v for s in R["STAGES_NEVER_RERUN_BY_A_RULE_CHANGE"]) for v in dep.values())
    assert quants()["GEOMETRY_STAGES_RERUN"] == 0
    assert reg("FREEZE_URBAN_OWNER_RULES_V1")["QORTUBA_GEOMETRY_CHANGED"] == "NONE"


def test_the_geometry_that_went_in_is_the_geometry_that_came_out():
    sk = reg("QORTUBA_QS01_SKIRTING_TAKEOFF", QS)["ROWS"]
    w = reg("QORTUBA_QS01_BLOCK_WALL_TAKEOFF", QS)
    rooms = {x["ROOM_ID"]: x for x in reg("QORTUBA_ROOM_FINISH_AND_SKIRTING_RECALC")["ROWS"]}
    assert len(rooms) == len(sk)
    for x in sk:
        assert abs(rooms[x["ROOM_ID"]]["GROSS_WALL_LINE_LM"] - x["GROSS_WALL_LINE_PERIMETER_LM"]) < 1e-9
    walls = reg("QORTUBA_BLOCKWORK_RECALC")["ROWS"]
    assert len(walls) == w["COUNT"]
    for a, b in zip(walls, w["ROWS"]):
        assert a["WALL_ID"] == b["WALL_ID"] and a["LENGTH_M"] == b["LENGTH_M"]


def test_the_floors_still_add_up_after_the_paintry_moved():
    s = reg("QORTUBA_QS01_SUMMARY_TOTALS", QS)
    q = by_id()
    total = q["Q-11"]["MEASURED_NET_QUANTITY"] + q["Q-12"]["MEASURED_NET_QUANTITY"] + q["Q-13"]["MEASURED_NET_QUANTITY"]
    assert abs(total - s["C_TOTAL_INTERNAL_APARTMENT_FLOOR_AREA_M2"]["VALUE"]) < 1e-6


# ------------------------------------------------------------------ §R the questions
def test_an_answered_question_is_closed_and_not_asked_again():
    ql = reg("QORTUBA_QUESTION_LEDGER")
    old = {d["DECISION_ID"] for d in reg("QORTUBA_OWNER_DECISIONS_REQUIRED")["ROWS"]}
    closed = {c["DECISION_ID"] for c in ql["CLOSED"]}
    # every D-xx closure answers a question that was actually asked; later closures carry their own ids
    assert {c for c in closed if c.startswith("D-")} <= old
    assert len(closed) >= 16
    for c in ql["CLOSED"]:
        assert c["ANSWERED_BY"]
    open_q = " ".join(x["QUESTION"].lower() for x in ql["STILL_OPEN"])
    for word in ("plaster height", "blockwork height", "tiling height", "deduction column"):
        assert word not in open_q, word


def test_every_open_question_still_blocks_something_real():
    for x in reg("QORTUBA_QUESTION_LEDGER")["STILL_OPEN"]:
        assert x["BLOCKS"] and x["WHY"]
        assert x["KIND"] in ("PROJECT_INPUT", "PROJECT_RULE", "DRAWING_REQUIRED", "SPEC_REQUIRED")


# ------------------------------------------------------------------ the frozen inputs and the workbook
def test_the_frozen_inputs_still_verify():
    checked = OR.verify_inputs()
    assert len(checked) == 2 and all(c["VERIFIED"] for c in checked)
    fz = reg("FREEZE_URBAN_OWNER_RULES_V1")
    assert fz["HISTORICAL_BOQ_REINTERPRETED"] is False and fz["NEW_TAKEOFF_PHASES_CREATED"] == 0
    for name, sha in fz["CONTENTS"].items():
        assert OR._sha(OUT / f"{name}.json") == sha, name


def test_the_workbook_keeps_the_fifteen_house_columns_and_no_rate():
    from openpyxl import load_workbook
    wb = load_workbook(WB.FILE, read_only=True)
    ws = wb["جدول الكميات BOQ"]
    head = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert head[:15] == ["البند", "الوصف", "وحدة التسعير", "كمية القياس", "وحدة القياس", "قاعدة التحويل",
                         "الكمية النهائية", "% الهالك", "كمية الشراء", "سعر الوحدة", "الإجمالي", "مصدر القياس",
                         "قاعدة القياس التجاري", "حالة الاعتماد", "ملاحظات"]
    for row in ws.iter_rows(min_row=5):
        for c in (8, 9, 10, 11):
            assert row[c - 1].value is None, (row[0].value, c)


def test_the_audit_sheet_carries_the_rule_and_the_parameter_behind_each_figure():
    from openpyxl import load_workbook
    ws = load_workbook(WB.FILE, read_only=True)["تدقيق الحساب Audit"]
    head = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert "RULE_ID" in head and "PARAMETER_SOURCE" in head
    ri, pi = head.index("RULE_ID"), head.index("PARAMETER_SOURCE")
    seen = 0
    for row in ws.iter_rows(min_row=5):
        if row[0].value and str(row[0].value).startswith("Q-"):
            assert row[ri].value and row[pi].value, row[0].value
            seen += 1
    assert seen >= 20


def test_the_rooms_sheet_is_the_last_sheet_and_the_bill_is_trade_based():
    from openpyxl import load_workbook
    wb = load_workbook(WB.FILE, read_only=True)
    assert wb.sheetnames[-1] == "حساب الغرف Rooms"
    assert wb.sheetnames[1] == "جدول الكميات BOQ"
