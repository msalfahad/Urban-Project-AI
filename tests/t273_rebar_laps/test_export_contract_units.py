"""T-273 part (3): reinforcement in metric tons at the export boundary (tools/boq_export_v1.py --contract-units).

Hand calculation (synthetic villa, not the development villa):
  STR.REBAR kg: R1 MEASURED 1,234.5678 kg -> 1.234568 t; R2 MEASURED 0.4 kg -> 0.000400 t (visible at 6 decimals);
  R3 BOUNDED_EXPOSURE 100 .. 150 kg (HELD) -> 0.100000 .. 0.150000 t.
  Line total measured = (1,234.5678 + 0.4) kg / 1000 = 1.234968 t (from the unrounded kg sum, converted once).
  Concrete stays m3 (factor 1, unit unchanged).
"""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from engine.reporting_v3 import units as U
from engine.source import boq_line as BL
from tools import boq_export_v1 as X

PID, REV = "SYNTH-VILLA-G", "a" * 64
CVER, BSHA = "ARCH_TAKEOFF_V3.0.TEST", "b" * 64


def ln(code, eid, ev="MEASURED", q=None, unit="kg", **kw):
    base = dict(project_id=PID, revision_set_sha256=REV, item_code=code, unit=unit, evidence_state=ev,
                measured_net=q, element_id=eid, formula=kw.pop("formula", "n*L*d2/162"), rule_id="R-SYN",
                rule_version="1", floor=kw.pop("floor", "GB"), inputs=kw.pop("inputs", {"n": 3, "L": 1.0}))
    base.update(kw)
    return BL.make_line(**base)


def villa():
    return [ln("STR.REBAR", "R1", q=1234.5678), ln("STR.REBAR", "R2", q=0.4),
            ln("STR.REBAR", "R3", ev="BOUNDED_EXPOSURE", bounded_range={"low": 100.0, "high": 150.0, "basis": "lap"},
               rfi={"question": "lap length not printed"}),
            ln("STR.CONC.FOOTING", "F1", q=1.2, unit="m3", inputs={"L": 1.5, "B": 1.0, "H": 0.8})]


def run(tmp_path, lines, name="o", **kw):
    m = X.export(lines, tmp_path / name, candidate_version=CVER, boq_lines_sha256=BSHA, **kw)
    return m, tmp_path / name


def by(rows, eid):
    return [r for r in rows if r["element_id"] == eid][0]


def test_default_export_is_byte_identical_and_kg():
    import tempfile
    from pathlib import Path
    with tempfile.TemporaryDirectory() as d:
        m1, p1 = run(Path(d), villa(), "a")
        m2, p2 = run(Path(d), villa(), "b", contract_units=False)
        assert (p1 / "BOQ.csv").read_bytes() == (p2 / "BOQ.csv").read_bytes()
        assert (p1 / "BOQ.xlsx").read_bytes() == (p2 / "BOQ.xlsx").read_bytes()
        assert m1["schema"] == X.MANIFEST_SCHEMA and "contract_units" not in m1
        assert not (p1 / "BOQ_CONTRACT_UNITS.csv").exists()
        assert by(X.read_csv(p1 / "BOQ.csv"), "R1")["unit"] == "kg"


def test_contract_units_convert_once_with_six_decimals(tmp_path):
    m, p = run(tmp_path, villa(), contract_units=True)
    assert m["schema"] == X.MANIFEST_SCHEMA_CONTRACT and m["contract_units"]["schema"] == X.CONTRACT_SCHEMA
    rows = X.read_csv(p / "BOQ_CONTRACT_UNITS.csv")
    r1, r2, r3, f1 = by(rows, "R1"), by(rows, "R2"), by(rows, "R3"), by(rows, "F1")
    assert (r1["unit"], r1["engine_unit"], r1["net"]) == ("t", "kg", "1.234568")
    assert r2["qty_measured_unverified"] == "0.000400"                       # a small quantity stays visible
    assert (r3["quantity_state"], r3["held_low"], r3["held_high"], r3["net"]) == ("HELD", "0.100000", "0.150000", "")
    assert (f1["unit"], f1["engine_unit"], f1["net"], f1["conversion"]) == ("m3", "m3", "1.200000", "none")
    # the canonical rows keep the engine unit (the unit every verdict binds to)
    assert by(X.read_csv(p / "BOQ.csv"), "R1")["net"] == "1234.567800"
    # the factor agrees with the reporting policy (one conversion, not two)
    assert X.CONTRACT_UNIT["kg"][0] == U.UNIT_POLICY["kg"][0] and 1 / X.CONTRACT_UNIT["kg"][1] == U.UNIT_POLICY["kg"][1]


def test_totals_from_unrounded_kg_not_from_rounded_t_rows(tmp_path):
    lines = [ln("STR.REBAR", f"R{i}", q=0.4004) for i in range(3)]
    m, p = run(tmp_path, lines, contract_units=True)
    rows = X.read_csv(p / "BOQ_CONTRACT_UNITS.csv")
    assert [r["net"] for r in rows] == ["0.000400"] * 3                 # rounded line sum would be 0.001200
    tot = {t["item_code"]: t for t in X.read_csv(p / "BOQ_CONTRACT_TOTALS.csv")}["STR.REBAR"]
    assert tot["measured_unverified_sum"] == "0.001201" and tot["measured_unverified_engine_sum"] == "1.201200"


def test_line_total_in_t(tmp_path):
    m, p = run(tmp_path, villa(), contract_units=True)
    tot = {t["item_code"]: t for t in X.read_csv(p / "BOQ_CONTRACT_TOTALS.csv")}
    assert tot["STR.REBAR"]["unit"] == "t" and tot["STR.REBAR"]["measured_unverified_sum"] == "1.234968"
    assert tot["STR.REBAR"]["held_low"] == "0.100000" and tot["STR.REBAR"]["rows_without_quantity"] == "1"
    assert tot["STR.CONC.FOOTING"]["unit"] == "m3" and tot["STR.CONC.FOOTING"]["measured_unverified_sum"] == "1.200000"


def test_kg_and_t_rows_in_one_line_refused(tmp_path):
    lines = villa() + [ln("STR.REBAR", "R9", q=0.5, unit="t")]
    with pytest.raises(X.ExportError, match="mixed units"):
        run(tmp_path, lines, contract_units=True)


def test_xlsx_carries_both_sheets_and_the_manifest_hashes_the_new_files(tmp_path):
    m, p = run(tmp_path, villa(), contract_units=True)
    wb = load_workbook(p / "BOQ.xlsx")
    assert wb.sheetnames == ["BOQ", "Trade summary", "BOQ (contract units)", "Contract totals", "Manifest"]
    ws = wb["BOQ (contract units)"]
    hdr = [c.value for c in ws[1]]
    vals = {r[hdr.index("element_id")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert vals["R2"][hdr.index("net")] == 0.0004 and vals["R2"][hdr.index("unit")] == "t"
    assert vals["R3"][hdr.index("net")] is None                              # held: an empty cell, never 0
    man = json.loads((p / "EXPORT_MANIFEST.json").read_text(encoding="utf-8"))
    assert set(man["outputs"]) == {"BOQ.csv", "BOQ.xlsx", "BOQ_TRADE_SUMMARY.csv", "BOQ_CONTRACT_UNITS.csv",
                                   "BOQ_CONTRACT_TOTALS.csv"}
    assert man["outputs"]["BOQ_CONTRACT_UNITS.csv"] == X.sha256_bytes((p / "BOQ_CONTRACT_UNITS.csv").read_bytes())


def test_a_v2_verdict_survives_the_unit_change(tmp_path):
    lines = villa()
    ls = [x for x in lines if x["location"]["element_id"] in ("R1", "R2", "R3")]
    rec = {"schema": X.V2_VERDICT_SCHEMA, "verdict_id": "V-REBAR-R1", "project_id": PID, "revision_set_sha256": REV,
           "candidate_version": CVER, "boq_lines_sha256": BSHA, "line_ids": [ls[0]["line_id"]],
           "element_scope": {"ids": ["R1"], "count": 1}, "unit": "kg", "engine_qty": 1234.5678,
           "checker_qty": 1236.0, "evidence": [{"sheet": "VG-S-03", "printed": "3 T16"}],
           "checker": {"run_token": "tok", "routine": "Verifier", "verdict_commit": "c0ffee"}, "scope": "COMPONENT",
           "time": "2026-10-24T17:00Z"}
    m0, p0 = run(tmp_path, lines, "kg", v2_records=[rec])
    m1, p1 = run(tmp_path, lines, "t", v2_records=[rec], contract_units=True)
    assert m0["v2_verdicts"] == m1["v2_verdicts"] and "COUNTED" in m1["v2_verdicts"][0]["status"]
    assert (p0 / "BOQ.csv").read_bytes() == (p1 / "BOQ.csv").read_bytes()     # binding rows unchanged
    r = by(X.read_csv(p1 / "BOQ_CONTRACT_UNITS.csv"), "R1")
    assert (r["quantity_state"], r["qty_verified_component"], r["checker_qty"], r["verdict_id"]) == \
        ("VERIFIED COMPONENT", "1.234568", "1.236000", "V-REBAR-R1")


def test_cli_flag(tmp_path):
    p = tmp_path / "lines.json"
    p.write_text(json.dumps(villa()), encoding="utf-8")
    assert X.main(["--lines", str(p), "--candidate-version", CVER, "--out", str(tmp_path / "o"),
                   "--code-commit", "abc123", "--contract-units"]) == 0
    assert (tmp_path / "o" / "BOQ_CONTRACT_UNITS.csv").exists()
