"""Builds a synthetic 'Alsenan Chalet'-like workbook with known defects.

A stand-in until the real Alsenan Chalet Excel is provided. It embeds one
instance of every defect class the auditor must catch, in a realistic
multi-sheet, Arabic+English layout, so the end-to-end loader + rules path is
proven. When the real file arrives, the same `audit_workbook()` runs on it
unchanged.

Note the distinction the auditor draws:
  - A **money** subtotal (SUM of the Amount column) across mixed-unit lines is
    NORMAL and must NOT be flagged.
  - A **quantity** subtotal (SUM of the Quantity column) across mixed units is
    the 295.44 defect and MUST be flagged.
"""

from __future__ import annotations

from openpyxl import Workbook


def build_alsenan_like(path: str) -> str:
    wb = Workbook()

    # ── Sheet 1: Takeoff (the 295.44 quantity-total defect) ──────────────
    ws = wb.active
    ws.title = "Takeoff"
    ws.append(["Alsenan Chalet — Takeoff", None, None, None, None])
    ws.append(["Description", "Unit", "Quantity", "Rate", "Amount"])    # row 2 header
    ws.append(["Wall plaster area", "m2", 180.0, None, None])           # 3
    ws.append(["Skirting length", "m", 100.0, None, None])              # 4
    ws.append(["Courtyard", "m2", 15.44, None, None])                   # 5
    # WRONG: a quantity SUM across m2 + m → 295.44 (mixed units).
    ws.append(["Combined quantity", None, "=SUM(C3:C5)", 15, None])     # 6: R05

    # ── Sheet 2: Structure (concrete + steel) ────────────────────────────
    ws1 = wb.create_sheet("Structure")
    ws1.append(["Alsenan Chalet — Structure", None, None, None, None])
    ws1.append(["Description", "Unit", "Quantity", "Rate", "Amount"])   # row 2 header
    ws1.append(["Reinforced concrete raft foundation C15", "m3", 120, 45, 5400])  # 3: R08
    ws1.append(["Reinforcement steel", "kg", 600, 0.3, 180])           # 4: R09 (5 kg/m3)
    ws1.append(["Blinding concrete", "m3", 30, 25, 750])               # 5
    # A legitimate MONEY subtotal across mixed-unit lines — must NOT flag R05.
    ws1.append(["Total Structure (KWD)", None, None, None, "=SUM(E3:E5)"])  # 6

    # ── Sheet 3: Finishes (blank price, formula error, #REF!, skipped sum) ─
    ws2 = wb.create_sheet("Finishes")
    ws2.append(["Alsenan Chalet — Finishes", None, None, None, None])
    ws2.append(["Description", "Unit", "Quantity", "Rate", "Amount"])   # row 2 header
    ws2.append(["Porcelain floor tiles", "m2", 200, 5.0, 1000])        # 3: material, no waste → R10
    ws2.append(["Marble skirting", "m", 100, 3.0, 300])                # 4
    ws2.append(["سيجما (sigma)", "m2", 1500, 0, 0])                     # 5: R02 blank price
    ws2.append(["Gypsum board", "m2", 120, 2.5, 500])                  # 6: R03 (should be 300)
    ws2.append(["Broken reference item", "m2", 50, 4, "#REF!"])        # 7: R01
    # Money subtotal that SKIPS the priced rows 5,6,7 → R07.
    ws2.append(["Subtotal finishes (KWD)", None, None, None, "=SUM(E3:E4)"])  # 8

    wb.save(path)
    return path
