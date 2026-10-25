"""T-298 part (1): the contractual export is the command line's DEFAULT; reinforcement in metric tons (owner 2026-10-24
14:23 unit, 2026-10-25 11:51 (0) 'correct tonne reporting').

Hand calculation (synthetic villa VG-298, not the development villa):
  STR.REBAR kg   R1 MEASURED    2,468.1357 kg -> 2.4681357 t -> 2.468136 t (6 decimals)
                 R2 PROVISIONAL   531.8643 kg -> 0.5318643 t -> 0.531864 t
                 R3 BOUNDED_EXPOSURE 200 .. 260 kg (HELD, blank) -> 0.200000 .. 0.260000 t
                 line measured total (2,468.1357 + 531.8643) kg = 3,000.0000 kg -> 3.000000 t (one conversion of
                 the unrounded kg sum; the rounded t rows also give 3.000000 here, the next test shows when they do not)
  STR.CONC.FOOTING m3  F1 1.80 x 1.50 x 0.60 = 1.620 m3 (factor none, unit unchanged)
"""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from engine.source import boq_line as BL
from tools import boq_export_v1 as X

PID, REV = "SYNTH-VILLA-VG298", "c" * 64
CVER = "ARCH_TAKEOFF_V3.0.TEST"


def ln(code, eid, ev="MEASURED", q=None, unit="kg", **kw):
    base = dict(project_id=PID, revision_set_sha256=REV, item_code=code, unit=unit, evidence_state=ev,
                measured_net=q, element_id=eid, formula=kw.pop("formula", "n*L*d2/162"), rule_id="R-SYN",
                rule_version="1", floor=kw.pop("floor", "GF"), inputs=kw.pop("inputs", {"n": 4, "L": 2.0}))
    base.update(kw)
    return BL.make_line(**base)


def villa():
    return [ln("STR.REBAR", "R1", q=2468.1357), ln("STR.REBAR", "R2", ev="PROVISIONAL", q=531.8643),
            ln("STR.REBAR", "R3", ev="BOUNDED_EXPOSURE", bounded_range={"low": 200.0, "high": 260.0, "basis": "lap"},
               rfi={"question": "lap length not printed"}),
            ln("STR.CONC.FOOTING", "F1", q=1.62, unit="m3", inputs={"L": 1.8, "B": 1.5, "H": 0.6},
               formula="L*B*H")]


def cli(tmp_path, lines, name, *extra):
    p = tmp_path / f"{name}.json"
    p.write_text(BL.dumps(lines), encoding="utf-8")
    assert X.main(["--lines", str(p), "--candidate-version", CVER, "--out", str(tmp_path / name),
                   "--code-commit", "abc1234", *extra]) == 0
    return tmp_path / name, json.loads((tmp_path / name / "EXPORT_MANIFEST.json").read_text(encoding="utf-8"))


def by(rows, eid):
    return [r for r in rows if r["element_id"] == eid][0]


def test_default_command_line_writes_tonnes_with_kg_audit_columns(tmp_path):
    d, m = cli(tmp_path, villa(), "c")
    assert m["schema"] == X.MANIFEST_SCHEMA_CONTRACTUAL == "URBAN_BOQ_EXPORT_MANIFEST_V3"
    assert m["export_schema"] == X.CONTRACTUAL_SCHEMA == "URBAN_BOQ_EXPORT_V2"
    rows = X.read_csv(d / "BOQ.csv")
    r1, r2, r3, f1 = by(rows, "R1"), by(rows, "R2"), by(rows, "R3"), by(rows, "F1")
    assert (r1["unit"], r1["engine_unit"], r1["net"], r1["engine_net"]) == ("t", "kg", "2.468136", "2468.135700")
    assert r1["conversion"] == "kg / 1000 (once, at export)"
    assert (r2["qty_measured_unverified"], r2["engine_qty_measured_unverified"]) == ("0.531864", "531.864300")
    assert (r3["quantity_state"], r3["net"], r3["engine_net"], r3["held_low"], r3["held_high"],
            r3["engine_held_low"]) == ("HELD", "", "", "0.200000", "0.260000", "200.000000")   # HELD blank, never 0
    assert (f1["unit"], f1["engine_unit"], f1["net"], f1["engine_net"], f1["conversion"]) == \
        ("m3", "m3", "1.620000", "1.620000", "none")
    assert m["contractual_units"]["rows_converted"] == 3


def test_engine_unit_files_are_the_v1_bytes_and_the_opt_out_restores_them(tmp_path):
    d, m = cli(tmp_path, villa(), "c")
    e, me = cli(tmp_path, villa(), "e", "--engine-units")
    v1 = X.export(villa(), tmp_path / "lib", candidate_version=CVER, boq_lines_sha256=m["identity"]["boq_lines_sha256"],
                  code_commit="abc1234")
    assert me["schema"] == X.MANIFEST_SCHEMA and v1["schema"] == X.MANIFEST_SCHEMA and "contractual_units" not in me
    assert (e / "BOQ.csv").read_bytes() == (tmp_path / "lib" / "BOQ.csv").read_bytes()
    assert (d / "BOQ_ENGINE_UNITS.csv").read_bytes() == (e / "BOQ.csv").read_bytes()
    assert (d / "BOQ_TRADE_SUMMARY_ENGINE_UNITS.csv").read_bytes() == (e / "BOQ_TRADE_SUMMARY.csv").read_bytes()
    assert by(X.read_csv(e / "BOQ.csv"), "R1")["unit"] == "kg"
    # the identity of the rows verdicts bind to does not move with the unit
    assert m["semantic_sha256"] == me["semantic_sha256"] and m["identity"] == me["identity"]
    assert m["row_counts"] == me["row_counts"]


def test_every_non_reinforcement_cell_is_unchanged(tmp_path):
    d, _m = cli(tmp_path, villa(), "c")
    eng = {r["element_id"]: r for r in X.read_csv(d / "BOQ_ENGINE_UNITS.csv")}
    for r in X.read_csv(d / "BOQ.csv"):
        e = eng[r["element_id"]]
        moved = {c for c in X.COLUMNS if r[c] != e[c]}
        if e["unit"] == "kg":
            assert moved <= {"unit"} | set(X.CONTRACT_QTY_COLUMNS)
            assert all(r[f"engine_{c}"] == e[c] for c in X.CONTRACT_QTY_COLUMNS)
        else:
            assert moved == set()


def test_trade_summary_in_tonnes_from_the_unrounded_kg_sums(tmp_path):
    d, _m = cli(tmp_path, villa(), "c")
    s = {r["item_code"]: r for r in X.read_csv(d / "BOQ_TRADE_SUMMARY.csv")}
    rb = s["STR.REBAR"]
    assert (rb["unit"], rb["engine_unit"], rb["measured_unverified_sum"], rb["held_low"], rb["held_high"]) == \
        ("t", "kg", "3.000000", "0.200000", "0.260000")
    assert rb["held_components"] == "R3 [0.200000 .. 0.260000]"
    assert (s["STR.CONC.FOOTING"]["unit"], s["STR.CONC.FOOTING"]["measured_unverified_sum"]) == ("m3", "1.620000")
    lines = [ln("STR.REBAR", f"Q{i}", q=0.4004) for i in range(3)]          # 3 x 0.4004 kg = 1.2012 kg
    d2, _m2 = cli(tmp_path, lines, "r")
    assert [r["net"] for r in X.read_csv(d2 / "BOQ.csv")] == ["0.000400"] * 3   # rounded rows would sum to 0.001200
    assert X.read_csv(d2 / "BOQ_TRADE_SUMMARY.csv")[0]["measured_unverified_sum"] == "0.001201"
    tot = X.read_csv(d2 / "BOQ_CONTRACT_TOTALS.csv")[0]
    assert (tot["measured_unverified_sum"], tot["measured_unverified_engine_sum"]) == ("0.001201", "1.201200")


def test_a_v2_verdict_binds_identically_in_both_units(tmp_path):
    lines = villa()
    r1 = [x for x in lines if x["location"]["element_id"] == "R1"][0]
    base = tmp_path / "base"
    _e, me = cli(tmp_path, lines, "e0", "--engine-units")
    rec = {"schema": X.V2_VERDICT_SCHEMA, "verdict_id": "V-R1", "project_id": PID, "revision_set_sha256": REV,
           "candidate_version": CVER, "boq_lines_sha256": me["identity"]["boq_lines_sha256"],
           "line_ids": [r1["line_id"]], "element_scope": {"ids": ["R1"], "count": 1}, "unit": "kg",
           "engine_qty": 2468.1357, "checker_qty": 2470.0, "evidence": [{"sheet": "VG-S-04", "printed": "4 T16"}],
           "checker": {"run_token": "tok", "routine": "Verifier", "verdict_commit": "c0ffee1"}, "scope": "COMPONENT",
           "time": "2026-10-25T13:00Z"}
    base.mkdir()
    vf = base / "V2.json"
    vf.write_text(json.dumps({"schema": X.V2_SET_SCHEMA, "verdicts": [rec]}), encoding="utf-8")
    e, me = cli(tmp_path, lines, "e", "--engine-units", "--verdicts-v2", str(vf))
    d, m = cli(tmp_path, lines, "c", "--verdicts-v2", str(vf))
    assert m["v2_verdicts"] == me["v2_verdicts"] and m["v2_verdicts"][0]["status"] == "COUNTED (COMPONENT)"
    r = by(X.read_csv(d / "BOQ.csv"), "R1")
    assert (r["quantity_state"], r["qty_verified_component"], r["engine_qty_verified_component"], r["checker_qty"]) == \
        ("VERIFIED COMPONENT", "2.468136", "2468.135700", "2.470000")


def test_xlsx_sheets_and_manifest_hashes(tmp_path):
    d, m = cli(tmp_path, villa(), "c")
    wb = load_workbook(d / "BOQ.xlsx")
    assert wb.sheetnames == ["BOQ", "Trade summary", "Contract totals", "BOQ (engine units)",
                             "Trade summary (engine units)", "Manifest"]
    ws = wb["BOQ"]
    hdr = [c.value for c in ws[1]]
    vals = {row[hdr.index("element_id")]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert vals["R1"][hdr.index("net")] == 2.468136 and vals["R1"][hdr.index("unit")] == "t"
    assert vals["R1"][hdr.index("engine_net")] == 2468.1357
    assert vals["R3"][hdr.index("net")] is None                              # an empty cell, never 0
    assert set(m["outputs"]) == {"BOQ.csv", "BOQ.xlsx", "BOQ_TRADE_SUMMARY.csv", "BOQ_ENGINE_UNITS.csv",
                                 "BOQ_TRADE_SUMMARY_ENGINE_UNITS.csv", "BOQ_CONTRACT_TOTALS.csv"}
    for f, h in m["outputs"].items():
        assert h == X.sha256_bytes((d / f).read_bytes())


# ---- adverse --------------------------------------------------------------------------------------------------------

def test_double_conversion_refused(tmp_path):
    res = X.build(villa(), candidate_version=CVER, boq_lines_sha256="b" * 64)
    krow, _t = X.contractual_rows(res["rows"])
    with pytest.raises(X.ExportError, match="double conversion"):
        X.contract_view(krow)


def test_kg_relabelled_as_t_refused(tmp_path):
    lines = [ln("STR.REBAR", "R1", q=2.468, unit="t"), ln("STR.CONC.FOOTING", "F1", q=1.62, unit="m3")]
    p = tmp_path / "L.json"
    p.write_text(BL.dumps(lines), encoding="utf-8")
    with pytest.raises(X.ExportError, match="pre-converted or relabelled"):
        X.main(["--lines", str(p), "--candidate-version", CVER, "--out", str(tmp_path / "o"), "--code-commit", "a1"])
    assert not (tmp_path / "o" / "BOQ.csv").exists()                         # refused: nothing written


def test_mixed_kg_and_t_line_refused(tmp_path):
    p = tmp_path / "L.json"
    p.write_text(BL.dumps(villa() + [ln("STR.REBAR", "R9", q=0.5, unit="t")]), encoding="utf-8")
    with pytest.raises(X.ExportError, match="mixed units"):
        X.main(["--lines", str(p), "--candidate-version", CVER, "--out", str(tmp_path / "o"), "--code-commit", "a1"])


def test_two_contract_views_in_one_export_refused(tmp_path):
    with pytest.raises(X.ExportError, match="already is the contract-unit view"):
        X.export(villa(), tmp_path / "o", candidate_version=CVER, boq_lines_sha256="b" * 64, contract_units=True,
                 contractual=True)
