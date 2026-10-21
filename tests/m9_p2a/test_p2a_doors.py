"""Phase 2A P2-03 / P2-04 + the owner door rules of 2026-10-13 (ARCH_TAKEOFF_V2.1): door element conservation,
states with typed evidence, three separate widths, height / area only from a source, the PVC / aluminium / main
entrance / lift classes, provisional doors kept out of released totals. Synthetic only (round 2 AT-Q2-01 .. 06)."""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
from engine.source import boq_line as BL                       # noqa: E402
from engine.source import door_model as DM                     # noqa: E402

import p2a_helpers as H                                         # noqa: E402

FACTS_H = {"schema": "URBAN_ARCH_FACTS_V1",
           "opening_heights": {"DOOR": {"value_m": 2.1, "authority": "synthetic door schedule rev A"}}}


def _facts(**heights):
    from engine.source import arch_measure as AM
    return {"schema": AM.FACTS_SCHEMA, "opening_heights": heights}


def _lines(model):
    return DM.boq_lines({"project_id": "P", "revision_set_sha256": "0" * 64, "documents": []}, model)


def _widths(model, cls):
    return sorted(round(el["structural_opening_width_m"], 6) for el in model["elements"]
                  if el["boq_class"] == cls and el["released"])


@pytest.fixture(scope="module")
def base():
    return H.door_run(H.door_fixture())


def test_every_door_gets_its_class(base):
    _S, _M, model = base
    assert _widths(model, DM.PVC) == H.DOOR_TRUTH["PVC"]
    assert _widths(model, DM.ALU) == H.DOOR_TRUTH["ALU"]
    assert _widths(model, DM.MAIN) == H.DOOR_TRUTH["MAIN"]
    lift = [el for el in model["elements"] if el["boq_class"] == DM.LIFT]
    assert [round(el["structural_opening_width_m"], 6) for el in lift] == H.DOOR_TRUTH["LIFT"]
    assert sum(1 for el in model["elements"] if el["state"] == DM.UNRESOLVED) == H.DOOR_TRUTH["UNRESOLVED"]
    assert model["main_entrance"]["state"] == "IDENTIFIED"


def test_swing_conservation(base):
    """AT-Q2-01: drawn swings = the leaves of CONFIRMED + PROVISIONAL + UNRESOLVED elements; none in two elements."""
    _S, _M, model = base
    (c,) = model["conservation"]
    assert c["ok"] and c["swings"] == H.DOOR_TRUTH["swings"] and sum(c["leaves"].values()) == c["swings"]
    assert Counter(o for el in model["elements"] for o in el["swings"]).most_common(1)[0][1] == 1


def test_a_double_leaf_door_is_one_door_assembly(base):
    """AT-Q2-02 (one insert, two arcs): one element, two leaves, one BOQ line."""
    _S, _M, model = base
    dd = [el for el in model["elements"] if el["leaf_count"] == 2]
    assert len(dd) == 1 and abs(dd[0]["structural_opening_width_m"] - H.DOOR_TRUTH["double_leaf_width"]) < 1e-9
    assert len(dd[0]["nominal_leaf_width_m"]) == 2
    assert sum(1 for ln in _lines(model) if ln["physical_element"] == dd[0]["element_id"]) == 1


def test_the_main_entrance_is_counted_once_and_never_as_pvc_or_aluminium(base):
    _S, _M, model = base
    lines = _lines(model)
    main = [ln for ln in lines if ln["item_code"] == DM.MAIN]
    assert len(main) == 1
    eid = main[0]["physical_element"]
    assert not [ln for ln in lines if ln["physical_element"] == eid and ln["item_code"] != DM.MAIN]
    assert main[0]["calculation"]["inputs"]["material"] == "NOT_ESTABLISHED"


def test_two_entrance_candidates_are_never_resolved_by_assumption():
    words = H.DOOR_WORDS
    ops = H.DOOR_OPENINGS + [("DOOR", "x", 0, 200, 1200, 2100, None)]            # a second hall door to the outside
    sw = H.DOOR_SWINGS + (((200, 1200), 900, 90, 180),)
    _S, _M, model = H.door_run(H.door_fixture(words=words, openings=ops, swings=sw))
    assert model["main_entrance"]["state"] == "REVIEW"
    assert any(f["code"] == "MAIN_ENTRANCE_AMBIGUOUS" for f in model["findings"])
    lines = _lines(model)
    assert not [ln for ln in lines if ln["item_code"] == DM.MAIN]
    cands = {el["element_id"] for el in model["elements"] if el["boq_class"] == "MAIN_ENTRANCE_CANDIDATE"}
    assert len(cands) == 2 and not [ln for ln in lines if ln["physical_element"] in cands]
    assert _widths(model, DM.ALU) == [0.9]                       # the kitchen door is not a candidate


def test_no_entrance_evidence_gives_a_review_and_no_assumed_main_door():
    words = tuple(("HALL", p) if w == "ENTRANCE HALL" else (w, p) for w, p in H.DOOR_WORDS)
    _S, _M, model = H.door_run(H.door_fixture(words=words))
    assert model["main_entrance"]["state"] == "REVIEW"
    assert any(f["code"] == "MAIN_ENTRANCE_UNRESOLVED" for f in model["findings"])
    assert not [ln for ln in _lines(model) if ln["item_code"] == DM.MAIN]
    assert _widths(model, DM.ALU) == [0.9]


def test_external_by_adjacency_not_by_swing_direction():
    """the kitchen door hung to open outwards is still an external (aluminium) door."""
    sw = tuple(((12000, 6000), 900, 0, 90) if s[0] == (11800, 6000) else s for s in H.DOOR_SWINGS)
    _S, _M, model = H.door_run(H.door_fixture(swings=sw))
    assert _widths(model, DM.ALU) == [0.9]


def test_missing_height_keeps_count_and_width_and_blocks_height_and_area(base):
    _S, _M, model = base
    for el in model["elements"]:
        if el["released"]:
            assert el["height"]["state"] == "BLOCKED" and el["height_m"] is None and el["area_m2"] is None
            assert "MISSING SOURCE" in el["height"]["source"]
    for ln in _lines(model):
        assert ln["measured_net"] == 1 and ln["calculation"]["inputs"]["width_m"] > 0
        assert ln["calculation"]["inputs"]["height_m"] is None and ln["calculation"]["inputs"]["area_m2"] is None
    assert sum(1 for f in model["findings"] if f["code"] == "DOOR_HEIGHT_MISSING") == \
        sum(1 for el in model["elements"] if el["released"])


def test_a_sourced_height_gives_the_area():
    _S, _M, model = H.door_run(H.door_fixture(), facts=_facts(DOOR={"value_m": 2.1, "authority": "schedule rev A"}))
    for el in model["elements"]:
        if el["released"]:
            assert el["height"]["state"] == "SOURCE" and el["area_m2"] == round(el["structural_opening_width_m"] * 2.1, 6)
    assert model["summary"]["released"][DM.PVC]["area_m2_confirmed"] == round(2.1 * sum(H.DOOR_TRUTH["PVC"]), 6)


def test_an_authorised_fallback_is_assumed_and_kept_out_of_confirmed_areas():
    _S, _M, model = H.door_run(H.door_fixture(), facts=_facts(DOOR_ASSUMED={"value_m": 2.1,
                                                                             "authority": "owner fallback, test"}))
    rel = [el for el in model["elements"] if el["released"]]
    assert rel and all(el["height"]["state"] == "ASSUMED" and el["area_state"] == "ASSUMED" for el in rel)
    assert all(v["area_m2_confirmed"] is None for v in model["summary"]["released"].values() if v["count"])


def test_lift_doors_are_outside_the_door_boq_and_the_floor_geometry_is_unchanged(base):
    S, _M, model = base
    lift = {el["element_id"] for el in model["elements"] if el["boq_class"] == DM.LIFT}
    assert lift and not [ln for ln in _lines(model) if ln["physical_element"] in lift]
    assert all(el["section"] == "LIFT_EXCLUDED" and not el["released"] for el in model["elements"]
               if el["element_id"] in lift)
    S2, _M2, _m2 = H.door_run(H.door_fixture(), method="V2")
    assert S["plate"] == S2["plate"]                             # the shaft keeps its gross floor geometry


def test_a_lift_door_is_never_measured_even_with_a_door_height_source():
    _S, _M, model = H.door_run(H.door_fixture(), facts=_facts(DOOR={"value_m": 2.1, "authority": "schedule rev A"}))
    (lift,) = [el for el in model["elements"] if el["boq_class"] == DM.LIFT]
    assert lift["height_m"] is None and lift["area_m2"] is None and lift["height"]["state"] == "NOT_APPLICABLE"
    assert not lift["released"] and lift["section"] == "LIFT_EXCLUDED"


def test_a_leaf_closed_door_is_provisional_with_the_closure_width_and_out_of_released_totals():
    """AT-Q2-04 / AT-Q2-06: width = the wall-to-wall closure, range [radius, closure]; never released."""
    from engine.source import arch_measure as AM
    from engine.source import arch_rooms as AR
    from tools import arch_synth as AS
    # a door hung beside a 60 mm wall nib (the reviewed nib case): the nib is below the 100 mm face piece, so the
    # jamb is not proven, no band opening forms, the leaf closes the gap wall to wall -> PROVISIONAL (no nib rule)
    rooms = {"LIVING": [(200, 200), (4200, 200), (4200, 3000), (200, 3000)],
             "BATH": [(200, 3200), (1600, 3200), (1600, 4400), (200, 4400)]}
    ents = AS._s4_ents(rooms, [(0, 0), (4400, 0), (4400, 4600), (0, 4600)], [("DOOR", "y", 3000, 3200, 260, 1060, None)],
                       texts=(("LIVING", (2000, 1500)), ("BATH", (900, 3800))), swings=(((260, 3000), 800, 270, 360),))
    S = AR.spaces(ents, H.DOOR_ROLES, 1e-3, region_id="R", tol="UNIT_RELATIVE", geom=H.RG.Geometry(1e-3))
    M = AM.measure(S, floor="GF", project_id="P", revision_set_sha256="0" * 64, documents=[])
    fm = DM.floor_model(S, M["rooms"], ents, 1e-3, "GF", wall_layers=["A-WALL"])
    model = DM.building_model([("GF", fm)], {"GF": ents["texts"]}, {"GF": {}}, 1e-3, entrance_floor="GF")
    (el,) = [e for e in model["elements"] if e["state"] == DM.PROVISIONAL]
    assert abs(el["structural_opening_width_m"] - 0.8) < 1e-6 and el["structural_opening_source"].startswith("CLOSURE")
    lo, hi = el["width_range_m"]
    assert lo == pytest.approx(0.8) and hi == pytest.approx(el["structural_opening_width_m"])
    assert not el["released"] and el["section"] == "PROVISIONAL_NOT_RELEASED" and not _lines(model)
    assert model["conservation_ok"]


def _hand_S(hinges, radii):
    doors = {f"I{k}|A{k}": {"hinge": h, "radius": r, "ends": [(h[0] + r, h[1]), (h[0], h[1] + r)],
                            "swing_part": f"A{k}", "layer": "A-DOOR"} for k, (h, r) in enumerate(zip(hinges, radii))}
    st = {k: {"state": "CLOSED", "width": radii[0] + radii[1], "closure_a": f"CLOSURE|{k}|A"} for k in doors}
    return {"doors": doors, "door_status": st, "openings": [], "faces": [], "repairs": []}


def test_an_exploded_double_leaf_pair_is_one_element_and_a_non_pair_is_two_with_a_review():
    """AT-Q2-02: hinges at both ends of one closure, r1 + r2 = span within 2 %: one element; otherwise two + REVIEW."""
    m = DM.elements(_hand_S([(0.0, 0.0), (1600.0, 0.0)], [800.0, 800.0]), [], {"segments": []}, 1e-3, "GF")
    assert [el["leaf_count"] for el in m["elements"]] == [2] and m["conservation"]["ok"]
    m = DM.elements(_hand_S([(0.0, 0.0), (1600.0, 0.0)], [700.0, 800.0]), [], {"segments": []}, 1e-3, "GF")
    assert sorted(el["leaf_count"] for el in m["elements"]) == [1, 1] and m["conservation"]["ok"]
    assert any(f["code"] == "DOUBLE_LEAF_UNCERTAIN" for f in m["findings"])


def test_no_silent_promotion(base):
    """AT-Q2-03: every CONFIRMED element carries confirming evidence; a swing never confirms; a DIMENSION does, and
    only when its promotion policy is switched on (off in the V2.1 run: owner decision pending)."""
    _S, _M, model = base
    for el in model["elements"]:
        if el["state"] == DM.CONFIRMED:
            assert DM.confirming_evidence(el)
    hand = _hand_S([(0.0, 0.0), (5000.0, 0.0)], [800.0, 800.0])
    m = DM.elements(hand, [], {"segments": []}, 1e-3, "GF")
    el = m["elements"][0]
    assert el["state"] == DM.PROVISIONAL
    with pytest.raises(ValueError):
        DM._transition(el, DM.CONFIRMED, next(e["id"] for e in el["evidence"] if e["kind"] == "SWING"))
    target = DM.stable_id("GF", (0.0, 0.0), 1600.0, 1e-3)
    dims = [{"element_id": target, "agrees": True, "value_mm": 1600.0}]
    off = DM.elements(_hand_S([(0.0, 0.0), (5000.0, 0.0)], [800.0, 800.0]), [], {"segments": []}, 1e-3, "GF",
                      dimensions=dims)
    assert all(e["state"] == DM.PROVISIONAL for e in off["elements"])
    assert any(ev["kind"] == "DIMENSION" for e in off["elements"] for ev in e["evidence"])
    on = DM.elements(_hand_S([(0.0, 0.0), (5000.0, 0.0)], [800.0, 800.0]), [], {"segments": []}, 1e-3, "GF",
                     dimensions=dims, promote_by_dimension=True)
    (p,) = [e for e in on["elements"] if e["element_id"] == target]
    assert p["state"] == DM.CONFIRMED and p["history"][-1] == {"from": DM.PROVISIONAL, "to": DM.CONFIRMED,
                                                               "evidence": f"EV-{target}-DIM"}


def test_the_three_widths_have_their_own_sources(base):
    """owner requirement 2: the structural opening, the frame and the leaf are separate fields; no BOQ width is a
    swing radius."""
    _S, _M, model = base
    for el in model["elements"]:
        if el["released"]:
            assert el["structural_opening_source"].startswith("BAND_OPENING")
            assert el["frame_source"] in ("NOT_DRAWN", "door block frame span")
    for ln in _lines(model):
        el = next(e for e in model["elements"] if e["element_id"] == ln["physical_element"])
        if el["leaf_count"] == 1:
            assert ln["calculation"]["inputs"]["width_m"] == pytest.approx(el["structural_opening_width_m"])


def test_door_ids_are_stable_under_translation_and_units():
    """P2-03: the same drawing moved (with its plan origin) or redrawn in metres keeps every door id."""
    _S, _M, a = H.door_run(H.door_fixture())
    d = (1e5, -3e4)
    _S, _M, b = H.door_run(H.transform(H.door_fixture(), lambda p: (p[0] + d[0], p[1] + d[1])), origin=d)
    _S, _M, c = H.door_run(H.transform(H.door_fixture(), lambda p: (p[0] * 1e-3, p[1] * 1e-3), scale=1e-3), unit_m=1.0)
    ids = sorted(el["element_id"] for el in a["elements"])
    assert len(set(ids)) == len(ids) == len(a["elements"])
    assert ids == sorted(el["element_id"] for el in b["elements"]) == sorted(el["element_id"] for el in c["elements"])


def test_released_door_totals_contain_released_doors_only(base):
    _S, _M, model = base
    tot = BL.totals(_lines(model), by=("item_code",))
    n = sum(cell["pcs"]["net"] for cell in tot.values())
    assert n == sum(1 for el in model["elements"] if el["released"])
    assert n == len(H.DOOR_TRUTH["PVC"]) + len(H.DOOR_TRUTH["ALU"]) + len(H.DOOR_TRUTH["MAIN"])


def test_v1_and_v2_never_emit_the_door_classes(tmp_path):
    from tools import arch_synth as AS
    from tools import urban_boq_m9 as M9
    p = tmp_path / "villa.dxf"
    AS.villa(str(p))
    for method in ("ARCH_TAKEOFF_V1", "ARCH_TAKEOFF_V2"):
        try:
            room, arch, _E = M9.arch_compute(str(p), method=method)
        except M9.CompletenessFailed as ex:
            room, arch = ex.room_boq, ex.arch
        codes = {ln["item_code"] for ln in room["lines"]}
        assert "ARCH.DOOR" in codes and not codes & {DM.PVC, DM.ALU, DM.MAIN}
        assert "doors" not in arch and "geometry" not in arch
    room, arch, _E = M9.arch_compute(str(p), method="ARCH_TAKEOFF_V2.1")
    codes = {ln["item_code"] for ln in room["lines"]}
    assert "ARCH.DOOR" not in codes and DM.PVC in codes and arch["doors"]["conservation_ok"]


def test_the_workbook_shows_released_provisional_and_lift_doors_apart_only_under_v2_1(tmp_path):
    from openpyxl import load_workbook
    from tools import arch_synth as AS
    from tools import urban_boq_m9 as M9
    p = tmp_path / "villa.dxf"
    AS.villa(str(p))
    for method, has in (("ARCH_TAKEOFF_V2", False), ("ARCH_TAKEOFF_V2.1", True)):
        out = tmp_path / method
        assert M9.main(["arch", "--drawings", str(p), "--out", str(out), "--method", method, "--no-svg"]) == 0
        names = load_workbook(out / "URBAN_ROOM_BOQ.xlsx").sheetnames
        assert ({"DOORS", "DOORS PROVISIONAL", "LIFT DOORS (EXCLUDED)"} <= set(names)) is has
        assert (out / "DOOR_SCHEDULE.json").exists() is has and (out / "ARCH_GEOMETRY_DECISIONS.json").exists() is has
    ws = load_workbook(tmp_path / "ARCH_TAKEOFF_V2.1" / "URBAN_ROOM_BOQ.xlsx")["DOORS"]
    totals = [r for r in ws.iter_rows(min_row=4, values_only=True) if str(r[0]).endswith("TOTAL")]
    assert len(totals) == 3 and all(str(r[11]).startswith("=IF(") for r in totals)   # never a 0 area while BLOCKED
