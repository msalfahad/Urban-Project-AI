"""OWNER-RULE APPLICATION AUDIT: the checks that stop a spacing from becoming a material.

Two faces with a measured distance between them will report a thickness whatever they are.  A column does it, a stair
shaft does it, and so does a drafting arrow.  These tests hold the line that only a masonry wall may produce a blockwork
quantity, and that a linear skirting deduction is answered by widths alone - so no missing height may ever be the reason
a skirting is unfinished.
"""

from __future__ import annotations

import json
from pathlib import Path

from research.qs_wall_treatment_01 import protocol as PR
from research.qs_wall_treatment_01.pa08.qortuba.boq import audit_owner_rules as AU, material_identity as MI

OUT = Path(PR.OUT_DIR) / "pa08_qortuba_boq"
QS = Path(PR.OUT_DIR) / "pa08_qortuba_qs01"


def reg(name, folder=OUT):
    return json.loads((folder / f"{name}.json").read_text("utf-8"))


def audit():
    return reg("QORTUBA_OWNER_RULE_APPLICATION_AUDIT")


# ------------------------------------------------------------------ §1 skirting proved segment by segment
def test_every_room_path_adds_back_up_to_its_gross_perimeter():
    s = audit()["SECTION_1_SKIRTING_AND_PROFILE"]
    assert s["EVERY_ROOM_PATH_RECONCILES"] is True
    # two rooms are called BED.ROOM, so the rows are matched in order rather than by name
    src = reg("QORTUBA_QS01_SKIRTING_TAKEOFF", QS)["ROWS"]
    assert len(s["ROWS"]) == len(src)
    for r, x in zip(s["ROWS"], src):
        assert r["ROOM"] == x["ROOM"]
        assert abs(r["GROSS_ELIGIBLE_WALL_PATH_LM"] - x["GROSS_WALL_LINE_PERIMETER_LM"]) < 1e-6


def test_the_skirting_total_is_the_sum_of_its_rooms_both_ways():
    s = audit()["SECTION_1_SKIRTING_AND_PROFILE"]
    assert abs(sum(r["NET_SKIRTING_LM"] for r in s["ROWS"]) - s["TOTAL_AS_ISSUED_LM"]) < 1e-6
    assert abs(sum(r["NET_IF_ALL_WINDOW_WIDTHS_DEDUCTED_LM"] for r in s["ROWS"])
               - s["TOTAL_IF_ALL_WINDOW_WIDTHS_DEDUCTED_LM"]) < 1e-6
    assert s["TOTAL_IF_ALL_WINDOW_WIDTHS_DEDUCTED_LM"] < s["TOTAL_AS_ISSUED_LM"]


def test_the_profile_path_equals_the_skirting_path_in_every_room():
    for r in audit()["SECTION_1_SKIRTING_AND_PROFILE"]["ROWS"]:
        assert r["NET_PROFILE_LM"] == r["NET_SKIRTING_LM"], r["ROOM"]


def test_no_opening_height_enters_the_linear_deduction():
    s = audit()["SECTION_1_SKIRTING_AND_PROFILE"]
    assert s["HEIGHT_PLAYED_NO_PART"] is True
    q = {x["QUANTITY_ID"]: x for x in reg("QORTUBA_RECALCULATED_QUANTITIES_V1")["ROWS"]}
    for qid in ("Q-01", "Q-02"):
        assert q[qid]["UNIT"] == "LM"
        assert q[qid]["USES_TEMPORARY_DEFAULT"] is False, "a width needs no door-height default"
        assert "m2" not in q[qid]["FORMULA"]


def test_a_window_is_attributed_to_one_room_and_not_to_a_shared_name():
    """This plan has two rooms called BED.ROOM. Attributing by name double-counts their windows."""
    g = reg("QORTUBA_GLAZED_ELEMENT_HOSTS")["ROWS"]
    ids = [x["ROOM_ID"] for x in g if x["ROOM_ID"]]
    assert len(ids) == len(g), "every glazed element resolves to exactly one room id"
    beds = [x for x in g if x["ROOM"] == "BED.ROOM"]
    assert len({x["ROOM_ID"] for x in beds}) == len(beds), "the two bedrooms must not share a window"
    rooms = reg("QORTUBA_ROOM_FINISH_AND_SKIRTING_RECALC")["ROWS"]
    total = sum(r["WINDOW_WITH_WALL_BELOW_WIDTH_LM"] for r in rooms)
    # each room total is published to the millimetre, so the sum carries at most one millimetre per room
    assert abs(total - sum(x["WIDTH_M"] for x in g if x["WALL_BELOW"])) <= 0.001 * len(rooms)


def test_a_ceramic_room_is_zero_whatever_its_windows_do():
    for r in audit()["SECTION_1_SKIRTING_AND_PROFILE"]["ROWS"]:
        if r["FINISH_CLASS"] == "CERAMIC_SERVICE_ROOM":
            assert r["NET_SKIRTING_LM"] == 0.0 and r["NET_IF_ALL_WINDOW_WIDTHS_DEDUCTED_LM"] == 0.0, r["ROOM"]


def test_the_column_face_stays_on_the_path_and_is_reported_as_such():
    rows = audit()["SECTION_1_SKIRTING_AND_PROFILE"]["ROWS"]
    cols = sum(r["COLUMN_FACE_INCLUDED_LM"] for r in rows if r["FINISH_CLASS"] == "DRY_ROOM")
    assert abs(cols - 1.550) < 1e-6, cols


# ------------------------------------------------------------------ §2 the blue elements
def test_every_blue_element_is_listed_with_a_room_a_width_and_a_verdict():
    b = audit()["SECTION_2_BLUE_OPENINGS"]
    assert b["COUNT"] == 7
    for x in b["ROWS"]:
        assert x["OPENING_ID"] and x["WIDTH_M"] > 0 and x["WHY"]
        assert x["SKIRTING_WIDTH_DEDUCTION_LM"] == x["PROFILE_WIDTH_DEDUCTION_LM"]
        assert x["ROOM"], x["OPENING_ID"]


def test_only_the_element_that_interrupts_the_wall_was_deducted():
    b = audit()["SECTION_2_BLUE_OPENINGS"]
    ded = [x for x in b["ROWS"] if x["SKIRTING_WIDTH_DEDUCTION_LM"] > 0]
    assert len(ded) == 1 and ded[0]["OPENING_ID"] == "BE-03"
    assert ded[0]["INTERRUPTS_THE_WALL_IN_PLAN"] is True
    for x in b["ROWS"]:
        if x["WALL_BELOW"]:
            assert x["SKIRTING_WIDTH_DEDUCTION_LM"] == 0.0, x["OPENING_ID"]


def test_type_uncertainty_does_not_change_this_linear_trade():
    b = audit()["SECTION_2_BLUE_OPENINGS"]
    assert b["DOES_TYPE_UNCERTAINTY_CHANGE_THE_LINEAR_QUANTITY"] is False
    assert b["UNRESOLVED_ELEMENTS"] == ["BE-03"]
    for x in b["ROWS"]:
        assert x["DOES_TYPE_UNCERTAINTY_CHANGE_THIS_TRADE"] is False
    assert b["WHERE_TYPE_UNCERTAINTY_DOES_MATTER"]


# ------------------------------------------------------------------ §3 object identity
def test_a_thickness_group_with_no_masonry_produces_no_blockwork():
    a = audit()["SECTION_3_BLOCKWORK_OBJECT_IDENTITY"]
    assert set(a["THICKNESS_GROUPS_WITH_NO_MASONRY_AT_ALL"]) == {219, 300, 350, 450, 550, 600}
    q = {x["QUANTITY_ID"]: x for x in reg("QORTUBA_RECALCULATED_QUANTITIES_V1")["ROWS"]}
    for t in (219, 300, 350, 450, 550, 600):
        x = q[f"Q-07-{t}"]
        assert x["STATUS"] == "NOT_APPLICABLE" and x["MEASURED_NET_QUANTITY"] is None, t


def test_the_drafting_arrow_never_reaches_a_bill():
    row = [r for r in reg("QORTUBA_WALL_OBJECT_IDENTITY")["ROWS"] if int(r["THICKNESS_MM"]) == 219][0]
    assert row["PHYSICAL_OBJECT_CLASS"] == "GEOMETRIC_PAIR_ONLY"
    assert row["CAD_LAYERS"] == ["arrow2"] and row["FACE_ROLES"] == ["UNKNOWN_GEOMETRY", "UNKNOWN_GEOMETRY"]
    assert row["BLOCKWORK_CONFIRMED"] is False


def test_a_column_band_is_never_billed_as_a_wall():
    for r in reg("QORTUBA_WALL_OBJECT_IDENTITY")["ROWS"]:
        if r["BAND_TYPE"] == "COLUMN_BAND" or r["FACE_ROLES"] == ["COLUMN_FACE", "COLUMN_FACE"]:
            assert r["PHYSICAL_OBJECT_CLASS"] in ("COLUMN", "PIER"), r["WALL_ID"]
            assert r["BLOCKWORK_CONFIRMED"] is False, r["WALL_ID"]


def test_stair_layer_geometry_is_not_apartment_masonry():
    for r in reg("QORTUBA_WALL_OBJECT_IDENTITY")["ROWS"]:
        if "STAIR" in r["CAD_LAYERS"]:
            assert r["PHYSICAL_OBJECT_CLASS"] == "SHAFT" and r["BLOCKWORK_CONFIRMED"] is False, r["WALL_ID"]


def test_every_band_carries_a_class_from_the_list_and_its_evidence():
    for r in reg("QORTUBA_WALL_OBJECT_IDENTITY")["ROWS"]:
        assert r["PHYSICAL_OBJECT_CLASS"] in MI.OBJECT_CLASSES, r["WALL_ID"]
        assert r["EVIDENCE"] and all(len(e) > 15 for e in r["EVIDENCE"]), r["WALL_ID"]
        assert r["STATUS"] in ("BLOCKWORK_ELIGIBLE", "BLOCKWORK_ELIGIBLE_THICKNESS_ITEM_AMBIGUOUS",
                               "GEOMETRIC_REFERENCE_ONLY")


def test_the_plan_lengths_are_untouched_by_the_reclassification():
    w = reg("QORTUBA_QS01_BLOCK_WALL_TAKEOFF", QS)
    ident = {r["WALL_ID"]: r for r in reg("QORTUBA_WALL_OBJECT_IDENTITY")["ROWS"]}
    assert len(ident) == w["COUNT"]
    for r in w["ROWS"]:
        assert ident[r["WALL_ID"]]["LENGTH_M"] == r["LENGTH_M"]
        assert ident[r["WALL_ID"]]["THICKNESS_MM"] == r["THICKNESS_MM"]
    a = audit()["SECTION_3_BLOCKWORK_OBJECT_IDENTITY"]
    assert abs(a["CONFIRMED_MASONRY_LENGTH_M"] + a["EXCLUDED_LENGTH_M"]
               - sum(r["LENGTH_M"] for r in w["ROWS"])) < 1e-6


# ------------------------------------------------------------------ §4 the formula
def test_the_blockwork_formula_holds_and_adds_no_reveal():
    f = audit()["SECTION_4_BLOCKWORK_FORMULA"]
    assert f["FORMULA_DEFECTS"] == []
    assert f["GROSS_IS_LENGTH_TIMES_3_00"] and f["DEDUCTION_IS_FULL_OPENING_AREA"]
    assert f["REVEALS_ADDED_TO_BLOCKWORK"] == []


def test_a_pending_opening_stays_on_its_own_thickness_group():
    q = {x["QUANTITY_ID"]: x for x in reg("QORTUBA_RECALCULATED_QUANTITIES_V1")["ROWS"]}
    glz = {g["BLUE_ELEMENT_ID"]: g for g in reg("QORTUBA_GLAZED_ELEMENT_HOSTS")["ROWS"]}
    for qid, x in q.items():
        if not qid.startswith("Q-07-"):
            continue
        t = int(qid.split("-")[-1])
        for o in x["RESIDUAL_OPENINGS"]:
            if o.get("OPENING_ID") in glz:
                assert int(glz[o["OPENING_ID"]]["HOST_WALL_THICKNESS_MM"]) == t, (qid, o["OPENING_ID"])


# ------------------------------------------------------------------ §5 ceramic
def test_tile_preparation_reconciles_with_wall_ceramic_exactly():
    c = audit()["SECTION_5_CERAMIC_AND_TILE_PREPARATION"]
    assert c["Q10_RECONCILES_WITH_WALL_CERAMIC"] is True
    assert abs(c["TOTAL"]["NET_M2"] - c["TILE_PREPARATION_REPORTED"]) < 1e-6
    assert abs(c["BATHROOMS"]["NET_M2"] + c["PAINTRY"]["NET_M2"] - c["TOTAL"]["NET_M2"]) < 1e-6
    assert c["WHY_THE_SAME_BASIS"]


def test_the_ceramic_rooms_carry_no_skirting_no_profile_and_no_paint():
    c = audit()["SECTION_5_CERAMIC_AND_TILE_PREPARATION"]
    assert c["SKIRTING_AND_PROFILE_IN_THESE_ROOMS"] == 0.0
    assert c["PAINT_IN_THESE_ROOMS"] == 0.0
    assert c["NO_CERAMIC_ROOM_IN_THE_PAINT_ITEM"] is True
    assert c["WALL_TILE_HEIGHT_M"] == 3.00 and c["HEIGHT_RULE"] == "QP-01"
    assert c["FLOOR_CERAMIC_INCLUDED"]["BOTH_PRESENT"] is True


# ------------------------------------------------------------------ §6 plaster and paint
def test_every_room_names_its_own_unresolved_openings_only():
    p = audit()["SECTION_6_PLASTER_AND_PAINT"]
    ops = {o["OPENING_ID"]: o for o in reg("QORTUBA_OPENING_REGISTER")["ROWS"]}
    rooms = {r["ROOM"]: r for r in reg("QORTUBA_ROOM_FINISH_AND_SKIRTING_RECALC")["ROWS"]}
    for r in p["ROWS"]:
        name = rooms[r["ROOM"]]["ROOM_NAME"]
        for o in r["OPENINGS_WITH_NO_HEIGHT"]:
            assert name in ops[o["OPENING_ID"]]["ROOMS"], (r["ROOM"], o["OPENING_ID"])
    assert p["ALL_KNOWN_DOOR_DIMENSIONS_DEDUCTED"] is True
    assert p["PLASTER_EQUALS_PAINT"] is True


def test_the_reveal_rule_is_unchanged():
    p = audit()["SECTION_6_PLASTER_AND_PAINT"]
    assert p["REVEAL_RULE_KEPT"] == "0.25 m left, right and top; no sill"
    assert p["REVEAL_M2_IN_THE_FIGURE"] > 0


# ------------------------------------------------------------------ §7 what survives
def test_the_demotions_are_named_with_an_exact_reason():
    s = audit()["SECTION_7_FINAL_STATUS"]
    demoted = {d["QUANTITY_ID"] for d in s["FINAL_QUANTITIES_DEMOTED"]}
    assert demoted == {"Q-01", "Q-02", "Q-07-219", "Q-07-300", "Q-07-350", "Q-07-450", "Q-07-550", "Q-07-600"}
    for d in s["FINAL_QUANTITIES_DEMOTED"]:
        assert len(d["EXACT_REASON"]) > 60 and d["NEW_STATUS"] != "FINAL_QUANTITY_AVAILABLE"
    assert s["FINAL_AFTER_AUDIT"] < s["FINAL_BEFORE_AUDIT"]


def test_nothing_survives_as_final_without_identity_and_rule():
    ident = {r["WALL_ID"]: r for r in reg("QORTUBA_WALL_OBJECT_IDENTITY")["ROWS"]}
    for x in reg("QORTUBA_RECALCULATED_QUANTITIES_V1")["ROWS"]:
        if x["STATUS"] != "FINAL_QUANTITY_AVAILABLE":
            continue
        assert x["RESIDUAL_OPENINGS"] == [], x["QUANTITY_ID"]
        assert x["MEASURED_NET_QUANTITY"] is not None
    for w in reg("QORTUBA_BLOCKWORK_RECALC")["ROWS"]:
        if not ident[w["WALL_ID"]]["BLOCKWORK_CONFIRMED"]:
            assert w["GROSS_AREA_M2"] is None and w["NET_AREA_M2"] is None, w["WALL_ID"]
            assert w["STATUS"] == "GEOMETRIC_REFERENCE_ONLY"


def test_the_audit_changed_no_geometry_and_added_no_rule():
    a = audit()
    assert a["GEOMETRY_MODIFIED"] == "NONE" and a["RULES_ADDED"] == 0 and a["PHASES_CREATED"] == 0
    fz = reg("FREEZE_QORTUBA_OWNER_RULE_APPLICATION_AUDIT")
    assert AU.OUT.joinpath("QORTUBA_OWNER_RULE_APPLICATION_AUDIT.json").exists()
    import hashlib
    for n, sha in fz["CONTENTS"].items():
        assert hashlib.sha256((OUT / f"{n}.json").read_bytes()).hexdigest() == sha
    rules = reg("URBAN_OWNER_RULES_V1")
    assert rules["COUNTS"] == {"URBAN_STANDARD": 10, "QORTUBA_PROJECT_RULE": 8,
                               "APPROVED_TEMPORARY_DEFAULT": 2, "SUPERSEDED": 7}


def test_the_workbook_shows_the_object_identity():
    from openpyxl import load_workbook
    from research.qs_wall_treatment_01.pa08.qortuba.boq import workbook_boq as WB
    wb = load_workbook(WB.FILE, read_only=True)
    assert "هوية الجدران Identity" in wb.sheetnames
    ws = wb["هوية الجدران Identity"]
    head = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert "الصنف / Object class" in head and "الدليل / Evidence" in head
