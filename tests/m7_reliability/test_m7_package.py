"""M7 package: every committed deliverable present, hygiene over the package and the M7 code, the engine freeze against
git, the held-out summary against its own records, the post-freeze changes exactly the three named fixes, the CSV
registers rebuilt byte for byte from the records, the workbook and the visual QA PDF (generated, never in git) rebuilt
from the records, and (with the private drawings) the ST7757 records rebuilt live."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import subprocess
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
PKG = ROOT / "research/urban_engine_m7_reliability"
DXF = ROOT / "data/inputs/by_sha256/9f9d1179a5d2a6635f9b412915a72184b7db1ff7729f4ab39a72dc3391738079.dxf"
ARCH = ROOT / "data/inputs/by_sha256/ab54dd554c31fe4a6a63ff773dc6d034159a736c7dd78158483b473a4ed31cc4.dxf"
DELIVERABLES = ("00_README.md", "01_EXECUTIVE_REPORT.md", "02_BASELINE_AND_ARCHITECTURE.md",
                "03_ENGINE_WEAKNESS_REGISTER.csv", "04_GEOMETRY_AUDIT.md", "05_DRAWING_INTERPRETATION.md",
                "06_COMPLETENESS_RECONCILIATION.csv", "07_MEASUREMENT_RELEASE_REGISTER.csv",
                "08_INDEPENDENT_VALIDATION.md", "09_SECOND_VILLA_BLIND_TEST.md", "10_BEFORE_AFTER_QUANTITIES.csv",
                "11_QUOTATION_RISK.md", "13_REPRODUCIBILITY.md", "15_CHATGPT_HANDOFF.md", "00_ENGINE_FREEZE.json")
# 12_REGRESSION_RESULTS.md and TEST_GATE_M7.json record the full-suite run of this package, so they come after it
HYGIENE = re.compile(r"YEARS 2013|7757-[A-Z]|[؀-ۿ]{3,}|(?i:als[ey]?n[ya]?an)")
M7_CODE = ("engine/source/quantity_release.py", "engine/source/boq_reconcile.py", "tools/m7_metamorphic.py",
           "tools/m7_evaluate.py", "tools/m7_independent.py", "tools/urban_project_run.py",
           "engine/profiles/URBAN_M7_GENERIC_CONVENTIONS_V5.json")
POST_FREEZE = {"engine/source/generic_route.py", "engine/source/concrete_completeness.py", "tools/m7_independent.py"}


def _module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_every_deliverable_is_present():
    missing = [f for f in DELIVERABLES if not (PKG / f).exists()]
    assert not missing, missing


def test_hygiene_over_the_package_and_the_m7_code():
    files = [p for p in PKG.rglob("*") if p.is_file() and p.suffix in (".md", ".csv", ".json", ".py")]
    files += [ROOT / f for f in M7_CODE]
    bad = {str(p.relative_to(ROOT)): sorted(set(HYGIENE.findall(p.read_text(encoding="utf-8"))))[:3]
           for p in files if HYGIENE.search(p.read_text(encoding="utf-8"))}
    assert not bad, bad
    assert not [p for p in PKG.rglob("*") if p.suffix.lower() in (".dxf", ".dwg", ".png", ".jpg")]
    tracked = subprocess.run(["git", "-C", str(ROOT), "ls-files", str(PKG.relative_to(ROOT))], capture_output=True,
                             text=True, check=True).stdout.split()
    assert not [p for p in tracked if p.lower().endswith((".pdf", ".xlsx", ".dxf", ".dwg", ".png", ".jpg", ".zip"))]


def test_the_freeze_holds_the_files_as_committed():
    fz = json.loads((PKG / "00_ENGINE_FREEZE.json").read_text(encoding="utf-8"))
    assert fz["commit"] == "af2822a" and all(fz["git_blob_matches_commit"].values())
    for f, h in fz["files_sha256"].items():
        blob = subprocess.run(["git", "-C", str(ROOT), "show", f"{fz['commit']}:{f}"], capture_output=True,
                              check=True).stdout
        assert hashlib.sha256(blob).hexdigest() == h, f
    changed = {f for f, h in fz["files_sha256"].items() if hashlib.sha256((ROOT / f).read_bytes()).hexdigest() != h}
    # exactly the post-freeze fixes P1-P3, nothing else - and, after M7, the files a later milestone's own freeze
    # records as changed since this one (M8: research/urban_engine_m8_full_scope/00_ENGINE_FREEZE.json, whose package
    # test checks their hashes)
    later = set()
    m8 = ROOT / "research" / "urban_engine_m8_full_scope" / "00_ENGINE_FREEZE.json"
    if m8.exists():
        later = set(json.loads(m8.read_text(encoding="utf-8"))["changed_since_m7_freeze"])
    assert POST_FREEZE <= changed and changed - POST_FREEZE <= later, changed


def test_the_heldout_summary_is_its_own_records():
    S = json.loads((PKG / "heldout" / "HELDOUT_SUMMARY.json").read_text(encoding="utf-8"))
    assert S["engine_commit"] == "af2822a" and S["seeds"] == [301, 340]
    for name, h in S["files_sha256"].items():
        assert hashlib.sha256((PKG / "heldout" / name).read_bytes()).hexdigest() == h, name
    for k, s in S["runs"].items():
        rec = json.loads((PKG / "heldout" / f"{k}.json").read_text(encoding="utf-8"))["summary"]
        assert {x: rec[x] for x in s if x != "silent"} == {x: s[x] for x in s if x != "silent"}, k
    r = S["runs"]
    assert r["Q_A_V5"]["silent_rows"] == r["Q_B_V5"]["silent_rows"] == 0
    assert r["C_A_V5"]["detected"] == r["C_A_V5"]["injected"] and r["C_B_V5"]["detected"] == r["C_B_V5"]["injected"]
    assert all(v["fail"] == 0 for v in r["P_V5"].values())


def test_the_csv_registers_rebuild_from_the_records():
    B = _module("m7_build_package_t", PKG / "build_package.py")
    for name, text in B.files().items():
        assert text == (PKG / name).read_text(encoding="utf-8"), name


def test_the_workbook_has_its_sheets_and_one_rate_input(tmp_path):
    from openpyxl import load_workbook
    B = _module("m7_build_package_w", PKG / "build_package.py")
    B.workbook(tmp_path / "w.xlsx")
    wb = load_workbook(tmp_path / "w.xlsx")
    assert wb.sheetnames == ["Dashboard", "Quantities_by_floor", "Family_reconciliation", "Measured_vs_unresolved",
                             "Detection_classification", "Validation", "Quantity_error", "Financial_exposure",
                             "Traceability"]
    assert wb["Dashboard"]["B2"].value == 26.5
    assert all(str(wb["Financial_exposure"][f"C{i}"].value).endswith("*Dashboard!$B$2") for i in range(2, 7))


def test_the_visual_qa_pdf_rebuilds_from_the_records(tmp_path):
    V = _module("m7_build_visuals_t", PKG / "build_visuals.py")
    out = V.build(tmp_path / "qa.pdf")
    data = out.read_bytes()
    assert data[:5] == b"%PDF-" and len(data) > 10_000
    assert len(re.findall(rb"/Type\s*/Page[^s]", data)) == 7                         # contents + six pages
    assert out.read_bytes() == V.build(tmp_path / "qa2.pdf").read_bytes()          # no date, no machine state


@pytest.mark.skipif(not (DXF.exists() and ARCH.exists()), reason="prerequisite: the private ST7757 drawings")
def test_live_rebuild_of_the_st7757_records():
    S = _module("m7_st7757_run_t", PKG / "st7757_run.py")
    files, _, _, _ = S.compute(cache=None, ablation=False)
    files.pop("ABLATION.csv")
    for name, text in files.items():
        assert text == (PKG / "st7757" / name).read_text(encoding="utf-8"), name
