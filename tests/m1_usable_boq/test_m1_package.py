"""M1 package (ST7757): the provisional planning estimate over the frozen S1 - S9A stages.

Re-derived here independently of the builder where the registers allow it: the verified totals from S9's own
reconciliation tables, the beam readings from S9A 05's B, H, L and t, the neck volumes from the printed founding rule,
the planning totals from the item registers. The convention set is exercised by override and restored. Passing tests
prove bookkeeping and arithmetic, not engineering approval; nothing here is released."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

import pytest

from engine.source import delta_release as DR
from engine.source import provisional_estimate as PE

ROOT = Path(__file__).resolve().parents[2]
PKG = ROOT / "research" / "alsenan_m1_usable_boq"
BOQ = PKG / "boq"
S9 = ROOT / "research" / "alsenan_structural_s9"
S9A = ROOT / "research" / "alsenan_structural_s9a"
MANIFEST = PKG / "06_M1_FREEZE_MANIFEST.json"
ST_DXF = ROOT / "data/inputs/by_sha256/9f9d1179a5d2a6635f9b412915a72184b7db1ff7729f4ab39a72dc3391738079.dxf"
ST_PDF = ROOT / "data/inputs/by_sha256/74da1523f05eb3c6a4fe3ddebaef2c3d841891f003efc03bc32a3fb4553337e3.pdf"
HYG = re.compile(r"YEARS 2013|7757-[A-Z]|[؀-ۿ]{3,}")
FIREWALL = re.compile(r"registers_v3b|R9_1|alsenan_arch_truth_05|post_freeze|pre_s8_structural_completeness|"
                      r"multi_engine|freelancer|christiannp|u_c4n|U-C4N", re.I)
BENCH = ("44.19", "352.436075", "77.6305", "75.09", "39.7515", "65.965875", "59.8174", "21.8328", "12.348",
         "337.812", "333.173", "22.619", "22.916")
RELEASED = {"RELEASED_FROZEN_STAGE", "RELEASED_S9_DELTA"}


def J(p):
    return json.loads(Path(p).read_text(encoding="utf-8"))


def rows(p):
    with open(p, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def f(v):
    return None if v in (None, "", "None") else float(v)


@pytest.fixture(scope="module")
def adapter():
    sys.path.insert(0, str(PKG))
    import alsenan_adapter as A
    return A


@pytest.fixture(scope="module")
def built(adapter):
    conv = PE.load_conventions(PKG / "conventions_alsenan.json")
    T = adapter.load()
    return adapter.build_items(conv, T)[0], T, conv


@pytest.fixture(scope="module")
def summary():
    return J(BOQ / "SUMMARY.json")


def test_freeze_manifest_and_every_frozen_stage_verify():
    m = J(MANIFEST)
    DR.verify_frozen(MANIFEST, ROOT)
    assert m["references_read"] == [] and len(m["frozen_manifests_verified"]) >= 30
    for p in m["frozen_manifests_verified"]:
        DR.verify_frozen(ROOT / p, ROOT)
    assert not [i for i in m["inputs"] if FIREWALL.search(i)]
    assert not [o for o in m["outputs"] if Path(o).name in m["not_frozen"]]
    assert m["verified_release"] == {"m3": 158.070256, "kg": 18845.961347}


def test_verified_is_exactly_s9s_release(summary):
    s9c, s9s = rows(S9 / "02_CONCRETE_BOQ_RECONCILIATION.csv"), rows(S9 / "03_REINFORCEMENT_BOQ_RECONCILIATION.csv")
    rel_c = [r for r in s9c if r["LANE"] in RELEASED]
    rel_s = [r for r in s9s if r["LANE"] == "RELEASED_FROZEN_STAGE"]
    vc = [r for r in rows(BOQ / "VERIFIED_REGISTER.csv") if r["TRADE"] == "CONCRETE"]
    vs = [r for r in rows(BOQ / "VERIFIED_REGISTER.csv") if r["TRADE"] == "REINFORCEMENT"]
    assert len(vc) == len(rel_c) == 126 and len(vs) == len(rel_s) == 986
    assert sum(float(r["RELEASED_M3"]) for r in rel_c) == pytest.approx(158.070256, abs=1e-6)
    assert sum(float(r["BASE"]) for r in vc) == pytest.approx(158.070256, abs=1e-6)
    assert sum(float(r["BASE"]) for r in vs) == pytest.approx(18845.961347, abs=1e-6)
    assert summary["totals"]["CONCRETE"]["VERIFIED"] == pytest.approx(158.070256, abs=1e-6)
    assert all(r["LOW"] == r["BASE"] == r["HIGH"] and r["ASSUMPTIONS"] == "" for r in vc + vs)


def test_planning_totals_from_the_registers_and_unresolved_never_added(summary):
    reg = {c: rows(BOQ / f"{c}_REGISTER.csv") for c in ("VERIFIED", "PROVISIONAL", "UNRESOLVED")}
    for t in ("CONCRETE", "REINFORCEMENT"):
        v = sum(float(r["BASE"]) for r in reg["VERIFIED"] if r["TRADE"] == t)
        p = {k: sum(float(r[k]) for r in reg["PROVISIONAL"] if r["TRADE"] == t) for k in ("LOW", "BASE", "HIGH")}
        tot = summary["totals"][t]
        assert tot["PLANNING_BASE"] == pytest.approx(v + p["BASE"], abs=1e-6)
        assert tot["PLANNING_LOW"] == pytest.approx(v + p["LOW"], abs=1e-6)
        assert tot["PLANNING_HIGH"] == pytest.approx(v + p["HIGH"], abs=1e-6)
        assert tot["N_UNRESOLVED"] == sum(1 for r in reg["UNRESOLVED"] if r["TRADE"] == t)
    assert all(r["LOW"] == r["BASE"] == r["HIGH"] == "" for r in reg["UNRESOLVED"])
    boq = rows(BOQ / "CONTRACTOR_BOQ.csv")
    assert sum(f(r["PLANNING_BASE"]) or 0 for r in boq if r["UNIT"] == "m3") == \
        pytest.approx(summary["totals"]["CONCRETE"]["PLANNING_BASE"], abs=1e-6)
    assert sum(f(r["PLANNING_BASE"]) or 0 for r in boq if r["UNIT"] == "kg") == \
        pytest.approx(summary["totals"]["REINFORCEMENT"]["PLANNING_BASE"], abs=1e-6)
    rc, pl = summary["concrete_rc"], summary["plain_concrete"]
    assert rc["PLANNING_BASE"] + pl["PLANNING_BASE"] == pytest.approx(summary["totals"]["CONCRETE"]["PLANNING_BASE"])


def test_every_provisional_item_is_ordered_assumed_and_traceable():
    conv_ids = {r["id"] for r in rows(BOQ / "ASSUMPTIONS.csv")}
    for r in rows(BOQ / "PROVISIONAL_REGISTER.csv"):
        lo, b, hi = float(r["LOW"]), float(r["BASE"]), float(r["HIGH"])
        assert -1e-12 <= lo <= b + 1e-12 and b <= hi + 1e-12, r["ITEM_ID"]
        asm = r["ASSUMPTIONS"].split(";")
        assert asm and all(a in conv_ids or a.startswith("FROZEN:") for a in asm), r["ITEM_ID"]
        assert r["SOURCE"] and r["FORMULA"] and r["SUPPORT"] not in ("RELEASED", "ASSUMPTION_DOMINATED"), r["ITEM_ID"]
        if r["TRADE"] == "REINFORCEMENT":
            assert f(r["DIA_MM"]), r["ITEM_ID"]
        assert "kg/m3" not in (r["FORMULA"] + r["EXPRESSION"]).replace(" ", "").lower()
    owners = Counter((r["COMPONENT_ID"], r["PART"]) for c in ("VERIFIED", "PROVISIONAL", "UNRESOLVED")
                     for r in rows(BOQ / f"{c}_REGISTER.csv") if r["TRADE"] == "CONCRETE")
    assert not [k for k, n in owners.items() if n > 1]                     # one owner per component


def test_every_s9_record_has_one_m1_disposition():
    s9 = rows(S9 / "02_CONCRETE_BOQ_RECONCILIATION.csv")
    cov = rows(BOQ / "COVERAGE.csv")
    assert [r["COMPONENT_ID"] for r in cov] == [r["COMPONENT_ID"] for r in s9]
    c = Counter(r["M1"] for r in cov)
    assert "MISSING" not in c and c["SUPERSEDED_BY_S9A_BAND"] == 7 and c["COVERED_BY_PL-BASE"] == 3
    transfers = {d["COMPONENT_ID"] for d in rows(S9A / "10_S9A_DELTA_REGISTER.csv")
                 if d["CHANGE_KIND"] == "OWNERSHIP_TRANSFER"}
    assert {r["COMPONENT_ID"] for r in cov if r["M1"] == "SUPERSEDED_BY_S9A_BAND"} == transfers
    owned = {r["COMPONENT_ID"] for c in ("VERIFIED", "PROVISIONAL", "UNRESOLVED")
             for r in rows(BOQ / f"{c}_REGISTER.csv")}
    assert not (transfers & owned)                                         # a re-owned span is never counted


def test_beam_readings_recomputed_from_b_h_l_t():
    tot = {r["SET"].split(" (")[0]: r for r in rows(S9A / "06_BEAM_CONVENTION_TOTALS.csv") if r["GROUP"] == "ALL"}
    b = rows(S9A / "05_BEAM_CONVENTION_COMPARISON.csv")
    cond = [r for r in b if r["S9A_STATE"].startswith(("CONDITIONAL", "BAND_IS"))]
    A = sum(float(r["B_MM"]) * float(r["H_MM"]) * float(r["L_M"]) / 1e6 for r in cond)
    B = sum(float(r["B_MM"]) * (float(r["H_MM"]) + float(r["SLAB_T_MM"])) * float(r["L_M"]) / 1e6 for r in cond)
    assert len(cond) == 103 and A == pytest.approx(float(tot["conditional"]["TOTAL_DEPTH_M3"]), abs=1e-5)
    assert B == pytest.approx(float(tot["conditional"]["BELOW_SLAB_M3"]), abs=1e-5)
    s = J(PKG / "02_BEAM_CONVENTIONS.json")
    assert s["beams_in_S9A_05"] == len(b) == 150 and s["still_unmeasurable"] == 27
    assert s["groups"]["CONDITIONAL"]["CONVENTION_A_M3"] == pytest.approx(A, abs=1e-5)
    assert s["groups"]["CONDITIONAL"]["CONVENTION_B_M3"] == pytest.approx(B, abs=1e-5)
    assert s["convention_A_total_m3"] == pytest.approx(A + float(tot["source conflict"]["TOTAL_DEPTH_M3"]) +
                                                       float(tot["curved bands"]["TOTAL_DEPTH_M3"]), abs=1e-5)
    assert s["groups"]["RE_OWNED_TO_BAND"]["N"] == 7
    assert sum(len(json.loads(k)) and n for k, n in s["groups"]["UNMEASURABLE"]["WHY"].items()) == 27


def test_neck_volumes_follow_the_printed_founding_rule(built):
    items, T, conv = built
    s9 = T["S9_CONC_BY_ID"]
    n = 0
    for i in items:
        if i["family"] != "COLUMNS" or i["storey"] != "FOUNDATION" or i["category"] != "PROVISIONAL":
            continue
        d = i["dimensions"]
        area = d["neck_area_m2"]["SCHEDULE"] if "COLUMN_MISMATCH_SECTION" not in i["assumptions"] else \
            d["neck_area_m2"][conv.base("COLUMN_MISMATCH_SECTION")]
        assert i["base"] == pytest.approx(area * (1.00 + 1.50 - d["footing_depth_m"]), abs=1e-9)
        p = s9[i["component_id"]]["CONDITIONAL_M3"]
        if p and "COLUMN_MISMATCH_SECTION" not in i["assumptions"]:
            assert i["base"] == pytest.approx(float(p), abs=1e-6), i["component_id"]
            n += 1
    assert n >= 25


def test_registers_rederive_and_overrides_are_reversible(adapter, built):
    items, T, conv = built
    reg = {r["ITEM_ID"]: r for r in rows(BOQ / "PROVISIONAL_REGISTER.csv")}
    prov = [i for i in items if i["category"] == "PROVISIONAL"]
    assert len(prov) == len(reg)
    for i in prov:
        r = reg[i["item_id"]]
        assert (float(r["LOW"]), float(r["BASE"]), float(r["HIGH"])) == pytest.approx((i["low"], i["base"],
                                                                                       i["high"]), abs=1e-9)
    beams = lambda its: sum(i["base"] for i in its if i["family"] == "BEAMS" and i["trade"] == "CONCRETE"  # noqa
                            and i["category"] == "PROVISIONAL")
    b = adapter.build_items(PE.load_conventions(PKG / "conventions_alsenan.json",
                                                {"BEAM_DEPTH_READING": "BELOW_SLAB"}), T)[0]
    s = J(PKG / "02_BEAM_CONVENTIONS.json")
    corner = next(i for i in b if i["item_id"] == "P-C:BEAM-CORNER-SQUARES")
    assert beams(b) == pytest.approx(s["convention_B_total_m3"] + corner["base"], abs=1e-5)
    back = adapter.build_items(PE.load_conventions(PKG / "conventions_alsenan.json"), T)[0]
    assert beams(back) == pytest.approx(beams(items), abs=1e-12)
    deeper = adapter.build_items(PE.load_conventions(PKG / "conventions_alsenan.json", {"FOUNDING_DEPTH_M": 2.0}),
                                 T)[0]
    d = sum(i["base"] for i in deeper if i["trade"] == "CONCRETE" and i["category"] == "PROVISIONAL") - \
        sum(i["base"] for i in items if i["trade"] == "CONCRETE" and i["category"] == "PROVISIONAL")
    sens = {(r["CONVENTION_ID"], r["ALTERNATIVE"]): r for r in rows(BOQ / "SENSITIVITY.csv")}
    assert d == pytest.approx(float(sens[("FOUNDING_DEPTH_M", "2.0")]["CONCRETE_M3_DELTA"]), abs=1e-6)
    assert sum(1 for i in items if i["category"] == "VERIFIED") == sum(1 for i in deeper if i["category"] ==
                                                                       "VERIFIED")


def test_diagnosis_matches_the_registers():
    d = J(PKG / "01_BOTTLENECK_DIAGNOSIS.json")
    prov = rows(BOQ / "PROVISIONAL_REGISTER.csv")
    for m in d["F_depending_on_manual_approval"]:
        sel = [r for r in prov if r["BLOCKER"] == m["BLOCKER"]]
        assert m["N_ITEMS"] == len(sel)
        assert m["M3_BASE"] == pytest.approx(sum(float(r["BASE"]) for r in sel if r["TRADE"] == "CONCRETE"), abs=1e-5)
    top = d["G_five_biggest_causes_concrete"]
    assert len(top) == 5 and [t["M3_UPPER"] for t in top] == sorted((t["M3_UPPER"] for t in top), reverse=True)
    assert top[0]["BLOCKER"] == "BEAM_DEPTH_CONVENTION"
    assert sum(v["n"] for v in d["A_fully_measured_VERIFIED"].values()) == 126 + 986
    u = J(PKG / "05_USABILITY.json")
    assert u["detection_coverage"]["missing_from_m1"] == 0 and u["manual_interventions_during_the_run"] == 0
    assert u["blind_test_on_a_different_villa"].startswith("NOT RUN")
    assert "INCOMPLETE" in (PKG / "00_README.md").read_text()
    assert "INCOMPLETE" in J(S9 / "14_RELEASE_SUMMARY.json")["gate_note"]


def test_hygiene_and_no_benchmark_figure_in_any_output():
    m = J(MANIFEST)
    for rel in m["outputs"]:
        txt = (PKG / rel).read_text(encoding="utf-8")
        assert not HYG.search(txt), rel
        assert not [x for x in BENCH if x in txt], rel
    for p in (PKG / "alsenan_adapter.py", PKG / "build_m1.py", PKG / "conventions_alsenan.json"):
        assert not [x for x in BENCH if x in p.read_text()], p


def test_registry_declares_the_m1_modules():
    reg = (ROOT / "tests/structural_comparison_engine/rebar_product_registry.py").read_text()
    for m in ("engine/source/provisional_estimate.py", "engine/source/drawing_inventory.py",
              "engine/source/profile_takeoff.py", "research/alsenan_m1_usable_boq/build_m1.py",
              "research/alsenan_m1_usable_boq/alsenan_adapter.py", "tools/urban_structural_boq.py"):
        assert f'"{m}"' in reg, m


def test_workbook_if_present():
    x = BOQ / "STRUCTURAL_BOQ.xlsx"
    if not x.exists():
        pytest.skip("STRUCTURAL_BOQ.xlsx is git-ignored; build_m1.py writes it")
    from openpyxl import load_workbook
    wb = load_workbook(str(x))
    assert {"README", "BOQ_CONCRETE", "BOQ_PLAIN_CONCRETE", "BOQ_REINFORCEMENT", "VERIFIED", "PROVISIONAL",
            "UNRESOLVED", "ASSUMPTIONS", "SENSITIVITY", "EXCEPTIONS"} <= set(wb.sheetnames)
    ws = wb["BOQ_CONCRETE"]
    assert str(ws.cell(ws.max_row, 9).value).startswith("=SUM(")


@pytest.mark.skipif(not (ST_DXF.exists() and ST_PDF.exists()), reason="private ST7757 drawings not present")
def test_rebuild_is_byte_identical():
    m = J(MANIFEST)
    before = dict(m["outputs"])
    subprocess.run([sys.executable, "-I", str(PKG / "build_m1.py")], check=True, capture_output=True, cwd=str(ROOT))
    after = J(MANIFEST)["outputs"]
    assert after == before
    assert all(hashlib.sha256((PKG / k).read_bytes()).hexdigest() == v for k, v in after.items())
