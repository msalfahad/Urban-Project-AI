"""Qortuba RC1 synthetic: canonical BOQ items (legacy aliases non-additive, double counts refused, measure pairs),
room-by-room breakdowns that must reconcile, distinct sites with equal labels, the consolidated opening register
(attribute bases, no passage counted as a door, no material without authority), object-footprint authority by
scope, and the owner workbook (one additive sheet, no formula, every quantity cell = its register value, aliases
never summary lines, byte-identical rebuild)."""

from __future__ import annotations

import datetime
import time

import pytest
from openpyxl import load_workbook

from engine import boq_rc1_xlsx as BX
from engine.source import boq_canonical as BC, footprint_authority as FA, opening_register as OR
from engine.source import room_matrix as RM

TOL = 1e-6


def it(cid, layer="FLOOR_TILE", unit="m2", qty=10.0, bd=(("S1", 6.0), ("S2", 4.0)), **kw):
    return BC.item(cid, trade="T", layer=layer, unit=unit, qty=qty, status=kw.pop("status", BC.COMPLETE),
                   description_ar="ع", description_en="e",
                   breakdown=[{"key": k, "label": k, "qty": v, "kind": "ROOM"} for k, v in bd], **kw)


# ------------------------------------------------------------------------------- canonical BOQ
def test_a_legacy_alias_lives_inside_its_item_and_must_agree_with_what_it_claims():
    m = BC.build([it("F-1", legacy=[{"id": "Q-11", "relation": BC.SAME, "value": 10.0},
                                    {"id": "Q-03", "relation": BC.MISLABELLED, "value": 6.0, "keys": ["S1"]}])],
                 run={})
    assert BC.validate(m)["state"] == "PASS"
    bad = BC.build([it("F-1", legacy=[{"id": "Q-11", "relation": BC.SAME, "value": 9.0}])], run={})
    assert [e["error"] for e in BC.validate(bad)["errors"]] == ["ALIAS_VALUE_DIFFERS"]
    bad2 = BC.build([it("F-1", legacy=[{"id": "Q-03", "relation": BC.COMPONENT, "value": 6.0, "keys": ["S2"]}])],
                    run={})
    assert [e["error"] for e in BC.validate(bad2)["errors"]] == ["ALIAS_COMPONENT_DIFFERS"]


def test_two_additive_items_on_the_same_layer_and_site_are_a_double_count():
    m = BC.build([it("F-1"), it("F-2", qty=6.0, bd=(("S1", 6.0),))], run={})
    e = BC.validate(m)["errors"]
    assert e[0]["error"] == "ADDITIVE_DOUBLE_COUNT" and e[0]["shared_keys"] == ["S1"]
    ok = BC.build([it("F-1"), it("W-1", layer="WATERPROOF_MEMBRANE")], run={})
    assert BC.validate(ok)["state"] == "PASS"     # different physical layer on the same floor: not a duplicate


def test_an_alias_cannot_sit_in_two_items_or_also_be_canonical():
    a = {"id": "Q-03", "relation": BC.SAME, "value": 10.0}
    m = BC.build([it("F-1", legacy=[a]), it("W-1", layer="WP", legacy=[a])], run={})
    assert "ALIAS_IN_TWO_ITEMS" in [e["error"] for e in BC.validate(m)["errors"]]
    m2 = BC.build([it("F-1", legacy=[{"id": "W-1", "relation": BC.SAME, "value": 10.0}]), it("W-1", layer="WP")],
                  run={})
    assert "ALIAS_IS_ALSO_CANONICAL" in [e["error"] for e in BC.validate(m2)["errors"]]


def test_a_measure_pair_is_one_item_measured_twice_and_must_be_reciprocal():
    a = it("M-A", layer="MARBLE", unit="m2", qty=0.5, bd=(("T1", 0.5),), measure_pair="M-L")
    b = it("M-L", layer="MARBLE", unit="lm", qty=3.0, bd=(("T1", 3.0),), measure_pair="M-A")
    v = BC.validate(BC.build([a, b], run={}))
    assert v["state"] == "PASS" and v["measure_pairs"] == [["M-A", "M-L"]]
    c = it("M-L", layer="MARBLE", unit="lm", qty=3.0, bd=(("T1", 3.0),))
    assert "MEASURE_PAIR_NOT_RECIPROCAL" in [e["error"] for e in BC.validate(BC.build([a, c], run={}))["errors"]]


def test_breakdown_must_reconcile_and_statuses_are_guarded():
    bad = BC.build([it("F-1", qty=10.5)], run={})
    assert BC.validate(bad)["errors"][0]["error"] == "BREAKDOWN_DOES_NOT_RECONCILE"
    nonadd = BC.build([it("F-1", qty=10.5, breakdown_additive=False)], run={})
    assert BC.validate(nonadd)["state"] == "PASS"
    blk = BC.build([it("F-1", status=BC.BLOCKED)], run={})
    assert "QTY_ON_BLOCKED" in [e["error"] for e in BC.validate(blk)["errors"]]
    x = it("F-1")
    x["approved_for_boq"] = True
    assert "APPROVED_WITHOUT_RELEASE" in [e["error"] for e in BC.validate(BC.build([x], run={}))["errors"]]


# ------------------------------------------------------------------------------- room matrix
def test_row_breakdown_keeps_inside_site_effects_on_their_site_and_door_strips_as_their_own_sites():
    used = [{"site": "S1", "zones": ["HALL"], "area_m2": 10.0}, {"site": "S2", "zones": ["BED"], "area_m2": 5.0}]
    audit = [{"strip": "OP-1", "location": "INSIDE_SITE", "sides": {"S1": "C"}, "contribution_m2": -0.25,
              "in_row_total": False, "state": "SOFFIT_EXCLUDED_FROM_CEILING"},
             {"strip": "TH-1", "location": "SEPARATE_SITE", "sides": {"S1": "P", "S2": "P"}, "contribution_m2": 0.15,
              "in_row_total": True, "state": "CONTINUOUS_SAME_FINISH", "area_m2": 0.15},
             {"strip": "TH-2", "location": "SEPARATE_SITE", "sides": {"S1": "P", "X": "W"}, "contribution_m2": 0.0,
              "in_row_total": False, "state": "MARBLE_THRESHOLD_EXPLICIT"}]
    b = RM.row_breakdown(used, audit, 14.9, tol=TOL)
    assert b["reconciles"] and [e["key"] for e in b["entries"]] == ["S1", "S2", "TH-1"]
    assert b["entries"][0]["qty"] == pytest.approx(9.75) and b["entries"][2]["sides"] == ["S1", "S2"]
    assert not RM.row_breakdown(used, audit, 15.0, tol=TOL)["reconciles"]


def test_equal_labels_never_merge_two_physical_sites():
    n = RM.display_names([{"site_id": "SITE-bbbbbbbbbbbb", "zones": ["BATH"]},
                          {"site_id": "SITE-aaaaaaaaaaaa", "zones": ["BATH"]},
                          {"site_id": "SITE-cccccccccccc", "zones": ["HALL"]}])
    assert n["SITE-aaaaaaaaaaaa"] == "BATH (1) [aaaaaaaa]" and n["SITE-bbbbbbbbbbbb"] == "BATH (2) [bbbbbbbb]"
    assert n["SITE-cccccccccccc"] == "HALL"


def test_matrix_columns_must_sum_to_their_items():
    rows = [{"key": "S1", "kind": "ROOM", "display": "A"}, {"key": "S2", "kind": "ROOM", "display": "B"}]
    m = RM.matrix(rows, [it("F-1")], tol=TOL)
    assert m["state"] == "PASS" and m["rows"][0]["cells"]["F-1"] == 6.0
    m2 = RM.matrix(rows[:1], [it("F-1")], tol=TOL)
    assert m2["state"] == "FAIL" and m2["unplaced"] == [{"item": "F-1", "key": "S2"}]


# ------------------------------------------------------------------------------- opening register
def obs(occ, site, kind="DOOR", w=0.8, h=2.2, hb="OWNER_PROJECT_PARAMETER", **kw):
    return dict({"occurrence": occ, "site": site, "kind": kind, "width_m": w, "width_basis": "SOURCE",
                 "height_m": h, "height_basis": hb, "height_authority": "QP-21", "sources": [f"{occ}|{site}"]}, **kw)


def test_one_record_per_occurrence_from_its_two_sites_with_attribute_bases():
    r = OR.consolidate([obs("D1", "S1"), obs("D1", "S2"), obs("W1", "S1", kind="WINDOW", w=1.2, h=1.5)],
                       scope_sites={"S1", "S2"}, tol=TOL)
    d = r[0]
    assert (d["from_site"], d["to_site"], d["state"]) == ("S1", "S2", OR.IN_SCOPE)
    assert d["area_m2"] == pytest.approx(1.76) and d["area_basis"] == "OWNER_PROJECT_PARAMETER"
    assert r[1]["to_site"] == OR.OUTSIDE and OR.validate(r)["state"] == "PASS"


def test_observations_that_disagree_are_a_conflict_not_an_average():
    r = OR.consolidate([obs("D1", "S1"), obs("D1", "S2", w=0.9)], scope_sites={"S1", "S2"}, tol=TOL)
    assert r[0]["state"] == OR.CONFLICT and "width_m" in r[0]["conflicts"]
    assert OR.validate(r)["state"] == "FAIL"


def test_an_open_passage_is_never_a_door_and_a_material_needs_authority():
    r = OR.consolidate([obs("P1", "S1", kind="PASSAGE_X")], scope_sites={"S1"}, tol=TOL)
    r[0]["kind"] = OR.PASSAGE
    r[0]["door_present"] = True
    assert "PASSAGE_WITH_DOOR" in [e["error"] for e in OR.validate(r)["errors"]]
    m = OR.consolidate([obs("W1", "S1", kind="WINDOW", material="ALUMINIUM")], scope_sites={"S1"}, tol=TOL)
    assert "MATERIAL_WITHOUT_AUTHORITY" in [e["error"] for e in OR.validate(m)["errors"]]


def test_schedules_count_doors_entrances_and_skip_out_of_scope():
    r = OR.consolidate([obs("D1", "S1"), obs("D2", "S1", role="ENTRANCE"), obs("D3", "Z"),
                        obs("P1", "S1", kind=OR.PASSAGE, h=None, hb="NOT_ESTABLISHED")],
                       scope_sites={"S1"}, tol=TOL)
    s = OR.schedules(r)
    assert s["internal_door_count"] == 1 and s["entrance_doors"] == ["D2"] and s["out_of_scope"] == ["D3"]
    assert s["open_passages"] == ["P1"] and "P1" not in s["internal_doors"]


# ------------------------------------------------------------------------------- footprint authority
DRY = {"policy_id": "F-DRY@v1", "trade": "FLOOR_FINISH", "space_classes": ["DRY_INTERNAL_ROOM"],
       "object_classes": [FA.ANY], "treatment": "FOOTPRINT_INCLUDED",
       "excludes": [{"space_class": "SERVICE_ROOM", "text": "not a rule for wet / service floors"}]}


def test_a_dry_room_fact_resolves_dry_rooms_only():
    o = [{"key": "H1", "object_class": "FURNITURE"}]
    assert FA.resolve(site="S", space_class="DRY_INTERNAL_ROOM", trade="FLOOR_FINISH", objects=o,
                      policies=[DRY])["state"] == FA.RESOLVED
    u = FA.resolve(site="S", space_class="SERVICE_ROOM", trade="FLOOR_FINISH",
                   objects=[{"key": "H2", "object_class": "SANITARY_FIXTURE"}], policies=[DRY])
    assert u["state"] == FA.UNRESOLVED and u["treatment"] is None and u["objects"][0]["key"] == "H2"
    assert any("do not include SERVICE_ROOM" in w for w in u["rejected"][0]["why"])
    assert FA.resolve(site="S", space_class="SERVICE_ROOM", trade="FLOOR_FINISH", objects=[],
                      policies=[])["state"] == FA.NOT_REQUIRED


def test_two_policies_with_different_treatments_conflict():
    other = dict(DRY, policy_id="F-X@v1", treatment="FOOTPRINT_EXCLUDED", excludes=[])
    r = FA.resolve(site="S", space_class="DRY_INTERNAL_ROOM", trade="FLOOR_FINISH",
                   objects=[{"key": "H1", "object_class": "FURNITURE"}], policies=[DRY, other])
    assert r["state"] == FA.CONFLICT


# ------------------------------------------------------------------------------- owner workbook
def model():
    return {"00_READ_ME": {"role": "INFO", "header": ["NOTE"], "rows": [["view only"]]},
            "01_BOQ_SUMMARY": {"role": "ADDITIVE_SUMMARY", "header": ["ITEM", "QTY", "STATUS", "APPROVED"],
                               "rows": [["F-1", 10.0, "COMPUTED_SHADOW_COMPLETE", "NO - SHADOW"],
                                        ["F-2", 4.5, "AUTHORISED_SUBTOTAL", "NO - SHADOW"]],
                               "qty_cols": [1], "status_col": 2},
            "02_ROOMS": {"role": "BREAKDOWN", "header": ["ROOM", "F-1"], "rows": [["A", 6.0], ["B", 4.0]],
                         "qty_cols": [1]}}


SRC = [{"sheet": "01_BOQ_SUMMARY", "row": 0, "col": 1, "value": 10.0, "ref": "reg.F-1"},
       {"sheet": "02_ROOMS", "row": 1, "col": 1, "value": 4.0, "ref": "reg.F-1.B"}]
NOW = datetime.datetime(2026, 10, 2)


def val(p, m, **kw):
    return BX.validate(p, m, sources=kw.pop("sources", SRC), approved_col=3, canonical_ids=["F-1", "F-2"],
                       legacy_ids=["Q-03"], **kw)


def test_the_workbook_reads_back_exactly_and_states_which_sheet_is_additive(tmp_path):
    p = tmp_path / "rc1.xlsx"
    BX.write(model(), p, created=NOW)
    v = val(p, model())
    assert v["state"] == "PASS" and v["quantity_cells_checked"] == 2
    ws = load_workbook(p)["02_ROOMS"]
    assert ws["A2"].value.startswith("BREAKDOWN - NOT ADDITIVE") and ws.freeze_panes == "B4"


def test_formula_alias_line_register_mismatch_and_approval_are_refused(tmp_path):
    p = tmp_path / "rc1.xlsx"
    BX.write(model(), p, created=NOW)
    wb = load_workbook(p)
    wb["02_ROOMS"]["B5"] = "=B4+1"
    wb.save(p)
    v = val(p, model())
    assert v["state"] == "FAIL" and v["formulas"]
    m = model()
    m["01_BOQ_SUMMARY"]["rows"].append(["Q-03", 6.0, "COMPUTED_SHADOW_COMPLETE", "YES"])
    BX.write(m, p, created=NOW)
    v2 = val(p, m)
    assert v2["alias_lines_in_summary"] == ["Q-03"] and v2["approved_without_release"] == ["Q-03"]
    BX.write(model(), p, created=NOW)
    v3 = val(p, model(), sources=[dict(SRC[0], value=10.5)])
    assert v3["state"] == "FAIL" and v3["register_mismatches"][0]["ref"] == "reg.F-1"


def test_exactly_one_additive_sheet_and_byte_identical_rebuild(tmp_path):
    m = model()
    m["02_ROOMS"]["role"] = "ADDITIVE_SUMMARY"
    with pytest.raises(ValueError):
        BX.write(m, tmp_path / "x.xlsx", created=NOW)
    a = BX.write(model(), tmp_path / "a.xlsx", created=NOW)
    time.sleep(1.1)                  # openpyxl stamps the save time; the writer must pin it
    b = BX.write(model(), tmp_path / "b.xlsx", created=NOW)
    assert a["file_sha256"] == b["file_sha256"] and a["content_digest"] == b["content_digest"]
