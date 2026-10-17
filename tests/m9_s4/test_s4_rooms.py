"""M9 S4 - physical vs functional floor: regression tests for the exact mechanisms that lost room area on the real
villa (each reproduced on hand-built plan entities; every expected value is derived by hand from the coordinates in
tools/arch_synth.py, never taken from the engine):

    door swing on the wrong opening    a door's hinge standing 0.14 m from an archway made the archway a DOOR
    open passage as a boundary         an archway / passage with nothing drawn in it split a room from its vestibule /
                                       lobby (the room lost the floor it owns)
    wall body leaking into a room      caps placed on a line's averaged, angle-rounded offset missed the drawn segment
    doorway piece in the room          a door closed on one face only (no band opening): the 0.8 x 0.2 piece joined
    void as porcelain                  an outlined void (VOID word + X-marked rectangle) stayed floor finish
    stair below seen as floor          a flight of the plan below, seen through the opening, stayed floor finish
    shaft as an unnamed room           the same doorless footprint on every storey
    the room sum as the floor area     every m2 of a plan is in one category; the room sum is a part of it
"""

from __future__ import annotations

import math

import pytest

from engine.source import arch_floor as AF
from engine.source import arch_plan as AP
from engine.source import arch_rooms as AR
from tools import arch_synth as AS

U = 0.001
TOL = 1e-6


def _spaces(ents, roles):
    return AR.spaces(ents, roles, U, region_id="T")


def _by_code(S):
    out = {}
    for f in S["spaces"]:
        out.setdefault("/".join(f["codes"]) or "UNLABELLED", []).append(f)
    return out


@pytest.fixture(scope="module")
def floor():
    ents, roles, truth = AS.s4_floor()
    return _spaces(ents, roles), truth


# ---- door swings: a swing serves the ONE opening it fits
def test_swing_near_an_archway_serves_its_own_door(floor):
    S, _ = floor
    kinds = sorted((round(o["width_m"], 2), o["kind"]) for o in S["openings"])
    assert (1.3, "PASSAGE") in kinds                          # the archway the vestibule door's hinge stands next to
    assert kinds.count((0.8, "DOOR")) == 2
    used = [o["door_symbol"] for o in S["openings"] if o["door_symbol"]]
    assert len(used) == len(set(used)) == 2                   # each swing serves exactly one opening


def test_match_swings_scores_fit_not_proximity():
    # one 0.80 m swing hinged at the shared corner of a 0.80 door gap and a 1.30 archway gap
    door = {"opening_id": "D", "u": (0.0, 1.0), "n": (-1.0, 0.0), "s0": 4500.0, "s1": 5300.0, "o0": -2400.0,
            "o1": -2200.0}
    arch = {"opening_id": "A", "u": (1.0, 0.0), "n": (0.0, 1.0), "s0": 800.0, "s1": 2100.0, "o0": 4200.0,
            "o1": 4400.0}
    sw = {"SW": {"hinge": (2200.0, 4500.0), "radius": 800.0, "ends": [(3000.0, 4500.0), (2200.0, 5300.0)]}}
    assert AR.match_swings([arch, door], sw, U) == {"D": ["SW"]}
    assert AR.match_swings([door, arch], sw, U) == {"D": ["SW"]}    # the order of the openings does not matter
    # a 1.40 gap whose jamb is 0.16 m from a 0.89 m swing that closes a perpendicular gap: not that gap's door
    wide = {"opening_id": "W", "u": (1.0, 0.0), "n": (0.0, 1.0), "s0": 0.0, "s1": 1400.0, "o0": -200.0, "o1": 0.0}
    sw2 = {"S2": {"hinge": (0.0, 160.0), "radius": 890.0, "ends": [(0.0, 1050.0), (-890.0, 160.0)]}}
    assert AR.match_swings([wide], sw2, U) == {}


# ---- open passages: a room owns the floor it opens onto, never a boundary invented or a space it shares
def test_vestibule_through_a_clear_archway_belongs_to_the_bedroom(floor):
    S, t = floor
    r = _by_code(S)["MASTER_BEDROOM"]
    assert len(r) == 1
    assert r[0]["area_m2"] == pytest.approx(t["MASTER_BEDROOM"], abs=TOL)
    assert [a["area_m2"] for a in r[0]["attached_cells"]] == pytest.approx([t["attached_m2"]["MASTER_BEDROOM"]])
    assert r[0]["attached_cells"][0]["rule"] == "ATTACHED_THROUGH_OPEN_PASSAGE"


def test_kitchen_takes_its_lobby_up_to_the_lobby_door(floor):
    S, t = floor
    k = _by_code(S)["KITCHEN"][0]
    assert k["area_m2"] == pytest.approx(t["KITCHEN"], abs=TOL)
    assert [a["area_m2"] for a in k["attached_cells"]] == pytest.approx([t["attached_m2"]["KITCHEN"]])


def test_marked_openings_keep_their_boundary(floor):
    S, t = floor
    by = _by_code(S)
    assert by["BEDROOM"][0]["area_m2"] == pytest.approx(t["BEDROOM"], abs=TOL)     # arch symbol block in the gap
    assert by["BATHROOM"][0]["area_m2"] == pytest.approx(t["BATHROOM"], abs=TOL)
    assert sorted(f["area_m2"] for f in by["UNLABELLED"]) == pytest.approx(t["UNLABELLED"])
    marks = [o["marks"] for o in S["openings"] if o["kind"] == "PASSAGE" and not o["marks"]["clear"]]
    assert len(marks) == 2
    assert any(m["blocks"] for m in marks)                                         # the arch symbol
    assert any(m["head_lines"] == 4 for m in marks)                                # four lines along the gap


def test_a_space_larger_than_half_its_room_is_not_attached(floor):
    S, t = floor
    sh = S["zoning"]["shared"]
    assert len(sh) == 1 and sh[0]["area_m2"] == pytest.approx(2.0 * 2.0)
    assert sh[0]["rule"].startswith("NOT_ATTACHED_LARGER_THAN_50%")


def test_physical_regions_join_cells_through_every_clear_passage(floor):
    S, t = floor
    regions = S["zoning"]["physical_regions"]
    areas = sorted(round(r["area_m2"], 6) for r in regions)
    # bedroom + archway piece + vestibule; kitchen + piece + lobby; bath + piece + closet (one floor, two zones)
    assert areas == pytest.approx(sorted([t["MASTER_BEDROOM"], t["KITCHEN"], 2.0 * 1.8 + 1.0 * 0.2 + 2.0 * 2.0]))


# ---- wall bodies and doorway pieces are never floor
def test_wall_body_caps_meet_the_drawn_segments():
    ents, roles, t = AS.s4_walls()
    S = _spaces(ents, roles)
    living = _by_code(S)["LIVING"][0]
    assert living["area_m2"] == pytest.approx(t["LIVING"], abs=1e-3)              # the partition body is wall
    stub = [s for s in ents["segments"] if abs(s["a"][0] - s["b"][0]) > 1 and abs(s["a"][1] - s["b"][1]) > 1]
    ends = [p for c in AR.wall_caps(ents, ["A-WALL"], U) for p in (c["a"], c["b"])
            if min(AR._dist_point_seg(p, s["a"], s["b"]) for s in stub) < 50]
    assert len(ends) == 4 and all(min(AR._dist_point_seg(p, s["a"], s["b"]) for s in stub) < 1e-6 for p in ends)
    # the root cause it fixes: the band's averaged, angle-rounded offset line misses the drawn segment by more than
    # the topology's join tolerance (0.5 mm) at this angle
    b = next(b for b in AP._bands([s for s in stub], U))
    (ux, uy), (nx, ny) = b["u"], b["n"]
    t_end = b["hi"]
    p_avg = (ux * t_end + nx * b["A"]["off"], uy * t_end + ny * b["A"]["off"])
    assert min(AR._dist_point_seg(p_avg, s["a"], s["b"]) for s in stub) > 0.5


def test_doorway_piece_is_its_own_face_not_the_rooms():
    ents, roles, t = AS.s4_walls()
    S = _spaces(ents, roles)
    assert _by_code(S)["BATHROOM"][0]["area_m2"] == pytest.approx(t["BATH"], abs=TOL)
    assert [round(f["area_m2"], 6) for f in S["faces"] if f["class"] == "DOORWAY"] == pytest.approx([t["doorway_m2"]])


def test_closure_b_joins_the_other_wall_face_at_both_jambs():
    # closure A along y = 3200 from x 1600 to 2400; the wall's other face (y = 3000) ends at both jamb stations
    walls = {"segments": [{"id": "W1", "layer": "A-WALL", "a": (200.0, 3000.0), "b": (1600.0, 3000.0)},
                          {"id": "W2", "layer": "A-WALL", "a": (2400.0, 3000.0), "b": (4200.0, 3000.0)}]}
    from engine.source import topology as TP
    a = TP.BoundaryItem("CLOSURE|D|D|A", "SEGMENT", (1600.0, 3200.0, 2400.0, 3200.0), "DOOR_CLOSURE", ())
    b = AR.closures_b([a], walls, ["A-WALL"], U)
    assert len(b) == 1 and b[0].geometry == pytest.approx((1600.0, 3000.0, 2400.0, 3000.0))
    # a matching face on BOTH sides is ambiguous: nothing is added
    walls["segments"] += [{"id": "W3", "layer": "A-WALL", "a": (200.0, 3400.0), "b": (1600.0, 3400.0)},
                          {"id": "W4", "layer": "A-WALL", "a": (2400.0, 3400.0), "b": (4200.0, 3400.0)}]
    assert AR.closures_b([a], walls, ["A-WALL"], U) == []


# ---- voids: deducted only inside a positive outline holding a VOID word
@pytest.mark.parametrize("case", ["X_RECT_HAND_DRAWN", "SYMBOL_NO_WORD", "WORD_NO_OUTLINE"])
def test_void_outline_rules(case):
    ents, roles, t = AS.s4_voids()[case]
    S = _spaces(ents, roles)
    f = S["spaces"][0]
    assert f["codes"] == t["codes"]
    assert (f.get("void_m2") or 0.0) == pytest.approx(t["void_m2"], abs=TOL)
    parts = AF.split_face(f, S["voids"], U, True)
    if case == "X_RECT_HAND_DRAWN":
        assert parts["VOID"] == pytest.approx(t["void_m2"]) and parts["ROOM_FLOOR"] == pytest.approx(t["floor_m2"])
        assert S["voids"][0]["evidence"] == "X_RECT"
    if case == "SYMBOL_NO_WORD":
        assert S["voids"][0]["state"] == "SYMBOL_WITHOUT_VOID_WORD_NOT_DEDUCTED"
        assert parts == {"ROOM_FLOOR": pytest.approx(t["floor_m2"])}
    if case == "WORD_NO_OUTLINE":
        assert f["class"] == "MIXED" and f["issue"].startswith("VOID_EXTENT_UNRESOLVED")


def test_x_rect_second_line_must_follow_the_other_diagonal():
    P, Q, C = (3000.0, 2000.0), (7000.0, 5000.0), (7000.0, 2000.0)
    far = (3000.0, 5000.0)
    for deg, found in ((3.0, True), (9.0, False)):
        tip = AS._rot((C[0] + 0.6 * (far[0] - C[0]), C[1] + 0.6 * (far[1] - C[1])), C, deg)
        segs = [{"id": "a", "layer": "X", "a": P, "b": Q}, {"id": "b", "layer": "X", "a": C, "b": tip}]
        x = AR.x_rect_at((5000.0, 3600.0), segs, U)
        assert (x is not None) == found, deg
        if found:
            assert x["area_m2"] == pytest.approx(12.0)


# ---- shafts: the same doorless footprint on every storey
def test_stacked_doorless_footprint_is_a_shaft():
    floors = []
    for fl, ents in AS.s4_shaft_plans():
        S = _spaces(ents, AS.S4_ROLES)
        floors.append((fl, S["column_boxes"], S["faces"], S["openings_by_id"]))
    stacks = AR.stacked_shafts(floors, U)
    assert len(stacks) == 1 and stacks[0]["floors"] == ["GF", "1F"] and stacks[0]["area_m2"] == pytest.approx(3.24)
    for _fl, _c, faces, _o in floors:
        shafts = [f for f in faces if f["codes"] == ["SHAFT"]]
        assert len(shafts) == 1 and shafts[0]["label_source"] == "STACKED_FOOTPRINT"
        # the 1.8 x 1.8 store with a door on the lower plan never stacks
        assert sum(1 for f in faces if f["class"] == "UNLABELLED" and abs(f["area_m2"] - 3.24) < 1e-6) == 1


# ---- the whole route: floor conservation, voids / stairs below never porcelain
@pytest.fixture(scope="module")
def villa(tmp_path_factory):
    from tools import urban_boq_m9 as M9
    p = tmp_path_factory.mktemp("s4villa") / "villa.dxf"
    AS.villa(p)
    room, arch, _ = M9.arch_compute(str(p), method="ARCH_TAKEOFF_V1")
    return room, arch


@pytest.fixture(scope="module")
def stairs(tmp_path_factory):
    from tools import urban_boq_m9 as M9
    p = tmp_path_factory.mktemp("s4stairs") / "stairs.dxf"
    AS.stairs(p)
    room, arch, _ = M9.arch_compute(str(p), method="ARCH_TAKEOFF_V1")
    return room, arch


def test_every_m2_of_each_plan_is_in_one_category(villa, stairs):
    for _room, arch in (villa, stairs):
        assert arch["floor_areas"]
        for fa in arch["floor_areas"]:
            assert {c["check"].split(" (")[0]: c["state"] for c in fa["checks"]} == {
                "faces = categories": "BALANCED", "plate = faces": "BALANCED",
                "floor finish lines = room + open-plan floor": "BALANCED"}, fa["floor"]
            cats = fa["categories_m2"]
            assert fa["internal_physical_floor_m2"] == pytest.approx(sum(cats[k] for k in AF.INTERNAL_FLOOR), abs=1e-3)
            assert fa["of_which"]["rooms_single_label_m2"] <= fa["internal_physical_floor_m2"] + 1e-9


def test_villa_floor_plate_is_the_sum_of_its_parts(villa):
    _room, arch = villa
    gf = next(f for f in arch["floor_areas"] if f["floor"] == "GF")
    assert gf["gross_floor_plate_m2"] == pytest.approx(12.4 * 9.4, abs=1e-3)
    assert sum(gf["categories_m2"].values()) == pytest.approx(12.4 * 9.4, abs=1e-3)


def test_stair_seen_through_the_opening_is_neither_floor_nor_counted_twice(stairs):
    room, arch = stairs
    up = next(f for f in arch["floor_areas"] if f["floor"] == "1F")
    gf = next(f for f in arch["floor_areas"] if f["floor"] == "GF")
    assert up["categories_m2"]["STAIR_BELOW_SEEN"] > 0                      # the fan of the plan below
    assert gf["categories_m2"]["STAIR_BELOW_SEEN"] == 0
    assert gf["categories_m2"]["STAIR_FLIGHTS"] > 0                         # counted once, on its own plan
    for ln in room["lines"]:
        if ln["item_code"].startswith("FIN.FLOOR") and ln["location"]["floor"] == "1F":
            assert not any("seen through" in d["what"] for d in ln.get("deductions") or ()) or ln["measured_net"] >= 0


def test_floor_areas_are_written_and_in_the_workbook(villa, tmp_path):
    from openpyxl import load_workbook
    from tools import urban_boq_m9 as M9
    room, arch = villa
    res, _ = M9.write_m9(tmp_path, room=room, project_id=room["project_id"], arch=arch)
    assert "ARCH_FLOOR_AREAS.json" in res["files"]
    ws = load_workbook(tmp_path / "URBAN_ROOM_BOQ.xlsx")["FLOOR AREAS"]
    cats = [ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1)]
    assert "INTERNAL PHYSICAL FLOOR TOTAL" in cats and "RESIDUAL (outer boundary - all faces)" in cats
    totals = [ws.cell(row=r, column=4).value for r in range(4, ws.max_row + 1)
              if ws.cell(row=r, column=2).value == "INTERNAL PHYSICAL FLOOR TOTAL"]
    assert all(isinstance(v, str) and v.startswith("=SUMIFS(") for v in totals)   # totals are formulas


def test_s4_files_carry_no_client_identity():
    import re
    from pathlib import Path
    root = Path(__file__).resolve().parents[2]
    pat = re.compile("YEARS" + r" 2013|7757-[A-Z]|[\u0600-\u06ff]{3,}|(?i:als[ey]?n[ya]?an)")
    for rel in ("engine/source/arch_floor.py", "tests/m9_s4/test_s4_rooms.py", "tests/m9_s4/test_s4_real_regression.py"):
        text = (root / rel).read_text(encoding="utf-8")
        assert not pat.search(text), rel
