"""QORTUBA_FINAL_PRICING_AUDIT: the rules that keep a measured length from passing as a priced quantity.

The matrix this audit corrects called a row ready when it held a number.  These tests hold the stricter line: a row is ready
only when what exists is in the unit and on the basis the bill prices that item by, and a raw length or count never counts.
"""

from __future__ import annotations

import json
from pathlib import Path

from research.qs_wall_treatment_01 import protocol as PR
from research.qs_wall_treatment_01.pa08.qortuba.pricing import audit as AD

OUT = Path(PR.OUT_DIR) / "pa08_qortuba_pricing"
QS = Path(PR.OUT_DIR) / "pa08_qortuba_qs01"


def reg(name):
    return json.loads((OUT / f"{name}.json").read_text("utf-8"))


def rows():
    return reg("QORTUBA_FINAL_PRICING_AUDIT")["ROWS"]


def test_every_row_has_exactly_one_class_from_the_list():
    for r in rows():
        assert r["CLASS"] in AD.CLASSES, (r["ITEM_NO"], r["CLASS"])


def test_the_three_quantity_columns_are_separate():
    """MEASURED_INPUT, FINAL_BOQ_QUANTITY and UNIT are three fields, so the sheet is useful before the final exists."""
    for r in rows():
        assert "MEASURED_INPUT" in r and "FINAL_BOQ_QUANTITY" in r and "UNIT" in r
        assert r["FINAL_BOQ_QUANTITY"]["UNIT"] == r["UNIT"]
        mi = r["MEASURED_INPUT"]
        assert set(mi) == {"VALUE", "UNIT"}
        if mi["VALUE"] is not None:
            assert mi["UNIT"], r["ITEM_NO"]


def test_a_length_is_never_a_final_quantity_where_the_unit_is_an_area():
    for r in rows():
        mi = r["MEASURED_INPUT"]
        if r["UNIT"] == "M2" and mi["VALUE"] is not None and mi["UNIT"] in ("LM", "M", "M (width)", "NR"):
            assert r["CLASS"] != "FINAL_PRICING_QUANTITY", r["ITEM_NO"]
            assert r["BLOCKER"], r["ITEM_NO"]


def test_no_row_is_final_without_a_final_quantity():
    for r in rows():
        if r["CLASS"] == "FINAL_PRICING_QUANTITY":
            assert r["FINAL_BOQ_QUANTITY"]["VALUE"] is not None, r["ITEM_NO"]
        else:
            assert r["FINAL_BOQ_QUANTITY"]["VALUE"] is None, r["ITEM_NO"]


def test_every_non_final_row_states_its_blocker_and_its_formula():
    for r in rows():
        if r["CLASS"] in ("MEASUREMENT_INPUT", "SCOPE_NOT_CONFIRMED", "SOURCE_REQUIRED", "SPEC_REQUIRED"):
            assert r["BLOCKER"] and len(r["BLOCKER"]) > 5, r["ITEM_NO"]
            assert r["FINAL_QUANTITY_FORMULA"] and len(r["FINAL_QUANTITY_FORMULA"]) > 8, r["ITEM_NO"]


def test_the_named_over_permissive_rows_were_demoted():
    """The rows the audit exists to correct, checked one by one against the brief's own examples."""
    by = {r["ITEM_NO"]: r for r in rows()}
    blockwork = [r for r in rows() if r["ITEM_NO"].startswith("B-") and r["MEASURED_INPUT"]["UNIT"] == "M"]
    assert blockwork and all(r["UNIT"] == "M2" and r["CLASS"] == "MEASUREMENT_INPUT" for r in blockwork)
    plaster = [r for r in rows() if r["ITEM_NO"].startswith("IP-") and r["MEASURED_INPUT"]["VALUE"] is not None]
    assert plaster and all(r["UNIT"] == "M2" and r["CLASS"] == "MEASUREMENT_INPUT" for r in plaster)
    paint = [r for r in rows() if r["ITEM_NO"].startswith("PT-") and r["MEASURED_INPUT"]["UNIT"] == "LM"
             and r["CLASS"] != "NOT_APPLICABLE"]
    assert paint and all(r["UNIT"] == "M2" for r in paint)
    alu = [r for r in rows() if r["ITEM_NO"].startswith("AL-")]
    assert alu and not any(r["CLASS"] == "FINAL_PRICING_QUANTITY" for r in alu)
    stair = [r for r in rows() if r["ITEM_NO"].startswith("MR-") and r["MEASURED_INPUT"]["VALUE"] is not None]
    assert stair and all("not a marble quantity" in r["FINAL_QUANTITY_FORMULA"] or r["CLASS"] == "MEASUREMENT_INPUT"
                         for r in stair)
    ceiling = [r for r in rows() if r["ITEM_NO"] in ("CL-01", "CL-02")]
    assert len(ceiling) == 2 and all(r["CLASS"] == "MEASUREMENT_INPUT" for r in ceiling)


def test_a_measured_input_is_never_silently_dropped():
    """Every figure the workpaper established still appears somewhere, demoted but not deleted."""
    s = json.loads((QS / "QORTUBA_QS01_SUMMARY_TOTALS.json").read_text("utf-8"))
    present = {r["MEASURED_INPUT"]["VALUE"] for r in rows() if r["MEASURED_INPUT"]["VALUE"] is not None}
    for key in ("A_DRY_INTERNAL_FLOOR_AREA_M2", "B_WET_INTERNAL_FLOOR_AREA_M2", "D_TOTAL_SKIRTING_LM",
                "E_TOTAL_BLACK_PROFILE_LM", "F_TOTAL_CEILING_GEOMETRY_M2", "N_TOTAL_PAINT_ELIGIBLE_FACE_LENGTH_LM"):
        assert s[key]["VALUE"] in present, key


def test_readiness_never_counts_a_raw_length_as_ready():
    rd = reg("QORTUBA_CORRECTED_READINESS")
    assert rd["RAW_LENGTHS_ARE_NOT_COUNTED"]
    for x in rd["ROWS"]:
        assert x["READINESS"] in AD.READINESS
        if x["READINESS"] == "READY_TO_PRICE":
            assert x["FINAL_PRICING_QUANTITIES"] > 0, x["TRADE_EN"]
        else:
            assert x["ONE_MINIMUM_INPUT"], x["TRADE_EN"]


def test_readiness_follows_the_rows_it_summarises():
    by = {}
    for r in rows():
        pre = r["ITEM_NO"].rsplit("-", 1)[0]
        by.setdefault(pre, []).append(r)
    pre_of = {n: p for n, p, _, _ in AD.SECTIONS}
    for x in reg("QORTUBA_CORRECTED_READINESS")["ROWS"]:
        tr = by[pre_of[x["SECTION_NO"]]]
        counted = [r for r in tr if r["CLASS"] != "NOT_APPLICABLE"]
        measured = [r for r in counted if r["MEASURED_INPUT"]["VALUE"] is not None]
        assert x["ITEMS"] == len(tr) and x["ITEMS_WITH_A_MEASURED_INPUT"] == len(measured)
        expected = ("READY_TO_PRICE" if x["FINAL_PRICING_QUANTITIES"] == len(counted) and counted
                    else "PARTIALLY_READY" if measured else "NOT_READY")
        assert x["READINESS"] == expected, (x["TRADE_EN"], x["READINESS"], expected)


def test_every_trade_gets_exactly_one_minimum_input():
    rd = reg("QORTUBA_CORRECTED_READINESS")["ROWS"]
    seen = {x["SECTION_NO"] for x in rd}
    assert seen == {n for n, _, _, _ in AD.SECTIONS}
    for x in rd:
        assert isinstance(x["ONE_MINIMUM_INPUT"], str) and x["ONE_MINIMUM_INPUT"] in x["ALL_BLOCKERS"]


def test_the_minimum_input_prefers_what_releases_a_measured_row():
    """Asking for a sanitary drawing before a tile height would waste the owner's first answer."""
    by_sec = {x["SECTION_NO"]: x for x in reg("QORTUBA_CORRECTED_READINESS")["ROWS"]}
    assert "tiling height" in by_sec["1"]["ONE_MINIMUM_INPUT"]
    assert by_sec["3"]["ONE_MINIMUM_INPUT"] == "the wall height"
    assert by_sec["8"]["ONE_MINIMUM_INPUT"] == "the opening heights"


def test_blockers_are_grouped_so_one_answer_shows_what_it_frees():
    b = reg("QORTUBA_BLOCKED_FINAL_QUANTITIES")
    seen = set()
    for x in b["ROWS"]:
        assert x["BLOCKER"] not in seen
        seen.add(x["BLOCKER"])
        assert x["BLOCKS_ITEMS"] == len(x["ITEMS"]) and x["TRADES"]
    assert b["ROWS"][0]["BLOCKS_ITEMS"] >= 5, "the biggest single unlock should be visible at the top"
    assert b["TOTAL_BLOCKED_ITEMS"] == sum(x["BLOCKS_ITEMS"] for x in b["ROWS"])


def test_nothing_is_priced_and_no_quantity_was_invented():
    fr = reg("FREEZE_QORTUBA_FINAL_PRICING_AUDIT")
    assert fr["PRICED"] is False and fr["RATES_SUPPLIED"] == 0
    assert fr["NEW_QUANTITIES_CREATED"] == 0
    assert fr["GEOMETRY_ENGINE_CHANGED"] == "NONE" and fr["NEW_GEOMETRY_PHASE"] == "NONE"
    for r in rows():
        assert r["RATE"] is None and r["AMOUNT"] is None
        assert r["WASTE_PERCENT"] is None and r["PURCHASE_QTY"] is None


def test_the_audit_left_the_workpaper_untouched():
    fr = json.loads((QS / "FREEZE_PA08_QORTUBA_ROOM_BY_ROOM_QS_01.json").read_text("utf-8"))
    for n, h in fr["CONTENTS"].items():
        assert AD._sha(QS / f"{n}.json") == h, f"{n} changed"
    assert reg("FREEZE_QORTUBA_FINAL_PRICING_AUDIT")["WORKPAPER_FREEZE_VERIFIED"] == fr["DIGEST"]


def test_the_workbook_carries_every_trade_sheet_and_no_price():
    from openpyxl import load_workbook
    wb = load_workbook(OUT / "QORTUBA_PRICING_INPUT_MATRIX.xlsx")
    for n, _, ar, _ in AD.SECTIONS:
        assert f"{n}- {ar}"[:31] in wb.sheetnames, ar
    assert any("حساب الغرف" in s for s in wb.sheetnames), "the room backup sheet is missing"
    ws = wb["1- سيراميك"]
    header = [c.value for c in ws[4]]
    assert "الكمية النهائية\nFinal BOQ Quantity" in header and "كمية القياس\nMeasured Input" in header
    final_col = header.index("الكمية النهائية\nFinal BOQ Quantity") + 1
    for row in ws.iter_rows(min_row=5):
        assert row[final_col - 1].value is None, "a final BOQ quantity was written into the workbook"
