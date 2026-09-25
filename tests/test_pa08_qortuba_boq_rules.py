"""URBAN_BOQ_RULE_REGISTRY: the rules that stop one project's habit from becoming a company law.

A precedent read off a single previous bill is evidence about that bill.  These tests hold three lines: nothing is promoted
to an Urban standard without evidence from more than one project, no rule crosses from the trade it was read from into
another, and the physical measurement, the commercial rule and the payable quantity stay three separate columns.
"""

from __future__ import annotations

import json
from pathlib import Path

from research.qs_wall_treatment_01 import protocol as PR
from research.qs_wall_treatment_01.pa08.qortuba.boq import rule_registry as RR, workbook_boq as WB

OUT = Path(PR.OUT_DIR) / "pa08_qortuba_boq"
QS = Path(PR.OUT_DIR) / "pa08_qortuba_qs01"


def reg(name, folder=OUT):
    return json.loads((folder / f"{name}.json").read_text("utf-8"))


def rules():
    return reg("URBAN_BOQ_RULE_REGISTRY")


def cmap():
    return reg("QORTUBA_BOQ_CONVERSION_MAP")


# ------------------------------------------------------------------ nothing is promoted
def test_no_rule_stands_at_urban_standard():
    R = rules()
    assert R["RULES_AT_URBAN_STANDARD"] == []
    assert R["BY_LEVEL"].get("URBAN_STANDARD", 0) == 0
    for r in R["ROWS"]:
        assert r["RULE_LEVEL"] in RR.RULE_LEVELS
        assert r["RULE_LEVEL"] != "URBAN_STANDARD", r["RULE_ID"]
    assert reg("QORTUBA_RULE_SAFETY_REGISTER")["PROMOTIONS_MADE_IN_THIS_PHASE"] == 0


def test_a_safe_candidate_is_still_only_a_candidate():
    """The owner's §T correction: a candidate never applies on its own, only an approved standard does."""
    for s in reg("QORTUBA_RULE_SAFETY_REGISTER")["ROWS"]:
        assert s["SAFETY_CLASS"] in RR.SAFETY
        assert s["MAY_BE_PROMOTED_TO_URBAN_STANDARD_NOW"] is False, s["SAFETY_CLASS"]
        assert s["MAY_BE_APPLIED_TO_QORTUBA_WITHOUT_OWNER"] is False, s["SAFETY_CLASS"]
        assert s["WHY_NOT_PROMOTED"] and s["WHY_NOT_APPLIED"]


def test_a_contractor_specific_rule_never_applies_to_qortuba_by_itself():
    for r in rules()["ROWS"]:
        if r["GENERALISATION_SAFETY"] == "CONTRACTOR_SPECIFIC":
            assert r["APPLIES_TO_QORTUBA"] is False, r["RULE_ID"]
            assert r["OWNER_APPROVAL_REQUIRED"] is True, r["RULE_ID"]


# ------------------------------------------------------------------ no rule crosses a trade line
def test_a_rule_cited_on_a_conversion_row_belongs_to_that_row_s_trade():
    by_id = {r["RULE_ID"]: r for r in rules()["ROWS"]}
    cited = 0
    for x in cmap()["ROWS"]:
        rule = x["COMMERCIAL_RULE_REQUIRED"] or ""
        for rid, r in by_id.items():
            if rule.startswith(rid):
                cited += 1
                assert r["TRADE"] == x["TRADE"], (x["MEASUREMENT_INPUT_ID"], rid, r["TRADE"], x["TRADE"])
    assert cited >= 10, "the map should be citing rules by id"


def test_the_plaster_half_deduction_reaches_no_other_trade():
    plaster = [r for r in rules()["ROWS"] if r["RULE_ID"] == "R-02"][0]
    assert "نصف" in (plaster["DEDUCTION_RULE"] or "")
    for x in cmap()["ROWS"]:
        if x["TRADE"] != plaster["TRADE"]:
            assert "R-02" not in (x["COMMERCIAL_RULE_REQUIRED"] or ""), x["MEASUREMENT_INPUT_ID"]


def test_paint_is_asked_its_own_deduction_question():
    """Paint and plaster carry the same half deduction on the same old bill.  That is one subcontractor, not a rule."""
    dec = {d["DECISION_ID"]: d for d in reg("QORTUBA_OWNER_DECISIONS_REQUIRED")["ROWS"]}
    assert any("paint" in d["QUESTION"].lower() for d in dec.values())
    assert any("plaster" in d["QUESTION"].lower() for d in dec.values())
    assert len({d["DECISION_ID"] for d in dec.values() if "half-deduction" in d["QUESTION"]}) == 2, [d["QUESTION"] for d in dec.values()]


# ------------------------------------------------------------------ the three layers stay three
def test_every_row_keeps_physical_commercial_and_final_apart():
    for x in cmap()["ROWS"]:
        assert set(("PHYSICAL_MEASUREMENT", "COMMERCIAL_MEASUREMENT_RULE", "FINAL_BOQ_QUANTITY")) <= set(x)
        assert x["PHYSICAL_MEASUREMENT"]["VALUE"] == x["MEASURED_INPUT"]
        assert x["CURRENT_STATUS"] in RR.STATUSES
        if x["CURRENT_STATUS"] != "FINAL_QUANTITY_AVAILABLE":
            assert x["FINAL_BOQ_QUANTITY"]["VALUE"] is None, x["MEASUREMENT_INPUT_ID"]


def test_a_final_quantity_only_exists_where_the_units_already_agree():
    finals = [x for x in cmap()["ROWS"] if x["CURRENT_STATUS"] == "FINAL_QUANTITY_AVAILABLE"]
    assert finals, "at least the wet-area waterproofing floor should convert"
    for x in finals:
        assert x["INPUT_UNIT"] == x["TARGET_PRICING_UNIT"], x["MEASUREMENT_INPUT_ID"]
        assert x["FINAL_BOQ_QUANTITY"]["VALUE"] == x["MEASURED_INPUT"]
        assert not x["MISSING_INPUT"] and not x["COMMERCIAL_RULE_REQUIRED"]


def test_a_commercial_rule_never_overwrites_the_physical_value():
    for x in cmap()["ROWS"]:
        assert x["COMMERCIAL_MEASUREMENT_RULE"]["VALUE"] is None
        assert x["PHYSICAL_MEASUREMENT"]["STATE"] in ("SOURCE_ESTABLISHED", "NOT_MEASURED")


# ------------------------------------------------------------------ nothing was measured again
def test_every_measured_value_is_the_frozen_workpaper_value():
    s = json.loads((QS / "QORTUBA_QS01_SUMMARY_TOTALS.json").read_text("utf-8"))
    w = json.loads((QS / "QORTUBA_QS01_BLOCK_WALL_TAKEOFF.json").read_text("utf-8"))
    by = {x["MEASUREMENT_INPUT_ID"]: x["MEASURED_INPUT"] for x in cmap()["ROWS"]}
    assert by["CM-C5"] == s["D_TOTAL_SKIRTING_LM"]["VALUE"]
    assert by["CM-PR1"] == s["E_TOTAL_BLACK_PROFILE_LM"]["VALUE"]
    assert by["CM-WP1"] == s["B_WET_INTERNAL_FLOOR_AREA_M2"]["VALUE"]
    assert by["CM-WP2"] == s["H_TOTAL_BATHROOM_HOST_WALL_LM"]["VALUE"]
    assert by["CM-C1"] == s["A_DRY_INTERNAL_FLOOR_AREA_M2"]["VALUE"]
    assert by["CM-PT1"] == s["N_TOTAL_PAINT_ELIGIBLE_FACE_LENGTH_LM"]["VALUE"]
    assert by["CM-B150"] == w["TOTAL_LENGTH_M_BY_THICKNESS"]["150"]
    assert by["CM-B200"] == w["TOTAL_LENGTH_M_BY_THICKNESS"]["200"]
    assert cmap()["QORTUBA_QUANTITIES_RECOMPUTED"] == 0
    assert cmap()["NEW_QUANTITIES_INVENTED"] == 0


def test_the_cornice_run_is_the_gross_wall_line_not_the_net_skirting():
    """A cornice runs over a doorway; a skirting stops at it.  Splitting them is the point of the separate key."""
    sk = json.loads((QS / "QORTUBA_QS01_SKIRTING_TAKEOFF.json").read_text("utf-8"))["ROWS"]
    gross = round(sum(x["GROSS_WALL_LINE_PERIMETER_LM"] for x in sk), 3)
    by = {x["MEASUREMENT_INPUT_ID"]: x["MEASURED_INPUT"] for x in cmap()["ROWS"]}
    assert abs(by["CM-CL2"] + by["CM-CL3"] - gross) < 1e-6
    assert by["CM-CL2"] + by["CM-CL3"] > by["CM-C5"], "the gross line must exceed the net skirting run"


def test_the_profile_inherits_the_skirting_path_and_does_not_invent_a_second_one():
    by = {x["MEASUREMENT_INPUT_ID"]: x for x in cmap()["ROWS"]}
    assert by["CM-PR1"]["MEASURED_INPUT"] == by["CM-C5"]["MEASURED_INPUT"]
    assert by["CM-PR1"]["TRADE"] != by["CM-C5"]["TRADE"], "one path, two trade ids"
    assert by["CM-PR1"]["COMMERCIAL_RULE_REQUIRED"], "sharing a geometric path is not sharing a payable path"


def test_an_unclassified_glazed_element_is_not_folded_into_the_windows_item():
    blue = json.loads((QS / "QORTUBA_QS01_BLUE_ELEMENT_SCHEDULE.json").read_text("utf-8"))
    by = {x["MEASUREMENT_INPUT_ID"]: x for x in cmap()["ROWS"]}
    wins = [b for b in blue["ROWS"] if b["TYPE"] == "WINDOW_WITH_WALL_BELOW"]
    assert abs(by["CM-AL1"]["MEASURED_INPUT"] - round(sum(b["WIDTH_MM"] for b in wins) / 1000, 3)) < 1e-9
    assert "CM-AL3" in by and by["CM-AL3"]["MEASURED_INPUT"] > 0


# ------------------------------------------------------------------ no reference-project number crosses over
def test_the_registry_stores_structure_and_never_a_historical_number():
    for r in rules()["ROWS"]:
        for k, v in r.items():
            assert not isinstance(v, (int, float)) or isinstance(v, bool), (r["RULE_ID"], k)
    assert reg("FREEZE_URBAN_BOQ_RULE_REGISTRY")["HISTORICAL_QUANTITIES_USED_AS_QORTUBA_INPUTS"] == 0


def test_every_historical_rule_names_the_sheet_it_was_read_from():
    for r in rules()["ROWS"]:
        if r["RULE_LEVEL"] == "HISTORICAL_PRECEDENT":
            assert r["SOURCE_WORKBOOK"] and r["SOURCE_SHEET"] and r["SOURCE_ITEM"], r["RULE_ID"]
        else:
            assert r["SOURCE_WORKBOOK"] is None, "a project rule must not borrow a historical citation"


def test_the_new_item_declares_that_it_has_no_precedent():
    c = reg("QORTUBA_CANDIDATE_BOQ_ITEMS")["ROWS"][0]
    assert c["ITEM_AR"] == "بروفايل أعلى النعلة"
    assert c["HISTORICAL_PRECEDENT"] is None and c["WHY_NO_PRECEDENT"]
    assert c["RULE_LEVEL"] == "QORTUBA_PROJECT_RULE"
    assert c["URBAN_STANDARD_CANDIDATE"] is True and c["PROMOTED_TO_URBAN_STANDARD"] is False
    assert c["OWNER_APPROVAL_REQUIRED"] is True
    assert reg("QORTUBA_CANDIDATE_BOQ_ITEMS")["PROMOTED_TO_URBAN_STANDARD"] == 0


# ------------------------------------------------------------------ the questions asked
def test_every_decision_changes_at_least_one_row():
    for d in reg("QORTUBA_OWNER_DECISIONS_REQUIRED")["ROWS"]:
        assert d["ROWS_AFFECTED"], d["DECISION_ID"]
        assert d["OPTIONS"] and len(d["OPTIONS"]) >= 2
        assert d["WHY_IT_MATTERS"] and len(d["WHY_IT_MATTERS"]) > 40
        assert d["DECISION_KIND"] in ("PROJECT_INPUT", "PROJECT_RULE")


def test_no_decision_asks_for_something_the_workpaper_already_holds():
    have = ("how long is the skirting", "what is the floor area", "how many doors", "what is the perimeter")
    for d in reg("QORTUBA_OWNER_DECISIONS_REQUIRED")["ROWS"]:
        q = d["QUESTION"].lower()
        assert not any(h in q for h in have), d["QUESTION"]


def test_a_missing_drawing_is_listed_as_a_document_and_not_as_a_decision():
    D = reg("QORTUBA_OWNER_DECISIONS_REQUIRED")
    assert D["DOCUMENTS_NOT_DECISIONS"]
    for d in D["ROWS"]:
        assert "schedule" not in d["QUESTION"].lower() or "tiled" in d["QUESTION"].lower()


def test_readiness_is_reported_per_trade_and_never_as_one_percentage():
    R = reg("QORTUBA_FINAL_BOQ_READINESS")
    assert R["NO_SINGLE_READINESS_PERCENTAGE"]
    blob = json.dumps(R, ensure_ascii=False)
    assert "%" not in blob.replace("NO_SINGLE_READINESS_PERCENTAGE", "")
    for r in R["ROWS"]:
        assert set(r) == {"TRADE", "BOQ_ITEMS", "FINAL_BOQ_QUANTITY_AVAILABLE", "BLOCKED", "WHAT_BLOCKS_IT",
                          "TRADE_READY_TO_PRICE"}
        if r["TRADE_READY_TO_PRICE"]:
            assert r["BLOCKED"] == 0 and r["FINAL_BOQ_QUANTITY_AVAILABLE"] > 0


# ------------------------------------------------------------------ the frozen inputs and the output shape
def test_the_frozen_inputs_still_verify():
    checked = RR.verify_inputs()
    assert len(checked) == 4 and all(c["VERIFIED"] for c in checked)
    fz = reg("FREEZE_URBAN_BOQ_RULE_REGISTRY")
    assert fz["QORTUBA_GEOMETRY_CHANGED"] == "NONE"
    assert fz["NEW_TAKEOFF_PHASES_CREATED"] == 0
    assert fz["RULES_PROMOTED_TO_URBAN_STANDARD"] == 0
    for name, sha in fz["CONTENTS"].items():
        assert RR._sha(OUT / f"{name}.json") == sha, name


def test_the_workbook_carries_the_fifteen_house_columns_in_order():
    from openpyxl import load_workbook
    assert WB.FILE.exists(), "run workbook_boq first"
    ws = load_workbook(WB.FILE, read_only=True)["جدول الكميات BOQ"]
    head = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    assert head[:15] == ["البند", "الوصف", "وحدة التسعير", "كمية القياس", "وحدة القياس", "قاعدة التحويل",
                         "الكمية النهائية", "% الهالك", "كمية الشراء", "سعر الوحدة", "الإجمالي", "مصدر القياس",
                         "قاعدة القياس التجاري", "حالة الاعتماد", "ملاحظات"]


def test_no_rate_or_amount_is_filled_anywhere():
    from openpyxl import load_workbook
    ws = load_workbook(WB.FILE, read_only=True)["جدول الكميات BOQ"]
    for row in ws.iter_rows(min_row=5):
        for c in (8, 9, 10, 11):
            assert row[c - 1].value is None, (row[0].value, c)
