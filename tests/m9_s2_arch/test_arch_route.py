"""M9 S2 architectural route on a SYNTHETIC two-floor villa (tools/arch_synth.py): every expected value below is
written by hand from the fixture coordinates (rectangle / L areas, perimeters, door widths), never taken from the
engine. The production path is exercised: urban_boq_m9.arch_compute and the `arch` CLI write the same files the real
run writes."""

from __future__ import annotations

import json
import math
from collections import Counter
from pathlib import Path

import pytest

ezdxf = pytest.importorskip("ezdxf")

from tools import arch_synth as AS                     # noqa: E402
from tools import urban_boq_m9 as M9                   # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
FACTS = json.loads((ROOT / "research/urban_m9_s2/SYNTHETIC_FACTS_S2.json").read_text(encoding="utf-8"))
METHOD = "ARCH_TAKEOFF_V1"
TOL = 1e-6


def _run(path, facts=None, **variant):
    AS.villa(path, **variant)
    room, arch, _E = M9.arch_compute(str(path), method=METHOD, facts=facts)
    return room, arch


@pytest.fixture(scope="module")
def base(tmp_path_factory):
    return _run(tmp_path_factory.mktemp("villa") / "villa.dxf")


@pytest.fixture(scope="module")
def with_facts(tmp_path_factory):
    return _run(tmp_path_factory.mktemp("villa_f") / "villa.dxf", FACTS)


def _rooms(arch, floor):
    return {"/".join(r["codes"]): r for r in arch["rooms"] if r["floor"] == floor}


def _lines(room, floor, code, item):
    return [ln for ln in room["lines"] if ln["location"]["floor"] == floor and ln["item_code"] == item and
            ln["location"]["room_id"] == code]


# ---- regions, roles, levels, plates
def test_two_plan_regions_by_title_with_levels(base):
    _room, arch = base
    regs = {r["floor"]: r for r in arch["regions"]}
    assert set(regs) == {"GF", "1F"}
    assert all(r["stage_state"] == "MEASURED" for r in regs.values())
    assert regs["GF"]["level_m"] == 0.0 and regs["GF"]["floor_above_level_m"] == 3.5
    assert regs["1F"]["level_m"] == 3.5 and regs["1F"]["floor_above_level_m"] is None


def test_layer_roles_are_inferred_from_geometry(base):
    regs = {r["floor"]: r for r in base[1]["regions"]}
    gf = {k: v for k, v in regs["GF"]["layer_roles"].items() if v != "OTHER"}
    assert gf == {"A-WALL": "WALL", "A-DOOR": "DOOR", "A-GLAZ": "WINDOW", "A-COL": "COLUMN"}
    assert regs["1F"]["layer_roles"]["A-STAIR"] == "OTHER"          # tread lines are never walls
    assert regs["GF"]["roles_source"] != "OVERRIDE"


def test_floor_plates_match_the_exterior_outline(base):
    regs = {r["floor"]: r for r in base[1]["regions"]}
    assert regs["GF"]["plate"]["largest_m2"] == pytest.approx(12.4 * 9.4, abs=TOL)
    assert regs["1F"]["plate"]["largest_m2"] == pytest.approx(9.4 * 8.0, abs=TOL)


# ---- rooms: area, gross, obstructions, boundary, openings
@pytest.mark.parametrize("floor", ["GF", "1F"])
def test_every_room_area_and_boundary(base, floor):
    got = _rooms(base[1], floor)
    truth = AS.TRUTH[floor]["rooms"]
    assert set(got) == set(truth)
    for code, t in truth.items():
        r = got[code]
        assert r["class"] == "ROOM", code
        assert r["area_m2"] == pytest.approx(t["net_m2"], abs=TOL), code
        assert r["gross_m2"] == pytest.approx(t["gross_m2"], abs=TOL), code
        assert r["obstructions"] == t["holes"], code
        assert r["runs_lm"]["boundary"] == pytest.approx(t["boundary_lm"], abs=TOL), code
        assert r["runs_lm"]["door"] == pytest.approx(t["doors_lm"], abs=TOL), code
        assert r["runs_lm"]["passage"] == pytest.approx(t["passages_lm"], abs=TOL), code
        assert r["runs_lm"]["window_threshold"] == pytest.approx(t["sill_windows_lm"], abs=TOL), code


def test_l_shaped_room_with_two_columns_is_measured_exactly(base):
    r = _rooms(base[1], "GF")["LIVING"]
    assert r["area_m2"] == pytest.approx(9.0 * 4.0 + 6.0 * 5.0 - 2 * 0.16, abs=TOL)
    assert r["rect"] is None and "polygon" in r["dims"]             # never a bounding rectangle


def test_arabic_room_word_names_the_room(base):
    r = _rooms(base[1], "GF")["BEDROOM"]
    assert r["name_ar"] and r["name_en"] == "Bedroom"


@pytest.mark.parametrize("floor", ["GF", "1F"])
def test_openings_counts_and_widths(base, floor):
    ops = [o for o in base[1]["openings"] if o["floor"] == floor]
    assert dict(Counter(o["kind"] for o in ops)) == {k: v for k, v in AS.TRUTH[floor]["openings"].items() if v}
    assert sorted(o["width_m"] for o in ops if o["kind"] == "DOOR") == pytest.approx(AS.TRUTH[floor]["door_widths_m"])
    assert all(o["wall_t_m"] == pytest.approx(0.2, abs=1e-6) for o in ops)


def test_the_mirrored_door_insert_is_a_door(base):
    # the LIVING | KITCHEN door is inserted with yscale -0.9 (mirrored): it must close its 0.90 gap as a DOOR
    ops = [o for o in base[1]["openings"] if o["floor"] == "GF" and o["kind"] == "DOOR"]
    kitchen = _rooms(base[1], "GF")["KITCHEN"]
    assert len(kitchen["doors"]) == 1
    d = next(o for o in ops if o["opening_id"] == kitchen["doors"][0])
    assert d["width_m"] == pytest.approx(0.9) and d["door_symbol"]


# ---- lines without facts: what is measured, what is blocked
def test_floor_ceiling_skirting_lines(base):
    room, arch = base
    for floor in ("GF", "1F"):
        rooms = _rooms(arch, floor)
        for code, t in AS.TRUTH[floor]["rooms"].items():
            rid = rooms[code]["room_id"]
            if t["kind"] == "STAIR":
                assert not _lines(room, floor, rid, "FIN.FLOOR.PORCELAIN") and _lines(room, floor, rid, "ARCH.SPACE")
                continue
            item = "FIN.FLOOR.WET_TILE" if t["kind"] == "WET" else "FIN.FLOOR.PORCELAIN"
            (fl,) = _lines(room, floor, rid, item)
            assert fl["measured_net"] == pytest.approx(t["net_m2"], abs=TOL)
            assert fl["gross"] - sum(d["m2"] for d in fl["deductions"]) == pytest.approx(fl["measured_net"], abs=TOL)
            (ce,) = _lines(room, floor, rid, "FIN.CEILING")
            assert ce["evidence_state"] == "PROVISIONAL" and ce["measured_net"] == pytest.approx(t["net_m2"], abs=TOL)
            (sk,) = _lines(room, floor, rid, "FIN.SKIRTING")
            if t["kind"] == "WET":
                assert sk["evidence_state"] == "BLOCKED_RFI" and sk["measured_net"] is None      # UP-CER-004
            else:
                assert sk["measured_net"] == pytest.approx(t["boundary_lm"] - t["doors_lm"] - t["passages_lm"],
                                                           abs=TOL)


def test_wet_room_waterproofing(base):
    room, arch = base
    for floor in ("GF", "1F"):
        for code, t in AS.TRUTH[floor]["rooms"].items():
            if t["kind"] != "WET":
                continue
            rid = _rooms(arch, floor)[code]["room_id"]
            (wf,) = _lines(room, floor, rid, "FIN.WP.FLOOR")
            dws = _lines(room, floor, rid, "FIN.WP.DOORWAY")
            (up,) = _lines(room, floor, rid, "FIN.WP.UPTURN")
            assert wf["measured_net"] == pytest.approx(t["net_m2"], abs=TOL)
            assert sorted(d["measured_net"] for d in dws) == pytest.approx(sorted(w * d for w, d in t["doorways"]))
            assert up["measured_net"] == pytest.approx(t["boundary_lm"] * 0.20, abs=TOL)   # doorway included


def test_wall_finishes_are_blocked_without_heights_and_name_the_level_interval(base):
    room, _ = base
    pl = [ln for ln in room["lines"] if ln["item_code"] == "FIN.PLASTER"]
    assert pl and all(ln["evidence_state"] == "BLOCKED_RFI" for ln in pl)
    gf = [ln for ln in pl if ln["location"]["floor"] == "GF"]
    assert all("3.50 m" in ln["rfi"]["question"] for ln in gf)       # the plan-level interval, never used as height


def test_stair_flight_is_recognised_and_blocked_without_a_riser(base):
    room, arch = base
    (st,) = [ln for ln in room["lines"] if ln["item_code"] == "FIN.STAIR.FLIGHT"]
    t = AS.TRUTH["1F"]["stair"]
    i = st["calculation"]["inputs"]
    assert (i["tread_lines"], i["going_m"], i["width_m"], i["run_m"]) == (t["tread_lines"], t["going_m"],
                                                                            t["width_m"], t["run_m"])
    assert st["evidence_state"] == "BLOCKED_RFI" and "riser" in st["rfi"]["question"]


def test_materials_are_blocked_without_a_selected_size(base):
    room, _ = base
    tiles = [m for m in room["materials"] if m["method"].startswith("TILE")]
    assert tiles and all(m["state"] == "BLOCKED" and "FLOOR_FINISH_SPEC_NOT_ESTABLISHED" in m["blocked_by"]
                         for m in tiles)
    assert not room["comparisons"]                                   # no joint width: nothing compared


# ---- with project facts (synthetic test parameters)
def _by(room, arch, floor, code, item):
    return _lines(room, floor, _rooms(arch, floor)[code]["room_id"], item)


def test_wall_finishes_with_heights_and_opening_heights(with_facts):
    room, arch = with_facts
    H, Hd, Hw, Ht = 3.0, 2.1, 1.5, 2.4
    t = AS.TRUTH["GF"]["rooms"]["LIVING"]
    (pl,) = _by(room, arch, "GF", "LIVING", "FIN.PLASTER")
    assert pl["measured_net"] == pytest.approx(t["boundary_lm"] * H - t["doors_lm"] * Hd - 1.5 * Hw, abs=TOL)
    (pt,) = _by(room, arch, "GF", "LIVING", "FIN.PAINT")
    assert pt["measured_net"] == pytest.approx(pl["measured_net"], abs=TOL)
    b = AS.TRUTH["GF"]["rooms"]["BATHROOM"]
    (wt,) = _by(room, arch, "GF", "BATHROOM", "FIN.WALL.TILE")
    assert wt["measured_net"] == pytest.approx(b["boundary_lm"] * Ht - 0.8 * Hd - 0.6 * Hw, abs=TOL)
    assert not _by(room, arch, "GF", "BATHROOM", "FIN.PAINT")       # a wet room is tiled, not painted
    assert "WINDOW_POSITION_ASSUMED" in wt["notes"]


def test_wet_skirting_follows_the_wall_tile_height(with_facts, tmp_path):
    room, arch = with_facts
    b = AS.TRUTH["GF"]["rooms"]["BATHROOM"]
    (sk,) = _by(room, arch, "GF", "BATHROOM", "FIN.SKIRTING")      # tile 2.40 below the 3.00 clear height
    assert sk["measured_net"] == pytest.approx(b["boundary_lm"] - b["doors_lm"], abs=TOL)
    full = json.loads(json.dumps(FACTS))
    full["wall_heights"]["*"]["WALL_TILE"]["value_m"] = 3.0          # full height: UP-CER-004, no skirting
    room2, arch2 = _run(tmp_path / "v.dxf", full)
    (sk2,) = _by(room2, arch2, "GF", "BATHROOM", "FIN.SKIRTING")
    assert sk2["evidence_state"] == "NOT_APPLICABLE_WITH_EVIDENCE" and sk2["measured_net"] is None
    (dry,) = _by(room2, arch2, "GF", "KITCHEN", "FIN.SKIRTING")      # a dry room is unaffected
    assert dry["measured_net"] == pytest.approx(18.4 - 0.9, abs=TOL)


def test_stair_flight_with_a_sourced_riser(with_facts):
    room, _ = with_facts
    (st,) = [ln for ln in room["lines"] if ln["item_code"] == "FIN.STAIR.FLIGHT"]
    assert st["evidence_state"] == "PROVISIONAL"
    assert st["measured_net"] == pytest.approx(math.hypot(9 * 0.30, 9 * 0.175), abs=1e-6)


def test_tiles_rect_and_polygon_rooms_use_the_layout_engines(with_facts):
    from engine.source import material_spec as MS
    from engine.source import tile_layout as TL
    room, arch = with_facts
    mats = {m["location"]["room_id"]: m for m in room["materials"] if m["method"].startswith("TILE")}
    c = dict(FACTS["tile_constraints"], wet_room=False)
    # a rectangle: exactly the S1 engine on the same dimensions
    k = _rooms(arch, "GF")["KITCHEN"]
    ref = TL.best({"length_m": k["rect"][0], "width_m": k["rect"][1]}, MS.tile("TILE-120x60"), c)["selected"]
    got = mats[k["room_id"]]
    assert got["method"] == TL.METHOD
    assert got["units"]["consumed_units"] == ref["units"]["consumed_units"]
    assert got["cutting_waste_m2"] == pytest.approx(ref["cutting_waste_m2"], abs=1e-9)
    # the L room with columns: the polygon engine, net = the BOQ line (waste is never inside the net)
    liv = _rooms(arch, "GF")["LIVING"]
    lm = mats[liv["room_id"]]
    assert lm["method"] == TL.POLY_METHOD and lm["state"] != "BLOCKED"
    (fl,) = _by(room, arch, "GF", "LIVING", "FIN.FLOOR.PORCELAIN")
    assert lm["net_installed_m2"] == pytest.approx(fl["measured_net"], abs=1e-6)
    assert lm["consumed_m2"] == pytest.approx(lm["units"]["consumed_units"] * 1.2 * 0.6, abs=1e-6)
    assert lm["consumed_m2"] >= lm["net_installed_m2"]


def test_owner_sizes_are_compared_but_the_selected_size_stays_the_basis(with_facts):
    room, arch = with_facts
    liv = _rooms(arch, "GF")["LIVING"]["room_id"]
    rows = [r for r in room["comparisons"] if r["room_id"] == liv]
    assert {r["spec_id"] for r in rows} == {"TILE-120x60", "TILE-60x60", "TILE-160x80"}
    sel = [r for r in rows if r["role"].startswith("SELECTED")]
    assert [r["spec_id"] for r in sel] == ["TILE-120x60"]
    assert all(r["recommendation"] == "" or r["recommendation"].startswith("RECOMMENDATION ONLY") for r in rows)


def test_membrane_of_a_wet_room_matches_the_rectangle_engine(with_facts):
    from engine.source import material_spec as MS
    from engine.source import membrane_layout as ML
    room, arch = with_facts
    b = _rooms(arch, "GF")["BATHROOM"]
    mem = next(m for m in room["materials"] if m["method"].startswith("MEMBRANE") and
               m["location"]["room_id"] == b["room_id"])
    ref = ML.plan({"length_m": b["rect"][0], "width_m": b["rect"][1]}, MS.membrane(),
                  constraints={"upturn_h_m": 0.20, "doorways": [{"door_id": "D", "width_m": 0.8, "depth_m": 0.2}],
                               "doorway_detail": "NOT_ESTABLISHED", "direction": "MIN_MATERIAL", "min_reusable_m": 0.3})
    for k in ("net_installed_m2", "overlap_m2", "consumed_m2", "cutting_waste_m2", "reusable_offcuts_m2"):
        assert mem[k] == pytest.approx(ref[k], abs=1e-9), k
    assert mem["net_installed_m2"] == pytest.approx(6.72 + 0.16 + 10.4 * 0.20, abs=1e-9)
    linked = {i for c in mem["components"] for i in c["line_ids"]}
    wp = {ln["line_id"] for ln in room["lines"] if ln["item_code"].startswith("FIN.WP.") and
          ln["location"]["room_id"] == b["room_id"]}
    assert linked == wp


# ---- metamorphic: moving, mirroring or turning the plan contents changes no quantity
@pytest.mark.parametrize("variant", [{"dx": 1234.5, "dy": -987.25}, {"mirror": True}, {"quarter_turns": 1},
                                     {"mirror": True, "quarter_turns": 3, "dx": 50.0}])
def test_metamorphic_variants_keep_every_room_quantity(base, tmp_path, variant):
    room2, arch2 = _run(tmp_path / "v.dxf", **variant)
    for floor in ("GF", "1F"):
        a, b = _rooms(base[1], floor), _rooms(arch2, floor)
        assert set(a) == set(b), (variant, floor)
        for code in a:
            for k in ("area_m2", "gross_m2", "perimeter_lm"):
                assert b[code][k] == pytest.approx(a[code][k], abs=1e-6), (variant, floor, code, k)
            for k in ("door", "passage", "window_threshold", "wall", "column"):
                assert b[code]["runs_lm"][k] == pytest.approx(a[code]["runs_lm"][k], abs=1e-6), (variant, code, k)
        oa = sorted((o["kind"], o["width_m"]) for o in base[1]["openings"] if o["floor"] == floor)
        ob = sorted((o["kind"], o["width_m"]) for o in arch2["openings"] if o["floor"] == floor)
        assert oa == ob, (variant, floor)
    assert Counter(ln["item_code"] for ln in room2["lines"]) == Counter(ln["item_code"] for ln in base[0]["lines"])


# ---- the CLI (the production command) and its workbook
def test_cli_requires_the_versioned_method(tmp_path):
    AS.villa(tmp_path / "v.dxf")
    with pytest.raises(SystemExit):
        M9.main(["arch", "--drawings", str(tmp_path / "v.dxf"), "--out", str(tmp_path / "o")])
    with pytest.raises(SystemExit):
        M9.main(["arch", "--drawings", str(tmp_path / "v.dxf"), "--out", str(tmp_path / "o"), "--method", "V0"])


def test_cli_writes_every_output_and_a_room_by_room_sheet(tmp_path):
    from openpyxl import load_workbook
    AS.villa(tmp_path / "v.dxf")
    facts = tmp_path / "facts.json"
    facts.write_text(json.dumps(FACTS), encoding="utf-8")
    out = tmp_path / "out"
    assert M9.main(["arch", "--drawings", str(tmp_path / "v.dxf"), "--out", str(out), "--method", METHOD,
                    "--facts", str(facts), "--reference-areas", "GF=116.56", "1F=75.2"]) == 0
    for f in ("BOQ_LINES.json", "BOQ_LINES.csv", "MATERIAL_RESULTS.json", "COVERAGE.json", "COVERAGE.csv",
              "RUN_RULES.json", "ARCH_PLAN.json", "ARCH_ROOMS.json", "ARCH_OPENINGS.json", "UNRESOLVED.json",
              "URBAN_ROOM_BOQ.xlsx", "OVERLAY_GF.svg", "OVERLAY_1F.svg"):
        assert (out / f).is_file(), f
    rr = json.loads((out / "RUN_RULES.json").read_text(encoding="utf-8"))
    assert rr["arch"]["method"] == METHOD
    plan = json.loads((out / "ARCH_PLAN.json").read_text(encoding="utf-8"))
    assert all(abs(r["reference"]["diff_m2"]) < 1e-6 for r in plan["regions"])
    wb = load_workbook(out / "URBAN_ROOM_BOQ.xlsx")
    ws = wb["ROOM BY ROOM BOQ"]
    head = [c.value for c in ws[3]][:14]
    assert head == ["FLOOR", "ROOM", "TRADE", "ITEM", "DIMENSIONS", "GROSS", "DEDUCTIONS", "NET QUANTITY", "UNIT",
                    "WASTE %", "WASTE QUANTITY", "PROCUREMENT", "SOURCE", "STATUS"]
    lines = {ln["line_id"]: ln for ln in json.loads((out / "BOQ_LINES.json").read_text(encoding="utf-8"))}
    col = {c.value: i for i, c in enumerate(ws[3])}
    seen_rooms, n_detail = set(), 0
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[col["ROW TYPE"]] != "DETAIL":
            continue
        n_detail += 1
        ln = lines[row[col["LINE ID"]]]
        seen_rooms.add(row[col["ROOM ID"]])
        net = row[col["NET QUANTITY"]]
        if isinstance(net, str) and net.startswith("="):            # NET = GROSS - DEDUCTIONS, as a formula
            assert row[col["GROSS"]] - row[col["DEDUCTIONS"]] == pytest.approx(ln["measured_net"], abs=1e-5)
        else:
            assert net == ln["measured_net"]
    assert n_detail == len(lines)                                    # every line once, room by room
    assert {r["room_id"] for r in json.loads((out / "ARCH_ROOMS.json").read_text(encoding="utf-8"))["rooms"]} <= \
        seen_rooms | {None}
    cov = {r["code"]: r for r in json.loads((out / "COVERAGE.json").read_text(encoding="utf-8"))["families"]}
    for code in ("B3-01", "B3-02", "B3-06", "B4-01", "B4-03", "B4-06", "B5-01", "B6-01", "B8-01", "B8-02", "B8-03",
                 "B8-05", "B7-01"):
        assert cov[code]["run_state"] == "QUANTIFIED", code


def test_the_workbook_module_does_not_import_engine_source():
    import ast
    tree = ast.parse((ROOT / "engine/boq_room_workbook.py").read_text(encoding="utf-8"))
    mods = [n.module for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)] + \
        [a.name for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names]
    assert mods and not [m for m in mods if m and m.startswith("engine.source")]
