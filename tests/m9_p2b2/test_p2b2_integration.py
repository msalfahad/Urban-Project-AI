"""Phase 2B-2 integration (ARCH_TAKEOFF_V2.4 through the CLI's arch_compute / main on the synthetic two-storey villa of
tools/arch_synth; synthetic geometry only).

    every V2.4 output is bound to the rules sheet it was computed with (embedded run record or OUTPUT_MANIFEST.json);
    the door jamb plaster uses the height of each door's category; V2.3 is not changed by V2.4 and does not follow the
    registry 1.1.0; a changed project rule moves the earlier outputs aside; V2.4 is deterministic, conserves its ledger
    and classifies every question; the before / after of the corrections (the doorway segment is no longer a wall upturn,
    the kitchen gets its membrane material, the main-entrance candidate's jamb is a range)."""

from __future__ import annotations

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from engine.source import boq_line as BL                      # noqa: E402
from engine.source import knowledge as K                      # noqa: E402
from engine.source import project_rules as PR                 # noqa: E402
from tools import arch_synth as AS                            # noqa: E402
from tools import urban_boq_m9 as M9                          # noqa: E402

V23, V24 = "ARCH_TAKEOFF_V2.3", "ARCH_TAKEOFF_V2.4"
V1_0_SHA256 = K._digest_version("1.0.0")


def _villa(tmp):
    p = Path(tmp) / "villa.dxf"
    if not p.exists():
        AS.villa(p)
    return p


def _run(dxf, method, **kw):
    try:
        room, arch = M9.arch_compute(str(dxf), method=method, **kw)[:2]
    except M9.CompletenessFailed as ex:                         # the run's outputs are still inspected
        room, arch = ex.room_boq, ex.arch
    return room, arch


def _rules(*rules):
    return {"schema": PR.SCHEMA, "project_id": "SYNTHETIC VILLA", "rules": list(rules)}


def _rule(rid, rule_id, parameter, value, scope, kind="DETAIL", ref="synthetic detail 1 / A-501"):
    return {"id": rid, "rule_id": rule_id, "parameter": parameter, "value": value, "scope": scope,
            "source": {"kind": kind, "ref": ref}}


def _total(room, code, state=None):
    return round(sum(ln["measured_net"] or 0 for ln in room["lines"] if ln["item_code"] == code and
                     (state is None or ln["evidence_state"] == state)), 6)


def _jambs(room):
    return {(ln["location"]["floor"], ln["location"]["room_id"].split("-OPENING-", 1)[-1]): ln for ln in room["lines"]
            if ln["item_code"] == "FIN.PLASTER" and ln.get("physical_element") == "door jambs and head"}


# ---------------------------------------------------------------------------------------------------------- the hash

def test_the_rules_sheet_hash_is_in_every_v2_4_output(tmp_path):
    """UP-RULES-001: the rules sheet exists before the BOQ and its hash travels with every output - embedded in each
    JSON document that carries the run record, and every file of the run (CSV, BOQ lines, workbook, overlays included)
    listed with its SHA-256 in OUTPUT_MANIFEST.json under the same rules hash."""
    import hashlib
    from openpyxl import load_workbook
    out = tmp_path / "out"
    assert M9.main(["arch", "--drawings", str(_villa(tmp_path)), "--out", str(out), "--method", V24]) == 0
    sheet = json.loads((out / "PROJECT_RULES_SHEET.json").read_text(encoding="utf-8"))["sheet"]
    h = sheet["rules_sha256"]
    man = json.loads((out / "OUTPUT_MANIFEST.json").read_text(encoding="utf-8"))
    assert man["rules_sheet"]["rules_sha256"] == h and man["rules_sheet"]["sheet_sha256"] == sheet["sheet_sha256"]
    assert man["knowledge"] == {"rule_set": "URBAN_KB_1.1.0", "digest": K._digest_version("1.1.0")}
    files = sorted(p.name for p in out.iterdir() if p.is_file() and p.name != "OUTPUT_MANIFEST.json")
    assert [f["name"] for f in man["files"]] == files
    for f in man["files"]:
        assert f["sha256"] == hashlib.sha256((out / f["name"]).read_bytes()).hexdigest(), f["name"]
    rr = json.loads((out / "RUN_RULES.json").read_text(encoding="utf-8"))
    assert rr["arch"]["v2_4"]["rules_sheet"]["rules_sha256"] == h
    for name in files:
        if name.endswith(".json") and name not in ("BOQ_LINES.json", "MATERIAL_RESULTS.json", "COVERAGE.json",
                                                   "DRAWING_DISCOVERY.json"):
            assert h in (out / name).read_text(encoding="utf-8"), f"{name} does not carry the rules hash"
    ws = load_workbook(out / "URBAN_ROOM_BOQ.xlsx", read_only=True)["SUMMARY"]
    cells = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert any(h in c for c in cells) and sheet["sheet_sha256"] in cells
    # the sheet precedes every quantity: its rule rows depend only on the inputs (no usage in the hash)
    assert {r["section"] for r in sheet["rows"]} <= set(PR.SECTIONS)
    assert sheet["invalidation"]["state"] == "FIRST_RUN"


def test_a_changed_project_rule_moves_the_earlier_outputs_aside(tmp_path):
    """the same OUT: a first run on the Urban standards, a second with a project upturn for bathrooms (the earlier
    outputs move to SUPERSEDED_1 and the change names the upturn item), a third with the same rules (UNCHANGED,
    nothing moved)."""
    dxf, out = _villa(tmp_path), tmp_path / "out"
    assert M9.main(["arch", "--drawings", str(dxf), "--out", str(out), "--method", V24]) == 0
    first = json.loads((out / "PROJECT_RULES_SHEET.json").read_text(encoding="utf-8"))["sheet"]["rules_sha256"]
    pr = tmp_path / "rules.json"
    pr.write_text(json.dumps(_rules(_rule("PR-1", "UP-WP-003", "upturn_h_m", 0.30,
                                          {"kind": "LOCATION_TYPE", "types": ["BATHROOM"]}))), encoding="utf-8")
    assert M9.main(["arch", "--drawings", str(dxf), "--out", str(out), "--method", V24,
                    "--project-rules", str(pr)]) == 0
    sheet = json.loads((out / "PROJECT_RULES_SHEET.json").read_text(encoding="utf-8"))["sheet"]
    inv = sheet["invalidation"]
    assert inv["state"] == "RULES_CHANGED" and inv["previous_rules_sha256"] == first
    assert inv["moved"]["moved_to"] == "SUPERSEDED_1" and "BOQ_LINES.json" in inv["moved"]["files"]
    assert "FIN.WP.WALL_UPTURN" in inv["affected_items"] and "WATERPROOFING" in inv["affected_trades"]
    old = json.loads((out / "SUPERSEDED_1" / "PROJECT_RULES_SHEET.json").read_text(encoding="utf-8"))["sheet"]
    assert old["rules_sha256"] == first
    lines = json.loads((out / "BOQ_LINES.json").read_text(encoding="utf-8"))
    bath = [ln for ln in lines if ln["item_code"] == "FIN.WP.WALL_UPTURN" and
            ln["calculation"]["inputs"].get("upturn_h_m") == 0.30]
    assert len(bath) == 2                                       # the two bathrooms; the kitchen keeps 0.20
    assert M9.main(["arch", "--drawings", str(dxf), "--out", str(out), "--method", V24,
                    "--project-rules", str(pr)]) == 0
    again = json.loads((out / "PROJECT_RULES_SHEET.json").read_text(encoding="utf-8"))["sheet"]
    assert again["invalidation"]["state"] == "UNCHANGED" and not (out / "SUPERSEDED_2").exists()


# ----------------------------------------------------------------------------------------------- jamb plaster height

def test_v2_4_jamb_plaster_uses_the_category_height(tmp_path):
    """UP-DOOR-HGT-001: each door's jamb / head plaster at its category height - internal and bathroom PVC 2.20 m
    (OWNER_STANDARD_ASSUMED), a main-entrance candidate a range 2.20 .. 2.50 m (never one assumed height), a project
    height for one category governs only that category."""
    dxf = _villa(tmp_path)
    room, arch = _run(dxf, V24)
    j = _jambs(room)
    cats = Counter(ln["calculation"]["inputs"]["category"] for ln in j.values())
    assert cats == {"INTERNAL_PVC": 3, "BATHROOM_PVC": 2, "MAIN_ENTRANCE_CANDIDATE": 1}
    for ln in j.values():
        i = ln["calculation"]["inputs"]
        t, w = i["wall_t_m"], i["width_m"]
        if i["category"] == "MAIN_ENTRANCE_CANDIDATE":
            assert ln["measured_net"] is None and ln["evidence_state"] == "BOUNDED_EXPOSURE"
            assert i["height_state"] == "CLASS_PENDING"
            assert ln["rfi"]["question"].startswith("JAMB_HEIGHT_CLASS_PENDING")
            lo, hi = 2 * t * 2.20 + t * w, 2 * t * 2.50 + t * w
            assert json.dumps(ln).count(str(round(lo, 6))) >= 1 and json.dumps(ln).count(str(round(hi, 6))) >= 1
        else:
            assert i["height_m"] == 2.20 and i["height_state"] == "OWNER_STANDARD_ASSUMED"
            assert ln["measured_net"] == pytest.approx(2 * t * 2.20 + t * w)
    # a project height for internal PVC doors only
    room2, _ = _run(dxf, V24, project_rules=_rules(_rule("PR-1", "UP-DOOR-HGT-001", "door_height_m", 2.40,
                                                         {"kind": "DOOR_CATEGORY", "categories": ["INTERNAL_PVC"]},
                                                         kind="SCHEDULE", ref="door schedule A-601")))
    for k, ln in _jambs(room2).items():
        i = ln["calculation"]["inputs"]
        if i["category"] == "INTERNAL_PVC":
            assert i["height_m"] == 2.40 and i["height_state"] == "PROJECT_SPECIFIED"
            assert ln["measured_net"] == pytest.approx(2 * i["wall_t_m"] * 2.40 + i["wall_t_m"] * i["width_m"])
        elif i["category"] == "BATHROOM_PVC":
            assert i["height_m"] == 2.20 and ln["measured_net"] == j[k]["measured_net"]
    # before: V2.3 took every jamb at 2.20 m, the main-entrance candidate included
    j23 = _jambs(_run(dxf, V23)[0])
    assert all(ln["calculation"]["inputs"].get("height_m") == 2.20 for ln in j23.values())
    main = next(k for k, ln in j.items() if ln["calculation"]["inputs"]["category"] == "MAIN_ENTRANCE_CANDIDATE")
    assert j23[main]["measured_net"] is not None


# ------------------------------------------------------------------------------------------------------- frozen V2.3

def test_v2_3_is_unchanged_by_v2_4_and_does_not_follow_the_new_registry(tmp_path, monkeypatch):
    """V2.3 reads the registry 1.0.0 (its digest unchanged) and none of the V2.4 items; a modified registry 1.1.0
    (upturn 0.25 m) changes the V2.4 upturn and leaves V2.3 byte-identical."""
    dxf = _villa(tmp_path)
    a23 = _run(dxf, V23)
    assert a23[1]["run"]["policies"]["knowledge"]["digest"] == V1_0_SHA256
    assert "v2_4" not in a23[1]["run"] and "rules_sheet" not in a23[1]
    codes = {ln["item_code"] for ln in a23[0]["lines"]}
    assert not codes & {"FIN.WP.WALL_UPTURN", "FIN.WP.DOORWAY_TRANSITION", "FIN.WP.TERMINATION", "FIN.WP.PENETRATION"}
    a24 = _run(dxf, V24)
    up24 = _total(a24[0], "FIN.WP.WALL_UPTURN")
    reg = json.loads(K.VERSIONS["1.1.0"].read_text(encoding="utf-8"))
    for r in reg["rules"]:
        if r["rule_id"] == "UP-WP-003":
            r["logic"]["upturn_h_m"] = 0.25
    alt = tmp_path / "URBAN_ENGINEERING_KNOWLEDGE_1.1.0.json"
    alt.write_text(json.dumps(reg, indent=1), encoding="utf-8")
    monkeypatch.setitem(K.VERSIONS, "1.1.0", alt)
    K._load_version.cache_clear()
    K._digest_version.cache_clear()
    try:
        b23, b24 = _run(dxf, V23), _run(dxf, V24)
    finally:
        monkeypatch.undo()
        K._load_version.cache_clear()
        K._digest_version.cache_clear()
    assert BL.dumps(b23[0]["lines"]) == BL.dumps(a23[0]["lines"])
    assert b23[1]["run"] == a23[1]["run"]
    assert _total(b24[0], "FIN.WP.WALL_UPTURN") == pytest.approx(up24 * 0.25 / 0.20)
    assert b24[1]["rules_sheet"]["rules_sha256"] != a24[1]["rules_sheet"]["rules_sha256"]


def test_v2_4_is_deterministic_and_conserves_its_ledger(tmp_path):
    """two runs give the same lines, materials, rules hash and door feasibility; the ledger conserves with no
    CONFIRMED_MISSING record; every question is classified and none is A (established)."""
    dxf = _villa(tmp_path)
    a, b = _run(dxf, V24), _run(dxf, V24)
    assert BL.dumps(a[0]["lines"]) == BL.dumps(b[0]["lines"])
    assert json.dumps(a[0]["materials"], sort_keys=True) == json.dumps(b[0]["materials"], sort_keys=True)
    assert a[1]["rules_sheet"]["rules_sha256"] == b[1]["rules_sheet"]["rules_sha256"]
    assert a[1]["door_feasibility"]["rows"] == b[1]["door_feasibility"]["rows"]
    led = a[1]["ledger"]
    assert led["summary"]["conservation"]["ok"] and led["summary"]["conservation"]["not_in_any_record"] == []
    assert led["summary"]["by_status"]["CONFIRMED_MISSING"] == 0
    assert led["table"]["wp_system"].startswith("ELIGIBILITY_V0_3_WP_SYSTEM")
    assert led["table"]["owner_scope"].endswith("(knowledge registry 1.1.0 rules)")
    q = a[1]["owner_questions"]["summary"]
    assert q["unclassified"] == 0 and q["instances"]["A_ESTABLISHED"] == 0 and q["established_but_still_raised"] == 0
    for ln in a[0]["lines"]:
        BL.validate(ln)                                         # raises LineError on a malformed line


# ------------------------------------------------------------------------------------------------- before / after

def test_before_after_waterproofing_and_doors_on_the_villa(tmp_path):
    """V2.3 -> V2.4 on the villa: the wall upturn excludes the doorway segments (7.84 -> 7.34 m2, the three doorways
    become transitions, unresolved without a detail), the kitchen gets its membrane material and layout, the floor
    membrane and the doorway continuation are unchanged (no double count), the roll count is a cutting plan per
    membrane system and stays PROVISIONAL without a specified system; no proposed modification is an as-drawn line."""
    dxf = _villa(tmp_path)
    r23, a23 = _run(dxf, V23)
    r24, a24 = _run(dxf, V24)
    assert _total(r23, "FIN.WP.UPTURN") == pytest.approx(7.84) and _total(r24, "FIN.WP.UPTURN") == 0
    assert _total(r24, "FIN.WP.WALL_UPTURN") == pytest.approx(7.34)
    assert _total(r24, "FIN.WP.FLOOR") == _total(r23, "FIN.WP.FLOOR") == pytest.approx(31.36)
    assert _total(r24, "FIN.WP.DOORWAY") == _total(r23, "FIN.WP.DOORWAY")
    tr = [ln for ln in r24["lines"] if ln["item_code"] == "FIN.WP.DOORWAY_TRANSITION"]
    assert len(tr) == 3 and all(ln["evidence_state"] == "BOUNDED_EXPOSURE" and ln["measured_net"] is None for ln in tr)
    mem23 = {(m["location"] or {}).get("room_id") for m in r23["materials"] if "MEMBRANE" in (m.get("method") or "")}
    mem24 = {(m["location"] or {}).get("room_id") for m in r24["materials"] if "MEMBRANE" in (m.get("method") or "")}
    kitchen = next(r["room_id"] for r in a24["rooms"] if "KITCHEN" in (r.get("codes") or ()))
    assert kitchen not in mem23 and kitchen in mem24 and mem23 < mem24
    (sysrow,) = a24["waterproofing"]["systems"]
    assert sysrow["system_id"] == "URBAN_DEFAULT_WET_AREA" and sysrow["roll_count_state"] == "PROVISIONAL"
    assert sysrow["rolls"] == 6 and sysrow["rolls_if_each_area_alone"] == 7
    assert sysrow["net_m2"] == pytest.approx(31.36)
    assert sysrow["net_m2"] < sysrow["installed_m2"] <= sysrow["consumed_m2"] <= sysrow["rolls"] * 10.0
    for ln in r24["lines"]:
        assert "PROPOSED" not in json.dumps(ln.get("calculation") or {}), ln["line_id"]
    df = a24["door_feasibility"]
    assert Counter(r["feasibility"] for r in df["rows"]) == {"ENLARGEMENT_GEOMETRICALLY_POSSIBLE": 3,
                                                              "ENLARGEMENT_REQUIRES_ENGINEERING_REVIEW": 2,
                                                              "INSUFFICIENT_DRAWING_EVIDENCE": 1}
    bs = a24["boq_scenarios"]
    assert {"AS_DRAWN", "OWNER_STANDARD", "PROPOSED_MODIFICATION", "WATERPROOFING_PROCUREMENT"} <= set(bs)
    assert "PROPOSED_MODIFICATION" not in json.dumps(bs["AS_DRAWN"])


def test_the_engineering_summary_answers_the_eight_questions(tmp_path):
    room, arch = _run(_villa(tmp_path), V24)
    s = arch["engineering_summary"]
    assert list(s)[2:] == ["what_the_drawing_shows", "what_urban_requires", "where_they_differ",
                           "can_the_modification_fit", "measured_assumed_proposed_unresolved", "material_needed",
                           "governing_project_rules", "approvals_needed"]
    assert s["rules_sha256"] == arch["rules_sheet"]["rules_sha256"]
    assert s["what_urban_requires"]["waterproofing"] == {"upturn_h_m": 0.20, "lap_m": 0.10, "roll_m": [1.0, 10.0]}
    assert s["where_they_differ"]["doorway_transitions_unresolved"] == 3
    assert s["approvals_needed"]["items"] == len(arch["approvals"]["items"]) > 0
    wd = defaultdict(int)
    for a in arch["approvals"]["items"]:
        wd[a["subject"]] += 1
    assert wd["DOOR"] >= 1 and wd["WATERPROOFING"] == 1
