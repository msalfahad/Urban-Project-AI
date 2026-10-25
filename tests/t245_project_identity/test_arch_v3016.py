"""T-245: ARCH_TAKEOFF_V3.0.16 (engine.source.arch_v3016): ONE project id and ONE revision set over the architectural
and structural routes of the candidate; the export (tools/boq_export_v1.py) accepts it, binds verdicts to the canonical
line content and states every drawing's revision state.

Hand-made synthetic villa N245 (not the development villa), written before running the code:
  architectural route: project id "villa-n245", plan file A (sha256 'a' x 64) -> route set sha256("aaa..a").
  structural route:    project id "N245 structural set", file S (sha 'b' x 64) -> route set sha256("bbb..b").
  lines: GF floor tiles R1 4.00 x 5.00 = 20.000 m2 (arch); footing F1 1.50 x 1.20 x 0.60 = 1.080 m3 (str);
         slab S1 BLOCKED_RFI (held, no quantity) (str).
  V3.0.16: every line project = {"villa-n245", sha256("aaa..a|bbb..b"), route {...}}; line ids, order, quantities,
  states unchanged (20.000 / 1.080 / held blank); the local file name of S ("S local copy.dxf") becomes the declared
  title "N245-ST.dxf".
  Normal: two runs whose local copy of S has different names give byte-identical V3.0.16 json (the V3.0.15-level json
  differs: the cause). The export runs, VERIFIED 0, structural revision UNCONFIRMED stated. A WHOLE footing verdict made
  on run 1 is COUNTED on run 2.
  Adverse: a line of a third project, two revisions of one sheet, a document no route uses, a route whose files do
  not give its set, a line on an undeclared file, an UNCONFIRMED revision without its reason, "new project", the
  unrestamped two-route lines in the export (probe 10 stays green), an identity record of another set; a one-line
  quantity change (F1 1.080 -> 1.100) leaves the run-1 verdict STALE.
"""

from __future__ import annotations

import copy
import hashlib
import json

import pytest

from engine.source import arch_v3016 as A
from engine.source import boq_line as BL
from tools import boq_export_v1 as X

PA, PS = "a" * 64, "b" * 64
PID, SPID = "villa-n245", "N245 structural set"
CVER = A.METHOD


def rset(*shas):
    return hashlib.sha256("|".join(sorted(shas)).encode("utf-8")).hexdigest()


def lines(local_s="S local copy.dxf", f1=1.08):
    ra, rs = rset(PA), rset(PS)
    da, ds = [{"file": "N245-AR.dxf", "sha256": PA}], [{"file": local_s, "sha256": PS}]
    return [
        BL.make_line(project_id=PID, revision_set_sha256=ra, item_code="FIN.FLOOR.PORCELAIN", unit="m2",
                     evidence_state="MEASURED", measured_net=20.0, element_id="R1", floor="GF", formula="L*W",
                     inputs={"L": 4.0, "W": 5.0}, rule_id="R-SYN", rule_version="1", documents=da),
        BL.make_line(project_id=SPID, revision_set_sha256=rs, item_code="STR.CONC.FOOTING", unit="m3",
                     evidence_state="MEASURED", measured_net=f1, element_id="F1", floor="FDN", formula="L*B*H",
                     inputs={"L": 1.5, "B": 1.2, "H": 0.6}, rule_id="R-SYN", rule_version="1", documents=ds),
        BL.make_line(project_id=SPID, revision_set_sha256=rs, item_code="STR.CONC.SLAB", unit="m3",
                     evidence_state="BLOCKED_RFI", element_id="S1", floor="GF", formula="A*t",
                     rfi={"question": "slab depth not printed"}, rule_id="R-SYN", rule_version="1", documents=ds)]


def inputs(**over):
    d = {"schema": A.INPUTS_SCHEMA,
         "project": {"project_id": PID, "evidence": "one project folder holds both sets; same drawing number"},
         "documents": [
             {"discipline": "ARCHITECTURAL", "title": "N245-AR.dxf", "sheet": "AR all", "revision": "B",
              "revision_state": "PRINTED", "sha256": PA},
             {"discipline": "STRUCTURAL", "title": "N245-ST.dxf", "sheet": "ST all", "revision": "UNCONFIRMED",
              "revision_state": "UNCONFIRMED", "reason": "two issue dates on the sheets", "sha256": PS}],
         "routes": [{"route_project_id": PID, "documents": [PA], "evidence": "architectural route"},
                    {"route_project_id": SPID, "documents": [PS], "evidence": "structural route"}]}
    d.update(over)
    return d


def v16(ls=None, inp=None):
    return A.integrate(lines() if ls is None else ls, A.load_inputs(inputs() if inp is None else inp))


# ---- normal ---------------------------------------------------------------------------------------------------------

def test_one_identity_over_both_routes_and_nothing_else_moves():
    l15 = lines()
    out, ident, local = v16(l15)
    want = rset(PA, PS)
    assert {(x["project"]["project_id"], x["project"]["revision_set_sha256"]) for x in out} == {(PID, want)}
    assert [x["line_id"] for x in out] == [x["line_id"] for x in l15]           # no line id churn
    assert [x["measured_net"] for x in out] == [20.0, 1.08, None]                # held stays blank, never 0
    assert [x["evidence_state"] for x in out] == ["MEASURED", "MEASURED", "BLOCKED_RFI"]
    assert out[1]["project"]["route"] == {"project_id": SPID, "revision_set_sha256": rset(PS)}
    assert out[1]["source"]["documents"] == [{"file": "N245-ST.dxf", "sha256": PS}]
    assert local == {"S local copy.dxf": 2} and ident["lines_by_route_project"] == {PID: 1, SPID: 2}
    for a, b in zip(l15, out):
        a2, b2 = copy.deepcopy(a), copy.deepcopy(b)
        for z in (a2, b2):
            z.pop("project")
            z["source"]["documents"] = [{"sha256": d["sha256"]} for d in z["source"]["documents"]]
        assert a2 == b2


def test_local_file_name_no_longer_changes_the_json_the_named_cause():
    a15, b15 = lines("S local copy.dxf"), lines("S_local_copy.dxf")
    assert BL.dumps(a15) != BL.dumps(b15)                                         # the cause, before V3.0.16
    oa, ob = v16(a15)[0], v16(b15)[0]
    assert BL.dumps(oa) == BL.dumps(ob) and X.lines_sha256(oa) == X.lines_sha256(ob)
    assert BL.dumps(oa) == BL.dumps(v16(a15)[0])                                   # repeated run, same bytes


def test_write_files_identity_record_and_diff(tmp_path):
    names, ident = A.write(tmp_path, lines(), A.load_inputs(inputs()))
    assert names == sorted(["V3_0_16_BOQ_LINES.json", "V3_0_16_BOQ_LINES.csv", "V3_0_16_IDENTITY.json",
                            "V3_0_16_DIFF.csv", "V3_0_16_RUN.json"])
    rec = json.loads((tmp_path / "V3_0_16_IDENTITY.json").read_text(encoding="utf-8"))
    assert rec["schema"] == A.IDENTITY_SCHEMA and rec["revision_set_sha256"] == rset(PA, PS)
    assert rec["disclosures"] == ["STRUCTURAL ST all: revision UNCONFIRMED (two issue dates on the sheets)"]
    assert ident["diff"]["vs_v3_0_15"] == {"IDENTITY_ONLY": 3} and ident["totals_moved_by_code"] == {}
    run = json.loads((tmp_path / "V3_0_16_RUN.json").read_text(encoding="utf-8"))
    assert run["released_default"] is False and "S local copy.dxf" in json.dumps(run["run_metadata"])
    assert "S local copy.dxf" not in (tmp_path / "V3_0_16_BOQ_LINES.json").read_text(encoding="utf-8")


def test_export_accepts_the_v3016_candidate_and_states_the_revision(tmp_path):
    out = tmp_path / "c"
    out.mkdir()
    A.write(out, lines(), A.load_inputs(inputs()))
    lp = out / "V3_0_16_BOQ_LINES.json"
    assert X.main(["--lines", str(lp), "--candidate-version", CVER, "--out", str(tmp_path / "x"), "--code-commit",
                   "abc", "--identity", str(out / "V3_0_16_IDENTITY.json")]) == 0
    m = json.loads((tmp_path / "x" / "EXPORT_MANIFEST.json").read_text(encoding="utf-8"))
    assert m["identity"]["project_id"] == PID and m["identity"]["revision_set_sha256"] == rset(PA, PS)
    assert m["identity"]["boq_lines_sha256"] == X.lines_sha256(json.loads(lp.read_text(encoding="utf-8")))
    assert m["row_counts"]["VERIFIED"] == 0 and m["row_counts"]["total"] == 3
    assert m["project_identity"]["disclosures"] == [
        "STRUCTURAL ST all: revision UNCONFIRMED (two issue dates on the sheets)"]
    assert m["project_identity"]["review_label"].startswith("INTERNAL REVIEW CANDIDATE")
    rows = X.read_csv(tmp_path / "x" / "BOQ.csv")
    held = [r for r in rows if r["element_id"] == "S1"][0]
    assert held["qty_measured_unverified"] == "" and held["net"] == ""             # HELD blank, never 0
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "x" / "BOQ.xlsx")["Manifest"]
    txt = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "revision UNCONFIRMED (two issue dates on the sheets)" in txt and "INTERNAL REVIEW CANDIDATE" in txt


def test_verdict_rebinds_on_a_rerun_and_a_quantity_change_is_stale(tmp_path):
    r1 = v16(lines("S local copy.dxf"))[0]
    ids = [x["line_id"] for x in r1 if x["item_code"] == "STR.CONC.FOOTING"]
    rec = X.verdict_add(r1, candidate_version=CVER, boq_lines_sha256=X.lines_sha256(r1), line_ids=ids, scope="WHOLE",
                        checker_qty=1.08, evidence=[{"sheet": "ST footing schedule", "printed_value": "F1 1.5x1.2x0.6"}],
                        run_token="20261025T050000-1-1", routine="Verifier", verdict_commit="c0ffee12",
                        time="2026-10-25T05:00:00Z")
    r2 = v16(lines("S_local_copy.dxf"))[0]                                         # the re-run, another local name
    m2 = X.export(r2, tmp_path / "r2", candidate_version=CVER, boq_lines_sha256=X.lines_sha256(r2), v2_records=[rec])
    st2 = {v["verdict_id"]: v["status"] for v in m2["v2_verdicts"]}[rec["verdict_id"]]
    assert st2.startswith("COUNTED") and "STALE" not in st2                       # bound, not stale (no completeness
    #                                                                               record here: not VERIFIED, by design)
    r3 = v16(lines("S local copy.dxf", f1=1.10))[0]                               # one genuine quantity change
    m3 = X.export(r3, tmp_path / "r3", candidate_version=CVER, boq_lines_sha256=X.lines_sha256(r3), v2_records=[rec])
    assert {v["verdict_id"]: v["status"] for v in m3["v2_verdicts"]}[rec["verdict_id"]].startswith("STALE")


# ---- adverse --------------------------------------------------------------------------------------------------------

def test_a_line_of_another_project_is_refused():
    ls = lines()
    ls.append(BL.make_line(project_id="other-villa", revision_set_sha256=rset(PA), item_code="FIN.FLOOR.PORCELAIN",
                           unit="m2", evidence_state="MEASURED", measured_net=5.0, element_id="R9", floor="GF",
                           formula="L*W", inputs={"L": 1.0, "W": 5.0}, rule_id="R-SYN", rule_version="1"))
    with pytest.raises(ValueError, match="no declared route"):
        v16(ls)


def test_a_drawing_changed_since_the_identity_is_refused():
    """the structural file is replaced by another revision (sha 'c'): its route set is no declared route."""
    ls = lines()
    for x in ls[1:]:
        x["project"]["revision_set_sha256"] = rset("c" * 64)
    with pytest.raises(ValueError, match="no declared route"):
        v16(ls)


def test_two_revisions_of_one_sheet_are_refused():
    d = inputs()["documents"]
    d.append(dict(d[1], sha256="c" * 64, revision="C", revision_state="PRINTED"))
    with pytest.raises(A.InputsError, match="two revisions of one sheet"):
        A.load_inputs(inputs(documents=d, routes=[inputs()["routes"][0],
                                                  dict(inputs()["routes"][1], documents=[PS, "c" * 64])]))


@pytest.mark.parametrize("change,needle", [
    (lambda d: d["documents"].append({"discipline": "SANITARY", "title": "x.pdf", "sheet": "SN", "revision": "A",
                                      "revision_state": "PRINTED", "sha256": "d" * 64}), "no route uses"),
    (lambda d: d["routes"][1].update(documents=["e" * 64]), "outside the revision set"),
    (lambda d: d["documents"][1].pop("reason"), "names why"),
    (lambda d: d["project"].update(project_id="new project"), "project id is missing"),
    (lambda d: d["project"].pop("evidence"), "source evidence"),
    (lambda d: d["documents"][0].update(sha256="A" * 64), "64 lower-case hex"),
    (lambda d: d["documents"][1].update(sha256=PA), "declared twice"),
    (lambda d: d["documents"][0].update(discipline="CIVIL"), "discipline"),
    (lambda d: d.update(schema="X"), "schema"),
])
def test_identity_inputs_refused(change, needle):
    d = inputs()
    change(d)
    with pytest.raises(A.InputsError, match=needle):
        A.load_inputs(d)


def test_a_line_on_an_undeclared_file_is_refused():
    ls = lines()
    ls[0]["source"]["documents"].append({"file": "extra.dxf", "sha256": "f" * 64})
    with pytest.raises(ValueError, match="outside the revision set"):
        v16(ls)


def test_export_still_refuses_two_projects_probe_10(tmp_path):
    ls = lines()
    with pytest.raises(X.ExportError, match="project ids"):
        X.export(ls, tmp_path / "p", candidate_version=CVER, boq_lines_sha256=X.lines_sha256(ls))


def test_export_refuses_an_identity_record_of_another_set(tmp_path):
    out, _i, _l = v16()
    rec = A.identity_record(A.load_inputs(inputs()))
    bad = copy.deepcopy(rec)
    bad["documents"][1]["sha256"] = "c" * 64
    with pytest.raises(X.ExportError, match="revision set"):
        X.export(out, tmp_path / "i", candidate_version=CVER, boq_lines_sha256=X.lines_sha256(out), identity_doc=bad)
    two = copy.deepcopy(rec)
    two["documents"].append(dict(two["documents"][1], sha256="c" * 64))
    with pytest.raises(X.ExportError, match="two revisions of one sheet"):
        X.export(out, tmp_path / "j", candidate_version=CVER, boq_lines_sha256=X.lines_sha256(out), identity_doc=two)
    other = dict(rec, project_id="other")
    with pytest.raises(X.ExportError, match="project"):
        X.export(out, tmp_path / "k", candidate_version=CVER, boq_lines_sha256=X.lines_sha256(out), identity_doc=other)
