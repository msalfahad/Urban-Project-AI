"""BOQ XLSX EXPORT V2 (Qortuba RC1 owner-review workbook) - a VIEW. It calculates NOTHING.

The caller hands a workbook model: {sheet name: {"role": "ADDITIVE_SUMMARY" | "BREAKDOWN" | "SCHEDULE" | "INFO",
"header": [...], "rows": [[...]], "qty_cols": [...], "status_col": n | None, "wrap_cols": [...]}}. Exactly ONE sheet
may be ADDITIVE_SUMMARY; every other sheet states on row 2 that it must not be added to the summary. Every quantity
cell is a plain number copied from an engine register row (no formula, no sum, no correction factor, no rounding of
the stored value - only the display format rounds). Lives outside engine/source because it uses openpyxl
(engine/source is stdlib-only).

validate() reads the saved file back and refuses: a missing sheet, any cell that differs from the model, any formula,
a quantity cell that differs from the register value it was copied from (cell sources), a status outside the
vocabulary, an approved line without RELEASED, a legacy alias code as a summary line, a canonical item that is not
exactly once in the summary, a row-count drop on any sheet.

The saved zip is re-packed with fixed entry timestamps, so the same model gives the same bytes.

Parameters (Alsenan Phase A; defaults reproduce the Qortuba RC1 workbook byte for byte): `banner` (row 1 of every
sheet - a project's name belongs to the caller, not to this module), `summary_name` (the additive sheet named in the
breakdown / schedule role lines) and SCOPED additive summaries: more than one ADDITIVE_SUMMARY sheet is accepted only
when every one declares a distinct non-empty "scope" (e.g. ARCHITECTURAL / STRUCTURAL); each line is then added within
its own scope only, and the role line says so.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile

POLICY_ID = "BOQ_XLSX_EXPORT_V2"
BANNER = ("SHADOW / RC1_REFERENCE - QORTUBA ARCHITECTURAL QTO - NOT APPROVED FOR TENDER OR CONTRACT - NO PRICES - "
          "VIEW OVER ENGINE REGISTERS - NO PRODUCTION MIGRATION")
ROLE_TEXT = {"ADDITIVE_SUMMARY": "ADDITIVE SUMMARY: one line per canonical BOQ item. The ONLY sheet whose lines may "
                                 "be added (and a MEASURE PAIR is one item measured twice - price one basis).",
             "BREAKDOWN": "BREAKDOWN - NOT ADDITIVE: these lines repeat summary quantities by room / element. "
                          "Never add them to 01_BOQ_SUMMARY.",
             "SCHEDULE": "SCHEDULE - NOT ADDITIVE: opening records and counts behind summary lines.",
             "INFO": "INFORMATION - no quantities to add."}
STATUSES = ("COMPUTED_SHADOW_COMPLETE", "AUTHORISED_SUBTOTAL", "BLOCKED", "RELEASED", "NOT_APPLICABLE",
            "IN_SCOPE", "OUT_OF_MEASURED_SCOPE", "PASS", "FAIL", "OPEN", "INFO")
FILL = {"COMPUTED_SHADOW_COMPLETE": "FFF2CC", "AUTHORISED_SUBTOTAL": "F8CBAD", "BLOCKED": "BFBFBF",
        "RELEASED": "C6EFCE", "OUT_OF_MEASURED_SCOPE": "D9D9D9", "FAIL": "FF9999", "OPEN": "F8CBAD"}
QTY_FORMAT = "#,##0.000"
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
HEADER_ROW, FIRST_DATA_ROW = 3, 4


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, (list, tuple, dict)):
        return json.dumps(v, ensure_ascii=False, sort_keys=True, default=str)
    return v


def _trim(row):
    r = ["" if v is None else v for v in row]
    while r and r[-1] == "":
        r.pop()
    return r


def content(model: dict) -> dict:
    return {name: {"role": s["role"], "header": [_norm(h) for h in s["header"]],
                   "rows": [[_norm(v) for v in r] for r in s["rows"]]} for name, s in model.items()}


def content_digest(model: dict) -> str:
    return hashlib.sha256(json.dumps(content(model), ensure_ascii=False, sort_keys=True,
                                     default=str).encode()).hexdigest()


def _repack(path):
    """Fixed zip entry times, and the core-properties 'modified' stamp (openpyxl writes the save time) pinned to
    'created', so the same model always gives the same bytes."""
    with open(path, "rb") as f:
        src = zipfile.ZipFile(io.BytesIO(f.read()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                txt = data.decode("utf-8")
                c = re.search(r"<dcterms:created[^>]*>([^<]*)</dcterms:created>", txt)
                if c:
                    txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                                 lambda m: m.group(1) + c.group(1) + m.group(2), txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, data)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


def role_text(sheet: dict, summary_name: str = "01_BOQ_SUMMARY") -> str:
    t = ROLE_TEXT[sheet["role"]].replace("01_BOQ_SUMMARY", summary_name)
    if sheet["role"] == "ADDITIVE_SUMMARY" and sheet.get("scope"):
        t += f" SCOPE: {sheet['scope']} - add these lines only to each other, never to another scope."
    return t


def check_additive(model: dict) -> None:
    add = [s for s in model.values() if s["role"] == "ADDITIVE_SUMMARY"]
    if len(add) == 1:
        return
    scopes = [s.get("scope") for s in add]
    if not add or not all(scopes) or len(set(scopes)) != len(scopes):
        raise ValueError("exactly one ADDITIVE_SUMMARY sheet is required, or several with distinct declared scopes")


def write(model: dict, path, *, created, banner: str = BANNER, summary_name: str = "01_BOQ_SUMMARY") -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    check_additive(model)
    c = content(model)
    wb = Workbook()
    wb.remove(wb.active)
    for name, s in model.items():
        ws = wb.create_sheet(name)
        ws.append([banner])
        ws.append([role_text(s, summary_name)])
        ws.append(c[name]["header"])
        for r in c[name]["rows"]:
            ws.append(r)
        ws["A1"].font = Font(bold=True, color="C00000")
        ws["A2"].font = Font(bold=True, color="1F4E79" if s["role"] == "ADDITIVE_SUMMARY" else "7F6000")
        for cell in ws[HEADER_ROW]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ncol = max(len(c[name]["header"]), 1)
        last = HEADER_ROW + len(c[name]["rows"])
        for q in s.get("qty_cols") or ():
            for i in range(FIRST_DATA_ROW, last + 1):
                ws.cell(i, q + 1).number_format = QTY_FORMAT
        sc = s.get("status_col")
        for i in range(FIRST_DATA_ROW, last + 1):
            st = ws.cell(i, sc + 1).value if sc is not None else None
            if st in FILL:
                ws.cell(i, sc + 1).fill = PatternFill("solid", fgColor=FILL[st])
                if st in ("BLOCKED", "OUT_OF_MEASURED_SCOPE"):
                    for j in range(1, ncol + 1):
                        ws.cell(i, j).font = Font(italic=True, color="595959")
            for w in s.get("wrap_cols") or ():
                ws.cell(i, w + 1).alignment = Alignment(wrap_text=True, vertical="top")
        widths = {int(k): v for k, v in (s.get("widths") or {}).items()}
        for j in range(ncol):
            ws.column_dimensions[get_column_letter(j + 1)].width = widths.get(j, 16 if j else 22)
        ws.freeze_panes = f"B{FIRST_DATA_ROW}" if s["role"] != "INFO" else f"A{FIRST_DATA_ROW}"
        if c[name]["rows"] and s["role"] != "INFO":
            ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(ncol)}{last}"
    wb.properties.creator = "Urban QTO (BOQ_XLSX_EXPORT_V2)"
    wb.properties.created = wb.properties.modified = created
    wb.save(str(path))
    _repack(path)
    with open(path, "rb") as f:
        data = f.read()
    return {"path": str(path), "sheets": list(model), "content_digest": content_digest(model),
            "file_sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}


def read(path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def validate(path, model: dict, *, sources=(), summary_sheet="01_BOQ_SUMMARY", item_col=0, approved_col=None,
             canonical_ids=(), legacy_ids=(), tol=0.0, banner: str = BANNER, summary_name: str | None = None) -> dict:
    """sources: [{"sheet", "row" (0-based data row), "col", "value", "ref"}] - the register value each quantity cell
    was copied from (fetched by the caller from the register, not from the model)."""
    got, exp = read(path), content(model)
    diffs, formulas, drops = [], [], []
    for name in model:
        rows = got.get(name)
        if rows is None:
            diffs.append({"sheet": name, "error": "missing"})
            continue
        if rows[0][:1] != [banner]:
            diffs.append({"sheet": name, "error": "banner"})
        if rows[1][:1] != [role_text(model[name], summary_name or summary_sheet)]:
            diffs.append({"sheet": name, "error": "role line"})
        body = [_trim(r) for r in rows[HEADER_ROW:]]
        while body and not body[-1]:
            body.pop()
        want = [_trim(r) for r in exp[name]["rows"]]
        if _trim(rows[HEADER_ROW - 1]) != _trim(exp[name]["header"]):
            diffs.append({"sheet": name, "error": "header"})
        if len(body) != len(want):
            drops.append({"sheet": name, "read": len(body), "model": len(want)})
        bad = next(({"row": i, "got": a, "want": b} for i, (a, b) in enumerate(zip(body, want)) if a != b), None)
        if bad:
            diffs.append({"sheet": name, "error": "content", "first": bad})
        formulas += [f"{name}!R{i + 1}" for i, r in enumerate(rows) for v in r if isinstance(v, str) and
                     v.startswith("=")]
        sc = model[name].get("status_col")
        if sc is not None:
            for i, r in enumerate(body):
                if len(r) > sc and r[sc] not in STATUSES:
                    diffs.append({"sheet": name, "error": "status", "row": i, "value": r[sc]})
    src_bad = []
    for s in sources:
        rows = got.get(s["sheet"]) or []
        i = HEADER_ROW + s["row"]
        v = rows[i][s["col"]] if i < len(rows) and s["col"] < len(rows[i]) else None
        if not (isinstance(v, (int, float)) and s["value"] is not None and abs(v - s["value"]) <= tol) and not (
                v is None and s["value"] is None):
            src_bad.append({"sheet": s["sheet"], "row": s["row"], "col": s["col"], "cell": v, "register": s["value"],
                            "ref": s["ref"]})
    summary = [_trim(r) for r in (got.get(summary_sheet) or [])[HEADER_ROW:]]
    codes = [r[item_col] for r in summary if r]
    alias_lines = sorted(set(codes) & set(legacy_ids))
    count_bad = sorted(c for c in canonical_ids if codes.count(c) != 1)
    approved_bad = []
    if approved_col is not None:
        sc = model[summary_sheet].get("status_col")
        approved_bad = [r[item_col] for r in summary if len(r) > approved_col and r[approved_col] == "YES" and
                        r[sc] != "RELEASED"]
    ok = not (diffs or formulas or src_bad or alias_lines or count_bad or approved_bad or drops)
    return {"state": "PASS" if ok else "FAIL", "differences": diffs, "formulas": formulas,
            "register_mismatches": src_bad, "quantity_cells_checked": len(sources),
            "alias_lines_in_summary": alias_lines, "canonical_not_exactly_once": count_bad,
            "approved_without_release": approved_bad, "row_drops": drops,
            "rows_per_sheet": {n: len(exp[n]["rows"]) for n in model}, "content_digest": content_digest(model)}


def policy_record() -> dict:
    rec = {"policy_id": POLICY_ID, "banner": BANNER, "roles": ROLE_TEXT, "statuses": list(STATUSES),
           "quantity_format": QTY_FORMAT, "zip_time": list(FIXED_ZIP_TIME),
           "never": ["a formula", "a total computed in the workbook", "a quantity not copied from a register row",
                     "a breakdown sheet presented as additive", "a legacy alias as a summary line",
                     "a SHADOW line marked approved", "a price"]}
    rec["digest"] = hashlib.sha256(json.dumps(rec, sort_keys=True).encode()).hexdigest()
    return rec
