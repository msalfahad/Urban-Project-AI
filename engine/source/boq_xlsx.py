"""BOQ XLSX EXPORT (R8.20) - a spreadsheet VIEW over BOQ_REPORT_LAYER_V1 report rows. It calculates NOTHING.

Sheets: 00_READ_ME (banner, status legend, rules), 01_BOQ_SUMMARY (TOTAL lines only), 02_ROOM_BREAKDOWN (room lines,
each marked as a breakdown of its summary item - NOT additive to the summary), 03_TRACEABILITY (one line per report
row), 04_BLOCKERS_NOTES, 05_RUN_INFO. Every quantity cell is the report row's QTY written as a plain number (no
formula, no sum, no rounding of the stored value); every line shows its STATUS; every sheet carries the SHADOW banner.
validate() reads the file back and refuses any cell that differs from its report row, any formula, any missing
status or trace. No prices.
"""

from __future__ import annotations

import datetime
import hashlib
import json

POLICY_ID = "BOQ_XLSX_EXPORT_V1"
SHEETS = ("00_READ_ME", "01_BOQ_SUMMARY", "02_ROOM_BREAKDOWN", "03_TRACEABILITY", "04_BLOCKERS_NOTES", "05_RUN_INFO")
COLUMNS = ("ITEM CODE", "TRADE", "DESCRIPTION AR", "DESCRIPTION EN", "FLOOR", "ROOM / ZONE", "QUANTITY", "UNIT",
           "STATUS", "SOURCE REVISION", "RULE / AUTHORITY", "BLOCKER / NOTE", "APPROVED FOR BOQ")
TRACE_COLUMNS = ("ITEM CODE", "ROW KIND", "ROOM / ZONE", "STATUS", "RUN ID", "SOURCE REVISION", "TRADE ROW", "SITES",
                 "SURFACE IDS", "AUTHORITY DIGEST", "RULE IDS", "OWNER FACTS", "RELEASE STATE", "EVIDENCE ROW")
BANNER = "SHADOW BOQ - NOT APPROVED FOR TENDER / CONTRACT - PRESENTATION VIEW OVER EVIDENCE - NO PRODUCTION MIGRATION"
FILL = {"COMPUTED_SHADOW_COMPLETE": "FFF2CC", "AUTHORISED_SUBTOTAL": "F8CBAD", "BLOCKED": "D9D9D9",
        "RELEASED": "C6EFCE"}
QTY_FORMAT = "0.000000"


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    return json.dumps(v, ensure_ascii=False, sort_keys=True, default=str)


def _line(r, *, breakdown=False):
    zone = r["ROOM_ZONE"] if isinstance(r["ROOM_ZONE"], str) else " / ".join(r["ROOM_ZONE"])
    cells = [r["ITEM_CODE"], r["TRADE"], r["DESCRIPTION_AR"], r["DESCRIPTION_EN"], r["LOCATION_FLOOR"], zone, r["QTY"],
             r["UNIT"], r["STATUS"], r["trace"]["source_revision"], _txt(r["RULE_AUTHORITY"]),
             _txt(r["NOTES_BLOCKERS"]), "YES" if r["approved_for_boq"] else "NO - SHADOW"]
    return ([f"BREAKDOWN OF {r['ITEM_CODE']} (not additive to 01_BOQ_SUMMARY)"] if breakdown else []) + cells


def _trace(r):
    t = r["trace"]
    zone = r["ROOM_ZONE"] if isinstance(r["ROOM_ZONE"], str) else " / ".join(r["ROOM_ZONE"])
    return [r["ITEM_CODE"], r["ROW_KIND"], zone, r["STATUS"], t["run_id"], t["source_revision"], t["trade_row"],
            _txt(t["sites"]), _txt(t["surface_ids"]), t["authority_digest"], _txt(t["rule_ids"]),
            _txt(t["owner_facts"]), t["release_state"], r["evidence_row"]]


def expected(rep: dict) -> dict:
    """The exact sheet contents (rows of cell values below the banner) derived from the report rows."""
    tot = [r for r in rep["rows"] if r["ROW_KIND"] == "TOTAL"]
    room = [r for r in rep["rows"] if r["ROW_KIND"] == "ROOM"]
    run = rep["run"]
    return {
        "00_READ_ME": [[BANNER], ["This workbook is a VIEW over the evidence rows of the Urban QTO engine. It is not an "
                                  "editable calculation source: every quantity is copied from one evidence row."],
                       ["01_BOQ_SUMMARY holds the TOTAL line of every item. 02_ROOM_BREAKDOWN repeats room lines of the "
                        "same items: a breakdown, NOT additional quantity - never add the two sheets."],
                       ["STATUS: COMPUTED_SHADOW_COMPLETE = complete shadow quantity, NOT approved; "
                        "AUTHORISED_SUBTOTAL = an open contributor or authority gap remains; BLOCKED = no quantity; "
                        "RELEASED = only with a release record (none exists)."],
                       ["No prices. No formulas. Production migration: NO."]],
        "01_BOQ_SUMMARY": [list(COLUMNS)] + [_line(r) for r in tot],
        "02_ROOM_BREAKDOWN": [["BREAKDOWN OF"] + list(COLUMNS)] + [_line(r, breakdown=True) for r in room],
        "03_TRACEABILITY": [list(TRACE_COLUMNS)] + [_trace(r) for r in rep["rows"]],
        "04_BLOCKERS_NOTES": [["ITEM CODE", "STATUS", "BLOCKER / NOTE"]] +
                             [[r["ITEM_CODE"], r["STATUS"], _txt(b)] for r in tot for b in r["NOTES_BLOCKERS"]],
        "05_RUN_INFO": [["KEY", "VALUE"]] + [[k, _txt(v)] for k, v in sorted(run.items())] +
                       [["report_policy", rep["policy"]], ["report_digest", rep["digest"]], ["export_policy", POLICY_ID],
                        ["pricing", "NONE"], ["calculates", "NO"]]}


def _trim(row):
    r = ["" if v is None else v for v in row]
    while r and r[-1] == "":
        r.pop()
    return r


def write(rep: dict, path, *, created: datetime.datetime) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    content = expected(rep)
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        ws.append([BANNER])
        ws["A1"].font = Font(bold=True, color="C00000")
        for row in content[name]:
            ws.append(row)
        if name in ("01_BOQ_SUMMARY", "02_ROOM_BREAKDOWN"):
            off = 1 if name == "02_ROOM_BREAKDOWN" else 0
            qcol, scol = COLUMNS.index("QUANTITY") + 1 + off, COLUMNS.index("STATUS") + 1 + off
            for i in range(3, ws.max_row + 1):
                ws.cell(i, qcol).number_format = QTY_FORMAT
                st = ws.cell(i, scol).value
                if st in FILL:
                    ws.cell(i, scol).fill = PatternFill("solid", fgColor=FILL[st])
            for c in ws[2]:
                c.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
        ws.freeze_panes = "A3"
        ws["A1"].alignment = Alignment(wrap_text=False)
    wb.properties.creator = "Urban QTO (BOQ_XLSX_EXPORT_V1)"
    wb.properties.created = wb.properties.modified = created
    wb.save(str(path))
    return {"path": str(path), "sheets": list(SHEETS), "content_digest": content_digest(content)}


def content_digest(content: dict) -> str:
    return hashlib.sha256(json.dumps(content, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def read(path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def validate(path, rep: dict) -> dict:
    """Read the workbook back: every sheet = banner + exactly the expected rows; no formulas; statuses present."""
    got, exp = read(path), expected(rep)
    diffs, formulas = [], []
    for name in SHEETS:
        rows = got.get(name)
        if rows is None:
            diffs.append({"sheet": name, "error": "missing"})
            continue
        if rows[0][:1] != [BANNER]:
            diffs.append({"sheet": name, "error": "banner"})
        body, want = [_trim(r) for r in rows[1:]], [_trim(r) for r in exp[name]]
        while body and not body[-1]:
            body.pop()
        if body != want:
            diffs.append({"sheet": name, "error": "content differs from the report rows",
                          "first": next(({"row": i, "got": a, "want": b} for i, (a, b) in enumerate(zip(body, want))
                                         if a != b), {"rows": [len(body), len(want)]})})
        formulas += [f"{name}!{i}" for i, r in enumerate(rows) for v in r if isinstance(v, str) and v.startswith("=")]
    summary = got.get("01_BOQ_SUMMARY", [])[2:]
    scol = COLUMNS.index("STATUS")
    no_status = [r[0] for r in summary if r[scol] not in FILL]
    approved = [r[0] for r in summary if r[COLUMNS.index("APPROVED FOR BOQ")] == "YES" and r[scol] != "RELEASED"]
    ok = not diffs and not formulas and not no_status and not approved
    return {"state": "PASS" if ok else "FAIL", "differences": diffs, "formulas": formulas,
            "missing_status": no_status, "approved_without_release": approved,
            "summary_lines": len(summary), "breakdown_lines": max(0, len(got.get("02_ROOM_BREAKDOWN", [])) - 2),
            "trace_lines": max(0, len(got.get("03_TRACEABILITY", [])) - 2), "content_digest": content_digest(exp)}


def policy_record() -> dict:
    rec = {"policy_id": POLICY_ID, "sheets": list(SHEETS), "columns": list(COLUMNS), "trace": list(TRACE_COLUMNS),
           "banner": BANNER, "quantity_format": QTY_FORMAT,
           "never": ["a formula", "a sum or subtotal computed in the workbook", "a quantity not copied from a report "
                     "row", "a summary and a breakdown presented as additive", "a SHADOW line marked approved",
                     "a price"]}
    rec["digest"] = hashlib.sha256(json.dumps(rec, sort_keys=True).encode()).hexdigest()
    return rec
