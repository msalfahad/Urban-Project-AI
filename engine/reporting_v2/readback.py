"""REPORTING V2 - read the saved XLSX back and validate it against the reporting model.

Refuses: a missing / reordered sheet, any formula cell, any value cell that differs from the model (quantity values
compared exactly - the stored value is the engine value, only the number format rounds), a quantity cell count that
differs from the model's quantity cells, a blocked quantity written as a number.
"""

from __future__ import annotations

from .layout import quantity_cells


def validate(path, model, cell_map) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    diffs = []
    want_sheets = [s["name"][:31] for s in model["sheets"]]
    if wb.sheetnames != want_sheets:
        diffs.append(["sheet_order", wb.sheetnames, want_sheets])
    formulas = 0
    for ws in wb:
        for rw in ws.iter_rows():
            for c in rw:
                if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                    formulas += 1
                    diffs.append(["formula", ws.title, c.coordinate])
    counts = {}
    for sheet, coord, kind, exp in cell_map:
        got = wb[sheet][coord].value
        counts[kind] = counts.get(kind, 0) + 1
        if kind in ("qty", "num", "count", "pct", "dim"):
            if isinstance(got, bool) or not isinstance(got, (int, float)) or got != exp:
                diffs.append([sheet, coord, kind, got, exp])
        elif got != exp:
            diffs.append([sheet, coord, kind, got, exp])
    mq = quantity_cells(model)
    num_model = sum(1 for *_, c in mq if c["q"] is not None)
    txt_model = sum(1 for *_, c in mq if c["q"] is None)
    if counts.get("qty", 0) != num_model:
        diffs.append(["quantity_cell_count", counts.get("qty", 0), num_model])
    if counts.get("qty_text", 0) != txt_model:
        diffs.append(["blocked_cell_count", counts.get("qty_text", 0), txt_model])
    for sheet, coord, kind, exp in cell_map:
        if kind == "qty_text" and exp == "BLOCKED" and isinstance(wb[sheet][coord].value, (int, float)):
            diffs.append(["blocked_written_as_number", sheet, coord])
    return {"state": "PASS" if not diffs else "FAIL", "differences": diffs[:50], "difference_count": len(diffs),
            "formula_cells": formulas, "validated_quantity_cells": counts.get("qty", 0),
            "validated_blocked_cells": counts.get("qty_text", 0), "validated_cells_by_kind": counts,
            "validated_cells": sum(counts.values()), "sheets": wb.sheetnames}
