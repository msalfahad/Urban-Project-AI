"""REPORTING V2 - XLSX renderer (openpyxl). Every Urban engine quantity is a plain value copied from the reporting
model (a locked cell, never a formula); only the number format rounds. Formula cells exist only where the model holds a
formula spec (report totals / subtotals, coverage counts, manual-check arithmetic, differences, % and reconciliation
status) - light blue; manual-check inputs are the only unlocked cells - light yellow. Every sheet is protected without a
password (filters, sorting and formatting stay allowed). Print-ready: A4 / A3 landscape, fit to width, repeated title
rows, section page breaks, footer (project - phase | page X of Y | Urban Projects). The saved zip is re-packed with
fixed entry times so the same model gives the same bytes.

render(model, path, logo=None) -> {"file_sha256", "cell_map", ...}; cell_map lists every cell written (sheet,
coordinate, kind, expected[, evaluated value]) for readback.py; kind "formula" carries the exact formula text.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import zipfile
from datetime import datetime
from pathlib import Path

from . import terms as T
from .model import is_f, is_in, is_q, qty_text

NAVY, NAVY_2, BAND, HEAD, GRID, GREY = "1F3A5F", "2C4E78", "E9EEF4", "DCE3EC", "C3CBD5", "6B7280"
STATUS_FILL = {"COMPUTED": "E2EFDA", "PARTIAL": "FFF2CC", "REVIEW": "FFF2CC", "BLOCKED": "F8D7DA", "INFO": "EDEDED"}
INPUT_FILL, FORMULA_FILL = "FFF9C4", "DDEBF7"
RSTATUS_FILL = {"MATCH": ("C6EFCE", "006100"), "CLOSE": ("FFEB9C", "7F6000"), "REVIEW": ("FFC7CE", "9C0006"),
                "ENGINE BLOCKED": ("EDEDED", "555555"), "NOT CHECKED": ("EDEDED", "555555"), "NOT COMPARABLE": ("EDEDED", "555555")}
DIFF_FMT = "+#,##0.000;-#,##0.000;0.000"
PCT_FMT = "+0.0%;-0.0%;0.0%"
FONT = "Arial"
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
CREATED = datetime(2026, 10, 3)
QTY_FMT = {"m3": "#,##0.000", "m2": "#,##0.000", "lm": "#,##0.000", "m": "#,##0.000", "nr": "#,##0"}
SUMMARY_FMT = {"m2": "#,##0.00"}
DEFAULT_W = {"text": 22, "ar": 22, "qty": 13, "count": 9, "dim": 10, "status": 11, "code": 26, "cls": 14, "pct": 9}
TITLE_ROWS = 5


def _fmt(colspec, sheet_role):
    if colspec["kind"] == "count":
        return "#,##0"
    if colspec["kind"] == "pct":
        return "0%"
    if colspec["kind"] == "dim":
        return "0.000"
    if colspec["kind"] == "num":
        return "General"
    u = colspec.get("unit")
    if sheet_role == "SUMMARY" and u in SUMMARY_FMT:
        return SUMMARY_FMT[u]
    return QTY_FMT.get(u, "#,##0.000")


def _repack(path):
    src = zipfile.ZipFile(io.BytesIO(Path(path).read_bytes()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                txt = data.decode("utf-8")
                c = re.search(r"<dcterms:created[^>]*>([^<]*)</dcterms:created>", txt)
                if c:
                    txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", lambda m: m.group(1) + c.group(1) + m.group(2), txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, data)
    Path(path).write_bytes(buf.getvalue())


def render(model, path, logo=None) -> dict:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.protection import SheetProtection

    thin = Side(style="thin", color=GRID)
    med = Side(style="medium", color=NAVY)
    thick = Side(style="thick", color=NAVY)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = lambda c: PatternFill("solid", fgColor=c)
    proj = model["project"]
    cell_map = []
    wb = Workbook()
    wb.remove(wb.active)
    for s in model["sheets"]:
        ws = wb.create_sheet(s["name"][:31])
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = GREY if s["role"] == "TECH" else NAVY
        grid, spans = _grid(s["sections"])
        ncols = len(grid)
        last = ncols + 1                                                   # column A is a margin
        ws.column_dimensions["A"].width = 2
        for j, w in enumerate(grid):
            ws.column_dimensions[get_column_letter(j + 2)].width = round(w, 2)
        total_w = sum(grid) + 2
        # ---------------- title block (rows 1-5, repeated on every printed page)
        ws.row_dimensions[1].height = 34
        if logo and Path(logo).exists():
            from openpyxl.drawing.image import Image
            im = Image(str(logo))
            im.width, im.height = 44, 44
            ws.add_image(im, "B1")
        ws.cell(1, 3, f"{proj['name_en']}  —  QTO REPORT  —  {proj['phase']}").font = Font(FONT, size=14, bold=True, color=NAVY)
        ws.cell(1, last, proj["release_state"]).font = Font(FONT, size=9, bold=True, color="9C2B2B")
        ws.cell(1, last).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(1, 3).alignment = Alignment(vertical="center")
        ws.cell(2, 3, proj["subtitle"]).font = Font(FONT, size=8.5, color=GREY)
        ws.cell(2, last, proj.get("name_ar", "")).font = Font(FONT, size=10, bold=True, color=NAVY)
        ws.cell(2, last).alignment = Alignment(horizontal="right", readingOrder=2)
        ws.row_dimensions[4].height = 30
        half = max(2, ncols // 2)
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=1 + half)
        ws.merge_cells(start_row=4, start_column=2 + half, end_row=4, end_column=last)
        ws.cell(4, 2, f"{s['name'][:2]}.  {s['title_en']}").font = Font(FONT, size=16, bold=True, color="FFFFFF")
        ws.cell(4, 2).alignment = Alignment(vertical="center", indent=1)
        ws.cell(4, 2 + half, s["title_ar"]).font = Font(FONT, size=16, bold=True, color="FFFFFF")
        ws.cell(4, 2 + half).alignment = Alignment(horizontal="right", vertical="center", readingOrder=2, indent=1)
        for j in range(2, last + 1):
            ws.cell(4, j).fill = fill(NAVY)
        if s.get("note"):
            ws.cell(5, 2, s["note"]).font = Font(FONT, size=8.5, italic=True, color="9C2B2B")
        r = TITLE_ROWS + 1
        page_used, breaks = 0.0, []
        rowpos, colpos, pending, cf = {}, {}, [], []
        scale = min(1.0, (277 if total_w <= 170 else 400) / (total_w * 1.9 + 1))
        cap = 420.0 / max(scale, 0.35)
        for si, sec in enumerate(s["sections"]):
            sp = spans[si]
            cols = [dict(c, width=sum(grid[a:b + 1])) for c, (a, b) in zip(sec["columns"], sp)]
            est = 46 + sum(_row_h(rw, cols) for rw in sec["rows"]) + (80 if sec["kind"] == "coverage" else 0)
            if page_used > 0 and page_used + est > cap and est <= cap:
                breaks.append(r - 1)
                page_used = 0.0
            r += 1
            # section band
            ws.row_dimensions[r].height = 22
            span = max(sp[-1][1] + 1, 2)
            for j in range(2, 2 + span):
                ws.cell(r, j).fill = fill(NAVY_2)
            ws.cell(r, 2, sec["title_en"]).font = Font(FONT, size=11, bold=True, color="FFFFFF")
            ws.cell(r, 2).alignment = Alignment(vertical="center", indent=1)
            ws.cell(r, 1 + span, sec["title_ar"]).font = Font(FONT, size=11, bold=True, color="FFFFFF")
            ws.cell(r, 1 + span).alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
            r += 1
            header_row = r
            if sec["kind"] not in ("notes", "info"):
                for j, c in enumerate(cols):
                    label = c["en"] + (f" ({c['unit']})" if c.get("unit") and c["kind"] == "qty" and c["en"] not in ("QTY",) else "")
                    if c.get("ar"):
                        label += "\n" + c["ar"]
                    a, b = sp[j]
                    for jj in range(a, b + 1):
                        ws.cell(r, 2 + jj).fill = fill(HEAD)
                        ws.cell(r, 2 + jj).border = box
                    x = ws.cell(r, 2 + a, label)
                    x.font = Font(FONT, size=8.5, bold=True, color=NAVY)
                    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    if b > a:
                        ws.merge_cells(start_row=r, start_column=2 + a, end_row=r, end_column=2 + b)
                ws.row_dimensions[r].height = max(30, 11.5 * max(
                    math.ceil(len(c["en"] + (f" ({c['unit']})" if c.get("unit") else "")) / max(cc["width"] * 1.05, 1))
                    + (1 if c.get("ar") else 0) for c, cc in zip(sec["columns"], cols)) + 4)
                r += 1
            for j, c in enumerate(cols):
                colpos[(sec["id"], c["key"])] = 2 + sp[j][0]
            first_data = r
            for i, rw in enumerate(sec["rows"]):
                rowpos[(sec["id"], rw.get("key", i))] = r
                is_group = rw["role"] == "NOTE" and rw.get("note") == "GROUP"
                bold = rw["role"] in ("SUBTOTAL", "TOTAL") or is_group
                for j, c in enumerate(cols):
                    if j >= len(rw["cells"]):
                        break
                    v = rw["cells"][j]
                    a, b = sp[j]
                    x = ws.cell(r, 2 + a)
                    kind = c["kind"]
                    if is_f(v):
                        u = v.get("fmt") or c.get("unit")
                        x.number_format = (DIFF_FMT if c["key"] == "diff" else PCT_FMT if c["key"] == "pct" and kind == "pct"
                                           else "0.0%" if kind == "pct" else "#,##0" if kind == "count"
                                           else _fmt(dict(c, unit=u), s["role"]) if kind == "qty" else "General")
                        x.fill = fill(FORMULA_FILL)
                        x.font = Font(FONT, size=9 if kind != "rstatus" else 8.5, bold=bold or kind == "rstatus")
                        x.alignment = Alignment(horizontal="center" if kind == "rstatus" else "right", vertical="center")
                        pending.append((x, v, [sec["id"], rw.get("key", i), c["key"]]))
                    elif is_in(v):
                        x.value = v["input"]
                        x.fill = fill(INPUT_FILL)
                        x.protection = Protection(locked=False)
                        if v["itype"] == "text":
                            x.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                        else:
                            x.number_format = "0.0%" if v["itype"] == "pct" else ("#,##0" if kind == "count" else
                                                                                  "0.000" if kind == "dim" else _fmt(c, s["role"]))
                            x.alignment = Alignment(horizontal="right", vertical="center")
                        x.font = Font(FONT, size=9, bold=bold)
                        cell_map.append([ws.title, x.coordinate, "input", v["input"], [sec["id"], rw.get("key", i), c["key"]]])
                    elif is_q(v):
                        if v["q"] is None:
                            x.value = qty_text(v)
                            x.font = Font(FONT, size=9, bold=bold, color="9C2B2B" if "blocked" in v else GREY)
                            x.alignment = Alignment(horizontal="center", vertical="center")
                            cell_map.append([ws.title, x.coordinate, "qty_text", qty_text(v)])
                        else:
                            x.value = v["q"]
                            u = v.get("unit") or c.get("unit")
                            x.number_format = (QTY_FMT.get(u, "#,##0.000") if kind == "qty" and v.get("unit") else
                                               _fmt(c, s["role"]) if kind in ("qty", "dim", "count", "pct") else QTY_FMT.get(u, "#,##0.000"))
                            x.font = Font(FONT, size=9, bold=bold)
                            x.alignment = Alignment(horizontal="right", vertical="center")
                            cell_map.append([ws.title, x.coordinate, "qty", v["q"]])
                    elif kind == "status" and v in T.STATUS_LABEL:
                        x.value = T.STATUS_LABEL[v][0]
                        x.fill = fill(STATUS_FILL[v])
                        x.font = Font(FONT, size=8.5, bold=True, color="333333")
                        x.alignment = Alignment(horizontal="center", vertical="center")
                        cell_map.append([ws.title, x.coordinate, "status", T.STATUS_LABEL[v][0]])
                    else:
                        x.value = v if not isinstance(v, (list, dict)) else None
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            x.number_format = _fmt(c, s["role"])
                            cell_map.append([ws.title, x.coordinate, kind if kind in ("count", "pct", "dim") else "num", v])
                        elif v not in (None, ""):
                            cell_map.append([ws.title, x.coordinate, "text", v])
                        ar = kind == "ar"
                        x.font = Font(FONT, size=10 if ar else (9 if not is_group else 9.5), bold=bold,
                                      color=NAVY if is_group else "222222")
                        x.alignment = Alignment(horizontal="right" if ar else ("right" if isinstance(v, (int, float)) else "left"),
                                                vertical="center", wrap_text=not isinstance(v, (int, float)),
                                                readingOrder=2 if ar else 0)
                    for jj in range(a, b + 1):
                        y = ws.cell(r, 2 + jj)
                        if sec["kind"] not in ("notes", "info"):
                            y.border = box
                        if is_group:
                            y.fill = fill(BAND)
                        elif rw["role"] in ("SUBTOTAL", "TOTAL"):
                            if kind != "status" and not (is_f(v) or is_in(v)):
                                y.fill = fill("F2F4F7")
                            y.border = Border(left=thin, right=thin, top=med, bottom=thin)
                        if is_group and sec.get("floor_blocks"):
                            y.border = Border(left=thin, right=thin, top=thick, bottom=thin)
                    if b > a and sec["kind"] not in ("notes", "info"):
                        ws.merge_cells(start_row=r, start_column=2 + a, end_row=r, end_column=2 + b)
                if sec["kind"] in ("notes",) and rw["cells"]:
                    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
                if sec["kind"] == "info" and len(rw["cells"]) == 2:
                    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=last)
                    ws.cell(r, 2).font = Font(FONT, size=9, bold=True, color=NAVY)
                ws.row_dimensions[r].height = _row_h(rw, cols if sec["kind"] not in ("notes",) else [{"width": total_w, "kind": "text"}])
                r += 1
            last_data = r - 1
            if sec.get("filter") and last_data >= first_data:
                ws.auto_filter.ref = f"B{header_row}:{get_column_letter(2 + sp[-1][1])}{last_data}"
            for j, c in enumerate(cols):
                if c["kind"] == "rstatus" and last_data >= first_data:
                    cf.append(f"{get_column_letter(2 + sp[j][0])}{first_data}:{get_column_letter(2 + sp[j][0])}{last_data}")
            if sec.get("note"):
                ws.cell(r, 2, sec["note"]).font = Font(FONT, size=8, italic=True, color=GREY)
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
                ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[r].height = 12 * max(1, math.ceil(len(sec["note"]) / (total_w * 1.15)))
                r += 1
            if sec["kind"] == "coverage":
                ch = BarChart()
                ch.type, ch.grouping, ch.overlap = "bar", "standard", 0
                n = len(sec["rows"])
                ch.add_data(Reference(ws, min_col=4, min_row=header_row + 1, max_row=header_row + n), titles_from_data=False)
                ch.set_categories(Reference(ws, min_col=2, min_row=header_row + 1, max_row=header_row + n))
                colours = {"COMPUTED": "5B9A3C", "PARTIAL": "D9822B", "REVIEW": "C9A227", "BLOCKED": "C0504D"}
                ser = ch.series[0]
                for k, rw in enumerate(sec["rows"]):
                    pt = DataPoint(idx=k)
                    pt.graphicalProperties.solidFill = colours.get(rw["status"], "BFBFBF")
                    ser.dPt.append(pt)
                ch.dataLabels = DataLabelList()
                ch.dataLabels.showVal = True
                ch.dataLabels.showSerName = False
                ch.dataLabels.showCatName = False
                ch.dataLabels.showLegendKey = False
                ch.legend = None
                ch.y_axis.majorGridlines = None
                ch.y_axis.delete = True
                ch.x_axis.scaling.orientation = "maxMin"
                ch.gapWidth = 40
                ch.height, ch.width = 4.2, 9.5
                ws.add_chart(ch, f"{get_column_letter(min(last, 2 + sp[-1][1] + 2))}{header_row}")
                r += 5                                                     # room for the chart below the table
            page_used = (page_used + est) % cap if page_used + est > cap else page_used + est
        for b in breaks:
            ws.row_breaks.append(Break(id=b))
        # ---------------- formulas (second pass: every reference resolves to a written cell of this sheet)
        for x, v, where in pending:
            text = formula_text(v["f"], rowpos, colpos)
            x.value = text
            cell_map.append([ws.title, x.coordinate, "formula", text, v["v"], where])
        for ref_ in cf:
            for label, (bg, fg) in RSTATUS_FILL.items():
                ws.conditional_formatting.add(ref_, CellIsRule(operator="equal", formula=[f'"{label}"'], fill=fill(bg),
                                                                font=Font(FONT, size=8.5, bold=True, color=fg)))
        ws.protection = SheetProtection(sheet=True, formatCells=False, formatColumns=False, formatRows=False,
                                        autoFilter=False, sort=False, selectLockedCells=False, selectUnlockedCells=False)
        # ---------------- print setup
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4 if total_w <= 170 else ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top, ws.page_margins.bottom = 0.5, 0.55
        ws.print_title_rows = f"1:{TITLE_ROWS}"
        ws.print_area = f"A1:{get_column_letter(last)}{r}"
        ws.oddFooter.left.text = f"{proj['name_en']} — QTO REPORT — {proj['phase']}"
        ws.oddFooter.center.text = "Page &P of &N"
        ws.oddFooter.right.text = "Urban Projects"
        for part in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
            part.size, part.font = 8, FONT
        if s["role"] in ("BREAKDOWN", "SCHEDULE"):
            ws.oddHeader.right.text = "Project totals: use 00_TOTAL_SUMMARY only"
            ws.oddHeader.right.size = 8
        ws.freeze_panes = f"A{TITLE_ROWS + 1}"
    wb.properties.creator = "Urban QTO Reporting V2"
    wb.properties.title = f"{proj['name_en']} - QTO REPORT - {proj['phase']}"
    wb.properties.created = wb.properties.modified = CREATED
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(str(path))
    _repack(path)
    data = Path(path).read_bytes()
    return {"file_sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data), "sheets": [s["name"] for s in model["sheets"]],
            "cell_map": cell_map}


def formula_text(f, rowpos, colpos) -> str:
    """The Excel text of a formula spec (see model.FORMULA_OPS). PARAMS references are absolute."""
    from openpyxl.utils import get_column_letter as L

    def a1(sec, rk, ck):
        r, c = rowpos[(sec, rk)], colpos[(sec, ck)]
        return f"${L(c)}${r}" if sec == "PARAMS" else f"{L(c)}{r}"

    def R(a):
        if "ref" in a:
            return a1(*a["ref"])
        if "sub" in a:
            return f"({R(a['sub'][0])}-{R(a['sub'][1])})"
        sec, k1, k2, ck = a["range"]
        return f"{a1(sec, k1, ck)}:{a1(sec, k2, ck)}"

    op, args = f["op"], f["args"]
    if op in ("SUM", "SUMNB", "COUNT"):
        refs = [R(a) for a in args]
        chunks = [refs[i:i + 200] for i in range(0, len(refs), 200)]
        body = ",".join(refs) if len(chunks) <= 1 else ",".join(f"SUM({','.join(c)})" for c in chunks)
        if op == "COUNT":
            return f"=COUNT({body})"
        if op == "SUM":
            return f"=SUM({body})"
        cnt = "+".join(f"COUNT({','.join(c)})" for c in chunks)
        return f'=IF({cnt}=0,"",SUM({body}))'
    if op == "DIFF":
        u, m = R(args[0]), R(args[1])
        return f'=IF(AND(ISNUMBER({u}),ISNUMBER({m})),ROUND({m}-{u},9),"")'
    if op == "PCT":
        u, m = R(args[0]), R(args[1])
        return f'=IF(AND(ISNUMBER({u}),ISNUMBER({m})),IF({u}<>0,ROUND(({m}-{u})/{u},9),""),"")'
    if op == "RSTATUS":
        u, m, t1, t2 = (R(a) for a in args)
        return (f'=IF({u}="BLOCKED","ENGINE BLOCKED",IF(NOT(ISNUMBER({u})),"NOT COMPARABLE",IF({m}="","NOT CHECKED",'
                f'IF(NOT(ISNUMBER({m})),"NOT COMPARABLE",IF({u}=0,IF(ABS({m})<=0.000000001,"MATCH","REVIEW"),'
                f'IF(ABS({m}/{u}-1)<={t1},"MATCH",IF(ABS({m}/{u}-1)<={t2},"CLOSE","REVIEW")))))))')
    if op == "PROD":
        *fs, div = args
        terms_ = [R(a) for a in fs]
        checks = [R(x) for a in fs for x in (a["sub"] if "sub" in a else [a])]
        prod = "*".join(terms_) + (f"/{div:g}" if div != 1 else "")
        return f'=IF(AND({",".join(f"ISNUMBER({c})" for c in checks)}),{prod},"")'
    if op == "COUNTIFS":
        return "=COUNTIFS(" + ",".join(f'{R(a)},"{crit}"' for a, crit in zip(args[0::2], args[1::2])) + ")"
    if op == "DIV":
        a, b = R(args[0]), R(args[1])
        return f'=IF(AND(ISNUMBER({a}),ISNUMBER({b})),IF({b}<>0,{a}/{b},""),"")'
    raise ValueError(op)


def _grid(sections):
    """One column grid per sheet: the table with the most columns is the anchor; every other table maps each of its
    columns onto one or more consecutive grid columns (merged) until its desired width is reached."""
    w = lambda c: c.get("width") or DEFAULT_W.get(c["kind"], 16)
    tables = [s for s in sections if s["kind"] not in ("notes", "info")]
    if not tables:
        return [20.0] * 6, [[(0, 5)] * len(s["columns"]) for s in sections]
    anchor = max(tables, key=lambda s: len(s["columns"]))
    grid = [float(w(c)) for c in anchor["columns"]]
    need = max(sum(w(c) for c in s["columns"]) for s in tables)
    if sum(grid) < need:
        grid = [g * need / sum(grid) for g in grid]
    spans = []
    for s in sections:
        if s is anchor:
            spans.append([(j, j) for j in range(len(grid))])
            continue
        if s["kind"] in ("notes", "info"):
            spans.append([(0, 0)] + [(1, len(grid) - 1)] * (len(s["columns"]) - 1) if len(s["columns"]) > 1 else [(0, len(grid) - 1)])
            continue
        out, g = [], 0
        desired = [w(c) for c in s["columns"]]
        if sum(desired) >= 0.6 * sum(grid):
            # proportional: snap the scaled column boundaries to the grid, every column at least one grid column,
            # the last column ends at the grid end
            k = sum(grid) / sum(desired)
            cum_g = [sum(grid[:i + 1]) for i in range(len(grid))]
            acc = 0.0
            for j, d in enumerate(desired):
                rest = len(desired) - j - 1
                acc += d * k
                if rest == 0:
                    b = len(grid) - 1
                else:
                    b = min(range(g, len(grid) - rest), key=lambda i: abs(cum_g[i] - acc))
                out.append((g, b))
                g = b + 1
            spans.append(out)
            continue
        for k, c in enumerate(s["columns"]):
            rest = len(s["columns"]) - k - 1
            a, acc = g, grid[g]
            g += 1
            while acc < 0.9 * w(c) and len(grid) - g > rest:
                acc += grid[g]
                g += 1
            out.append((a, g - 1))
        spans.append(out)
    return grid, spans


def _row_h(rw, cols) -> float:
    lines = 1
    for j, c in enumerate(cols):
        if j >= len(rw["cells"]):
            break
        v = rw["cells"][j]
        if isinstance(v, str) and v:
            w = (c.get("width") or DEFAULT_W.get(c.get("kind"), 16)) * 1.15
            lines = max(lines, sum(math.ceil(max(len(p), 1) / max(w, 1)) for p in v.split("\n")))
    return 13.0 * lines + 3
