"""REPORTING V3 - readback: every engine value and formula is where the cell map says, and after a LibreOffice
recalculation every formula evaluates to the value the model expects."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

TOL = 1e-6


def _same(got, want):
    if want in (None, ""):
        return got in (None, "")
    if isinstance(want, str):
        return got == want
    try:
        return abs(float(got) - float(want)) <= TOL * max(1.0, abs(float(want)))
    except (TypeError, ValueError):
        return False


def validate(path, cell_map) -> dict:
    wb = load_workbook(path)
    bad = []
    counts = {"value": 0, "formula": 0, "input": 0}
    for c in cell_map:
        v = wb[c["sheet"]][c["cell"]].value
        counts[c["kind"]] += 1
        if c["kind"] == "formula":
            if v != c["formula"]:
                bad.append({**c, "got": v})
        elif not _same(v, c["expected"]):
            bad.append({**c, "got": v})
    return {"state": "PASS" if not bad else "FAIL", "checked": counts, "problems": bad[:20]}


def recalc(path, cell_map) -> dict:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return {"state": "NOT_RUN", "why": "LibreOffice not installed"}
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "out"
        out.mkdir()
        subprocess.run([exe, "--headless", "--norestore", f"-env:UserInstallation=file://{td}/profile",
                        "--convert-to", "xlsx", "--outdir", str(out), str(path)], capture_output=True, timeout=240)
        f = out / Path(path).name
        if not f.exists():
            return {"state": "FAIL", "why": "LibreOffice produced no file"}
        wb = load_workbook(f, data_only=True)
        bad, n = [], 0
        for c in cell_map:
            if c["kind"] != "formula":
                continue
            n += 1
            v = wb[c["sheet"]][c["cell"]].value
            if not _same(v, c["expected"]):
                bad.append({"sheet": c["sheet"], "cell": c["cell"], "formula": c["formula"], "expected": c["expected"],
                            "got": v})
        errs = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row
                   if isinstance(cell.value, str) and cell.value.startswith("#") and cell.value.endswith(("!", "?", "A")))
    return {"state": "PASS" if not bad and not errs else "FAIL", "formulas": n, "mismatches": bad[:20],
            "error_cells": errs}
