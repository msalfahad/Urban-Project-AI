"""REPORTING V3b - workbook layouts over the two-view line model (TECHNICAL_QTO / COMMERCIAL_BOQ).

Trade workbook:  TOTAL | GF | 1F | 2F+ROOF | OTHER+EXTERNAL | DETAIL | TRACE (+ ROOMS for 05_FLOORING)
  level sheet  one row per BOQ line: technical class / qty / IN TECH, commercial class / qty / IN COMM, confidence,
               range, procurement eligibility, method / assumption, waste method / % / qty / procurement; subtotals are
               SUMIFS per summary item and unit, one for each view (IN TECH = Y, IN COMM = Y)
  TOTAL        per summary item: technical by level (links) + TOTAL, commercial by level (links) + TOTAL, of which
               provisional, budget allowance (beside, never inside), procurement-eligible part (frozen values)
Master:        MASTER_SUMMARY block A (NET / TECHNICAL) and block B (COMMERCIAL) | COVERAGE | BLOCKERS | EXPECTED_SCOPE |
               OWNER_METHODS | DUAL_MEASUREMENT | REBAR_POPULATION
Display units: m³, m², lm, m, No., t (engine kg / 1000) - never a cross-unit total.
"""

from __future__ import annotations

from collections import defaultdict

from openpyxl.styles import PatternFill
from openpyxl.worksheet.pagebreak import Break

from .book import QTY_FMT, STATUS_FILL, Book

LEVEL_SHEET = {"GF": "GF", "1F": "1F", "2F_ROOF": "2F+ROOF", "OTHER_EXTERNAL": "OTHER+EXTERNAL"}
DISPLAY = {"m3": ("m³", 1.0), "m2": ("m²", 1.0), "lm": ("lm", 1.0), "m": ("m", 1.0), "nr": ("No.", 1.0),
           "kg": ("t", 0.001)}
for k in ("DERIVED", "MEASURED", "RASTER_DERIVED", "SOURCE_RULE", "CODE_METHOD", "OWNER_PROJECT_FACT", "URBAN_STANDARD"):
    STATUS_FILL.setdefault(k, ("C6EFCE", "006100"))
for k in ("PROVISIONAL_SOURCE_DERIVED", "PROVISIONAL_SOURCE_RANGE", "PROVISIONAL_GEOMETRIC_INFERENCE",
          "PROVISIONAL_CODE_METHOD", "PROVISIONAL_URBAN_FALLBACK", "PROVISIONAL_OWNER_METHOD", "OWNER_APPROVED_PROVISIONAL"):
    STATUS_FILL.setdefault(k, ("FCE4D6", "843C0C"))
STATUS_FILL.setdefault("BUDGET_ESTIMATE", ("F8CBAD", "833C0B"))
STATUS_FILL.setdefault("PENDING", ("EDEDED", "555555"))
STATUS_FILL.setdefault("BLOCKED_SOURCE_CONFLICT", ("FFC7CE", "9C0006"))
STATUS_FILL.setdefault("NOT_IN_SCOPE", ("EDEDED", "555555"))


def _q(u):
    return QTY_FMT.get(u, "#,##0.000")


def _ref(sheet, cell):
    return f"'{sheet}'!{cell}"


def _rnd(v, n=6):
    return None if v is None else round(v, n)


def du(unit):
    return DISPLAY.get(unit, (unit, 1.0))


def _s(v, f):
    return None if v is None else v * f


def in_tech(x):
    return bool(x["release"]["technical"]["in_total"] and x.get("sumrow")
                and x.get("no_total") in (None, False, "ALTERNATIVE_BASIS"))


def in_comm(x):
    return bool(x["release"]["commercial"]["in_commercial_total"] and x.get("sumrow")
                and x.get("no_total") in (None, False, "ALTERNATIVE_BASIS"))


def trade_book(model, code, trade, trade_ar, out_dir, meta, rooms=None) -> dict:
    lines = [x for x in model["lines"] if x["trade"] == trade]
    names = model["level_names"]
    b = Book(out_dir / f"{code}_{trade}.xlsx", project=meta["project"], phase=meta["phase"])
    tot = b.sheet("TOTAL", [42, 28, 7] + [12] * 10 + [13, 13, 13, 12])
    sheets = {lv: b.sheet(LEVEL_SHEET[lv], [5, 18, 42, 30, 40, 7, 16, 12, 6, 18, 12, 6, 6, 11, 11, 7, 40, 26, 8, 10, 11,
                                            28, 8]) for lv in model["levels"]}
    subs = {}
    for lv in model["levels"]:
        ws = sheets[lv]
        en, ar = names[lv]
        b.title(ws, f"{trade.replace('_', ' ')} - {en}", f"{trade_ar} - {ar}",
                f"{meta['project']} | {meta['phase']} | engine values frozen; light blue = subtotal formulas; "
                f"TECH = technical QTO, COMM = commercial BOQ (labelled provisional)")
        b.header(ws, ["#", "CODE", "DESCRIPTION", "الوصف", "CALCULATION / FORMULA / REASON", "UNIT", "TECH CLASS",
                      "TECH QTY", "IN TECH", "COMM CLASS", "COMM QTY", "IN COMM", "CONF", "LOW", "HIGH", "PROCURE",
                      "METHOD / ASSUMPTION", "WASTE METHOD", "WASTE %", "WASTE QTY", "PROCUREMENT", "SUMMARY ITEM", "LINE"])
        xs = [x for x in lines if x["level"] == lv]
        first = ws._row
        if not xs:
            b.note(ws, "No item of this trade at this level.")
        groups = defaultdict(list)
        for x in xs:
            groups[x["group"]].append(x)
        n = 0
        for g, gx in groups.items():
            b.band(ws, g)
            for x in gx:
                n += 1
                u, f = du(x["unit"])
                t, c, w = x["release"]["technical"], x["release"]["commercial"], x["waste"]
                ma = " | ".join(s for s in (c.get("method"), c.get("assumption")) if s and s != "AS TECHNICAL")
                if x.get("no_total"):
                    ma = (f"[{x['no_total']} - not summed] " + ma).strip()
                b.row(ws, [n, x["code"], x["desc_en"], x["desc_ar"], x["formula"], u, t["class"], _s(t["qty"], f),
                           "Y" if in_tech(x) else "N", c["class"], _s(c["qty"], f), "Y" if in_comm(x) else "N",
                           c.get("confidence"), _s(c.get("low"), f), _s(c.get("high"), f),
                           "Y" if c.get("procurement_eligible") else "N", ma, w.get("WASTE_METHOD"), w.get("WASTE_PCT"),
                           _s(w.get("WASTE_QTY"), f), _s(w.get("PROCUREMENT"), f), x.get("sumrow") or "-", x["line_id"]],
                      kinds=[None] * 7 + ["qty", None, None, "qty"], fmts=[None] * 7 + [_q(u), None, None, _q(u), None,
                                                                                         None, _q(u), _q(u), None, None,
                                                                                         None, "0.00", _q(u), _q(u)],
                      status_col=6, ar_cols=(3,), wrap_cols=(2, 4, 16))
                bg = STATUS_FILL.get(c["class"], ("FFFFFF", "000000"))[0]
                ws.cell(row=ws._row - 1, column=10).fill = PatternFill("solid", start_color=bg, end_color=bg)
        last = ws._row - 1
        b.divider(ws)
        st, sc = defaultdict(float), defaultdict(float)
        for x in xs:
            if not x.get("sumrow"):
                continue
            u, f = du(x["unit"])
            st[(x["sumrow"], u)] += (x["release"]["technical"]["qty"] or 0.0) * f if in_tech(x) else 0.0
            sc[(x["sumrow"], u)] += (x["release"]["commercial"]["qty"] or 0.0) * f if in_comm(x) else 0.0
        keys = sorted(set(st) | set(sc))
        if keys:
            b.band(ws, f"SUBTOTALS - {en}   |   المجاميع - {ar}", color="2C4E78")
            for item, u in keys:
                ft = (f'=SUMIFS($H${first}:$H${last},$V${first}:$V${last},"{item}",$F${first}:$F${last},"{u}",'
                      f'$I${first}:$I${last},"Y")')
                fc = (f'=SUMIFS($K${first}:$K${last},$V${first}:$V${last},"{item}",$F${first}:$F${last},"{u}",'
                      f'$L${first}:$L${last},"Y")')
                r = b.row(ws, ["", "", f"SUBTOTAL - {item}", "", "IN TECH = Y / IN COMM = Y rows above", u, "",
                               {"f": ft, "expect": _rnd(st[(item, u)])}, "", "", {"f": fc, "expect": _rnd(sc[(item, u)])},
                               "", "", "", "", "", "", "", "", "", "", item, ""],
                          fmts=[None] * 7 + [_q(u), None, None, _q(u)], bold=True)
                subs[(lv, item, u)] = (f"H{r}", st[(item, u)], f"K{r}", sc[(item, u)])
        ws.row_breaks.append(Break(id=ws._row))
    # TOTAL
    b.title(tot, f"{trade.replace('_', ' ')} - TRADE TOTAL (A: TECHNICAL | B: COMMERCIAL)", f"{trade_ar} - الإجمالي",
            f"{meta['project']} | level cells link to the level subtotals; budget allowances are shown beside the "
            f"commercial total, never inside it")
    b.header(tot, ["SUMMARY ITEM", "البند", "UNIT", "A GF", "A 1F", "A 2F+R", "A OTH", "A TOTAL", "B GF", "B 1F", "B 2F+R",
                   "B OTH", "B TOTAL", "OF WHICH PROVISIONAL", "BUDGET (beside)", "PROCUREMENT ELIGIBLE", "STATUS"])
    for r in [m for m in model["master"]["rows"] if m["trade"] == trade]:
        u, f = du(r["unit"])
        vals = [r["item"], r["item_ar"], u]
        rr = tot._row
        for blk, ci, vi in (("technical", 0, 1), ("commercial", 2, 3)):
            for lv in model["levels"]:
                k = (lv, r["item"], u)
                if k in subs:
                    vals.append({"f": "=" + _ref(LEVEL_SHEET[lv], subs[k][ci]), "expect": _rnd(subs[k][vi])})
                else:
                    vals.append(0)
            col0 = "D" if blk == "technical" else "I"
            col1 = "G" if blk == "technical" else "L"
            vals.append({"f": f"=SUM({col0}{rr}:{col1}{rr})", "expect": _rnd(r[blk]["total"] * f)})
        c = r["commercial"]
        vals += [_s(c["provisional_part"], f), _s(c["budget"], f), _s(c["procurement_eligible"], f), _status(r)]
        b.row(tot, vals, fmts=[None, None, None] + [_q(u)] * 13, status_col=16, ar_cols=(1,))
    b.skip(tot)
    b.note(tot, "A = technical classes only (MEASURED / DERIVED / RASTER_DERIVED / SOURCE_RULE / CODE_METHOD / OWNER fact / "
                "URBAN standard / PARTIAL measured part). B adds the labelled provisional classes; BUDGET_ESTIMATE is "
                "beside, never inside. Different units are never added; rebar purchased (BBS) is its own basis.")
    # DETAIL / TRACE
    det = b.sheet("DETAIL", [8, 22, 34, 70, 13, 9, 22])
    b.title(det, f"{trade.replace('_', ' ')} - DETAILED CALCULATION", f"{trade_ar} - تفاصيل الحساب")
    b.header(det, ["LINE", "CODE", "REF", "FORMULA", "QTY (engine unit)", "ENGINE UNIT", "STATUS"])
    for x in lines:
        for d in x.get("details") or []:
            b.row(det, [x["line_id"], x["code"], d.get("ref"), d.get("formula"), d.get("qty"), x["unit"], d.get("status")],
                  fmts=[None, None, None, None, "#,##0.000"], wrap_cols=(3,))
    tr = b.sheet("TRACE", [8, 22, 10, 22, 22, 60, 40, 30])
    b.title(tr, f"{trade.replace('_', ' ')} - TRACEABILITY", f"{trade_ar} - التتبع")
    b.header(tr, ["LINE", "CODE", "LEVEL", "TECH CLASS", "COMM CLASS", "AUTHORITY", "SUPERSEDES (V3a)", "SUMMARY ITEM"])
    for x in lines:
        b.row(tr, [x["line_id"], x["code"], x["level"], x["release"]["technical"]["class"],
                   x["release"]["commercial"]["class"], x["authority"], ", ".join(x.get("supersedes") or [])[:200],
                   x.get("sumrow") or "-"], status_col=3, wrap_cols=(5, 6))
    if rooms is not None:
        ro = b.sheet("ROOMS", [10, 7, 26, 12, 12, 10, 8, 14, 12, 14, 30, 26, 8, 10, 12])
        b.title(ro, "FLOORING - ONE ROW PER ROOM", "الأرضيات - غرفة غرفة",
                "PRIMARY = topology polygon; CHECK = DXF DIMENSION chain (same drawing: calculation-independent, not "
                "source-independent); FINAL NET = the primary (never two payable quantities)")
        b.header(ro, ["ROOM", "FLOOR", "NAME", "PRIMARY m²", "CHECK m²", "DIFF m²", "DIFF %", "RECON", "FINAL NET m²",
                      "MATERIAL", "AUTHORITY", "WASTE METHOD", "WASTE %", "WASTE m²", "PROCUREMENT m²"])
        for x in rooms:
            rr = ro._row
            chk = x.get("check")
            b.row(ro, [x["room"], x["floor"], x["name"], x["primary"], chk,
                       {"f": f'=IF(E{rr}="","",D{rr}-E{rr})', "expect": "" if chk is None else _rnd(x["primary"] - chk)},
                       {"f": f'=IF(OR(E{rr}="",E{rr}=0),"",(D{rr}-E{rr})/E{rr})',
                        "expect": "" if not chk else _rnd((x["primary"] - chk) / chk)},
                       x["recon"], x["final"], x["material"], x["authority"], x["waste_method"], x["waste_pct"],
                       x["waste_qty"], x["procurement"]],
                  kinds=[None, None, None, "qty", "qty"], fmts=[None, None, None, "#,##0.00", "#,##0.00", "#,##0.00",
                                                              "0.00%", None, "#,##0.00"], wrap_cols=(2, 10, 11))
    return b.save()


def _status(r):
    t, c = r["technical"]["total"], r["commercial"]["total"]
    if t and abs(t - c) < 1e-6 and not r["commercial"]["budget"]:
        return "COMPUTED"
    if c or r["commercial"]["budget"]:
        return "PARTIAL" if t else "PROVISIONAL_SOURCE_DERIVED"
    return "PENDING"


def master_book(model, out_dir, meta, extra) -> dict:
    b = Book(out_dir / "00_URBAN_BOQ_MASTER_SUMMARY.xlsx", project=meta["project"], phase=meta["phase"])
    ws = b.sheet("MASTER_SUMMARY", [26, 44, 28, 7] + [12] * 5 + [3] + [12] * 5 + [13, 12, 13, 14])
    b.title(ws, "URBAN BOQ - MASTER SUMMARY (A: NET / TECHNICAL | B: COMMERCIAL)", "ملخص جدول الكميات - فني / تجاري",
            f"{meta['project']} | {meta['phase']} | {model['release_label']} | values are the frozen engine totals; "
            f"TOTAL = SUM of the levels")
    b.header(ws, ["TRADE", "SUMMARY ITEM", "البند", "UNIT", "A GF", "A 1F", "A 2F+R", "A OTH", "A TOTAL", "", "B GF", "B 1F",
                  "B 2F+R", "B OTH", "B TOTAL", "OF WHICH PROV.", "BUDGET (beside)", "PROCURE ELIG.", "STATUS"])
    trades = {t[1]: (t[0], t[2]) for t in model["trades"]}
    cur = None
    for r in model["master"]["rows"]:
        if r["trade"] != cur:
            cur = r["trade"]
            b.band(ws, f"{trades[cur][0]}  {cur.replace('_', ' ')}   |   {trades[cur][1]}")
        u, f = du(r["unit"])
        rr = ws._row
        A = [_s(r["technical"][lv], f) for lv in model["levels"]]
        B = [_s(r["commercial"][lv], f) for lv in model["levels"]]
        c = r["commercial"]
        b.row(ws, [r["trade"], r["item"], r["item_ar"], u] + A +
              [{"f": f"=SUM(E{rr}:H{rr})", "expect": _rnd(r["technical"]["total"] * f)}, ""] + B +
              [{"f": f"=SUM(K{rr}:N{rr})", "expect": _rnd(c["total"] * f)}, _s(c["provisional_part"], f),
               _s(c["budget"], f), _s(c["procurement_eligible"], f), _status(r)],
              kinds=[None] * 4 + ["qty"] * 4 + [None, None] + ["qty"] * 4,
              fmts=[None] * 4 + [_q(u)] * 5 + [None] + [_q(u)] * 8, status_col=18, ar_cols=(2,))
    b.skip(ws)
    b.note(ws, model["master"]["rule"])
    b.note(ws, "Units: concrete m³; reinforcement t (engine kg / 1000); areas m²; lengths lm; counts No. Rebar net (straight "
               "+ hooks) and rebar purchased from 12 m stock (BBS cut) are separate bases, never added together.")
    cov = b.sheet("COVERAGE", [26, 22, 8, 14, 14, 12, 40, 70])
    b.title(cov, "COVERAGE BY TRADE (TECHNICAL vs COMMERCIAL)", "التغطية حسب البند")
    b.header(cov, ["TRADE", "البند", "ITEMS", "TECHNICAL %", "COMMERCIAL %", "PROVISIONAL", "TECHNICAL CLASSES",
                   "NOT RELEASED (commercial)"])
    for r in extra["coverage"]["rows"]:
        b.row(cov, [r["trade"], r["trade_ar"], r["items"], r["technical_coverage_pct"], r["commercial_coverage_pct"],
                    r["provisional_items"], ", ".join(f"{k} {v}" for k, v in r["technical_classes"].items()),
                    ", ".join(r["not_released"])], ar_cols=(1,), wrap_cols=(6, 7))
    b.note(cov, extra["coverage"]["rule"])
    bl = b.sheet("BLOCKERS", [6, 30, 70, 40, 50])
    b.title(bl, "BLOCKER RESOLUTION REGISTER", "سجل حل المعوقات")
    b.header(bl, ["ID", "BLOCKER", "ROUTES", "RELEASE CLASS", "REMAINING RISK"])
    for r in extra["blockers"]["rows"]:
        b.row(bl, [r["id"], r["blocker"], r["routes"], r["release_class"], r["remaining_risk"]], wrap_cols=(1, 2, 3, 4))
    ex = b.sheet("EXPECTED_SCOPE", [14, 40, 8, 16, 16, 60])
    b.title(ex, "EXPECTED SCOPE (item names only)", "النطاق المتوقع")
    b.header(ex, ["GROUP", "ITEM", "LINES", "TECHNICAL", "COMMERCIAL", "SCOPE FILTER"])
    for r in extra["expected"]["rows"]:
        b.row(ex, [r["group"], r["item"], r["urban_lines"], r["technical"], r["commercial"], r["scope_filter"]],
              wrap_cols=(5,))
    om = b.sheet("OWNER_METHODS", [12, 36, 26, 90, 24])
    b.title(om, "OWNER METHOD REGISTER", "سجل طرق المالك")
    b.header(om, ["ID", "METHOD", "TOPIC", "DECISION / RULE", "PROMOTION CLASS"])
    for r in extra["owner"]["rows"]:
        b.row(om, [r["id"], r["method_id"], r["topic"], r["value"], r["promotion_class"]], wrap_cols=(2, 3))
    dm = b.sheet("DUAL_MEASUREMENT", [30, 34, 40, 30, 16, 26, 50])
    b.title(dm, "DUAL MEASUREMENT", "القياس المزدوج")
    b.header(dm, ["ITEM", "ROUTE A", "ROUTE B", "ROUTE C", "CALC. INDEP.", "SOURCE INDEP.", "RESULT"])
    for r in extra["dual"]["rows"]:
        b.row(dm, [r["item"], r["route_a"], r["route_b"], r["route_c"], r["calculation_independence"],
                   r["source_independence"], r["result"]], wrap_cols=(1, 2, 3, 6))
    rp = b.sheet("REBAR_POPULATION", [22, 8, 8, 13, 13, 12, 12, 12, 12, 9, 50])
    b.title(rp, "REBAR POPULATION REGISTER (t)", "سجل مجموعات التسليح")
    b.header(rp, ["POPULATION", "SETS", "BLOCKED", "TECHNICAL t", "NET t", "HOOKS t", "LAPS t", "LOW t", "HIGH t",
                  "SHARE %", "CLASSES"])
    for r in extra["rebar"]["rows"]:
        b.row(rp, [r["population"], r["sets"], r["blocked"], r["technical_kg"] / 1000, r["net_kg"] / 1000,
                   r["hook_kg"] / 1000, r["lap_kg"] / 1000, r["low_kg"] / 1000, r["high_kg"] / 1000,
                   r["weight_share_pct"], ", ".join(r["classes"])],
              fmts=[None, None, None] + ["#,##0.000"] * 6, wrap_cols=(10,))
    return b.save()


def recon_book(model, out_dir, meta) -> dict:
    b = Book(out_dir / "10_QS_RECONCILIATION.xlsx", project=meta["project"], phase=meta["phase"])
    p = b.sheet("PARAMETERS", [40, 14, 60])
    b.title(p, "RECONCILIATION PARAMETERS", "معايير المطابقة")
    b.header(p, ["PARAMETER", "VALUE", "MEANING"], repeat=False)
    m_row = b.row(p, ["MATCH tolerance", {"input": 0.02}, "|difference| <= tolerance x Urban value -> MATCH"],
                  fmts=[None, "0.0%"])
    c_row = b.row(p, ["CLOSE tolerance", {"input": 0.10}, "<= this -> CLOSE; above -> REVIEW"], fmts=[None, "0.0%"])
    for t in model["trades"]:
        rows = [r for r in model["master"]["rows"] if r["trade"] == t[1]]
        if not rows:
            continue
        ws = b.sheet(f"{t[0]}_{t[1][:20]}", [44, 30, 8, 14, 14, 14, 14, 10, 16])
        b.title(ws, f"QS RECONCILIATION - {t[1].replace('_', ' ')}", f"مطابقة الكميات - {t[2]}",
                "enter the manual / check quantity in the yellow cells IN THE UNIT SHOWN; difference against the "
                "commercial value, % and status are formulas")
        b.header(ws, ["SUMMARY ITEM", "البند", "UNIT", "URBAN TECHNICAL", "URBAN COMMERCIAL", "MANUAL / CHECK",
                      "DIFFERENCE", "%", "STATUS"])
        for r in rows:
            u, f = du(r["unit"])
            rr = ws._row
            b.row(ws, [r["item"], r["item_ar"], u, _s(r["technical"]["total"], f), _s(r["commercial"]["total"], f),
                       {"input": None}, {"f": f'=IF(F{rr}="","",F{rr}-E{rr})', "expect": ""},
                       {"f": f'=IF(OR(F{rr}="",E{rr}=0),"",(F{rr}-E{rr})/E{rr})', "expect": ""},
                       {"f": f'=IF(F{rr}="","NOT CHECKED",IF(E{rr}=0,"REVIEW",IF(ABS(H{rr})<=PARAMETERS!$B${m_row},"MATCH",'
                             f'IF(ABS(H{rr})<=PARAMETERS!$B${c_row},"CLOSE","REVIEW"))))', "expect": "NOT CHECKED"}],
                  kinds=[None, None, None, "qty", "qty"], fmts=[None, None, None, _q(u), _q(u), _q(u), _q(u), "0.0%"],
                  ar_cols=(1,))
    return b.save()
