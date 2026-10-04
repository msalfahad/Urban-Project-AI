"""REPORTING V3 - openpyxl workbook writer in the agreed mock-up style.

Book(path) collects a cell map while it writes: every engine quantity ("value"), every formula ("formula", with the
value the model expects) and every manual-check input ("input"). save() writes the file, re-packs the zip with fixed
entry times (same content -> same bytes) and returns {"file_sha256", "cell_map"}.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

NAVY, GREEN, LIGHT_BLUE, HEAD, GRID, WHITE = "1F3A5F", "548235", "DDEBF7", "DCE3EC", "C3CBD5", "FFFFFF"
STATUS_FILL = {"COMPUTED": ("C6EFCE", "006100"), "PARTIAL": ("FFEB9C", "7F6000"), "REVIEW": ("FFEB9C", "7F6000"),
               "BLOCKED": ("FFC7CE", "9C0006"), "NOT_IN_SOURCE": ("EDEDED", "555555"), "INFO": ("EDEDED", "555555")}
STATUS_AR = {"COMPUTED": "محسوب", "PARTIAL": "جزئي", "REVIEW": "مراجعة", "BLOCKED": "متوقف", "NOT_IN_SOURCE": "غير موجود بالمصدر"}
INPUT_FILL = "FFF9C4"
FONT = "Arial"
CREATED = datetime(2026, 10, 4)
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
QTY_FMT = {"m3": "#,##0.000", "m2": "#,##0.00", "lm": "#,##0.00", "m": "#,##0.000", "nr": "#,##0", "kg": "#,##0.0"}
THIN = Side(style="thin", color=GRID)
THICK = Side(style="thick", color=NAVY)


def _fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


class Book:
    def __init__(self, path, *, project, phase):
        self.path = Path(path)
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.wb.properties.creator = "Urban Projects"
        self.wb.properties.created = CREATED
        self.wb.properties.modified = CREATED
        self.project, self.phase = project, phase
        self.cell_map = []

    # ------------------------------------------------------------------ sheets
    def sheet(self, name, widths, *, landscape=True):
        ws = self.wb.create_sheet(name)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.6)
        ws.oddFooter.left.text = f"{self.project} - {self.phase}"
        ws.oddFooter.center.text = "Page &P of &N"
        ws.oddFooter.right.text = "Urban Projects"
        ws._row = 1
        ws._ncol = len(widths)
        return ws

    def title(self, ws, en, ar, sub=None):
        r = ws._row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ws._ncol)
        c = ws.cell(r, 1, f"{en}   |   {ar}")
        c.font = Font(name=FONT, bold=True, size=14, color=WHITE)
        c.fill = _fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 26
        r += 1
        if sub:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ws._ncol)
            c = ws.cell(r, 1, sub)
            c.font = Font(name=FONT, italic=True, size=9, color="333333")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.row_dimensions[r].height = 28
            r += 1
        for col in range(1, ws._ncol + 1):
            ws.cell(r, col).border = Border(bottom=THICK)
        ws.row_dimensions[r].height = 4
        ws._row = r + 2

    def header(self, ws, cols, repeat=True):
        r = ws._row
        for i, v in enumerate(cols, 1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT, bold=True, size=9, color=WHITE)
            c.fill = _fill(NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.row_dimensions[r].height = 30
        if repeat:
            ws.print_title_rows = f"{r}:{r}"
            ws.freeze_panes = ws.cell(r + 1, 1)
        ws._row = r + 1
        return r

    def band(self, ws, text, color=GREEN, size=10):
        r = ws._row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ws._ncol)
        c = ws.cell(r, 1, text)
        c.font = Font(name=FONT, bold=True, size=size, color=WHITE)
        c.fill = _fill(color)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 18
        ws._row = r + 1
        return r

    def row(self, ws, values, *, kinds=None, fmts=None, status_col=None, fill=None, bold=False, ar_cols=(), wrap_cols=(),
            expect=None, height=None):
        """values: list; a dict {"f": formula, "expect": value} writes a formula; {"input": v} a manual input cell."""
        r = ws._row
        for i, v in enumerate(values, 1):
            c = ws.cell(r, i)
            kind = None
            if isinstance(v, dict) and "f" in v:
                c.value = v["f"]
                kind = "formula"
                self.cell_map.append({"sheet": ws.title, "cell": c.coordinate, "kind": "formula", "formula": v["f"],
                                      "expected": v.get("expect")})
                c.fill = _fill(LIGHT_BLUE)
            elif isinstance(v, dict) and "input" in v:
                c.value = v["input"]
                kind = "input"
                c.fill = _fill(INPUT_FILL)
                self.cell_map.append({"sheet": ws.title, "cell": c.coordinate, "kind": "input", "expected": v["input"]})
            else:
                c.value = v
                if kinds and i - 1 < len(kinds) and kinds[i - 1] == "qty" and v is not None:
                    kind = "value"
                    self.cell_map.append({"sheet": ws.title, "cell": c.coordinate, "kind": "value", "expected": v})
            f = fmts[i - 1] if fmts and i - 1 < len(fmts) else None
            if f:
                c.number_format = f
            c.font = Font(name=FONT, size=9, bold=bold)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            al = Alignment(vertical="top", wrap_text=(i - 1) in wrap_cols or (i - 1) in ar_cols)
            if (i - 1) in ar_cols:
                al = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
            c.alignment = al
            if fill and kind != "input":
                c.fill = _fill(fill)
        if status_col is not None:
            st = values[status_col]
            if st in STATUS_FILL:
                bg, fg = STATUS_FILL[st]
                c = ws.cell(r, status_col + 1)
                c.fill = _fill(bg)
                c.font = Font(name=FONT, size=9, bold=True, color=fg)
                c.alignment = Alignment(horizontal="center", vertical="top")
        if height:
            ws.row_dimensions[r].height = height
        ws._row = r + 1
        return r

    def divider(self, ws):
        r = ws._row
        for col in range(1, ws._ncol + 1):
            ws.cell(r, col).border = Border(top=THICK)
        ws._row = r + 1

    def skip(self, ws, n=1):
        ws._row += n

    def note(self, ws, text, italic=True):
        r = ws._row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ws._ncol)
        c = ws.cell(r, 1, text)
        c.font = Font(name=FONT, size=8, italic=italic, color="444444")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = max(14, 12 * (1 + len(text) // 160))
        ws._row = r + 1

    # ------------------------------------------------------------------ save
    def save(self) -> dict:
        self.wb.save(self.path)
        _repack(self.path)
        data = self.path.read_bytes()
        return {"file": self.path.name, "file_sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data),
                "cell_map": self.cell_map}


def _repack(path):
    src = zipfile.ZipFile(io.BytesIO(Path(path).read_bytes()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                txt = data.decode("utf-8")
                txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                             lambda m: m.group(1) + CREATED.strftime("%Y-%m-%dT%H:%M:%SZ") + m.group(2), txt)
                txt = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)",
                             lambda m: m.group(1) + CREATED.strftime("%Y-%m-%dT%H:%M:%SZ") + m.group(2), txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            dst.writestr(zi, data)
    Path(path).write_bytes(buf.getvalue())
