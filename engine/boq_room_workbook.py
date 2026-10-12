"""ROOM-BY-ROOM BOQ WORKBOOK (M9 S1) - a freelancer-style view over canonical lines and material results.

    write(path, *, room_boq=None, structural_lines=(), census_rows=(), meta, created) -> {"path", "sheets", "sha256"}
    expected_totals(room_boq)                                                         -> the Python totals the
                                                                                         formulas must reproduce

Trade sheets carry FLOOR | ROOM (AR/EN) | ITEM | MATERIAL SIZE | NET M² | WASTE % | WASTE M² | MATERIAL M² |
BOXES/ROLLS | SOURCE | STATUS, then FORMULA | METHOD | UNIT | PROCUREMENT and the traceability columns (joint,
overlap, reusable offcut, packaging surplus, units consumed, unit area, balance check, line id, concerns, row type).
Room, floor and project subtotals are formulas (SUMIFS over the DETAIL / MATERIAL rows, so a subtotal is never summed
twice); WASTE % = WASTE M² / NET M²; MATERIAL M² = units consumed x unit area (- untouched roll tails for a membrane);
the balance check must read 0. A trade the S1 engine does not measure lists its families with their run state and
reason, and no number. No price, rate or amount anywhere; no kg/m3 ratio. Lives outside engine/source (openpyxl).
"""

from __future__ import annotations

import hashlib
from collections import OrderedDict

from engine.boq_rc1_xlsx import _repack

POLICY_ID = "URBAN_ROOM_BOQ_XLSX_V1"
FONT = "Arial"
BANNER = ("NOT QUOTATION-READY - machine quantities (MACHINE_ONLY), gate INCOMPLETE; no prices; synthetic / input "
          "geometry only where stated")
HEAD = ["FLOOR", "ROOM (AR/EN)", "ITEM", "MATERIAL SIZE", "NET M²", "WASTE %", "WASTE M²", "MATERIAL M²",
        "BOXES/ROLLS", "SOURCE", "STATUS", "FORMULA", "METHOD", "UNIT", "PROCUREMENT", "JOINT M²", "OVERLAP M²",
        "REUSABLE OFFCUT M²", "PACKAGING SURPLUS M²", "UNITS CONSUMED", "UNIT AREA M²", "BALANCE CHECK M²", "LINE ID",
        "CONCERNS", "ROOM ID", "ROW TYPE"]
C = {h: i + 1 for i, h in enumerate(HEAD)}                     # 1-based column numbers
SUMMED = ("NET M²", "WASTE M²", "MATERIAL M²", "JOINT M²", "OVERLAP M²", "REUSABLE OFFCUT M²", "PACKAGING SURPLUS M²")
FIRST = 4                                                       # first data row (banner, note, header on row 3)
TRADE_SHEETS = OrderedDict([
    ("WALL PORCELAIN", ("B4-01", "B4-02", "B4-09", "B4-10")),
    ("PLASTER & PAINT", ("B4-03", "B4-04", "B4-05", "B4-06", "B4-07", "B4-08")),
    ("CEILINGS", ("B5-01", "B5-02", "B5-03", "B5-04", "B5-05")),
    ("SKIRTING", ("B3-06", "B3-07", "B7-07", "B7-08")),
    ("MARBLE", ("B3-03", "B3-04", "B7-02", "B7-03", "B7-04", "B7-05", "B7-06", "B6-06", "B6-07")),
    ("STAIRS & RAILINGS", ("B7-01", "B7-09", "B7-10", "B7-11", "B7-13", "B7-14")),
    ("DOORS & WINDOWS", ("B6-01", "B6-02", "B6-03", "B6-04", "B6-05")),
    ("ALUMINIUM", ("B6-04", "B6-05", "B13-07")),
])


def _L(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def _q(name):
    return "'" + name.replace("'", "''") + "'"


def _room_label(loc):
    ar, en = loc.get("room_name_ar") or loc.get("name_ar"), loc.get("room_name_en") or loc.get("name_en")
    return " / ".join(x for x in (ar, en) if x) or (loc.get("room_id") or "")


QUANTITY_STATES = ("MEASURED", "SOURCE_DERIVED", "PROVISIONAL")     # boq_line.QUANTITY_STATES (no engine.source import:
                                                                   # the workbook is a view over the lines it is given)


def _item_label(ln):
    return f"{ln['item_ar']} / {ln['item_en']}"


def _size(spec):
    if not spec:
        return "NOT ESTABLISHED"
    if spec.get("kind") == "TILE" and spec.get("length_m"):
        return f"{spec['length_m'] * 100:g}×{spec['width_m'] * 100:g} cm ({spec['spec_id']})"
    if spec.get("kind") == "MEMBRANE_SHEET" and spec.get("roll_width_m"):
        return (f"roll {spec['roll_width_m']:.2f}×{spec['roll_length_m']:.2f} m, lap {spec['min_side_lap_m']:.2f} m "
                f"({spec['spec_id']})")
    return spec.get("spec_id") or ""


class _Sheet:
    def __init__(self, wb, name, note):
        from openpyxl.styles import Alignment, Font, PatternFill
        self.Font, self.Fill, self.Al = Font, PatternFill, Alignment
        self.ws = wb.create_sheet(name)
        self.name = name
        self.ws.append([BANNER])
        self.ws["A1"].font = Font(name=FONT, bold=True, color="C00000")
        self.ws.append([note])
        self.ws["A2"].font = Font(name=FONT, italic=True, color="595959")
        self.ws.append(HEAD)
        for c in self.ws[3]:
            c.font = Font(name=FONT, bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        self.r = 3

    def row(self, vals, kind, bold=False, fill=None):
        self.r += 1
        for h, v in vals.items():
            self.ws.cell(self.r, C[h], v)
        self.ws.cell(self.r, C["ROW TYPE"], kind)
        for j in range(1, len(HEAD) + 1):
            c = self.ws.cell(self.r, j)
            formula = isinstance(c.value, str) and c.value.startswith("=")
            c.font = self.Font(name=FONT, bold=bold, color="000000" if formula or not isinstance(c.value, (int, float))
                               else "0000FF")
            if HEAD[j - 1].endswith("M²"):
                c.number_format = "0.000"
            if HEAD[j - 1] == "WASTE %":
                c.number_format = "0.00%"
            if fill:
                c.fill = self.Fill("solid", fgColor=fill)
        return self.r

    def finish(self):
        widths = {"FLOOR": 10, "ROOM (AR/EN)": 22, "ITEM": 30, "MATERIAL SIZE": 26, "SOURCE": 26, "STATUS": 22,
                  "FORMULA": 40, "METHOD": 26, "CONCERNS": 60, "LINE ID": 24, "BOXES/ROLLS": 16, "PROCUREMENT": 18}
        for h, j in C.items():
            self.ws.column_dimensions[_L(j)].width = widths.get(h, 12)
        self.ws.freeze_panes = "C4"
        self.ws.auto_filter.ref = f"A3:{_L(len(HEAD))}{max(self.r, 4)}"

    # ---- formulas
    def rng(self, h):
        return f"${_L(C[h])}${FIRST}:${_L(C[h])}${self.r_last}"

    def sumifs(self, h, crit):
        """SUMIFS over DETAIL + MATERIAL rows with extra criteria [(col header, value)]."""
        parts = []
        for kind in ("DETAIL", "MATERIAL"):
            cr = "".join(f',{self.rng(ch)},"{v}"' for ch, v in crit)
            parts.append(f'SUMIFS({self.rng(h)},{self.rng("ROW TYPE")},"{kind}"{cr})')
        return "=" + "+".join(parts)


def _subtotals(sh, row_no, crit, label, kind):
    """a subtotal is a number only when every row it covers is resolved; otherwise the text
    "PARTIAL: <sum> (<n> UNRESOLVED)" - a missing net or material is never read as zero."""
    vals = {"FLOOR": label[0], "ROOM (AR/EN)": label[1], "ITEM": label[2]}
    box = sh.rng("BOXES/ROLLS")
    cr = "".join(f',{sh.rng(ch)},"{v}"' for ch, v in crit)
    open_net = f'COUNTIFS({sh.rng("ROW TYPE")},"DETAIL",{sh.rng("STATUS")},"BLOCKED_RFI*"{cr})'
    open_mat = f'COUNTIFS({box},"UNRESOLVED (material*"{cr})'
    for h in SUMMED:
        total = sh.sumifs(h, crit)[1:]
        cnt = open_net if h == "NET M²" else open_mat
        vals[h] = f'=IF({cnt}>0,"PARTIAL: "&TEXT({total},"0.000")&" ("&{cnt}&" UNRESOLVED)",{total})'
    n, w = _L(C["NET M²"]), _L(C["WASTE M²"])
    vals["WASTE %"] = (f'=IF(AND(ISNUMBER({n}{row_no}),ISNUMBER({w}{row_no})),IF({n}{row_no}>0,{w}{row_no}/{n}{row_no},""),'
                       f'"PARTIAL")')
    # whole packs: a subtotal is a number only when every room's pack count is resolved
    unres = f'COUNTIFS({box},"UNRESOLVED*"{cr})'
    tot = "+".join(f'SUMIFS({box},{sh.rng("ROW TYPE")},"{k}"{cr})' for k in ("DETAIL", "MATERIAL"))
    vals["BOXES/ROLLS"] = f'=IF({unres}>0,"INCOMPLETE: "&{unres}&" UNRESOLVED",{tot})'
    vals["STATUS"] = kind
    return vals


def _row_formulas(sh, r, vals):
    n, w = _L(C["NET M²"]), _L(C["WASTE M²"])
    ws = sh.ws
    if vals.get("_material_kind") == "TILE":
        u, a = _L(C["UNITS CONSUMED"]), _L(C["UNIT AREA M²"])
        j, o = _L(C["JOINT M²"]), _L(C["REUSABLE OFFCUT M²"])
        ws.cell(r, C["MATERIAL M²"], f"={u}{r}*{a}{r}")
        ws.cell(r, C["BALANCE CHECK M²"], f"=ROUND({_L(C['MATERIAL M²'])}{r}-({n}{r}-{j}{r}+{w}{r}+{o}{r}),6)")
        ws.cell(r, C["WASTE %"], f"=IF(AND(ISNUMBER({n}{r}),ISNUMBER({w}{r})),IF({n}{r}>0,{w}{r}/{n}{r},\"\"),\"\")")
    elif vals.get("_material_kind") == "MEMBRANE":
        u, a, s = _L(C["UNITS CONSUMED"]), _L(C["UNIT AREA M²"]), _L(C["PACKAGING SURPLUS M²"])
        o, ov = _L(C["REUSABLE OFFCUT M²"]), _L(C["OVERLAP M²"])
        first, last = vals["_comp_rows"]
        m = _L(C["MATERIAL M²"])
        ws.cell(r, C["MATERIAL M²"], f"={u}{r}*{a}{r}-{s}{r}")
        ws.cell(r, C["BALANCE CHECK M²"],
                f"=ROUND({m}{r}-(SUM({n}{first}:{n}{last})+SUM({ov}{first}:{ov}{last})+{ov}{r}+{w}{r}+{o}{r}),6)")
        ws.cell(r, C["WASTE %"], f"=IF(SUM({n}{first}:{n}{last})>0,{w}{r}/SUM({n}{first}:{n}{last}),\"\")")
        ws.cell(r, C["BOXES/ROLLS"], f"={u}{r}")
    for c in (C["MATERIAL M²"], C["BALANCE CHECK M²"], C["WASTE %"]):
        ws.cell(r, c).font = sh.Font(name=FONT, color="000000")
        ws.cell(r, c).number_format = "0.00%" if c == C["WASTE %"] else "0.000"


def _status(line, res=None):
    s = f"{line['evidence_state']} / {line['approval_state']}"
    if res is not None:
        s += f" / material {res['state']}"
    return s


def _floor_rows(room_boq):
    """FLOOR PORCELAIN groups: one DETAIL row per room floor-finish line."""
    mats = {m["line_id"]: m for m in room_boq["materials"] if m.get("line_id")}
    groups = OrderedDict()
    for ln in room_boq["lines"]:
        if ln["item_code"] not in ("FIN.FLOOR.PORCELAIN", "FIN.FLOOR.WET_TILE"):
            continue
        res = mats.get(ln["line_id"])
        if res is not None and res.get("method", "").startswith("MEMBRANE"):
            res = None
        loc = ln["location"]
        v = {"FLOOR": loc["floor"], "ROOM (AR/EN)": _room_label(loc), "ITEM": _item_label(ln),
             "NET M²": ln["measured_net"], "SOURCE": f"{ln['source']['geometry_source']}; {ln['rule']['rule_id']}@{ln['rule']['version']}",
             "STATUS": _status(ln, res), "FORMULA": ln["calculation"]["formula"], "UNIT": ln["unit"], "LINE ID": ln["line_id"],
             "ROOM ID": loc["room_id"]}
        if ln["evidence_state"] == "BLOCKED_RFI":
            v.update({"MATERIAL SIZE": _size((res or {}).get("material_spec")), "CONCERNS": ln["rfi"]["question"],
                      "METHOD": "-", "BOXES/ROLLS": "UNRESOLVED (material: no geometry)", "PROCUREMENT": "UNRESOLVED"})
        elif res is None or res["state"] == "BLOCKED":
            v.update({"MATERIAL SIZE": _size((res or {}).get("material_spec")),
                      "METHOD": (res or {}).get("method", "-"), "CONCERNS": (res or {}).get("blocked_by") or "",
                      "BOXES/ROLLS": "UNRESOLVED (material blocked)", "PROCUREMENT": "UNRESOLVED"})
        else:
            p, lay = res["packaging"], res["layout"]
            v.update({"MATERIAL SIZE": _size(res["material_spec"]),
                      "METHOD": f"{res['method']}: {lay['orientation']}deg {lay['align_x']}/{lay['align_y']} "
                                f"({lay['full_tiles']} full + {lay['cut_pieces']} cut)",
                      "WASTE M²": res["cutting_waste_m2"], "JOINT M²": res["joint_m2"], "OVERLAP M²": 0.0,
                      "REUSABLE OFFCUT M²": res["reusable_offcuts_m2"], "UNITS CONSUMED": res["units"]["consumed_units"],
                      "UNIT AREA M²": res["units"]["unit_area_m2"],
                      "PACKAGING SURPLUS M²": p["packaging_surplus_m2"],
                      "BOXES/ROLLS": p["packs"] if p["state"] == "RESOLVED" else "UNRESOLVED (no box data)",
                      "PROCUREMENT": (f"{p['procurement_packs']} boxes" if p["state"] == "RESOLVED"
                                      else f"{p['procurement_units']} pcs (boxes UNRESOLVED)"),
                      "CONCERNS": "; ".join(res["concerns"]), "_material_kind": "TILE"})
        groups.setdefault(loc["floor"], OrderedDict()).setdefault(loc["room_id"], []).append((v, "DETAIL"))
    return groups


def _wp_rows(room_boq):
    groups = OrderedDict()
    mats = [m for m in room_boq["materials"] if m.get("method", "").startswith("MEMBRANE")]
    by_room = {}
    for m in mats:
        by_room[(m.get("location") or {}).get("room_id")] = m
    for ln in room_boq["lines"]:
        if not ln["item_code"].startswith("FIN.WP."):
            continue
        loc = ln["location"]
        res = by_room.get(loc["room_id"])
        comp = None
        if res and res["state"] != "BLOCKED":
            comp = next((c for c in res["components"] if ln["line_id"] in c.get("line_ids", ())), None)
        v = {"_item_code": ln["item_code"], "FLOOR": loc["floor"], "ROOM (AR/EN)": _room_label(loc),
             "ITEM": _item_label(ln) + (f" ({ln['physical_element']})" if ln["item_code"] == "FIN.WP.DOORWAY" else ""),
             "MATERIAL SIZE": _size((res or {}).get("material_spec")), "NET M²": ln["measured_net"],
             "SOURCE": f"{ln['source']['geometry_source']}; {ln['rule']['rule_id']}@{ln['rule']['version']}",
             "STATUS": _status(ln), "FORMULA": ln["calculation"]["formula"], "UNIT": ln["unit"],
             "LINE ID": ln["line_id"], "ROOM ID": loc["room_id"], "CONCERNS": ln.get("notes") or
             ((ln.get("rfi") or {}).get("question") or ""), "METHOD": "net surface"}
        if comp is not None and ln["item_code"] != "FIN.WP.DOORWAY":
            v["OVERLAP M²"] = comp["overlap_m2"]
        if comp is not None and ln["item_code"] == "FIN.WP.DOORWAY":
            n_d = len([x for x in room_boq["lines"] if x["location"]["room_id"] == loc["room_id"]
                       and x["item_code"] == "FIN.WP.DOORWAY"])
            v["OVERLAP M²"] = round(comp["overlap_m2"] / n_d, 6) if n_d == 1 else None
            if n_d > 1:
                v["CONCERNS"] = (v["CONCERNS"] + "; doorway laps shown on the material row").strip("; ")
        groups.setdefault(loc["floor"], OrderedDict()).setdefault(loc["room_id"], []).append((v, "DETAIL"))
    for (rid, res) in by_room.items():
        fl = (res.get("location") or {}).get("floor")
        rows = groups.get(fl, {}).get(rid)
        if rows is None:
            continue
        loc = res.get("location") or {}
        v = {"FLOOR": fl, "ROOM (AR/EN)": rows[0][0]["ROOM (AR/EN)"], "ITEM": "رولات عازل / Membrane rolls (room cutting plan)",
             "MATERIAL SIZE": _size(res.get("material_spec")), "UNIT": "roll", "ROOM ID": rid,
             "SOURCE": "; ".join(f"{x['rule_id']}@{x['version']}" for x in res.get("rule_refs", ())) or "-"}
        if res["state"] == "BLOCKED":
            v.update({"STATUS": "material BLOCKED", "CONCERNS": res["blocked_by"], "BOXES/ROLLS": "UNRESOLVED (material blocked)",
                      "PROCUREMENT": "UNRESOLVED", "METHOD": res["method"]})
            rows.append((v, "MATERIAL"))
            continue
        lay, p = res["layout"], res["packaging"]
        dw_ov = sum(c["overlap_m2"] for c in res["components"] if c["component"] == "doorway")
        n_d = len([1 for vv, _ in rows if vv.get("_item_code") == "FIN.WP.DOORWAY"])
        v.update({"STATUS": f"material {res['state']}", "WASTE M²": res["cutting_waste_m2"],
                  "REUSABLE OFFCUT M²": res["reusable_offcuts_m2"], "UNITS CONSUMED": res["units"]["consumed_units"],
                  "UNIT AREA M²": res["units"]["unit_area_m2"], "PACKAGING SURPLUS M²": p["packaging_surplus_m2"],
                  "OVERLAP M²": round(dw_ov, 6) if n_d > 1 else None,
                  "PROCUREMENT": f"{p['procurement_packs']} rolls (spare {p['spare']['state']})",
                  "METHOD": f"{res['method']}: strips {lay['direction']} x{lay['floor_strips']}, band {lay['band_width_m']:g} m, "
                            f"{lay['seams']['upturn_end_laps']['count']} band joints",
                  "FORMULA": "rolls x roll area - untouched roll tails; waste = consumed - net - overlap - offcuts",
                  "CONCERNS": "; ".join(res["concerns"]), "_material_kind": "MEMBRANE"})
        rows.append((v, "MATERIAL"))
    return groups


def _info(wb, name, lines):
    from openpyxl.styles import Alignment, Font
    ws = wb.create_sheet(name)
    ws.append([BANNER])
    ws["A1"].font = Font(name=FONT, bold=True, color="C00000")
    for ln in lines:
        ws.append(ln if isinstance(ln, (list, tuple)) else [ln])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, bold=bool(c.value) and isinstance(c.value, str) and c.value.isupper())
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 28
    return ws


def _table(wb, name, note, header, rows, fmts=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(name)
    ws.append([BANNER])
    ws["A1"].font = Font(name=FONT, bold=True, color="C00000")
    ws.append([note])
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    ws.append(header)
    for c in ws[3]:
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=4):
        for c in row:
            formula = isinstance(c.value, str) and c.value.startswith("=")
            c.font = Font(name=FONT, bold=isinstance(c.value, str) and c.value.endswith("TOTAL"),
                          color="000000" if formula or not isinstance(c.value, (int, float)) else "0000FF")
            f = (fmts or {}).get(c.column)
            if f:
                c.number_format = f
    for j in range(1, len(header) + 1):
        ws.column_dimensions[_L(j)].width = 16 if j > 1 else 22
    ws.freeze_panes = "B4"
    if rows:
        ws.auto_filter.ref = f"A3:{_L(len(header))}{3 + len(rows)}"
    return ws


def _layout_membrane(sh, groups, detail_rows):
    """lay out DETAIL / MATERIAL rows grouped by floor and room, then room, floor and project subtotal rows; a
    membrane MATERIAL row is given the row range of its room's surface rows (its balance check)."""
    plan = []
    for fl, rooms in groups.items():
        for rid, rows in rooms.items():
            for vals, kind in rows:
                plan.append(("ROW", vals, kind, rid))
            plan.append(("ROOM", fl, rid, rows[0][0].get("ROOM (AR/EN)")))
        plan.append(("FLOOR", fl))
    plan.append(("PROJECT",))
    sh.r_last = FIRST + len(plan) - 1
    comp_rows = {}
    nxt = FIRST
    for p in plan:                                              # pre-compute row numbers of each room's components
        if p[0] == "ROW" and p[2] == "DETAIL":
            comp_rows.setdefault(p[3], []).append(nxt)
        nxt += 1
    for p in plan:
        if p[0] == "ROW":
            vals = dict(p[1])
            if vals.get("_material_kind") == "MEMBRANE":
                rr = comp_rows[p[3]]
                vals["_comp_rows"] = (min(rr), max(rr))
            r = sh.row({k: v for k, v in vals.items() if not k.startswith("_")}, p[2])
            detail_rows.append(r)
            _row_formulas(sh, r, vals)
        elif p[0] == "ROOM":
            r = sh.r + 1
            sh.row(_subtotals(sh, r, [("ROOM ID", p[2])], (p[1], p[3], "ROOM SUBTOTAL"), "ROOM_SUBTOTAL"),
                   "ROOM_SUBTOTAL", bold=True, fill="F2F2F2")
        elif p[0] == "FLOOR":
            r = sh.r + 1
            sh.row(_subtotals(sh, r, [("FLOOR", p[1])], (p[1], "", "FLOOR SUBTOTAL"), "FLOOR_SUBTOTAL"),
                   "FLOOR_SUBTOTAL", bold=True, fill="E2EFDA")
        else:
            r = sh.r + 1
            sh.row(_subtotals(sh, r, [], ("PROJECT", "", "PROJECT TOTAL"), "PROJECT_TOTAL"), "PROJECT_TOTAL",
                   bold=True, fill="FFF2CC")


def _structural(wb, lines):
    head = ["FLOOR", "ITEM (AR/EN)", "ELEMENT", "UNIT", "NET", "RANGE LOW", "RANGE HIGH", "EXPOSURE LOW",
            "EXPOSURE HIGH", "STATUS", "SOURCE", "FORMULA", "LINE ID", "ROW TYPE"]
    rows = []
    floors = OrderedDict()
    for ln in lines:
        floors.setdefault(ln["location"]["floor"] or "-", []).append(ln)
    n_detail = sum(len(v) for v in floors.values())
    units = sorted({ln["unit"] for ln in lines})
    n_rows = n_detail + len(floors) * len(units) + len(units)
    last = 3 + n_rows
    R = lambda col: f"${col}$4:${col}${last}"            # noqa: E731
    for fl, ls in floors.items():
        for ln in ls:
            br = ln.get("bounded_range") or {}
            prov = ln["evidence_state"] in QUANTITY_STATES
            rows.append([fl, _item_label(ln), ln.get("subtrade") or ln.get("physical_element"), ln["unit"],
                         ln["measured_net"], br.get("low") if prov else None, br.get("high") if prov else None,
                         None if prov else br.get("low"), None if prov else br.get("high"),
                         f"{ln['evidence_state']} / {ln['approval_state']}", ln["rule"]["rule_id"],
                         ln["calculation"]["formula"], ln["line_id"], "DETAIL"])
        for u in units:
            rows.append([fl, f"FLOOR SUBTOTAL ({u})", "", u] +
                        [f'=SUMIFS({R(c)},{R("A")},"{fl}",{R("D")},"{u}",{R("N")},"DETAIL")' for c in "EFGHI"] +
                        ["FLOOR_SUBTOTAL", "", "never a kg/m3 ratio; unresolved findings only as exposure", "", "FLOOR_SUBTOTAL"])
    for u in units:
        rows.append(["PROJECT", f"PROJECT TOTAL ({u})", "", u] +
                    [f'=SUMIFS({R(c)},{R("D")},"{u}",{R("N")},"DETAIL")' for c in "EFGHI"] +
                    ["PROJECT_TOTAL", "", "", "", "PROJECT_TOTAL"])
    if not lines:
        rows = [["-", "NOT_RUN: no structural drawing in this run", "", "", None, None, None, None, None, "NOT_RUN", "",
                 "", "", "INFO"]]
    fm = {c: "0.000" for c in range(5, 10)}
    return _table(wb, "STRUCTURAL", "route G canonical lines (urban_boq_m2 items adapted, numbers unchanged); like "
                  "units only; exposure = bounded unresolved findings, never added to NET", head, rows, fm)


# ---- M9 S2: the room-by-room view of an architectural-plan run (arch=...); without it the S1 workbook is unchanged
RBR_HEAD = ["FLOOR", "ROOM", "TRADE", "ITEM", "DIMENSIONS", "GROSS", "DEDUCTIONS", "NET QUANTITY", "UNIT", "WASTE %",
            "WASTE QUANTITY", "PROCUREMENT", "SOURCE", "STATUS", "NOTE / QUESTION", "OVERLAP", "REUSABLE OFFCUT",
            "MATERIAL CONSUMED", "LINE ID", "ROOM ID", "ROW TYPE"]
RC = {h: i + 1 for i, h in enumerate(RBR_HEAD)}
RBR_ORDER = ("FIN.FLOOR.PORCELAIN", "FIN.FLOOR.WET_TILE", "ARCH.SPACE", "ARCH.EXTERIOR", "FIN.SKIRTING", "FIN.CEILING",
             "FIN.WP.FLOOR", "FIN.WP.DOORWAY", "FIN.WP.UPTURN", "FIN.WALL.TILE", "FIN.PLASTER", "FIN.PAINT",
             "FIN.STAIR.FLIGHT", "ARCH.DOOR", "ARCH.WINDOW", "ARCH.PASSAGE")
TRADE_NAMES = {"FLOOR_FINISH": "أرضيات / Flooring", "CEILING": "الأسقف / Ceilings", "WATERPROOFING": "عازل / Waterproofing",
               "WALL_FINISH": "سيراميك جدران / Wall tiling", "PLASTER": "لياسة / Plaster", "PAINT": "دهان / Paint",
               "OPENINGS": "أبواب - شبابيك / Doors & windows", "EXTERNAL_WORKS": "أعمال خارجية / External works",
               "STAIR_FINISH": "الدرج / Stairs", "RAILING": "درابزين / Railings"}     # controlled vocabulary only
ARCH_TRADE_SHEETS = OrderedDict([("WALL PORCELAIN", ("FIN.WALL.TILE",)), ("PLASTER & PAINT", ("FIN.PLASTER", "FIN.PAINT")),
                                 ("CEILINGS", ("FIN.CEILING",)), ("SKIRTING", ("FIN.SKIRTING",)),
                                 ("MARBLE", ()), ("STAIRS & RAILINGS", ("FIN.STAIR.FLIGHT",)),
                                 ("DOORS & WINDOWS", ("ARCH.DOOR", "ARCH.WINDOW", "ARCH.PASSAGE")), ("ALUMINIUM", ())])


def _ded_sum(ln):
    return sum(v for d in ln.get("deductions") or () for k, v in d.items() if k != "what" and isinstance(v, (int, float)))


def _dims(ln):
    inp = ln["calculation"]["inputs"]
    ic = ln["item_code"]
    if "dims" in inp:
        return inp["dims"]
    if ic == "FIN.SKIRTING":
        return (f"boundary {inp.get('boundary', 0):.2f} lm: wall {inp.get('wall', 0):.2f}, column {inp.get('column', 0):.2f}, "
                f"door {inp.get('door', 0):.2f}, passage {inp.get('passage', 0):.2f}, sill window "
                f"{inp.get('window_threshold', 0):.2f}")
    if ic == "FIN.WP.DOORWAY":
        return f"{ln['physical_element']}: {inp['width_m']:.2f} x {inp['depth_m']:.2f} m"
    if ic == "FIN.WP.UPTURN":
        return f"{inp['perimeter_lm']:.2f} lm x {inp['upturn_h_m']:.2f} m"
    if ic.startswith("ARCH.") and "width_m" in inp:
        h = inp.get("height_m")
        return f"{ln['physical_element']}: width {inp['width_m']:.2f} m, wall {inp['wall_t_m']:.2f} m, height " + \
            (f"{h:.2f} m" if h else "not established")
    if "wall_run_lm" in inp:
        return f"run {inp['wall_run_lm']:.2f} lm" + (f" x {inp['height_m']:.2f} m" if inp.get("height_m") else
                                                     " x height (not established)")
    return ""


def _rbr_detail(ln, mat):
    loc = ln["location"]
    v = {"FLOOR": loc["floor"], "ROOM": _room_label(loc), "TRADE": TRADE_NAMES.get(ln["trade"], ln["trade"]),
         "ITEM": _item_label(ln), "DIMENSIONS": _dims(ln), "UNIT": ln["unit"],
         "SOURCE": f"{ln['source']['geometry_source']}; {ln['rule']['rule_id']}@{ln['rule']['version']}",
         "STATUS": _status(ln, mat), "LINE ID": ln["line_id"], "ROOM ID": loc["room_id"],
         "NOTE / QUESTION": ln.get("notes") or ""}
    br = ln.get("bounded_range")
    if ln["evidence_state"] in QUANTITY_STATES:
        net = ln["measured_net"]
        if ln.get("gross") is not None and abs(ln["gross"] - _ded_sum(ln) - net) <= 1e-5:
            v["GROSS"], v["DEDUCTIONS"], v["_net_formula"] = ln["gross"], _ded_sum(ln), True
            if ln.get("deductions"):
                v["NOTE / QUESTION"] = "; ".join(x for x in [v["NOTE / QUESTION"], "deductions: " + ", ".join(
                    f"{d['what']} {[x for k, x in d.items() if k != 'what'][0]:.3f}" for d in ln["deductions"])] if x)
        else:
            v["NET QUANTITY"] = net
            if ln.get("gross") is not None:
                v["GROSS"] = ln["gross"]
    elif br:
        v["GROSS"] = ln.get("gross")
        v["DEDUCTIONS"] = _ded_sum(ln) if ln.get("deductions") else None
        v["NOTE / QUESTION"] = (f"RANGE {br['low']:.3f} ... {br['high']:.3f} {ln['unit']} ({br['basis']}); " +
                                ((ln.get("rfi") or {}).get("question") or "")).strip("; ")
    else:
        v["NOTE / QUESTION"] = (ln.get("rfi") or {}).get("question") or ln.get("notes") or ""
    if mat is not None and ln["item_code"] in ("FIN.FLOOR.PORCELAIN", "FIN.FLOOR.WET_TILE"):
        if mat["state"] == "BLOCKED":
            v["PROCUREMENT"] = "UNRESOLVED: " + (mat.get("blocked_by") or "").split(":")[0]
            v["NOTE / QUESTION"] = "; ".join(x for x in [v["NOTE / QUESTION"], "material: " + (mat.get("blocked_by") or "")] if x)
        else:
            p = mat["packaging"]
            v.update({"WASTE QUANTITY": mat["cutting_waste_m2"], "_waste_formula": True,
                      "REUSABLE OFFCUT": mat["reusable_offcuts_m2"], "MATERIAL CONSUMED": mat["consumed_m2"],
                      "PROCUREMENT": (f"{p['procurement_packs']} boxes" if p["state"] == "RESOLVED" else
                                      f"{p['procurement_units']} pcs {_size(mat['material_spec'])} (boxes UNRESOLVED)")})
            if mat.get("concerns"):
                v["NOTE / QUESTION"] = "; ".join(x for x in [v["NOTE / QUESTION"]] + mat["concerns"] if x)
    if mat is not None and ln["item_code"].startswith("FIN.WP.") and mat["state"] != "BLOCKED":
        comp = next((c for c in mat.get("components") or () if ln["line_id"] in c.get("line_ids", ())), None)
        n_same = len(comp["line_ids"]) if comp else 0
        if comp is not None and n_same == 1:
            v["OVERLAP"] = comp["overlap_m2"]
    return v


def _rbr_membrane_row(rid, floor, room_label, mat):
    v = {"FLOOR": floor, "ROOM": room_label, "TRADE": TRADE_NAMES["WATERPROOFING"],
         "ITEM": "رولات عازل / Membrane rolls (room roll plan)", "UNIT": "m2", "ROOM ID": rid,
         "SOURCE": "; ".join(f"{x['rule_id']}@{x['version']}" for x in mat.get("rule_refs", ())) or mat.get("method", "-"),
         "STATUS": f"material {mat['state']}", "DIMENSIONS": _size(mat.get("material_spec"))}
    if mat["state"] == "BLOCKED":
        v.update({"PROCUREMENT": "UNRESOLVED: " + (mat.get("blocked_by") or "").split(":")[0],
                  "NOTE / QUESTION": mat.get("blocked_by") or ""})
        return v
    p, lay = mat["packaging"], mat["layout"]
    v.update({"NET QUANTITY": mat["net_installed_m2"], "OVERLAP": mat["overlap_m2"], "WASTE QUANTITY": mat["cutting_waste_m2"],
              "_waste_formula": True, "REUSABLE OFFCUT": mat["reusable_offcuts_m2"], "MATERIAL CONSUMED": mat["consumed_m2"],
              "PROCUREMENT": f"{p['procurement_packs']} rolls ({p['packaging_surplus_m2']:.3f} m2 untouched roll tail)",
              "NOTE / QUESTION": (f"installed {mat['net_installed_m2']:.3f} m2, overlaps {mat['overlap_m2']:.3f}, cutting "
                                  f"{mat['cutting_waste_m2']:.3f}, reusable {mat['reusable_offcuts_m2']:.3f}, rolls "
                                  f"{p['procurement_packs']}; strips {lay['direction']} x{lay['floor_strips']}; " +
                                  "; ".join(mat.get("concerns") or ()))})
    return v


def _rbr_plan(rb, arch, codes=None):
    """DETAIL / MATERIAL rows in floor (plan) order, then room (as measured), then item order."""
    order = {c: i for i, c in enumerate(RBR_ORDER)}
    floors = [r["floor"] for r in arch["regions"]]
    room_rank = {r["room_id"]: i for i, r in enumerate(arch["rooms"])}
    tile = {m["line_id"]: m for m in rb["materials"] if m.get("line_id") and not m.get("method", "").startswith("MEMBRANE")}
    mem = {(m.get("location") or {}).get("room_id"): m for m in rb["materials"] if m.get("method", "").startswith("MEMBRANE")}
    ls = [ln for ln in rb["lines"] if codes is None or ln["item_code"] in codes]
    ls.sort(key=lambda ln: (floors.index(ln["location"]["floor"]) if ln["location"]["floor"] in floors else 99,
                            room_rank.get(ln["location"]["room_id"], 10 ** 6), ln["location"]["room_id"] or "",
                            order.get(ln["item_code"], 99), ln["physical_element"] or ""))
    rows, done = [], set()
    for i, ln in enumerate(ls):
        rid = ln["location"]["room_id"]
        mat = tile.get(ln["line_id"]) if ln["item_code"].startswith("FIN.FLOOR") else \
            mem.get(rid) if ln["item_code"].startswith("FIN.WP.") else None
        rows.append((_rbr_detail(ln, mat), "DETAIL"))
        nxt = ls[i + 1] if i + 1 < len(ls) else None
        last_wp = ln["item_code"].startswith("FIN.WP.") and not (nxt and nxt["location"]["room_id"] == rid and
                                                               nxt["item_code"].startswith("FIN.WP."))
        if last_wp and rid in mem and rid not in done:
            done.add(rid)
            rows.append((_rbr_membrane_row(rid, ln["location"]["floor"], _room_label(ln["location"]), mem[rid]), "MATERIAL"))
    return rows, ls


def _rbr_sheet(wb, name, note, rb, arch, codes=None, info_rows=()):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(name)
    ws.append([BANNER])
    ws["A1"].font = Font(name=FONT, bold=True, color="C00000")
    ws.append([note])
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    ws.append(RBR_HEAD)
    for c in ws[3]:
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    rows, ls = _rbr_plan(rb, arch, codes)
    keys = []                                                   # (floor, item label, unit) in first-seen order
    for ln in ls:
        k = (ln["location"]["floor"], _item_label(ln), ln["unit"])
        if k not in keys:
            keys.append(k)
    items = []
    for _, it, u in keys:
        if (it, u) not in items:
            items.append((it, u))
    n_total = len(keys) + len(items)
    floor_labels = {_item_label(ln) for ln in ls if ln["item_code"].startswith("FIN.FLOOR.")}
    last = FIRST + len(rows) + n_total + (1 if rows else 0) - 1
    rng = lambda h: f"${_L(RC[h])}${FIRST}:${_L(RC[h])}${max(last, FIRST)}"     # noqa: E731
    L = lambda h: _L(RC[h])                                                     # noqa: E731
    r = FIRST - 1
    for vals, kind in rows:
        r += 1
        for h, v in vals.items():
            if not h.startswith("_"):
                ws.cell(r, RC[h], v)
        if vals.get("_net_formula"):
            ws.cell(r, RC["NET QUANTITY"], f"=ROUND({L('GROSS')}{r}-{L('DEDUCTIONS')}{r},6)")
        if vals.get("_waste_formula"):
            n, w = f"{L('NET QUANTITY')}{r}", f"{L('WASTE QUANTITY')}{r}"
            ws.cell(r, RC["WASTE %"], f'=IF(AND(ISNUMBER({n}),ISNUMBER({w})),IF({n}>0,{w}/{n},""),"")')
        ws.cell(r, RC["ROW TYPE"], kind)
    if rows:
        r += 1
        ws.cell(r, 1, "TOTALS (formulas over the DETAIL rows; a total with an unresolved row reads PARTIAL)")
        ws.cell(r, RC["ROW TYPE"], "HEADING")
    for fl, it, u in keys + [("PROJECT", it, u) for it, u in items]:
        r += 1
        crit = (f',{rng("ITEM")},"{it}",{rng("UNIT")},"{u}"' + ("" if fl == "PROJECT" else f',{rng("FLOOR")},"{fl}"'))
        tot = f'SUMIFS({rng("NET QUANTITY")},{rng("ROW TYPE")},"DETAIL"{crit})'
        opn = f'COUNTIFS({rng("ROW TYPE")},"DETAIL",{rng("STATUS")},"BLOCKED_RFI*"{crit})+' \
              f'COUNTIFS({rng("ROW TYPE")},"DETAIL",{rng("STATUS")},"BOUNDED_EXPOSURE*"{crit})'
        wst = f'SUMIFS({rng("WASTE QUANTITY")},{rng("ROW TYPE")},"DETAIL"{crit})'
        wopen = f'COUNTIFS({rng("ROW TYPE")},"DETAIL",{rng("PROCUREMENT")},"UNRESOLVED*"{crit})'
        ws.cell(r, RC["FLOOR"], fl)
        ws.cell(r, RC["ROOM"], "FLOOR TOTAL" if fl != "PROJECT" else "PROJECT TOTAL")
        ws.cell(r, RC["ITEM"], it)
        ws.cell(r, RC["UNIT"], u)
        ws.cell(r, RC["NET QUANTITY"], f'=IF({opn}>0,"PARTIAL: "&TEXT({tot},"0.000")&" ("&({opn})&" UNRESOLVED)",{tot})')
        if it in floor_labels:
            ws.cell(r, RC["WASTE QUANTITY"], f'=IF({wopen}>0,"PARTIAL: "&TEXT({wst},"0.000")&" ("&{wopen}&" UNRESOLVED)",{wst})')
        ws.cell(r, RC["ROW TYPE"], "FLOOR_TOTAL" if fl != "PROJECT" else "PROJECT_TOTAL")
    for vals in info_rows:
        r += 1
        for h, v in vals.items():
            ws.cell(r, RC[h], v)
        ws.cell(r, RC["ROW TYPE"], "INFO")
    for row in ws.iter_rows(min_row=FIRST, max_row=max(r, FIRST)):
        kind = row[RC["ROW TYPE"] - 1].value
        for c in row:
            formula = isinstance(c.value, str) and c.value.startswith("=")
            c.font = Font(name=FONT, bold=kind in ("FLOOR_TOTAL", "PROJECT_TOTAL", "HEADING"),
                          color="000000" if formula or not isinstance(c.value, (int, float)) else "0000FF")
            h = RBR_HEAD[c.column - 1]
            if h in ("GROSS", "DEDUCTIONS", "NET QUANTITY", "WASTE QUANTITY", "OVERLAP", "REUSABLE OFFCUT",
                     "MATERIAL CONSUMED"):
                c.number_format = "0.000"
            if h == "WASTE %":
                c.number_format = "0.00%"
            if kind in ("FLOOR_TOTAL", "PROJECT_TOTAL"):
                c.fill = PatternFill("solid", fgColor="E2EFDA" if kind == "FLOOR_TOTAL" else "FFF2CC")
    widths = {"FLOOR": 8, "ROOM": 26, "TRADE": 22, "ITEM": 32, "DIMENSIONS": 34, "PROCUREMENT": 26, "SOURCE": 22,
              "STATUS": 26, "NOTE / QUESTION": 70, "LINE ID": 24, "ROOM ID": 12}
    for h, j in RC.items():
        ws.column_dimensions[_L(j)].width = widths.get(h, 11)
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{_L(len(RBR_HEAD))}{max(r, 4)}"
    return {"sheet": name, "detail_rows": sum(1 for _, k in rows if k == "DETAIL"),
            "material_rows": sum(1 for _, k in rows if k == "MATERIAL"), "total_rows": n_total, "last_row": r}


def _arch_sheets(wb, arch):
    run = arch["run"]
    regs = arch["regions"]
    _table(wb, "PLAN REGIONS", "one row per plan region (floor title inside a sheet frame) and what each stage did",
           ["REGION", "FLOOR", "STAGE", "LAYER ROLES", "ROLES SOURCE", "LEVEL", "FLOOR ABOVE LEVEL", "FACES",
            "SPACES", "OPENINGS", "DOOR SYMBOLS", "PLATE M²", "REFERENCE M²", "PLATE - REFERENCE M²", "LINES",
            "WARNINGS"],
           [[r["region_id"], r["floor"], r["stage_state"], ", ".join(f"{k}={v}" for k, v in sorted(r["layer_roles"].items())
                                                                     if v != "OTHER"),
             r["roles_source"], r["level_m"], r["floor_above_level_m"], str(r.get("face_counts") or ""),
             str(r.get("space_counts") or ""), str(r.get("openings") or ""), str(r.get("door_status") or ""),
             (r.get("plate") or {}).get("largest_m2"), (r.get("reference") or {}).get("area_m2"),
             (r.get("reference") or {}).get("diff_m2"), r.get("lines"), "; ".join(r.get("warnings") or ()) +
             (("; " + r["warning"]) if r.get("warning") else "")] for r in regs],
           {12: "0.000", 13: "0.000", 14: "0.000"})
    rows = []
    for rm in arch["rooms"]:
        ru = rm["runs_lm"]
        rows.append([rm["floor"], rm["room_id"], f"{rm['name_ar']} / {rm['name_en']}", rm["class"], rm["kind"] or "-",
                     ", ".join(rm["codes"]), rm["area_m2"], rm["gross_m2"], rm["obstruction_m2"], ru.get("boundary"),
                     ru["wall"], ru["column"], ru["door"], ru["passage"], ru["window_threshold"], ru["glazing_edge"],
                     rm["dims"], rm["level_m"], len(rm["doors"]) + rm["door_closures"], len(rm["windows"]),
                     len(rm["passages"]), rm.get("issue") or ""])
    _table(wb, "ARCH SPACES", f"every measured space of the plans ({run['method']}); ROOM / MIXED of one kind carry "
           "finish lines, MERGED / UNLABELLED / MIXED of different kinds carry their area and a question",
           ["FLOOR", "ROOM ID", "ROOM (AR/EN)", "CLASS", "KIND", "ROOM WORDS", "NET AREA M²", "GROSS M²",
            "OBSTRUCTIONS M²", "BOUNDARY LM", "WALL LM", "COLUMN LM", "DOOR LM", "PASSAGE LM", "SILL WINDOW LM",
            "GLAZING LM", "DIMENSIONS", "LEVEL", "DOORS", "WINDOWS", "PASSAGES", "ISSUE"], rows,
           {c: "0.000" for c in range(7, 17)})
    _table(wb, "OPENINGS", "every jambed wall-band opening; DOOR = a door symbol hinged at the gap, WINDOW = glazing "
           "inside the gap, PASSAGE = neither", ["FLOOR", "OPENING ID", "KIND", "WIDTH M", "WALL THICKNESS M", "ROOMS",
                                                 "DOOR SYMBOL", "JAMBS"],
           [[o["floor"], o["opening_id"], o["kind"], o["width_m"], o["wall_t_m"], ", ".join(o["rooms"]),
             o.get("door_symbol") or "", " / ".join(str(j) for j in o["jambs"])] for o in arch["openings"]],
           {4: "0.000", 5: "0.000"})


def write(path, *, room_boq=None, structural_lines=(), census_rows=(), meta, created, arch=None):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    rb = room_boq or {"lines": [], "materials": [], "comparisons": [], "unresolved": [], "assumptions": []}
    summary = [f"URBAN ROOM-BY-ROOM BOQ ({POLICY_ID})", "",
               ["project", meta.get("project_id")], ["rule set", f"{meta.get('rule_set')} ({meta.get('rule_set_digest')})"],
               ["room geometry source", rb.get("geometry_source") or "no room input"],
               ["structural lines", len(structural_lines)], ["room lines", len(rb["lines"])],
               ["material results", len(rb["materials"])], ["unresolved items", len(rb["unresolved"])], "",
               "STATES", ["evidence", "MEASURED > SOURCE_DERIVED > PROVISIONAL > BOUNDED_EXPOSURE > BLOCKED_RFI"],
               ["approval", "MACHINE_ONLY (no line here is reviewed or approved)"], "",
               "DEFINITIONS (owner, M9 S1 addendum)",
               ["cutting waste", "material consumed - installed (placed) - overlap - reusable offcuts retained"],
               ["waste %", "cutting waste / net installed area"],
               ["packaging surplus", "whole boxes / rolls x pack area - material consumed (kept apart)"],
               ["spare allowance", "NOT_SET unless the owner sets one with its source"],
               ["balance check", "must read 0.000 on every row"], "",
               "LIMITS",
               ["S1 layouts", "rectangles only; a polygon room keeps its net area, its layout is BLOCKED (S2)"],
               ["offcuts", "retained offcuts are reported per room, never credited to another room (S2)"],
               ["not measured in S1", "wall tile, plaster / paint, ceilings, skirting, marble, stairs / railings, doors / "
                                      "windows, aluminium: their sheets list the families and why, with no number"]]
    if arch is not None:
        run = arch["run"]
        summary[0] = f"URBAN ROOM-BY-ROOM BOQ ({POLICY_ID}, architectural route {run['method']})"
        cut = summary.index("LIMITS")
        summary = summary[:cut] + [
            "LIMITS (architectural route)",
            ["rooms", "faces of the wall network closed at jambed openings and door symbols, named from the room words "
                      "inside; MERGED / UNLABELLED / mixed-kind spaces carry their area and a question, no finish"],
            ["wall finishes", "BLOCKED until the finish heights are established (source finish height, owner fact, or "
                              "the plan-level interval minus every sourced component); never floor-to-floor"],
            ["openings", "widths from the plan; heights only from the facts (else the wall finish net is a range)"],
            ["materials", "a tile size / joint and the membrane spec come only from the facts; without them the "
                          "material is UNRESOLVED and nothing is assumed"],
            ["layouts", "rectangular rooms are laid exactly; polygon rooms use the polygon layout when available"],
            ["not measured", "stairs / railings, marble, aluminium, external works: their families keep their census "
                             "state and reason"]]
    _info(wb, "SUMMARY", summary)
    if arch is not None:
        _rbr_sheet(wb, "ROOM BY ROOM BOQ", "every room of every floor, one row per item; NET = GROSS - DEDUCTIONS "
                   "(formula) where both are measured; waste stays apart from net; totals are formulas", rb, arch)
        _arch_sheets(wb, arch)
    asm = [["scope", "name", "value", "source"]] + [[a["scope"], a["name"], str(a["value"]), a["source"]]
                                                   for a in rb["assumptions"]]
    seen = set()
    for m in rb["materials"]:
        for a in m.get("assumptions") or ():
            key = ((m.get("location") or {}).get("room_id"), a)
            if key not in seen:
                seen.add(key)
                asm.append([(m.get("location") or {}).get("room_id"), m.get("method"), a, "material calculation"])
    _info(wb, "ASSUMPTIONS", asm)
    det = []
    sh = _Sheet(wb, "FLOOR PORCELAIN", "one row per room floor finish; material from the room's own layout "
                "(selected size); alternatives on TILE COMPARISON are recommendations only")
    fg = _floor_rows(rb)
    if fg:
        _layout_membrane(sh, fg, det)
    sh.finish()
    sh = _Sheet(wb, "WATERPROOFING", "net surfaces per room (floor, doorway pieces, 0.20 m upturn incl. doorway) and "
                "the room's roll plan on its own MATERIAL row; overlap shown per surface")
    wg = _wp_rows(rb)
    if wg:
        _layout_membrane(sh, wg, det)
    sh.finish()
    cen = {r["code"]: r for r in census_rows}
    for name, codes in TRADE_SHEETS.items():
        if arch is not None:
            info = [{"FLOOR": "-", "ROOM": "-", "ITEM": f"{cen[c]['item_ar']} / {cen[c]['item_en']} ({c})",
                     "STATUS": cen[c]["run_state"], "UNIT": cen[c]["unit"], "SOURCE": f"capability {cen[c]['capability']}",
                     "NOTE / QUESTION": cen[c]["reason"]} for c in codes if c in cen]
            _rbr_sheet(wb, name, "measured lines of this trade (room by room), then every family of the trade with its "
                       "census state", rb, arch, codes=ARCH_TRADE_SHEETS[name], info_rows=info)
            continue
        sh = _Sheet(wb, name, "not measured by the S1 engine: each family with its run state and reason; no number")
        sh.r_last = FIRST + len(codes) - 1
        for code in codes:
            r = cen.get(code)
            if r is None:
                continue
            sh.row({"FLOOR": "-", "ROOM (AR/EN)": "-", "ITEM": f"{r['item_ar']} / {r['item_en']} ({code})",
                    "SOURCE": f"capability {r['capability']}", "STATUS": r["run_state"], "UNIT": r["unit"],
                    "CONCERNS": r["reason"], "METHOD": f"next slice {r['next_slice']}"}, "INFO")
        sh.finish()
    _structural(wb, list(structural_lines))
    # procurement: formulas over the finish sheets
    fp, wp = _q("FLOOR PORCELAIN"), _q("WATERPROOFING")
    sizes = []
    for m in rb["materials"]:
        sp = m.get("material_spec") or {}
        if m["state"] != "BLOCKED" and sp.get("spec_id") and (sp["spec_id"], sp.get("kind")) not in sizes:
            sizes.append((sp["spec_id"], sp.get("kind")))
    prow = []
    col = lambda sheet, h: f"{sheet}!${_L(C[h])}:${_L(C[h])}"           # noqa: E731
    for sid, kind in sizes:
        spec = next(m["material_spec"] for m in rb["materials"] if (m.get("material_spec") or {}).get("spec_id") == sid)
        sheet = fp if kind == "TILE" else wp
        crit = f',{col(sheet, "MATERIAL SIZE")},"{_size(spec)}"'
        rt = "DETAIL" if kind == "TILE" else "MATERIAL"
        f = lambda h: f'=SUMIFS({col(sheet, h)},{col(sheet, "ROW TYPE")},"{rt}"{crit})'      # noqa: E731
        same = {(m.get("material_spec") or {}).get("pieces_per_box") for m in rb["materials"]
                if (m.get("material_spec") or {}).get("spec_id") == sid and m["state"] != "BLOCKED"}
        ppb = (same.pop() if len(same) == 1 else None) if kind == "TILE" else 1
        r = len(prow) + 4
        prow.append([_size(spec), "pcs" if kind == "TILE" else "rolls", f("UNITS CONSUMED"), f("MATERIAL M²"),
                     f("WASTE M²"), f("REUSABLE OFFCUT M²"), f("PACKAGING SURPLUS M²"),
                     f'=IF(COUNTIFS({col(sheet, "BOXES/ROLLS")},"UNRESOLVED*"{crit})>0,"UNRESOLVED",'
                     f'SUMIFS({col(sheet, "BOXES/ROLLS")},{col(sheet, "ROW TYPE")},"{rt}"{crit}))',
                     (f"=CEILING(C{r}/{ppb},1)" if ppb else "UNRESOLVED (no box data)") if kind == "TILE"
                     else "S2 (pooling roll tails across rooms)",
                     "per-room packs summed; pooled = whole boxes for the summed pieces (offcut reuse across rooms: S2)"
                     if kind == "TILE" else "per-room roll plans summed; untouched tails are packaging surplus"])
    _table(wb, "PROCUREMENT", "procurement views (never the BOQ net); spare allowance NOT_SET unless owner-approved",
           ["MATERIAL SIZE", "UNIT", "UNITS CONSUMED", "MATERIAL M²", "WASTE M²", "REUSABLE OFFCUT M²",
            "PACKAGING SURPLUS M²", "PACKS (PER-ROOM SUM)", "PACKS (POOLED)", "NOTE"], prow,
           {c: "0.000" for c in (4, 5, 6, 7)})
    crow = [[c["floor"], c["room_id"], c["spec_id"], c["role"], c["layout"], c["net_m2"], c["full"], c["cut"],
             c["consumed"], c["material_consumed_m2"], c["reusable_offcuts_m2"], c["cutting_waste_m2"],
             (c["cutting_waste_pct"] or 0) / 100.0 if c["cutting_waste_pct"] is not None else None,
             c["boxes"] if c["boxes"] is not None else f"UNRESOLVED ({c['procurement_pcs']} pcs)",
             c["packaging_surplus_m2"], "YES" if c["feasible"] else "NO", c["recommendation"], "; ".join(c["concerns"])]
            for c in rb["comparisons"]]
    _table(wb, "TILE COMPARISON", "per room, every candidate size under the same constraints; the SELECTED size is "
           "the BOQ basis, alternatives are recommendations only", ["FLOOR", "ROOM ID", "SIZE", "ROLE", "LAYOUT",
                                                                     "NET M²", "FULL", "CUT", "CONSUMED (pcs)",
                                                                     "MATERIAL M²", "REUSABLE OFFCUT M²", "WASTE M²",
                                                                     "WASTE %", "BOXES", "PACKAGING SURPLUS M²",
                                                                     "FEASIBLE", "RECOMMENDATION", "CONCERNS"], crow,
           {6: "0.000", 10: "0.000", 11: "0.000", 12: "0.000", 13: "0.00%", 15: "0.000"})
    urow = [[u["room_id"], u["item"], u["reason"]] for u in rb["unresolved"]]
    urow += [[ln["location"].get("floor") or "-", ln["item_code"], (ln.get("rfi") or {}).get("question") or "bounded exposure"]
             for ln in structural_lines if ln["evidence_state"] in ("BLOCKED_RFI", "BOUNDED_EXPOSURE")]
    _table(wb, "UNRESOLVED", "every question that blocks a quantity, a pack count or a detail", ["ROOM / FLOOR", "ITEM",
                                                                                                  "QUESTION / REASON"], urow)
    _table(wb, "COVERAGE", "the 99 BOQ families: engine capability and this run's terminal state",
           ["CODE", "GROUP", "ITEM (AR)", "ITEM (EN)", "UNIT", "CAPABILITY", "RUN STATE", "LINES", "REASON", "NEXT SLICE"],
           [[r["code"], r["group"], r["item_ar"], r["item_en"], r["unit"], r["capability"], r["run_state"], r["lines"],
             r["reason"], r["next_slice"]] for r in census_rows])
    lrow = [[ln["line_id"], ln["location"].get("floor"), ln["location"].get("room_id"), ln["item_code"], ln["unit"],
             ln["measured_net"], ln["evidence_state"], ln["approval_state"], ln["rule"]["rule_id"], ln["rule"]["version"]]
            for ln in list(structural_lines) + rb["lines"]]
    _table(wb, "LINES", "every canonical line (traceability: the finish sheets and the material results name these ids)",
           ["LINE ID", "FLOOR", "ROOM ID", "ITEM CODE", "UNIT", "MEASURED NET", "EVIDENCE", "APPROVAL", "RULE", "VERSION"],
           lrow, {6: "0.000"})
    wb.properties.creator = f"Urban QTO ({POLICY_ID})"
    wb.properties.created = wb.properties.modified = created
    wb.save(str(path))
    _repack(path)
    data = open(path, "rb").read()
    return {"path": str(path), "sheets": wb.sheetnames, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}


def expected_totals(room_boq):
    """the Python totals per room / floor / project the workbook formulas must reproduce (tests)."""
    out = {"floor_finish": {}, "waterproofing": {}}
    for ln in room_boq["lines"]:
        if ln["measured_net"] is None:
            continue
        key = "floor_finish" if ln["item_code"].startswith("FIN.FLOOR") else \
            "waterproofing" if ln["item_code"].startswith("FIN.WP") else None
        if key is None:
            continue
        loc = ln["location"]
        for g in (("ROOM", loc["room_id"]), ("FLOOR", loc["floor"]), ("PROJECT", "")):
            out[key][g] = out[key].get(g, 0.0) + ln["measured_net"]
    return out
