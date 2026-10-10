"""M1 PLANNING-ESTIMATE WORKBOOK (structural BOQ) - a view over the M1 registers.

The caller hands a workbook model, {sheet name: sheet}, in order. A sheet is
    {"kind": "INFO", "lines": [text, ...]}                                        a read-me / notes sheet
    {"kind": "TABLE", "header": [...], "rows": [[...]], "formats": {col: fmt}, "wrap_cols": [...], "widths": {...},
     "add": {col: [col_a, col_b]}, "scale": {col: (src_col, factor)}, "total": "TOTAL" | None, "note": text}
Every quantity in "rows" is a plain number copied from an M1 register row. The only formulas are the visible
arithmetic of the estimate: an "add" column is =A + B on each row (planning = verified + provisional), a "scale"
column is =A x factor (t = kg / 1000), and a "total" row is =SUM over the data rows of every numeric column. A
`None` in a quantity column stays an empty cell, never 0. Fonts: Arial; numbers copied from a register are blue,
formulas black (the usual input / calculation convention).

write() saves with fixed zip entry times, so the same model gives the same bytes; validate() reads the file back
and refuses a cell that differs from the model or a formula other than the ones the model declares. Lives outside
engine/source because it uses openpyxl.
"""

from __future__ import annotations

import hashlib
import json

from engine.boq_rc1_xlsx import _repack

POLICY_ID = "BOQ_M1_XLSX_V1"
HEADER_ROW = 3
FONT = "Arial"
BLUE, BLACK, GREY, RED = "0000FF", "000000", "595959", "C00000"
FILLS = {"VERIFIED": "E2EFDA", "PROVISIONAL": "FFF2CC", "UNRESOLVED": "EDEDED"}


class WorkbookError(ValueError):
    pass


def _norm(v):
    if isinstance(v, (list, tuple, dict)):
        return json.dumps(v, ensure_ascii=False, sort_keys=True, default=str)
    return v


def _col(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i + 1)


def _formulas(s, nrows):
    """{(row offset, col): formula} for the derived cells of a TABLE sheet."""
    out = {}
    first = HEADER_ROW + 1
    for k in range(nrows):
        r = first + k
        for c, (a, b) in (s.get("add") or {}).items():
            out[(k, int(c))] = f"={_col(a)}{r}+{_col(b)}{r}"
        for c, (a, f) in (s.get("scale") or {}).items():
            out[(k, int(c))] = f"={_col(a)}{r}*{f}"
    if s.get("total") and nrows:
        last = first + nrows - 1
        for c in sorted(set(int(x) for x in (s.get("formats") or {}))):
            out[(nrows, c)] = f"=SUM({_col(c)}{first}:{_col(c)}{last})"
    return out


def content_digest(model):
    return hashlib.sha256(json.dumps({n: s for n, s in model.items()}, sort_keys=True, default=str,
                                     ensure_ascii=False).encode()).hexdigest()


def write(model, path, *, created, banner):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for name, s in model.items():
        ws = wb.create_sheet(name[:31])
        ws.append([banner])
        ws["A1"].font = Font(name=FONT, bold=True, color=RED)
        if s["kind"] == "INFO":
            ws.append([s.get("note", "")])
            for line in s["lines"]:
                ws.append([line])
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.font = Font(name=FONT, bold=bool(c.value) and str(c.value).isupper())
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions["A"].width = 140
            continue
        ws.append([s.get("note", "")])
        ws["A2"].font = Font(name=FONT, italic=True, color=GREY)
        ws.append([_norm(h) for h in s["header"]])
        forms = _formulas(s, len(s["rows"]))
        for k, r in enumerate(s["rows"]):
            ws.append([forms.get((k, j), _norm(v)) for j, v in enumerate(r)])
        if s.get("total") and s["rows"]:
            ws.append([s["total"]] + [forms.get((len(s["rows"]), j)) for j in range(1, len(s["header"]))])
        fmts = {int(k): v for k, v in (s.get("formats") or {}).items()}
        cat = s.get("category_col")
        for cell in ws[HEADER_ROW]:
            cell.font = Font(name=FONT, bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        last = ws.max_row
        for i in range(HEADER_ROW + 1, last + 1):
            is_total = s.get("total") and i == last
            fill = None
            if cat is not None and not is_total:
                fill = FILLS.get(ws.cell(i, cat + 1).value)
            for j in range(1, len(s["header"]) + 1):
                c = ws.cell(i, j)
                formula = isinstance(c.value, str) and c.value.startswith("=")
                c.font = Font(name=FONT, bold=bool(is_total),
                              color=BLACK if formula or (j - 1) not in fmts else BLUE)
                if (j - 1) in fmts:
                    c.number_format = fmts[j - 1]
                if (j - 1) in (s.get("wrap_cols") or ()):
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                if fill:
                    c.fill = PatternFill("solid", fgColor=fill)
        widths = {int(k): v for k, v in (s.get("widths") or {}).items()}
        for j in range(len(s["header"])):
            ws.column_dimensions[_col(j)].width = widths.get(j, 14 if j else 26)
        ws.freeze_panes = f"B{HEADER_ROW + 1}"
        if s["rows"]:
            ws.auto_filter.ref = f"A{HEADER_ROW}:{_col(len(s['header']) - 1)}{HEADER_ROW + len(s['rows'])}"
    wb.properties.creator = f"Urban QTO ({POLICY_ID})"
    wb.properties.created = wb.properties.modified = created
    wb.save(str(path))
    _repack(path)
    with open(path, "rb") as f:
        data = f.read()
    return {"path": str(path), "sheets": list(model), "content_digest": content_digest(model),
            "file_sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}


def validate(path, model):
    """Every model cell is in the file unchanged; the only formulas are the declared ones."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    for name, s in model.items():
        if name[:31] not in wb.sheetnames:
            raise WorkbookError(f"sheet {name} missing")
        ws = wb[name[:31]]
        if s["kind"] == "INFO":
            continue
        forms = _formulas(s, len(s["rows"]))
        for k, r in enumerate(s["rows"]):
            for j, v in enumerate(r):
                got = ws.cell(HEADER_ROW + 1 + k, j + 1).value
                want = forms.get((k, j), _norm(v))
                if want in ("", None) and got in ("", None):
                    continue
                if isinstance(want, float) and isinstance(got, (int, float)):
                    if abs(got - want) > 1e-9 * max(1.0, abs(want)):
                        raise WorkbookError(f"{name}!{_col(j)}{HEADER_ROW + 1 + k}: {got} != {want}")
                elif got != want:
                    raise WorkbookError(f"{name}!{_col(j)}{HEADER_ROW + 1 + k}: {got!r} != {want!r}")
        for row in ws.iter_rows(min_row=HEADER_ROW + 1):
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    k, j = c.row - HEADER_ROW - 1, c.column - 1
                    if forms.get((k, j)) != c.value:
                        raise WorkbookError(f"{name}!{c.coordinate}: undeclared formula {c.value}")
    return True


def policy_record():
    return {"policy": POLICY_ID, "formulas": "planning = verified + provisional; t = kg / 1000; SUM totals only",
            "rule": "a view: every quantity is copied from an M1 register row; unresolved never enters a total"}
