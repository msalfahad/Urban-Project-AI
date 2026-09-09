"""Load a real .xlsx workbook into normalized BoqRow objects.

Reads the workbook twice: once for cached values (data_only=True) and once for
formulas (data_only=False). That lets the auditor check both the numbers and the
structure — a SUM's range, a #REF!, a hand-typed total.

Column detection is by header keyword (English + Arabic) so it copes with the
one-workbook-per-trade layout the QS uses without hard-coding cell positions. A
`ColumnMap` can be supplied to override detection when a sheet is unusual.
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from .model import BoqRow

# Header synonyms → logical column.
HEADER_SYNONYMS = {
    "description": ["description", "item", "desc", "بيان", "الوصف", "البند", "وصف", "الأعمال", "التفاصيل"],
    "unit": ["unit", "uom", "الوحدة", "وحدة"],
    "quantity": ["quantity", "qty", "الكمية", "كمية", "العدد"],
    "unit_rate": ["rate", "unit rate", "unit price", "price", "سعر الوحدة", "السعر", "سعر", "فئة"],
    "amount": ["amount", "total", "value", "الاجمالي", "الإجمالي", "المجموع", "القيمة", "الاجمالى"],
}

TOTAL_KEYWORDS = ["total", "subtotal", "sub-total", "grand total", "المجموع", "الإجمالي",
                  "الاجمالي", "اجمالي", "المجموع الكلي", "الاجمالى"]

EXCEL_ERROR_TOKENS = ["#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"]


@dataclass
class ColumnMap:
    description: int | None = None
    unit: int | None = None
    quantity: int | None = None
    unit_rate: int | None = None
    amount: int | None = None


def _to_number(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s == "" or any(tok in s for tok in EXCEL_ERROR_TOKENS):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _detect_header(ws_vals) -> tuple[int, ColumnMap] | None:
    """Find the header row and map columns from its labels."""
    for row in ws_vals.iter_rows(min_row=1, max_row=min(20, ws_vals.max_row)):
        labels = {}
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value).strip().lower()
            for logical, syns in HEADER_SYNONYMS.items():
                if logical in labels:
                    continue
                if any(s in text for s in syns):
                    labels[logical] = cell.column
        # A header row must at least locate a description and one numeric column.
        if "description" in labels and ({"quantity", "amount", "unit_rate"} & set(labels)):
            return row[0].row, ColumnMap(
                description=labels.get("description"),
                unit=labels.get("unit"),
                quantity=labels.get("quantity"),
                unit_rate=labels.get("unit_rate"),
                amount=labels.get("amount"),
            )
    return None


_CELL_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_RANGE_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")


def _parse_sum_rows(formula: str, amount_col: int | None) -> list[int]:
    """Extract the row numbers a formula references in the amount column.

    Handles =SUM(G5:G9) and =G5+G6 style totals. Restricts to the amount column
    so unrelated references don't count.
    """
    if not formula:
        return []
    rows: set[int] = set()
    col_letter = get_column_letter(amount_col) if amount_col else None

    for m in _RANGE_RE.finditer(formula):
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        if col_letter and (c1 != col_letter or c2 != col_letter):
            continue
        rows.update(range(min(r1, r2), max(r1, r2) + 1))
    # single cell references (strip out the ones already inside ranges is fine —
    # they resolve to the same rows)
    for m in _CELL_RE.finditer(formula):
        col, r = m.group(1), int(m.group(2))
        if col_letter and col != col_letter:
            continue
        rows.add(r)
    return sorted(rows)


def _cell_errors(*values) -> list[str]:
    errs = []
    for v in values:
        if v is None:
            continue
        s = str(v)
        for tok in EXCEL_ERROR_TOKENS:
            if tok in s and tok not in errs:
                errs.append(tok)
    return errs


def load_rows(path: str, column_map: ColumnMap | None = None) -> list[BoqRow]:
    """Load every sheet of a workbook into BoqRow objects."""
    wb_vals = load_workbook(path, data_only=True, read_only=False)
    wb_forms = load_workbook(path, data_only=False, read_only=False)

    rows: list[BoqRow] = []
    for name in wb_vals.sheetnames:
        ws_vals = wb_vals[name]
        ws_forms = wb_forms[name]

        detected = None if column_map else _detect_header(ws_vals)
        if column_map:
            header_row, cmap = 1, column_map
        elif detected:
            header_row, cmap = detected
        else:
            continue  # no recognisable BOQ table on this sheet

        for r_idx in range(header_row + 1, ws_vals.max_row + 1):
            def val(col):
                return ws_vals.cell(row=r_idx, column=col).value if col else None

            def form(col):
                if not col:
                    return ""
                f = ws_forms.cell(row=r_idx, column=col).value
                return f if isinstance(f, str) and f.startswith("=") else ""

            desc = val(cmap.description)
            desc_s = "" if desc is None else str(desc).strip()
            qty = _to_number(val(cmap.quantity))
            rate = _to_number(val(cmap.unit_rate))
            amount = _to_number(val(cmap.amount))
            amount_formula = form(cmap.amount)
            quantity_formula = form(cmap.quantity)

            # skip fully empty rows
            if not desc_s and qty is None and rate is None and amount is None and not amount_formula:
                continue

            is_total = (
                any(k in desc_s.lower() for k in TOTAL_KEYWORDS)
                or (amount_formula.upper().startswith("=SUM(") and qty is None and rate is None)
                or quantity_formula.upper().startswith("=SUM(")
            )

            errs = _cell_errors(
                val(cmap.amount), val(cmap.quantity), val(cmap.unit_rate),
                amount_formula, quantity_formula,
            )

            rows.append(BoqRow(
                row_index=r_idx,
                sheet=name,
                description=desc_s,
                unit="" if val(cmap.unit) is None else str(val(cmap.unit)).strip(),
                quantity=qty,
                unit_rate=rate,
                amount=amount,
                amount_formula=amount_formula,
                quantity_formula=quantity_formula,
                cell_errors=errs,
                is_total_row=is_total,
                amount_sum_rows=_parse_sum_rows(amount_formula, cmap.amount),
                qty_sum_rows=_parse_sum_rows(quantity_formula, cmap.quantity),
                section=name,
            ))

    wb_vals.close()
    wb_forms.close()
    return rows
